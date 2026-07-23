#!/usr/bin/env python3
"""Generate 4P Index coding spreadsheet for Decree 2016-1804 (English only).

Fishing-gear provisions with plastic/marine-debris relevance only; strict plastics-relevance filter.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT = "/workspace/4P_Index_Decret_2016-1804_Code_Peche_Maritime.xlsx"

COLUMNS = [
    ("A", "policy_name"), ("B", "policy_url"), ("C", "policy_year"),
    ("D", "policy_objective"), ("E", "policy_target"), ("F", "policy_target_text"),
    ("G", "policy_type"), ("H", "policy_type_justification"),
    ("I", "policy_integration"), ("J", "policy_sectors_list"),
    ("K", "policy_circularity"), ("L", "policy_lifecycle_phases_list"),
    ("M", "policy_budget"), ("N", "policy_budget_text"),
    ("O", "policy_score"), ("P", "instrument_type"),
    ("Q", "instrument_lifecycle_stage"), ("R", "instrument_description"),
    ("S", "instrument_in_force"), ("T", "instrument_implementation"),
    ("U", "instrument_implementation_text"), ("V", "instrument_score"),
    ("W", "comments"),
]

POLICY = {
    "policy_name": (
        "Decree No. 2016-1804 of 22 November 2016 implementing the Maritime Fisheries Code"
    ),
    "policy_url": "https://www.ditp.gouv.sn/download/file/fid/76",
    "policy_year": 2016,
    "policy_objective": (
        "Set modalities for applying Law No. 2015-18 (Maritime Fisheries Code), including fishing-gear "
        "control, net-mesh rules and marine environment protection — relevant to synthetic-polymer fishing "
        "gear and ghost-gear reduction, without plastic marine litter-specific measures."
    ),
    "policy_target": 0.75,
    "policy_target_text": (
        "Art. 25: ban on drift gillnets for shrimp fishing in all Senegalese waters; "
        "Art. 35: ban on devices obstructing mesh or reducing selective action; "
        "Art. 37: Minister may require selective gear to preserve resources and the marine environment."
    ),
    "policy_type": 0.75,
    "policy_type_justification": (
        "Implementing decree adopted on 22 November 2016 to apply Law No. 2015-18. "
        "Sub-legislative executive implementing regulation (≠1.0), not parliamentary legislation."
    ),
    "policy_integration": 1.0,
    "policy_sectors_list": (
        "maritime fisheries, artisanal fishing, industrial fishing, marine environment, maritime surveillance"
    ),
    "policy_circularity": 0.50,
    "policy_lifecycle_phases_list": "use, disposal, marine environment",
    "policy_budget": 0,
    "policy_budget_text": (
        "No dedicated budget allocation in decree; licensing fees governed separately under Law 2015-18 Art. 64. "
        "Enforcement via fisheries surveillance agents and observer programme (Ch. 6)."
    ),
}

INSTRUMENTS = [
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Use",
        "instrument_description": (
            "Art. 25: prohibits use of drift gillnets for shrimp fishing in all Senegalese waters, "
            "alongside beach-seine restrictions in specified zones and nets encircling rock-fish habitats."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            "Art. 26: any gear not listed in Art. 24 requires prior ministerial authorisation (+0.25 authority); "
            "Arts. 59-62: observers may inspect fishing gear on board (+0.25 monitoring); "
            "Law 2015-18 Arts. 125-126: confiscation and destruction of prohibited gear (+0.25 enforcement)."
        ),
        "comments": "Drift gillnet ban — reduces ghost gear and non-selective catch (nylon/polyester nets).",
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Use",
        "instrument_description": (
            "Art. 35: for all fishing gear types, it is prohibited to use means or devices obstructing mesh "
            "or reducing selective action (with limited trawl protection-panel exception)."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            "Arts. 31-34: detailed mesh-measurement procedures for enforcement agents (+0.25 monitoring); "
            "implements Law 2015-18 Art. 125(e) (+0.25 enforcement); "
            "Art. 23: Minister may tighten conservation measures by order (+0.25 authority)."
        ),
        "comments": "Ban on selectivity-reducing devices — complements Law 2015-18 monofilament net ban.",
    },
    {
        "instrument_type": 0.75,
        "instrument_lifecycle_stage": "Marine environment",
        "instrument_description": (
            "Art. 37: Fisheries Minister may take necessary measures on use of any device or rigging that "
            "destroys natural habitats to ensure preservation of resources and the marine environment; "
            "may promote or require selective gear for biodiversity and stock management."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            "Art. 23: Minister may supplement conservation measures by order (+0.25 authority); "
            "Art. 42: temporary zone closures possible (+0.25 monitoring). "
            "No article-specific penalty for marine litter."
        ),
        "comments": "Marine environment and habitat protection — framework for anti-pollution/gear measures.",
    },
]


def build_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "4P Index Coding"
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    auto_fill = PatternFill("solid", fgColor="E2EFDA")

    for col_idx, (col_letter, col_name) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=f"{col_letter} — {col_name}")
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for row_idx, instrument in enumerate(INSTRUMENTS, start=2):
        row_data = {**POLICY, **instrument}
        for col_idx, (_, key) in enumerate(COLUMNS, start=1):
            if key in ("policy_score", "instrument_score"):
                continue
            ws.cell(row=row_idx, column=col_idx, value=row_data.get(key, ""))

        g_col, i_col, k_col, m_col = (get_column_letter(n) for n in (7, 9, 11, 13))
        p_col, t_col, o_col, v_col = (get_column_letter(n) for n in (16, 20, 15, 22))
        ws[f"{o_col}{row_idx}"] = f"=AVERAGE({g_col}{row_idx},{i_col}{row_idx},{k_col}{row_idx},{m_col}{row_idx},{p_col}{row_idx})"
        ws[f"{v_col}{row_idx}"] = f"=AVERAGE({p_col}{row_idx},{t_col}{row_idx})"
        ws[f"{o_col}{row_idx}"].fill = auto_fill
        ws[f"{v_col}{row_idx}"].fill = auto_fill

    widths = {
        "A": 48, "B": 52, "C": 10, "D": 52, "E": 12, "F": 60,
        "G": 10, "H": 52, "I": 14, "J": 40, "K": 14, "L": 36,
        "M": 12, "N": 52, "O": 12, "P": 14, "Q": 22, "R": 60,
        "S": 14, "T": 18, "U": 52, "V": 14, "W": 52,
    }
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = "A2"
    wb.save(OUTPUT)
    print(f"Saved {OUTPUT} ({len(INSTRUMENTS)} instrument rows)")


if __name__ == "__main__":
    build_workbook()
