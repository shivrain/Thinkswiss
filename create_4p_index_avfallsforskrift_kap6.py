#!/usr/bin/env python3
"""Generate 4P Index Excel for Avfallsforskriften Kap. 6 (beverage DRS/pant)."""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

POLICY = {
    "policy_name": (
        "Avfallsforskriften – Kap. 6: Retursystemer for emballasje til drikkevarer (pant) "
        "(Waste Regulation – Ch. 6: Return systems for beverage packaging (deposit-return scheme))"
    ),
    "policy_url": "https://lovdata.no/dokument/SF/forskrift/2004-06-01-930/KAPITTEL_6",
    "policy_year": 2025,
    "policy_objective": (
        "To establish effective deposit-return systems (panteordning) for beverage inner packaging—including "
        "plastic bottles and other containers distributed to consumers—with high return rates, mandatory retailer "
        "take-back, approved return-system requirements, deposit marking and fixed deposit levels, linking "
        "achieved return rates to excise relief under særavgiftsforskriften Kap. 3-5, to prevent littering and "
        "reduce beverage-packaging waste."
    ),
    "policy_target": 1,
    "policy_target_text": (
        '"The purpose … is to contribute to effective return systems with a high return rate for beverage inner '
        'packaging." (§6-2); "A precondition for approval is that the return system is expected to achieve a '
        'minimum 25 % return rate." (§6-4); return-rate decisions under §6-5 feed excise reduction in '
        "særavgiftsforskriften Kap. 3-5 (full miljøavgift exemption at ≥95 % return via §3-5-3(2))"
    ),
    "policy_type": 0.75,
    "policy_type_justification": (
        "Executive regulation (forskrift) issued 1 June 2004 under produktkontrolloven §4 and forurensningsloven; "
        "implements EU Packaging Directive elements. Substantially revised May 2022 (§6-5a PRO cross-duties) and "
        "June 2025 (FOR-2025-06-18-1128, in force 1 July 2025). Deposit rates set in §6-8; chapter last amended "
        "2025."
    ),
    "policy_integration": 1,
    "policy_sectors_list": (
        "beverage, retail, waste management, producers, recycling, consumption, food & beverage, environmental leakage"
    ),
    "policy_circularity": 1,
    "policy_lifecycle_phases_list": (
        "consumption, end-of-life, recycling, environmental leakage, litter/pollution"
    ),
    "policy_budget": 1,
    "policy_budget_text": (
        '"A deposit scheme means an arrangement where consumer and point of sale pay a certain amount (deposit) '
        'for packaging, refunded on return of empty packaging." (§6-3); fixed deposit rates NOK 2.00 (≤0.5 L) and '
        "NOK 3.00 (>0.5 L) per unit (§6-8); return systems pay approval and annual Treasury fees (§6-11, §6-11a)"
    ),
}

INSTRUMENTS = [
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Consumption",
        "instrument_description": (
            "§6-1: Scope—chapter applies to return systems for inner packaging of beverages (bottles, cans etc.) "
            "used in distribution all the way to the consumer, including plastic beverage bottles."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            '"The provisions in this chapter apply to return systems for inner packaging of beverages." (§6-1); '
            '"only inner packaging used in distribution all the way to the consumer" (§6-1)'
        ),
        "comments": "Defines national mandatory DRS scope for consumer beverage containers.",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Litter/Pollution",
        "instrument_description": (
            "§6-2: Purpose—to contribute to effective return systems with high return rates for beverage inner "
            "packaging so systems help prevent littering and reduce waste volumes from such packaging."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            '"The purpose … is to contribute to effective return systems with a high return rate for inner '
            'packaging of beverages so that return systems help prevent littering and reduce waste volumes from '
            'such inner packaging." (§6-2)'
        ),
        "comments": "Qualitative high-return objective; operational systems exceed 95 % nationally.",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Consumption",
        "instrument_description": (
            "§6-3: Definitions—return system (free consumer return for reuse/recycling); deposit scheme (refundable "
            "deposit on packaging); inner packaging (bottle/can unit filled with beverage); recovery includes "
            "material recovery."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            '"A return system means an arrangement where the consumer can deliver empty packaging free of charge '
            'for return for reuse or recovery." (§6-3); "A deposit scheme means … deposit refunded on return of '
            'empty packaging." (§6-3)'
        ),
        "comments": "Defines core DRS/pant concepts for plastic and other beverage containers.",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§6-4: Return-system approval—producers/importers may establish or join return systems; Miljødirektoratet "
            "approves if system is expected to achieve minimum 25 % return and packaging goes to environmentally "
            "sound reuse/recovery; must demonstrate ability to meet §6-5a duties; register as legal entity."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            '"A precondition for approval is that the return system is expected to achieve a minimum 25 % return '
            'rate, and that packaging goes to environmentally sound reuse or recovery." (§6-4); approved systems '
            "must demonstrate ability to fulfil §6-5a duties (§6-4)"
        ),
        "comments": "Gatekeeper for approved PRO/return operators (e.g. Infinitum); Miljødirektoratet avgjør godkjenning.",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": (
            "§6-5: Return-rate setting—Miljødirektoratet sets expected return rate annually (typically 3-year "
            "average of achieved rates); decision forms basis for environmental-tax reduction under "
            "særavgiftsforskriften Kap. 3-5 (including ≥95 % full exemption threshold)."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            '"The Norwegian Environment Agency sets which return rate the return system can be expected to achieve." '
            "(§6-5); "
            '"Decisions on return rate form the basis for reduction of tax under the excise regulation Kap. 3-5 on '
            'beverage packaging." (§6-5)'
        ),
        "comments": "Links collection performance to excise incentives; operational DRS targets >95 %.",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": (
            "§6-5a: Extended PRO duties—approved return systems must likewise comply with Avfallsforskriften Kap. 7 "
            "duties on waste prevention reporting, cost coverage, material recovery (§7-10), recovery calculation "
            "(§7-11), non-profit operation, equal treatment, reporting/documentation, financial reserves and public "
            "information (§§7-7, 7-8, 7-10, 7-11, 7-14–7-17, 7-19–7-20)."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            '"The duties and requirements in avfallsforskriften §§7-7 … 7-10 (duty to material recovery) … 7-16 '
            '(duty to report and document) … apply correspondingly to approved return systems under §6-4." (§6-5a)'
        ),
        "comments": "Incorporates material-recovery targets for plastic packaging via Kap. 7 cross-reference.",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Consumption",
        "instrument_description": (
            "§6-6: Deposit marking—inner packaging in the deposit scheme shall be marked with a deposit symbol "
            "showing the deposit rate; minimum mark size 9 mm × 9 mm on packaging or label."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            '"Inner packaging included in a deposit scheme shall be marked with a deposit mark showing the deposit '
            'rate." (§6-6); minimum size 9 mm × 9 mm (§6-6)'
        ),
        "comments": "Mandatory marking (skal) enabling consumer identification of deposit containers.",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§6-7: Retailer take-back—points of sale of deposit-scheme beverages must accept reasonable quantities "
            "of empty inner packaging they sell; deposit must be paid out in cash on delivery."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            '"Points of sale of beverages in packaging included in a deposit scheme have a duty to accept in return '
            'reasonable quantities of empty inner packaging that they sell." (§6-7); deposit payable in cash on '
            "delivery (§6-7)"
        ),
        "comments": "Core retailer collection obligation; mandatory plikt.",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Consumption",
        "instrument_description": (
            "§6-8: Deposit levels—point of sale and consumer shall pay deposit of NOK 2.00 per unit (filled volume "
            "≤0.5 L) and NOK 3.00 per unit (>0.5 L); Miljødirektoratet may set higher rates for low-return or "
            "high-cost packaging types on application."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            '"For inner packaging in a deposit scheme, point of sale and consumer shall pay deposit at the following '
            'rates: a. … ≤50 cl: NOK 2.00 per unit b. … over 50 cl: NOK 3.00 per unit." (§6-8)'
        ),
        "comments": "Binding national deposit rates (skal); economic incentive for returning plastic bottles.",
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Consumption",
        "instrument_description": (
            "§6-9: Packaging ban power—Miljødirektoratet may prohibit use of inner packaging that prevents proper "
            "operation of established deposit schemes."
        ),
        "instrument_in_force": 0,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": (
            '"The Norwegian Environment Agency may prohibit use of inner packaging that prevents proper '
            'implementation of established deposit schemes." (§6-9)'
        ),
        "comments": "Enabling power (kan forby); not a standing ban.",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§6-10: Supervision—Miljødirektoratet supervises compliance with chapter provisions and decisions made "
            "under the chapter."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            '"The Norwegian Environment Agency supervises that the provisions in this chapter and decisions made '
            'under the provisions in this chapter are complied with." (§6-10)'
        ),
        "comments": "Enforcement framework for approved return systems and retailers.",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§6-11: Approval fee—return systems shall pay Treasury fee for Miljødirektoratet approval processing "
            "(tiered up to NOK 203,900 depending on case complexity)."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            '"The return system shall pay a fee to the Treasury for the Norwegian Environment Agency\'s case '
            'processing in connection with approval or change of approval." (§6-11)'
        ),
        "comments": "Producer-financed approval costs; added 2022, fee rates updated 2024–2025.",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§6-11a: Annual fee—approved return systems shall pay annual Treasury fee (NOK 20,400) for "
            "Miljødirektoratet reporting oversight."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            '"Approved return systems shall pay an annual fee to the Treasury for the Norwegian Environment '
            'Agency\'s work on annual reporting." (§6-11a); NOK 20,400 per system (§6-11a)'
        ),
        "comments": "Economic instrument financing DRS administration; added 2024.",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§6-11b: Data-system costs—approved return systems shall cover Miljødirektoratet costs for data systems "
            "necessary to ensure compliance with chapter producer-responsibility requirements."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            '"Approved return systems shall cover the Norwegian Environment Agency\'s costs for development and '
            'operation of data systems necessary to ensure that requirements linked to producer responsibility in '
            'this chapter are complied with." (§6-11b)'
        ),
        "comments": "Producer-financed EPR data infrastructure; Miljødirektoratet fastsetter gebyr.",
    },
]

HEADERS = [
    ("A", "policy_name"),
    ("B", "policy_url"),
    ("C", "policy_year"),
    ("D", "policy_objective"),
    ("E", "policy_target"),
    ("F", "policy_target_text"),
    ("G", "policy_type"),
    ("H", "policy_type_justification"),
    ("I", "policy_integration"),
    ("J", "policy_sectors_list"),
    ("K", "policy_circularity"),
    ("L", "policy_lifecycle_phases_list"),
    ("M", "policy_budget"),
    ("N", "policy_budget_text"),
    ("O", "policy_score"),
    ("P", "instrument_type"),
    ("Q", "instrument_lifecycle_stage"),
    ("R", "instrument_description"),
    ("S", "instrument_in_force"),
    ("T", "instrument_implementation"),
    ("U", "instrument_implementation_text"),
    ("V", "instrument_score"),
    ("W", "comments"),
]


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "4P Index - Avfallsforskr Kap6"

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col_idx, (col_letter, header) in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=f"{col_letter} — {header}")
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, inst in enumerate(INSTRUMENTS, start=2):
        policy_score = (
            POLICY["policy_type"]
            + POLICY["policy_integration"]
            + POLICY["policy_circularity"]
            + POLICY["policy_budget"]
            + inst["instrument_type"]
        ) / 5
        instrument_score = (inst["instrument_type"] + inst["instrument_implementation"]) / 2

        row_values = [
            POLICY["policy_name"],
            POLICY["policy_url"],
            POLICY["policy_year"],
            POLICY["policy_objective"],
            POLICY["policy_target"],
            POLICY["policy_target_text"],
            POLICY["policy_type"],
            POLICY["policy_type_justification"],
            POLICY["policy_integration"],
            POLICY["policy_sectors_list"],
            POLICY["policy_circularity"],
            POLICY["policy_lifecycle_phases_list"],
            POLICY["policy_budget"],
            POLICY["policy_budget_text"],
            round(policy_score, 3),
            inst["instrument_type"],
            inst["instrument_lifecycle_stage"],
            inst["instrument_description"],
            inst["instrument_in_force"],
            inst["instrument_implementation"],
            inst["instrument_implementation_text"],
            round(instrument_score, 3),
            inst["comments"],
        ]

        for col_idx, value in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30

    widths = {
        "A": 55, "B": 42, "C": 12, "D": 45, "E": 12, "F": 42, "G": 12, "H": 40,
        "I": 16, "J": 42, "K": 16, "L": 38, "M": 14, "N": 45, "O": 12, "P": 14,
        "Q": 20, "R": 55, "S": 14, "T": 18, "U": 45, "V": 14, "W": 42,
    }
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width

    output_path = "/workspace/4P_Index_Avfallsforskriften_Kap6.xlsx"
    wb.save(output_path)
    print(f"Saved {output_path} ({len(INSTRUMENTS)} instrument rows)")


if __name__ == "__main__":
    main()
