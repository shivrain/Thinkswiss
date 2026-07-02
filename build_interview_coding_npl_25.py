"""Generate interview coding spreadsheet for Nepal interview NPL_25.

Coding follows the Overview All Interviews codebook and is based on the
third ordinary-household respondent interview of 17 February 2026 —
interview guideline only (no recording; no separate transcript). Uses
expanded Impacts taxonomy.
"""

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUTPUT_PATH = "/workspace/interview_coding_NPL_25.xlsx"

DARK_BLUE = "1F3864"
MID_BLUE = "2E75B6"
LIGHT_BLUE = "D9E2F3"
LIGHT_GREEN = "E2EFDA"
ORANGE = "FCE4D6"
WHITE = "FFFFFF"

SOURCES = [
  ("Interview guideline", "Ordinary household respondent #3 — field notes, 17 Feb 2026"),
  ("Codebook", "Overview All Interviews — NEW Codebook (expanded Impacts)"),
  (
    "Coding note",
    "No audio recording (consent given, no recording). Codes verified against "
    "guideline/field notes only. Q4 follow-up, Q6 follow-up, Q9, and Q10 not "
    "asked. Same fieldwork day as NPL_21–NPL_24.",
  ),
]

STAKEHOLDER = {
  "Interview ID": "NPL_25",
  "Country": "NEPAL",
  "Interview number": "Household / citizen fieldwork series (#3, 17 Feb 2026)",
  "Interview date": "17 February 2026",
  "Interviewee": "Unknown — ordinary household resident",
  "Affiliation / role": "Ordinary person / household member",
  "Organization": "NA",
  "Stakeholder list mapping": "household",
  "Coded actor type (codebook)": "household",
  "Actor classification rationale": (
    "Lay citizen/household respondent — not government, retailer, or waste "
    "operator; describes personal concern, municipal collection reliance, and "
    "cultural-religious littering norms."
  ),
  "Perspective": (
    "Highly concerned citizen — health impacts and non-degradable plastics; "
    "rising packaging, bottles, and diapers; government action limited; "
    "residents resist segregation for comfort; advocates sacred-site taboos and "
    "municipal god monuments to deter open dumping"
  ),
  "Project context": (
    "Daily municipal waste-collection vehicle; mayor/local government responsible; "
    "municipalities place god monuments to create saint places where waste "
    "disposal is culturally prohibited"
  ),
}

CODING = {
  "Country": "NPL",
  "Country name": "NEPAL",
  "ID": "NPL_25",
  "Actortype": "household",
  "Problem_awareness_pop": "medium",
  "Problem_awareness_pol": "high",
  "Problem_severity": "high",
  "Problem_littering": "yes",
  "Problem_consumption": "yes",
  "Problem_recycling": "no",
  "Problem_waste_mgmt": "yes",
  "Problem_production": "no",
  "Problem_alternatives": "yes",
  "Problem_waste_segregation": "yes",
  "Problem_import": "no",
  "Impacts": "health",
  "Governance": "yes",
  "Relevance_international_pol": "no",
  "Coordination_sectoral": "no",
  "Coordination_levels": "no",
  "Unclear_responsibilities": "no",
  "Actors": "yes",
  "Federal government": "no",
  "Provincial government": "no",
  "Local government": "yes",
  "Students": "no",
  "Private_sector": "no",
  "Civil_society": "no",
  "Science": "no",
  "Households": "yes",
  "Education institutions": "no",
  "Private_companies(hotels, shops, etc.)": "no",
  "Implementation issues": "yes",
  "Monitoring": "no",
  "Financial_resources": "no",
  "Research": "no",
  "Infrastructure": "no",
  "Enforcement": "yes",
  "Policies in place": "yes",
  "Pol_epr": "no",
  "Pol_import": "no",
  "Pol_awarness": "yes",
  "Pol_education": "no",
  "Pol_capacity": "no",
  "Pol_RD": "no",
  "Pol_tax": "no",
  "Pol_ban": "no",
  "Pol_subsitutes": "no",
  "Pol_clean_up": "yes",
  "Pol_upcycling": "no",
  "Pol_recycling": "no",
  "Pol_waste_collection": "yes",
  "Solutions": "yes",
  "Sol_lead_agency": "no",
  "Sol_responsibilities": "no",
  "Sol_epr": "no",
  "Sol_awarness": "yes",
  "Sol_segregation": "no",
  "Sol_upcycling": "no",
  "Sol_recycling": "no",
  "Sol_education": "no",
  "Sol_capacity": "no",
  "Sol_RD": "no",
  "Sol_tax": "no",
  "Sol_ban": "no",
  "Sol_finance": "no",
  "Sol_infrastructure": "no",
  "Sol_subsitutes": "no",
  "Sol_clean_up": "yes",
  "Sol_enforcement": "yes",
  "Sol_monitoring": "no",
  "Traditions to build on (free-hand)": (
    "Sacred-site/taboo tradition — waste must not be thrown near temples or "
    "saint places; municipalities install god monuments to create such protected "
    "zones (material traditions rejected — plastics preferred for comfort)"
  ),
}

CODEBOOK = [
  ("Country", "Country ISO code"),
  ("Country name", "Country name"),
  ("ID", "InterviewID: ISO_Nr"),
  ("Actortype", "governmental, pol_party, science, cso, igo, private_sector, shop_owner, hotel_owner, household"),
  ("Problem_awareness_pop", "high, medium, low, NA — lack of awareness among population mentioned"),
  ("Problem_awareness_pol", "high, medium, low, NA — lack of awareness among policy makers mentioned"),
  ("Problem_severity", "high, medium, low, NA"),
  ("Problem_littering", "Littering is a problem; was mentioned: yes/no"),
  ("Problem_consumption", "High consumption is a problem; was mentioned: yes/no"),
  ("Problem_recycling", "No or too little recycling is a problem; was mentioned: yes/no"),
  ("Problem_waste_mgmt", "Ill waste management is a problem; was mentioned: yes/no"),
  ("Problem_production", "High production rates; was mentioned: yes/no"),
  ("Problem_alternatives", "No good alternatives; was mentioned: yes/no"),
  ("Problem_waste_segregation", "Lack of segregation was mentioned: yes/no"),
  ("Problem_import", "Plastic imports are a problem; was mentioned: yes/no"),
  ("Impacts", "Impacts mentioned on water, cities, biodiversity, health, etc. or NA"),
  ("Governance", "Mentioned: yes/no"),
  ("Relevance_international_pol", "International treaties are relevant for national level policy"),
  ("Coordination_sectoral", "Lack of coordination across sectors"),
  ("Coordination_levels", "Lack of coordination across levels"),
  ("Unclear_responsibilities", "Unclear responsibilities among the involved actors"),
  ("Actors", "Mentioned: yes/no"),
  ("Federal government", "Mentioned: yes/no"),
  ("Provincial government", "Mentioned: yes/no"),
  ("Local government", "Mentioned: yes/no"),
  ("Students", "Mentioned: yes/no"),
  ("Private_sector", "Mentioned: yes/no"),
  ("Civil_society", "Mentioned: yes/no"),
  ("Science", "Mentioned: yes/no"),
  ("Households", "Mentioned: yes/no"),
  ("Education institutions", "Mentioned: yes/no"),
  ("Private_companies(hotels, shops, etc.)", "Mentioned: yes/no"),
  ("Implementation issues", "Mentioned: yes/no"),
  ("Monitoring", "Lack of monitoring"),
  ("Financial_resources", "Lack of financial resources"),
  ("Research", "Lack of research"),
  ("Infrastructure", "Lack of infrastructure"),
  ("Enforcement", "Lack of enforcement"),
  ("Policies in place", "Mentioned: yes/no"),
  ("Pol_epr", "Extended producer responsibility"),
  ("Pol_import", "Import regulations"),
  ("Pol_awarness", "Awareness campaigns"),
  ("Pol_education", "Education programs"),
  ("Pol_capacity", "Capacity building"),
  ("Pol_RD", "Research & development"),
  ("Pol_tax", "Levies or taxes on specific products (often bags)"),
  ("Pol_ban", "Bans of specific products"),
  ("Pol_subsitutes", "Alternatives like reusable bags, or banana leaf plates"),
  ("Pol_clean_up", "Clean-up campaigns"),
  ("Pol_upcycling", "Making a new product, like blacktop"),
  ("Pol_recycling", "Making a new raw material"),
  ("Pol_waste_collection", "Waste collection"),
  ("Solutions", "Mentioned: yes/no"),
  ("Sol_lead_agency", "Procedural measures to establish lead agency"),
  ("Sol_responsibilities", "Clear responsibilities"),
  ("Sol_epr", "Extended producer responsibility"),
  ("Sol_awarness", "Awareness raising measures"),
  ("Sol_segregation", "Segregation of waste"),
  ("Sol_upcycling", "Upcycling of plastics"),
  ("Sol_recycling", "New raw materials"),
  ("Sol_education", "Education programs at schools"),
  ("Sol_capacity", "Capacity-building programs"),
  ("Sol_RD", "Research programs"),
  ("Sol_tax", "Taxes or levies on products"),
  ("Sol_ban", "Ban of products"),
  ("Sol_finance", "Financial measures"),
  ("Sol_infrastructure", "Infrastructural measures"),
  ("Sol_subsitutes", "Alternatives"),
  ("Sol_clean_up", "Clean-up campaigns"),
  ("Sol_enforcement", "Enforcement procedures"),
  ("Sol_monitoring", "Monitoring procedures"),
  ("Traditions to build on (free-hand)", "Freehand or NA"),
]

NOTES = {
  "Actortype": (
    "Unknown ordinary household resident (third respondent, 17 Feb 2026). "
    "Stakeholder list: household. Codebook: household."
  ),
  "Problem_awareness_pop": (
    "People want comfortable life and resist segregation efforts — behavioural "
    "compliance gap despite concern about plastics."
  ),
  "Problem_awareness_pol": (
    "Not much done yet by government — limited policy action and implementation."
  ),
  "Problem_severity": (
    "Very much concerned — plastics non-degradable while other waste decomposes; "
    "health impacts with different diseases."
  ),
  "Problem_littering": (
    "Open-space dumping problem — addressed culturally by prohibiting waste near "
    "temples/saint places."
  ),
  "Problem_consumption": (
    "More plastics in packaging and bottles; increased diaper use by babies and "
    "elderly people."
  ),
  "Problem_waste_mgmt": (
    "Plastics do not decompose unlike other waste; government has not done much; "
    "daily collection alone insufficient."
  ),
  "Problem_alternatives": (
    "People do not want to return to traditions — plastics make life more "
    "comfortable."
  ),
  "Problem_waste_segregation": (
    "No wish to make efforts for segregation — comfort prioritized over "
    "source separation."
  ),
  "Impacts": (
    "Health: different diseases appear from plastic pollution."
  ),
  "Local government": (
    "Mayor and municipality responsible; daily collection vehicle; municipalities "
    "place god monuments to create saint places deterring open dumping."
  ),
  "Households": (
    "Residents resist segregation and prefer plastic convenience."
  ),
  "Enforcement": (
    "Weak formal government action; informal enforcement through sacred-site "
    "taboo against littering near temples."
  ),
  "Policies in place": (
    "Daily municipal waste-collection vehicle; cultural-religious littering "
    "restrictions near saint places/temples; god monuments installed by "
    "municipalities."
  ),
  "Pol_awarness": (
    "Sacred-site norm communicated — people must not throw waste near temples; "
    "many saint places exist for this reason."
  ),
  "Pol_clean_up": (
    "Municipal god monuments create saint places to reduce open-space dumping."
  ),
  "Pol_waste_collection": (
    "Daily waste pick-up organised by municipality."
  ),
  "Sol_awarness": (
    "Leverage cultural-religious awareness that waste cannot be thrown near "
    "saint places and temples."
  ),
  "Sol_clean_up": (
    "Municipalities install god monuments to designate saint places and deter "
    "open dumping."
  ),
  "Sol_enforcement": (
    "Sacred-site taboo as informal enforcement mechanism complementing weak "
    "government action."
  ),
  "Traditions to build on (free-hand)": (
    "Religious tradition: prohibition on discarding waste near temples/saint "
    "places; municipalities replicate via god monuments. Material traditions "
    "rejected — respondents prefer plastic convenience over returning to older "
    "practices."
  ),
}


def hfill(hex_):
  return PatternFill("solid", fgColor=hex_)


def hfont(bold=False, size=10, color="000000", italic=False):
  return Font(bold=bold, size=size, color=color, italic=italic)


thin = Side(style="thin", color="BFBFBF")


def thin_border():
  return Border(left=thin, right=thin, top=thin, bottom=thin)


WRAP_ALIGN = Alignment(wrap_text=True, vertical="top", horizontal="left")
CENTER_ALIGN = Alignment(wrap_text=True, vertical="center", horizontal="center")


def build_workbook():
  wb = openpyxl.Workbook()

  ws = wb.active
  ws.title = "Interview Coding"

  ws.merge_cells("A1:D1")
  title = ws["A1"]
  title.value = "Plastic Pollution Governance — Interview Coding (Codebook)"
  title.fill = hfill(DARK_BLUE)
  title.font = hfont(bold=True, size=13, color=WHITE)
  title.alignment = CENTER_ALIGN

  ws.merge_cells("A2:D2")
  meta = ws["A2"]
  meta.value = (
    "Interview: NPL_25  |  Country: NEPAL  |  Ordinary household resident (#3)  |  "
    "Actor: household  |  17 Feb 2026"
  )
  meta.fill = hfill(MID_BLUE)
  meta.font = hfont(size=9, color=WHITE)
  meta.alignment = CENTER_ALIGN

  headers = ["Variable", "Code", "Codebook definition", "Coding notes"]
  for col, header in enumerate(headers, start=1):
    cell = ws.cell(row=4, column=col, value=header)
    cell.fill = hfill(LIGHT_BLUE)
    cell.font = hfont(bold=True, size=10, color=DARK_BLUE)
    cell.alignment = CENTER_ALIGN
    cell.border = thin_border()

  row = 5
  for variable, definition in CODEBOOK:
    values = [
      variable,
      CODING[variable],
      definition,
      NOTES.get(variable, ""),
    ]
    fill = hfill(LIGHT_GREEN) if row % 2 == 0 else hfill(WHITE)
    for col, value in enumerate(values, start=1):
      cell = ws.cell(row=row, column=col, value=value)
      cell.fill = fill
      cell.font = hfont(size=10)
      cell.alignment = WRAP_ALIGN
      cell.border = thin_border()
    row += 1

  ws.column_dimensions["A"].width = 34
  ws.column_dimensions["B"].width = 28
  ws.column_dimensions["C"].width = 52
  ws.column_dimensions["D"].width = 58
  ws.freeze_panes = "A5"

  wide = wb.create_sheet("Wide Format")
  wide_headers = [variable for variable, _ in CODEBOOK]
  for col, header in enumerate(wide_headers, start=1):
    cell = wide.cell(row=1, column=col, value=header)
    cell.fill = hfill(LIGHT_BLUE)
    cell.font = hfont(bold=True, size=9, color=DARK_BLUE)
    cell.alignment = CENTER_ALIGN
    cell.border = thin_border()
  for col, header in enumerate(wide_headers, start=1):
    cell = wide.cell(row=2, column=col, value=CODING[header])
    cell.fill = hfill(ORANGE)
    cell.alignment = WRAP_ALIGN
    cell.border = thin_border()
    wide.column_dimensions[get_column_letter(col)].width = max(14, min(28, len(header) + 2))

  ref = wb.create_sheet("Codebook Reference")
  ref.merge_cells("A1:B1")
  ref["A1"].value = "Codebook Variable Definitions"
  ref["A1"].fill = hfill(DARK_BLUE)
  ref["A1"].font = hfont(bold=True, size=12, color=WHITE)
  ref["A1"].alignment = CENTER_ALIGN
  for idx, (variable, definition) in enumerate(CODEBOOK, start=3):
    ref.cell(row=idx, column=1, value=variable).font = hfont(bold=True)
    ref.cell(row=idx, column=2, value=definition).alignment = WRAP_ALIGN
  ref.column_dimensions["A"].width = 34
  ref.column_dimensions["B"].width = 70

  stakeholder = wb.create_sheet("Stakeholder Analysis")
  stakeholder.merge_cells("A1:B1")
  stakeholder["A1"].value = "Stakeholder Analysis — NPL_25"
  stakeholder["A1"].fill = hfill(DARK_BLUE)
  stakeholder["A1"].font = hfont(bold=True, size=12, color=WHITE)
  stakeholder["A1"].alignment = CENTER_ALIGN
  for idx, (key, value) in enumerate(STAKEHOLDER.items(), start=3):
    stakeholder.cell(row=idx, column=1, value=key).font = hfont(bold=True)
    cell = stakeholder.cell(row=idx, column=2, value=value)
    cell.alignment = WRAP_ALIGN
  stakeholder.column_dimensions["A"].width = 28
  stakeholder.column_dimensions["B"].width = 90

  sources = wb.create_sheet("Sources")
  sources.merge_cells("A1:B1")
  sources["A1"].value = "Source Documents Used for Verification"
  sources["A1"].fill = hfill(DARK_BLUE)
  sources["A1"].font = hfont(bold=True, size=12, color=WHITE)
  sources["A1"].alignment = CENTER_ALIGN
  for idx, (doc, detail) in enumerate(SOURCES, start=3):
    sources.cell(row=idx, column=1, value=doc).font = hfont(bold=True)
    sources.cell(row=idx, column=2, value=detail).alignment = WRAP_ALIGN
  sources.column_dimensions["A"].width = 42
  sources.column_dimensions["B"].width = 60

  wb.save(OUTPUT_PATH)
  print(f"Saved {OUTPUT_PATH}")


if __name__ == "__main__":
  build_workbook()
