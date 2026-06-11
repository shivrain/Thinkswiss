from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUTPUT = Path("plastic_pollution_index_internship_presentation.pptx")

BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)
NEAR_BLACK = RGBColor(20, 20, 20)
DARK_GRAY = RGBColor(70, 70, 70)
MID_GRAY = RGBColor(140, 140, 140)
LIGHT_GRAY = RGBColor(238, 238, 238)


def set_run(run, size=18, bold=False, color=BLACK):
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_textbox(slide, text, left, top, width, height, size=18, bold=False,
                color=BLACK, align=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    if align is not None:
        p.alignment = align
    for run in p.runs:
        set_run(run, size=size, bold=bold, color=color)
    return box


def add_header(slide, prs, title, eyebrow=None):
    if eyebrow:
        add_textbox(slide, eyebrow.upper(), Inches(0.65), Inches(0.35),
                    Inches(11.7), Inches(0.28), size=8.5, bold=True,
                    color=DARK_GRAY)
    add_textbox(slide, title, Inches(0.65), Inches(0.68), Inches(11.7),
                Inches(0.62), size=27, bold=True, color=BLACK)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65),
                                  Inches(1.38), Inches(12.0), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = BLACK
    line.line.fill.background()


def add_footer(slide, slide_no):
    add_textbox(slide, "Plastic Pollution Index | Internship Presentation",
                Inches(0.65), Inches(7.08), Inches(6.5), Inches(0.25),
                size=8.5, color=MID_GRAY)
    add_textbox(slide, f"{slide_no:02d}", Inches(12.25), Inches(7.08),
                Inches(0.45), Inches(0.25), size=8.5, color=MID_GRAY,
                align=PP_ALIGN.RIGHT)


def add_bullets(slide, items, left, top, width, height, size=16,
                bullet_indent=0, line_spacing=0.88):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p.text = f"- {text}"
        p.level = level
        p.space_after = Pt(6 * line_spacing)
        p.font.name = "Aptos"
        p.font.size = Pt(size - (level * 1.5))
        p.font.color.rgb = BLACK if level == 0 else DARK_GRAY
        if bullet_indent:
            p.margin_left = Pt(bullet_indent + level * 18)
            p.margin_first_line = Pt(-10)
    return box


def add_numbered_card(slide, number, title, body, left, top, width, height):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top,
                                  width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = BLACK
    card.line.width = Pt(1.1)

    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, left + Inches(0.18),
                                    top + Inches(0.18), Inches(0.38),
                                    Inches(0.38))
    circle.fill.solid()
    circle.fill.fore_color.rgb = BLACK
    circle.line.fill.background()
    add_textbox(slide, str(number), left + Inches(0.18), top + Inches(0.225),
                Inches(0.38), Inches(0.18), size=8.5, bold=True,
                color=WHITE, align=PP_ALIGN.CENTER)

    add_textbox(slide, title, left + Inches(0.68), top + Inches(0.16),
                width - Inches(0.85), Inches(0.35), size=11.5, bold=True)
    add_textbox(slide, body, left + Inches(0.18), top + Inches(0.70),
                width - Inches(0.36), height - Inches(0.80), size=9.4,
                color=DARK_GRAY)


def add_section_slide(prs, section_no, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = WHITE
    add_textbox(slide, f"PART {section_no}", Inches(0.75), Inches(1.65),
                Inches(3.0), Inches(0.30), size=11, bold=True,
                color=DARK_GRAY)
    add_textbox(slide, title, Inches(0.75), Inches(2.05), Inches(9.7),
                Inches(0.85), size=35, bold=True)
    add_textbox(slide, subtitle, Inches(0.78), Inches(3.05), Inches(9.8),
                Inches(0.40), size=16, color=DARK_GRAY)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.78),
                                  Inches(3.75), Inches(5.2), Inches(0.035))
    line.fill.solid()
    line.fill.fore_color.rgb = BLACK
    line.line.fill.background()
    return slide


def add_table(slide, headers, rows, left, top, width, height, col_widths=None,
              font_size=10.2):
    table = slide.shapes.add_table(len(rows) + 1, len(headers), left, top,
                                   width, height).table
    if col_widths:
        for idx, col_width in enumerate(col_widths):
            table.columns[idx].width = col_width
    for idx, header in enumerate(headers):
        cell = table.cell(0, idx)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = NEAR_BLACK
        for p in cell.text_frame.paragraphs:
            p.font.name = "Aptos"
            p.font.size = Pt(font_size)
            p.font.bold = True
            p.font.color.rgb = WHITE
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r_idx % 2 else LIGHT_GRAY
            for p in cell.text_frame.paragraphs:
                p.font.name = "Aptos"
                p.font.size = Pt(font_size)
                p.font.color.rgb = BLACK
    return table


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide_no = 1

    # Title
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    add_textbox(slide, "Plastic Pollution Index", Inches(0.75), Inches(1.25),
                Inches(11.5), Inches(0.8), size=40, bold=True)
    add_textbox(slide, "Policy coding workflow, open questions, and suggestions",
                Inches(0.78), Inches(2.18), Inches(10.8), Inches(0.42),
                size=17, color=DARK_GRAY)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.78),
                                  Inches(3.05), Inches(5.3), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = BLACK
    line.line.fill.background()
    add_textbox(slide, "Internship Presentation", Inches(0.78),
                Inches(3.35), Inches(3.8), Inches(0.35), size=14,
                bold=True)
    add_textbox(slide, "Nepal policy table and coding process",
                Inches(0.78), Inches(3.77), Inches(4.8), Inches(0.32),
                size=12.5, color=DARK_GRAY)
    add_footer(slide, slide_no)
    slide_no += 1

    # Agenda
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, prs, "Presentation structure", "Overview")
    agenda = [
        ("01", "What I did", "Workflow from policy identification to the final coded table."),
        ("02", "Questions", "Methodological decisions to align before finalizing the index."),
        ("03", "Suggestions", "Process improvements for search, coding, quality control, and handoff."),
    ]
    for idx, (num, title, body) in enumerate(agenda):
        top = Inches(2.02 + idx * 1.36)
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.82), top,
                                     Inches(11.7), Inches(0.95))
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = BLACK
        add_textbox(slide, num, Inches(1.03), top + Inches(0.24),
                    Inches(0.58), Inches(0.30), size=13, bold=True)
        add_textbox(slide, title, Inches(1.85), top + Inches(0.17),
                    Inches(3.0), Inches(0.32), size=18, bold=True)
        add_textbox(slide, body, Inches(1.86), top + Inches(0.52),
                    Inches(9.5), Inches(0.25), size=11.5, color=DARK_GRAY)
    add_footer(slide, slide_no)
    slide_no += 1

    add_section_slide(prs, "01", "What I did",
                      "From policy search and source checks to the final coded table")
    slide_no += 1

    # Workflow
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, prs, "Workflow: how I arrived at the table", "What I did")
    workflow = [
        ("Scoping and inventory",
         "Defined the working policy universe and expanded the list beyond the initial entries, including relevant provincial and municipal policies."),
        ("Source collection",
         "Located policy documents, prioritized official sources where available, and captured links or document references for the tracker."),
        ("AI-assisted extraction",
         "Uploaded policy documents to Claude to extract code-relevant details and draft comments for the policy-coding sheet."),
        ("Manual coding",
         "Manually coded selected policies where the document was ambiguous, high-impact, or where automated output needed verification."),
        ("Cross-checking",
         "Checked Claude outputs against original policy text, corrected hallucinated or unsupported claims, and added clarifying comments."),
        ("Table integration",
         "Entered codes and notes into the compiled sheet, crossed completed policies off the master list, and downloaded coded policy documents for handoff."),
    ]
    lefts = [0.70, 2.78, 4.86, 6.94, 9.02, 11.10]
    for idx, (title, body) in enumerate(workflow, start=1):
        add_numbered_card(slide, idx, title, body, Inches(lefts[idx - 1]),
                          Inches(2.0), Inches(1.76), Inches(3.05))
        if idx < len(workflow):
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                           Inches(lefts[idx - 1] + 1.68),
                                           Inches(3.28), Inches(0.45),
                                           Inches(0.28))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = MID_GRAY
            arrow.line.fill.background()
    add_textbox(slide,
                "Result: a coded policy table with comments, source trail, and flagged methodological questions for team review.",
                Inches(0.78), Inches(5.65), Inches(11.6), Inches(0.45),
                size=14, bold=True, color=BLACK)
    add_footer(slide, slide_no)
    slide_no += 1

    # Comprehensive work completed
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, prs, "Comprehensive list of work completed", "What I did")
    left_items = [
        "Built and refined the Nepal policy list used for coding.",
        "Added relevant provincial and local policies where they affected plastic or waste governance.",
        "Identified border cases: main Acts, directives, local plans, reports, reviews, and consultation papers.",
        "Uploaded policy documents to Claude for first-pass extraction and coding support.",
        "Reviewed Claude-generated comments and corrected hallucinated or unsupported details.",
        "Manually coded selected policies when automation was not reliable enough.",
    ]
    right_items = [
        "Cross-checked coded entries against policy documents and available source links.",
        "Added comments explaining uncertain or interpretive coding decisions.",
        "Entered final code values into the compiled table and updated the master tracking list.",
        "Marked completed policies by crossing them off the main list.",
        "Downloaded and organized coded policy documents for auditability.",
        "Flagged unresolved questions by table column for supervisor/team alignment.",
    ]
    add_bullets(slide, left_items, Inches(0.85), Inches(1.82),
                Inches(5.55), Inches(4.95), size=14)
    add_bullets(slide, right_items, Inches(6.85), Inches(1.82),
                Inches(5.55), Inches(4.95), size=14)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.55),
                                  Inches(1.75), Inches(0.02), Inches(4.85))
    line.fill.solid()
    line.fill.fore_color.rgb = LIGHT_GRAY
    line.line.fill.background()
    add_footer(slide, slide_no)
    slide_no += 1

    # Quality controls
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, prs, "Quality controls used before finalizing entries",
               "What I did")
    rows = [
        ("Source check", "Verified policy title, date/year, jurisdiction, and document/source link where possible."),
        ("Coding support", "Used Claude for extraction, but treated it as a draft rather than a final authority."),
        ("Manual review", "Re-coded or checked policies with ambiguous scope, unclear legal status, or important index implications."),
        ("Comment trail", "Added comments where the code depended on interpretation or where the policy text was not explicit."),
        ("Completion tracking", "Updated the compiled sheet and crossed completed policies off the main list to avoid duplication."),
    ]
    add_table(slide, ["Check", "Purpose"], rows, Inches(0.85), Inches(1.75),
              Inches(11.65), Inches(4.65), col_widths=[Inches(2.7),
                                                       Inches(8.95)],
              font_size=11)
    add_footer(slide, slide_no)
    slide_no += 1

    add_section_slide(prs, "02", "Questions",
                      "Decisions needed to make coding consistent across policies")
    slide_no += 1

    # Inclusion questions
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, prs, "Questions: inclusion of policies", "Questions")
    rows = [
        ("Provincial / municipal policies",
         "Should local-level policies be included when they contain plastic or waste provisions that would otherwise be missing?"),
        ("Main Acts vs. directives",
         "Should smaller directives be coded separately or documented under the parent Act/strategy? Example: directives linked to broader biodiversity strategy documents."),
        ("Lake, biodiversity, tourism, and preservation plans",
         "Should sectoral plans that mention plastics be included if they are enacted centrally but are not plastic-specific?"),
        ("Reports, reviews, and consultation papers",
         "Should government-published reviews, implementation reports, and consultation papers be coded, used as context, or excluded?"),
    ]
    add_table(slide, ["Decision area", "Question for alignment"], rows,
              Inches(0.75), Inches(1.73), Inches(11.85), Inches(4.55),
              col_widths=[Inches(3.4), Inches(8.45)], font_size=10.4)
    add_textbox(slide,
                "Decision needed: define a consistent inclusion rule before adding or removing borderline documents.",
                Inches(0.82), Inches(6.43), Inches(11.45), Inches(0.32),
                size=12.5, bold=True)
    add_footer(slide, slide_no)
    slide_no += 1

    # Coding questions
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, prs, "Questions: coding rules by table column", "Questions")
    rows = [
        ("F: Quantifiable targets",
         "Should targets count only when they are plastic-specific, or can broader waste targets count when plastics are clearly within scope?"),
        ("G: Type of policy",
         "How should legislation be distinguished from orders, rules, and regulations? Does plastic-specific focus affect the score?"),
        ("I: Policy integration",
         "If an environment, tourism, or lake-preservation policy mentions plastic, how integrated must plastic be to score highly?"),
        ("N: Plastics life cycle",
         "If a policy covers several life-cycle stages but omits recycling or end-of-life treatment, should it receive a full or partial score?"),
        ("P: Budget",
         "Should only plastic/waste-specific budget lines count, or can a general policy/program budget count when plastics are part of implementation?"),
    ]
    add_table(slide, ["Column", "Question for alignment"], rows,
              Inches(0.75), Inches(1.68), Inches(11.85), Inches(4.95),
              col_widths=[Inches(2.45), Inches(9.40)], font_size=10.1)
    add_footer(slide, slide_no)
    slide_no += 1

    add_section_slide(prs, "03", "Suggestions",
                      "Ways to make policy search, coding, and review easier to audit")
    slide_no += 1

    # Suggestions
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, prs, "Suggestions to improve the coding workflow",
               "Suggestions")
    suggestions = [
        ("Search prompt", "Revise the Claude policy-search prompt to prioritize official government sources and return source URLs."),
        ("Evidence standard", "Require a quote, page, section, or paragraph reference for every coded value."),
        ("Chronology", "Ask Claude to arrange laws chronologically and group amendments with the principal Act."),
        ("Decision log", "Maintain a short log of inclusion and scoring decisions for recurring border cases."),
        ("Provenance fields", "Track source type, official/non-official status, manual check status, and reviewer initials/date."),
        ("Review pass", "Run a final consistency pass across high-risk columns after the inclusion rules are agreed."),
    ]
    for idx, (title, body) in enumerate(suggestions):
        col = idx % 2
        row = idx // 2
        left = Inches(0.82 + col * 6.0)
        top = Inches(1.77 + row * 1.45)
        rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top,
                                      Inches(5.45), Inches(1.05))
        rect.fill.solid()
        rect.fill.fore_color.rgb = WHITE
        rect.line.color.rgb = BLACK
        add_textbox(slide, title, left + Inches(0.20), top + Inches(0.16),
                    Inches(5.0), Inches(0.24), size=12.5, bold=True)
        add_textbox(slide, body, left + Inches(0.20), top + Inches(0.48),
                    Inches(5.0), Inches(0.38), size=10.4, color=DARK_GRAY)
    add_footer(slide, slide_no)
    slide_no += 1

    # Discussion
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, prs, "Discussion points for supervisor/team review",
               "Close")
    items = [
        "Confirm whether provincial and municipal policies should be included in the final table.",
        "Confirm how to treat parent Acts, directives, amendments, and strategy-linked instruments.",
        "Decide whether reports, reviews, and consultation papers are codable policy documents or supporting evidence.",
        "Finalize scoring rules for targets, policy type, integration, life-cycle coverage, and budget.",
        "Apply the agreed rules consistently across the coded table before final handoff.",
    ]
    add_bullets(slide, items, Inches(1.1), Inches(1.95), Inches(10.9),
                Inches(3.75), size=16)
    add_textbox(slide, "Thank you", Inches(1.1), Inches(6.15),
                Inches(4.0), Inches(0.45), size=24, bold=True)
    add_textbox(slide, "Questions and feedback", Inches(1.12),
                Inches(6.55), Inches(4.0), Inches(0.25), size=12,
                color=DARK_GRAY)
    add_footer(slide, slide_no)

    prs.save(OUTPUT)


if __name__ == "__main__":
    build_deck()
    print(f"Created {OUTPUT}")
