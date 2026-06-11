#!/usr/bin/env python3
"""Generate PPI Nepal internship presentation — black & white, professional."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── Palette ──────────────────────────────────────────────────────────────────
BLK = RGBColor(0,   0,   0)
WHT = RGBColor(255, 255, 255)
DG  = RGBColor(50,  50,  50)    # dark gray  — body text
MG  = RGBColor(130, 130, 130)   # mid gray   — labels / captions
LG  = RGBColor(245, 245, 245)   # light gray — box fill
BDR = RGBColor(195, 195, 195)   # border gray

W = 13.33
H = 7.5

prs = Presentation()
prs.slide_width  = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]   # blank layout


# ── Primitives ────────────────────────────────────────────────────────────────

def rect(sl, l, t, w, h, fill=LG, lc=BDR, lw=0.75):
    shp = sl.shapes.add_shape(
        1,  # MSO_AUTO_SHAPE_TYPE.RECTANGLE
        Inches(l), Inches(t), Inches(w), Inches(h)
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = lc
    shp.line.width = Pt(lw)
    return shp


def tb(sl, l, t, w, h):
    return sl.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))


def p(shp, txt, sz=11, b=False, i=False, c=DG, a=PP_ALIGN.LEFT):
    """Add text to a textbox shape, reusing the first empty paragraph."""
    tf = shp.text_frame if hasattr(shp, "text_frame") else shp
    tf.word_wrap = True
    pars = tf.paragraphs
    if len(pars) == 1 and not pars[0].runs:
        par = pars[0]
    else:
        par = tf.add_paragraph()
    par.alignment = a
    r = par.add_run()
    r.text       = txt
    r.font.name  = "Calibri"
    r.font.size  = Pt(sz)
    r.font.bold  = b
    r.font.italic = i
    r.font.color.rgb = c


def bars(sl):
    """Top and bottom thin black accent bars."""
    rect(sl, 0, 0,      W, 0.07, fill=BLK, lc=BLK, lw=0)
    rect(sl, 0, H - .07, W, 0.07, fill=BLK, lc=BLK, lw=0)


def header(sl, title, sub=None):
    """Standard slide header.  Returns y-offset where content area starts."""
    t = tb(sl, .5, .22, 12.33, .60)
    p(t, title, sz=21, b=True, c=BLK)
    rect(sl, .5, .88, 12.33, .025, fill=BLK, lc=BLK, lw=0)
    if sub:
        s = tb(sl, .5, .93, 12.33, .40)
        p(s, sub, sz=9.5, i=True, c=MG)
        return 1.44
    return 1.20


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 1 — TITLE
# ═══════════════════════════════════════════════════════════════════════════════
s1 = prs.slides.add_slide(BLANK)
bars(s1)
# Vertical black strip on the left
rect(s1, 0, 0, .28, H, fill=BLK, lc=BLK, lw=0)

t = tb(s1, 1.1, 1.95, 11.4, 1.25)
p(t, "PLASTIC POLLUTION INDEX — NEPAL", sz=34, b=True, c=BLK)

rect(s1, 1.1, 3.28, 7.5, .045, fill=BLK, lc=BLK, lw=0)

t = tb(s1, 1.1, 3.40, 11.4, .55)
p(t, "Policy Research & Coding  ·  Internship Work Summary", sz=14, c=DG)

t = tb(s1, 1.1, 4.10, 4.0, .35)
p(t, "June 2026", sz=11, c=MG)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 2 — WHAT I DID  (9-step snake-flow grid)
# ═══════════════════════════════════════════════════════════════════════════════
s2 = prs.slides.add_slide(BLANK)
bars(s2)
yc = header(s2, "WHAT I DID", "Policy Research & Coding Workflow")

STEPS = [
    ("01", "Initial Policy Discovery",
     "Claude AI + Ecolex; first scan for national & provincial Nepal policies"),
    ("02", "Official Source Search",
     "Nepali government websites only; keyword + site-type + PDF filetype filtering"),
    ("03", "Policy List Compilation",
     "Curated list including provincial level; arranged chronologically per Keerthi's suggestion"),
    ("04", "AI-Assisted Coding",
     "Uploaded policy documents to Claude using the provided coding prompt"),
    ("05", "Review & Annotation",
     "Flagged hallucinations with comments; compiled emerging questions for supervisor review"),
    ("06", "Manual Coding",
     "Directly coded select policies; left uncoded any not found on official sources"),
    ("07", "Master Sheet Update",
     "Added all finalised codes to the compiled spreadsheet"),
    ("08", "Cross-Checking",
     "Verified against the main policy list; provincial-level doubts flagged as comments"),
    ("09", "Documentation & Upload",
     "Coded sheets + original policy documents uploaded to SwitchDrive"),
]

BW = 3.68;  BH = 1.54;  CG = 0.285;  RG = 0.21
SX = 0.45;  SY = yc + 0.04
AW = 0.32;  AH = 0.40   # arrow textbox size

for i, (num, title, desc) in enumerate(STEPS):
    row = i // 3
    base_col = i % 3
    col = (2 - base_col) if row % 2 == 1 else base_col
    x = SX + col * (BW + CG)
    y = SY + row * (BH + RG)

    rect(s2, x, y, BW, BH, fill=LG, lc=BDR, lw=0.75)

    # Step number (large, top-left)
    t = tb(s2, x + .12, y + .07, .58, .50)
    p(t, num, sz=20, b=True, c=BLK)

    # Step title
    t = tb(s2, x + .73, y + .07, BW - .85, .46)
    p(t, title, sz=10, b=True, c=BLK)

    # Step description
    t = tb(s2, x + .12, y + .57, BW - .24, .88)
    p(t, desc, sz=8.5, c=DG)

# Horizontal arrows (same physical x for every row, direction flips)
for row in range(3):
    arrow_txt = "→" if row % 2 == 0 else "←"
    ay = SY + row * (BH + RG) + BH / 2 - AH / 2
    for gap in range(2):
        ax_c = SX + (gap + 1) * (BW + CG) - CG / 2
        t = tb(s2, ax_c - AW / 2, ay, AW, AH)
        p(t, arrow_txt, sz=14, c=MG, a=PP_ALIGN.CENTER)

# Vertical arrows at the turning corners
for row in range(2):
    vc = 2 if row % 2 == 0 else 0   # col 2 turns row 0→1; col 0 turns row 1→2
    vx_c = SX + vc * (BW + CG) + BW / 2
    vy = SY + (row + 1) * (BH + RG) - RG / 2 - AH / 2
    t = tb(s2, vx_c - AW / 2, vy, AW, AH)
    p(t, "↓", sz=14, c=MG, a=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 3 — Q1: INCLUSION OF POLICIES
# ═══════════════════════════════════════════════════════════════════════════════
s3 = prs.slides.add_slide(BLANK)
bars(s3)
y3 = header(s3, "Q1 — INCLUSION OF POLICIES",
            "Which policies should be included in the database?")

COLS3 = [
    ("Provincial &\nMunicipal Level", [
        "Many plastic-related policies exist at sub-national level — omitting them risks significant coverage gaps in the index.",
        "Should provincial/municipal-level acts and rules be systematically included?",
    ]),
    ("Main Acts &\nSub-Instruments", [
        "Do we list directives under main Acts separately, or only the parent Act?",
        "Key case: National Biodiversity & Strategy Action Plan and its sub-directives.",
        "Should lake/biodiversity master plans (enacted centrally) be included if they explicitly mention plastic?",
    ]),
    ("Border Cases", [
        "Reports — implementation or otherwise (many at local level).",
        "Officially published government reviews, e.g. Review of E & Solar PV Waste Management in Nepal (covers recycling, export to India).",
        "Consultation papers — e.g. Consultation Paper on Regulatory Framework for E-Waste Management.",
    ]),
]

CW3 = (12.33 - 2 * .20) / 3   # ≈ 3.977"
SX3 = 0.5

for ci, (col_title, bullets) in enumerate(COLS3):
    cx3 = SX3 + ci * (CW3 + .20)

    # Column header (black box, white text)
    rect(s3, cx3, y3, CW3, .60, fill=BLK, lc=BLK, lw=0)
    t = tb(s3, cx3 + .10, y3 + .06, CW3 - .20, .52)
    p(t, col_title, sz=10, b=True, c=WHT, a=PP_ALIGN.CENTER)

    # Bullets
    by = y3 + .72
    for btext in bullets:
        t = tb(s3, cx3 + .10, by, CW3 - .20, 1.12)
        p(t, "•  " + btext, sz=9, c=DG)
        by += 1.18


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 4 — Q2: CODING
# ═══════════════════════════════════════════════════════════════════════════════
s4 = prs.slides.add_slide(BLANK)
bars(s4)
y4 = header(s4, "Q2 — CODING",
            "Clarifications needed on specific spreadsheet columns")

ROWS4 = [
    ("Col F", "Quantifiable Targets",
     "Should targets be plastic-specific only, or does general waste management also count?"),
    ("Col G", "Type of Policy",
     "What distinguishes legislation from orders/rules/regulations? Is plastic-specific legislation required for a score of 1?"),
    ("Col I", "Policy Integration",
     "If a policy broadly covers environment without explicit plastic mention, is it 'well-integrated'? How do we score a tourism or lake preservation plan?"),
    ("Col N", "Plastics Life Cycle",
     "Which life-cycle stages are in scope? E.g. the EPA covers almost all stages but does not mention recycling — would it still score 1?"),
    ("Col P", "Budget",
     "Should the budget be explicitly for waste/plastic, or does a general policy budget (e.g. tourism promotion) also count?"),
]

RH4 = .90
SY4 = y4 + .05
REF_W  = .82
NAME_W = 1.60
Q_W    = 12.33 - .20 - REF_W - NAME_W - .38

for ri, (ref, name, question) in enumerate(ROWS4):
    ry = SY4 + ri * (RH4 + .07)
    bg = LG if ri % 2 == 0 else WHT
    rect(s4, .50, ry, 12.33, RH4, fill=bg, lc=BDR, lw=.5)

    # Column reference label
    t = tb(s4, .60, ry + .09, REF_W, RH4 - .18)
    p(t, ref, sz=9, b=True, c=MG, a=PP_ALIGN.CENTER)

    # Column name
    t = tb(s4, .60 + REF_W + .08, ry + .09, NAME_W, RH4 - .18)
    p(t, name, sz=9.5, b=True, c=BLK)

    # Thin vertical divider
    rect(s4, .60 + REF_W + NAME_W + .20, ry + .15,
         .015, RH4 - .30, fill=BDR, lc=BDR, lw=0)

    # Question
    t = tb(s4, .60 + REF_W + NAME_W + .24, ry + .09, Q_W, RH4 - .18)
    p(t, question, sz=9, c=DG)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 5 — SUGGESTIONS
# ═══════════════════════════════════════════════════════════════════════════════
s5 = prs.slides.add_slide(BLANK)
bars(s5)
y5 = header(s5, "SUGGESTIONS")

SUGGS = [
    ("01", "Refine the Policy-Search Claude Prompt",
     "Revise the prompt to prioritise official government websites as primary sources, "
     "reducing hallucinated or unverifiable policy references."),
    ("02", "Arrange Legislation Chronologically",
     "As recommended by Keerthi — edit the Claude prompt so that Acts and their amendments "
     "are returned together in chronological order, making the coding timeline easier to follow."),
    ("03", "Link Acts with Their Amendments",
     "Explore whether the Claude prompt can automatically pair each principal Act with its "
     "associated amendments, rather than listing them as separate entries in the database."),
]

SSX = 0.70
SSY = y5 + .35

for si, (num, title, desc) in enumerate(SUGGS):
    sy = SSY + si * 1.88

    # Number badge (filled black square)
    rect(s5, SSX, sy, .60, .60, fill=BLK, lc=BLK, lw=0)
    t = tb(s5, SSX, sy, .60, .60)
    p(t, num, sz=11, b=True, c=WHT, a=PP_ALIGN.CENTER)

    # Title
    t = tb(s5, SSX + .78, sy, 11.6, .46)
    p(t, title, sz=13, b=True, c=BLK)

    # Description
    t = tb(s5, SSX + .78, sy + .50, 11.6, 1.15)
    p(t, desc, sz=10, c=DG)

    # Separator rule between items
    if si < len(SUGGS) - 1:
        rect(s5, SSX, sy + 1.70, 11.85, .015, fill=BDR, lc=BDR, lw=0)


# ── Save ──────────────────────────────────────────────────────────────────────
OUT = "/workspace/PPI_Nepal_Presentation.pptx"
prs.save(OUT)
print(f"Saved → {OUT}")
