#!/usr/bin/env python3
"""Generate 4P Index coding spreadsheet for Loi 2013-10 CGCT (English only).

Source text extracted via OCR from the official dri.gouv.sn PDF (image-based, 68 pages).
Waste-competence provisions only; strict plastics-relevance filter.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT = "/workspace/4P_Index_Loi_2013-10_CGCT.xlsx"

COLUMNS = [
    ("A", "policy_name"),
    ("B", "policy_url"),
    ("C", "policy_year"),
    ("D", "policy_objective"),
    ("E", "policy_target"),
    ("F", "policy_target_text"),
    ("G", "policy_type"),
    ("H", "policy_type_justification"),
    ("I", "policy_integration"),
    ("J", "policy_sectors_list"),
    ("K", "policy_circularity"),
    ("L", "policy_lifecycle_phases_list"),
    ("M", "policy_budget"),
    ("N", "policy_budget_text"),
    ("O", "policy_score"),
    ("P", "instrument_type"),
    ("Q", "instrument_lifecycle_stage"),
    ("R", "instrument_description"),
    ("S", "instrument_in_force"),
    ("T", "instrument_implementation"),
    ("U", "instrument_implementation_text"),
    ("V", "instrument_score"),
    ("W", "comments"),
]

POLICY = {
    "policy_name": (
        "Law No. 2013-10 of 28 December 2013 establishing the General Code of Local Authorities (CGCT)"
    ),
    "policy_url": (
        "https://www.dri.gouv.sn/sites/default/files/an-documents/"
        "LOI%20N%202013%2010%20DU%2028%20DECEMBRE%202013.pdf"
    ),
    "policy_year": 2013,
    "policy_objective": (
        "Refound the decentralisation framework by defining competences of departments and communes "
        "(local authorities), including municipal cleanliness, household waste management, hygiene and public "
        "health — key subnational framework for solid waste, including plastics, without plastic-specific measures."
    ),
    "policy_target": 0,
    "policy_target_text": "",
    "policy_type": 1.0,
    "policy_type_justification": (
        "General Code adopted by the National Assembly on 19 December 2013 (urgency procedure) and promulgated "
        "on 28 December 2013 as Law No. 2013-10. Parliamentary legislation (Code), not executive implementing "
        "regulation (≠0.75)."
    ),
    "policy_integration": 1.0,
    "policy_sectors_list": (
        "local government, sanitation, environment, local governance, local public finance"
    ),
    "policy_circularity": 0.50,
    "policy_lifecycle_phases_list": "disposal, recycling",
    "policy_budget": 0.5,
    "policy_budget_text": (
        "Art. 195: household waste collection tax and street-sweeping tax among commune direct local taxes; "
        "Art. 202: obligatory commune operating expenditure on street maintenance/cleaning and local "
        "disinfection/hygiene services."
    ),
}

INSTRUMENTS = [
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Waste management",
        "instrument_description": (
            "Arts. 170 and 205: waste management and the fight against insalubrity are transferred "
            "competences — at city (ville) level for urban agglomerations (Art. 170) and at commune level "
            "under Livre IV environment competences (Art. 205), alongside communal environmental action plans."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            "Art. 203: obligatory expenses from competence transfers (+0.25 authority); "
            "Art. 4: competence transfers must include concomitant State resource transfer (+0.25 monitoring). "
            "No plastic-specific enforcement mechanism in these articles."
        ),
        "comments": (
            "Core subnational waste-governance pillar for solid waste including plastics; no explicit plastic mention."
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Waste management",
        "instrument_description": (
            "Arts. 119 and 202(11): municipal police must ensure public salubrity including cleaning, "
            "lighting and removal of obstructions on public streets; commune operating budgets must fund "
            "maintenance and cleaning of streets, roads and public squares on local authority territory."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            "Art. 118: mayor responsible for municipal police under State representative control (+0.25 authority); "
            "Art. 124: State may act if municipal authorities fail (+0.25 enforcement); "
            "Art. 201–202: obligatory cleaning expenditure must appear in commune budget (+0.25 monitoring)."
        ),
        "comments": (
            "Direct framework for public-space waste removal including plastic litter; complements Art. 205 competence."
        ),
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Waste management",
        "instrument_description": (
            "Art. 195(5): commune direct local taxes include the household waste collection tax "
            "and the street-sweeping tax to finance municipal waste services."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            "Assessment and collection modalities set by law (+0.25 authority); "
            "municipal councils may create additional centimes within legal maxima by deliberation (+0.25 monitoring). "
            "No article-specific penalty for non-payment in this provision."
        ),
        "comments": (
            "Local fiscal instrument financing household waste collection, including plastic fractions."
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Waste management",
        "instrument_description": (
            "Art. 202(14): obligatory commune operating expenditure on local disinfection and hygiene services "
            "under applicable regulations, linked to sworn hygiene agents (Art. 126) and commune prevention "
            "measures (Art. 307)."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": (
            "Art. 201–202: obligatory operating expenses must appear in budget (+0.25 authority). "
            "Enforcement mechanisms indirect; no article-specific sanction."
        ),
        "comments": (
            "Mandatory funding of local hygiene services supporting waste/insalubrity management."
        ),
    },
]


def build_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "4P Index Coding"

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    auto_fill = PatternFill("solid", fgColor="E2EFDA")

    for col_idx, (col_letter, col_name) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=f"{col_letter} — {col_name}")
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for row_idx, instrument in enumerate(INSTRUMENTS, start=2):
        row_data = {**POLICY, **instrument}
        for col_idx, (_, key) in enumerate(COLUMNS, start=1):
            if key in ("policy_score", "instrument_score"):
                continue
            ws.cell(row=row_idx, column=col_idx, value=row_data.get(key, ""))

        g_col, i_col, k_col, m_col = (get_column_letter(n) for n in (7, 9, 11, 13))
        p_col, t_col, o_col, v_col = (get_column_letter(n) for n in (16, 20, 15, 22))
        ws[f"{o_col}{row_idx}"] = f"=AVERAGE({g_col}{row_idx},{i_col}{row_idx},{k_col}{row_idx},{m_col}{row_idx},{p_col}{row_idx})"
        ws[f"{v_col}{row_idx}"] = f"=AVERAGE({p_col}{row_idx},{t_col}{row_idx})"
        ws[f"{o_col}{row_idx}"].fill = auto_fill
        ws[f"{v_col}{row_idx}"].fill = auto_fill

    widths = {
        "A": 48, "B": 52, "C": 10, "D": 52, "E": 12, "F": 60,
        "G": 10, "H": 52, "I": 14, "J": 40, "K": 14, "L": 36,
        "M": 12, "N": 52, "O": 12, "P": 14, "Q": 22, "R": 60,
        "S": 14, "T": 18, "U": 52, "V": 14, "W": 52,
    }
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = "A2"
    wb.save(OUTPUT)
    print(f"Saved {OUTPUT} ({len(INSTRUMENTS)} instrument rows)")


if __name__ == "__main__":
    build_workbook()
