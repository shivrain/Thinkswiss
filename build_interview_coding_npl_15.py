"""Generate interview coding spreadsheet for Nepal interview NPL_15.

Coding follows the Overview All Interviews codebook and is based on the
Suna Maya Margen (Kathmandu Metropolitan City Environment Inspector) interview
of 23 June 2025 — interview guideline and field notes. Uses expanded Impacts
taxonomy.
"""

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUTPUT_PATH = "/workspace/interview_coding_NPL_15.xlsx"

DARK_BLUE = "1F3864"
MID_BLUE = "2E75B6"
LIGHT_BLUE = "D9E2F3"
LIGHT_GREEN = "E2EFDA"
ORANGE = "FCE4D6"
WHITE = "FFFFFF"

SOURCES = [
  ("Interview notes", "Interview no. 2 — KTM, Suna Maya Margen, 23 June 2025"),
  ("Interview guideline", "KTM Environment Inspector structured guideline (Themes A–D)"),
  ("Codebook", "Overview All Interviews — NEW Codebook (expanded Impacts)"),
]

STAKEHOLDER = {
  "Interview ID": "NPL_15",
  "Country": "NEPAL",
  "Interview number": "2 (KTM Metropolitan City series)",
  "Interview date": "23 June 2025",
  "Interviewee": "Ms. Suna Maya Margen (30 years)",
  "Affiliation / role": (
    "Environment Inspector — Kathmandu Metropolitan City (KTM); public "
    "education; household, community, and school composting programmes"
  ),
  "Organization": (
    "Kathmandu Metropolitan City — Environment Division; manages landfill for "
    "23 municipalities across 4 districts; Doko Recyclers MoU (Cluster 7)"
  ),
  "Stakeholder list mapping": "local gov",
  "Coded actor type (codebook)": "governmental",
  "Actor classification rationale": (
    "Municipal government environment inspector implementing KTM waste policy, "
    "private-sector MoUs, and public education — local government actor."
  ),
  "Perspective": (
    "Municipal implementer — emphasizes land scarcity, MRF opposition, "
    "registered private-sector partnerships, multilayer plastic growth, and "
    "need for federal–provincial–local coordination on land"
  ),
  "Project context": (
    "Multilayer plastics rose from ~8% to ~15% of waste (no recycling value); "
    "Doko Recyclers collects Cluster 7; river-bed debris excavation; UNDP MoU "
    "on plastics recovery/study; single-use ban planned"
  ),
}

CODING = {
  "Country": "NPL",
  "Country name": "NEPAL",
  "ID": "NPL_15",
  "Actortype": "governmental",
  "Problem_awareness_pop": "medium",
  "Problem_awareness_pol": "medium",
  "Problem_severity": "high",
  "Problem_littering": "yes",
  "Problem_consumption": "yes",
  "Problem_recycling": "yes",
  "Problem_waste_mgmt": "yes",
  "Problem_production": "yes",
  "Problem_alternatives": "yes",
  "Problem_waste_segregation": "yes",
  "Problem_import": "yes",
  "Impacts": (
    "visual pollution, rivers, drainage blockage, health, wildlife, microplastics"
  ),
  "Governance": "yes",
  "Relevance_international_pol": "no",
  "Coordination_sectoral": "yes",
  "Coordination_levels": "yes",
  "Unclear_responsibilities": "yes",
  "Actors": "yes",
  "Federal government": "yes",
  "Provincial government": "yes",
  "Local government": "yes",
  "Students": "yes",
  "Private_sector": "yes",
  "Civil_society": "no",
  "Science": "yes",
  "Households": "yes",
  "Education institutions": "yes",
  "Private_companies(hotels, shops, etc.)": "yes",
  "Implementation issues": "yes",
  "Monitoring": "yes",
  "Financial_resources": "yes",
  "Research": "yes",
  "Infrastructure": "yes",
  "Enforcement": "yes",
  "Policies in place": "yes",
  "Pol_epr": "no",
  "Pol_import": "yes",
  "Pol_awarness": "yes",
  "Pol_education": "yes",
  "Pol_capacity": "yes",
  "Pol_RD": "yes",
  "Pol_tax": "no",
  "Pol_ban": "yes",
  "Pol_subsitutes": "yes",
  "Pol_clean_up": "yes",
  "Pol_upcycling": "no",
  "Pol_recycling": "yes",
  "Pol_waste_collection": "yes",
  "Solutions": "yes",
  "Sol_lead_agency": "yes",
  "Sol_responsibilities": "yes",
  "Sol_epr": "no",
  "Sol_awarness": "yes",
  "Sol_segregation": "yes",
  "Sol_upcycling": "no",
  "Sol_recycling": "yes",
  "Sol_education": "yes",
  "Sol_capacity": "yes",
  "Sol_RD": "yes",
  "Sol_tax": "no",
  "Sol_ban": "yes",
  "Sol_finance": "yes",
  "Sol_infrastructure": "yes",
  "Sol_subsitutes": "yes",
  "Sol_clean_up": "yes",
  "Sol_enforcement": "yes",
  "Sol_monitoring": "yes",
  "Traditions to build on (free-hand)": (
    "Newari festival practice of minimizing plastics — lapti (leaf plates); "
    "cloth bags at stores instead of polythene (noted as existing in some shops)"
  ),
}

CODEBOOK = [
  ("Country", "Country ISO code"),
  ("Country name", "Country name"),
  ("ID", "InterviewID: ISO_Nr"),
  ("Actortype", "governmental, pol_party, science, cso, igo, private_sector, shop_owner, hotel_owner, household"),
  ("Problem_awareness_pop", "high, medium, low, NA — lack of awareness among population mentioned"),
  ("Problem_awareness_pol", "high, medium, low, NA — lack of awareness among policy makers mentioned"),
  ("Problem_severity", "high, medium, low, NA"),
  ("Problem_littering", "Littering is a problem; was mentioned: yes/no"),
  ("Problem_consumption", "High consumption is a problem; was mentioned: yes/no"),
  ("Problem_recycling", "No or too little recycling is a problem; was mentioned: yes/no"),
  ("Problem_waste_mgmt", "Ill waste management is a problem; was mentioned: yes/no"),
  ("Problem_production", "High production rates; was mentioned: yes/no"),
  ("Problem_alternatives", "No good alternatives; was mentioned: yes/no"),
  ("Problem_waste_segregation", "Lack of segregation was mentioned: yes/no"),
  ("Problem_import", "Plastic imports are a problem; was mentioned: yes/no"),
  ("Impacts", "Impacts mentioned on water, cities, biodiversity, health, etc. or NA"),
  ("Governance", "Mentioned: yes/no"),
  ("Relevance_international_pol", "International treaties are relevant for national level policy"),
  ("Coordination_sectoral", "Lack of coordination across sectors"),
  ("Coordination_levels", "Lack of coordination across levels"),
  ("Unclear_responsibilities", "Unclear responsibilities among the involved actors"),
  ("Actors", "Mentioned: yes/no"),
  ("Federal government", "Mentioned: yes/no"),
  ("Provincial government", "Mentioned: yes/no"),
  ("Local government", "Mentioned: yes/no"),
  ("Students", "Mentioned: yes/no"),
  ("Private_sector", "Mentioned: yes/no"),
  ("Civil_society", "Mentioned: yes/no"),
  ("Science", "Mentioned: yes/no"),
  ("Households", "Mentioned: yes/no"),
  ("Education institutions", "Mentioned: yes/no"),
  ("Private_companies(hotels, shops, etc.)", "Mentioned: yes/no"),
  ("Implementation issues", "Mentioned: yes/no"),
  ("Monitoring", "Lack of monitoring"),
  ("Financial_resources", "Lack of financial resources"),
  ("Research", "Lack of research"),
  ("Infrastructure", "Lack of infrastructure"),
  ("Enforcement", "Lack of enforcement"),
  ("Policies in place", "Mentioned: yes/no"),
  ("Pol_epr", "Extended producer responsibility"),
  ("Pol_import", "Import regulations"),
  ("Pol_awarness", "Awareness campaigns"),
  ("Pol_education", "Education programs"),
  ("Pol_capacity", "Capacity building"),
  ("Pol_RD", "Research & development"),
  ("Pol_tax", "Levies or taxes on specific products (often bags)"),
  ("Pol_ban", "Bans of specific products"),
  ("Pol_subsitutes", "Alternatives like reusable bags, or banana leaf plates"),
  ("Pol_clean_up", "Clean-up campaigns"),
  ("Pol_upcycling", "Making a new product, like blacktop"),
  ("Pol_recycling", "Making a new raw material"),
  ("Pol_waste_collection", "Waste collection"),
  ("Solutions", "Mentioned: yes/no"),
  ("Sol_lead_agency", "Procedural measures to establish lead agency"),
  ("Sol_responsibilities", "Clear responsibilities"),
  ("Sol_epr", "Extended producer responsibility"),
  ("Sol_awarness", "Awareness raising measures"),
  ("Sol_segregation", "Segregation of waste"),
  ("Sol_upcycling", "Upcycling of plastics"),
  ("Sol_recycling", "New raw materials"),
  ("Sol_education", "Education programs at schools"),
  ("Sol_capacity", "Capacity-building programs"),
  ("Sol_RD", "Research programs"),
  ("Sol_tax", "Taxes or levies on products"),
  ("Sol_ban", "Ban of products"),
  ("Sol_finance", "Financial measures"),
  ("Sol_infrastructure", "Infrastructural measures"),
  ("Sol_subsitutes", "Alternatives"),
  ("Sol_clean_up", "Clean-up campaigns"),
  ("Sol_enforcement", "Enforcement procedures"),
  ("Sol_monitoring", "Monitoring procedures"),
  ("Traditions to build on (free-hand)", "Freehand or NA"),
]

NOTES = {
  "Actortype": (
    "Suna Maya Margen, KTM Environment Inspector. Stakeholder list: local gov. "
    "Codebook: governmental."
  ),
  "Problem_awareness_pop": (
    "Visual pollution linked to tourists and lifestyle; public opposition to "
    "waste facilities — people don't want waste everywhere."
  ),
  "Problem_awareness_pol": (
    "Ministry of Forest and Environment policies not implemented and under "
    "revision; 40-micron notification impact assessment needed."
  ),
  "Problem_severity": (
    "Biggest visual problem since 1998; multilayer plastics grew from ~8% to "
    "~15% of waste composition (no recycling value)."
  ),
  "Problem_littering": (
    "Dumped at river-banks; debris near river beds excavated and transferred."
  ),
  "Problem_consumption": (
    "Visual pollution attributed to tourists and lifestyle changes in Kathmandu."
  ),
  "Problem_recycling": (
    "Multilayer plastics have no value; technology/recycling factory needed; "
    "Doko Recyclers MoU for Cluster 7 collection."
  ),
  "Problem_waste_mgmt": (
    "Sufficient waste volume but unable to procure land; 23 municipalities in "
    "4 districts use KTM landfill site management."
  ),
  "Problem_production": (
    "Multilayer plastic share increased from ~8% to ~15%; import of plastics "
    "raised as open policy question."
  ),
  "Problem_alternatives": (
    "Lapti (Newari leaf plates) and other alternatives expensive; government "
    "subsidies for alternative-producing companies proposed."
  ),
  "Problem_waste_segregation": (
    "Can segregate in principle but no space/land within Kathmandu city for "
    "MRF and processing infrastructure."
  ),
  "Problem_import": (
    "Import of plastics noted as open policy question requiring assessment."
  ),
  "Impacts": (
    "Visual pollution: biggest visible problem since 1998. Rivers: dumping at "
    "river-banks. Drainage blockage: plastic in drains. Health: microplastics "
    "impacts need assessment. Wildlife: microplastics effects on environment/"
    "wildlife cited. Microplastics: explicit research need."
  ),
  "Coordination_sectoral": (
    "KTM works only with registered private sector (not informal); Doko "
    "Recyclers MoU; UNDP donor MoU on plastics recovery."
  ),
  "Coordination_levels": (
    "Land authority with federal government but implementation by LG — "
    "difficult for municipalities; need central/provincial land provision "
    "within metro/valley; waste transferred between LGs."
  ),
  "Unclear_responsibilities": (
    "Federal land authority vs local implementation; 23 municipalities rely "
    "on KTM landfill management across 4 districts."
  ),
  "Federal government": (
    "Ministry of Forest and Environment (policy under revision); land authority; "
    "subsidies to producers and alternative companies proposed."
  ),
  "Provincial government": (
    "Should provide land within metro city or valley; better coordination with "
    "district and LG needed."
  ),
  "Local government": (
    "KTM city primary responsibility; Environment Act 2077 BS; Natural Resource "
    "Management Act 2007; Mayors Forum (18) plastic provisions."
  ),
  "Students": (
    "School-level composting programmes run by environment inspector."
  ),
  "Private_sector": (
    "Doko Recyclers MoU Cluster 7; only registered private operators engaged — "
    "unregistered/informal sector not worked with."
  ),
  "Science": (
    "UNDP plastic study; need assessment of 40-micron notification; "
    "microplastics health/environment research."
  ),
  "Households": (
    "Household-level composting programmes; stores should offer cloth bags not "
    "polythene."
  ),
  "Education institutions": (
    "School-level composting as part of public education remit."
  ),
  "Private_companies(hotels, shops, etc.)": (
    "Retail stores — cloth bags instead of polythene; producer subsidies "
    "proposed."
  ),
  "Monitoring": (
    "40-micron ban notification impact assessment needed; pollution standards "
    "under Environment Act 2077."
  ),
  "Financial_resources": (
    "Sufficient waste but cannot procure land; donor support via UNDP MoU; "
    "subsidies for producers and alternative companies proposed."
  ),
  "Research": (
    "UNDP plastic study; microplastics impacts on health/environment/wildlife; "
    "40-micron notification assessment."
  ),
  "Infrastructure": (
    "Lack of land for MRF within Kathmandu; public opposition to MRF near "
    "neighbouring municipalities; recycling factory/technology needed."
  ),
  "Enforcement": (
    "MoFE policies not implemented; only registered private sector legalized — "
    "previously private operators not legalized."
  ),
  "Policies in place": (
    "Environment Act 2077 BS (air, water pollution standards); Natural "
    "Resource Management Act 2007; Mayors Forum (18) plastic/SWM provisions; "
    "KTM SWM private-sector MoU framework; 40-micron notification."
  ),
  "Pol_import": "Import of plastics raised as policy question requiring regulation.",
  "Pol_awarness": "Public education on reducing plastics — core inspector role.",
  "Pol_education": (
    "Public education; household, community, and school composting programmes."
  ),
  "Pol_capacity": (
    "MRF infrastructure; technology and recycling factory capacity needed."
  ),
  "Pol_RD": "UNDP plastic study; microplastics impact research needed.",
  "Pol_ban": "Single-use plastics ban — planning stage.",
  "Pol_subsitutes": (
    "Government subsidies for companies producing alternatives; lapti leaf plates "
    "at Newari festivals."
  ),
  "Pol_clean_up": (
    "River-bed debris excavation and transfer collaboration."
  ),
  "Pol_recycling": (
    "Doko Recyclers MoU Cluster 7; plastics recovery via UNDP MoU."
  ),
  "Pol_waste_collection": (
    "Solid waste management at city level; private-sector MoU collection clusters."
  ),
  "Sol_lead_agency": (
    "Better federal–provincial–district–LG coordination on land and rules."
  ),
  "Sol_responsibilities": (
    "Central/provincial governments provide land; LG implements; KTM manages "
    "regional landfill for 23 municipalities."
  ),
  "Sol_awarness": (
    "Continued public education; assessment of 40-micron notification to inform "
    "public understanding."
  ),
  "Sol_segregation": (
    "Segregation possible but requires land/MRF space — central/provincial "
    "support for sites within valley."
  ),
  "Sol_recycling": (
    "Recycling factory/technology; Doko Recyclers model expansion; UNDP "
    "plastics recovery."
  ),
  "Sol_education": (
    "School and community composting; public education on plastic reduction."
  ),
  "Sol_capacity": (
    "MRF facilities; registered private-sector capacity via MoUs."
  ),
  "Sol_RD": (
    "Microplastics health/environment/wildlife studies; 40-micron ban "
    "impact assessment."
  ),
  "Sol_ban": "Single-use plastics ban in planning.",
  "Sol_finance": (
    "Subsidies to producers and alternative-product companies; UNDP donor "
    "support."
  ),
  "Sol_infrastructure": (
    "MRF land provided by central/provincial government within metro/valley; "
    "recycling factory technology."
  ),
  "Sol_subsitutes": (
    "Subsidize alternative producers; promote lapti and cloth bags at stores."
  ),
  "Sol_clean_up": "River-bed debris excavation and transfer programmes.",
  "Sol_enforcement": (
    "Rules and regulations within KTM areas; work only with registered private "
    "operators; implement revised MoFE policies."
  ),
  "Sol_monitoring": (
    "Assess 40-micron notification effectiveness; microplastics monitoring "
    "research."
  ),
  "Traditions to build on (free-hand)": (
    "Newari festival practice minimizing plastics via lapti (leaf plates); "
    "some stores already offer cloth bags instead of polythene — alternatives "
    "expensive without subsidies."
  ),
}


def hfill(hex_):
  return PatternFill("solid", fgColor=hex_)


def hfont(bold=False, size=10, color="000000", italic=False):
  return Font(bold=bold, size=size, color=color, italic=italic)


thin = Side(style="thin", color="BFBFBF")


def thin_border():
  return Border(left=thin, right=thin, top=thin, bottom=thin)


WRAP_ALIGN = Alignment(wrap_text=True, vertical="top", horizontal="left")
CENTER_ALIGN = Alignment(wrap_text=True, vertical="center", horizontal="center")


def build_workbook():
  wb = openpyxl.Workbook()

  ws = wb.active
  ws.title = "Interview Coding"

  ws.merge_cells("A1:D1")
  title = ws["A1"]
  title.value = "Plastic Pollution Governance — Interview Coding (Codebook)"
  title.fill = hfill(DARK_BLUE)
  title.font = hfont(bold=True, size=13, color=WHITE)
  title.alignment = CENTER_ALIGN

  ws.merge_cells("A2:D2")
  meta = ws["A2"]
  meta.value = (
    "Interview: NPL_15  |  Country: NEPAL  |  Suna Maya Margen (KTM)  |  "
    "Actor: governmental  |  23 June 2025"
  )
  meta.fill = hfill(MID_BLUE)
  meta.font = hfont(size=9, color=WHITE)
  meta.alignment = CENTER_ALIGN

  headers = ["Variable", "Code", "Codebook definition", "Coding notes"]
  for col, header in enumerate(headers, start=1):
    cell = ws.cell(row=4, column=col, value=header)
    cell.fill = hfill(LIGHT_BLUE)
    cell.font = hfont(bold=True, size=10, color=DARK_BLUE)
    cell.alignment = CENTER_ALIGN
    cell.border = thin_border()

  row = 5
  for variable, definition in CODEBOOK:
    values = [
      variable,
      CODING[variable],
      definition,
      NOTES.get(variable, ""),
    ]
    fill = hfill(LIGHT_GREEN) if row % 2 == 0 else hfill(WHITE)
    for col, value in enumerate(values, start=1):
      cell = ws.cell(row=row, column=col, value=value)
      cell.fill = fill
      cell.font = hfont(size=10)
      cell.alignment = WRAP_ALIGN
      cell.border = thin_border()
    row += 1

  ws.column_dimensions["A"].width = 34
  ws.column_dimensions["B"].width = 28
  ws.column_dimensions["C"].width = 52
  ws.column_dimensions["D"].width = 58
  ws.freeze_panes = "A5"

  wide = wb.create_sheet("Wide Format")
  wide_headers = [variable for variable, _ in CODEBOOK]
  for col, header in enumerate(wide_headers, start=1):
    cell = wide.cell(row=1, column=col, value=header)
    cell.fill = hfill(LIGHT_BLUE)
    cell.font = hfont(bold=True, size=9, color=DARK_BLUE)
    cell.alignment = CENTER_ALIGN
    cell.border = thin_border()
  for col, header in enumerate(wide_headers, start=1):
    cell = wide.cell(row=2, column=col, value=CODING[header])
    cell.fill = hfill(ORANGE)
    cell.alignment = WRAP_ALIGN
    cell.border = thin_border()
    wide.column_dimensions[get_column_letter(col)].width = max(14, min(28, len(header) + 2))

  ref = wb.create_sheet("Codebook Reference")
  ref.merge_cells("A1:B1")
  ref["A1"].value = "Codebook Variable Definitions"
  ref["A1"].fill = hfill(DARK_BLUE)
  ref["A1"].font = hfont(bold=True, size=12, color=WHITE)
  ref["A1"].alignment = CENTER_ALIGN
  for idx, (variable, definition) in enumerate(CODEBOOK, start=3):
    ref.cell(row=idx, column=1, value=variable).font = hfont(bold=True)
    ref.cell(row=idx, column=2, value=definition).alignment = WRAP_ALIGN
  ref.column_dimensions["A"].width = 34
  ref.column_dimensions["B"].width = 70

  stakeholder = wb.create_sheet("Stakeholder Analysis")
  stakeholder.merge_cells("A1:B1")
  stakeholder["A1"].value = "Stakeholder Analysis — NPL_15"
  stakeholder["A1"].fill = hfill(DARK_BLUE)
  stakeholder["A1"].font = hfont(bold=True, size=12, color=WHITE)
  stakeholder["A1"].alignment = CENTER_ALIGN
  for idx, (key, value) in enumerate(STAKEHOLDER.items(), start=3):
    stakeholder.cell(row=idx, column=1, value=key).font = hfont(bold=True)
    cell = stakeholder.cell(row=idx, column=2, value=value)
    cell.alignment = WRAP_ALIGN
  stakeholder.column_dimensions["A"].width = 28
  stakeholder.column_dimensions["B"].width = 90

  sources = wb.create_sheet("Sources")
  sources.merge_cells("A1:B1")
  sources["A1"].value = "Source Documents Used for Verification"
  sources["A1"].fill = hfill(DARK_BLUE)
  sources["A1"].font = hfont(bold=True, size=12, color=WHITE)
  sources["A1"].alignment = CENTER_ALIGN
  for idx, (doc, detail) in enumerate(SOURCES, start=3):
    sources.cell(row=idx, column=1, value=doc).font = hfont(bold=True)
    sources.cell(row=idx, column=2, value=detail).alignment = WRAP_ALIGN
  sources.column_dimensions["A"].width = 42
  sources.column_dimensions["B"].width = 60

  wb.save(OUTPUT_PATH)
  print(f"Saved {OUTPUT_PATH}")


if __name__ == "__main__":
  build_workbook()
