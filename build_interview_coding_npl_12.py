"""Generate interview coding spreadsheet for Nepal interview NPL_12.

Coding follows the Overview All Interviews codebook and is based on the
Ministry of Urban Development (MoUD) interview — transcript and interview
guideline. National policy formulation for solid waste management. Uses
expanded Impacts taxonomy.
"""

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUTPUT_PATH = "/workspace/interview_coding_NPL_12.xlsx"

DARK_BLUE = "1F3864"
MID_BLUE = "2E75B6"
LIGHT_BLUE = "D9E2F3"
LIGHT_GREEN = "E2EFDA"
ORANGE = "FCE4D6"
WHITE = "FFFFFF"

SOURCES = [
  ("Interview transcript", "MoUD — Ministry for Urban Planning / Urban Development"),
  ("Interview guideline", "MoUD structured guideline (Themes A–D, full coding notes)"),
  ("Codebook", "Overview All Interviews — NEW Codebook (expanded Impacts)"),
]

STAKEHOLDER = {
  "Interview ID": "NPL_12",
  "Country": "NEPAL",
  "Interview number": "MoUD national policy series",
  "Interview date": "June 2025 (fieldwork; exact time not recorded)",
  "Interviewee": "MoUD official — solid waste management policy and governance",
  "Affiliation / role": (
    "Ministry of Urban Development (MoUD) — formulating policies, acts, and "
    "regulations for solid waste management under cabinet-delegated authority"
  ),
  "Organization": (
    "Ministry of Urban Development; Department of Urban Development and "
    "Building Construction (DUDBC) training centre"
  ),
  "Stakeholder list mapping": "national gov",
  "Coded actor type (codebook)": "governmental",
  "Actor classification rationale": (
    "Central/national government ministry responsible for policy, act, and "
    "regulation formulation — supports local implementation but does not "
    "directly operate municipal waste collection."
  ),
  "Perspective": (
    "National policy architect — confident in draft SWM Act; emphasizes source "
    "segregation, circular economy, EPR, PPP/FDI, three-tier coordination, and "
    "future separate plastics regulatory framework"
  ),
  "Project context": (
    "Draft Solid Waste Management Act (post-2015 constitution); Banchare Danda "
    "landfill filled in 3–4 years (designed 20); Bagmati campaign ~10 years; "
    "40-micron ban; separate plastic policy planned"
  ),
}

CODING = {
  "Country": "NPL",
  "Country name": "NEPAL",
  "ID": "NPL_12",
  "Actortype": "governmental",
  "Problem_awareness_pop": "medium",
  "Problem_awareness_pol": "low",
  "Problem_severity": "high",
  "Problem_littering": "yes",
  "Problem_consumption": "yes",
  "Problem_recycling": "yes",
  "Problem_waste_mgmt": "yes",
  "Problem_production": "yes",
  "Problem_alternatives": "yes",
  "Problem_waste_segregation": "yes",
  "Problem_import": "no",
  "Impacts": (
    "water, rivers, cities, health, tourism, agriculture, drainage blockage, "
    "wildlife, dumping sites, foul smell, visual pollution"
  ),
  "Governance": "yes",
  "Relevance_international_pol": "yes",
  "Coordination_sectoral": "yes",
  "Coordination_levels": "yes",
  "Unclear_responsibilities": "yes",
  "Actors": "yes",
  "Federal government": "yes",
  "Provincial government": "yes",
  "Local government": "yes",
  "Students": "no",
  "Private_sector": "yes",
  "Civil_society": "no",
  "Science": "yes",
  "Households": "yes",
  "Education institutions": "yes",
  "Private_companies(hotels, shops, etc.)": "yes",
  "Implementation issues": "yes",
  "Monitoring": "yes",
  "Financial_resources": "no",
  "Research": "no",
  "Infrastructure": "no",
  "Enforcement": "yes",
  "Policies in place": "yes",
  "Pol_epr": "yes",
  "Pol_import": "no",
  "Pol_awarness": "yes",
  "Pol_education": "yes",
  "Pol_capacity": "yes",
  "Pol_RD": "yes",
  "Pol_tax": "no",
  "Pol_ban": "yes",
  "Pol_subsitutes": "yes",
  "Pol_clean_up": "yes",
  "Pol_upcycling": "no",
  "Pol_recycling": "yes",
  "Pol_waste_collection": "yes",
  "Solutions": "yes",
  "Sol_lead_agency": "yes",
  "Sol_responsibilities": "yes",
  "Sol_epr": "yes",
  "Sol_awarness": "yes",
  "Sol_segregation": "yes",
  "Sol_upcycling": "no",
  "Sol_recycling": "yes",
  "Sol_education": "yes",
  "Sol_capacity": "yes",
  "Sol_RD": "yes",
  "Sol_tax": "no",
  "Sol_ban": "yes",
  "Sol_finance": "yes",
  "Sol_infrastructure": "yes",
  "Sol_subsitutes": "yes",
  "Sol_clean_up": "no",
  "Sol_enforcement": "yes",
  "Sol_monitoring": "yes",
  "Traditions to build on (free-hand)": (
    "Leaf plates instead of plastic; natural consumption–land cycle broken by "
    "plastic intervention; past organic waste absorbed by nature at low population "
    "density; emerging home bins for decomposable vs non-decomposable waste"
  ),
}

CODEBOOK = [
  ("Country", "Country ISO code"),
  ("Country name", "Country name"),
  ("ID", "InterviewID: ISO_Nr"),
  ("Actortype", "governmental, pol_party, science, cso, igo, private_sector, shop_owner, hotel_owner, household"),
  ("Problem_awareness_pop", "high, medium, low, NA — lack of awareness among population mentioned"),
  ("Problem_awareness_pol", "high, medium, low, NA — lack of awareness among policy makers mentioned"),
  ("Problem_severity", "high, medium, low, NA"),
  ("Problem_littering", "Littering is a problem; was mentioned: yes/no"),
  ("Problem_consumption", "High consumption is a problem; was mentioned: yes/no"),
  ("Problem_recycling", "No or too little recycling is a problem; was mentioned: yes/no"),
  ("Problem_waste_mgmt", "Ill waste management is a problem; was mentioned: yes/no"),
  ("Problem_production", "High production rates; was mentioned: yes/no"),
  ("Problem_alternatives", "No good alternatives; was mentioned: yes/no"),
  ("Problem_waste_segregation", "Lack of segregation was mentioned: yes/no"),
  ("Problem_import", "Plastic imports are a problem; was mentioned: yes/no"),
  ("Impacts", "Impacts mentioned on water, cities, biodiversity, health, etc. or NA"),
  ("Governance", "Mentioned: yes/no"),
  ("Relevance_international_pol", "International treaties are relevant for national level policy"),
  ("Coordination_sectoral", "Lack of coordination across sectors"),
  ("Coordination_levels", "Lack of coordination across levels"),
  ("Unclear_responsibilities", "Unclear responsibilities among the involved actors"),
  ("Actors", "Mentioned: yes/no"),
  ("Federal government", "Mentioned: yes/no"),
  ("Provincial government", "Mentioned: yes/no"),
  ("Local government", "Mentioned: yes/no"),
  ("Students", "Mentioned: yes/no"),
  ("Private_sector", "Mentioned: yes/no"),
  ("Civil_society", "Mentioned: yes/no"),
  ("Science", "Mentioned: yes/no"),
  ("Households", "Mentioned: yes/no"),
  ("Education institutions", "Mentioned: yes/no"),
  ("Private_companies(hotels, shops, etc.)", "Mentioned: yes/no"),
  ("Implementation issues", "Mentioned: yes/no"),
  ("Monitoring", "Lack of monitoring"),
  ("Financial_resources", "Lack of financial resources"),
  ("Research", "Lack of research"),
  ("Infrastructure", "Lack of infrastructure"),
  ("Enforcement", "Lack of enforcement"),
  ("Policies in place", "Mentioned: yes/no"),
  ("Pol_epr", "Extended producer responsibility"),
  ("Pol_import", "Import regulations"),
  ("Pol_awarness", "Awareness campaigns"),
  ("Pol_education", "Education programs"),
  ("Pol_capacity", "Capacity building"),
  ("Pol_RD", "Research & development"),
  ("Pol_tax", "Levies or taxes on specific products (often bags)"),
  ("Pol_ban", "Bans of specific products"),
  ("Pol_subsitutes", "Alternatives like reusable bags, or banana leaf plates"),
  ("Pol_clean_up", "Clean-up campaigns"),
  ("Pol_upcycling", "Making a new product, like blacktop"),
  ("Pol_recycling", "Making a new raw material"),
  ("Pol_waste_collection", "Waste collection"),
  ("Solutions", "Mentioned: yes/no"),
  ("Sol_lead_agency", "Procedural measures to establish lead agency"),
  ("Sol_responsibilities", "Clear responsibilities"),
  ("Sol_epr", "Extended producer responsibility"),
  ("Sol_awarness", "Awareness raising measures"),
  ("Sol_segregation", "Segregation of waste"),
  ("Sol_upcycling", "Upcycling of plastics"),
  ("Sol_recycling", "New raw materials"),
  ("Sol_education", "Education programs at schools"),
  ("Sol_capacity", "Capacity-building programs"),
  ("Sol_RD", "Research programs"),
  ("Sol_tax", "Taxes or levies on products"),
  ("Sol_ban", "Ban of products"),
  ("Sol_finance", "Financial measures"),
  ("Sol_infrastructure", "Infrastructural measures"),
  ("Sol_subsitutes", "Alternatives"),
  ("Sol_clean_up", "Clean-up campaigns"),
  ("Sol_enforcement", "Enforcement procedures"),
  ("Sol_monitoring", "Monitoring procedures"),
  ("Traditions to build on (free-hand)", "Freehand or NA"),
]

NOTES = {
  "Actortype": (
    "MoUD official — solid waste management policy formulation. Stakeholder list: "
    "national gov. Codebook: governmental."
  ),
  "Problem_awareness_pop": (
    "Public closely attached to plastic for convenience; don't consider future "
    "public health implications; no segregation culture. Newer generation more "
    "concerned — separate home bins emerging but not widespread."
  ),
  "Problem_awareness_pol": (
    "MoUD actively drafting SWM Act with stakeholder consultation; confident in "
    "implementable approaches; separate plastic policy planned — policymakers "
    "highly engaged (low lack of pol awareness)."
  ),
  "Problem_severity": (
    "Plastic disturbs society most; remarkable volume in waste composition; "
    "major nuisance to old cities and emerging towns; Bagmati campaign 10 years "
    "but plastic still majority of problem."
  ),
  "Problem_littering": (
    "Plastic visible everywhere in urban areas; waste thrown haphazardly without "
    "control; composite waste thrown in streets."
  ),
  "Problem_consumption": (
    "People's lives closely attached to plastic for generations; convenient for "
    "carrying goods."
  ),
  "Problem_recycling": (
    "Circular economy promoted but blocked without source segregation; recycling "
    "in future plastic-specific framework."
  ),
  "Problem_waste_mgmt": (
    "SWM in totality challenged; sanitary landfills operated as dumping sites; "
    "Banchare Danda filled in 3–4 years (designed 20)."
  ),
  "Problem_production": (
    "Plastic has remarkable volume in overall waste composition; policies focus "
    "more on plastic because of significant share."
  ),
  "Problem_alternatives": (
    "No specific plastic-waste policy yet; separate framework and community "
    "guidelines planned; leaf plates as past alternative."
  ),
  "Problem_waste_segregation": (
    "Main challenge — no culture of source segregation; heterogeneous waste "
    "streams mixed together makes municipal management very difficult."
  ),
  "Impacts": (
    "Water/rivers: obstructs rivers and canals, blocks water lines. Cities: "
    "visible everywhere in urban areas, old and emerging towns. Health: "
    "low-quality black plastic for meat/yogurt without hazard awareness. "
    "Tourism: aesthetic pollution negatively impacts tourism. Agriculture: "
    "productivity decreasing. Drainage blockage: sewer and water line "
    "blockages. Wildlife: livestock consume plastic. Dumping sites: landfills "
    "degrade to open dumps — leachate, bad smells. Foul smell: at dumping "
    "sites. Visual pollution: unhygienic open dumping feared by communities."
  ),
  "Relevance_international_pol": (
    "SDG framework and human rights integrated; waste management as "
    "environmental justice; Human Rights Commission reporting requirements."
  ),
  "Coordination_sectoral": (
    "Wide stakeholder consultation — universities, private sector, mayors, "
    "researchers; producer vs public responsibility tension."
  ),
  "Coordination_levels": (
    "Three-tier roles being defined in new Act; provincial/central to assist "
    "LG land acquisition; inter-municipality waste conflicts."
  ),
  "Unclear_responsibilities": (
    "Municipalities confused on private-sector onboarding and donor support; "
    "local authorities implement but follow central framework."
  ),
  "Federal government": (
    "MoUD develops policies, acts, regulations; coordinates international "
    "donors; 40-micron ban; draft SWM Act with EPR."
  ),
  "Provincial government": (
    "Will assist local governments unable to acquire landfill land under new Act."
  ),
  "Local government": (
    "Constitutional responsibility under Local Government Operation Act 2017; "
    "must manage collection, landfills, implementation; own regulations within "
    "central framework."
  ),
  "Private_sector": (
    "PPP provisions; FDI and foreign technology welcome; sample municipal–"
    "private agreement documents prepared."
  ),
  "Science": (
    "Universities consulted in act drafting; welcome international expertise "
    "for circular economy."
  ),
  "Households": (
    "Source segregation habits must develop; separate decomposable/non-"
    "decomposable bins at home emerging."
  ),
  "Education institutions": (
    "Universities among stakeholders consulted for act formulation."
  ),
  "Private_companies(hotels, shops, etc.)": (
    "Manufacturing industries and producers under EPR; domestic and "
    "international private waste operators."
  ),
  "Monitoring": (
    "Insufficient monitoring and technology for segregation patterns; planned "
    "municipal software tracking population, production, segregation, disposal."
  ),
  "Infrastructure": (
    "Transfer stations and segregation facilities exist but sanitary landfills "
    "misoperated — operational failure not absence of facilities."
  ),
  "Enforcement": (
    "Polluter's pay principle weak in implementation; burning prohibition with "
    "penalties proposed in draft act."
  ),
  "Policies in place": (
    "General SWM policy (not plastic-specific); 40-micron ban; Bagmati campaign; "
    "draft SWM Act with EPR, PPP, burning ban."
  ),
  "Pol_epr": (
    "EPR proposed in draft act — producers responsible for disposal; fee to "
    "local body if unable to manage."
  ),
  "Pol_awarness": "Bagmati Cleaning Campaign (~10 years); public informed on segregation.",
  "Pol_education": (
    "Province/district orientation and dissemination planned; DUDBC training "
    "centre for local government capacity."
  ),
  "Pol_capacity": (
    "DUDBC training centre; guidelines post-act; enhance local government "
    "capacity on landfill distance management (5km/10km)."
  ),
  "Pol_RD": (
    "Circular economy; foreign technology investment; separate plastic policy "
    "covering recycling and manufacturing."
  ),
  "Pol_ban": (
    "Plastic thinner than 40 microns banned; strict burning prohibition with "
    "penalties in draft act."
  ),
  "Pol_subsitutes": "Traditional leaf plates cited as pre-plastic alternative.",
  "Pol_clean_up": "Bagmati Cleaning Campaign ongoing ~10 years.",
  "Pol_recycling": "Segregation and recycling in draft act; future plastic-specific framework.",
  "Pol_waste_collection": (
    "Act covers collection, transportation, and disposal; private sector entry "
    "via PPP."
  ),
  "Sol_lead_agency": (
    "MoUD coordinating; central government donor coordination; three-tier "
    "role clarification."
  ),
  "Sol_responsibilities": (
    "Define central, provincial, local roles; EPR for producers; regional "
    "sanitary landfill sites covering multiple municipalities."
  ),
  "Sol_epr": (
    "Extended producer responsibility in draft act; industry cooperation after "
    "enactment; software data to inform producer discussions."
  ),
  "Sol_awarness": (
    "Province/district consultations and orientation — not just website posting; "
    "public source-segregation awareness scaling beyond Kathmandu."
  ),
  "Sol_segregation": (
    "Source segregation as primary solution — enables circular economy, reduces "
    "volume, generates income."
  ),
  "Sol_recycling": "Circular economy as core approach; future plastic recycling framework.",
  "Sol_education": (
    "Training/orientation packages via DUDBC; municipal official engagement at "
    "provincial and district levels."
  ),
  "Sol_capacity": (
    "Technical backing for elected municipal leaders; regulations and guidelines "
    "post-act; worker health and safety protections."
  ),
  "Sol_RD": (
    "Separate plastic waste policy and community guidelines planned; welcome "
    "international expertise."
  ),
  "Sol_ban": "Strict burning prohibition with penalties in draft act.",
  "Sol_finance": (
    "PPP model; FDI; EPR fees; source segregation linked to income generation "
    "and circular economy sustainability."
  ),
  "Sol_infrastructure": (
    "Transfer stations; proper sanitary landfill operation; provincial/central "
    "land-acquisition support; foreign technology."
  ),
  "Sol_subsitutes": (
    "Revive leaf plates and natural cycle; reduce waste through segregation."
  ),
  "Sol_enforcement": (
    "Burning prohibition with penalties; EPR fee mechanism; stronger polluter's "
    "pay implementation."
  ),
  "Sol_monitoring": (
    "Municipal software tracking population, production rate, segregation "
    "activities, transport, disposal, and landfill status."
  ),
  "Traditions to build on (free-hand)": (
    "Leaf plates instead of plastic; natural consumption–land cycle (pet-to-khet "
    "concept); past organic waste practices when population was low; emerging "
    "home dustbin segregation culture."
  ),
}


def hfill(hex_):
  return PatternFill("solid", fgColor=hex_)


def hfont(bold=False, size=10, color="000000", italic=False):
  return Font(bold=bold, size=size, color=color, italic=italic)


thin = Side(style="thin", color="BFBFBF")


def thin_border():
  return Border(left=thin, right=thin, top=thin, bottom=thin)


WRAP_ALIGN = Alignment(wrap_text=True, vertical="top", horizontal="left")
CENTER_ALIGN = Alignment(wrap_text=True, vertical="center", horizontal="center")


def build_workbook():
  wb = openpyxl.Workbook()

  ws = wb.active
  ws.title = "Interview Coding"

  ws.merge_cells("A1:D1")
  title = ws["A1"]
  title.value = "Plastic Pollution Governance — Interview Coding (Codebook)"
  title.fill = hfill(DARK_BLUE)
  title.font = hfont(bold=True, size=13, color=WHITE)
  title.alignment = CENTER_ALIGN

  ws.merge_cells("A2:D2")
  meta = ws["A2"]
  meta.value = (
    "Interview: NPL_12  |  Country: NEPAL  |  MoUD  |  "
    "Actor: governmental  |  June 2025"
  )
  meta.fill = hfill(MID_BLUE)
  meta.font = hfont(size=9, color=WHITE)
  meta.alignment = CENTER_ALIGN

  headers = ["Variable", "Code", "Codebook definition", "Coding notes"]
  for col, header in enumerate(headers, start=1):
    cell = ws.cell(row=4, column=col, value=header)
    cell.fill = hfill(LIGHT_BLUE)
    cell.font = hfont(bold=True, size=10, color=DARK_BLUE)
    cell.alignment = CENTER_ALIGN
    cell.border = thin_border()

  row = 5
  for variable, definition in CODEBOOK:
    values = [
      variable,
      CODING[variable],
      definition,
      NOTES.get(variable, ""),
    ]
    fill = hfill(LIGHT_GREEN) if row % 2 == 0 else hfill(WHITE)
    for col, value in enumerate(values, start=1):
      cell = ws.cell(row=row, column=col, value=value)
      cell.fill = fill
      cell.font = hfont(size=10)
      cell.alignment = WRAP_ALIGN
      cell.border = thin_border()
    row += 1

  ws.column_dimensions["A"].width = 34
  ws.column_dimensions["B"].width = 28
  ws.column_dimensions["C"].width = 52
  ws.column_dimensions["D"].width = 58
  ws.freeze_panes = "A5"

  wide = wb.create_sheet("Wide Format")
  wide_headers = [variable for variable, _ in CODEBOOK]
  for col, header in enumerate(wide_headers, start=1):
    cell = wide.cell(row=1, column=col, value=header)
    cell.fill = hfill(LIGHT_BLUE)
    cell.font = hfont(bold=True, size=9, color=DARK_BLUE)
    cell.alignment = CENTER_ALIGN
    cell.border = thin_border()
  for col, header in enumerate(wide_headers, start=1):
    cell = wide.cell(row=2, column=col, value=CODING[header])
    cell.fill = hfill(ORANGE)
    cell.alignment = WRAP_ALIGN
    cell.border = thin_border()
    wide.column_dimensions[get_column_letter(col)].width = max(14, min(28, len(header) + 2))

  ref = wb.create_sheet("Codebook Reference")
  ref.merge_cells("A1:B1")
  ref["A1"].value = "Codebook Variable Definitions"
  ref["A1"].fill = hfill(DARK_BLUE)
  ref["A1"].font = hfont(bold=True, size=12, color=WHITE)
  ref["A1"].alignment = CENTER_ALIGN
  for idx, (variable, definition) in enumerate(CODEBOOK, start=3):
    ref.cell(row=idx, column=1, value=variable).font = hfont(bold=True)
    ref.cell(row=idx, column=2, value=definition).alignment = WRAP_ALIGN
  ref.column_dimensions["A"].width = 34
  ref.column_dimensions["B"].width = 70

  stakeholder = wb.create_sheet("Stakeholder Analysis")
  stakeholder.merge_cells("A1:B1")
  stakeholder["A1"].value = "Stakeholder Analysis — NPL_12"
  stakeholder["A1"].fill = hfill(DARK_BLUE)
  stakeholder["A1"].font = hfont(bold=True, size=12, color=WHITE)
  stakeholder["A1"].alignment = CENTER_ALIGN
  for idx, (key, value) in enumerate(STAKEHOLDER.items(), start=3):
    stakeholder.cell(row=idx, column=1, value=key).font = hfont(bold=True)
    cell = stakeholder.cell(row=idx, column=2, value=value)
    cell.alignment = WRAP_ALIGN
  stakeholder.column_dimensions["A"].width = 28
  stakeholder.column_dimensions["B"].width = 90

  sources = wb.create_sheet("Sources")
  sources.merge_cells("A1:B1")
  sources["A1"].value = "Source Documents Used for Verification"
  sources["A1"].fill = hfill(DARK_BLUE)
  sources["A1"].font = hfont(bold=True, size=12, color=WHITE)
  sources["A1"].alignment = CENTER_ALIGN
  for idx, (doc, detail) in enumerate(SOURCES, start=3):
    sources.cell(row=idx, column=1, value=doc).font = hfont(bold=True)
    sources.cell(row=idx, column=2, value=detail).alignment = WRAP_ALIGN
  sources.column_dimensions["A"].width = 42
  sources.column_dimensions["B"].width = 60

  wb.save(OUTPUT_PATH)
  print(f"Saved {OUTPUT_PATH}")


if __name__ == "__main__":
  build_workbook()
