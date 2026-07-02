"""Generate interview coding spreadsheet for Nepal interview NPL_20.

Coding follows the Overview All Interviews codebook and is based on the
municipal environment officer interview (Chitwan-area small municipality;
Safa Urja contractor) — transcript and interview guideline. Uses expanded
Impacts taxonomy.
"""

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUTPUT_PATH = "/workspace/interview_coding_NPL_20.xlsx"

DARK_BLUE = "1F3864"
MID_BLUE = "2E75B6"
LIGHT_BLUE = "D9E2F3"
LIGHT_GREEN = "E2EFDA"
ORANGE = "FCE4D6"
WHITE = "FFFFFF"

SOURCES = [
  ("Interview transcript", "Municipal environment officer — Marlene Kammerer / Deep / Ramdevi, June 2025"),
  ("Interview guideline", "Plastic Pollution Governance in Nepal — municipal officer structured guideline"),
  ("Codebook", "Overview All Interviews — NEW Codebook (expanded Impacts)"),
  (
    "Coding note",
    "Respondent name and exact municipality not recorded in transcript; context "
    "indicates Chitwan-area small municipality contracting Safa Urja (cf. NPL_11). "
    "Dhulikhel mayor example cited by interpreter as comparator, not respondent's "
    "own municipality.",
  ),
]

STAKEHOLDER = {
  "Interview ID": "NPL_20",
  "Country": "NEPAL",
  "Interview number": "Chitwan-area municipal fieldwork series",
  "Interview date": "June 2025 (fieldwork; exact date not recorded)",
  "Interviewee": "Municipal environment officer — name not recorded",
  "Affiliation / role": (
    "Environment officer, small municipality (Chitwan area) — oversees waste/plastic "
    "management coordination with contracted private operator"
  ),
  "Organization": (
    "Municipality — Environment & Disaster Risk Management budget line; contracts "
    "Safa Urja for household collection, partial segregation, and recycling; "
    "coordinates with Ratnanagar, Kalika, Khairahani, Rapti on integrated landfill PPP"
  ),
  "Stakeholder list mapping": "governmental",
  "Coded actor type (codebook)": "governmental",
  "Actor classification rationale": (
    "Municipal government environment officer describing local implementation, "
    "budget allocation, enforcement, and contractor oversight — not private sector "
    "or CSO."
  ),
  "Perspective": (
    "Local-government implementer — very concerned; emphasizes visual pollution, "
    "cloth-bag campaigns, weak national enforcement of bans, public-awareness gaps, "
    "municipal land constraints, and need for cross-municipality PPP landfill and "
    "technology (shredders, green roads)"
  ),
  "Project context": (
    "Dustbins at corners; cloth bags distributed; Safa Urja household collection; "
    "no municipal engineered landfill (private sector has land/MRF); SWMA bill "
    "to Council of Ministers; Natural Resource Protection Act penalties; 40-micron "
    "and black-plastic production ban partially effective"
  ),
}

CODING = {
  "Country": "NPL",
  "Country name": "NEPAL",
  "ID": "NPL_20",
  "Actortype": "governmental",
  "Problem_awareness_pop": "medium",
  "Problem_awareness_pol": "medium",
  "Problem_severity": "high",
  "Problem_littering": "yes",
  "Problem_consumption": "yes",
  "Problem_recycling": "yes",
  "Problem_waste_mgmt": "yes",
  "Problem_production": "yes",
  "Problem_alternatives": "yes",
  "Problem_waste_segregation": "yes",
  "Problem_import": "no",
  "Impacts": (
    "visual pollution, tourism, agriculture, health, air pollution, cities"
  ),
  "Governance": "yes",
  "Relevance_international_pol": "no",
  "Coordination_sectoral": "yes",
  "Coordination_levels": "yes",
  "Unclear_responsibilities": "no",
  "Actors": "yes",
  "Federal government": "yes",
  "Provincial government": "no",
  "Local government": "yes",
  "Students": "no",
  "Private_sector": "yes",
  "Civil_society": "yes",
  "Science": "no",
  "Households": "yes",
  "Education institutions": "no",
  "Private_companies(hotels, shops, etc.)": "yes",
  "Implementation issues": "yes",
  "Monitoring": "yes",
  "Financial_resources": "yes",
  "Research": "no",
  "Infrastructure": "yes",
  "Enforcement": "yes",
  "Policies in place": "yes",
  "Pol_epr": "no",
  "Pol_import": "no",
  "Pol_awarness": "yes",
  "Pol_education": "yes",
  "Pol_capacity": "yes",
  "Pol_RD": "no",
  "Pol_tax": "no",
  "Pol_ban": "yes",
  "Pol_subsitutes": "yes",
  "Pol_clean_up": "yes",
  "Pol_upcycling": "yes",
  "Pol_recycling": "yes",
  "Pol_waste_collection": "yes",
  "Solutions": "yes",
  "Sol_lead_agency": "no",
  "Sol_responsibilities": "no",
  "Sol_epr": "no",
  "Sol_awarness": "yes",
  "Sol_segregation": "yes",
  "Sol_upcycling": "yes",
  "Sol_recycling": "yes",
  "Sol_education": "yes",
  "Sol_capacity": "yes",
  "Sol_RD": "no",
  "Sol_tax": "no",
  "Sol_ban": "yes",
  "Sol_finance": "yes",
  "Sol_infrastructure": "yes",
  "Sol_subsitutes": "yes",
  "Sol_clean_up": "no",
  "Sol_enforcement": "yes",
  "Sol_monitoring": "yes",
  "Traditions to build on (free-hand)": "NA",
}

CODEBOOK = [
  ("Country", "Country ISO code"),
  ("Country name", "Country name"),
  ("ID", "InterviewID: ISO_Nr"),
  ("Actortype", "governmental, pol_party, science, cso, igo, private_sector, shop_owner, hotel_owner, household"),
  ("Problem_awareness_pop", "high, medium, low, NA — lack of awareness among population mentioned"),
  ("Problem_awareness_pol", "high, medium, low, NA — lack of awareness among policy makers mentioned"),
  ("Problem_severity", "high, medium, low, NA"),
  ("Problem_littering", "Littering is a problem; was mentioned: yes/no"),
  ("Problem_consumption", "High consumption is a problem; was mentioned: yes/no"),
  ("Problem_recycling", "No or too little recycling is a problem; was mentioned: yes/no"),
  ("Problem_waste_mgmt", "Ill waste management is a problem; was mentioned: yes/no"),
  ("Problem_production", "High production rates; was mentioned: yes/no"),
  ("Problem_alternatives", "No good alternatives; was mentioned: yes/no"),
  ("Problem_waste_segregation", "Lack of segregation was mentioned: yes/no"),
  ("Problem_import", "Plastic imports are a problem; was mentioned: yes/no"),
  ("Impacts", "Impacts mentioned on water, cities, biodiversity, health, etc. or NA"),
  ("Governance", "Mentioned: yes/no"),
  ("Relevance_international_pol", "International treaties are relevant for national level policy"),
  ("Coordination_sectoral", "Lack of coordination across sectors"),
  ("Coordination_levels", "Lack of coordination across levels"),
  ("Unclear_responsibilities", "Unclear responsibilities among the involved actors"),
  ("Actors", "Mentioned: yes/no"),
  ("Federal government", "Mentioned: yes/no"),
  ("Provincial government", "Mentioned: yes/no"),
  ("Local government", "Mentioned: yes/no"),
  ("Students", "Mentioned: yes/no"),
  ("Private_sector", "Mentioned: yes/no"),
  ("Civil_society", "Mentioned: yes/no"),
  ("Science", "Mentioned: yes/no"),
  ("Households", "Mentioned: yes/no"),
  ("Education institutions", "Mentioned: yes/no"),
  ("Private_companies(hotels, shops, etc.)", "Mentioned: yes/no"),
  ("Implementation issues", "Mentioned: yes/no"),
  ("Monitoring", "Lack of monitoring"),
  ("Financial_resources", "Lack of financial resources"),
  ("Research", "Lack of research"),
  ("Infrastructure", "Lack of infrastructure"),
  ("Enforcement", "Lack of enforcement"),
  ("Policies in place", "Mentioned: yes/no"),
  ("Pol_epr", "Extended producer responsibility"),
  ("Pol_import", "Import regulations"),
  ("Pol_awarness", "Awareness campaigns"),
  ("Pol_education", "Education programs"),
  ("Pol_capacity", "Capacity building"),
  ("Pol_RD", "Research & development"),
  ("Pol_tax", "Levies or taxes on specific products (often bags)"),
  ("Pol_ban", "Bans of specific products"),
  ("Pol_subsitutes", "Alternatives like reusable bags, or banana leaf plates"),
  ("Pol_clean_up", "Clean-up campaigns"),
  ("Pol_upcycling", "Making a new product, like blacktop"),
  ("Pol_recycling", "Making a new raw material"),
  ("Pol_waste_collection", "Waste collection"),
  ("Solutions", "Mentioned: yes/no"),
  ("Sol_lead_agency", "Procedural measures to establish lead agency"),
  ("Sol_responsibilities", "Clear responsibilities"),
  ("Sol_epr", "Extended producer responsibility"),
  ("Sol_awarness", "Awareness raising measures"),
  ("Sol_segregation", "Segregation of waste"),
  ("Sol_upcycling", "Upcycling of plastics"),
  ("Sol_recycling", "New raw materials"),
  ("Sol_education", "Education programs at schools"),
  ("Sol_capacity", "Capacity-building programs"),
  ("Sol_RD", "Research programs"),
  ("Sol_tax", "Taxes or levies on products"),
  ("Sol_ban", "Ban of products"),
  ("Sol_finance", "Financial measures"),
  ("Sol_infrastructure", "Infrastructural measures"),
  ("Sol_subsitutes", "Alternatives"),
  ("Sol_clean_up", "Clean-up campaigns"),
  ("Sol_enforcement", "Enforcement procedures"),
  ("Sol_monitoring", "Monitoring procedures"),
  ("Traditions to build on (free-hand)", "Freehand or NA"),
]

NOTES = {
  "Actortype": (
    "Municipal environment officer (Chitwan-area small municipality). Stakeholder "
    "list: governmental. Codebook: governmental."
  ),
  "Problem_awareness_pop": (
    "Mixed: residents use cloth bags and avoid burning (1–2 complaints/year), but "
    "packet littering (gutka, noodles) persists and 'society is not very aware'; "
    "nationally people behave unconsciously without enforcement."
  ),
  "Problem_awareness_pol": (
    "Laws exist but are weak, broad, and flexible; penalties infrequent; 40-micron "
    "ban not 100% effective; federal level struggles with policy clarity and "
    "coordination (guideline)."
  ),
  "Problem_severity": (
    "Very concerned — plastics non-biodegradable, no value/not recycled, small packets "
    "persist in soil and farmlands; black/sub-40-micron plastics still found."
  ),
  "Problem_littering": (
    "Gutka and noodle wrappers discarded despite dustbins; littering visible on "
    "roads and in environment."
  ),
  "Problem_consumption": (
    "Day-to-day plastic use; Kurkure, chips, noodles packets without reuse value."
  ),
  "Problem_recycling": (
    "Some plastics have no value and are not recycled; Safa Urja recycles most but "
    "low-value fractions remain problematic."
  ),
  "Problem_waste_mgmt": (
    "Plastics not properly managed; no municipal landfill; reliance on private "
    "sector; residual dumping at riverbank."
  ),
  "Problem_production": (
    "Below-40-micron and black-plastic production banned but not fully effective — "
    "plastics still found."
  ),
  "Problem_alternatives": (
    "No proper alternatives for banned plastics; pollution continues because "
    "substitutes unavailable."
  ),
  "Problem_waste_segregation": (
    "Safa Urja segregates only partially; no transfer station makes segregation "
    "difficult; municipality lacks segregation infrastructure."
  ),
  "Impacts": (
    "Visual pollution: primary local concern — gutka/noodle wrappers reduce aesthetics "
    "for residents and tourists. Agriculture: packets persist in farmlands. Health/air: "
    "burning releases toxins/global-warming gases (limited locally). Cities: litter "
    "on streets despite dustbins."
  ),
  "Coordination_sectoral": (
    "Municipality contracts Safa Urja but limited linkage ('do not have link to "
    "municipality'); private sector handles most processing while municipality oversees."
  ),
  "Coordination_levels": (
    "Challenges differ by level — local operational/resource gaps vs federal policy "
    "clarity/coordination struggles; neighbouring municipalities coordinating integrated "
    "waste management."
  ),
  "Federal government": (
    "40-micron/black-plastic production ban; Solid Waste Management Act bill to Council "
    "of Ministers; Natural Resource Protection Act penalties."
  ),
  "Local government": (
    "Municipality formulates rules, provides dustbins, Environment & Disaster budget, "
    "cloth-bag campaigns, contractor oversight."
  ),
  "Private_sector": (
    "Safa Urja contracted for household collection, partial segregation, shredding, "
    "and recycling on private land; pays municipality; Arna Beer, pharmaceuticals "
    "in industrial zone."
  ),
  "Civil_society": (
    "NGO/non-governmental sectors mentioned as actors who must be involved alongside "
    "municipality and government."
  ),
  "Households": (
    "Primary collection point via Safa Urja; cloth-bag adoption increasing; inconsistent "
    "compliance with dustbin use."
  ),
  "Private_companies(hotels, shops, etc.)": (
    "Small restaurants, cinema hall, one municipal hospital (autoclave policy not "
    "invested in); simple business owners targeted with households."
  ),
  "Monitoring": (
    "Weak enforcement; penalties for open dumping/burning exist but infrequent; "
    "only 1–2 burning complaints this year."
  ),
  "Financial_resources": (
    "Municipality makes minimal direct investment; Environment & Disaster budget "
    "allocated but private sector bears operational costs; high initial infrastructure "
    "investment needed."
  ),
  "Infrastructure": (
    "No municipal engineered landfill; no transfer station; lack of shredders and "
    "advanced technology; land unavailability for municipality."
  ),
  "Enforcement": (
    "Laws weak and flexible; penalties not strict or frequent; ban below 40 microns "
    "not 100% effective."
  ),
  "Policies in place": (
    "40-micron/black-plastic production ban; Natural Resource Protection Act "
    "penalties; Local Governance Act; SWMA bill pending; municipal littering rules "
    "and dustbins; cloth-bag distribution."
  ),
  "Pol_awarness": (
    "Municipality campaigns against road littering; promote reduce/reuse/recycle and "
    "cloth bags."
  ),
  "Pol_education": (
    "Public informed about laws and environmental friendliness alongside awareness "
    "campaigns."
  ),
  "Pol_capacity": (
    "Local government has capacity to formulate policies and administer rules; "
    "implementation depends on public compliance."
  ),
  "Pol_ban": (
    "National ban on plastic production below 40 microns and black plastics — partially "
    "effective."
  ),
  "Pol_subsitutes": (
    "Cloth bags promoted and distributed (including Bhatbhateni-style model); fiber "
    "bags advocated."
  ),
  "Pol_clean_up": (
    "Campaigns to keep streets cleaner; efforts to prevent plastics scattered on roads."
  ),
  "Pol_upcycling": (
    "Private sector shreds plastics and old clothes into particles for reuse; plastic "
    "can be used to make bags."
  ),
  "Pol_recycling": (
    "Safa Urja recycles most waste at MRF; buys useful plastic particles; green-road "
    "potential discussed."
  ),
  "Pol_waste_collection": (
    "Safa Urja household collection under contract; dustbins at corners; municipality "
    "Environment & Disaster budget."
  ),
  "Sol_awarness": (
    "Public awareness necessary alongside laws — both must go together; citizen "
    "compliance improves with awareness and alternatives."
  ),
  "Sol_segregation": (
    "Transfer stations and better segregation infrastructure; shredders to prevent "
    "plastic spreading."
  ),
  "Sol_upcycling": (
    "Reduce/reuse/recycle — make other items from plastic (e.g. bags); shred for "
    "reuse particles."
  ),
  "Sol_recycling": (
    "Expand recycling technology and green-road applications; private-sector "
    "recycling model scaled."
  ),
  "Sol_education": (
    "Inform citizens about laws and environmentally friendly behaviour."
  ),
  "Sol_capacity": (
    "Cross-municipality PPP integrated waste management (Ratnanagar, Kalika, "
    "Khairahani, Rapti common landfill)."
  ),
  "Sol_ban": (
    "Close below-40-micron production at root level; stricter national guidelines."
  ),
  "Sol_finance": (
    "PPP model for shared landfill; initial investment costly but reduces long-term "
    "costs; government provision of transfer-station facilities."
  ),
  "Sol_infrastructure": (
    "Engineered landfill, transfer stations, shredders, green roads, recycling "
    "facilities — land acquisition critical."
  ),
  "Sol_subsitutes": (
    "Promote cloth/fiber bags; minimize plastic at source."
  ),
  "Sol_enforcement": (
    "Stronger laws and more frequent penalties — if law made stronger, visible "
    "littering would reduce."
  ),
  "Sol_monitoring": (
    "Stricter enforcement and penalty application for open dumping and burning."
  ),
  "Traditions to build on (free-hand)": (
    "NA — solutions emphasize cloth bags, 3R/4R, and technology rather than traditional "
    "Nepali practices; local-traditions follow-up not answered in transcript."
  ),
}


def hfill(hex_):
  return PatternFill("solid", fgColor=hex_)


def hfont(bold=False, size=10, color="000000", italic=False):
  return Font(bold=bold, size=size, color=color, italic=italic)


thin = Side(style="thin", color="BFBFBF")


def thin_border():
  return Border(left=thin, right=thin, top=thin, bottom=thin)


WRAP_ALIGN = Alignment(wrap_text=True, vertical="top", horizontal="left")
CENTER_ALIGN = Alignment(wrap_text=True, vertical="center", horizontal="center")


def build_workbook():
  wb = openpyxl.Workbook()

  ws = wb.active
  ws.title = "Interview Coding"

  ws.merge_cells("A1:D1")
  title = ws["A1"]
  title.value = "Plastic Pollution Governance — Interview Coding (Codebook)"
  title.fill = hfill(DARK_BLUE)
  title.font = hfont(bold=True, size=13, color=WHITE)
  title.alignment = CENTER_ALIGN

  ws.merge_cells("A2:D2")
  meta = ws["A2"]
  meta.value = (
    "Interview: NPL_20  |  Country: NEPAL  |  Municipal environment officer  |  "
    "Actor: governmental  |  June 2025"
  )
  meta.fill = hfill(MID_BLUE)
  meta.font = hfont(size=9, color=WHITE)
  meta.alignment = CENTER_ALIGN

  headers = ["Variable", "Code", "Codebook definition", "Coding notes"]
  for col, header in enumerate(headers, start=1):
    cell = ws.cell(row=4, column=col, value=header)
    cell.fill = hfill(LIGHT_BLUE)
    cell.font = hfont(bold=True, size=10, color=DARK_BLUE)
    cell.alignment = CENTER_ALIGN
    cell.border = thin_border()

  row = 5
  for variable, definition in CODEBOOK:
    values = [
      variable,
      CODING[variable],
      definition,
      NOTES.get(variable, ""),
    ]
    fill = hfill(LIGHT_GREEN) if row % 2 == 0 else hfill(WHITE)
    for col, value in enumerate(values, start=1):
      cell = ws.cell(row=row, column=col, value=value)
      cell.fill = fill
      cell.font = hfont(size=10)
      cell.alignment = WRAP_ALIGN
      cell.border = thin_border()
    row += 1

  ws.column_dimensions["A"].width = 34
  ws.column_dimensions["B"].width = 28
  ws.column_dimensions["C"].width = 52
  ws.column_dimensions["D"].width = 58
  ws.freeze_panes = "A5"

  wide = wb.create_sheet("Wide Format")
  wide_headers = [variable for variable, _ in CODEBOOK]
  for col, header in enumerate(wide_headers, start=1):
    cell = wide.cell(row=1, column=col, value=header)
    cell.fill = hfill(LIGHT_BLUE)
    cell.font = hfont(bold=True, size=9, color=DARK_BLUE)
    cell.alignment = CENTER_ALIGN
    cell.border = thin_border()
  for col, header in enumerate(wide_headers, start=1):
    cell = wide.cell(row=2, column=col, value=CODING[header])
    cell.fill = hfill(ORANGE)
    cell.alignment = WRAP_ALIGN
    cell.border = thin_border()
    wide.column_dimensions[get_column_letter(col)].width = max(14, min(28, len(header) + 2))

  ref = wb.create_sheet("Codebook Reference")
  ref.merge_cells("A1:B1")
  ref["A1"].value = "Codebook Variable Definitions"
  ref["A1"].fill = hfill(DARK_BLUE)
  ref["A1"].font = hfont(bold=True, size=12, color=WHITE)
  ref["A1"].alignment = CENTER_ALIGN
  for idx, (variable, definition) in enumerate(CODEBOOK, start=3):
    ref.cell(row=idx, column=1, value=variable).font = hfont(bold=True)
    ref.cell(row=idx, column=2, value=definition).alignment = WRAP_ALIGN
  ref.column_dimensions["A"].width = 34
  ref.column_dimensions["B"].width = 70

  stakeholder = wb.create_sheet("Stakeholder Analysis")
  stakeholder.merge_cells("A1:B1")
  stakeholder["A1"].value = "Stakeholder Analysis — NPL_20"
  stakeholder["A1"].fill = hfill(DARK_BLUE)
  stakeholder["A1"].font = hfont(bold=True, size=12, color=WHITE)
  stakeholder["A1"].alignment = CENTER_ALIGN
  for idx, (key, value) in enumerate(STAKEHOLDER.items(), start=3):
    stakeholder.cell(row=idx, column=1, value=key).font = hfont(bold=True)
    cell = stakeholder.cell(row=idx, column=2, value=value)
    cell.alignment = WRAP_ALIGN
  stakeholder.column_dimensions["A"].width = 28
  stakeholder.column_dimensions["B"].width = 90

  sources = wb.create_sheet("Sources")
  sources.merge_cells("A1:B1")
  sources["A1"].value = "Source Documents Used for Verification"
  sources["A1"].fill = hfill(DARK_BLUE)
  sources["A1"].font = hfont(bold=True, size=12, color=WHITE)
  sources["A1"].alignment = CENTER_ALIGN
  for idx, (doc, detail) in enumerate(SOURCES, start=3):
    sources.cell(row=idx, column=1, value=doc).font = hfont(bold=True)
    sources.cell(row=idx, column=2, value=detail).alignment = WRAP_ALIGN
  sources.column_dimensions["A"].width = 42
  sources.column_dimensions["B"].width = 60

  wb.save(OUTPUT_PATH)
  print(f"Saved {OUTPUT_PATH}")


if __name__ == "__main__":
  build_workbook()
