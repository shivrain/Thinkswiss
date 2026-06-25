#!/usr/bin/env python3
"""Generate 4P Index Excel for Avfallsforskriften Kap. 7 (packaging and packaging waste)."""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

POLICY = {
    "policy_name": (
        "Avfallsforskriften – Kap. 7: Emballasje og emballasjeavfall "
        "(Waste Regulation – Ch. 7: Packaging and packaging waste)"
    ),
    "policy_url": "https://lovdata.no/dokument/SF/forskrift/2004-06-01-930/KAPITTEL_7",
    "policy_year": 2025,
    "policy_objective": (
        "To reduce environmental impacts of packaging—including plastic packaging—through waste prevention, "
        "design optimisation, extended producer responsibility (EPR), nationwide separate collection, material "
        "recovery with binding recycling targets, essential design requirements (Annex I), plastic carrier-bag "
        "reporting, and approved producer-responsibility organisations (PROs), implementing the EU Packaging "
        "Directive framework."
    ),
    "policy_target": 1,
    "policy_target_text": (
        '"From 2025 at least 47 % of plastic packaging placed on the market shall be materially recovered." '
        "(§7-10(b)); "
        '"From 2030 at least 52 % of plastic packaging … shall be materially recovered." (§7-10(c)); '
        "until 2024: 30 % plastic packaging excluding expanded polystyrene (§7-10(a))"
    ),
    "policy_type": 0.75,
    "policy_type_justification": (
        "Executive regulation (forskrift) under forurensningsloven and produktkontrolloven; implements EU Packaging "
        "Directive 1994/62/EC and amendments (2018/852, 2019/904). Chapter reintroduced August 2017; major "
        "restructuring June 2025 (FOR-2025-06-18-1129, in force 1 July 2025). §7-5a on plastic drink cups "
        "added November 2025."
    ),
    "policy_integration": 1,
    "policy_sectors_list": (
        "packaging, retail, manufacturing, waste management, producers, recycling, food & beverage, consumption"
    ),
    "policy_circularity": 1,
    "policy_lifecycle_phases_list": (
        "production, design, consumption, end-of-life, recycling, environmental leakage"
    ),
    "policy_budget": 1,
    "policy_budget_text": (
        '"Producers shall cover necessary costs of separate collection and subsequent transport and treatment of '
        'waste from packaging they place on the Norwegian market." (§7-8); PROs shall maintain six months\' '
        "financial reserves (§7-17); approval and annual fees to Treasury (§7-24, §7-24a)"
    ),
}

INSTRUMENTS = [
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Design",
        "instrument_description": (
            "§7-1: Purpose—to reduce environmental problems from packaging in use, increase reuse and material "
            "recovery, reduce packaging-waste impacts through minimisation, optimisation, and ensuring collection, "
            "reuse and material recovery of used packaging and packaging waste."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            '"The purpose of this chapter is to reduce the environmental problems packaging causes when it is used, '
            'increase reuse and material recovery and reduce environmental problems from packaging waste." (§7-1)'
        ),
        "comments": "Policy objective framing national packaging EPR and circularity.",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": (
            "§7-2: Scope—regulates packaging placed on the Norwegian market and separate collection, reuse, material "
            "recovery and treatment of packaging waste; also covers single-use plastic drink cups/lids not classified "
            "as packaging."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            '"The provisions regulate unused packaging brought into circulation in the Norwegian market, as well as '
            'separate collection, reuse, material recovery and other treatment of used packaging and packaging waste." '
            "(§7-2); also applies to single-use plastic drink cups (§7-2)"
        ),
        "comments": "Defines EPR scope including plastic packaging streams.",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Design",
        "instrument_description": (
            "§7-4 + Annex I: Essential requirements—packaging may only be marketed if it meets Annex I basic "
            "requirements on minimising volume/weight, reusability/recyclability, hazardous-substance reduction, "
            "and material-recoverable design; presumed met if harmonised standards complied with."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            '"Packaging may only be marketed in the Norwegian market if it meets the basic requirements in Annex I." '
            "(§7-4); Annex I: packaging shall be designed and manufactured for reuse or recovery including material "
            "recovery (Annex I §1)"
        ),
        "comments": "EU essential requirements / design-for-recycling gate; mandatory market access condition (kan kun).",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§7-5: Producer EPR membership—producers shall fulfil §§7-7–7-11 duties by joining or establishing an "
            "approved PRO (§7-13), or for beverage inner packaging via approved Kap. 6 return system; producer "
            "remains liable if PRO/return system fails."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            '"The producer shall fulfil its duties under §§7-7 to 7-11 by establishing or joining an approved '
            'producer-responsibility organisation." (§7-5); beverage inner packaging may use approved return system '
            "under §6-4 (§7-5)"
        ),
        "comments": "Core EPR financing and compliance structure; mandatory (skal).",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§7-5a: EPR for non-packaging plastic drink cups—anyone placing single-use plastic drink cups/lids on "
            "the market must ensure separate collection, prevent landfill/energy recovery of recyclables, cover "
            "collection/treatment costs, and join an approved PRO."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            '"Anyone who commercially places on the market single-use drink cups wholly or partly made of plastic … '
            'which are not packaging, shall: a. ensure separate collection … c. cover necessary costs of separate '
            'collection and subsequent transport and treatment." (§7-5a); duties via approved PRO (§7-5a)'
        ),
        "comments": "Extends plastic EPR to SUP drink cups; added November 2025.",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Design",
        "instrument_description": (
            "§7-6–§7-7: Waste prevention—producers shall work on waste prevention; annual report on prevention "
            "efforts and compliance with Annex I manufacturing/composition requirements, including packaging-volume "
            "trends."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            '"The producer shall work on waste prevention." (§7-6); '
            '"The producer shall … prepare an annual report on producers\' efforts and results on waste prevention" '
            "and compliance with Annex I requirements (§7-7)"
        ),
        "comments": "Upstream prevention and design reporting; Miljødirektoratet kan fastsette retningslinjer (enabling).",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§7-8: Cost coverage—producers shall cover necessary costs of separate collection and subsequent "
            "transport and treatment of waste from packaging they place on the market, net of recovery revenues."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            '"The producer shall cover necessary costs of separate collection and subsequent transport and treatment '
            'of waste from the packaging it places on the Norwegian market." (§7-8)'
        ),
        "comments": "Binding producer financing for collection and treatment infrastructure.",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": (
            "§7-9: Nationwide separate collection—producers must ensure continuous nationwide separate collection "
            "from municipalities and businesses, sufficient volumes for §7-10 recovery, accept sorted waste from "
            "collectors, and ensure material-recovery fractions are not landfilled or energy-recovered without "
            "justification."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            '"The producer shall ensure: a. separate collection of packaging waste from business and municipalities … '
            'collection shall take place continuously and nationwide … c. that packaging waste sorted for material '
            'recovery is not energy-recovered or landfilled" (§7-9)'
        ),
        "comments": "Operative collection system; Miljødirektoratet kan pålegge innsamling.",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": (
            "§7-10: Material recovery targets—producers must ensure material recovery of packaging waste at minimum "
            "rates by material: from 2025 ≥47 % plastic packaging, ≥60 % carton, ≥80 % paper, ≥70 % ferrous metal, "
            "≥50 % aluminium, ≥70 % glass, ≥25 % wood; from 2030 ≥52 % plastic."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            '"The producer shall ensure material recovery of packaging waste." (§7-10); '
            '"From 2025 … at least 47 % plastic packaging … at least 60 % packaging carton …" (§7-10(b)); '
            '"From 2030 … at least 52 % plastic packaging" (§7-10(c))'
        ),
        "comments": "Binding quantitative recycling targets; key plastics-relevant instrument.",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": (
            "§7-11 + Annex II: Recovery calculation—material recovery rate = recovered weight / packaging placed on "
            "market per calendar year, per material; measurement points defined for plastic (sorted polymer before "
            "pelletising/extrusion) and other materials in Annex II."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            '"Material recovery rate is calculated as the ratio between materially recovered weight of packaging '
            'waste and weight of packaging placed on the market in a calendar year." (§7-11); Annex II: plastic '
            "measurement point before pelletising/extrusion/moulding (Annex II A)"
        ),
        "comments": "Methodological rules for verifying §7-10 compliance.",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§7-13: PRO approval—producer-responsibility organisations must be approved by Miljødirektoratet, "
            "demonstrate ability to meet §§7-7–7-11 duties, register as legal entity; plastic/carton PROs also cover "
            "plastic drink cups; biennial compliance documentation in first two years."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            '"A producer-responsibility organisation shall be approved by the Norwegian Environment Agency." (§7-13); '
            '"must demonstrate … ability to fulfil requirements in §§7-7 to 7-11" (§7-13)'
        ),
        "comments": "Governance gatekeeper for PRO operators (e.g. Grønt Punkt, Infinitum-linked entities).",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§7-16: Annual PRO reporting—by 1 April report to Miljødirektoratet on placed packaging, collected "
            "waste, recovery rates and treatment routes; for plastic PROs include plastic carrier-bag counts split "
            "by thickness (<15 µm / ≥15 µm) and recycled-plastic share in single-use plastic drink bottles "
            "(produktforskriften §2b-7); third-party verification of rates every two years."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            '"The producer-responsibility organisation shall annually, before 1 April, report to the Norwegian '
            'Environment Agency." (§7-16); for plastic PRO: "number of plastic carrier bags … divided by thickness '
            'over and under 15 micrometres" (§7-16(i)); recycled plastic in drink bottles (§7-16(j))'
        ),
        "comments": "Plastic carrier-bag and recycled-content reporting; mandatory (skal).",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Consumption",
        "instrument_description": (
            "§7-19–§7-20: Information duties—PROs shall inform consumers and businesses on packaging-waste "
            "handling (≥1 campaign/year per packaging type) and publish owners, members, fees, and material-recovery "
            "performance against §7-10."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            '"The producer-responsibility organisation shall ensure consumers and business actors receive '
            'information on handling of packaging waste … at least one information campaign each year." (§7-19); '
            '"achievement of material recovery requirements under §7-10 for their members" shall be publicly '
            "available (§7-20(d))"
        ),
        "comments": "Transparency and behaviour-change instruments.",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§7-14–§7-15: PRO governance—organisations must be non-profit (no dividends to owners) and provide "
            "open, non-discriminatory membership access to all obliged producers."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            '"The producer-responsibility organisation shall be without economic purpose (non-profit)." (§7-14); '
            '"Everyone with a duty of membership … shall have access to participate." (§7-15)'
        ),
        "comments": "Structural PRO requirements ensuring fair EPR market.",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": (
            "§7-22: Third-party collector reporting—collectors recovering >1,500 kg/year of packaging waste outside "
            "PRO agreements must report quantities by material to a PRO or Miljødirektoratet (exempting Kap. 6 DRS "
            "beverage packaging)."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            '"Collectors of packaging waste for material recovery shall annually report the quantity of packaging '
            'waste … exceeding 1,500 kg and not collected under agreement with an approved PRO." (§7-22)'
        ),
        "comments": "Prevents free-riding outside PRO system; monitoring instrument.",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§7-25: Administrative fines—Miljødirektoratet may impose infringement fines for breach of PRO "
            "membership duty (§7-5) or plastic drink-cup PRO membership (§7-5a); governed by Chapter 18B."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            '"The Norwegian Environment Agency may, in the event of a breach of §7-5 first paragraph and §7-5a second '
            'paragraph, impose an infringement fine on the person responsible." (§7-25)'
        ),
        "comments": "Enforcement instrument (kan ilegge); added June 2025.",
    },
]

HEADERS = [
    ("A", "policy_name"),
    ("B", "policy_url"),
    ("C", "policy_year"),
    ("D", "policy_objective"),
    ("E", "policy_target"),
    ("F", "policy_target_text"),
    ("G", "policy_type"),
    ("H", "policy_type_justification"),
    ("I", "policy_integration"),
    ("J", "policy_sectors_list"),
    ("K", "policy_circularity"),
    ("L", "policy_lifecycle_phases_list"),
    ("M", "policy_budget"),
    ("N", "policy_budget_text"),
    ("O", "policy_score"),
    ("P", "instrument_type"),
    ("Q", "instrument_lifecycle_stage"),
    ("R", "instrument_description"),
    ("S", "instrument_in_force"),
    ("T", "instrument_implementation"),
    ("U", "instrument_implementation_text"),
    ("V", "instrument_score"),
    ("W", "comments"),
]


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "4P Index - Avfallsforskr Kap7"

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col_idx, (col_letter, header) in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=f"{col_letter} — {header}")
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, inst in enumerate(INSTRUMENTS, start=2):
        policy_score = (
            POLICY["policy_type"]
            + POLICY["policy_integration"]
            + POLICY["policy_circularity"]
            + POLICY["policy_budget"]
            + inst["instrument_type"]
        ) / 5
        instrument_score = (inst["instrument_type"] + inst["instrument_implementation"]) / 2

        row_values = [
            POLICY["policy_name"],
            POLICY["policy_url"],
            POLICY["policy_year"],
            POLICY["policy_objective"],
            POLICY["policy_target"],
            POLICY["policy_target_text"],
            POLICY["policy_type"],
            POLICY["policy_type_justification"],
            POLICY["policy_integration"],
            POLICY["policy_sectors_list"],
            POLICY["policy_circularity"],
            POLICY["policy_lifecycle_phases_list"],
            POLICY["policy_budget"],
            POLICY["policy_budget_text"],
            round(policy_score, 3),
            inst["instrument_type"],
            inst["instrument_lifecycle_stage"],
            inst["instrument_description"],
            inst["instrument_in_force"],
            inst["instrument_implementation"],
            inst["instrument_implementation_text"],
            round(instrument_score, 3),
            inst["comments"],
        ]

        for col_idx, value in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30

    widths = {
        "A": 52, "B": 42, "C": 12, "D": 45, "E": 12, "F": 42, "G": 12, "H": 40,
        "I": 16, "J": 42, "K": 16, "L": 38, "M": 14, "N": 45, "O": 12, "P": 14,
        "Q": 20, "R": 55, "S": 14, "T": 18, "U": 45, "V": 14, "W": 42,
    }
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width

    output_path = "/workspace/4P_Index_Avfallsforskriften_Kap7.xlsx"
    wb.save(output_path)
    print(f"Saved {output_path} ({len(INSTRUMENTS)} instrument rows)")


if __name__ == "__main__":
    main()
