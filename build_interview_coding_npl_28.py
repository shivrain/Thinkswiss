"""Generate interview coding spreadsheet for Nepal interview NPL_28.

Coding follows the Overview All Interviews codebook and is based on the
Sandip (Ratnanagar/Kalika/Khairahani PPP waste-management) interview —
interview guideline. Uses expanded Impacts taxonomy.
"""

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUTPUT_PATH = "/workspace/interview_coding_NPL_28.xlsx"

DARK_BLUE = "1F3864"
MID_BLUE = "2E75B6"
LIGHT_BLUE = "D9E2F3"
LIGHT_GREEN = "E2EFDA"
ORANGE = "FCE4D6"
WHITE = "FFFFFF"

SOURCES = [
  ("Interview guideline", "Sandip — structured guideline (Ratnanagar PPP context)"),
  ("Codebook", "Overview All Interviews — NEW Codebook (expanded Impacts)"),
  (
    "Coding note",
    "Interview date and exact affiliation not recorded in guideline header. "
    "Content strongly overlaps NPL_11 (Safa Urja / Ratnanagar–Kalika–Khairahani "
    "PPP): 21 tractors/day, NPR 66 lakh collection (+10%/yr), littering penalty "
    "NPR 10,000 — likely same waste-management ecosystem; coded as distinct "
  "respondent (Sandip).",
  ),
]

STAKEHOLDER = {
  "Interview ID": "NPL_28",
  "Country": "NEPAL",
  "Interview number": "Ratnanagar / Chitwan PPP waste-management series",
  "Interview date": "Not recorded in guideline",
  "Interviewee": "Sandip (surname/affiliation recorded as 'regiments' in guideline — likely transcription error)",
  "Affiliation / role": (
    "Waste-management practitioner / PPP operator perspective — Ratnanagar, "
    "Kalika, and Khairahani municipalities"
  ),
  "Organization": (
    "PPP solid-waste modality (Suraha cooperation programme); private collectors "
    "under municipal contract; ~21 tractor-loads/day from Ratnanagar alone"
  ),
  "Stakeholder list mapping": "private_sector",
  "Coded actor type (codebook)": "private_sector",
  "Actor classification rationale": (
    "Discusses private-sector waste collection (66 lakh NPR fee, +10%/yr), PPP "
    "application process, enterprise rate schedules, and tractor operations — "
    "private implementer rather than household or municipal officer."
  ),
  "Perspective": (
    "Practitioner-operator — plastic is a real long-term problem; monsoon "
    "worsens collection and agricultural contamination; advocates multi-level "
    "bans, alternatives, reusable/recyclable plastics, landfill upgrades, and "
    "household responsibility despite private collection"
  ),
  "Project context": (
    "School awareness campaigns; Suraha cooperation SWM intervals; Kalika/"
    "Ratnanagar/Khairahani PPP; NPR 10,000 littering penalty (Ratnanagar); "
    "household segregation encouraged; international experts involved; winter "
    "household burning reduces volume"
  ),
}

CODING = {
  "Country": "NPL",
  "Country name": "NEPAL",
  "ID": "NPL_28",
  "Actortype": "private_sector",
  "Problem_awareness_pop": "high",
  "Problem_awareness_pol": "medium",
  "Problem_severity": "high",
  "Problem_littering": "yes",
  "Problem_consumption": "yes",
  "Problem_recycling": "yes",
  "Problem_waste_mgmt": "yes",
  "Problem_production": "no",
  "Problem_alternatives": "yes",
  "Problem_waste_segregation": "yes",
  "Problem_import": "no",
  "Impacts": (
    "agriculture, visual pollution"
  ),
  "Governance": "yes",
  "Relevance_international_pol": "yes",
  "Coordination_sectoral": "no",
  "Coordination_levels": "yes",
  "Unclear_responsibilities": "yes",
  "Actors": "yes",
  "Federal government": "yes",
  "Provincial government": "yes",
  "Local government": "yes",
  "Students": "yes",
  "Private_sector": "yes",
  "Civil_society": "yes",
  "Science": "no",
  "Households": "yes",
  "Education institutions": "yes",
  "Private_companies(hotels, shops, etc.)": "yes",
  "Implementation issues": "yes",
  "Monitoring": "no",
  "Financial_resources": "yes",
  "Research": "no",
  "Infrastructure": "yes",
  "Enforcement": "yes",
  "Policies in place": "yes",
  "Pol_epr": "no",
  "Pol_import": "no",
  "Pol_awarness": "yes",
  "Pol_education": "yes",
  "Pol_capacity": "yes",
  "Pol_RD": "no",
  "Pol_tax": "yes",
  "Pol_ban": "yes",
  "Pol_subsitutes": "yes",
  "Pol_clean_up": "no",
  "Pol_upcycling": "no",
  "Pol_recycling": "yes",
  "Pol_waste_collection": "yes",
  "Solutions": "yes",
  "Sol_lead_agency": "no",
  "Sol_responsibilities": "yes",
  "Sol_epr": "no",
  "Sol_awarness": "yes",
  "Sol_segregation": "yes",
  "Sol_upcycling": "no",
  "Sol_recycling": "yes",
  "Sol_education": "yes",
  "Sol_capacity": "yes",
  "Sol_RD": "no",
  "Sol_tax": "yes",
  "Sol_ban": "yes",
  "Sol_finance": "yes",
  "Sol_infrastructure": "yes",
  "Sol_subsitutes": "yes",
  "Sol_clean_up": "no",
  "Sol_enforcement": "yes",
  "Sol_monitoring": "no",
  "Traditions to build on (free-hand)": (
    "Leaf plates and paper cups used in some places as alternatives to plastic"
  ),
}

CODEBOOK = [
  ("Country", "Country ISO code"),
  ("Country name", "Country name"),
  ("ID", "InterviewID: ISO_Nr"),
  ("Actortype", "governmental, pol_party, science, cso, igo, private_sector, shop_owner, hotel_owner, household"),
  ("Problem_awareness_pop", "high, medium, low, NA — lack of awareness among population mentioned"),
  ("Problem_awareness_pol", "high, medium, low, NA — lack of awareness among policy makers mentioned"),
  ("Problem_severity", "high, medium, low, NA"),
  ("Problem_littering", "Littering is a problem; was mentioned: yes/no"),
  ("Problem_consumption", "High consumption is a problem; was mentioned: yes/no"),
  ("Problem_recycling", "No or too little recycling is a problem; was mentioned: yes/no"),
  ("Problem_waste_mgmt", "Ill waste management is a problem; was mentioned: yes/no"),
  ("Problem_production", "High production rates; was mentioned: yes/no"),
  ("Problem_alternatives", "No good alternatives; was mentioned: yes/no"),
  ("Problem_waste_segregation", "Lack of segregation was mentioned: yes/no"),
  ("Problem_import", "Plastic imports are a problem; was mentioned: yes/no"),
  ("Impacts", "Impacts mentioned on water, cities, biodiversity, health, etc. or NA"),
  ("Governance", "Mentioned: yes/no"),
  ("Relevance_international_pol", "International treaties are relevant for national level policy"),
  ("Coordination_sectoral", "Lack of coordination across sectors"),
  ("Coordination_levels", "Lack of coordination across levels"),
  ("Unclear_responsibilities", "Unclear responsibilities among the involved actors"),
  ("Actors", "Mentioned: yes/no"),
  ("Federal government", "Mentioned: yes/no"),
  ("Provincial government", "Mentioned: yes/no"),
  ("Local government", "Mentioned: yes/no"),
  ("Students", "Mentioned: yes/no"),
  ("Private_sector", "Mentioned: yes/no"),
  ("Civil_society", "Mentioned: yes/no"),
  ("Science", "Mentioned: yes/no"),
  ("Households", "Mentioned: yes/no"),
  ("Education institutions", "Mentioned: yes/no"),
  ("Private_companies(hotels, shops, etc.)", "Mentioned: yes/no"),
  ("Implementation issues", "Mentioned: yes/no"),
  ("Monitoring", "Lack of monitoring"),
  ("Financial_resources", "Lack of financial resources"),
  ("Research", "Lack of research"),
  ("Infrastructure", "Lack of infrastructure"),
  ("Enforcement", "Lack of enforcement"),
  ("Policies in place", "Mentioned: yes/no"),
  ("Pol_epr", "Extended producer responsibility"),
  ("Pol_import", "Import regulations"),
  ("Pol_awarness", "Awareness campaigns"),
  ("Pol_education", "Education programs"),
  ("Pol_capacity", "Capacity building"),
  ("Pol_RD", "Research & development"),
  ("Pol_tax", "Levies or taxes on specific products (often bags)"),
  ("Pol_ban", "Bans of specific products"),
  ("Pol_subsitutes", "Alternatives like reusable bags, or banana leaf plates"),
  ("Pol_clean_up", "Clean-up campaigns"),
  ("Pol_upcycling", "Making a new product, like blacktop"),
  ("Pol_recycling", "Making a new raw material"),
  ("Pol_waste_collection", "Waste collection"),
  ("Solutions", "Mentioned: yes/no"),
  ("Sol_lead_agency", "Procedural measures to establish lead agency"),
  ("Sol_responsibilities", "Clear responsibilities"),
  ("Sol_epr", "Extended producer responsibility"),
  ("Sol_awarness", "Awareness raising measures"),
  ("Sol_segregation", "Segregation of waste"),
  ("Sol_upcycling", "Upcycling of plastics"),
  ("Sol_recycling", "New raw materials"),
  ("Sol_education", "Education programs at schools"),
  ("Sol_capacity", "Capacity-building programs"),
  ("Sol_RD", "Research programs"),
  ("Sol_tax", "Taxes or levies on products"),
  ("Sol_ban", "Ban of products"),
  ("Sol_finance", "Financial measures"),
  ("Sol_infrastructure", "Infrastructural measures"),
  ("Sol_subsitutes", "Alternatives"),
  ("Sol_clean_up", "Clean-up campaigns"),
  ("Sol_enforcement", "Enforcement procedures"),
  ("Sol_monitoring", "Monitoring procedures"),
  ("Traditions to build on (free-hand)", "Freehand or NA"),
]

NOTES = {
  "Actortype": (
    "Sandip — PPP/private waste-collection operator context (Ratnanagar/Kalika/"
    "Khairahani). Stakeholder list: private_sector. Codebook: private_sector."
  ),
  "Problem_awareness_pop": (
    "Lack of awareness on environmental consequences; haphazard plastic use; "
    "people careless because private sector manages waste for them."
  ),
  "Problem_awareness_pol": (
    "Long-term solutions needed — not achievable at once; new ban programmes "
    "underway but household segregation still not practised."
  ),
  "Problem_severity": (
    "Real problem in Nepal — concerned; long-term issue; maximum plastic use "
    "during four-month monsoon."
  ),
  "Problem_littering": (
    "Plastic bag littering banned in Ratnanagar with NPR 10,000 penalty."
  ),
  "Problem_consumption": (
    "Haphazard use; plastic use maximum during monsoon (4 months); minimum use "
    "among people needed."
  ),
  "Problem_recycling": (
    "Need longer-lasting, reusable, and recyclable plastics; thick plastics "
    "priced higher."
  ),
  "Problem_waste_mgmt": (
    "Not proper management; monsoon makes waste collection difficult; less "
    "landfill accumulation due to operations but systemic gaps remain."
  ),
  "Problem_alternatives": (
    "Alternatives to plastics needed at municipal, provincial, and central "
    "policy levels."
  ),
  "Problem_waste_segregation": (
    "Household segregation encouraged but not in practice; segregated waste "
    "collected when provided."
  ),
  "Impacts": (
    "Agriculture: during monsoon plastic reaches agricultural land and reduces "
    "productivity. Visual pollution: improper management visible; local "
    "awareness programmes address long-term problem."
  ),
  "Relevance_international_pol": (
    "International experts involved in waste-management programmes."
  ),
  "Coordination_levels": (
    "Policy on banning plastics and alternatives needed at municipality, "
    "provincial, and central levels — multi-level coordination gap."
  ),
  "Unclear_responsibilities": (
    "Municipality must fulfill responsibilities; households also responsible — "
    "shared duties blurred when private sector handles collection."
  ),
  "Federal government": (
    "Central-level policy on banning plastics and alternatives advocated."
  ),
  "Provincial government": (
    "Provincial-level plastic-ban and alternative policy needed."
  ),
  "Local government": (
    "Ratnanagar, Kalika, Khairahani municipalities; PPP waste modality; "
    "littering penalty; municipality must fulfill responsibilities."
  ),
  "Students": (
    "Schools responsible alongside households and shopkeepers; regular school "
    "awareness campaigns."
  ),
  "Private_sector": (
    "Private collectors under PPP; fee rates fixed per enterprises; people rely "
    "on private management and become careless."
  ),
  "Civil_society": (
    "Suraha cooperation solid-waste management programme at regular intervals."
  ),
  "Households": (
    "Each household member responsible; winter burning reduces waste volume; "
    "segregation not practised."
  ),
  "Education institutions": (
    "Regular awareness campaigns at school level; positive local attitude from "
    "programmes."
  ),
  "Private_companies(hotels, shops, etc.)": (
    "Shopkeepers named as responsible actors; enterprise fee schedules for "
    "private collectors."
  ),
  "Financial_resources": (
    "Waste collection fee NPR 66 lakh with 10% annual increase; PPP application "
    "process; rates fixed per enterprise type."
  ),
  "Infrastructure": (
    "Improved landfill sites needed; 21 tractors/day from Ratnanagar alone; "
    "less waste piled at landfill due to current operations."
  ),
  "Enforcement": (
    "NPR 10,000 penalty for plastic-bag littering in Ratnanagar; compliance "
    "undermined by reliance on private collection."
  ),
  "Policies in place": (
    "New programmes banning plastic use; minimize plastic bags; school awareness "
    "campaigns; Suraha cooperation SWM; Kalika/Ratnanagar/Khairahani PPP; "
    "Ratnanagar littering fine; household segregation encouragement; "
    "international expert involvement."
  ),
  "Pol_awarness": (
    "Regular awareness campaigns at schools and local level; positive community "
    "attitude reported."
  ),
  "Pol_education": (
    "School-level awareness campaigns on plastic pollution."
  ),
  "Pol_capacity": (
    "International experts involved in waste-management capacity."
  ),
  "Pol_tax": (
    "NPR 10,000 littering penalty; higher rates for thick plastics; enterprise "
    "fee schedules for collectors (66 lakh collection revenue)."
  ),
  "Pol_ban": (
    "New programmes banning plastic use; Ratnanagar ban on littering plastic "
    "bags; multi-level ban policy advocated."
  ),
  "Pol_subsitutes": (
    "Minimize plastic bags; leaves and paper cups in some places; alternatives "
    "at all government tiers."
  ),
  "Pol_recycling": (
    "Reusable and recyclable plastics promoted; segregated fractions collected."
  ),
  "Pol_waste_collection": (
    "21 tractors/day from Ratnanagar; Suraha cooperation intervals; PPP collection "
    "modality; segregated waste collected."
  ),
  "Sol_responsibilities": (
    "Responsibility among all people; municipality must fulfill duties; "
    "households, schools, shopkeepers accountable."
  ),
  "Sol_awarness": (
    "Continued awareness programmes at local and school level."
  ),
  "Sol_segregation": (
    "Household waste segregation to be practised — encouraged but not yet "
    "routine."
  ),
  "Sol_recycling": (
    "Use longer-lasting, reusable, and recyclable plastics; price thick "
    "plastics higher to reduce use."
  ),
  "Sol_education": (
    "School awareness campaigns scaled as part of long-term solution."
  ),
  "Sol_capacity": (
    "International expert involvement in waste-management programmes."
  ),
  "Sol_tax": (
    "Higher rates for thick plastics to discourage consumption."
  ),
  "Sol_ban": (
    "Municipal, provincial, and central policies banning plastics and mandating "
    "alternatives — not achievable overnight."
  ),
  "Sol_finance": (
    "PPP model via application letter; NPR 66 lakh collection fee rising 10% "
    "annually; enterprise-based collector rates."
  ),
  "Sol_infrastructure": (
    "Improved landfill sites required for long-term management."
  ),
  "Sol_subsitutes": (
    "Leaves and paper cups; policy-mandated alternatives to plastics at all "
    "levels."
  ),
  "Sol_enforcement": (
    "Ratnanagar NPR 10,000 littering penalty model; stronger compliance needed "
    "alongside collection services."
  ),
  "Traditions to build on (free-hand)": (
    "Leaf plates and paper cups still used in some places as plastic substitutes."
  ),
}


def hfill(hex_):
  return PatternFill("solid", fgColor=hex_)


def hfont(bold=False, size=10, color="000000", italic=False):
  return Font(bold=bold, size=size, color=color, italic=italic)


thin = Side(style="thin", color="BFBFBF")


def thin_border():
  return Border(left=thin, right=thin, top=thin, bottom=thin)


WRAP_ALIGN = Alignment(wrap_text=True, vertical="top", horizontal="left")
CENTER_ALIGN = Alignment(wrap_text=True, vertical="center", horizontal="center")


def build_workbook():
  wb = openpyxl.Workbook()

  ws = wb.active
  ws.title = "Interview Coding"

  ws.merge_cells("A1:D1")
  title = ws["A1"]
  title.value = "Plastic Pollution Governance — Interview Coding (Codebook)"
  title.fill = hfill(DARK_BLUE)
  title.font = hfont(bold=True, size=13, color=WHITE)
  title.alignment = CENTER_ALIGN

  ws.merge_cells("A2:D2")
  meta = ws["A2"]
  meta.value = (
    "Interview: NPL_28  |  Country: NEPAL  |  Sandip (PPP waste operator)  |  "
    "Actor: private_sector"
  )
  meta.fill = hfill(MID_BLUE)
  meta.font = hfont(size=9, color=WHITE)
  meta.alignment = CENTER_ALIGN

  headers = ["Variable", "Code", "Codebook definition", "Coding notes"]
  for col, header in enumerate(headers, start=1):
    cell = ws.cell(row=4, column=col, value=header)
    cell.fill = hfill(LIGHT_BLUE)
    cell.font = hfont(bold=True, size=10, color=DARK_BLUE)
    cell.alignment = CENTER_ALIGN
    cell.border = thin_border()

  row = 5
  for variable, definition in CODEBOOK:
    values = [
      variable,
      CODING[variable],
      definition,
      NOTES.get(variable, ""),
    ]
    fill = hfill(LIGHT_GREEN) if row % 2 == 0 else hfill(WHITE)
    for col, value in enumerate(values, start=1):
      cell = ws.cell(row=row, column=col, value=value)
      cell.fill = fill
      cell.font = hfont(size=10)
      cell.alignment = WRAP_ALIGN
      cell.border = thin_border()
    row += 1

  ws.column_dimensions["A"].width = 34
  ws.column_dimensions["B"].width = 28
  ws.column_dimensions["C"].width = 52
  ws.column_dimensions["D"].width = 58
  ws.freeze_panes = "A5"

  wide = wb.create_sheet("Wide Format")
  wide_headers = [variable for variable, _ in CODEBOOK]
  for col, header in enumerate(wide_headers, start=1):
    cell = wide.cell(row=1, column=col, value=header)
    cell.fill = hfill(LIGHT_BLUE)
    cell.font = hfont(bold=True, size=9, color=DARK_BLUE)
    cell.alignment = CENTER_ALIGN
    cell.border = thin_border()
  for col, header in enumerate(wide_headers, start=1):
    cell = wide.cell(row=2, column=col, value=CODING[header])
    cell.fill = hfill(ORANGE)
    cell.alignment = WRAP_ALIGN
    cell.border = thin_border()
    wide.column_dimensions[get_column_letter(col)].width = max(14, min(28, len(header) + 2))

  ref = wb.create_sheet("Codebook Reference")
  ref.merge_cells("A1:B1")
  ref["A1"].value = "Codebook Variable Definitions"
  ref["A1"].fill = hfill(DARK_BLUE)
  ref["A1"].font = hfont(bold=True, size=12, color=WHITE)
  ref["A1"].alignment = CENTER_ALIGN
  for idx, (variable, definition) in enumerate(CODEBOOK, start=3):
    ref.cell(row=idx, column=1, value=variable).font = hfont(bold=True)
    ref.cell(row=idx, column=2, value=definition).alignment = WRAP_ALIGN
  ref.column_dimensions["A"].width = 34
  ref.column_dimensions["B"].width = 70

  stakeholder = wb.create_sheet("Stakeholder Analysis")
  stakeholder.merge_cells("A1:B1")
  stakeholder["A1"].value = "Stakeholder Analysis — NPL_28"
  stakeholder["A1"].fill = hfill(DARK_BLUE)
  stakeholder["A1"].font = hfont(bold=True, size=12, color=WHITE)
  stakeholder["A1"].alignment = CENTER_ALIGN
  for idx, (key, value) in enumerate(STAKEHOLDER.items(), start=3):
    stakeholder.cell(row=idx, column=1, value=key).font = hfont(bold=True)
    cell = stakeholder.cell(row=idx, column=2, value=value)
    cell.alignment = WRAP_ALIGN
  stakeholder.column_dimensions["A"].width = 28
  stakeholder.column_dimensions["B"].width = 90

  sources = wb.create_sheet("Sources")
  sources.merge_cells("A1:B1")
  sources["A1"].value = "Source Documents Used for Verification"
  sources["A1"].fill = hfill(DARK_BLUE)
  sources["A1"].font = hfont(bold=True, size=12, color=WHITE)
  sources["A1"].alignment = CENTER_ALIGN
  for idx, (doc, detail) in enumerate(SOURCES, start=3):
    sources.cell(row=idx, column=1, value=doc).font = hfont(bold=True)
    sources.cell(row=idx, column=2, value=detail).alignment = WRAP_ALIGN
  sources.column_dimensions["A"].width = 42
  sources.column_dimensions["B"].width = 60

  wb.save(OUTPUT_PATH)
  print(f"Saved {OUTPUT_PATH}")


if __name__ == "__main__":
  build_workbook()
