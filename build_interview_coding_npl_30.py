"""Generate interview coding spreadsheet for Nepal interview NPL_30.

Coding follows the Overview All Interviews codebook and is based on the
household woman (Ward 23 Basantpur, Kathmandu) interview of 18 February 2026 —
interview guideline (recorded). Uses expanded Impacts taxonomy.
"""

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUTPUT_PATH = "/workspace/interview_coding_NPL_30.xlsx"

DARK_BLUE = "1F3864"
MID_BLUE = "2E75B6"
LIGHT_BLUE = "D9E2F3"
LIGHT_GREEN = "E2EFDA"
ORANGE = "FCE4D6"
WHITE = "FFFFFF"

SOURCES = [
  ("Interview guideline", "Household woman Ward 23 Basantpur — guideline, 18 Feb 2026"),
  ("Codebook", "Overview All Interviews — NEW Codebook (expanded Impacts)"),
  (
    "Coding note",
    "Consent given with recording. Respondent name not recorded. Also chemist at "
    "private Kathmandu water-quality lab — chemistry perspective on water "
    "contaminants noted. Questions 5, 6, and traditions follow-up not asked.",
  ),
]

STAKEHOLDER = {
  "Interview ID": "NPL_30",
  "Country": "NEPAL",
  "Interview number": "Kathmandu household / Ward 23 fieldwork series",
  "Interview date": "18 February 2026",
  "Interviewee": "Unknown — household woman, Ward 23 Basantpur",
  "Affiliation / role": (
    "Household resident (Ward 23 Basantpur, Kathmandu); chemist at private "
    "water-quality laboratory"
  ),
  "Organization": (
    "Private water-quality lab (Kathmandu); Ward 23 Basantpur household"
  ),
  "Stakeholder list mapping": "household",
  "Coded actor type (codebook)": "household",
  "Actor classification rationale": (
    "Primary perspective is household resident in metropolitan Kathmandu ward — "
    "chemist profession informs water-quality comments but interview centres on "
    "domestic waste behaviour and municipal collection."
  ),
  "Perspective": (
    "Concerned Kathmandu household member — plastic cheap/convenient; neighbours "
    "litter from windows; past dry/wet separate collection discontinued; daily "
    "municipal collection breeds complacency; wants metropolitan-led segregation "
    "restart and national policy"
  ),
  "Project context": (
    "Kathmandu: daily waste collection by municipality; former separate-day dry/wet "
    "collection; local club initiatives unsustainable; water supply iron/nitrate/"
    "ammonium issues; contrast with areas lacking collection where concern higher"
  ),
}

CODING = {
  "Country": "NPL",
  "Country name": "NEPAL",
  "ID": "NPL_30",
  "Actortype": "household",
  "Problem_awareness_pop": "high",
  "Problem_awareness_pol": "high",
  "Problem_severity": "medium",
  "Problem_littering": "yes",
  "Problem_consumption": "yes",
  "Problem_recycling": "no",
  "Problem_waste_mgmt": "yes",
  "Problem_production": "no",
  "Problem_alternatives": "yes",
  "Problem_waste_segregation": "yes",
  "Problem_import": "no",
  "Impacts": (
    "visual pollution, air pollution, rivers, water, health"
  ),
  "Governance": "yes",
  "Relevance_international_pol": "no",
  "Coordination_sectoral": "no",
  "Coordination_levels": "no",
  "Unclear_responsibilities": "yes",
  "Actors": "yes",
  "Federal government": "yes",
  "Provincial government": "no",
  "Local government": "yes",
  "Students": "no",
  "Private_sector": "yes",
  "Civil_society": "yes",
  "Science": "yes",
  "Households": "yes",
  "Education institutions": "no",
  "Private_companies(hotels, shops, etc.)": "yes",
  "Implementation issues": "yes",
  "Monitoring": "no",
  "Financial_resources": "no",
  "Research": "no",
  "Infrastructure": "yes",
  "Enforcement": "no",
  "Policies in place": "yes",
  "Pol_epr": "no",
  "Pol_import": "no",
  "Pol_awarness": "yes",
  "Pol_education": "no",
  "Pol_capacity": "no",
  "Pol_RD": "no",
  "Pol_tax": "no",
  "Pol_ban": "no",
  "Pol_subsitutes": "no",
  "Pol_clean_up": "no",
  "Pol_upcycling": "no",
  "Pol_recycling": "no",
  "Pol_waste_collection": "yes",
  "Solutions": "yes",
  "Sol_lead_agency": "no",
  "Sol_responsibilities": "yes",
  "Sol_epr": "no",
  "Sol_awarness": "yes",
  "Sol_segregation": "yes",
  "Sol_upcycling": "no",
  "Sol_recycling": "no",
  "Sol_education": "yes",
  "Sol_capacity": "no",
  "Sol_RD": "no",
  "Sol_tax": "no",
  "Sol_ban": "no",
  "Sol_finance": "no",
  "Sol_infrastructure": "yes",
  "Sol_subsitutes": "no",
  "Sol_clean_up": "no",
  "Sol_enforcement": "no",
  "Sol_monitoring": "no",
  "Traditions to build on (free-hand)": "NA",
}

CODEBOOK = [
  ("Country", "Country ISO code"),
  ("Country name", "Country name"),
  ("ID", "InterviewID: ISO_Nr"),
  ("Actortype", "governmental, pol_party, science, cso, igo, private_sector, shop_owner, hotel_owner, household"),
  ("Problem_awareness_pop", "high, medium, low, NA — lack of awareness among population mentioned"),
  ("Problem_awareness_pol", "high, medium, low, NA — lack of awareness among policy makers mentioned"),
  ("Problem_severity", "high, medium, low, NA"),
  ("Problem_littering", "Littering is a problem; was mentioned: yes/no"),
  ("Problem_consumption", "High consumption is a problem; was mentioned: yes/no"),
  ("Problem_recycling", "No or too little recycling is a problem; was mentioned: yes/no"),
  ("Problem_waste_mgmt", "Ill waste management is a problem; was mentioned: yes/no"),
  ("Problem_production", "High production rates; was mentioned: yes/no"),
  ("Problem_alternatives", "No good alternatives; was mentioned: yes/no"),
  ("Problem_waste_segregation", "Lack of segregation was mentioned: yes/no"),
  ("Problem_import", "Plastic imports are a problem; was mentioned: yes/no"),
  ("Impacts", "Impacts mentioned on water, cities, biodiversity, health, etc. or NA"),
  ("Governance", "Mentioned: yes/no"),
  ("Relevance_international_pol", "International treaties are relevant for national level policy"),
  ("Coordination_sectoral", "Lack of coordination across sectors"),
  ("Coordination_levels", "Lack of coordination across levels"),
  ("Unclear_responsibilities", "Unclear responsibilities among the involved actors"),
  ("Actors", "Mentioned: yes/no"),
  ("Federal government", "Mentioned: yes/no"),
  ("Provincial government", "Mentioned: yes/no"),
  ("Local government", "Mentioned: yes/no"),
  ("Students", "Mentioned: yes/no"),
  ("Private_sector", "Mentioned: yes/no"),
  ("Civil_society", "Mentioned: yes/no"),
  ("Science", "Mentioned: yes/no"),
  ("Households", "Mentioned: yes/no"),
  ("Education institutions", "Mentioned: yes/no"),
  ("Private_companies(hotels, shops, etc.)", "Mentioned: yes/no"),
  ("Implementation issues", "Mentioned: yes/no"),
  ("Monitoring", "Lack of monitoring"),
  ("Financial_resources", "Lack of financial resources"),
  ("Research", "Lack of research"),
  ("Infrastructure", "Lack of infrastructure"),
  ("Enforcement", "Lack of enforcement"),
  ("Policies in place", "Mentioned: yes/no"),
  ("Pol_epr", "Extended producer responsibility"),
  ("Pol_import", "Import regulations"),
  ("Pol_awarness", "Awareness campaigns"),
  ("Pol_education", "Education programs"),
  ("Pol_capacity", "Capacity building"),
  ("Pol_RD", "Research & development"),
  ("Pol_tax", "Levies or taxes on specific products (often bags)"),
  ("Pol_ban", "Bans of specific products"),
  ("Pol_subsitutes", "Alternatives like reusable bags, or banana leaf plates"),
  ("Pol_clean_up", "Clean-up campaigns"),
  ("Pol_upcycling", "Making a new product, like blacktop"),
  ("Pol_recycling", "Making a new raw material"),
  ("Pol_waste_collection", "Waste collection"),
  ("Solutions", "Mentioned: yes/no"),
  ("Sol_lead_agency", "Procedural measures to establish lead agency"),
  ("Sol_responsibilities", "Clear responsibilities"),
  ("Sol_epr", "Extended producer responsibility"),
  ("Sol_awarness", "Awareness raising measures"),
  ("Sol_segregation", "Segregation of waste"),
  ("Sol_upcycling", "Upcycling of plastics"),
  ("Sol_recycling", "New raw materials"),
  ("Sol_education", "Education programs at schools"),
  ("Sol_capacity", "Capacity-building programs"),
  ("Sol_RD", "Research programs"),
  ("Sol_tax", "Taxes or levies on products"),
  ("Sol_ban", "Ban of products"),
  ("Sol_finance", "Financial measures"),
  ("Sol_infrastructure", "Infrastructural measures"),
  ("Sol_subsitutes", "Alternatives"),
  ("Sol_clean_up", "Clean-up campaigns"),
  ("Sol_enforcement", "Enforcement procedures"),
  ("Sol_monitoring", "Monitoring procedures"),
  ("Traditions to build on (free-hand)", "Freehand or NA"),
]

NOTES = {
  "Actortype": (
    "Household woman, Ward 23 Basantpur (also chemist, private water-quality "
    "lab). Stakeholder list: household. Codebook: household."
  ),
  "Problem_awareness_pop": (
    "Neighbours don't care — use bags and throw waste through windows; rare "
    "discussions; people concerned but see no options."
  ),
  "Problem_awareness_pol": (
    "No idea about national policies and campaigns; expects metropolitan "
    "government to act; need policy from central government."
  ),
  "Problem_severity": (
    "Concerned — plastic non-decomposable and everywhere; cheap and convenient "
    "at every shop."
  ),
  "Problem_littering": (
    "Neighbours throw plastic/waste through windows; visual pollution primary "
    "concern; plastic everywhere."
  ),
  "Problem_consumption": (
    "Plastic cheap and convenient — available at every shop and taken home."
  ),
  "Problem_waste_mgmt": (
    "Kathmandu daily collection makes people not care; former dry/wet separate "
    "collection discontinued; uncertainty whether municipality mixes segregated "
    "household waste."
  ),
  "Problem_alternatives": (
    "People concerned but have no options — affordable substitutes lacking."
  ),
  "Problem_waste_segregation": (
    "Households separate segregable and non-segregable waste but system collapsed; "
    "municipality previously collected dry/wet on separate days — no longer working."
  ),
  "Impacts": (
    "Visual pollution: primary concern. Air pollution: open burning in localities. "
    "Rivers: clogs pipelines. Water/health: chemist notes iron, nitrate, ammonium "
    "in Kathmandu water; municipal supply good where available."
  ),
  "Unclear_responsibilities": (
    "People expect municipalities to act; little initiative from residents; "
    "metropolitan office should be responsible."
  ),
  "Federal government": (
    "Need of policy from the government — national framework absent from "
    "respondent's knowledge."
  ),
  "Local government": (
    "Metropolitan city and municipality responsible; daily Kathmandu collection; "
    "should restart separate household collection."
  ),
  "Private_sector": (
    "Respondent works at private water-quality laboratory — technical water "
    "analysis perspective."
  ),
  "Civil_society": (
    "Local clubs ran initiatives in the past but were not sustainable."
  ),
  "Science": (
    "Chemist testing water quality — iron, nitrate, ammonium in Kathmandu water."
  ),
  "Households": (
    "Household segregation attempted; neighbours litter; complacency where daily "
    "collection exists."
  ),
  "Private_companies(hotels, shops, etc.)": (
    "Plastic obtained cheaply at every shop — retail availability drives uptake."
  ),
  "Implementation issues": (
    "Dry/wet separate-day collection stopped — households may not have separated "
    "properly or municipality couldn't handle it; local club initiatives not "
    "sustained."
  ),
  "Infrastructure": (
    "Metropolitan city should manage waste infrastructure; separate collection "
    "from households needed."
  ),
  "Policies in place": (
    "Former municipal practice: separate dry and wet waste collected on separate "
    "days — discontinued. Daily municipal waste collection currently. Local club "
    "awareness initiatives (past)."
  ),
  "Pol_awarness": (
    "Local clubs undertook initiatives — not sustainable; awareness needed so "
    "people don't throw plastics everywhere."
  ),
  "Pol_waste_collection": (
    "Daily municipal collection in Kathmandu; formerly separate-day dry/wet "
    "collection (defunct)."
  ),
  "Sol_responsibilities": (
    "Metropolitan office responsible for future plastic/waste policy and management."
  ),
  "Sol_awarness": (
    "People need awareness not to throw plastics everywhere; educate community "
    "members as models for others."
  ),
  "Sol_segregation": (
    "Municipality should restart collecting dry and wet waste separately from "
    "households and manage it properly."
  ),
  "Sol_education": (
    "Deep's question: more educated community members could model behaviour for "
    "less aware neighbours."
  ),
  "Sol_infrastructure": (
    "Metropolitan city to manage separate household collection and waste handling."
  ),
  "Traditions to build on (free-hand)": (
    "NA — traditions follow-up not asked."
  ),
}


def hfill(hex_):
  return PatternFill("solid", fgColor=hex_)


def hfont(bold=False, size=10, color="000000", italic=False):
  return Font(bold=bold, size=size, color=color, italic=italic)


thin = Side(style="thin", color="BFBFBF")


def thin_border():
  return Border(left=thin, right=thin, top=thin, bottom=thin)


WRAP_ALIGN = Alignment(wrap_text=True, vertical="top", horizontal="left")
CENTER_ALIGN = Alignment(wrap_text=True, vertical="center", horizontal="center")


def build_workbook():
  wb = openpyxl.Workbook()

  ws = wb.active
  ws.title = "Interview Coding"

  ws.merge_cells("A1:D1")
  title = ws["A1"]
  title.value = "Plastic Pollution Governance — Interview Coding (Codebook)"
  title.fill = hfill(DARK_BLUE)
  title.font = hfont(bold=True, size=13, color=WHITE)
  title.alignment = CENTER_ALIGN

  ws.merge_cells("A2:D2")
  meta = ws["A2"]
  meta.value = (
    "Interview: NPL_30  |  Country: NEPAL  |  Ward 23 Basantpur household  |  "
    "Actor: household  |  18 Feb 2026"
  )
  meta.fill = hfill(MID_BLUE)
  meta.font = hfont(size=9, color=WHITE)
  meta.alignment = CENTER_ALIGN

  headers = ["Variable", "Code", "Codebook definition", "Coding notes"]
  for col, header in enumerate(headers, start=1):
    cell = ws.cell(row=4, column=col, value=header)
    cell.fill = hfill(LIGHT_BLUE)
    cell.font = hfont(bold=True, size=10, color=DARK_BLUE)
    cell.alignment = CENTER_ALIGN
    cell.border = thin_border()

  row = 5
  for variable, definition in CODEBOOK:
    values = [
      variable,
      CODING[variable],
      definition,
      NOTES.get(variable, ""),
    ]
    fill = hfill(LIGHT_GREEN) if row % 2 == 0 else hfill(WHITE)
    for col, value in enumerate(values, start=1):
      cell = ws.cell(row=row, column=col, value=value)
      cell.fill = fill
      cell.font = hfont(size=10)
      cell.alignment = WRAP_ALIGN
      cell.border = thin_border()
    row += 1

  ws.column_dimensions["A"].width = 34
  ws.column_dimensions["B"].width = 28
  ws.column_dimensions["C"].width = 52
  ws.column_dimensions["D"].width = 58
  ws.freeze_panes = "A5"

  wide = wb.create_sheet("Wide Format")
  wide_headers = [variable for variable, _ in CODEBOOK]
  for col, header in enumerate(wide_headers, start=1):
    cell = wide.cell(row=1, column=col, value=header)
    cell.fill = hfill(LIGHT_BLUE)
    cell.font = hfont(bold=True, size=9, color=DARK_BLUE)
    cell.alignment = CENTER_ALIGN
    cell.border = thin_border()
  for col, header in enumerate(wide_headers, start=1):
    cell = wide.cell(row=2, column=col, value=CODING[header])
    cell.fill = hfill(ORANGE)
    cell.alignment = WRAP_ALIGN
    cell.border = thin_border()
    wide.column_dimensions[get_column_letter(col)].width = max(14, min(28, len(header) + 2))

  ref = wb.create_sheet("Codebook Reference")
  ref.merge_cells("A1:B1")
  ref["A1"].value = "Codebook Variable Definitions"
  ref["A1"].fill = hfill(DARK_BLUE)
  ref["A1"].font = hfont(bold=True, size=12, color=WHITE)
  ref["A1"].alignment = CENTER_ALIGN
  for idx, (variable, definition) in enumerate(CODEBOOK, start=3):
    ref.cell(row=idx, column=1, value=variable).font = hfont(bold=True)
    ref.cell(row=idx, column=2, value=definition).alignment = WRAP_ALIGN
  ref.column_dimensions["A"].width = 34
  ref.column_dimensions["B"].width = 70

  stakeholder = wb.create_sheet("Stakeholder Analysis")
  stakeholder.merge_cells("A1:B1")
  stakeholder["A1"].value = "Stakeholder Analysis — NPL_30"
  stakeholder["A1"].fill = hfill(DARK_BLUE)
  stakeholder["A1"].font = hfont(bold=True, size=12, color=WHITE)
  stakeholder["A1"].alignment = CENTER_ALIGN
  for idx, (key, value) in enumerate(STAKEHOLDER.items(), start=3):
    stakeholder.cell(row=idx, column=1, value=key).font = hfont(bold=True)
    cell = stakeholder.cell(row=idx, column=2, value=value)
    cell.alignment = WRAP_ALIGN
  stakeholder.column_dimensions["A"].width = 28
  stakeholder.column_dimensions["B"].width = 90

  sources = wb.create_sheet("Sources")
  sources.merge_cells("A1:B1")
  sources["A1"].value = "Source Documents Used for Verification"
  sources["A1"].fill = hfill(DARK_BLUE)
  sources["A1"].font = hfont(bold=True, size=12, color=WHITE)
  sources["A1"].alignment = CENTER_ALIGN
  for idx, (doc, detail) in enumerate(SOURCES, start=3):
    sources.cell(row=idx, column=1, value=doc).font = hfont(bold=True)
    sources.cell(row=idx, column=2, value=detail).alignment = WRAP_ALIGN
  sources.column_dimensions["A"].width = 42
  sources.column_dimensions["B"].width = 60

  wb.save(OUTPUT_PATH)
  print(f"Saved {OUTPUT_PATH}")


if __name__ == "__main__":
  build_workbook()
