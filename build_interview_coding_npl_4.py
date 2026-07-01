"""Generate interview coding spreadsheet for Nepal interview NPL_4.

Coding follows the Overview All Interviews codebook and is based on the
Ashok Kumar Byanju Shrestha (Mayor, Dhulikhel Municipality) interview of
22 June 2025 — transcript and interview notes.
"""

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUTPUT_PATH = "/workspace/interview_coding_NPL_4.xlsx"

DARK_BLUE = "1F3864"
MID_BLUE = "2E75B6"
LIGHT_BLUE = "D9E2F3"
LIGHT_GREEN = "E2EFDA"
ORANGE = "FCE4D6"
WHITE = "FFFFFF"

SOURCES = [
  ("Interview transcript", "Dhulikhel Mayor — 22 June 2025 (2:55–3:40pm)"),
  ("Interview notes", "Kathmandu interview no. 4 — guideline notes"),
  ("Interview guideline", "Plastic Pollution Governance in Nepal — Municipality"),
  ("Codebook", "Overview All Interviews — NEW Codebook"),
]

STAKEHOLDER = {
  "Interview ID": "NPL_4",
  "Country": "NEPAL",
  "Interview number": "4 (Kathmandu series)",
  "Interview date": "22 June 2025 (2:55pm – 3:40pm)",
  "Interviewee": "Ashok Kumar Byanju Shrestha",
  "Affiliation / role": "Mayor — Dhulikhel Municipality",
  "Stakeholder list mapping": "local gov",
  "Coded actor type (codebook)": "governmental",
  "Actor classification rationale": (
    "Elected head of local government constitutionally responsible for sanitation, "
    "waste management, and environmental protection. Coded governmental (local "
    "government) per codebook; distinct from pol_party (national political "
    "leadership, cf. NPL_2) and municipal officer/civil servant (cf. NPL_3)."
  ),
  "Role in plastic governance": (
    "Municipal policy and budget (NPR 1.7M + 10 lakh for plastics); green-city "
    "initiative; school/community green clubs; LAPA project; Environment and "
    "Disaster Management Unit"
  ),
  "Perspective": (
    "Local-government leader — autonomous under constitution but lacking federal "
    "technical/financial support; advocates inter-municipal and private-sector "
    "partnerships"
  ),
  "Key partners mentioned": (
    "LAPA project, Kathmandu University, Dhulikhel Hospital, NGOs/INGOs, Bill & "
    "Melinda Gates Foundation, UCLG, SPARC, waste-to-energy company (Dharan MoU)"
  ),
}

CODING = {
  "Country": "NPL",
  "Country name": "NEPAL",
  "ID": "NPL_4",
  "Actortype": "governmental",
  "Problem_awareness_pop": "high",
  "Problem_awareness_pol": "medium",
  "Problem_severity": "high",
  "Problem_littering": "yes",
  "Problem_consumption": "yes",
  "Problem_recycling": "yes",
  "Problem_waste_mgmt": "yes",
  "Problem_production": "yes",
  "Problem_alternatives": "yes",
  "Problem_waste_segregation": "yes",
  "Problem_import": "no",
  "Impacts": "water, cities, health, agriculture",
  "Governance": "yes",
  "Relevance_international_pol": "no",
  "Coordination_sectoral": "yes",
  "Coordination_levels": "yes",
  "Unclear_responsibilities": "yes",
  "Actors": "yes",
  "Federal government": "yes",
  "Provincial government": "yes",
  "Local government": "yes",
  "Students": "yes",
  "Private_sector": "yes",
  "Civil_society": "yes",
  "Science": "yes",
  "Households": "yes",
  "Education institutions": "yes",
  "Private_companies(hotels, shops, etc.)": "yes",
  "Implementation issues": "yes",
  "Monitoring": "yes",
  "Financial_resources": "yes",
  "Research": "yes",
  "Infrastructure": "yes",
  "Enforcement": "yes",
  "Policies in place": "yes",
  "Pol_epr": "no",
  "Pol_import": "no",
  "Pol_awarness": "yes",
  "Pol_education": "yes",
  "Pol_capacity": "yes",
  "Pol_RD": "yes",
  "Pol_tax": "no",
  "Pol_ban": "yes",
  "Pol_subsitutes": "yes",
  "Pol_clean_up": "yes",
  "Pol_upcycling": "yes",
  "Pol_recycling": "yes",
  "Pol_waste_collection": "yes",
  "Solutions": "yes",
  "Sol_lead_agency": "yes",
  "Sol_responsibilities": "yes",
  "Sol_epr": "no",
  "Sol_awarness": "yes",
  "Sol_segregation": "yes",
  "Sol_upcycling": "yes",
  "Sol_recycling": "yes",
  "Sol_education": "yes",
  "Sol_capacity": "yes",
  "Sol_RD": "yes",
  "Sol_tax": "no",
  "Sol_ban": "no",
  "Sol_finance": "yes",
  "Sol_infrastructure": "yes",
  "Sol_subsitutes": "yes",
  "Sol_clean_up": "yes",
  "Sol_enforcement": "yes",
  "Sol_monitoring": "yes",
  "Traditions to build on (free-hand)": (
    "Women's groups and handicraft producers repurposing plastic into goods "
    "(small-scale community practice)"
  ),
}

CODEBOOK = [
  ("Country", "Country ISO code"),
  ("Country name", "Country name"),
  ("ID", "InterviewID: ISO_Nr"),
  ("Actortype", "governmental, pol_party, science, cso, igo, private_sector, shop_owner, hotel_owner, household"),
  ("Problem_awareness_pop", "high, medium, low, NA — lack of awareness among population mentioned"),
  ("Problem_awareness_pol", "high, medium, low, NA — lack of awareness among policy makers mentioned"),
  ("Problem_severity", "high, medium, low, NA"),
  ("Problem_littering", "Littering is a problem; was mentioned: yes/no"),
  ("Problem_consumption", "High consumption is a problem; was mentioned: yes/no"),
  ("Problem_recycling", "No or too little recycling is a problem; was mentioned: yes/no"),
  ("Problem_waste_mgmt", "Ill waste management is a problem; was mentioned: yes/no"),
  ("Problem_production", "High production rates; was mentioned: yes/no"),
  ("Problem_alternatives", "No good alternatives; was mentioned: yes/no"),
  ("Problem_waste_segregation", "Lack of segregation was mentioned: yes/no"),
  ("Problem_import", "Plastic imports are a problem; was mentioned: yes/no"),
  ("Impacts", "Impacts mentioned on water, cities, biodiversity, health, etc. or NA"),
  ("Governance", "Mentioned: yes/no"),
  ("Relevance_international_pol", "International treaties are relevant for national level policy"),
  ("Coordination_sectoral", "Lack of coordination across sectors"),
  ("Coordination_levels", "Lack of coordination across levels"),
  ("Unclear_responsibilities", "Unclear responsibilities among the involved actors"),
  ("Actors", "Mentioned: yes/no"),
  ("Federal government", "Mentioned: yes/no"),
  ("Provincial government", "Mentioned: yes/no"),
  ("Local government", "Mentioned: yes/no"),
  ("Students", "Mentioned: yes/no"),
  ("Private_sector", "Mentioned: yes/no"),
  ("Civil_society", "Mentioned: yes/no"),
  ("Science", "Mentioned: yes/no"),
  ("Households", "Mentioned: yes/no"),
  ("Education institutions", "Mentioned: yes/no"),
  ("Private_companies(hotels, shops, etc.)", "Mentioned: yes/no"),
  ("Implementation issues", "Mentioned: yes/no"),
  ("Monitoring", "Lack of monitoring"),
  ("Financial_resources", "Lack of financial resources"),
  ("Research", "Lack of research"),
  ("Infrastructure", "Lack of infrastructure"),
  ("Enforcement", "Lack of enforcement"),
  ("Policies in place", "Mentioned: yes/no"),
  ("Pol_epr", "Extended producer responsibility"),
  ("Pol_import", "Import regulations"),
  ("Pol_awarness", "Awareness campaigns"),
  ("Pol_education", "Education programs"),
  ("Pol_capacity", "Capacity building"),
  ("Pol_RD", "Research & development"),
  ("Pol_tax", "Levies or taxes on specific products (often bags)"),
  ("Pol_ban", "Bans of specific products"),
  ("Pol_subsitutes", "Alternatives like reusable bags, or banana leaf plates"),
  ("Pol_clean_up", "Clean-up campaigns"),
  ("Pol_upcycling", "Making a new product, like blacktop"),
  ("Pol_recycling", "Making a new raw material"),
  ("Pol_waste_collection", "Waste collection"),
  ("Solutions", "Mentioned: yes/no"),
  ("Sol_lead_agency", "Procedural measures to establish lead agency"),
  ("Sol_responsibilities", "Clear responsibilities"),
  ("Sol_epr", "Extended producer responsibility"),
  ("Sol_awarness", "Awareness raising measures"),
  ("Sol_segregation", "Segregation of waste"),
  ("Sol_upcycling", "Upcycling of plastics"),
  ("Sol_recycling", "New raw materials"),
  ("Sol_education", "Education programs at schools"),
  ("Sol_capacity", "Capacity-building programs"),
  ("Sol_RD", "Research programs"),
  ("Sol_tax", "Taxes or levies on products"),
  ("Sol_ban", "Ban of products"),
  ("Sol_finance", "Financial measures"),
  ("Sol_infrastructure", "Infrastructural measures"),
  ("Sol_subsitutes", "Alternatives"),
  ("Sol_clean_up", "Clean-up campaigns"),
  ("Sol_enforcement", "Enforcement procedures"),
  ("Sol_monitoring", "Monitoring procedures"),
  ("Traditions to build on (free-hand)", "Freehand or NA"),
]

NOTES = {
  "Actortype": (
    "Mayor of Dhulikhel Municipality — local government (stakeholder list: local gov). "
    "Founder of Kathmandu University; former deputy mayor (20 years ago)."
  ),
  "Problem_awareness_pop": (
    "Lack of household awareness on reuse, recycling, and segregation cited as key "
    "barrier; all waste goes directly to landfill."
  ),
  "Problem_awareness_pol": (
    "National government concerned but lacks action-oriented programs; ministry has "
    "not monitored Dhulikhel in 20+ years; SWM Commission collapsed."
  ),
  "Problem_severity": (
    "Very concerned; plastic dangerous to environment and human health; municipal "
    "council allocated dedicated plastic-management budget."
  ),
  "Problem_littering": (
    "Plastics on roadsides, rivers, forests; highway travellers (Araniko & BP highways) "
    "throw garbage."
  ),
  "Problem_consumption": (
    "Widespread plastic use internally and externally across households, industries, "
    "and businesses."
  ),
  "Problem_recycling": (
    "No plastic recycling company currently in Dhulikhel; lacks technical machinery; "
    "provincial-supported recycling existed in Hetauda but not locally."
  ),
  "Problem_waste_mgmt": (
    "Despite WHO LDCT/healthy-city status, no solid waste management system in place."
  ),
  "Problem_production": "Industries and businesses use plastic; nearby pipe/water-tank factories.",
  "Problem_alternatives": (
    "Municipality seeking alternatives — collection and creating new goods from "
    "recycled plastics; LAPA technology development."
  ),
  "Problem_waste_segregation": (
    "Segregation is a major issue; all garbage goes to landfill without sorting."
  ),
  "Impacts": (
    "Agriculture fields, rivers, roadsides, forests polluted. Health risks from "
    "plastic. Aquatic pollution act cited (water bodies/habitat)."
  ),
  "Coordination_sectoral": (
    "Need collaboration with KU, hospital, environmental experts, NGOs, private sector; "
    "surrounding municipalities face similar problems."
  ),
  "Coordination_levels": (
    "Cannot work with federal/provincial currently; DCC lacks coordination mandate "
    "beyond border disputes; 293 municipalities need inter-municipal cooperation."
  ),
  "Unclear_responsibilities": (
    "Ministry approves acts but LG implements/monitors; SWM Commission dissolved; "
    "no clear bylaws for how to manage/sell/produce/recycle plastic."
  ),
  "Federal government": (
    "Controls public land; landfill permission needed; should provide equipment, "
    "budget, technical support — currently not providing."
  ),
  "Provincial government": (
    "Prioritizing sector (Hetauda project); provincial budget supported Hetauda "
    "metropolitan recycling collaboration."
  ),
  "Local government": (
    "Constitutionally responsible for sanitation/waste; Dhulikhel + 293 municipalities; "
    "DCC for border coordination; Environment & Disaster Management Unit."
  ),
  "Students": (
    "School-based green clubs; KU student interns; professors collaborated with municipality."
  ),
  "Private_sector": (
    "PPP, contract, and MoU models; waste-to-energy company (CNG/biogas); DPR prepared."
  ),
  "Civil_society": (
    "NGOs, INGOs, social workers, LAPA project, women's handicraft groups, green clubs."
  ),
  "Science": (
    "Kathmandu University, environmental experts, LAPA Indian colleagues developing "
    "plastic-reduction technology, technical survey personnel."
  ),
  "Households": (
    "Household awareness and segregation as first priority; model ward for plastic reduction."
  ),
  "Education institutions": (
    "Kathmandu University (within municipality); school green clubs in all wards planned."
  ),
  "Private_companies(hotels, shops, etc.)": (
    "Tushithani Hotel (5-star, own WWTP); pipe/water-tank factories in nearby "
    "municipalities; former local recycling company now closed."
  ),
  "Monitoring": (
    "Ministry never monitored in 20 years; DCC monitoring role limited to borders; "
    "municipality must self-monitor."
  ),
  "Financial_resources": (
    "Lack of budget and technology; NPR 1.7M + 10 lakh allocated but insufficient "
    "for large projects; need federal funding."
  ),
  "Research": (
    "Lack of technical knowledge for landfill management; LAPA R&D on plastic "
    "reduction technology."
  ),
  "Infrastructure": (
    "No technical machines; landfill not technically managed; no recycler in Dhulikhel; "
    "need waste-to-energy facilities."
  ),
  "Enforcement": (
    "20/40 micron restrictions exist but markets sell all plastic types; aquatic "
    "pollution act and draft SWM act weakly implemented."
  ),
  "Policies in place": (
    "Green City initiative (energy efficiency, carbon neutrality); Environment & "
    "Natural Resource policies; 85 municipal bylaws; micron notifications; LAPA/LAPA budget."
  ),
  "Pol_awarness": (
    "School and community green clubs; household awareness programmes; model ward initiative."
  ),
  "Pol_education": (
    "School green clubs; household segregation education; one model ward (of 12) "
    "for plastic reduction."
  ),
  "Pol_capacity": (
    "Municipal capacity-building through KU collaboration; seeking federal technical "
    "equipment and training support."
  ),
  "Pol_RD": "LAPA project developing plastic-reduction technology with Indian colleagues.",
  "Pol_ban": "National/municipal 20 and 40 micron plastic restrictions (poorly enforced).",
  "Pol_subsitutes": (
    "Exploring alternatives and creating new goods from collected plastics."
  ),
  "Pol_clean_up": (
    "Plastic collection and management initiatives; LAPA project; community-led green clubs."
  ),
  "Pol_upcycling": (
    "Women's groups/handicrafts repurposing plastic; making new goods from recycled materials."
  ),
  "Pol_recycling": (
    "Municipal budget for reduce-recycle-reuse; collaboration with NGOs on recycling."
  ),
  "Pol_waste_collection": (
    "Waste management system with private sector, NGOs, and social organizations."
  ),
  "Sol_lead_agency": (
    "Strengthen District Coordination Committee for inter-municipal waste/environment "
    "collaboration with budget and technical mandate."
  ),
  "Sol_responsibilities": (
    "National bylaws needed for how to manage, sell, produce, and recycle plastic; "
    "clarify municipal price-setting authority over private-sector sales."
  ),
  "Sol_awarness": "Household-level awareness as first step; community green clubs.",
  "Sol_segregation": "Household waste segregation before landfill disposal.",
  "Sol_upcycling": "Expand handicraft/women's group plastic repurposing; new products from waste.",
  "Sol_recycling": "3Rs prioritization; inter-municipal recycling facilities (cf. Hetauda model).",
  "Sol_education": (
    "Two capacity types: community awareness + technical personnel training; "
    "school green clubs."
  ),
  "Sol_capacity": (
    "More manpower and technical assistance; federal equipment and knowledge transfer."
  ),
  "Sol_RD": "LAPA technology for plastic neutralization; DPR for waste management system.",
  "Sol_finance": (
    "Federal budget and equipment; municipal NPR 1.7M + 10 lakh; PPP/contract/MoU "
    "with private sector."
  ),
  "Sol_infrastructure": (
    "Waste-to-energy (CNG/biogas) via private partner; technical landfill management; "
    "shared regional facilities across neighbouring municipalities."
  ),
  "Sol_subsitutes": "Alternative product development from collected plastics.",
  "Sol_clean_up": (
    "Inter-municipal collaboration (2nd, 3rd, 4th municipalities) for plastic and "
    "water-related waste management."
  ),
  "Sol_enforcement": "Proper bylaws and enforcement of micron restrictions at market level.",
  "Sol_monitoring": (
    "Federal ministry should support monitoring capacity; strengthen DCC coordination role."
  ),
  "Traditions to build on (free-hand)": (
    "Women's groups and handicraft producers repurposing plastic into goods — "
    "small-scale but existing community practice."
  ),
}


def hfill(hex_):
  return PatternFill("solid", fgColor=hex_)


def hfont(bold=False, size=10, color="000000", italic=False):
  return Font(bold=bold, size=size, color=color, italic=italic)


thin = Side(style="thin", color="BFBFBF")


def thin_border():
  return Border(left=thin, right=thin, top=thin, bottom=thin)


WRAP_ALIGN = Alignment(wrap_text=True, vertical="top", horizontal="left")
CENTER_ALIGN = Alignment(wrap_text=True, vertical="center", horizontal="center")


def build_workbook():
  wb = openpyxl.Workbook()

  ws = wb.active
  ws.title = "Interview Coding"

  ws.merge_cells("A1:D1")
  title = ws["A1"]
  title.value = "Plastic Pollution Governance — Interview Coding (Codebook)"
  title.fill = hfill(DARK_BLUE)
  title.font = hfont(bold=True, size=13, color=WHITE)
  title.alignment = CENTER_ALIGN

  ws.merge_cells("A2:D2")
  meta = ws["A2"]
  meta.value = (
    "Interview: NPL_4  |  Country: NEPAL  |  Mayor — Dhulikhel Municipality  |  "
    "Actor: governmental (local gov)  |  22 June 2025"
  )
  meta.fill = hfill(MID_BLUE)
  meta.font = hfont(size=9, color=WHITE)
  meta.alignment = CENTER_ALIGN

  headers = ["Variable", "Code", "Codebook definition", "Coding notes"]
  for col, header in enumerate(headers, start=1):
    cell = ws.cell(row=4, column=col, value=header)
    cell.fill = hfill(LIGHT_BLUE)
    cell.font = hfont(bold=True, size=10, color=DARK_BLUE)
    cell.alignment = CENTER_ALIGN
    cell.border = thin_border()

  row = 5
  for variable, definition in CODEBOOK:
    values = [
      variable,
      CODING[variable],
      definition,
      NOTES.get(variable, ""),
    ]
    fill = hfill(LIGHT_GREEN) if row % 2 == 0 else hfill(WHITE)
    for col, value in enumerate(values, start=1):
      cell = ws.cell(row=row, column=col, value=value)
      cell.fill = fill
      cell.font = hfont(size=10)
      cell.alignment = WRAP_ALIGN
      cell.border = thin_border()
    row += 1

  ws.column_dimensions["A"].width = 34
  ws.column_dimensions["B"].width = 28
  ws.column_dimensions["C"].width = 52
  ws.column_dimensions["D"].width = 58
  ws.freeze_panes = "A5"

  wide = wb.create_sheet("Wide Format")
  wide_headers = [variable for variable, _ in CODEBOOK]
  for col, header in enumerate(wide_headers, start=1):
    cell = wide.cell(row=1, column=col, value=header)
    cell.fill = hfill(LIGHT_BLUE)
    cell.font = hfont(bold=True, size=9, color=DARK_BLUE)
    cell.alignment = CENTER_ALIGN
    cell.border = thin_border()
  for col, header in enumerate(wide_headers, start=1):
    cell = wide.cell(row=2, column=col, value=CODING[header])
    cell.fill = hfill(ORANGE)
    cell.alignment = WRAP_ALIGN
    cell.border = thin_border()
    wide.column_dimensions[get_column_letter(col)].width = max(14, min(28, len(header) + 2))

  ref = wb.create_sheet("Codebook Reference")
  ref.merge_cells("A1:B1")
  ref["A1"].value = "Codebook Variable Definitions"
  ref["A1"].fill = hfill(DARK_BLUE)
  ref["A1"].font = hfont(bold=True, size=12, color=WHITE)
  ref["A1"].alignment = CENTER_ALIGN
  for idx, (variable, definition) in enumerate(CODEBOOK, start=3):
    ref.cell(row=idx, column=1, value=variable).font = hfont(bold=True)
    ref.cell(row=idx, column=2, value=definition).alignment = WRAP_ALIGN
  ref.column_dimensions["A"].width = 34
  ref.column_dimensions["B"].width = 70

  stakeholder = wb.create_sheet("Stakeholder Analysis")
  stakeholder.merge_cells("A1:B1")
  stakeholder["A1"].value = "Stakeholder Analysis — NPL_4"
  stakeholder["A1"].fill = hfill(DARK_BLUE)
  stakeholder["A1"].font = hfont(bold=True, size=12, color=WHITE)
  stakeholder["A1"].alignment = CENTER_ALIGN
  for idx, (key, value) in enumerate(STAKEHOLDER.items(), start=3):
    stakeholder.cell(row=idx, column=1, value=key).font = hfont(bold=True)
    cell = stakeholder.cell(row=idx, column=2, value=value)
    cell.alignment = WRAP_ALIGN
  stakeholder.column_dimensions["A"].width = 28
  stakeholder.column_dimensions["B"].width = 90

  sources = wb.create_sheet("Sources")
  sources.merge_cells("A1:B1")
  sources["A1"].value = "Source Documents Used for Verification"
  sources["A1"].fill = hfill(DARK_BLUE)
  sources["A1"].font = hfont(bold=True, size=12, color=WHITE)
  sources["A1"].alignment = CENTER_ALIGN
  for idx, (doc, detail) in enumerate(SOURCES, start=3):
    sources.cell(row=idx, column=1, value=doc).font = hfont(bold=True)
    sources.cell(row=idx, column=2, value=detail).alignment = WRAP_ALIGN
  sources.column_dimensions["A"].width = 42
  sources.column_dimensions["B"].width = 60

  wb.save(OUTPUT_PATH)
  print(f"Saved {OUTPUT_PATH}")


if __name__ == "__main__":
  build_workbook()
