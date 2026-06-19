#!/usr/bin/env python3
"""Generate 4P Index coding and English translation for Mozambique Decree 45/2006."""

import os
from datetime import date

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT_DIR = "/workspace/output/mozambique"
os.makedirs(OUTPUT_DIR, exist_ok=True)

POLICY_FIELDS = {
    "policy_name": (
        "Regulamento para Prevenção da Poluição e Proteção do Ambiente Marinho e Costeiro "
        "(Decreto n.º 45/2006)"
    ),
    "policy_url": "https://faolex.fao.org/docs/pdf/moz111422.pdf",
    "policy_year": 2006,
    "policy_objective": (
        "Prevent and limit pollution from illegal discharges by ships, platforms and land-based "
        "sources along Mozambique's coast; protect maritime public domain, beaches and fragile "
        "ecosystems; and control marine litter including plastics discharged at sea or deposited "
        "along the coast."
    ),
    "policy_target": 0,
    "policy_target_text": "",
    "policy_type": 0.75,
    "policy_type_justification": (
        "Executive regulation (Regulamento) approved by Council of Ministers Decree n.º 45/2006 "
        "under Articles 33 and 9(1) and Articles 12–14 of the Environment Law (Law n.º 20/97) — "
        "sub-legislative instrument, not parliamentary legislation."
    ),
    "policy_integration": 0.75,
    "policy_sectors_list": (
        "fisheries, waste management, industry, tourism, municipalities, water, packaging"
    ),
    "policy_circularity": 0.75,
    "policy_lifecycle_phases_list": (
        "consumption, disposal, environmental leakage"
    ),
    "policy_budget": 0.5,
    "policy_budget_text": (
        'Art. 10: port reception waste fees set by maritime authority; Art. 38: litter/garbage fines '
        'up to 10,000,000 MT; Art. 85: fine revenues allocated 60% INAMAR, 10% State Budget, 30% FUNAB.'
    ),
    "policy_score": "[auto]",
}

# Eight distinct instruments — consolidated to avoid overcoding procedural/enforcement duplicates.
INSTRUMENTS = [
    {
        "instrument_type": 0.80,
        "instrument_lifecycle_stage": "Waste management",
        "instrument_description": (
            "Art. 5: Ports, port installations, platforms and coastal emitter facilities must provide "
            "adequate waste collection and treatment installations; characteristics defined through "
            "environmental impact assessment including capacity, location, effluent standards and "
            "emergency equipment."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            'Art. 5(1): "deverão dispor obrigatoriamente de instalações ou meios adequados para a '
            'recolha e tratamento dos diversos tipos de resíduos"; EIA requirements in Art. 5(2); '
            'Art. 7: environmental audits; Art. 38 fines for violations.'
        ),
        "instrument_score": "[auto]",
        "comments": (
            "Core infrastructure obligation for ship and port waste streams including plastics. "
            "Not coded separately: Art. 6 internal manuals and Art. 7 contingency plans (procedural "
            "requirements supporting this system)."
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "End of life",
        "instrument_description": (
            "Art. 8–9: Ship commanders calling at national ports must deposit all ship-generated waste "
            "at port reception facilities before departure (with limited voyage-capacity exception); "
            "must submit accurate waste quantity/type forms on arrival."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            'Art. 8(1): "deverá depositar todos os resíduos gerados no seu navio num meio portuário de '
            'recepção, antes de deixar o porto"; Art. 8(3): port authority may compel delivery; '
            'Art. 9: mandatory waste data form; Art. 38(3): litter discharge fines.'
        ),
        "instrument_score": "[auto]",
        "comments": "MARPOL-aligned port reception delivery obligation applicable to all ship waste including plastics.",
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Environmental leakage",
        "instrument_description": (
            "Art. 18(2): Prohibits discharge of any waste in national jurisdictional waters outside port "
            "facilities, explicitly including synthetic ropes, synthetic fishing nets and plastic bags."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            '"Não será permitida a descarga de qualquer tipo de resíduos, inclusive cabos sintéticos, '
            'redes sintéticas de pesca e sacos plásticos, nas águas jurisdicionais nacionais, fora de '
            'instalações portuárias"; Art. 38(2)-(3): litter fines; Art. 21: incident reporting; '
            'Art. 26-27: maritime inspection powers.'
        ),
        "instrument_score": "[auto]",
        "comments": (
            "Only provision explicitly naming plastics (sacos plásticos) and synthetic fishing gear. "
            "Highest-priority plastic-specific instrument in this regulation."
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Environmental leakage",
        "instrument_description": (
            "Art. 15–18(1): Prohibits or strictly conditions at-sea discharge of high-risk hazardous "
            "substances, non-hazardous wastes, oils and oily mixtures in national waters, with "
            "ecologically sensitive area restrictions and approved procedures."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            'Art. 15: "É proibida a descarga" of high-risk substances; Art. 17–18(1): conditional '
            'prohibitions for other wastes and oils; Art. 14: MARPOL certificates required; '
            'Art. 38: graduated fines.'
        ),
        "instrument_score": "[auto]",
        "comments": (
            "Bundled vessel discharge standard — separate from Art. 18(2) plastic litter ban. "
            "Annex I incorporates MARPOL 73/78 rules referenced in Art. 2."
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Environmental leakage",
        "instrument_description": (
            "Art. 51–52: Prohibits dumping toxic/nocive effluents and non-biodegradable residues into "
            "coastal, marine, river and beach areas; prohibits depositing waste outside proper "
            "receptacles along the coast (with beach-user carry-off obligation where bins absent); "
            "bans open defecation and coastal dumps/skips/landfills in fragile ecosystems."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            'Art. 51(1): prohibition on "resíduos, especialmente de carácter não biodegradável"; '
            'Art. 52(1): "proibida a deposição de resíduos... fora dos receptáculos próprios"; '
            'Art. 38(2)-(3): specific litter fines; Art. 74: MICOA/INAMAR/municipal inspection.'
        ),
        "instrument_score": "[auto]",
        "comments": (
            "Primary land-based marine litter controls. Art. 53–61 (beach use rules) revoked by "
            "Decree 98/2020 but coastal litter prohibitions in Art. 51–52 remain operative."
        ),
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Waste management",
        "instrument_description": (
            "Art. 10: Port reception waste fees charged to ships; fee system must not incentivise "
            "marine discharge — ships must contribute significantly even if not using reception "
            "facilities, with possible reductions for ships demonstrating reduced waste generation."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            'Art. 10(2): "Os sistemas de recuperação dos custos... não devem constituir um incentivo '
            'à descarga dos resíduos no mar"; differentiated fees by ship category/type/size.'
        ),
        "instrument_score": "[auto]",
        "comments": "Economic instrument supporting Art. 8 port delivery obligation.",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Environmental leakage",
        "instrument_description": (
            "Art. 43: Ministry of Environmental Coordination must prevent and control discharge or "
            "spillage into the sea of harmful substances, garbage, wastewater and sewage from coastal "
            "establishments and land-based sources per Decrees 30/2003 and 18/2004."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            'Art. 43(1): "tomará medidas adequadas para prevenir e controlar a descarga ou o '
            'derramamento no mar de substâncias nocivas e perigosas, lixos ou águas residuais"; '
            'cross-reference to effluent standards Decree 18/2004.'
        ),
        "instrument_score": "[auto]",
        "comments": (
            "Overarching governance/coordination for land-to-sea pollution including solid waste/litter. "
            "Not coded separately: Art. 45–50 authorization and action programmes (implementation details)."
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Waste management",
        "instrument_description": (
            "Art. 38(2)–(3): Administrative fines for toxic/hazardous litter dumping (1,000,000–10,000,000 MT) "
            "and for garbage discharge/littering (5,000–75,000 MT depending on intent/negligence), plus "
            "subsidiary sanctions in Art. 39 (vessel seizure, activity suspension)."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            'Art. 38(2): "violação das disposições relativas à prevenção do lançamento de lixos tóxicos '
            'ou perigosos será punida com multa de 1.000.000,00 Mtn a 10.000.000,00 Mtn"; Art. 38(3): '
            'garbage discharge fines; Art. 36: infringement investigation; Art. 85: FUNAB allocation.'
        ),
        "instrument_score": "[auto]",
        "comments": (
            "Single enforcement instrument for litter/garbage violations; not duplicated per fine tier. "
            "Title III coastal fines in Annex VII coded under same enforcement framework."
        ),
    },
]

COLUMNS = [
    "policy_name", "policy_url", "policy_year", "policy_objective", "policy_target",
    "policy_target_text", "policy_type", "policy_type_justification", "policy_integration",
    "policy_sectors_list", "policy_circularity", "policy_lifecycle_phases_list",
    "policy_budget", "policy_budget_text", "policy_score",
    "instrument_type", "instrument_lifecycle_stage", "instrument_description",
    "instrument_in_force", "instrument_implementation", "instrument_implementation_text",
    "instrument_score", "comments",
]

COLUMN_LABELS = {
    "policy_name": "A — policy_name",
    "policy_url": "B — policy_url",
    "policy_year": "C — policy_year",
    "policy_objective": "D — policy_objective",
    "policy_target": "E — policy_target",
    "policy_target_text": "F — policy_target_text",
    "policy_type": "G — policy_type",
    "policy_type_justification": "H — policy_type_justification",
    "policy_integration": "I — policy_integration",
    "policy_sectors_list": "J — policy_sectors_list",
    "policy_circularity": "K — policy_circularity",
    "policy_lifecycle_phases_list": "L — policy_lifecycle_phases_list",
    "policy_budget": "M — policy_budget",
    "policy_budget_text": "N — policy_budget_text",
    "policy_score": "O — policy_score [auto]",
    "instrument_type": "P — instrument_type",
    "instrument_lifecycle_stage": "Q — instrument_lifecycle_stage",
    "instrument_description": "R — instrument_description",
    "instrument_in_force": "S — instrument_in_force",
    "instrument_implementation": "T — instrument_implementation",
    "instrument_implementation_text": "U — instrument_implementation_text",
    "instrument_score": "V — instrument_score [auto]",
    "comments": "W — comments",
}


def build_dataframe():
    return pd.DataFrame([{**POLICY_FIELDS, **i} for i in INSTRUMENTS], columns=COLUMNS)


def export_csv(df, basename):
    path = os.path.join(OUTPUT_DIR, f"{basename}.csv")
    df.rename(columns=COLUMN_LABELS).to_csv(path, index=False, encoding="utf-8-sig")
    return path


def export_xlsx(df, basename):
    path = os.path.join(OUTPUT_DIR, f"{basename}.xlsx")
    labelled = df.rename(columns=COLUMN_LABELS)
    labelled.to_excel(path, index=False, sheet_name="4P Index Coding")
    wb = load_workbook(path)
    ws = wb.active
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = min(45, max(12, len(str(ws.cell(1, col).value or "")) * 0.9))
    ws.freeze_panes = "A2"
    wb.save(path)
    return path


def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_para(doc, text):
    doc.add_paragraph(text)


def create_translation_doc():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading(
        "Regulation on Prevention of Marine and Coastal Pollution",
        0,
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("Decree No. 45/2006 of 30 November")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(
        "English translation of: Regulamento para Prevenção da Poluição e Proteção do "
        "Ambiente Marinho e Costeiro\n"
    ).italic = True
    meta.add_run(f"Translated for 4P Index policy analysis — {date.today().isoformat()}\n")
    meta.add_run("Source: https://faolex.fao.org/docs/pdf/moz111422.pdf")

    sections = [
        ("Decree No. 45/2006", (
            "The Council of Ministers decrees:\n"
            "Article 1. The Regulation on Prevention of Pollution and Protection of the Marine and "
            "Coastal Environment is approved as an annex forming an integral part of this Decree.\n"
            "Article 2. Decree No. 495/73 of 6 October is repealed.\n"
            "Article 3. This Regulation enters into force sixty days after publication.\n"
            "Approved by the Council of Ministers on 10 October 2006."
        )),
        ("Regulation — Title I. General provisions", (
            "Article 2 (Object). To prevent and limit pollution from illegal discharges by ships, "
            "platforms or land-based sources off the Mozambican coast, and to establish legal bases "
            "for protection of maritime, lacustrine and fluvial public domain, beaches and fragile ecosystems.\n\n"
            "Article 3 (Scope). Applies to all persons and entities whose activities may negatively "
            "impact the environment in maritime public domain areas, including discharges of harmful "
            "substances by ships, ports, coastal installations and land-based sources in internal waters, "
            "territorial sea, EEZ and high sea (for national vessels). Harmful substances classified per "
            "hazardous waste regulation.\n\n"
            "Key definitions (Art. 1): Discharge; ship-generated waste; cargo waste; dumping; fragile "
            "ecosystems; coastal zone; port reception facilities."
        )),
        ("Title II. Ships and platforms — Prevention and control", (
            "Article 5. Ports, port facilities, platforms and coastal emitter installations must have "
            "adequate waste collection and treatment means; EIA defines dimensions, capacity, standards "
            "and emergency equipment.\n\n"
            "Article 6. Operators must prepare internal pollution-risk and waste management manuals "
            "approved by the environmental authority.\n\n"
            "Article 7. Individual oil/hazardous substance contingency plans updated every five years; "
            "submitted to INAMAR and MICOA.\n\n"
            "Article 8. Ship commanders must deposit all ship-generated waste at port reception facilities "
            "before leaving port, unless sufficient storage capacity exists for the next delivery port; "
            "port authority may compel delivery if next port is inadequate.\n\n"
            "Article 9. Ships (except small fishing/recreational vessels) must submit accurate waste "
            "quantity and type data on arrival.\n\n"
            "Article 10. Port reception waste fees must not incentivise discharge at sea; ships must "
            "contribute significantly to costs; fee reductions possible for low-waste vessels.\n\n"
            "Articles 11–14. Oil and hazardous substance record books, packaging rules, insurance and "
            "MARPOL pollution prevention certificates required for applicable vessels."
        )),
        ("Title II — Chapter III. Discharges (plastic-relevant provisions)", (
            "Article 15. Discharge of high-risk harmful substances in national waters is prohibited.\n\n"
            "Article 17. Discharge of non-hazardous wastes prohibited except under cumulative conditions "
            "including international norms, outside ecologically sensitive areas, and MICOA-approved procedures.\n\n"
            "Article 18.\n"
            "1. Discharge of oils, oily mixtures and residues prohibited except as permitted internationally "
            "and with MICOA-approved procedures outside sensitive areas.\n"
            "2. Discharge of any waste — including synthetic ropes, synthetic fishing nets and plastic bags — "
            "in national jurisdictional waters outside port facilities is not permitted.\n\n"
            "Article 19. Force majeure exceptions with liability for environmental damage.\n\n"
            "Article 21. Immediate incident notification to INAMAR and MICOA."
        )),
        ("Title III. Land-based sources", (
            "Article 43. MICOA shall take adequate measures to prevent and control discharge or spillage "
            "into the sea of harmful substances, garbage, wastewater and sewage from coastal establishments, "
            "per Decrees 30/2003 and 18/2004.\n\n"
            "Article 51.\n"
            "1. Dumping toxic or harmful effluents and any non-biodegradable residues into coastal areas, "
            "territorial sea, ports, rivers, lakes, beaches and maritime administration areas is prohibited.\n"
            "2. Vessels may not discharge persistent oils or environmentally harmful substances without "
            "legal compliance.\n\n"
            "Article 52.\n"
            "1. Depositing waste outside proper receptacles along the coast and fragile ecosystems is prohibited.\n"
            "2. Open defecation in regulated areas is prohibited.\n"
            "3. Where receptacles are absent, beach users must collect and remove their waste.\n"
            "4. Installing dumps, skips, landfills or toxic material storage along the coast in fragile "
            "ecosystems is prohibited."
        )),
        ("Sanctions (selected)", (
            "Article 38.\n"
            "2. Violation of toxic/hazardous litter prevention: fine 1,000,000–10,000,000 MT.\n"
            "3. Garbage discharge/litter violations: fine 5,000–75,000 MT (by intent/negligence); "
            "omission of incident communication also fined.\n\n"
            "Article 85. Fine revenues: 60% INAMAR, 10% State Budget, 30% Environment Fund (FUNAB)."
        )),
        ("Note on subsequent amendments", (
            "Decree No. 98/2020 revoked Articles 53–61 of this Regulation (beach management provisions). "
            "Core pollution prevention, ship waste, discharge prohibitions and coastal litter rules remain "
            "in force. Annex I incorporates MARPOL 73/78 discharge rules; Annex VII lists Title III fines."
        )),
    ]

    for heading, body in sections:
        add_heading(doc, heading, level=1)
        add_para(doc, body)

    doc.add_page_break()
    add_heading(doc, "4P Index Coding Summary", level=1)
    add_para(
        doc,
        f"Policy: {POLICY_FIELDS['policy_name']}\n"
        f"Instruments coded: {len(INSTRUMENTS)} (consolidated to avoid overcoding)\n"
        f"See: 4P_Index_Mozambique_Decree_45_2006.xlsx",
    )

    path = os.path.join(OUTPUT_DIR, "Mozambique_Decree_45_2006_English_Translation.docx")
    doc.save(path)
    return path


def create_index_html(basename):
    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <title>Mozambique Decree 45/2006 — 4P Index</title>
  <style>
    body {{ font-family: system-ui, sans-serif; max-width: 800px; margin: 2rem auto; padding: 0 1rem; }}
    .card {{ border: 1px solid #ddd; border-radius: 8px; padding: 1rem; margin: 1rem 0; }}
    a.download {{ display: inline-block; background: #1F4E79; color: #fff; padding: 0.6rem 1.2rem;
      border-radius: 4px; text-decoration: none; margin-top: 0.5rem; }}
  </style>
</head>
<body>
  <h1>4P Index — Decree 45/2006 (Marine & Coastal Pollution)</h1>
  <p>Generated {date.today().isoformat()}. Eight consolidated instrument rows.</p>
  <div class="card">
    <a class="download" href="{basename}.xlsx" download>Download Excel</a><br><br>
    <a class="download" href="{basename}.csv" download>Download CSV</a><br><br>
    <a class="download" href="Mozambique_Decree_45_2006_English_Translation.docx" download>Download English translation</a>
  </div>
</body>
</html>"""
    path = os.path.join(OUTPUT_DIR, "index_decree_45_2006.html")
    with open(path, "w", encoding="utf-8") as f:
        f.write(html)
    return path


def main():
    basename = "4P_Index_Mozambique_Decree_45_2006"
    df = build_dataframe()
    paths = [export_csv(df, basename), export_xlsx(df, basename), create_translation_doc(), create_index_html(basename)]
    print("Generated:")
    for p in paths:
        print(f"  {p}")


if __name__ == "__main__":
    main()
