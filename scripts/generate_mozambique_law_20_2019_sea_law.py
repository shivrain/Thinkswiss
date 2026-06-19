#!/usr/bin/env python3
"""Generate 4P Index coding and English translation for Mozambique Law 20/2019 (Sea Law)."""

import os
from datetime import date

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT_DIR = "/workspace/output/mozambique"
os.makedirs(OUTPUT_DIR, exist_ok=True)

POLICY_FIELDS = {
    "policy_name": "Lei do Mar (Lei n.º 20/2019)",
    "policy_url": "https://sibmoz.gov.mz/content/uploads/2022/01/Lei-do-Mar.pdf",
    "policy_year": 2019,
    "policy_objective": (
        "Establish the legal regime applicable to the exercise of sovereignty and jurisdiction "
        "over the national maritime space, the exploitation of living and non-living marine resources, "
        "and the use of the maritime public domain; revise Law 4/96 to align with international "
        "instruments and ensure conservation, preservation and sustainable development of maritime space."
    ),
    "policy_target": 0,
    "policy_target_text": "",
    "policy_type": 1.0,
    "policy_type_justification": (
        "Parliamentary Law n.º 20/2019 enacted by the Assembleia da República under Articles 6(2) "
        "and 178(1) of the Constitution — primary binding legislation revoking Law n.º 4/96. "
        "Establishes rights, duties, prohibitions, crimes and contraventions enforceable in national "
        "maritime jurisdiction."
    ),
    "policy_integration": 1.0,
    "policy_sectors_list": (
        "fisheries, waste management, industry, tourism, municipalities, water, packaging"
    ),
    "policy_circularity": 0.75,
    "policy_lifecycle_phases_list": (
        "consumption, disposal, environmental leakage"
    ),
    "policy_budget": 0.50,
    "policy_budget_text": (
        "Art. 87: private use of maritime space requires payment of usage fees; Art. 89: Government "
        "may create financing mechanisms for maritime economy, fiscalization, security and marine "
        "environment protection and monitoring; Art. 5g–i polluter-pays and user-pays principles."
    ),
    "policy_score": "[auto]",
}

# Six consolidated instruments — jurisdictional chapters not overcoded line-by-line.
INSTRUMENTS = [
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Environmental leakage",
        "instrument_description": (
            "Art. 5(g–i), Art. 87: Establishes polluter-pays principle — polluters bear costs of "
            "restoring polluted marine environment from economic activities; user-pays principle "
            "requiring fees for access and use of marine and coastal resources; private maritime "
            "space use titles require payment of applicable usage fees."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            'Art. 5g: "do poluidor pagador" — polluter assumes costs of restoring polluted marine '
            'environment; Art. 5i: user-pays fee for marine/coastal resource use; Art. 87(1): '
            '"pagamento da respectiva taxa" as counterpart for private maritime space use titles.'
        ),
        "instrument_score": "[auto]",
        "comments": (
            "Binding statutory principles in primary law — stronger than POLMAR policy direction. "
            "Not coded separately: Art. 89 general maritime financing mechanisms."
        ),
    },
    {
        "instrument_type": 0.80,
        "instrument_lifecycle_stage": "Environmental leakage",
        "instrument_description": (
            "Art. 42, Art. 53: State adopts international legislation to prevent, reduce and control "
            "marine pollution; Government prohibitions include emission of toxic/non-degradable "
            "substances from terrestrial or vessel sources, vessel pollution including ballasting, "
            "intentional or accidental garbage (lixo) discharge, pollution from offshore installations "
            "and structures, and unauthorized dumping on the continental shelf."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            'Art. 42(2a): ban on emission of "substâncias tóxicas... especialmente, as não degradáveis"; '
            'Art. 42(2c): "descargas intencionais ou não de lixo"; Art. 53(3b): non-degradable '
            'substances from atmosphere or vessels or by dumping; Art. 53(3c): unauthorized dumping.'
        ),
        "instrument_score": "[auto]",
        "comments": (
            "Primary plastic-relevant prohibitions — explicit garbage discharge ban and non-degradable "
            "substances. Complements Decree 45/2006 operational rules. Not coded separately: "
            "geophysical survey mammal-protection rules (Art. 42.2f–g)."
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Environmental leakage",
        "instrument_description": (
            "Art. 93(a): Maritime crime of polluting national maritime space or degrading marine "
            "environment from any source without legal compliance, including violation of MARPOL "
            "(International Convention for the Prevention of Pollution from Ships); penalty of "
            "2–8 years imprisonment and corresponding fine."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            'Art. 93(1a): pollution or degradation of marine environment without observance of legal, '
            'regulatory or authority-imposed obligations "bem como da Convenção Internacional para '
            'a Prevenção da Poluição por Navios (MARPOL)" — "pena de prisão de dois a oito anos '
            'e multa correspondente".'
        ),
        "instrument_score": "[auto]",
        "comments": (
            "Explicit MARPOL incorporation as criminal offence — covers vessel garbage/plastics "
            "annexes. Not coded separately: Art. 93g–h toxic substance discharge crimes (bundled "
            "under general pollution crime)."
        ),
    },
    {
        "instrument_type": 0.80,
        "instrument_lifecycle_stage": "Environmental leakage",
        "instrument_description": (
            "Art. 91(6–7): Maritime fiscalization in environmental preservation includes marine "
            "pollution prevention and control of maritime waters; public-health fiscalization "
            "includes marine litter (lixo marinho) control and mitigation and sanitation of the "
            "maritime shoreline (orla marítima)."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            'Art. 91(6b): "garantia do controlo e da prevenção da poluição das águas no espaço '
            'marítimo"; Art. 91(7b): "controlo e mitigação do lixo marinho"; Art. 91(7c): '
            '"saneamento da orla marítima".'
        ),
        "instrument_score": "[auto]",
        "comments": (
            "Only explicit 'lixo marinho' reference in the law — operational mandate for integrated "
            "maritime fiscalization. Links to CEFMAR coordination (Art. 92)."
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Environmental leakage",
        "instrument_description": (
            "Art. 90–92: Maritime space fiscalization covers control, monitoring and surveillance "
            "of maritime activities including vessel/platform inspection and criminal/contravention "
            "sanctioning; Art. 91(1a) represses environmental crimes and marine pollution; "
            "CEFMAR coordinates integrated fiscalization across all competent entities."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            'Art. 90(1): fiscalization includes "autuação e sancionamento das infracções... '
            'criminal e contravencional"; Art. 91(1a): "repressão... dos crimes ambientais e da '
            'poluição no mar"; Art. 92(1): CEFMAR integrates all maritime fiscalization entities.'
        ),
        "instrument_score": "[auto]",
        "comments": (
            "Governance/enforcement architecture — operationalised with Decree 51/2024 inspection. "
            "Not coded separately: port security supervision (Art. 91.3d), piracy/immigration (91.1a)."
        ),
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Environmental leakage",
        "instrument_description": (
            "Art. 95(c–d), Art. 42(3), Art. 58: Contraventions for acts or omissions contrary to "
            "marine environment protection (Art. 42) and pollution control (Art. 58); sanctions "
            "include pecuniary penalties, restrictive measures and forfeiture of vessels/equipment "
            "to the State; Government establishes contravention sanctioning framework."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            'Art. 95(1c): contraventions for violations of Art. 42(2) and Art. 58; Art. 95(2): '
            'sanctions "pecuniária, restritivas do exercício de direitos, incluindo a perda de '
            'embarcações e/ou equipamentos a favor do Estado"; Art. 42(3): vessel inspection and '
            'detention powers for pollution offences.'
        ),
        "instrument_score": "[auto]",
        "comments": (
            "Administrative/civil enforcement layer below Art. 93 crimes. Art. 58 continental-shelf "
            "pollution jurisdiction bundled here. Not coded: innocent passage contraventions (95.1a)."
        ),
    },
]

COLUMNS = [
    "policy_name", "policy_url", "policy_year", "policy_objective", "policy_target",
    "policy_target_text", "policy_type", "policy_type_justification", "policy_integration",
    "policy_sectors_list", "policy_circularity", "policy_lifecycle_phases_list",
    "policy_budget", "policy_budget_text", "policy_score",
    "instrument_type", "instrument_lifecycle_stage", "instrument_description",
    "instrument_in_force", "instrument_implementation", "instrument_implementation_text",
    "instrument_score", "comments",
]

COLUMN_LABELS = {
    "policy_name": "A — policy_name",
    "policy_url": "B — policy_url",
    "policy_year": "C — policy_year",
    "policy_objective": "D — policy_objective",
    "policy_target": "E — policy_target",
    "policy_target_text": "F — policy_target_text",
    "policy_type": "G — policy_type",
    "policy_type_justification": "H — policy_type_justification",
    "policy_integration": "I — policy_integration",
    "policy_sectors_list": "J — policy_sectors_list",
    "policy_circularity": "K — policy_circularity",
    "policy_lifecycle_phases_list": "L — policy_lifecycle_phases_list",
    "policy_budget": "M — policy_budget",
    "policy_budget_text": "N — policy_budget_text",
    "policy_score": "O — policy_score [auto]",
    "instrument_type": "P — instrument_type",
    "instrument_lifecycle_stage": "Q — instrument_lifecycle_stage",
    "instrument_description": "R — instrument_description",
    "instrument_in_force": "S — instrument_in_force",
    "instrument_implementation": "T — instrument_implementation",
    "instrument_implementation_text": "U — instrument_implementation_text",
    "instrument_score": "V — instrument_score [auto]",
    "comments": "W — comments",
}


def build_dataframe():
    return pd.DataFrame([{**POLICY_FIELDS, **i} for i in INSTRUMENTS], columns=COLUMNS)


def export_csv(df, basename):
    path = os.path.join(OUTPUT_DIR, f"{basename}.csv")
    df.rename(columns=COLUMN_LABELS).to_csv(path, index=False, encoding="utf-8-sig")
    return path


def export_xlsx(df, basename):
    path = os.path.join(OUTPUT_DIR, f"{basename}.xlsx")
    labelled = df.rename(columns=COLUMN_LABELS)
    labelled.to_excel(path, index=False, sheet_name="4P Index Coding")
    wb = load_workbook(path)
    ws = wb.active
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = min(45, max(12, len(str(ws.cell(1, col).value or "")) * 0.9))
    ws.freeze_panes = "A2"
    wb.save(path)
    return path


def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_para(doc, text):
    doc.add_paragraph(text)


def create_translation_doc():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("Sea Law — Lei do Mar", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("Law No. 20/2019 of 8 November")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("English translation of: Lei do Mar (Lei n.º 20/2019)\n").italic = True
    meta.add_run(f"Translated for 4P Index policy analysis — {date.today().isoformat()}\n")
    meta.add_run("Source: https://sibmoz.gov.mz/content/uploads/2022/01/Lei-do-Mar.pdf")

    sections = [
        ("Law No. 20/2019", (
            "The Assembly of the Republic resolves to revise Law No. 4/96 of 4 January — Sea Law — "
            "to align with international instruments and ensure conservation, preservation and "
            "sustainable use of the national maritime space.\n\n"
            "Article 1 (Object): This Law establishes the legal regime for exercising sovereignty "
            "and jurisdiction over the national maritime space, exploitation of living and non-living "
            "marine resources, and use of the maritime public domain."
        )),
        ("Scope of application (Art. 2)", (
            "The Law applies to: (a) the national maritime space including navigable waters of "
            "public lacustrine and fluvial domains, seabed and subsoil, and adjacent public domain; "
            "(b) national and foreign vessels wherever located; (c) maritime objects, cables, pipelines, "
            "artificial islands and structures; (d) maritime infrastructure and equipment."
        )),
        ("Guiding principles (Art. 5, selected)", (
            "Ecosystem approach — preserve good environmental status of marine and coastal areas.\n\n"
            "Polluter pays — polluters bear costs of restoring the polluted marine environment.\n\n"
            "User pays — access and use of marine and coastal resources require payment of fees.\n\n"
            "Precautionary principle — adopt measures to prevent acts harmful to the marine environment."
        )),
        ("Marine environment protection (Arts. 42, 53)", (
            "The State adopts international legislation to prevent, reduce and control marine pollution. "
            "Government prohibitions include:\n"
            "a) emission of toxic, prejudicial or harmful substances, especially non-degradable ones, "
            "from land, atmosphere or vessels or by dumping;\n"
            "b) pollution from vessels including ballasting and transfer of dangerous cargoes;\n"
            "c) intentional or accidental garbage discharge;\n"
            "d) pollution from installations exploiting seabed resources;\n"
            "e) pollution from devices operating in the marine environment.\n\n"
            "On the continental shelf, unauthorized dumping and pollution from any source are prohibited."
        )),
        ("Maritime fiscalization (Arts. 90–92)", (
            "Maritime fiscalization covers control, monitoring and sanctioning of maritime activities. "
            "Environmental fiscalization includes pollution prevention and control of maritime waters. "
            "Public-health fiscalization includes marine litter control and mitigation and shoreline "
            "sanitation.\n\n"
            "CEFMAR (Maritime Fiscalization Operations Coordination Centre) integrates all entities "
            "with maritime fiscalization functions for coordinated enforcement."
        )),
        ("Crimes and contraventions (Arts. 93, 95)", (
            "Maritime crimes include polluting the national maritime space or degrading the marine "
            "environment without legal compliance, including violation of MARPOL — penalty of 2 to 8 "
            "years imprisonment and fine.\n\n"
            "Contraventions include acts contrary to marine environment protection under Art. 42 and "
            "pollution control under Art. 58, with pecuniary sanctions and possible forfeiture of "
            "vessels and equipment to the State."
        )),
        ("Note on plastics", (
            "The Sea Law does not mention plastics explicitly. Relevant provisions are: prohibitions "
            "on non-degradable substance discharge and garbage (lixo) dumping; explicit marine litter "
            "(lixo marinho) control mandate; MARPOL criminal liability covering vessel waste. "
            "Operational detail is in Decree 45/2006 and related regulations."
        )),
    ]

    for heading, body in sections:
        add_heading(doc, heading, level=1)
        add_para(doc, body)

    doc.add_page_break()
    add_heading(doc, "4P Index Coding Summary", level=1)
    add_para(
        doc,
        f"Policy: {POLICY_FIELDS['policy_name']}\n"
        f"Instruments coded: {len(INSTRUMENTS)} (consolidated — jurisdictional chapters not overcoded)\n"
        f"See: 4P_Index_Mozambique_Law_20_2019_Sea_Law.xlsx",
    )

    path = os.path.join(OUTPUT_DIR, "Mozambique_Law_20_2019_Sea_Law_English_Translation.docx")
    doc.save(path)
    return path


def create_index_html(basename):
    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <title>Mozambique Law 20/2019 Sea Law — 4P Index</title>
  <style>
    body {{ font-family: system-ui, sans-serif; max-width: 800px; margin: 2rem auto; padding: 0 1rem; }}
    .card {{ border: 1px solid #ddd; border-radius: 8px; padding: 1rem; margin: 1rem 0; }}
    a.download {{ display: inline-block; background: #1F4E79; color: #fff; padding: 0.6rem 1.2rem;
      border-radius: 4px; text-decoration: none; margin-top: 0.5rem; }}
  </style>
</head>
<body>
  <h1>4P Index — Law 20/2019 (Sea Law)</h1>
  <p>Generated {date.today().isoformat()}. Six consolidated instrument rows.</p>
  <div class="card">
    <a class="download" href="{basename}.xlsx" download>Download Excel</a><br><br>
    <a class="download" href="{basename}.csv" download>Download CSV</a><br><br>
    <a class="download" href="Mozambique_Law_20_2019_Sea_Law_English_Translation.docx" download>Download English translation</a>
  </div>
</body>
</html>"""
    path = os.path.join(OUTPUT_DIR, "index_law_20_2019_sea_law.html")
    with open(path, "w", encoding="utf-8") as f:
        f.write(html)
    return path


def main():
    basename = "4P_Index_Mozambique_Law_20_2019_Sea_Law"
    df = build_dataframe()
    paths = [
        export_csv(df, basename),
        export_xlsx(df, basename),
        create_translation_doc(),
        create_index_html(basename),
    ]
    print("Generated:")
    for p in paths:
        print(f"  {p}")


if __name__ == "__main__":
    main()
