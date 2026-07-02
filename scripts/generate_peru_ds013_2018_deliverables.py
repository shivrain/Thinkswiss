#!/usr/bin/env python3
"""Generate 4P Index Excel and English translation for Peru DS 013-2018-MINAM."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

OUTPUT_DIR = Path("/workspace/output/peru-ds-013-2018-minam")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

BRANCH = "cursor/peru-ds-013-2018-4p-index-d350"
REPO_RAW = f"https://github.com/shivrain/Thinkswiss/raw/{BRANCH}"

POLICY_NAME_ES = (
    "Decreto Supremo N.° 013-2018-MINAM – Reducción del plástico de un solo uso y "
    "consumo responsable del plástico en las entidades del Poder Ejecutivo"
)
POLICY_NAME_EN = (
    "Supreme Decree No. 013-2018-MINAM – Plastic Reduction and Responsible Plastic "
    "Consumption in Executive Branch Entities"
)
POLICY_URL = (
    "https://sinia.minam.gob.pe/normas/se-aprueba-reduccion-plastico-un-solo-uso-promueve-consumo-responsable"
)


def bi(es: str, en: str) -> str:
    return f"{es} / {en}"


POLICY = {
    "policy_name": bi(POLICY_NAME_ES, POLICY_NAME_EN),
    "policy_url": POLICY_URL,
    "policy_year": 2018,
    "policy_objective": bi(
        "Decreto supremo que obliga a todas las entidades del Poder Ejecutivo a reducir "
        "progresivamente el plástico de un solo uso y promover consumo responsable del "
        "plástico en la administración pública, mediante prohibiciones de adquisición, "
        "ingreso y uso de bolsas plásticas de un solo uso, sorbetes plásticos y envases "
        "de tecnopor, y sustitución por alternativas reutilizables, biodegradables o "
        "valorizables sin microplásticos — medida de compras públicas verdes/ecoeficientes.",
        "Supreme decree requiring all Executive Branch entities to progressively reduce "
        "single-use plastic and promote responsible plastic consumption in public "
        "administration, through prohibitions on acquisition, entry and use of single-use "
        "plastic bags, plastic straws and styrofoam containers, and replacement with "
        "reusable, biodegradable or valorizable alternatives without microplastics — a "
        "green/eco-efficient public procurement measure.",
    ),
    "policy_target": 1,
    "policy_target_text": bi(
        "Art. 3.1 y 3.2: prohibiciones en plazo no mayor de ciento ochenta (180) días "
        "hábiles desde la publicación. Art. 3.3: ingreso y uso prohibido en áreas "
        "protegidas, patrimonio cultural/natural y museos en treinta (30) días hábiles. "
        "Art. 4: prohibición similar en 180 días hábiles. Art. 2.4: bolsa de un solo uso "
        "con superficie menor de 30 x 30 cm y/o espesor menor a cincuenta micras (50 µm). "
        "DCT Primera: dotación existente utilizable hasta un (01) año desde la prohibición.",
        "Art. 3.1 and 3.2: prohibitions within no more than one hundred eighty (180) "
        "business days of publication. Art. 3.3: entry and use prohibited in protected "
        "areas, cultural/natural heritage sites and museums within thirty (30) business "
        "days. Art. 4: similar prohibition within 180 business days. Art. 2.4: single-use "
        "bag with surface area under 30 x 30 cm and/or thickness under fifty microns "
        "(50 µm). First Transitory Provision: existing stock may be used for up to one "
        "(01) year from prohibition entry into force.",
    ),
    "policy_type": 0.75,
    "policy_type_justification": bi(
        "Decreto Supremo N.° 013-2018-MINAM aprobado por el Poder Ejecutivo (2 nov. 2018, "
        "publicado 5 nov. 2018), refrendado por la Ministra del Ambiente. No enmendado. "
        "Instrumento reglamentario ejecutivo de alcance en entidades del Poder Ejecutivo. "
        "Clasificado 0.75 (decreto supremo/regulación ejecutiva).",
        "Supreme Decree No. 013-2018-MINAM approved by the Executive (2 Nov 2018, published "
        "5 Nov 2018), countersigned by the Minister of Environment. Not amended. Executive "
        "regulatory instrument applying to Executive Branch entities. Classified 0.75 "
        "(supreme decree/executive regulation).",
    ),
    "policy_integration": 0.50,
    "policy_sectors_list": bi(
        "administración pública, ambiente, consumo, compras públicas, cultura, turismo, salud",
        "public administration, environment, consumption, public procurement, culture, "
        "tourism, health",
    ),
    "policy_circularity": 0.75,
    "policy_lifecycle_phases_list": bi(
        "producción, consumo, reciclaje, fuga ambiental",
        "production, consumption, recycling, environmental leakage",
    ),
    "policy_budget": 0.5,
    "policy_budget_text": bi(
        "Art. 9: 'La implementación de la presente norma se financia con cargo al presupuesto "
        "institucional de las entidades involucradas, sin demandar recursos adicionales al "
        "Tesoro Público.'",
        "Art. 9: 'Implementation of this norm is financed from the institutional budgets of "
        "the entities involved, without requiring additional resources from the Public "
        "Treasury.'",
    ),
}

INSTRUMENTS = [
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Consumption",
        "instrument_description": bi(
            "Art. 1: promover consumo responsable del plástico y reducir uso de plástico de "
            "un solo uso en entidades del Poder Ejecutivo, reemplazándolos progresivamente "
            "por plástico reutilizable, biodegradable u otros sin microplásticos que "
            "aseguren valorización.",
            "Art. 1: promote responsible plastic consumption and reduce single-use plastic in "
            "Executive Branch entities, progressively replacing with reusable plastic, "
            "biodegradable or other alternatives that do not generate microplastic pollution "
            "and ensure valorization.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "'Debiendo reemplazarlos progresivamente' — obligación orientadora. MINAM asistencia "
            "técnica Art. 8. Sin sanción específica en Art. 1.",
            "'Must progressively replace' — guiding obligation. MINAM technical assistance "
            "Art. 8. No specific penalty in Art. 1.",
        ),
        "comments": bi(
            "Objetivo general de compras públicas sostenibles aplicable a plásticos.",
            "General sustainable public-procurement objective applicable to plastics.",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Consumption",
        "instrument_description": bi(
            "Art. 2: definiciones de plástico, microplásticos (<5 mm), bolsas de un solo uso "
            "(superficie <30 x 30 cm y/o espesor <50 µm), plástico reutilizable y "
            "biodegradabilidad según NTP 900.080.",
            "Art. 2: definitions of plastic, microplastics (<5 mm), single-use bags (surface "
            "<30 x 30 cm and/or thickness <50 µm), reusable plastic and biodegradability per "
            "Peruvian Technical Standard 900.080.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Umbrales técnicos cuantificados en definiciones. Base normativa para prohibiciones "
            "Art. 3–4. DCT Segunda establece verificación de conformidad.",
            "Quantified technical thresholds in definitions. Regulatory basis for Arts. 3–4 "
            "prohibitions. Second Transitory Provision sets conformity verification.",
        ),
        "comments": bi(
            "Estándares técnicos para identificar plásticos de un solo uso prohibidos.",
            "Technical standards to identify prohibited single-use plastics.",
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Consumption",
        "instrument_description": bi(
            "Art. 3.1: 'Prohíbase la adquisición' por entidades del Poder Ejecutivo de bolsas "
            "plásticas de un solo uso, bolsas con aditivos que generen microplásticos, "
            "sorbetes plásticos y envases de tecnopor para alimentos/bebidas — plazo 180 días "
            "hábiles.",
            "Art. 3.1: 'Acquisition is prohibited' by Executive Branch entities of single-use "
            "plastic bags, bags with microplastic-generating additives, plastic straws and "
            "styrofoam containers for food/beverages — within 180 business days.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "'Prohíbase la adquisición'. Plazo 180 días. Entidades y oficinas de administración/"
            "RR.HH. velan cumplimiento (Art. 7). Medidas administrativas Art. 7.1. Excepciones "
            "Art. 5.",
            "'Acquisition is prohibited'. 180-day deadline. Entities and admin/HR offices "
            "oversee compliance (Art. 7). Administrative measures Art. 7.1. Exceptions Art. 5.",
        ),
        "comments": bi(
            "Prohibición directa de compra pública de plásticos de un solo uso.",
            "Direct ban on public procurement of single-use plastics.",
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Consumption",
        "instrument_description": bi(
            "Art. 3.2: 'Prohíbase el uso' de bolsas de un solo uso, sorbetes plásticos y "
            "envases de tecnopor por entidades del Poder Ejecutivo — plazo 180 días hábiles.",
            "Art. 3.2: 'Use is prohibited' of single-use bags, plastic straws and styrofoam "
            "containers by Executive Branch entities — within 180 business days.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "'Prohíbase el uso'. Misma estructura de implementación que Art. 3.1. Compra pública "
            "verde vinculada a DL 1278 Art. 11.",
            "'Use is prohibited'. Same implementation structure as Art. 3.1. Green procurement "
            "linked to DL 1278 Art. 11.",
        ),
        "comments": bi(
            "Prohibición de uso interno en administración pública.",
            "Ban on in-house use within public administration.",
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Environmental leakage",
        "instrument_description": bi(
            "Art. 3.3: 'Prohíbase el ingreso y el uso' de bolsas, sorbetes y tecnopor en "
            "Áreas Naturales Protegidas, Patrimonio Cultural, Patrimonio Natural de la "
            "Humanidad y museos del Poder Ejecutivo — plazo 30 días hábiles.",
            "Art. 3.3: 'Entry and use are prohibited' of bags, straws and styrofoam in Protected "
            "Natural Areas, Cultural Heritage sites, World Natural Heritage sites and "
            "Executive Branch museums — within 30 business days.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "'Prohíbase el ingreso y el uso'. Plazo más estricto 30 días. Aplica en sitios "
            "sensibles con riesgo de fuga de plásticos al ambiente.",
            "'Entry and use are prohibited'. Stricter 30-day deadline. Applies at sensitive "
            "sites with plastic leakage risk.",
        ),
        "comments": bi(
            "Control de plásticos en áreas protegidas y patrimonio.",
            "Plastic control in protected areas and heritage sites.",
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Consumption",
        "instrument_description": bi(
            "Art. 4: 'Prohíbase el uso' de bolsas o envoltorios plásticos para entrega de "
            "información impresa a administrados y adquisición de prensa en bolsas plásticas "
            "— plazo 180 días hábiles.",
            "Art. 4: 'Use is prohibited' of plastic bags or wrappers for delivery of printed "
            "information to citizens and acquisition of press products in plastic bags — "
            "within 180 business days.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "'Prohíbase el uso' y prohibición de adquisición de diarios/revistas en bolsas "
            "plásticas. Art. 7 supervisión.",
            "'Use is prohibited' and ban on acquiring newspapers/magazines in plastic bags. "
            "Art. 7 oversight.",
        ),
        "comments": bi(
            "Reduce envoltorios plásticos en servicios al ciudadano.",
            "Reduces plastic wrappers in citizen-facing services.",
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Consumption",
        "instrument_description": bi(
            "Art. 5.1: excepciones autorizadas a prohibiciones — bolsas para asepsia/alimentos "
            "a granel, bolsas para residuos, necesidades de higiene/salud, sorbetes por "
            "necesidad médica o para niños, personas con discapacidad y adultos mayores.",
            "Art. 5.1: authorized exceptions to prohibitions — bags for asepsis/bulk food, waste "
            "disposal bags, hygiene/health needs, straws for medical need or for children, "
            "persons with disabilities and older adults.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "'Autorízase, excepcionalmente'. Limita alcance de prohibiciones Art. 3–4. Criterios "
            "RM MINAM Art. 5.2 (poder habilitante).",
            "'Exceptionally authorized'. Limits scope of Arts. 3–4 prohibitions. MINAM RM "
            "criteria Art. 5.2 (enabling power).",
        ),
        "comments": bi(
            "Excepciones reducen incondicionalidad de prohibiciones de plástico.",
            "Exceptions reduce unconditional nature of plastic prohibitions.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Consumption",
        "instrument_description": bi(
            "Art. 5.2: MINAM 'puede establecer, mediante Resolución Ministerial, los criterios "
            "para la aplicación de las excepciones' del Art. 5.1.",
            "Art. 5.2: MINAM 'may establish, by Ministerial Resolution, the criteria for applying' "
            "the Art. 5.1 exceptions.",
        ),
        "instrument_in_force": 0,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "Poder habilitante ('puede establecer'). MINAM designado. Sin RM citada en el DS "
            "como ya ejercida — score in_force 0 salvo RM posterior.",
            "Enabling power ('may establish'). MINAM designated. No RM cited in decree as "
            "already exercised — in_force 0 unless subsequent MR issued.",
        ),
        "comments": bi(
            "Criterios de excepción para plásticos aún por reglamentar vía RM.",
            "Exception criteria for plastics still to be regulated by ministerial resolution.",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Consumption",
        "instrument_description": bi(
            "Art. 6: entidades participan en campañas de difusión sobre consumo responsable del "
            "plástico y promoción del plástico reutilizable impulsadas por MINAM, con asistencia "
            "de la DGECIA.",
            "Art. 6: entities participate in dissemination campaigns on responsible plastic "
            "consumption and promotion of reusable plastic led by MINAM, with DGECIA technical "
            "support.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "'Participan en las campañas'. MINAM y DGECIA designados. Sin sanciones por "
            "incumplimiento de participación.",
            "'Participate in campaigns'. MINAM and DGECIA designated. No penalties for "
            "non-participation.",
        ),
        "comments": bi(
            "Instrumento de sensibilización sobre plástico reutilizable.",
            "Awareness instrument on reusable plastic.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Consumption",
        "instrument_description": bi(
            "Art. 7.1: entidades dictan medidas administrativas para garantizar cumplimiento del "
            "DS. Art. 7.2: oficinas de administración y RR.HH. velan por cumplimiento de "
            "prohibiciones.",
            "Art. 7.1: entities issue administrative measures to ensure decree compliance. "
            "Art. 7.2: administration and HR offices oversee compliance with prohibitions.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "'Dictan las medidas administrativas' y 'velan por el cumplimiento'. Autoridad "
            "designada en cada entidad. Mecanismo de supervisión interna.",
            "'Issue administrative measures' and 'oversee compliance'. Authority designated "
            "in each entity. Internal oversight mechanism.",
        ),
        "comments": bi(
            "Gobernanza de cumplimiento de prohibiciones de plástico en el Poder Ejecutivo.",
            "Compliance governance for plastic prohibitions in the Executive Branch.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Consumption",
        "instrument_description": bi(
            "Art. 8: MINAM (DGGRS, DGCA, DGECIA) brinda asistencia técnica para aplicación del "
            "decreto a entidades del Poder Ejecutivo.",
            "Art. 8: MINAM (DGGRS, DGCA, DGECIA) provides technical assistance for decree "
            "implementation to Executive Branch entities.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "MINAM y direcciones generales designadas. 'Brinda asistencia técnica'. Sin "
            "mecanismo de sanción.",
            "MINAM and general directorates designated. 'Provides technical assistance'. No "
            "sanction mechanism.",
        ),
        "comments": bi(
            "Coordinación nacional MINAM para reducción de plástico en administración pública.",
            "National MINAM coordination for plastic reduction in public administration.",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": bi(
            "DCT Segunda: evaluación de conformidad de biodegradabilidad de bolsas según Anexo "
            "A NTP 900.080; mientras no haya organismos acreditados, fichas técnicas del "
            "productor como verificación.",
            "Second Transitory Provision: biodegradability conformity assessment of bags per "
            "NTP 900.080 Annex A; until accredited bodies exist, producer technical datasheets "
            "as verification.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Estándar NTP 900.080 referenciado. Mecanismo de verificación transitorio por "
            "fichas técnicas. Sin organismo acreditador nacional aún.",
            "NTP 900.080 standard referenced. Transitional verification via datasheets. No "
            "national accredited body yet.",
        ),
        "comments": bi(
            "Verificación de alternativas biodegradables a plástico de un solo uso.",
            "Verification of biodegradable alternatives to single-use plastic.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Consumption",
        "instrument_description": bi(
            "DCT Primera: entidades con procedimientos de selección en curso o dotación de "
            "bienes prohibidos pueden continuar hasta agotarlos, máximo un (01) año desde "
            "entrada en vigencia de prohibición Art. 3.1.",
            "First Transitory Provision: entities with ongoing selection procedures or stock of "
            "prohibited goods may continue until depleted, maximum one (01) year from Art. "
            "3.1 prohibition entry into force.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Plazo cuantificado 1 año para uso de existencias. Flexibiliza transición de "
            "plásticos de un solo uso adquiridos previamente.",
            "Quantified 1-year deadline for using existing stock. Eases transition from "
            "previously acquired single-use plastics.",
        ),
        "comments": bi(
            "Disposición transitoria para stocks de bolsas y tecnopor.",
            "Transitory provision for bag and styrofoam stocks.",
        ),
    },
]

COLUMNS = [
    ("A", "policy_name"),
    ("B", "policy_url"),
    ("C", "policy_year"),
    ("D", "policy_objective"),
    ("E", "policy_target"),
    ("F", "policy_target_text"),
    ("G", "policy_type"),
    ("H", "policy_type_justification"),
    ("I", "policy_integration"),
    ("J", "policy_sectors_list"),
    ("K", "policy_circularity"),
    ("L", "policy_lifecycle_phases_list"),
    ("M", "policy_budget"),
    ("N", "policy_budget_text"),
    ("O", "policy_score"),
    ("P", "instrument_type"),
    ("Q", "instrument_lifecycle_stage"),
    ("R", "instrument_description"),
    ("S", "instrument_in_force"),
    ("T", "instrument_implementation"),
    ("U", "instrument_implementation_text"),
    ("V", "instrument_score"),
    ("W", "comments"),
]


def policy_score_row(instrument_type: float) -> float:
    return round(
        (
            POLICY["policy_type"]
            + POLICY["policy_integration"]
            + POLICY["policy_circularity"]
            + POLICY["policy_budget"]
            + instrument_type
        )
        / 5,
        3,
    )


def instrument_score_row(instrument_type: float, implementation: float) -> float:
    return round((instrument_type + implementation) / 2, 3)


def build_excel(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "4P Index Coding"
    ws.append([f"{col} — {name}" for col, name in COLUMNS])
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for inst in INSTRUMENTS:
        ws.append(
            [
                POLICY["policy_name"],
                POLICY["policy_url"],
                POLICY["policy_year"],
                POLICY["policy_objective"],
                POLICY["policy_target"],
                POLICY["policy_target_text"],
                POLICY["policy_type"],
                POLICY["policy_type_justification"],
                POLICY["policy_integration"],
                POLICY["policy_sectors_list"],
                POLICY["policy_circularity"],
                POLICY["policy_lifecycle_phases_list"],
                POLICY["policy_budget"],
                POLICY["policy_budget_text"],
                policy_score_row(inst["instrument_type"]),
                inst["instrument_type"],
                inst["instrument_lifecycle_stage"],
                inst["instrument_description"],
                inst["instrument_in_force"],
                inst["instrument_implementation"],
                inst["instrument_implementation_text"],
                instrument_score_row(inst["instrument_type"], inst["instrument_implementation"]),
                inst["comments"],
            ]
        )

    widths = {
        "A": 40, "B": 38, "C": 8, "D": 44, "E": 8, "F": 44, "G": 8, "H": 40,
        "I": 10, "J": 40, "K": 10, "L": 36, "M": 10, "N": 42, "O": 10, "P": 10,
        "Q": 18, "R": 50, "S": 10, "T": 12, "U": 44, "V": 10, "W": 40,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(path)


def add_article(doc: Document, es: str, en: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(es)
    r.bold = True
    doc.add_paragraph(en).paragraph_format.space_after = Pt(8)


def build_word(path: Path) -> None:
    doc = Document()
    for margin in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(doc.sections[0], margin, Inches(1))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t = title.add_run(f"{POLICY_NAME_EN}\n{POLICY_NAME_ES}")
    t.bold = True
    t.font.size = Pt(14)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(
        "Approved: 2 November 2018 | Published: 5 November 2018 | Not amended\n"
        f"Source: {POLICY_URL}"
    )

    doc.add_heading("English Translation of Key Provisions", 1)
    doc.add_paragraph(
        "English translation of principal provisions of Supreme Decree No. 013-2018-MINAM on "
        "single-use plastic reduction and responsible plastic consumption in Executive Branch "
        "entities."
    )

    sections = [
        (
            "Article 1 — Purpose",
            "promover el consumo responsable del plástico y reducir el uso plástico de un solo "
            "uso en las entidades del Poder Ejecutivo, debiendo reemplazarlos progresivamente "
            "por plástico reutilizable, biodegradable u otros...",
            "Promote responsible plastic consumption and reduce single-use plastic in Executive "
            "Branch entities, progressively replacing with reusable plastic, biodegradable or "
            "other alternatives that do not generate microplastic or hazardous-substance "
            "pollution and ensure valorization.",
        ),
        (
            "Article 2 — Key definitions",
            "Microplásticos: menos de 5 mm... Bolsas de plástico de un solo uso: superficie "
            "menor de 30 x 30 cm y/o espesor menor a cincuenta micras (50 µm)...",
            "Microplastics: under 5 mm diameter. Single-use plastic bags: surface area under "
            "30 x 30 cm and/or thickness under fifty microns (50 µm). Reusable plastic and "
            "biodegradability per NTP 900.080.",
        ),
        (
            "Article 3.1 — Acquisition ban",
            "Prohíbase la adquisición... bolsas de plástico de un solo uso; sorbetes plásticos; "
            "envases de tecnopor... en un plazo no mayor de ciento ochenta (180) días hábiles...",
            "Acquisition is prohibited of single-use plastic bags, plastic straws and styrofoam "
            "food/beverage containers by Executive Branch entities within 180 business days of "
            "publication.",
        ),
        (
            "Article 3.2–3.3 — Use and entry bans",
            "Prohíbase el uso... 180 días hábiles. Prohíbase el ingreso y el uso... Áreas "
            "Naturales Protegidas... 30 días hábiles.",
            "Use prohibited within 180 business days. Entry and use prohibited in protected "
            "natural areas, cultural heritage, world natural heritage sites and Executive "
            "museums within 30 business days.",
        ),
        (
            "Article 4 — Printed information wrappers",
            "Prohíbase el uso de bolsas o envoltorios plásticos para la entrega de información "
            "impresa... 180 días hábiles.",
            "Use of plastic bags or wrappers for delivery of printed information to citizens, "
            "and acquisition of press products in plastic bags, is prohibited within 180 "
            "business days.",
        ),
        (
            "Article 5 — Exceptions",
            "Autorízase, excepcionalmente... asepsia, residuos, higiene/salud, necesidad "
            "médica... El Ministerio del Ambiente puede establecer... criterios...",
            "Exceptional authorization for asepsis/bulk food bags, waste bags, hygiene/health "
            "needs and medical/disability/child straws. MINAM may set exception criteria by "
            "ministerial resolution.",
        ),
        (
            "Articles 6–9 — Implementation and financing",
            "campañas de difusión... medidas administrativas... asistencia técnica... se financia "
            "con cargo al presupuesto institucional...",
            "Dissemination campaigns, administrative compliance measures, MINAM technical "
            "assistance; implementation financed from institutional budgets without additional "
            "treasury resources.",
        ),
        (
            "Transitory provisions",
            "dotación de bienes... hasta por un (01) año... evaluación de conformidad de la "
            "biodegradabilidad... NTP 900.080...",
            "Existing stock may be used up to one year. Biodegradability conformity assessed per "
            "NTP 900.080; producer datasheets accepted until accredited bodies exist.",
        ),
    ]

    for heading, es, en in sections:
        doc.add_heading(heading, 2)
        add_article(doc, es, en)

    doc.add_paragraph()
    note = doc.add_paragraph(
        "Note: DS 013-2018-MINAM applies only to Executive Branch public entities. It "
        "complements DL 1278 public procurement principles (Art. 11) and MINAM's national "
        "waste-management role. It does not extend to regional/local governments or the "
        "private sector except as suppliers to public entities."
    )
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(path)


def main() -> None:
    excel = OUTPUT_DIR / "Peru_DS_013_2018_MINAM_4P_Index_Coding.xlsx"
    word = OUTPUT_DIR / "Peru_DS_013_2018_MINAM_English_Translation.docx"
    build_excel(excel)
    build_word(word)
    print(f"Created: {excel}")
    print(f"Created: {word}")
    print(f"Instruments: {len(INSTRUMENTS)}")
    print()
    print("Download links:")
    print(f"  Excel: {REPO_RAW}/output/peru-ds-013-2018-minam/Peru_DS_013_2018_MINAM_4P_Index_Coding.xlsx")
    print(
        f"  Word:  {REPO_RAW}/output/peru-ds-013-2018-minam/Peru_DS_013_2018_MINAM_English_Translation.docx"
    )


if __name__ == "__main__":
    main()
