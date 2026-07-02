#!/usr/bin/env python3
"""Generate 4P Index Excel and English translation for Peru Law 32212."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

OUTPUT_DIR = Path("/workspace/output/peru-ley-32212")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

BRANCH = "cursor/peru-ley-32212-4p-index-d350"
REPO_RAW = f"https://github.com/shivrain/Thinkswiss/raw/{BRANCH}"

POLICY_NAME_ES = (
    "Ley N.° 32212 – Ley que modifica el Decreto Legislativo 1278 y la Ley 26793 para "
    "fortalecer la gestión y el manejo de residuos sólidos"
)
POLICY_NAME_EN = (
    "Law No. 32212 – Law Amending Legislative Decree No. 1278 and Law No. 26793 to "
    "Strengthen Solid-Waste Management"
)
POLICY_URL = "https://busquedas.elperuano.pe/dispositivo/NL/2356497-1"


def bi(es: str, en: str) -> str:
    return f"{es} / {en}"


POLICY = {
    "policy_name": bi(POLICY_NAME_ES, POLICY_NAME_EN),
    "policy_url": POLICY_URL,
    "policy_year": 2024,
    "policy_objective": bi(
        "Reforma amplia del Congreso que modifica 27 artículos y el anexo del DL 1278 e "
        "incorpora Arts. 66-A–85-A: reclasifica la limpieza pública como servicio esencial "
        "(incl. valorización y transferencia), exige recipientes con código de colores, "
        "aclara reglas de material de descarte, fortalece requisitos técnicos de EO-RS, "
        "financia iniciativas vía FONAM (Ley 26793) y PROFONANPE, y refuerza segregación, "
        "valorización y economía circular — aplicable a flujos de plásticos post-consumo.",
        "Broad congressional reform amending 27 articles and the annex of DL 1278 and "
        "incorporating Arts. 66-A–85-A: reclassifies public cleaning as an essential service "
        "(incl. valorization and transfer), mandates colour-coded bins, clarifies discarded-"
        "material rules, strengthens EO-RS technical requirements, funds initiatives via FONAM "
        "(Law 26793) and PROFONANPE, and reinforces segregation, valorization and circular "
        "economy — applicable to post-consumer plastic streams.",
    ),
    "policy_target": 1,
    "policy_target_text": bi(
        "DCF Primera/Segunda: adecuación reglamentaria en 120 días calendario. Art. 36: "
        "implementación progresiva de recipientes con Código de Colores. DCF Octava: "
        "ampliación un año plazos reconversión áreas degradadas (DS 001-2022-MINAM). "
        "Art. 70: descuentos en arbitrios por segregación en fuente. DCF Séptima: "
        "fideicomisos municipales para infraestructura de residuos.",
        "DCF Primera/Segunda: regulatory adaptation within 120 calendar days. Art. 36: "
        "progressive implementation of bins under Colour Code. DCF Octava: one-year extension "
        "for degraded-area reconversion deadlines (DS 001-2022-MINAM). Art. 70: cleaning-fee "
        "discounts for source segregation. DCF Séptima: municipal trusts for waste "
        "infrastructure.",
    ),
    "policy_type": 1.0,
    "policy_type_justification": bi(
        "Ley N.° 32212 aprobada por el Congreso (insistencia 31 oct. 2024, publicada 21 dic. "
        "2024 en El Peruano). Modifica DL 1278 y Ley 26793 (FONAM). No enmendada. "
        "Clasificado 1.0 (legislación parlamentaria).",
        "Law No. 32212 enacted by Congress (insistence 31 Oct 2024, published 21 Dec 2024 in "
        "El Peruano). Amends DL 1278 and Law 26793 (FONAM). Not amended. Classified 1.0 "
        "(parliamentary legislation).",
    ),
    "policy_integration": 1,
    "policy_sectors_list": bi(
        "ambiente, gobiernos municipales y regionales, industria, minería no metálica, "
        "inversión pública/privada, reciclaje, economía circular, FONAM/PROFONANPE",
        "environment, municipal and regional governments, industry, non-metallic mining, "
        "public/private investment, recycling, circular economy, FONAM/PROFONANPE",
    ),
    "policy_circularity": 1,
    "policy_lifecycle_phases_list": bi(
        "producción, consumo, recolección, reciclaje, valorización, disposición, fuga ambiental",
        "production, consumption, collection, recycling, valorization, disposal, environmental "
        "leakage",
    ),
    "policy_budget": 0.75,
    "policy_budget_text": bi(
        "Art. 3 Ley 26793(g): FONAM recibe reparaciones civiles por delitos ambientales en "
        "manejo de residuos. Art. 70-A: recaudación arbitrios vía servicios públicos. DCF "
        "Quinta: transferencia gratuita predios SBN. DCF Séptima: fideicomisos municipales. "
        "Art. 3 DL 1278: incentivos inversión pública/privada en limpieza pública.",
        "Art. 3 Law 26793(g): FONAM receives civil reparations for environmental crimes in "
        "waste management. Art. 70-A: cleaning-fee collection via public-service bills. "
        "DCF Quinta: free SBN land transfers. DCF Séptima: municipal trusts. Art. 3 DL 1278: "
        "public/private investment incentives in public cleaning.",
    ),
}

INSTRUMENTS = [
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "Art. 1: modifica 27 artículos y el anexo del DL 1278; incorpora Arts. 66-A, 66-B, "
            "66-C, 70-A y 85-A.",
            "Art. 1: amends 27 articles and the annex of DL 1278; incorporates Arts. 66-A, "
            "66-B, 66-C, 70-A and 85-A.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Acto habilitante de la reforma amplia. Vigencia desde publicación 21 dic. 2024.",
            "Enabling act for broad reform. Effective from publication 21 Dec 2024.",
        ),
        "comments": bi(
            "Reforma estructural del marco legal de residuos sólidos/plásticos.",
            "Structural reform of solid-waste/plastics legal framework.",
        ),
    },
    {
        "instrument_type": 0.75,
        "instrument_lifecycle_stage": "Collection",
        "instrument_description": bi(
            "Art. 3 DL 1278 modificado: servicio de limpieza pública esencial que comprende "
            "recolección, transporte, valorización, transferencia y disposición final; "
            "garantizado por gobiernos locales con incentivos a inversión pública y privada.",
            "Amended DL 1278 Art. 3: essential public cleaning service comprising collection, "
            "transport, valorization, transfer and final disposal; guaranteed by local "
            "governments with public and private investment incentives.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Reclasificación normativa como servicio público esencial. Aplica a todo el ciclo "
            "incl. valorización de plásticos.",
            "Regulatory reclassification as essential public service. Applies to full cycle "
            "incl. plastics valorization.",
        ),
        "comments": bi(
            "Incluye expresamente valorización y transferencia en limpieza pública.",
            "Expressly includes valorization and transfer in public cleaning.",
        ),
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": bi(
            "Art. 9 DL 1278 modificado: define material de descarte; excluye reingreso al mismo "
            "proceso y materiales de posconsumo o actividad complementaria.",
            "Amended DL 1278 Art. 9: defines discard material; excludes re-entry to same "
            "process and post-consumption or complementary-activity materials.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Clarifica frontera material de descarte/residuo sólido (relevante para plásticos "
            "industriales vs. post-consumo).",
            "Clarifies discard-material/solid-waste boundary (relevant for industrial vs. "
            "post-consumer plastics).",
        ),
        "comments": bi(
            "Post-consumo explícitamente excluido de material de descarte.",
            "Post-consumption explicitly excluded from discard material.",
        ),
    },
    {
        "instrument_type": 0.75,
        "instrument_lifecycle_stage": "Consumption",
        "instrument_description": bi(
            "Art. 33 DL 1278 modificado: segregación obligatoria en fuente de generación, áreas "
            "de acondicionamiento o infraestructuras de valorización autorizadas; prohibida en "
            "áreas de disposición final directa.",
            "Amended DL 1278 Art. 33: mandatory segregation at source, conditioning areas or "
            "authorized valorization infrastructure; prohibited in direct final-disposal areas.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Obligación de segregación aplicable a inorgánicos/plásticos. Fiscalización "
            "municipal/OEFA.",
            "Segregation obligation applicable to inorganics/plastics. Municipal/OEFA "
            "enforcement.",
        ),
        "comments": bi(
            "Base legal reforzada para segregación de plásticos post-consumo.",
            "Strengthened legal basis for post-consumer plastics segregation.",
        ),
    },
    {
        "instrument_type": 0.75,
        "instrument_lifecycle_stage": "Collection",
        "instrument_description": bi(
            "Art. 36 DL 1278 (6º párrafo): municipalidades obligadas a colocar recipientes "
            "con Código de Colores para Almacenamiento de Residuos Sólidos en parques, "
            "plazas, paraderos, mercados y centros comerciales; implementación progresiva.",
            "DL 1278 Art. 36 (6th paragraph): municipalities must place bins under the Solid-"
            "Waste Storage Colour Code in parks, squares, bus stops, markets and shopping "
            "centers; progressive implementation.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "Obligación municipal con implementación progresiva. Sin plazo específico en la "
            "ley.",
            "Municipal obligation with progressive implementation. No specific deadline in law.",
        ),
        "comments": bi(
            "Recipientes codificados por color para segregación de plásticos y otros residuos.",
            "Colour-coded bins for plastics and other waste segregation.",
        ),
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": bi(
            "Art. 36-A DL 1278 modificado: acondicionamiento de residuos orgánicos e "
            "inorgánicos (segregación, trituración, compactación, peletizado, etc.) en áreas "
            "de acondicionamiento o infraestructuras de valorización.",
            "Amended DL 1278 Art. 36-A: conditioning of organic and inorganic waste "
            "(segregation, shredding, compaction, pelletizing, etc.) in conditioning areas or "
            "valorization infrastructure.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Habilita acondicionamiento de plásticos para reciclaje industrial. Requiere "
            "autorización ambiental.",
            "Enables plastics conditioning for industrial recycling. Requires environmental "
            "authorization.",
        ),
        "comments": bi(
            "Incluye inorgánicos (PET, PE, PP, etc.) en cadena de acondicionamiento.",
            "Includes inorganics (PET, PE, PP, etc.) in conditioning chain.",
        ),
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": bi(
            "Art. 48 DL 1278 modificado: valorización material incluye reutilización, reciclaje, "
            "compostaje, recuperación de componentes/materiales para reincorporación en "
            "actividades productivas.",
            "Amended DL 1278 Art. 48: material valorization includes reuse, recycling, "
            "composting, recovery of components/materials for reincorporation in productive "
            "activities.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Definición ampliada de reciclaje/valorización material. Marco para industrialización.",
            "Expanded material recycling/valorization definition. Framework for "
            "industrialization.",
        ),
        "comments": bi(
            "Reciclaje explícito como forma de valorización material.",
            "Recycling explicitly as material valorization form.",
        ),
    },
    {
        "instrument_type": 0.75,
        "instrument_lifecycle_stage": "Collection",
        "instrument_description": bi(
            "Art. 60 DL 1278 modificado: EO-RS deben registrarse ante MINAM, contar con "
            "profesional especializado en gestión de residuos como director técnico, y contar "
            "con equipos/plantas/infraestructuras autorizadas.",
            "Amended DL 1278 Art. 60: EO-RS must register with MINAM, employ a waste-"
            "management specialist as technical director, and have authorized equipment/plants/"
            "infrastructure.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Requisitos técnicos reforzados para operadores que manejan plásticos y otros "
            "residuos. Reglamento desarrolla requisitos.",
            "Strengthened technical requirements for operators handling plastics and other "
            "waste. Regulation develops requirements.",
        ),
        "comments": bi(
            "Fortalecimiento EO-RS relevante para reciclaje post-consumo.",
            "EO-RS strengthening relevant to post-consumer recycling.",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Consumption",
        "instrument_description": bi(
            "Art. 70 DL 1278 modificado: municipalidades pueden otorgar descuentos en arbitrios "
            "de limpieza pública como incentivo a generadores que cumplan segregación en la "
            "fuente.",
            "Amended DL 1278 Art. 70: municipalities may grant cleaning-fee discounts as "
            "incentive to generators complying with source segregation.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "Incentivo tributario municipal voluntario. Vinculado a segregación de plásticos.",
            "Voluntary municipal tax incentive. Linked to plastics segregation.",
        ),
        "comments": bi(
            "Mecanismo de cambio de comportamiento para segregación domiciliaria.",
            "Behavior-change mechanism for household segregation.",
        ),
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Collection",
        "instrument_description": bi(
            "Art. 70-A incorporado: municipalidades pueden recaudar arbitrios de limpieza "
            "pública a través del comprobante mensual de cualquier servicio público; fondos "
            "destinados exclusivamente al servicio de limpieza e inversiones en residuos.",
            "Incorporated Art. 70-A: municipalities may collect cleaning fees through monthly "
            "public-service bills; funds exclusively for cleaning service and waste investments.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Excepción a DL 25988 para sostenibilidad financiera del servicio. Convenios con "
            "empresas prestadoras (Arts. 23(p), 24(b)).",
            "Exception to DL 25988 for service financial sustainability. Agreements with "
            "utility companies (Arts. 23(p), 24(b)).",
        ),
        "comments": bi(
            "Financia infraestructura de recolección/valorización de plásticos municipales.",
            "Finances municipal plastics collection/valorization infrastructure.",
        ),
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": bi(
            "Art. 3 Ley 26793 modificado (literal g): FONAM recibe montos de reparaciones "
            "civiles y acuerdos reparatorios por delitos ambientales en manejo de residuos "
            "sólidos, destinados a recuperación de áreas degradadas, rellenos sanitarios y "
            "cierre de brechas de infraestructura.",
            "Amended Law 26793 Art. 3 (lit. g): FONAM receives civil reparation and settlement "
            "amounts for environmental crimes in solid-waste management, for degraded-area "
            "recovery, landfills and infrastructure-gap closure.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Nueva fuente de financiamiento FONAM vinculada a infracciones en gestión de "
            "residuos/plásticos.",
            "New FONAM funding source linked to waste/plastics management violations.",
        ),
        "comments": bi(
            "Financia iniciativas de gestión integral incluyendo flujos plásticos.",
            "Funds integrated-management initiatives including plastic streams.",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "Art. 15(w) DL 1278 incorporado: MINAM aprueba Hoja de Ruta Nacional de Economía "
            "Circular del Perú con Plan Nacional de Acción de seguimiento.",
            "Incorporated DL 1278 Art. 15(w): MINAM approves National Circular Economy "
            "Roadmap of Peru with follow-up National Action Plan.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "Instrumento de planificación nacional. Desarrollo reglamentario pendiente.",
            "National planning instrument. Regulatory development pending.",
        ),
        "comments": bi(
            "Marco de economía circular aplicable a materiales plásticos.",
            "Circular-economy framework applicable to plastic materials.",
        ),
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Collection",
        "instrument_description": bi(
            "Art. 24.2(h) DL 1278 incorporado: municipalidades emiten constancia de "
            "formalización y administran Registro Municipal de Organizaciones de Recicladores "
            "para programas de segregación y recolección selectiva.",
            "Incorporated DL 1278 Art. 24.2(h): municipalities issue formalization certificates "
            "and administer Municipal Registry of Recycler Organizations for segregation and "
            "selective-collection programs.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Competencia municipal codificada en ley. Vinculada a Ley 29419 y DS 001-2022.",
            "Municipal competence codified in law. Linked to Law 29419 and DS 001-2022.",
        ),
        "comments": bi(
            "Formalización de recicladores de plásticos post-consumo.",
            "Formalization of post-consumer plastics recyclers.",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": bi(
            "Art. 42 DL 1278 modificado: importación de residuos no peligrosos permitida para "
            "valorización o acondicionamiento; exportación para valorización o disposición; "
            "MINAM autoriza conforme Convenio de Basilea.",
            "Amended DL 1278 Art. 42: import of non-hazardous waste permitted for valorization "
            "or conditioning; export for valorization or disposal; MINAM authorizes per Basel "
            "Convention.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Marco transfronterizo de residuos plásticos reciclables. EO-RS y generadores NM.",
            "Transboundary framework for recyclable plastic waste. EO-RS and NM generators.",
        ),
        "comments": bi(
            "Aplica a flujos de plásticos PE, PP, PET para reciclaje industrial.",
            "Applies to PE, PP, PET flows for industrial recycling.",
        ),
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "DCF Primera: Poder Ejecutivo adecúa Reglamento DL 1278 (DS 014-2017-MINAM) en "
            "120 días calendario, exonerado de AIR Ex Ante.",
            "DCF Primera: Executive adapts DL 1278 Regulation (DS 014-2017-MINAM) within 120 "
            "calendar days, exempt from Ex Ante Regulatory Impact Analysis.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "Plazo vencido (abr. 2025). Desarrollo reglamentario pendiente de verificación.",
            "Deadline elapsed (Apr 2025). Regulatory development subject to verification.",
        ),
        "comments": bi(
            "Mandato de operacionalización de la reforma en reglamento.",
            "Mandate to operationalize reform in regulation.",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Disposal",
        "instrument_description": bi(
            "DCF Séptima: gobiernos locales pueden constituir fideicomisos para garantizar "
            "continuidad del servicio de limpieza pública, operación y reinversión en "
            "infraestructura de residuos y áreas degradadas.",
            "DCF Séptima: local governments may establish trusts to guarantee continuity of "
            "public cleaning service, operation and reinvestment in waste infrastructure and "
            "degraded areas.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "Mecanismo financiero voluntario municipal. Sin plazo.",
            "Voluntary municipal financial mechanism. No deadline.",
        ),
        "comments": bi(
            "Financia infraestructura de todo el ciclo incluyendo valorización de plásticos.",
            "Finances whole-cycle infrastructure including plastics valorization.",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Collection",
        "instrument_description": bi(
            "Art. 85-A incorporado: en emergencias que cesen limpieza pública, gobiernos "
            "regionales y autoridades sectoriales pueden ejecutar acciones inmediatas para "
            "garantizar continuidad del servicio integral de limpieza pública.",
            "Incorporated Art. 85-A: when emergencies halt public cleaning, regional "
            "governments and sector authorities may take immediate action to ensure continuity "
            "of integrated public cleaning service.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "Mecanismo de contingencia. Aplica en desastres/riesgos.",
            "Contingency mechanism. Applies in disasters/risks.",
        ),
        "comments": bi(
            "Protege cadena de recolección/valorización ante interrupciones.",
            "Protects collection/valorization chain against interruptions.",
        ),
    },
]

COLUMNS = [
    ("A", "policy_name"),
    ("B", "policy_url"),
    ("C", "policy_year"),
    ("D", "policy_objective"),
    ("E", "policy_target"),
    ("F", "policy_target_text"),
    ("G", "policy_type"),
    ("H", "policy_type_justification"),
    ("I", "policy_integration"),
    ("J", "policy_sectors_list"),
    ("K", "policy_circularity"),
    ("L", "policy_lifecycle_phases_list"),
    ("M", "policy_budget"),
    ("N", "policy_budget_text"),
    ("O", "policy_score"),
    ("P", "instrument_type"),
    ("Q", "instrument_lifecycle_stage"),
    ("R", "instrument_description"),
    ("S", "instrument_in_force"),
    ("T", "instrument_implementation"),
    ("U", "instrument_implementation_text"),
    ("V", "instrument_score"),
    ("W", "comments"),
]


def policy_score_row(instrument_type: float) -> float:
    return round(
        (
            POLICY["policy_type"]
            + POLICY["policy_integration"]
            + POLICY["policy_circularity"]
            + POLICY["policy_budget"]
            + instrument_type
        )
        / 5,
        3,
    )


def instrument_score_row(instrument_type: float, implementation: float) -> float:
    return round((instrument_type + implementation) / 2, 3)


def build_excel(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "4P Index Coding"
    ws.append([f"{col} — {name}" for col, name in COLUMNS])
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for inst in INSTRUMENTS:
        ws.append(
            [
                POLICY["policy_name"],
                POLICY["policy_url"],
                POLICY["policy_year"],
                POLICY["policy_objective"],
                POLICY["policy_target"],
                POLICY["policy_target_text"],
                POLICY["policy_type"],
                POLICY["policy_type_justification"],
                POLICY["policy_integration"],
                POLICY["policy_sectors_list"],
                POLICY["policy_circularity"],
                POLICY["policy_lifecycle_phases_list"],
                POLICY["policy_budget"],
                POLICY["policy_budget_text"],
                policy_score_row(inst["instrument_type"]),
                inst["instrument_type"],
                inst["instrument_lifecycle_stage"],
                inst["instrument_description"],
                inst["instrument_in_force"],
                inst["instrument_implementation"],
                inst["instrument_implementation_text"],
                instrument_score_row(inst["instrument_type"], inst["instrument_implementation"]),
                inst["comments"],
            ]
        )

    widths = {
        "A": 40, "B": 38, "C": 8, "D": 44, "E": 8, "F": 44, "G": 8, "H": 40,
        "I": 10, "J": 40, "K": 10, "L": 36, "M": 10, "N": 42, "O": 10, "P": 10,
        "Q": 18, "R": 50, "S": 10, "T": 12, "U": 44, "V": 10, "W": 40,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(path)


def add_article(doc: Document, es: str, en: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(es)
    r.bold = True
    doc.add_paragraph(en).paragraph_format.space_after = Pt(8)


def build_word(path: Path) -> None:
    doc = Document()
    for margin in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(doc.sections[0], margin, Inches(1))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t = title.add_run(f"{POLICY_NAME_EN}\n{POLICY_NAME_ES}")
    t.bold = True
    t.font.size = Pt(14)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(
        "Congress insistence: 31 October 2024 | Published: 21 December 2024 (El Peruano)\n"
        "Amends DL 1278 and Law 26793 (FONAM)\n"
        f"Source: {POLICY_URL}"
    )

    doc.add_heading("English Translation of Key Provisions", 1)
    doc.add_paragraph(
        "English translation of principal provisions of Law No. 32212 strengthening solid-waste "
        "management."
    )

    sections = [
        (
            "Art. 3 DL 1278 — Essential public cleaning",
            "Limpieza pública es servicio esencial que comprende recolección, transporte, "
            "valorización, transferencia y disposición final.",
            "Public cleaning is an essential service comprising collection, transport, "
            "valorization, transfer and final disposal.",
        ),
        (
            "Art. 36 — Colour-coded bins",
            "Municipalidades obligadas a colocar recipientes con Código de Colores en espacios "
            "públicos; implementación progresiva.",
            "Municipalities must place bins under the Colour Code in public spaces; progressive "
            "implementation.",
        ),
        (
            "Art. 60 — EO-RS requirements",
            "EO-RS deben contar con profesional especializado, equipos e infraestructuras "
            "autorizadas.",
            "EO-RS must have specialist professional, equipment and authorized infrastructure.",
        ),
        (
            "Art. 3 Law 26793 — FONAM funding",
            "FONAM recibe reparaciones civiles por delitos ambientales en manejo de residuos "
            "sólidos.",
            "FONAM receives civil reparations for environmental crimes in solid-waste management.",
        ),
        (
            "Art. 70-A — Cleaning-fee collection",
            "Arbitrios de limpieza pública recaudables vía comprobante de servicios públicos.",
            "Cleaning fees collectible through public-service bills.",
        ),
        (
            "DCF Primera — Regulation deadline",
            "Adecuación del Reglamento DL 1278 en 120 días calendario.",
            "Adaptation of DL 1278 Regulation within 120 calendar days.",
        ),
    ]

    for heading, es, en in sections:
        doc.add_heading(heading, 2)
        add_article(doc, es, en)

    doc.add_paragraph()
    note = doc.add_paragraph(
        "Note: Law 32212 is a broad reform of Peru's flagship solid-waste law (DL 1278), "
        "strengthening public cleaning as an essential service, segregation infrastructure, "
        "EO-RS technical standards and FONAM financing — all applicable to plastic waste streams."
    )
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(path)


def main() -> None:
    excel = OUTPUT_DIR / "Peru_Ley_32212_4P_Index_Coding.xlsx"
    word = OUTPUT_DIR / "Peru_Ley_32212_English_Translation.docx"
    build_excel(excel)
    build_word(word)
    print(f"Created: {excel}")
    print(f"Created: {word}")
    print(f"Instruments: {len(INSTRUMENTS)}")
    print()
    print("Download links:")
    print(
        f"  Excel: {REPO_RAW}/output/peru-ley-32212/"
        "Peru_Ley_32212_4P_Index_Coding.xlsx"
    )
    print(
        f"  Word:  {REPO_RAW}/output/peru-ley-32212/"
        "Peru_Ley_32212_English_Translation.docx"
    )


if __name__ == "__main__":
    main()
