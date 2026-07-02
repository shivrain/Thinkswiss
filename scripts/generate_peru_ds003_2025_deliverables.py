#!/usr/bin/env python3
"""Generate 4P Index Excel and English translation for Peru DS 003-2025-MINAM."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

OUTPUT_DIR = Path("/workspace/output/peru-ds-003-2025-minam")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

BRANCH = "cursor/peru-ds-003-2025-4p-index-d350"
REPO_RAW = f"https://github.com/shivrain/Thinkswiss/raw/{BRANCH}"

POLICY_NAME_ES = (
    "Decreto Supremo N.° 003-2025-MINAM – Hoja de Ruta Nacional de Economía Circular "
    "al 2030"
)
POLICY_NAME_EN = (
    "Supreme Decree No. 003-2025-MINAM – National Circular Economy Roadmap to 2030"
)
POLICY_URL = (
    "https://www.gob.pe/institucion/minam/normas-legales/6507211-003-2025-minam"
)


def bi(es: str, en: str) -> str:
    return f"{es} / {en}"


POLICY = {
    "policy_name": bi(POLICY_NAME_ES, POLICY_NAME_EN),
    "policy_url": POLICY_URL,
    "policy_year": 2025,
    "policy_objective": bi(
        "Decreto Supremo que aprueba la Hoja de Ruta Nacional de Economía Circular al "
        "2030 (HRNEC) como instrumento estratégico multisectorial (10 sectores del "
        "Ejecutivo) para transición a modelos circulares de producción y consumo; "
        "articula políticas sectoriales, cubre cadenas de valor de plásticos, dirige "
        "hojas de ruta sectoriales y monitoreo MINAM; vinculado PNCP 2024-2030 Medida "
        "9.2 y DL 1278 Art. 15(w).",
        "Supreme Decree approving the National Circular Economy Roadmap to 2030 (HRNEC) "
        "as a multisectoral strategic instrument (10 Executive Branch sectors) for "
        "transition to circular production and consumption models; articulates sectoral "
        "policies, covers plastics value chains, directs sectoral roadmaps and MINAM "
        "monitoring; linked to PNCP 2024-2030 Measure 9.2 and DL 1278 Art. 15(w).",
    ),
    "policy_target": 1,
    "policy_target_text": bi(
        "Horizonte 2030. DCF Segunda: Plan Nacional de Acción de monitoreo en 120 días "
        "calendario. DCF Tercera: lineamientos de implementación en 180 días. HRNEC "
        "(Anexo): metas de corto, mediano y largo plazo; proyección incremento 2% PBI "
        "(S/13 908 millones) y 306 000+ empleos sostenibles al 2030; reducción consumo "
        "75,3 Mt materiales (según HRNEC aprobada).",
        "2030 horizon. DCF Segunda: National Monitoring Action Plan within 120 calendar "
        "days. DCF Tercera: implementation guidelines within 180 days. HRNEC (Annex): "
        "short-, medium- and long-term targets; projected 2% GDP increase (PEN 13,908 "
        "million) and 306,000+ sustainable jobs by 2030; 75.3 Mt materials consumption "
        "reduction (per approved HRNEC).",
    ),
    "policy_type": 0.5,
    "policy_type_justification": bi(
        "Decreto Supremo N.° 003-2025-MINAM aprobado por el Poder Ejecutivo (24 feb. "
        "2025, publicado 25 feb. 2025 en El Peruano), refrendado por PCM y 9 ministerios. "
        "Aprueba hoja de ruta nacional estratégica (Anexo integral). Vinculado DL 1278 "
        "Art. 15(w) y PNCP Medida 9.2. Clasificado 0.5 (decreto supremo/estrategia "
        "nacional).",
        "Supreme Decree No. 003-2025-MINAM approved by the Executive (24 Feb 2025, "
        "published 25 Feb 2025 in El Peruano), countersigned by PCM and 9 ministries. "
        "Approves national strategic roadmap (integral Annex). Linked to DL 1278 Art. "
        "15(w) and PNCP Measure 9.2. Classified 0.5 (supreme decree/national strategy).",
    ),
    "policy_integration": 1,
    "policy_sectors_list": bi(
        "transversal (10 sectores: ambiente, economía, producción, energía y minas, "
        "vivienda, comercio exterior, agricultura, trabajo, educación, PCM); plásticos, "
        "construcción, alimentación, residuos sólidos, transporte",
        "cross-cutting (10 sectors: environment, economy, production, energy & mining, "
        "housing, foreign trade, agriculture, labour, education, PCM); plastics, "
        "construction, food, solid waste, transport",
    ),
    "policy_circularity": 1,
    "policy_lifecycle_phases_list": bi(
        "producción, consumo, recolección, reciclaje, valorización, disposición "
        "(marco estratégico ciclo de vida completo)",
        "production, consumption, collection, recycling, valorization, disposal "
        "(whole-life-cycle strategic framework)",
    ),
    "policy_budget": 0.5,
    "policy_budget_text": bi(
        "Art. 4: financiamiento con presupuesto institucional de entidades involucradas, "
        "sin recursos adicionales del Tesoro Público. HRNEC proyecta ~S/2 800 millones/año "
        "en inversiones públicas y público-privadas hacia 2030 (Anexo).",
        "Art. 4: financing from institutional budgets of involved entities, without "
        "additional Treasury resources. HRNEC projects ~PEN 2,800 million/year in public "
        "and public-private investments toward 2030 (Annex).",
    ),
}

INSTRUMENTS = [
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "Art. 1: aprueba la 'Hoja de Ruta Nacional de Economía Circular al 2030' como "
            "Anexo integrante del DS.",
            "Art. 1: approves the 'National Circular Economy Roadmap to 2030' as integral "
            "Annex to the supreme decree.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Acto habilitante de la HRNEC. Vigencia desde publicación 25 feb. 2025.",
            "Enabling act for HRNEC. Effective from publication 25 Feb 2025.",
        ),
        "comments": bi(
            "Instrumento estratégico nacional de economía circular.",
            "National circular-economy strategic instrument.",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "Art. 2: HRNEC aplicable a entidades del sector público que incorporen economía "
            "circular en normas, instrumentos e intervenciones, en marco de competencias.",
            "Art. 2: HRNEC applicable to public-sector entities incorporating circular "
            "economy in norms, instruments and interventions within their competencies.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "Ámbito declarativo de aplicación sector público. Sin sanciones en el artículo.",
            "Declaratory public-sector scope. No sanctions in the article.",
        ),
        "comments": bi(
            "Marco transversal para integración de circularidad en políticas sectoriales.",
            "Cross-cutting framework for circularity integration in sector policies.",
        ),
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "Art. 3: MINAM responsable de monitoreo y evaluación de indicadores y acciones "
            "estratégicas de la HRNEC; implementa mecanismos de coordinación y capacitación.",
            "Art. 3: MINAM responsible for monitoring and evaluation of HRNEC indicators "
            "and strategic actions; implements coordination and training mechanisms.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Función rectora MINAM sobre implementación HRNEC. Vinculada DL 1278 Art. "
            "15(w) y Ley 32212.",
            "MINAM steering function on HRNEC implementation. Linked to DL 1278 Art. 15(w) "
            "and Law 32212.",
        ),
        "comments": bi(
            "Monitoreo de avances en circularidad de plásticos y otros materiales.",
            "Monitoring progress on plastics and other materials circularity.",
        ),
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "Art. 3 (2º párrafo): entidades implementadoras reportan avances al MINAM "
            "conforme a lineamientos que este apruebe.",
            "Art. 3 (2nd paragraph): implementing entities report progress to MINAM per "
            "guidelines it approves.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Obligación de reporte sectorial. Lineamientos vía DCF Tercera (180 días).",
            "Sectoral reporting obligation. Guidelines via DCF Tercera (180 days).",
        ),
        "comments": bi(
            "Accountability de 10 sectores del Ejecutivo refrendantes.",
            "Accountability of 10 countersigning Executive Branch sectors.",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "HRNEC Anexo — Marco estratégico: orienta transición hacia producción y "
            "consumo sostenible; optimiza uso de materiales, elimina generación de "
            "residuos y regenera ecosistemas; objetivos y acciones a corto, mediano y "
            "largo plazo.",
            "HRNEC Annex — Strategic framework: orients transition to sustainable "
            "production and consumption; optimizes material use, eliminates waste "
            "generation and regenerates ecosystems; short-, medium- and long-term "
            "objectives and actions.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "Contenido del Anexo aprobado por Art. 1. Horizonte 2030.",
            "Annex content approved by Art. 1. 2030 horizon.",
        ),
        "comments": bi(
            "Marco estratégico de ciclo de vida completo.",
            "Whole-life-cycle strategic framework.",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "HRNEC Anexo — Priorización multisectorial (10 sectores refrendantes): "
            "articula esfuerzos sectoriales incluyendo plásticos/envases, residuos "
            "sólidos, construcción, alimentación, transporte y producción.",
            "HRNEC Annex — Multisector prioritization (10 countersigning sectors): "
            "articulates sectoral efforts including plastics/packaging, solid waste, "
            "construction, food, transport and production.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "Sectores priorizados en Anexo HRNEC. Hojas de ruta sectoriales vía DCF "
            "Primera.",
            "Priority sectors in HRNEC Annex. Sectoral roadmaps via DCF Primera.",
        ),
        "comments": bi(
            "Cubre explícitamente cadenas de valor de plásticos.",
            "Explicitly covers plastics value chains.",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "HRNEC Anexo — Metas proyectadas al 2030: incremento 2% PBI (S/13 908 "
            "millones), 306 000+ empleos sostenibles, reducción 75,3 Mt consumo de "
            "materiales.",
            "HRNEC Annex — Projected 2030 targets: 2% GDP increase (PEN 13,908 million), "
            "306,000+ sustainable jobs, 75.3 Mt materials consumption reduction.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "Metas del Anexo (no en cuerpo del DS). Monitoreo vía Art. 3 y PANM&E.",
            "Annex targets (not in DS body). Monitoring via Art. 3 and National Action Plan.",
        ),
        "comments": bi(
            "Impacto económico proyectado incluye reciclaje/eficiencia de plásticos.",
            "Projected economic impact includes plastics recycling/efficiency.",
        ),
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "DCF Primera: Hojas de Ruta de Economía Circular Sectoriales elaboradas "
            "según principios HRNEC, con opinión favorable MINAM; aprobadas por DS a "
            "propuesta del ministerio sectorial.",
            "DCF Primera: Sectoral Circular Economy Roadmaps prepared per HRNEC "
            "principles, with favorable MINAM opinion; approved by supreme decree upon "
            "sector ministry proposal.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "Mandato de desarrollo sectorial. Incluye hoja de ruta sectorial de plásticos "
            "prevista.",
            "Sectoral development mandate. Includes planned plastics sector roadmap.",
        ),
        "comments": bi(
            "Deriva hojas de ruta sectoriales complementarias a legislación de plásticos.",
            "Derives sectoral roadmaps complementing plastics legislation.",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "DCF Segunda: MINAM aprueba por RM el 'Plan Nacional de Acción para el "
            "Monitoreo y Evaluación de la implementación de la HRNEC' en 120 días "
            "calendario.",
            "DCF Segunda: MINAM approves by ministerial resolution the 'National Action "
            "Plan for Monitoring and Evaluation of HRNEC implementation' within 120 "
            "calendar days.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "Plazo: hasta ~jun. 2025 desde vigencia. Instrumento de seguimiento nacional.",
            "Deadline: ~Jun 2025 from entry into force. National follow-up instrument.",
        ),
        "comments": bi(
            "Operacionaliza monitoreo de indicadores de circularidad/plásticos.",
            "Operationalizes circularity/plastics indicator monitoring.",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "DCF Tercera: MINAM aprueba por RM lineamientos para implementación de "
            "indicadores y acciones estratégicas HRNEC, en coordinación con entidades "
            "implementadoras, en 180 días calendario.",
            "DCF Tercera: MINAM approves by ministerial resolution guidelines for "
            "implementation of HRNEC indicators and strategic actions, coordinating with "
            "implementing entities, within 180 calendar days.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "Plazo: hasta ~ago. 2025. Base para reportes de avance Art. 3.",
            "Deadline: ~Aug 2025. Basis for Art. 3 progress reports.",
        ),
        "comments": bi(
            "Lineamientos de implementación sectorial.",
            "Sectoral implementation guidelines.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "DCF Cuarta: INEI brinda soporte en generación de estadísticas e indicadores "
            "para monitoreo y evaluación de la HRNEC.",
            "DCF Cuarta: INEI provides support in generating statistics and indicators "
            "for HRNEC monitoring and evaluation.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "Apoyo estadístico nacional. Sin plazo específico.",
            "National statistical support. No specific deadline.",
        ),
        "comments": bi(
            "Datos para seguimiento de materiales/plásticos en economía circular.",
            "Data for materials/plastics tracking in circular economy.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "Considerando DS: vinculación PNCP 2024-2030 Medida 9.2 'Gestión integral de "
            "residuos sólidos y transición hacia economía circular' — HRNEC articula "
            "esfuerzos sectoriales de circularidad.",
            "DS considering clause: link to PNCP 2024-2030 Measure 9.2 'Integrated solid-"
            "waste management and transition to circular economy' — HRNEC articulates "
            "sectoral circularity efforts.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Vinculación política nacional de competitividad. Complementa DL 1278/PLANRES.",
            "National competitiveness policy linkage. Complements DL 1278/PLANRES.",
        ),
        "comments": bi(
            "Integra residuos sólidos/plásticos en agenda de competitividad nacional.",
            "Integrates solid waste/plastics in national competitiveness agenda.",
        ),
    },
]

COLUMNS = [
    ("A", "policy_name"),
    ("B", "policy_url"),
    ("C", "policy_year"),
    ("D", "policy_objective"),
    ("E", "policy_target"),
    ("F", "policy_target_text"),
    ("G", "policy_type"),
    ("H", "policy_type_justification"),
    ("I", "policy_integration"),
    ("J", "policy_sectors_list"),
    ("K", "policy_circularity"),
    ("L", "policy_lifecycle_phases_list"),
    ("M", "policy_budget"),
    ("N", "policy_budget_text"),
    ("O", "policy_score"),
    ("P", "instrument_type"),
    ("Q", "instrument_lifecycle_stage"),
    ("R", "instrument_description"),
    ("S", "instrument_in_force"),
    ("T", "instrument_implementation"),
    ("U", "instrument_implementation_text"),
    ("V", "instrument_score"),
    ("W", "comments"),
]


def policy_score_row(instrument_type: float) -> float:
    return round(
        (
            POLICY["policy_type"]
            + POLICY["policy_integration"]
            + POLICY["policy_circularity"]
            + POLICY["policy_budget"]
            + instrument_type
        )
        / 5,
        3,
    )


def instrument_score_row(instrument_type: float, implementation: float) -> float:
    return round((instrument_type + implementation) / 2, 3)


def build_excel(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "4P Index Coding"
    ws.append([f"{col} — {name}" for col, name in COLUMNS])
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for inst in INSTRUMENTS:
        ws.append(
            [
                POLICY["policy_name"],
                POLICY["policy_url"],
                POLICY["policy_year"],
                POLICY["policy_objective"],
                POLICY["policy_target"],
                POLICY["policy_target_text"],
                POLICY["policy_type"],
                POLICY["policy_type_justification"],
                POLICY["policy_integration"],
                POLICY["policy_sectors_list"],
                POLICY["policy_circularity"],
                POLICY["policy_lifecycle_phases_list"],
                POLICY["policy_budget"],
                POLICY["policy_budget_text"],
                policy_score_row(inst["instrument_type"]),
                inst["instrument_type"],
                inst["instrument_lifecycle_stage"],
                inst["instrument_description"],
                inst["instrument_in_force"],
                inst["instrument_implementation"],
                inst["instrument_implementation_text"],
                instrument_score_row(inst["instrument_type"], inst["instrument_implementation"]),
                inst["comments"],
            ]
        )

    widths = {
        "A": 40, "B": 38, "C": 8, "D": 44, "E": 8, "F": 44, "G": 8, "H": 40,
        "I": 10, "J": 40, "K": 10, "L": 36, "M": 10, "N": 42, "O": 10, "P": 10,
        "Q": 18, "R": 50, "S": 10, "T": 12, "U": 44, "V": 10, "W": 40,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(path)


def add_article(doc: Document, es: str, en: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(es)
    r.bold = True
    doc.add_paragraph(en).paragraph_format.space_after = Pt(8)


def build_word(path: Path) -> None:
    doc = Document()
    for margin in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(doc.sections[0], margin, Inches(1))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t = title.add_run(f"{POLICY_NAME_EN}\n{POLICY_NAME_ES}")
    t.bold = True
    t.font.size = Pt(14)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(
        "Signed: 24 February 2025 | Published: 25 February 2025 (El Peruano)\n"
        "Countersigned by 10 Executive Branch sectors\n"
        f"Source: {POLICY_URL}"
    )

    doc.add_heading("English Translation of Key Provisions", 1)
    doc.add_paragraph(
        "English translation of principal provisions of Supreme Decree No. 003-2025-MINAM "
        "approving the National Circular Economy Roadmap to 2030."
    )

    sections = [
        (
            "Art. 1 — Approval",
            "Apruébase la 'Hoja de Ruta Nacional de Economía Circular al 2030' como Anexo.",
            "The 'National Circular Economy Roadmap to 2030' is approved as Annex.",
        ),
        (
            "Art. 2 — Scope",
            "Aplicable a entidades del sector público que incorporen economía circular.",
            "Applicable to public-sector entities incorporating circular economy.",
        ),
        (
            "Art. 3 — Monitoring",
            "MINAM monitorea y evalúa indicadores y acciones estratégicas; entidades "
            "reportan avances.",
            "MINAM monitors and evaluates indicators and strategic actions; entities report "
            "progress.",
        ),
        (
            "DCF Primera — Sectoral roadmaps",
            "Hojas de Ruta Sectoriales con opinión favorable MINAM, aprobadas por DS.",
            "Sectoral Roadmaps with favorable MINAM opinion, approved by supreme decree.",
        ),
        (
            "DCF Segunda — Action Plan",
            "Plan Nacional de Acción de monitoreo en 120 días calendario.",
            "National Monitoring Action Plan within 120 calendar days.",
        ),
        (
            "HRNEC Annex — Strategic goals",
            "Transición a producción/consumo sostenible; metas 2030 incl. 2% PBI y 306 000 "
            "empleos; prioriza plásticos y 10 sectores.",
            "Transition to sustainable production/consumption; 2030 targets incl. 2% GDP and "
            "306,000 jobs; prioritizes plastics and 10 sectors.",
        ),
    ]

    for heading, es, en in sections:
        doc.add_heading(heading, 2)
        add_article(doc, es, en)

    doc.add_paragraph()
    note = doc.add_paragraph(
        "Note: DS 003-2025-MINAM approves the HRNEC annex as the strategic multisector "
        "framework guiding Peru's circular transition, explicitly covering plastics value "
        "chains and complementing solid-waste and plastics legislation."
    )
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(path)


def main() -> None:
    excel = OUTPUT_DIR / "Peru_DS_003_2025_MINAM_4P_Index_Coding.xlsx"
    word = OUTPUT_DIR / "Peru_DS_003_2025_MINAM_English_Translation.docx"
    build_excel(excel)
    build_word(word)
    print(f"Created: {excel}")
    print(f"Created: {word}")
    print(f"Instruments: {len(INSTRUMENTS)}")
    print()
    print("Download links:")
    print(
        f"  Excel: {REPO_RAW}/output/peru-ds-003-2025-minam/"
        "Peru_DS_003_2025_MINAM_4P_Index_Coding.xlsx"
    )
    print(
        f"  Word:  {REPO_RAW}/output/peru-ds-003-2025-minam/"
        "Peru_DS_003_2025_MINAM_English_Translation.docx"
    )


if __name__ == "__main__":
    main()
