#!/usr/bin/env python3
"""Generate 4P Index Excel and English translation for Peru HRNEC 2030 document."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

OUTPUT_DIR = Path("/workspace/output/peru-hrnec-2030")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

BRANCH = "cursor/peru-hrnec-2030-4p-index-d350"
REPO_RAW = f"https://github.com/shivrain/Thinkswiss/raw/{BRANCH}"

POLICY_NAME_ES = "Hoja de Ruta Nacional de Economía Circular al 2030 (HRNEC)"
POLICY_NAME_EN = "National Circular Economy Roadmap to 2030 (HRNEC)"
POLICY_URL = (
    "https://cdn.www.gob.pe/uploads/document/file/9548248/"
    "6507211-hoja-de-ruta-nacional-de-economia-circular-al-2030.pdf"
)


def bi(es: str, en: str) -> str:
    return f"{es} / {en}"


POLICY = {
    "policy_name": bi(POLICY_NAME_ES, POLICY_NAME_EN),
    "policy_url": POLICY_URL,
    "policy_year": 2025,
    "policy_objective": bi(
        "Instrumento estratégico nacional de planificación y orientación multisectorial "
        "(aprobado por DS 003-2025-MINAM) para transición del Perú hacia economía circular "
        "al 2030; estandariza definición, principios (NTP-ISO 59004:2024), visión, cuatro "
        "objetivos estratégicos, 12 indicadores (Tabla 1), 35 acciones estratégicas (AE) "
        "y gobernanza MINAM/CNEC; articula HREC-S sectoriales y territorios; vinculado "
        "PNCP 2024-2030 Medida 9.2, PNA 2030 y DL 1278 Art. 15(w).",
        "National multisectoral strategic planning and guidance instrument (approved by "
        "Supreme Decree 003-2025-MINAM) for Peru's circular economy transition by 2030; "
        "standardizes definition, principles (NTP-ISO 59004:2024), vision, four strategic "
        "objectives, 12 indicators (Table 1), 35 strategic actions (AE) and MINAM/CNEC "
        "governance; articulates sectoral HREC-S and territories; linked to PNCP 2024-2030 "
        "Measure 9.2, National Environmental Policy 2030 and DL 1278 Art. 15(w).",
    ),
    "policy_target": 1,
    "policy_target_text": bi(
        "Horizonte de medición 2030 (tres fases: puesta en marcha 2025-2026, desarrollo "
        "2027-2028, consolidación 2029-2030). Metas Tabla 1: VPM 2,29; %PBIec 2,0%; "
        "%EMPec 1,6%; OE1-1 %MCCec 27,0%; OE1-2 ΔINVec 4,1% anual; OE2-2 MUYU 758 "
        "organizaciones; OE3-1 SENec 382 719 personas; OE3-2 DHGRS 0,070; OE4-2 %VAL "
        "45,2%. Visión país en EC trasciende 2030.",
        "2030 measurement horizon (three phases: launch 2025-2026, development 2027-2028, "
        "consolidation 2029-2030). Table 1 targets: VPM 2.29; %GDP from CE 2.0%; %employment "
        "from CE 1.6%; OE1-1 %NDC measures with CE 27.0%; OE1-2 Δpublic CE investment "
        "4.1% annual; OE2-2 MUYU label 758 organizations; OE3-1 SENec 382,719 persons; "
        "OE3-2 DHGRS 0.070; OE4-2 %municipalities valorizing waste 45.2%. Country CE "
        "vision extends beyond 2030.",
    ),
    "policy_type": 0.5,
    "policy_type_justification": bi(
        "Hoja de ruta nacional estratégica aprobada mediante DS 003-2025-MINAM (24 feb. "
        "2025, publicada 25 feb. 2025). Primera edición MINAM mayo 2025. Instrumento de "
        "política pública de planificación (no ley ni reglamento). Clasificado 0.5 "
        "(estrategia nacional / hoja de ruta).",
        "National strategic roadmap approved by Supreme Decree 003-2025-MINAM (24 Feb 2025, "
        "published 25 Feb 2025). First MINAM edition May 2025. Public policy planning "
        "instrument (not law or regulation). Classified 0.5 (national strategy / roadmap).",
    ),
    "policy_integration": 1,
    "policy_sectors_list": bi(
        "transversal (10 sectores Ejecutivo: PCM, MINAM, MEF, PRODUCE, MINEM, VIVIENDA, "
        "MINCETUR, MIDAGRI, MTPE, MINEDU); plásticos/envases, residuos sólidos, construcción, "
        "alimentación, transporte, industria, pesca, agua y saneamiento, agrario",
        "cross-cutting (10 Executive Branch sectors: PCM, MINAM, MEF, PRODUCE, MINEM, "
        "HOUSING, MINCETUR, MIDAGRI, MTPE, MINEDU); plastics/packaging, solid waste, "
        "construction, food, transport, industry, fisheries, water & sanitation, agriculture",
    ),
    "policy_circularity": 1,
    "policy_lifecycle_phases_list": bi(
        "diseño, producción, consumo, recolección, reciclaje, valorización, disposición, "
        "regeneración de ecosistemas (marco NTP-ISO 59004:2024; Cuadro 2)",
        "design, production, consumption, collection, recycling, valorization, disposal, "
        "ecosystem regeneration (NTP-ISO 59004:2024 framework; Table 2)",
    ),
    "policy_budget": 0.5,
    "policy_budget_text": bi(
        "AE 1.4: cartera de inversión pública en EC (~S/2 800 millones/año hacia 2030, "
        "inversiones públicas y público-privadas). AE 4.5: cartera regional. Financiamiento "
        "institucional de entidades implementadoras (DS Art. 4).",
        "AE 1.4: public investment portfolio in CE (~PEN 2,800 million/year toward 2030, "
        "public and public-private investments). AE 4.5: regional portfolio. Institutional "
        "financing of implementing entities (DS Art. 4).",
    ),
}

INSTRUMENTS = [
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "§1.1: definición estandarizada de economía circular como sistema económico que "
            "mantiene/recupera valor de productos, materiales y recursos y regenera "
            "ecosistemas; alineada NTP-ISO 59004:2024.",
            "§1.1: standardized circular economy definition as economic system maintaining/"
            "recovering value of products, materials and resources and regenerating "
            "ecosystems; aligned with NTP-ISO 59004:2024.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Marco conceptual base de la HRNEC. Vigente desde aprobación DS 003-2025.",
            "Conceptual foundation of HRNEC. Effective since DS 003-2025 approval.",
        ),
        "comments": bi(
            "Definición común para sectores y territorios.",
            "Common definition for sectors and territories.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "§2.1 Cuadro 3 — Visión del país en EC: motor principal del desarrollo sostenible; "
            "cadenas de valor circulares innovadoras, competitivas e inclusivas; consumo "
            "sostenible; prevención contaminación y conservación biodiversidad.",
            "§2.1 Table 3 — Country CE vision: main driver of sustainable development; "
            "innovative, competitive, inclusive circular value chains; sustainable "
            "consumption; pollution prevention and biodiversity conservation.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "Visión de largo plazo que trasciende 2030; orienta PEDN y ODS.",
            "Long-term vision extending beyond 2030; guides PEDN and SDGs.",
        ),
        "comments": bi(
            "Aspiración nacional de circularidad.",
            "National circularity aspiration.",
        ),
    },
    {
        "instrument_type": 0.30,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "OE 1 — Gobernanza y políticas circulares: fortalecer marco político, "
            "institucional y regulatorio nacional para gobernanza, inversión, tecnologías "
            "digitales y financiamiento propicio a la EC.",
            "OE 1 — Governance and circular policies: strengthen national political, "
            "institutional and regulatory framework for governance, investment, digital "
            "technologies and CE-conducive financing.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "11 AE bajo OE1 (1.1-1.11). Indicadores OE1-1 y OE1-2.",
            "11 AEs under OE1 (1.1-1.11). Indicators OE1-1 and OE1-2.",
        ),
        "comments": bi(
            "Pilar de gobernanza y HREC-S.",
            "Governance and HREC-S pillar.",
        ),
    },
    {
        "instrument_type": 0.30,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "OE 2 — Innovación y negocios circulares: impulsar innovación, digitalización, "
            "escalamiento de modelos de negocio circulares, mercados de materia prima "
            "secundaria y eco/bionegocios.",
            "OE 2 — Innovation and circular business: drive innovation, digitalization, "
            "scaling of circular business models, secondary raw material markets and "
            "eco/biobusinesses.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "7 AE bajo OE2 (2.1-2.7). Indicadores OE2-1 y OE2-2 (distintivo MUYU).",
            "7 AEs under OE2 (2.1-2.7). Indicators OE2-1 and OE2-2 (MUYU label).",
        ),
        "comments": bi(
            "Mercados secundarios y financiamiento verde.",
            "Secondary markets and green finance.",
        ),
    },
    {
        "instrument_type": 0.30,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "OE 3 — Consumo sostenible y cultura circular: sensibilización, competencias, "
            "prácticas responsables, consumo justo y estilos de vida sostenibles.",
            "OE 3 — Sustainable consumption and circular culture: awareness, skills, "
            "responsible practices, fair consumption and sustainable lifestyles.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "9 AE bajo OE3 (3.1-3.9). Indicadores OE3-1 (SENec) y OE3-2 (DHGRS).",
            "9 AEs under OE3 (3.1-3.9). Indicators OE3-1 (SENec) and OE3-2 (DHGRS).",
        ),
        "comments": bi(
            "Cultura circular y reducción residuos.",
            "Circular culture and waste reduction.",
        ),
    },
    {
        "instrument_type": 0.30,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "OE 4 — Territorios y ciudades circulares: integrar EC en políticas y prácticas "
            "regionales/locales; sostenibilidad de cadenas de valor priorizadas y "
            "articulación de actores en el territorio.",
            "OE 4 — Circular territories and cities: integrate CE in regional/local policies "
            "and practices; sustainability of prioritized value chains and territorial "
            "actor coordination.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "8 AE bajo OE4 (4.1-4.8). Indicadores OE4-1, OE4-2 y OE4-3.",
            "8 AEs under OE4 (4.1-4.8). Indicators OE4-1, OE4-2 and OE4-3.",
        ),
        "comments": bi(
            "Gestión municipal de residuos y corredores.",
            "Municipal waste management and corridors.",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "Tabla 1 — Indicadores de impacto I1-I3: variación productividad material (VPM "
            "1,2→2,29); % aporte EC al PBI (2,0%); % contribución EC al empleo (1,6%).",
            "Table 1 — Impact indicators I1-I3: material productivity variation (VPM "
            "1.2→2.29); % CE contribution to GDP (2.0%); % CE contribution to employment "
            "(1.6%).",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "Monitoreo MINAM; cálculo INEI/MINAM-MTPE. Fichas técnicas Anexo 3.",
            "MINAM monitoring; INEI/MINAM-MTPE calculation. Technical sheets Annex 3.",
        ),
        "comments": bi(
            "Indicadores macro de transición circular.",
            "Macro indicators of circular transition.",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "Tabla 1 OE1-1: % medidas NDC que incorporan EC (%MCCec) — línea base 0% (2020), "
            "meta 27,0% al 2030; sectores con NDC y HREC-S reportan.",
            "Table 1 OE1-1: % NDC measures incorporating CE (%MCCec) — baseline 0% (2020), "
            "target 27.0% by 2030; NDC and HREC-S sectors report.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "Vincula clima (NDC) y circularidad. Responsable cálculo MINAM.",
            "Links climate (NDC) and circularity. MINAM responsible for calculation.",
        ),
        "comments": bi(
            "Integración EC en compromisos climáticos.",
            "CE integration in climate commitments.",
        ),
    },
    {
        "instrument_type": 0.50,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "Tabla 1 OE3-2: diferencial de contribución HRNEC en generación de residuos "
            "sólidos (DHGRS) — meta 0,070 al 2030; gobiernos locales reportan.",
            "Table 1 OE3-2: HRNEC contribution differential in solid waste generation "
            "(DHGRS) — target 0.070 by 2030; local governments report.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "Indicador clave de reducción generación RS. Línea base 2022.",
            "Key indicator for MSW generation reduction. 2022 baseline.",
        ),
        "comments": bi(
            "Relevante plásticos y residuos municipales.",
            "Relevant for plastics and municipal waste.",
        ),
    },
    {
        "instrument_type": 0.50,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "Tabla 1 OE4-2: % municipalidades que valorizan RRSS (%VAL) — línea base 33,3% "
            "(2023), meta 45,2% al 2030; MINAM DGGRS responsable.",
            "Table 1 OE4-2: % municipalities valorizing MSW (%VAL) — baseline 33.3% (2023), "
            "target 45.2% by 2030; MINAM DGGRS responsible.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "Meta valorización residuos a nivel local. Réplica territorial (Cuadro 4).",
            "Local waste valorization target. Territorial replication (Table 4).",
        ),
        "comments": bi(
            "Reciclaje y valorización plásticos municipales.",
            "Recycling and valorization of municipal plastics.",
        ),
    },
    {
        "instrument_type": 0.25,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "Fases de implementación HRNEC: puesta en marcha (2025-2026), desarrollo "
            "(2027-2028), consolidación (2029-2030); AE con plazos corto (C), mediano (M) "
            "y largo (L).",
            "HRNEC implementation phases: launch (2025-2026), development (2027-2028), "
            "consolidation (2029-2030); AEs with short (C), medium (M) and long (L) "
            "timelines.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Calendario operativo de 35 AE. Monitoreo continuo MINAM.",
            "Operational schedule for 35 AEs. Continuous MINAM monitoring.",
        ),
        "comments": bi(
            "Secuenciación de la transición circular.",
            "Sequencing of circular transition.",
        ),
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "AE 1.3: diseñar e implementar sistema de monitoreo, evaluación y aprendizaje "
            "(MEA) de la HRNEC — indicadores impacto, resultado y proceso; sistema digital "
            "público; interoperabilidad INEI/SGTD.",
            "AE 1.3: design and implement HRNEC monitoring, evaluation and learning (MEA) "
            "system — impact, outcome and process indicators; public digital system; "
            "INEI/SGTD interoperability.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "En desarrollo. Incluye indicadores PRODUCE (materiales circulares, residuos "
            "no municipales valorizados).",
            "Under development. Includes PRODUCE indicators (circular materials, non-municipal "
            "waste valorized).",
        ),
        "comments": bi(
            "Sistema de seguimiento de políticas de plásticos/residuos.",
            "Tracking system for plastics/waste policies.",
        ),
    },
    {
        "instrument_type": 0.50,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "AE 1.11: aprobar e implementar guía para diseño y elaboración de HREC-S; "
            "alineación con HRNEC, Anexo 4 lineamientos; sectores con mandato PNCP 203-2024-EF.",
            "AE 1.11: approve and implement guide for design and elaboration of sectoral "
            "CE roadmaps (HREC-S); alignment with HRNEC, Annex 4 guidelines; sectors "
            "mandated under PNCP DS 203-2024-EF.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "HREC-S obligatorias sectores productivos PNCP. Guía vía RM MINAM.",
            "Mandatory HREC-S for PNCP productive sectors. Guide via MINAM ministerial "
            "resolution.",
        ),
        "comments": bi(
            "HREC-S plásticos/industria vinculadas a HRNEC.",
            "Plastics/industry HREC-S linked to HRNEC.",
        ),
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "AE 4.3: promover corredores de reciclaje integrando EC — recolección selectiva, "
            "transporte, acondicionamiento y valorización RS aprovechables; Programa Recicla; "
            "excluye incineración/gasificación como EC.",
            "AE 4.3: promote recycling corridors integrating CE — selective collection, "
            "transport, conditioning and valorization of recoverable MSW; Recicla Programme; "
            "excludes incineration/gasification as CE actions.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "Réplica territorial (Cuadro 4). Indicadores: n.° corredores, t reciclables/mes.",
            "Territorial replication (Table 4). Indicators: no. corridors, t recyclables/month.",
        ),
        "comments": bi(
            "Instrumento clave para plásticos y RS municipales.",
            "Key instrument for plastics and municipal solid waste.",
        ),
    },
    {
        "instrument_type": 0.50,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "AE 3.2: incorporar EC en compras públicas sostenibles — criterios en Bases "
            "Estándar, Plan Nacional Compras Públicas Sostenibles (CM-CPS MINAM).",
            "AE 3.2: incorporate CE in sustainable public procurement — criteria in Standard "
            "Bidding Documents, National Sustainable Public Procurement Action Plan "
            "(MINAM multisectoral commission).",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "8 fichas homologación sostenibilidad existentes; expansión criterios EC.",
            "8 existing sustainability homologation sheets; expansion of CE criteria.",
        ),
        "comments": bi(
            "Demanda pública para productos circulares/plásticos reciclados.",
            "Public demand for circular/recycled plastic products.",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "§4 Gobernanza: MINAM lidera despliegue HRNEC; entidades implementadoras "
            "ejecutan AE; CNEC (Coalición Nacional EC 'Perú País Circular') como espacio "
            "diálogo público-privado (PNCP Medida 9.2).",
            "§4 Governance: MINAM leads HRNEC deployment; implementing entities execute "
            "AEs; CNEC (National CE Coalition 'Peru Circular Country') as public-private "
            "dialogue space (PNCP Measure 9.2).",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "CNEC vía AE 1.6. Coordinación multisectorial activa.",
            "CNEC via AE 1.6. Active multisectoral coordination.",
        ),
        "comments": bi(
            "Gobernanza participativa de transición circular.",
            "Participatory governance of circular transition.",
        ),
    },
    {
        "instrument_type": 0.25,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "Anexo 1 — Definición 'plástico de un solo uso': bien de plástico para un solo "
            "uso y corta vida útil, o cuya composición dificulta biodegradabilidad y/o "
            "valorización (también descartable).",
            "Annex 1 — Definition of 'single-use plastic': plastic good designed for single "
            "use and short lifespan, or whose composition hinders biodegradability and/or "
            "valorization (also disposable).",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Glosario estandarizado HRNEC. Referencia para políticas sectoriales.",
            "Standardized HRNEC glossary. Reference for sector policies.",
        ),
        "comments": bi(
            "Definición SUP para alineación normativa plásticos.",
            "SUP definition for plastics regulatory alignment.",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Planning",
        "instrument_description": bi(
            "§2.4 / Anexo 4: HREC-S obligatorias sectores PNCP 203-2024-EF; réplica/adaptación "
            "indicadores y AE sectoriales (Cuadro 4); sectores con HREC aprobadas: industria, "
            "pesca, agua potable y saneamiento, agrario.",
            "§2.4 / Annex 4: mandatory HREC-S for PNCP 203-2024-EF sectors; sectoral "
            "indicator and AE replication/adaptation (Table 4); sectors with approved HREC-S: "
            "industry, fisheries, drinking water & sanitation, agriculture.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "HREC-S en aprobación/actualización. Voluntarias otros sectores.",
            "HREC-S under approval/update. Voluntary for other sectors.",
        ),
        "comments": bi(
            "Despliegue sectorial incluye cadenas plásticas.",
            "Sectoral rollout includes plastics chains.",
        ),
    },
]

COLUMNS = [
    ("A", "policy_name"),
    ("B", "policy_url"),
    ("C", "policy_year"),
    ("D", "policy_objective"),
    ("E", "policy_target"),
    ("F", "policy_target_text"),
    ("G", "policy_type"),
    ("H", "policy_type_justification"),
    ("I", "policy_integration"),
    ("J", "policy_sectors_list"),
    ("K", "policy_circularity"),
    ("L", "policy_lifecycle_phases_list"),
    ("M", "policy_budget"),
    ("N", "policy_budget_text"),
    ("O", "policy_score"),
    ("P", "instrument_type"),
    ("Q", "instrument_lifecycle_stage"),
    ("R", "instrument_description"),
    ("S", "instrument_in_force"),
    ("T", "instrument_implementation"),
    ("U", "instrument_implementation_text"),
    ("V", "instrument_score"),
    ("W", "comments"),
]


def policy_score_row(instrument_type: float) -> float:
    return round(
        (
            POLICY["policy_type"]
            + POLICY["policy_target"]
            + POLICY["policy_integration"]
            + POLICY["policy_circularity"]
            + POLICY["policy_budget"]
            + instrument_type
        )
        / 6,
        3,
    )


def instrument_score_row(instrument_type: float, implementation: float) -> float:
    return round((instrument_type + implementation) / 2, 3)


def build_excel(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "4P Index Coding"
    ws.append([f"{col} — {name}" for col, name in COLUMNS])
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for inst in INSTRUMENTS:
        ws.append(
            [
                POLICY["policy_name"],
                POLICY["policy_url"],
                POLICY["policy_year"],
                POLICY["policy_objective"],
                POLICY["policy_target"],
                POLICY["policy_target_text"],
                POLICY["policy_type"],
                POLICY["policy_type_justification"],
                POLICY["policy_integration"],
                POLICY["policy_sectors_list"],
                POLICY["policy_circularity"],
                POLICY["policy_lifecycle_phases_list"],
                POLICY["policy_budget"],
                POLICY["policy_budget_text"],
                policy_score_row(inst["instrument_type"]),
                inst["instrument_type"],
                inst["instrument_lifecycle_stage"],
                inst["instrument_description"],
                inst["instrument_in_force"],
                inst["instrument_implementation"],
                inst["instrument_implementation_text"],
                instrument_score_row(inst["instrument_type"], inst["instrument_implementation"]),
                inst["comments"],
            ]
        )

    widths = {
        "A": 40, "B": 38, "C": 8, "D": 44, "E": 8, "F": 44, "G": 8, "H": 40,
        "I": 10, "J": 40, "K": 10, "L": 36, "M": 10, "N": 42, "O": 10, "P": 10,
        "Q": 18, "R": 50, "S": 10, "T": 12, "U": 44, "V": 10, "W": 40,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(path)


def add_article(doc: Document, es: str, en: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(es)
    r.bold = True
    doc.add_paragraph(en).paragraph_format.space_after = Pt(8)


def build_word(path: Path) -> None:
    doc = Document()
    for margin in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(doc.sections[0], margin, Inches(1))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t = title.add_run(f"{POLICY_NAME_EN}\n{POLICY_NAME_ES}")
    t.bold = True
    t.font.size = Pt(14)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(
        "Approved by DS 003-2025-MINAM (24 Feb 2025) | First edition: May 2025 (MINAM)\n"
        f"Source: {POLICY_URL}"
    )

    doc.add_heading("English Translation of Key Provisions", 1)
    doc.add_paragraph(
        "English translation of principal provisions of the National Circular Economy "
        "Roadmap to 2030 (HRNEC) document."
    )

    sections = [
        (
            "§1.1 — Circular economy definition",
            "Sistema económico que mantiene/recupera valor de productos, materiales y "
            "recursos y regenera ecosistemas (NTP-ISO 59004:2024).",
            "Economic system maintaining/recovering value of products, materials and "
            "resources and regenerating ecosystems (NTP-ISO 59004:2024).",
        ),
        (
            "§2.1 — Country vision",
            "La economía circular es el principal motor del desarrollo sostenible del Perú.",
            "Circular economy is the main driver of Peru's sustainable development.",
        ),
        (
            "OE 1-4 — Strategic objectives",
            "Cuatro pilares: gobernanza, innovación/negocios, consumo/cultura, territorios/ciudades.",
            "Four pillars: governance, innovation/business, consumption/culture, territories/cities.",
        ),
        (
            "Tabla 1 — 2030 targets",
            "12 indicadores: VPM 2,29; %PBIec 2,0%; %VAL 45,2%; DHGRS 0,070; entre otros.",
            "12 indicators: VPM 2.29; %GDP from CE 2.0%; %VAL 45.2%; DHGRS 0.070; among others.",
        ),
        (
            "AE 4.3 — Recycling corridors",
            "Corredores de reciclaje para RS aprovechables; Programa Recicla; excluye incineración.",
            "Recycling corridors for recoverable MSW; Recicla Programme; excludes incineration.",
        ),
        (
            "AE 1.11 — Sectoral roadmaps",
            "Guía para HREC-S alineadas a HRNEC; sectores obligados PNCP.",
            "Guide for HREC-S aligned with HRNEC; PNCP-mandated sectors.",
        ),
        (
            "§4 — Governance",
            "MINAM lidera implementación; CNEC como espacio de diálogo (PNCP Medida 9.2).",
            "MINAM leads implementation; CNEC as dialogue space (PNCP Measure 9.2).",
        ),
        (
            "Anexo 1 — Single-use plastic",
            "Plástico de un solo uso: bien diseñado para un solo uso y corta vida útil.",
            "Single-use plastic: good designed for single use and short lifespan.",
        ),
    ]

    for heading, es, en in sections:
        doc.add_heading(heading, 2)
        add_article(doc, es, en)

    doc.add_paragraph()
    note = doc.add_paragraph(
        "Note: HRNEC is the strategic document approved as Annex to DS 003-2025-MINAM. "
        "It complements sectoral plastics and solid-waste legislation through HREC-S, "
        "recycling corridors (AE 4.3) and municipal valorization targets (OE4-2)."
    )
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(path)


def main() -> None:
    excel = OUTPUT_DIR / "Peru_HRNEC_2030_4P_Index_Coding.xlsx"
    word = OUTPUT_DIR / "Peru_HRNEC_2030_English_Translation.docx"
    build_excel(excel)
    build_word(word)
    print(f"Created: {excel}")
    print(f"Created: {word}")
    print(f"Instruments: {len(INSTRUMENTS)}")
    print()
    print("Download links:")
    print(
        f"  Excel: {REPO_RAW}/output/peru-hrnec-2030/"
        "Peru_HRNEC_2030_4P_Index_Coding.xlsx"
    )
    print(
        f"  Word:  {REPO_RAW}/output/peru-hrnec-2030/"
        "Peru_HRNEC_2030_English_Translation.docx"
    )


if __name__ == "__main__":
    main()
