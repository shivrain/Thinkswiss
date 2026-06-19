#!/usr/bin/env python3
"""Generate 4P Index coding and English translation for Mozambique Decree 51/2024."""

import os
from datetime import date

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT_DIR = "/workspace/output/mozambique"
os.makedirs(OUTPUT_DIR, exist_ok=True)

POLICY_FIELDS = {
    "policy_name": (
        "Regulamento sobre a Actividade de Fiscalização Ambiental (Decreto n.º 51/2024)"
    ),
    "policy_url": "https://faolex.fao.org/docs/pdf/moz228476.pdf",
    "policy_year": 2024,
    "policy_objective": (
        "Establish national mechanisms for environmental inspection of public and private activities "
        "that may harm the environment; verify compliance with environmental protection and quality "
        "norms (including waste, EIA mitigation and monitoring obligations) and enable administrative "
        "and coercive sanctions for violations of substantive environmental law."
    ),
    "policy_target": 0,
    "policy_target_text": "",
    "policy_type": 0.75,
    "policy_type_justification": (
        "Executive regulation (Regulamento) approved by Council of Ministers Decree n.º 51/2024 "
        "under Article 33 of the Environment Law (Law n.º 20/97) — sub-legislative instrument, "
        "not parliamentary legislation. Revokes Decree n.º 11/2006."
    ),
    "policy_integration": 0.75,
    "policy_sectors_list": (
        "waste management, industry, municipalities, fisheries, tourism, water, agriculture, packaging"
    ),
    "policy_circularity": 0.50,
    "policy_lifecycle_phases_list": (
        "production, consumption, disposal, environmental leakage"
    ),
    "policy_budget": 0.50,
    "policy_budget_text": (
        "Art. 21: 15-day voluntary fine payment window; Art. 22: fine revenues split 40% State, "
        "60% inspecting entity, channelled to Single Treasury Account."
    ),
    "policy_score": "[auto]",
}

# Four consolidated instruments — horizontal inspection framework; not overcoded per article.
INSTRUMENTS = [
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Environmental leakage",
        "instrument_description": (
            "Art. 2–5, 3, 12, 14: Defines environmental inspection (compliance verification, audit/monitoring "
            "oversight, EIA mitigation checks); Ministry conducts ordinary and extraordinary inspections "
            "nationwide with unannounced access, sample collection and imaging; may issue infringement reports, "
            "refer cases to the Public Prosecutor for embargo, destruction of works, equipment seizure or "
            "activity cancellation, and report criminal offences."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            'Art. 5(2): "A Fiscalização Ambiental não carece de autorização da entidade fiscalizada"; '
            'Art. 3(c)-(d): coercive referral to Ministério Público; Art. 2(c): EIA mitigation verification; '
            'Art. 4: ordinary vs extraordinary inspection types.'
        ),
        "instrument_score": "[auto]",
        "comments": (
            "Horizontal enforcement multiplier for substantive plastic/waste laws (e.g. Decrees 45/2006, 94/2014). "
            "Not coded separately: principles (Art. 6), conflicts of interest (Art. 7, 17), inspector ID (Art. 15), "
            "secrecy duty (Art. 16)."
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Waste management",
        "instrument_description": (
            "Art. 8–11: Environmental inspectors issue infringement notices (auto de notícia) for violations "
            "of environmental norms; may notify entities for clarification within 10 days; where irregularities "
            "can be remedied, a correction period of up to 30 days may be granted before fines apply."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            'Art. 8(1): auto de notícia on detecting "infração, transgressão ou irregularidade"; '
            'Art. 11(1): correction period "de no máximo 30 dias"; Art. 9: notice content and voluntary '
            'payment pathway; Art. 10: refusal-to-sign procedure.'
        ),
        "instrument_score": "[auto]",
        "comments": (
            "Procedural instrument enabling enforcement of underlying waste/plastic offences. "
            "Fines themselves are set by the substantive law infringed, not this regulation."
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Waste management",
        "instrument_description": (
            "Art. 18, 20–22: Inspectors apply fines and accessory penalties per the infringed legal instrument; "
            "graduated by offence severity and environmental damage; accessory penalties may include activity "
            "suspension and equipment seizure until legal conformity; 15-day voluntary payment; coercive "
            "collection if unpaid; revenues allocated 40% State / 60% inspecting entity."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            'Art. 18(2): accessory penalties include "suspensão da actividade" and "apreensão de equipamentos"; '
            'Art. 20: gradation by "magnitude e consequências... do dano sobre o ambiente"; '
            'Art. 21(1): 15-day voluntary payment; Art. 22: revenue split.'
        ),
        "instrument_score": "[auto]",
        "comments": (
            "Sanctions instrument — applies fines from substantive laws (marine litter, solid waste, effluent "
            "rules) rather than setting new plastic-specific penalties."
        ),
    },
    {
        "instrument_type": 0.80,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": (
            "Art. 13, 19: Inspected entities must provide and keep on-site all relevant environmental "
            "documentation — environmental licence, EIA and environmental management plan, audit and "
            "monitoring reports, exploration plans, territorial planning instruments — under administrative "
            "liability for non-compliance or refusal to cooperate."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            'Art. 13(1): mandatory documentation including "Estudo de Impacto Ambiental", "Plano de Gestão '
            'Ambiental", "Relatório de Monitorização Ambiental"; Art. 13(3): refusal triggers legal referral; '
            'Art. 19: on-site availability obligation.'
        ),
        "instrument_score": "[auto]",
        "comments": (
            "Documentation duty supports verification of waste-management and pollution-control compliance "
            "during inspections. Art. 13 and Art. 19 are substantively overlapping — coded once."
        ),
    },
]

COLUMNS = [
    "policy_name", "policy_url", "policy_year", "policy_objective", "policy_target",
    "policy_target_text", "policy_type", "policy_type_justification", "policy_integration",
    "policy_sectors_list", "policy_circularity", "policy_lifecycle_phases_list",
    "policy_budget", "policy_budget_text", "policy_score",
    "instrument_type", "instrument_lifecycle_stage", "instrument_description",
    "instrument_in_force", "instrument_implementation", "instrument_implementation_text",
    "instrument_score", "comments",
]

COLUMN_LABELS = {
    "policy_name": "A — policy_name",
    "policy_url": "B — policy_url",
    "policy_year": "C — policy_year",
    "policy_objective": "D — policy_objective",
    "policy_target": "E — policy_target",
    "policy_target_text": "F — policy_target_text",
    "policy_type": "G — policy_type",
    "policy_type_justification": "H — policy_type_justification",
    "policy_integration": "I — policy_integration",
    "policy_sectors_list": "J — policy_sectors_list",
    "policy_circularity": "K — policy_circularity",
    "policy_lifecycle_phases_list": "L — policy_lifecycle_phases_list",
    "policy_budget": "M — policy_budget",
    "policy_budget_text": "N — policy_budget_text",
    "policy_score": "O — policy_score [auto]",
    "instrument_type": "P — instrument_type",
    "instrument_lifecycle_stage": "Q — instrument_lifecycle_stage",
    "instrument_description": "R — instrument_description",
    "instrument_in_force": "S — instrument_in_force",
    "instrument_implementation": "T — instrument_implementation",
    "instrument_implementation_text": "U — instrument_implementation_text",
    "instrument_score": "V — instrument_score [auto]",
    "comments": "W — comments",
}


def build_dataframe():
    return pd.DataFrame([{**POLICY_FIELDS, **i} for i in INSTRUMENTS], columns=COLUMNS)


def export_csv(df, basename):
    path = os.path.join(OUTPUT_DIR, f"{basename}.csv")
    df.rename(columns=COLUMN_LABELS).to_csv(path, index=False, encoding="utf-8-sig")
    return path


def export_xlsx(df, basename):
    path = os.path.join(OUTPUT_DIR, f"{basename}.xlsx")
    labelled = df.rename(columns=COLUMN_LABELS)
    labelled.to_excel(path, index=False, sheet_name="4P Index Coding")
    wb = load_workbook(path)
    ws = wb.active
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = min(45, max(12, len(str(ws.cell(1, col).value or "")) * 0.9))
    ws.freeze_panes = "A2"
    wb.save(path)
    return path


def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_para(doc, text):
    doc.add_paragraph(text)


def create_translation_doc():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("Regulation on Environmental Inspection", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("Decree No. 51/2024 of 17 July")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(
        "English translation of: Regulamento sobre a Actividade de Fiscalização Ambiental\n"
    ).italic = True
    meta.add_run(f"Translated for 4P Index policy analysis — {date.today().isoformat()}\n")
    meta.add_run("Source: https://faolex.fao.org/docs/pdf/moz228476.pdf")

    sections = [
        ("Decree No. 51/2024", (
            "The Council of Ministers decrees:\n"
            "Article 1. The Regulation on Environmental Inspection Activity is approved as an annex "
            "forming an integral part of this Decree.\n"
            "Article 2. The Minister responsible for the environment shall approve complementary norms "
            "for operationalising this Decree.\n"
            "Article 3. Decree No. 11/2006 of 15 June is repealed.\n"
            "Article 4. This Decree enters into force on the date of its publication.\n"
            "Approved by the Council of Ministers on 21 May 2024."
        )),
        ("Regulation — Chapter I. General provisions", (
            "Article 1 (Object). To regulate environmental inspection activity for compliance with "
            "national environmental protection and quality norms.\n\n"
            "Article 2 (Definition). Environmental Inspection includes:\n"
            "a) verifying conformity of any activity with environmental protection and quality norms;\n"
            "b) supervising environmental audit and monitoring actions;\n"
            "c) verifying compliance with EIA mitigation measures; and\n"
            "d) supervising land use and territorial planning/resettlement instruments.\n\n"
            "Article 3 (Competencies). The Ministry of Environment shall:\n"
            "a) conduct environmental inspections nationwide;\n"
            "b) issue infringement reports for sanctions;\n"
            "c) trigger legal mechanisms before the Public Prosecutor to embargo, order destruction of "
            "works, seize equipment or cancel degrading activities; and\n"
            "d) refer criminal environmental offences to the Public Prosecutor."
        )),
        ("Chapter II. Inspection action", (
            "Article 4. Inspection may be ordinary (under the activity plan) or extraordinary (targeting "
            "specific activities threatening environmental balance).\n\n"
            "Article 5.\n"
            "1. Inspection is exercised by credentialed Environmental Inspectors or Environmental Technicians.\n"
            "2. Environmental inspection does not require authorisation from the inspected entity.\n"
            "3. Inspectors must inform the responsible person, with access to premises, documentation, "
            "sample collection and image recording.\n"
            "4. Upon completion, inspectors communicate findings via a signed environmental inspection form."
        )),
        ("Chapter III. Principles, guarantees and impediments", (
            "Article 6. Principles: independence, impartiality, legality, transparency.\n\n"
            "Article 7. State guarantees resources and security; conflict-of-interest rules apply.\n\n"
            "Article 8. On detecting any environmental infringement, inspectors issue an infringement notice "
            "(auto de notícia) using an approved form; may issue a 10-day clarification notice.\n\n"
            "Article 9. Infringement notice requirements (offender ID, facts, legal provision infringed, "
            "penalty, aggravating/attenuating circumstances, voluntary payment/defence deadline).\n\n"
            "Article 10. Procedure when the offender refuses to sign.\n\n"
            "Article 11. Correctable irregularities: up to 30-day compliance period before fine application."
        )),
        ("Chapter IV. Rights, duties and prerogatives", (
            "Article 12. Inspector rights: ID card, free access, adequate facilities, inter-agency "
            "coordination, police/Public Prosecutor assistance.\n\n"
            "Article 13. Inspected entities must provide environmental documentation (licence, DUAT, "
            "EIA, management plan, audit/monitoring reports, exploration plans, territorial instruments, "
            "resettlement plan) and cooperate fully.\n\n"
            "Article 14. Inspectors may request civil/police authority assistance.\n\n"
            "Articles 15–17. Credentials, secrecy duty, incompatibilities.\n\n"
            "Article 18.\n"
            "1. Fines follow sanctions in the underlying legal instruments.\n"
            "2. Accessory penalties may include activity suspension and equipment seizure until conformity.\n\n"
            "Article 19. Entities must keep required environmental information on-site.\n\n"
            "Article 20. Fine gradation based on severity and environmental damage magnitude.\n\n"
            "Article 21. 15-day voluntary fine payment or defence right.\n\n"
            "Article 22. Fine revenues: 40% State, 60% inspecting entity (Single Treasury Account)."
        )),
        ("Note on plastic pollution relevance", (
            "This regulation does not mention plastics explicitly. It is a horizontal enforcement framework "
            "that enables inspection and sanctioning of compliance with substantive environmental laws "
            "including marine litter rules (Decree 45/2006), urban solid waste management (Decree 94/2014), "
            "and effluent standards (Decree 18/2004). Replaces Decree 11/2006."
        )),
    ]

    for heading, body in sections:
        add_heading(doc, heading, level=1)
        add_para(doc, body)

    doc.add_page_break()
    add_heading(doc, "4P Index Coding Summary", level=1)
    add_para(
        doc,
        f"Policy: {POLICY_FIELDS['policy_name']}\n"
        f"Instruments coded: {len(INSTRUMENTS)} (consolidated — horizontal inspection framework)\n"
        f"See: 4P_Index_Mozambique_Decree_51_2024.xlsx",
    )

    path = os.path.join(OUTPUT_DIR, "Mozambique_Decree_51_2024_English_Translation.docx")
    doc.save(path)
    return path


def create_index_html(basename):
    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <title>Mozambique Decree 51/2024 — 4P Index</title>
  <style>
    body {{ font-family: system-ui, sans-serif; max-width: 800px; margin: 2rem auto; padding: 0 1rem; }}
    .card {{ border: 1px solid #ddd; border-radius: 8px; padding: 1rem; margin: 1rem 0; }}
    a.download {{ display: inline-block; background: #1F4E79; color: #fff; padding: 0.6rem 1.2rem;
      border-radius: 4px; text-decoration: none; margin-top: 0.5rem; }}
  </style>
</head>
<body>
  <h1>4P Index — Decree 51/2024 (Environmental Inspection)</h1>
  <p>Generated {date.today().isoformat()}. Four consolidated instrument rows.</p>
  <div class="card">
    <a class="download" href="{basename}.xlsx" download>Download Excel</a><br><br>
    <a class="download" href="{basename}.csv" download>Download CSV</a><br><br>
    <a class="download" href="Mozambique_Decree_51_2024_English_Translation.docx" download>Download English translation</a>
  </div>
</body>
</html>"""
    path = os.path.join(OUTPUT_DIR, "index_decree_51_2024.html")
    with open(path, "w", encoding="utf-8") as f:
        f.write(html)
    return path


def main():
    basename = "4P_Index_Mozambique_Decree_51_2024"
    df = build_dataframe()
    paths = [
        export_csv(df, basename),
        export_xlsx(df, basename),
        create_translation_doc(),
        create_index_html(basename),
    ]
    print("Generated:")
    for p in paths:
        print(f"  {p}")


if __name__ == "__main__":
    main()
