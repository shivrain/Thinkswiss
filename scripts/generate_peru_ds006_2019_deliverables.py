#!/usr/bin/env python3
"""Generate 4P Index Excel and English translation for Peru DS 006-2019-MINAM."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

OUTPUT_DIR = Path("/workspace/output/peru-ds-006-2019-minam")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

BRANCH = "cursor/peru-ds-006-2019-4p-index-d350"
REPO_RAW = f"https://github.com/shivrain/Thinkswiss/raw/{BRANCH}"

POLICY_NAME_ES = (
    "Decreto Supremo N.° 006-2019-MINAM – Reglamento de la Ley N.° 30884"
)
POLICY_NAME_EN = (
    "Supreme Decree No. 006-2019-MINAM – Regulation of Law No. 30884 "
    "(Single-Use Plastic Law)"
)
POLICY_URL = (
    "https://www.gob.pe/institucion/minam/normas-legales/290809-006-2019-minam"
)


def bi(es: str, en: str) -> str:
    return f"{es} / {en}"


POLICY = {
    "policy_name": bi(POLICY_NAME_ES, POLICY_NAME_EN),
    "policy_url": POLICY_URL,
    "policy_year": 2019,
    "policy_objective": bi(
        "Decreto supremo que aprueba el reglamento de la Ley N.° 30884, definiendo los bienes "
        "de plástico regulados (bolsas, sorbetes, envases de poliestireno expandido, vajilla "
        "plástica, botellas PET e insumos), plazos de cumplimiento, registro de "
        "fabricantes/importadores/distribuidores, educación y gestión integral de residuos "
        "plásticos, criterios para ampliar controles a otros bienes poliméricos, régimen "
        "compartido de fiscalización y sanción (OEFA, PRODUCE, INDECOPI, SERNANP, Cultura, "
        "municipios) y anexo de infracciones.",
        "Supreme decree approving the regulation of Law No. 30884, defining regulated plastic "
        "goods (bags, straws, expanded-polystyrene containers, plastic tableware, PET bottles "
        "and inputs), compliance timelines, manufacturer/importer/distributor registry, "
        "education and integrated plastic-waste management, criteria for extending controls "
        "to other polymer goods, shared oversight and sanction regime (OEFA, PRODUCE, "
        "INDECOPI, SERNANP, Culture, municipalities) and infractions annex.",
    ),
    "policy_target": 1,
    "policy_target_text": bi(
        "Implementa plazos Ley 30884: Art. 2.1 reemplazo bolsas 36 meses; Art. 3.1 "
        "prohibiciones 120 días en áreas protegidas/playas/administración; Art. 3.2 a 12 "
        "meses bolsas <900 cm²/<50 µm y sorbetes; Art. 3.3 a 36 meses bolsas no "
        "reutilizables, vajilla y tecnopor; Art. 10 PET-PCR 15% en 3 años. DCF Séptima: "
        "fiscalización OEFA bolsas/sorbetes desde 20 dic. 2019; tecnopor desde 20 dic. 2021. "
        "Art. 25.4: recolección selectiva mínimo los miércoles (Día del Reciclaje del "
        "Plástico).",
        "Implements Law 30884 timelines: Art. 2.1 bag replacement 36 months; Art. 3.1 "
        "prohibitions 120 days in protected areas/beaches/state admin; Art. 3.2 at 12 months "
        "bags <900 cm²/<50 µm and straws; Art. 3.3 at 36 months non-reusable bags, tableware "
        "and styrofoam; Art. 10 15% PET-PCR within 3 years. Seventh FCP: OEFA enforcement "
        "bags/straws from 20 Dec 2019; styrofoam from 20 Dec 2021. Art. 25.4: selective "
        "collection at minimum on Wednesdays (Plastic Recycling Day).",
    ),
    "policy_type": 0.75,
    "policy_type_justification": bi(
        "Decreto Supremo N.° 006-2019-MINAM aprobado por el Poder Ejecutivo (22 ago. 2019, "
        "publicado 23 ago. 2019 en El Peruano), refrendado por PCM, MINAM, MINSA, PRODUCE, "
        "MEF, MINEDU, MINCETUR y Cultura. No enmendado. Reglamento ejecutivo de implementación "
        "de Ley 30884. Clasificado 0.75 (decreto supremo/regulación ejecutiva).",
        "Supreme Decree No. 006-2019-MINAM approved by the Executive (22 Aug 2019, published "
        "23 Aug 2019 in El Peruano), countersigned by PCM, MINAM, MINSA, PRODUCE, MEF, "
        "MINEDU, MINCETUR and Culture. Not amended. Executive implementing regulation of "
        "Law 30884. Classified 0.75 (supreme decree/executive regulation).",
    ),
    "policy_integration": 1,
    "policy_sectors_list": bi(
        "comercio minorista, manufactura, importación, ambiente, producción, protección al "
        "consumidor (INDECOPI), aduanas (SUNAT), salud (MINSA), educación, turismo, cultura, "
        "gobiernos locales, reciclaje, empaques PET",
        "retail commerce, manufacturing, imports, environment, production, consumer protection "
        "(INDECOPI), customs (SUNAT), health (MINSA), education, tourism, culture, local "
        "governments, recycling, PET packaging",
    ),
    "policy_circularity": 1.0,
    "policy_lifecycle_phases_list": bi(
        "producción, diseño, consumo, reciclaje, compostaje, disposición, fuga ambiental",
        "production, design, consumption, recycling, composting, disposal, environmental leakage",
    ),
    "policy_budget": 0.5,
    "policy_budget_text": bi(
        "DS Art. 3: financiamiento con cargo al presupuesto institucional de los pliegos "
        "involucrados, sin demandar recursos adicionales al Tesoro Público. Art. 33: "
        "incentivos y mecanismos PRODUCE para MYPE (sin asignación presupuestal específica "
        "en el reglamento).",
        "DS Art. 3: financing from institutional budgets of involved entities, without "
        "requiring additional Public Treasury resources. Art. 33: incentives and PRODUCE "
        "mechanisms for MSMEs (no specific budget allocation in the regulation).",
    ),
}

INSTRUMENTS = [
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": bi(
            "DS Art. 1: aprueba el Reglamento de la Ley N.° 30884 — 6 capítulos, 33 "
            "artículos, 7 DCF, 1 DCT, 1 DCD y Anexo de infracciones, parte integrante del DS.",
            "DS Art. 1: approves the Regulation of Law No. 30884 — 6 chapters, 33 articles, "
            "7 final complementary provisions, 1 transitory, 1 repealing provision and "
            "infractions annex, integral part of the supreme decree.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "'Apruébase' — entrada en vigor del reglamento implementador. MINAM lidera. "
            "Publicado 23 ago. 2019.",
            "'Is approved' — implementing regulation enters into force. MINAM leads. "
            "Published 23 Aug 2019.",
        ),
        "comments": bi(
            "Acto habilitante del reglamento de la ley marco de plásticos.",
            "Enabling act for the flagship plastics law regulation.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": bi(
            "Reglamento Art. 1–2: objeto reglamentar Ley 30884; finalidad contribuir al "
            "derecho a ambiente equilibrado y orientar el plástico hacia economía circular "
            "(reutilizable, retornable, reciclable o degradación sin microplásticos).",
            "Regulation Arts. 1–2: purpose to regulate Law 30884; aim to contribute to the "
            "right to a balanced environment and steer plastics toward a circular economy "
            "(reusable, returnable, recyclable or degradation without microplastics).",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "Objetivo y finalidad declarativa. Marco para instrumentos operativos Capítulos "
            "II–VI.",
            "Declaratory purpose and aim. Framework for operative Chapters II–VI.",
        ),
        "comments": bi(
            "Enuncia explícitamente economía circular del plástico.",
            "Explicitly states plastic circular economy.",
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": bi(
            "Reglamento Art. 3: ámbito obligatorio para quien fabrique, importe, distribuya, "
            "comercialice, entregue, use o consuma: (a) bolsas plásticas; (b) bolsas/envoltorios "
            "en prensa/recibos/información; (c) sorbetes; (d) envases/vasos de poliestireno "
            "expandido; (e) vajilla plástica; (f) botellas PET; (g) insumos PET; (h) otros "
            "bienes poliméricos vía DS.",
            "Regulation Art. 3: mandatory scope for those who manufacture, import, distribute, "
            "sell, deliver, use or consume: (a) plastic bags; (b) bags/wrappers in press/receipts/"
            "information; (c) straws; (d) expanded-polystyrene containers/cups; (e) plastic "
            "tableware; (f) PET bottles; (g) PET inputs; (h) other polymer goods via supreme "
            "decree.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "'Cumplimiento obligatorio' para actores en cadena de valor. Define bienes regulados "
            "vinculados a Ley 30884 Arts. 2–3 y 10.",
            "'Mandatory compliance' for value-chain actors. Defines regulated goods linked to "
            "Law 30884 Arts. 2–3 and 10.",
        ),
        "comments": bi(
            "Definición central de bienes plásticos regulados y actores sujetos.",
            "Central definition of regulated plastic goods and subject actors.",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": bi(
            "Reglamento Art. 4–5: definiciones (bolsas de base polimérica, reutilizable, "
            "retornable, fabricante, importador, distribuidor, comercializador) y clasificación "
            "por usos (un solo uso/reutilizable/retornable), biodegradabilidad y "
            "valorización (reciclable/compostable).",
            "Regulation Arts. 4–5: definitions (polymer-based bags, reusable, returnable, "
            "manufacturer, importer, distributor, retailer) and classification by uses "
            "(single-use/reusable/returnable), biodegradability and valorization "
            "(recyclable/compostable).",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Glosario reglamentario + glosario Ley 30884/DL 1278/DS 014-2017. Base técnica "
            "para prohibiciones y registro.",
            "Regulatory glossary + Law 30884/DL 1278/DS 014-2017 glossary. Technical basis "
            "for prohibitions and registry.",
        ),
        "comments": bi(
            "Estándares de definición y clasificación de bienes poliméricos.",
            "Definition and classification standards for polymer goods.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": bi(
            "Reglamento Art. 6.1: fabricantes, importadores y distribuidores de bolsas de "
            "plástico 'se inscriben' en el Registro (Cap. III) y brindan información anual "
            "según Art. 12.1.",
            "Regulation Art. 6.1: manufacturers, importers and distributors of plastic bags "
            "'must register' in the Registry (Ch. III) and provide annual information per "
            "Art. 12.1.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "'Se inscriben' y reporte anual jurado (Art. 12.1). MINAM administra Registro. "
            "Sanciones Anexo 5.1–5.6.",
            "'Must register' and annual sworn report (Art. 12.1). MINAM administers Registry. "
            "Annex sanctions 5.1–5.6.",
        ),
        "comments": bi(
            "Registro obligatorio de productores/importadores de bolsas plásticas.",
            "Mandatory registry of plastic-bag producers/importers.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": bi(
            "Reglamento Art. 6.2: fabricantes/importadores/distribuidores de otros bienes "
            "plásticos 'pueden optar' por inscribirse en el Registro y reportar información "
            "anual facultativa.",
            "Regulation Art. 6.2: manufacturers/importers/distributors of other plastic goods "
            "'may opt' to register in the Registry and file optional annual information.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Inscripción facultativa ('pueden optar'). Reporte Art. 12.2 incluye % material "
            "reciclado si corresponde.",
            "Optional registration ('may opt'). Art. 12.2 reporting includes recycled content "
            "% where applicable.",
        ),
        "comments": bi(
            "Registro voluntario para bienes plásticos distintos a bolsas.",
            "Voluntary registry for plastic goods other than bags.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Consumption",
        "instrument_description": bi(
            "Reglamento Art. 6.3–6.4: comercializadores realizan educación/sensibilización "
            "sobre uso responsable de bienes plásticos; fabricantes/importadores/distribuidores/"
            "comercializadores participan en acciones de MINAM, MINEDU, PRODUCE y gobiernos "
            "descentralizados.",
            "Regulation Arts. 6.3–6.4: retailers carry out education/awareness on responsible "
            "plastic-goods use; manufacturers/importers/distributors/retailers participate in "
            "MINAM, MINEDU, PRODUCE and decentralized-government actions.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "'Realizan acciones' y 'participan de las acciones'. Sin sanción específica en "
            "Art. 6. Pertinencia cultural y lingüística.",
            "'Carry out actions' and 'participate in actions'. No specific penalty in Art. 6. "
            "Cultural and linguistic relevance.",
        ),
        "comments": bi(
            "Obligaciones de educación del sector privado sobre plásticos.",
            "Private-sector education obligations on plastics.",
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": bi(
            "Reglamento Art. 7: control aduanero de importación de bolsas y sorbetes (Art. 7.1, "
            "MINAM absuelve consultas SUNAT); fiscalización en mercado de bienes importados "
            "prohibidos (Art. 7.2); responsabilidad del importador (Art. 7.3).",
            "Regulation Art. 7: customs control of bag and straw imports (Art. 7.1, MINAM "
            "answers SUNAT technical queries); market surveillance of prohibited imported "
            "goods (Art. 7.2); importer liability (Art. 7.3).",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "Control en Aduanas + fiscalización mercado. SUNAT/MINAM coordinación. DCF Cuarta "
            "listado subpartidas.",
            "Customs control + market surveillance. SUNAT/MINAM coordination. Fourth FCP "
            "tariff subheading list.",
        ),
        "comments": bi(
            "Mecanismo de control de importaciones de plásticos regulados.",
            "Controlled-import mechanism for regulated plastics.",
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Consumption",
        "instrument_description": bi(
            "Reglamento Art. 8: a bolsas pequeñas/delgadas y sorbetes les aplican excepciones "
            "Ley 30884 Art. 4 'según disponga' MINSA en coordinación con MINAM.",
            "Regulation Art. 8: small/thin bags and straws are subject to Law 30884 Art. 4 "
            "exceptions 'as MINSA provides' in coordination with MINAM.",
        ),
        "instrument_in_force": 0,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "Remisión a acto MINSA ('según disponga'). Excepciones no autoejecutables en "
            "reglamento — habilitación.",
            "Referral to MINSA act ('as provides'). Exceptions not self-executing in "
            "regulation — enabling.",
        ),
        "comments": bi(
            "Verificar listado MINSA DCF Sexta (higiene/salud) para excepciones operativas.",
            "Verify MINSA Sixth FCP list (hygiene/health) for operative exceptions.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Consumption",
        "instrument_description": bi(
            "Reglamento Art. 9: deberes de consumidores — no generación/minimización de "
            "residuos plásticos, optar por reutilizables/reciclables/sin microplásticos, "
            "segregación adecuada.",
            "Regulation Art. 9: consumer duties — avoid/minimize plastic-waste generation, "
            "choose reusable/recyclable/non-microplastic alternatives, proper segregation.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "'Procurar', 'Minimizar', 'Optar', 'Realizar' — obligaciones orientadoras sin "
            "sanción directa en Art. 9.",
            "'Seek', 'Minimize', 'Choose', 'Carry out' — guiding duties without direct "
            "penalty in Art. 9.",
        ),
        "comments": bi(
            "Deberes ciudadanos de prevención y segregación de residuos plásticos.",
            "Citizen duties for plastic-waste prevention and segregation.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": bi(
            "Reglamento Art. 10–11: Registro de fabricantes/importadores/distribuidores "
            "(informativo, MINAM/SINIA, no constitutivo de derechos); inscripción directa en "
            "plataforma virtual.",
            "Regulation Arts. 10–11: manufacturer/importer/distributor Registry (informative, "
            "MINAM/SINIA, non-constitutive of rights); direct enrollment on virtual platform.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "Plataforma SINIA operativa. 'Se registran directamente'. No requisito habilitante "
            "pero obligatorio para bolsas (Art. 6.1).",
            "SINIA platform operational. 'Register directly'. Not a licensing requirement but "
            "mandatory for bags (Art. 6.1).",
        ),
        "comments": bi(
            "Registro nacional de cadena de suministro de plásticos.",
            "National supply-chain registry for plastics.",
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": bi(
            "Reglamento Art. 12.1: fabricantes/importadores/distribuidores de bolsas informan "
            "hasta último día hábil de marzo (declaración jurada, datos mensualizados) cantidad "
            "y peso por tipo de bien, resina, destino y agente de mercado.",
            "Regulation Art. 12.1: bag manufacturers/importers/distributors report by last "
            "business day of March (sworn declaration, monthly data) quantity and weight by "
            "good type, resin, destination and market agent.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "Plazo anual fijo y datos cuantificados. Carácter jurado. OEFA fiscaliza "
            "veracidad/reporte (Art. 28.1e-f).",
            "Fixed annual deadline and quantified data. Sworn character. OEFA enforces "
            "truthfulness/reporting (Art. 28.1e-f).",
        ),
        "comments": bi(
            "Monitoreo estadístico de producción/comercio de bolsas plásticas.",
            "Statistical monitoring of plastic-bag production/trade.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": bi(
            "Reglamento Art. 13–15: indicadores/estadísticas del Registro de acceso público; "
            "interoperabilidad PIDE (PCM/MINAM/PRODUCE/SUNAT); datos abiertos en "
            "datosabiertos.gob.pe.",
            "Regulation Arts. 13–15: public-access Registry indicators/statistics; PIDE "
            "interoperability (PCM/MINAM/PRODUCE/SUNAT); open data on datosabiertos.gob.pe.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "MINAM 'habilita datos en formatos abiertos'. PCM asistencia técnica PIDE. "
            "Excepciones datos personales Ley 29733.",
            "MINAM 'enables open-format data'. PCM technical assistance for PIDE. Personal "
            "data exceptions Law 29733.",
        ),
        "comments": bi(
            "Gobernanza de datos sobre plásticos regulados.",
            "Data governance on regulated plastics.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Consumption",
        "instrument_description": bi(
            "Reglamento Art. 16–19, 21: acciones semestrales de educación/sensibilización "
            "por MINAM (incl. formalización recicladores), MINEDU (currículo transversal), "
            "INDECOPI (consumidor), PRODUCE (industria/MYPE), administración pública y "
            "gobiernos locales.",
            "Regulation Arts. 16–19, 21: semiannual education/awareness actions by MINAM "
            "(incl. recycler formalization), MINEDU (cross-cutting curriculum), INDECOPI "
            "(consumers), PRODUCE (industry/MSMEs), public administration and local "
            "governments.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "'Realiza semestralmente' / 'incorpora' / 'desarrollan'. DCF Tercera: lineamientos "
            "MINAM en 60 días.",
            "'Carries out semiannually' / 'incorporates' / 'develop'. Third FCP: MINAM "
            "guidelines within 60 days.",
        ),
        "comments": bi(
            "Marco multiactor de educación sobre plásticos de un solo uso.",
            "Multi-actor education framework on single-use plastics.",
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Environmental leakage",
        "instrument_description": bi(
            "Reglamento Art. 20: entidades de áreas protegidas, patrimonio, museos, playas, "
            "ríos/lagos amazónicos deben educar visitantes; proveedores turísticos comunican "
            "prohibición Ley 30884 Art. 3.1(a) e incluyen información en material publicitario.",
            "Regulation Art. 20: entities in protected areas, heritage sites, museums, beaches, "
            "Amazon rivers/lakes must educate visitors; tourism providers communicate Law 30884 "
            "Art. 3.1(a) prohibition and include information in advertising material.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "'Deben realizar' educación permanente. Proveedores 'comunican' prohibición en "
            "áreas sensibles. SERNANP/Cultura fiscalizan Art. 28.4–28.5.",
            "'Must carry out' permanent education. Providers 'communicate' prohibition in "
            "sensitive areas. SERNANP/Culture enforce Arts. 28.4–28.5.",
        ),
        "comments": bi(
            "Control de plásticos en turismo y áreas de fuga ambiental.",
            "Plastic control in tourism and environmental-leakage areas.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": bi(
            "Reglamento Art. 22: MINAM/PRODUCE/MINEDU/CONCYTEC promueven I+D+i para mitigar "
            "impacto de residuos plásticos, optimizar gestión integral, reducir plástico de un "
            "solo uso y desarrollar sustitutos sostenibles.",
            "Regulation Art. 22: MINAM/PRODUCE/MINEDU/CONCYTEC promote R&D for mitigating "
            "plastic-waste impact, optimizing integrated management, reducing single-use "
            "plastic and developing sustainable substitutes.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "'Promueven proyectos' — función de fomento sin metas cuantificadas en Art. 22.",
            "'Promote projects' — promotional function without quantified targets in Art. 22.",
        ),
        "comments": bi(
            "Innovación y sustitutos al plástico de un solo uso.",
            "Innovation and single-use plastic substitutes.",
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Disposal",
        "instrument_description": bi(
            "Reglamento Art. 23: gestión integral de residuos de bienes plásticos según DL "
            "1278 — prevalece minimización sobre valorización y valorización sobre disposición "
            "final; prioriza bienes reutilizables.",
            "Regulation Art. 23: integrated management of plastic-goods waste per DL 1278 — "
            "minimization prevails over valorization and valorization over final disposal; "
            "prioritizes reusable goods.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "Jerarquía de residuos explícita para plásticos. Vinculado DL 1278 y DS "
            "014-2017-MINAM.",
            "Explicit waste hierarchy for plastics. Linked to DL 1278 and DS 014-2017-MINAM.",
        ),
        "comments": bi(
            "Instrumento de jerarquía de residuos aplicable a flujos plásticos.",
            "Waste-hierarchy instrument applicable to plastic streams.",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": bi(
            "Reglamento Art. 24.1: criterios para ampliar sustitución/prohibición de otros "
            "bienes plásticos, contenido reciclado y porcentajes — economía circular, "
            "representatividad en residuos/basura marina/microplásticos, avance tecnológico, "
            "no reciclables, interferencia en reciclabilidad, riesgo salud/ambiente.",
            "Regulation Art. 24.1: criteria for expanding substitution/prohibition of other "
            "plastic goods, recycled content and percentages — circular economy, "
            "representativeness in waste/marine litter/microplastics, technological advance, "
            "non-recyclables, recyclability interference, health/environment risk.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Criterios listados para futuras restricciones (DCF Sexta Ley 30884). No aplican "
            "solas — requieren DS Art. 24.2.",
            "Listed criteria for future restrictions (Law 30884 Sixth FCP). Not self-applied — "
            "require DS Art. 24.2.",
        ),
        "comments": bi(
            "Criterios para extender controles a otros plásticos.",
            "Criteria for extending controls to other plastics.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": bi(
            "Reglamento Art. 24.2: acciones de ampliación 'se establecen mediante Decreto "
            "Supremo', previo cumplimiento normativa OMC y CAN sobre obstáculos técnicos al "
            "comercio.",
            "Regulation Art. 24.2: expansion actions 'are established by Supreme Decree', "
            "after WTO and Andean Community technical-barriers-to-trade compliance.",
        ),
        "instrument_in_force": 0,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "Poder habilitante ('se establecen mediante DS'). MINAM/PRODUCE. in_force 0 salvo "
            "DS emitido.",
            "Enabling power ('established by supreme decree'). MINAM/PRODUCE. in_force 0 "
            "unless decree issued.",
        ),
        "comments": bi(
            "Habilitación para ampliar restricciones plásticas vía reglamento.",
            "Enabling expansion of plastic restrictions via regulation.",
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": bi(
            "Reglamento Art. 25: gobiernos locales con fabricantes/distribuidores/"
            "comercializadores implementan segregación en fuente y recolección selectiva de "
            "residuos plásticos; recolección selectiva mínimo los miércoles (Día del Reciclaje "
            "del Plástico); recicladores formalizados integrados al sistema.",
            "Regulation Art. 25: local governments with manufacturers/distributors/retailers "
            "implement source segregation and selective collection of plastic waste; selective "
            "collection at minimum on Wednesdays (Plastic Recycling Day); formalized recyclers "
            "integrated into the system.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "'Implementan sistemas' y 'debe ser selectiva'. Día miércoles obligatorio. DL "
            "1278 Programas de Segregación.",
            "'Implement systems' and 'must be selective'. Wednesday collection mandatory. DL "
            "1278 Segregation Programs.",
        ),
        "comments": bi(
            "Gestión municipal de residuos plásticos y recolección selectiva.",
            "Municipal plastic-waste management and selective collection.",
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": bi(
            "Reglamento Art. 26: valorización de residuos plásticos reutilizables/reciclables "
            "prioritariamente por reciclaje; biodegradables por compostaje; MINAM/gobiernos "
            "locales promueven infraestructura y regímenes especiales.",
            "Regulation Art. 26: valorization of reusable/recyclable plastic waste primarily "
            "through recycling; biodegradable through composting; MINAM/local governments "
            "promote infrastructure and special regimes.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "'Se realiza prioritariamente a través del reciclaje/compostaje'. Promoción "
            "infraestructura sin metas cuantificadas.",
            "'Is carried out primarily through recycling/composting'. Infrastructure promotion "
            "without quantified targets.",
        ),
        "comments": bi(
            "Valorización de flujos plásticos reciclables y biodegradables.",
            "Valorization of recyclable and biodegradable plastic streams.",
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": bi(
            "Reglamento Art. 28: reparto de fiscalización — OEFA (no fabricar/distribuir bolsas "
            "<900 cm²/<50 µm, sorbetes, tecnopor; registro); PRODUCE (reglamentos técnicos); "
            "INDECOPI (información al consumidor); SERNANP (ANP); Cultura (patrimonio/museos); "
            "municipios (Art. 2 y 3 Ley 30884 en comercio).",
            "Regulation Art. 28: enforcement allocation — OEFA (no manufacture/distribute "
            "bags <900 cm²/<50 µm, straws, styrofoam; registry); PRODUCE (technical "
            "regulations); INDECOPI (consumer information); SERNANP (protected areas); Culture "
            "(heritage/museums); municipalities (Law 30884 Arts. 2 and 3 in commerce).",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "'No fabricar' / 'No distribuir' — prohibiciones manufactura explícitas OEFA. "
            "Multi-autoridad coordinada. DCF Quinta/Séptima fechas inicio.",
            "'Do not manufacture' / 'Do not distribute' — explicit OEFA manufacturing "
            "prohibitions. Multi-authority coordination. Fifth/Seventh FCP start dates.",
        ),
        "comments": bi(
            "Régimen compartido de fiscalización de plásticos de un solo uso.",
            "Shared enforcement regime for single-use plastics.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": bi(
            "Reglamento Art. 29: entidades competentes 'pueden realizar' fiscalizaciones "
            "orientativas (sin fines punitivos salvo daños/riesgos significativos) para micro "
            "empresas o administrados no fiscalizados.",
            "Regulation Art. 29: competent entities 'may carry out' orientative inspections "
            "(non-punitive unless significant harm/risk) for micro enterprises or previously "
            "uninspected parties.",
        ),
        "instrument_in_force": 0,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "'Pueden realizar' — facultativo. Mecanismo de cumplimiento asistido, no "
            "obligación directa.",
            "'May carry out' — discretionary. Assisted-compliance mechanism, not direct "
            "obligation.",
        ),
        "comments": bi(
            "Fiscalización orientativa para MYPE — habilitación, no mandato.",
            "Orientative MSME inspections — enabling, not mandate.",
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Consumption",
        "instrument_description": bi(
            "Reglamento Art. 30–31 y Anexo: tipificación de infracciones (bolsas, sorbetes, "
            "tecnopor, registro) con sanciones graduadas hasta 130 UIT según tamaño empresarial; "
            "graduación según normas especiales y Ley 27444.",
            "Regulation Arts. 30–31 and Annex: typification of infractions (bags, straws, "
            "styrofoam, registry) with graduated sanctions up to 130 UIT by company size; "
            "graduation per special norms and Law 27444.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": bi(
            "Anexo con 20+ supuestos infraccionales cuantificados vinculados a Ley 30884. "
            "Multas por entrega gratuita bolsas, comercialización prohibida, incumplimiento "
            "registro.",
            "Annex with 20+ quantified infraction types linked to Law 30884. Fines for free "
            "bag delivery, prohibited sales, registry non-compliance.",
        ),
        "comments": bi(
            "Régimen sancionador detallado para incumplimientos plásticos.",
            "Detailed sanction regime for plastic non-compliance.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": bi(
            "Reglamento Art. 32: entidades de fiscalización reportan trimestralmente acciones "
            "de fiscalización y sanción en plataforma virtual del Registro.",
            "Regulation Art. 32: enforcement entities report quarterly enforcement and "
            "sanction actions on the Registry virtual platform.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "'Reportan trimestralmente' — obligación de reporte periódico a MINAM.",
            "'Report quarterly' — periodic reporting obligation to MINAM.",
        ),
        "comments": bi(
            "Monitoreo de cumplimiento multi-autoridad.",
            "Multi-authority compliance monitoring.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": bi(
            "Reglamento Art. 33: incentivos públicos a actores de cadena plástica; PRODUCE "
            "mecanismos MYPE para sustitución progresiva de tecnopor (formalización, "
            "financiamiento, instructivos, articulación); MINAM listado desempeño ambiental y "
            "Acuerdos de Producción Limpia; municipios pueden establecer incentivos.",
            "Regulation Art. 33: public incentives for plastic-chain actors; PRODUCE MSME "
            "mechanisms for progressive styrofoam substitution (formalization, financing, "
            "guides, linkage); MINAM environmental-performance list and Cleaner Production "
            "Agreements; municipalities may establish incentives.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "'Promueven', 'establece', 'puede otorgar' — mix obligatorio/facultativo. "
            "PRODUCE emite normas complementarias Art. 33.2.",
            "'Promote', 'establish', 'may grant' — mix of mandatory/discretionary. PRODUCE "
            "issues complementary norms Art. 33.2.",
        ),
        "comments": bi(
            "Incentivos y transición MYPE para prohibición de tecnopor.",
            "Incentives and MSME transition for styrofoam ban.",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": bi(
            "DCF Primera: reglamentos técnicos de bienes regulados se aprueban por DS con "
            "fiscalización y sanción; NTP voluntarias en concordancia.",
            "First FCP: technical regulations for regulated goods approved by supreme decree "
            "with enforcement and sanctions; voluntary Peruvian technical standards in "
            "concordance.",
        ),
        "instrument_in_force": 0,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "'Se aprueban mediante Decreto Supremo' — habilitación para estándares técnicos "
            "futuros (INACAL Ley 30884 Art. 5).",
            "'Are approved by supreme decree' — enabling future technical standards (INACAL "
            "Law 30884 Art. 5).",
        ),
        "comments": bi(
            "Marco para reglamentos técnicos de diseño de plásticos.",
            "Framework for plastic design technical regulations.",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": bi(
            "DCF Segunda: certificado de biodegradabilidad exigible a importadores/productores "
            "se regula por reglamento técnico respectivo (Ley 30884 Art. 11).",
            "Second FCP: biodegradability certificate required of importers/producers is "
            "regulated by the respective technical regulation (Law 30884 Art. 11).",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Remisión a reglamento técnico específico. Vincula exenciones ICBP y prohibiciones "
            "oxodegradables.",
            "Referral to specific technical regulation. Links ICBP exemptions and "
            "oxo-degradable prohibitions.",
        ),
        "comments": bi(
            "Estándar de certificación para bolsas biodegradables.",
            "Certification standard for biodegradable bags.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Consumption",
        "instrument_description": bi(
            "DCF Tercera: lineamientos MINAM para educación/comunicación en 60 días calendario "
            "desde vigencia, considerando diversidad cultural/lingüística.",
            "Third FCP: MINAM guidelines for education/communication within 60 calendar days "
            "of entry into force, considering cultural/linguistic diversity.",
        ),
        "instrument_in_force": 0,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "MINAM 'apruebe' lineamientos en plazo — habilitación. Educación Arts. 16–21 "
            "depende de lineamientos.",
            "MINAM 'approves' guidelines within deadline — enabling. Education Arts. 16–21 "
            "depends on guidelines.",
        ),
        "comments": bi(
            "Verificar si lineamientos MINAM fueron aprobados post-2019.",
            "Verify whether MINAM guidelines were approved post-2019.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": bi(
            "DCF Cuarta: MINAM con SUNAT elabora listado de bienes regulados con subpartidas "
            "arancelarias, aprobado por Resolución Ministerial.",
            "Fourth FCP: MINAM with SUNAT prepares list of regulated goods with tariff "
            "subheadings, approved by Ministerial Resolution.",
        ),
        "instrument_in_force": 0,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "'Elabora' / 'será aprobado por RM' — acto administrativo posterior requerido "
            "para control aduanero Art. 7.",
            "'Prepares' / 'will be approved by ministerial resolution' — subsequent "
            "administrative act required for Art. 7 customs control.",
        ),
        "comments": bi(
            "Listado aduanero pendiente de RM — verificar emisión.",
            "Customs list pending ministerial resolution — verify issuance.",
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": bi(
            "DCF Quinta y Séptima: fiscalización de bienes sin reglamento técnico inicia día "
            "siguiente al vencimiento de plazos Ley 30884; OEFA exigible para bolsas/sorbetes "
            "desde 20 dic. 2019 y tecnopor desde 20 dic. 2021.",
            "Fifth and Seventh FCPs: enforcement for goods without technical regulation starts "
            "the day after Law 30884 deadline expiry; OEFA enforceable for bags/straws from "
            "20 Dec 2019 and styrofoam from 20 Dec 2021.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": bi(
            "Fechas de inicio de fiscalización cuantificadas y vinculantes. Alineado calendario "
            "Ley 30884.",
            "Quantified binding enforcement start dates. Aligned with Law 30884 calendar.",
        ),
        "comments": bi(
            "Plazos operativos de cumplimiento y fiscalización.",
            "Operative compliance and enforcement timelines.",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Consumption",
        "instrument_description": bi(
            "DCF Sexta: MINSA con MINAM aprueba en 30 días listado de bolsas usadas por "
            "razones de higiene o salud.",
            "Sixth FCP: MINSA with MINAM approves within 30 days list of bags used for "
            "hygiene or health reasons.",
        ),
        "instrument_in_force": 0,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "MINSA 'aprueba' listado en plazo — habilitación para excepciones Art. 8 y Ley "
            "30884 Art. 4.",
            "MINSA 'approves' list within deadline — enabling for Art. 8 and Law 30884 Art. 4 "
            "exceptions.",
        ),
        "comments": bi(
            "Listado higiene/salud requerido para exenciones operativas.",
            "Hygiene/health list required for operative exemptions.",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": bi(
            "DCT Única: fabricantes/envasadores de botellas PET e importadores de insumos "
            "'pueden solicitar' asistencia técnica MINAM para implementar progresividad del "
            "15% PET-PCR (Ley 30884 Art. 10).",
            "Single Transitory Provision: PET bottle manufacturers/packagers and input "
            "importers 'may request' MINAM technical assistance to implement 15% post-consumer "
            "recycled PET progressivity (Law 30884 Art. 10).",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Asistencia técnica facultativa ('pueden solicitar'). Apoya cumplimiento contenido "
            "reciclado mínimo.",
            "Optional technical assistance ('may request'). Supports minimum recycled-content "
            "compliance.",
        ),
        "comments": bi(
            "Mecanismo de apoyo al requisito PET-PCR 15%.",
            "Support mechanism for 15% post-consumer recycled PET requirement.",
        ),
    },
]

COLUMNS = [
    ("A", "policy_name"),
    ("B", "policy_url"),
    ("C", "policy_year"),
    ("D", "policy_objective"),
    ("E", "policy_target"),
    ("F", "policy_target_text"),
    ("G", "policy_type"),
    ("H", "policy_type_justification"),
    ("I", "policy_integration"),
    ("J", "policy_sectors_list"),
    ("K", "policy_circularity"),
    ("L", "policy_lifecycle_phases_list"),
    ("M", "policy_budget"),
    ("N", "policy_budget_text"),
    ("O", "policy_score"),
    ("P", "instrument_type"),
    ("Q", "instrument_lifecycle_stage"),
    ("R", "instrument_description"),
    ("S", "instrument_in_force"),
    ("T", "instrument_implementation"),
    ("U", "instrument_implementation_text"),
    ("V", "instrument_score"),
    ("W", "comments"),
]


def policy_score_row(instrument_type: float) -> float:
    return round(
        (
            POLICY["policy_type"]
            + POLICY["policy_integration"]
            + POLICY["policy_circularity"]
            + POLICY["policy_budget"]
            + instrument_type
        )
        / 5,
        3,
    )


def instrument_score_row(instrument_type: float, implementation: float) -> float:
    return round((instrument_type + implementation) / 2, 3)


def build_excel(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "4P Index Coding"
    ws.append([f"{col} — {name}" for col, name in COLUMNS])
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for inst in INSTRUMENTS:
        ws.append(
            [
                POLICY["policy_name"],
                POLICY["policy_url"],
                POLICY["policy_year"],
                POLICY["policy_objective"],
                POLICY["policy_target"],
                POLICY["policy_target_text"],
                POLICY["policy_type"],
                POLICY["policy_type_justification"],
                POLICY["policy_integration"],
                POLICY["policy_sectors_list"],
                POLICY["policy_circularity"],
                POLICY["policy_lifecycle_phases_list"],
                POLICY["policy_budget"],
                POLICY["policy_budget_text"],
                policy_score_row(inst["instrument_type"]),
                inst["instrument_type"],
                inst["instrument_lifecycle_stage"],
                inst["instrument_description"],
                inst["instrument_in_force"],
                inst["instrument_implementation"],
                inst["instrument_implementation_text"],
                instrument_score_row(inst["instrument_type"], inst["instrument_implementation"]),
                inst["comments"],
            ]
        )

    widths = {
        "A": 40, "B": 38, "C": 8, "D": 44, "E": 8, "F": 44, "G": 8, "H": 40,
        "I": 10, "J": 40, "K": 10, "L": 36, "M": 10, "N": 42, "O": 10, "P": 10,
        "Q": 18, "R": 50, "S": 10, "T": 12, "U": 44, "V": 10, "W": 40,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(path)


def add_article(doc: Document, es: str, en: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(es)
    r.bold = True
    doc.add_paragraph(en).paragraph_format.space_after = Pt(8)


def build_word(path: Path) -> None:
    doc = Document()
    for margin in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(doc.sections[0], margin, Inches(1))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t = title.add_run(f"{POLICY_NAME_EN}\n{POLICY_NAME_ES}")
    t.bold = True
    t.font.size = Pt(14)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(
        "Signed: 22 August 2019 | Published: 23 August 2019 (El Peruano)\n"
        "Not amended\n"
        f"Source: {POLICY_URL}"
    )

    doc.add_heading("English Translation of Key Provisions", 1)
    doc.add_paragraph(
        "English translation of principal provisions of Supreme Decree No. 006-2019-MINAM "
        "regulating Law No. 30884 on single-use plastics."
    )

    sections = [
        (
            "Supreme Decree Art. 1 — Approval",
            "Apruébase el Reglamento de la Ley N° 30884... seis (6) Capítulos, treinta y tres "
            "(33) Artículos... y un (1) Anexo.",
            "The Regulation of Law No. 30884 is approved... six (6) Chapters, thirty-three "
            "(33) Articles... and one (1) Annex.",
        ),
        (
            "Regulation Arts. 1–2 — Purpose",
            "Reglamentar la Ley N° 30884... orientar el uso del plástico hacia una economía "
            "circular... reutilizables, retornables... reciclables.",
            "Regulate Law No. 30884... steer plastic use toward a circular economy... "
            "reusable, returnable... recyclable.",
        ),
        (
            "Regulation Art. 3 — Scope",
            "Cumplimiento obligatorio... bolsas de plástico, sorbetes, recipientes de "
            "poliestireno expandido, vajilla plástica, botellas PET, insumos PET y otros "
            "bienes poliméricos.",
            "Mandatory compliance... plastic bags, straws, expanded-polystyrene containers, "
            "plastic tableware, PET bottles, PET inputs and other polymer goods.",
        ),
        (
            "Regulation Arts. 6, 10–12 — Registry",
            "Fabricantes, importadores y distribuidores de bolsas se inscriben en el Registro "
            "SINIA y reportan anualmente cantidad y peso por tipo de resina.",
            "Bag manufacturers, importers and distributors enroll in the SINIA Registry and "
            "annually report quantity and weight by resin type.",
        ),
        (
            "Regulation Art. 25 — Selective collection",
            "Gobiernos locales implementan segregación y recolección selectiva de residuos "
            "plásticos... como mínimo los días miércoles.",
            "Local governments implement segregation and selective collection of plastic "
            "waste... at minimum on Wednesdays.",
        ),
        (
            "Regulation Art. 28 — Enforcement",
            "OEFA fiscaliza prohibiciones de fabricación/distribución; PRODUCE reglamentos "
            "técnicos; INDECOPI información al consumidor; SERNANP, Cultura y municipios en "
            "sus ámbitos.",
            "OEFA enforces manufacturing/distribution prohibitions; PRODUCE technical "
            "regulations; INDECOPI consumer information; SERNANP, Culture and municipalities "
            "in their scopes.",
        ),
        (
            "Annex — Infractions",
            "Tipificación de infracciones por bolsas, sorbetes, tecnopor y registro con "
            "sanciones hasta 130 UIT.",
            "Typification of infractions for bags, straws, styrofoam and registry with "
            "sanctions up to 130 UIT.",
        ),
        (
            "Final provisions — Timelines",
            "Fiscalización OEFA bolsas/sorbetes desde 20 dic. 2019; tecnopor desde 20 dic. "
            "2021.",
            "OEFA enforcement for bags/straws from 20 Dec 2019; styrofoam from 20 Dec 2021.",
        ),
    ]

    for heading, es, en in sections:
        doc.add_heading(heading, 2)
        add_article(doc, es, en)

    doc.add_paragraph()
    note = doc.add_paragraph(
        "Note: DS 006-2019-MINAM is the main implementing regulation of Law 30884. It "
        "complements DS 244-2019-EF (plastic-bag consumption tax) and links to DL 1278 and "
        "DS 014-2017-MINAM on integrated solid-waste management."
    )
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(path)


def main() -> None:
    excel = OUTPUT_DIR / "Peru_DS_006_2019_MINAM_4P_Index_Coding.xlsx"
    word = OUTPUT_DIR / "Peru_DS_006_2019_MINAM_English_Translation.docx"
    build_excel(excel)
    build_word(word)
    print(f"Created: {excel}")
    print(f"Created: {word}")
    print(f"Instruments: {len(INSTRUMENTS)}")
    print()
    print("Download links:")
    print(
        f"  Excel: {REPO_RAW}/output/peru-ds-006-2019-minam/"
        "Peru_DS_006_2019_MINAM_4P_Index_Coding.xlsx"
    )
    print(
        f"  Word:  {REPO_RAW}/output/peru-ds-006-2019-minam/"
        "Peru_DS_006_2019_MINAM_English_Translation.docx"
    )


if __name__ == "__main__":
    main()
