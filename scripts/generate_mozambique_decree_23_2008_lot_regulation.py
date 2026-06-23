#!/usr/bin/env python3
"""Generate 4P Index coding and English translation for Mozambique Decree 23/2008 (LOT Regulation)."""

import os
from datetime import date

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT_DIR = "/workspace/output/mozambique"
os.makedirs(OUTPUT_DIR, exist_ok=True)

POLICY_FIELDS = {
    "policy_name": (
        "Regulamento da Lei de Ordenamento do Território "
        "(Decreto n.º 23/2008)"
    ),
    "policy_url": (
        "https://sibmoz.gov.mz/content/uploads/2025/11/"
        "Decreto-n.o-23_2008_Regulamento-da-Lei-de-Ordenamento-do-Territorio.pdf"
    ),
    "policy_year": 2008,
    "policy_objective": (
        "Approve the Regulation of the Territorial Planning Law (Law 19/2007), "
        "establishing the legal regime for territorial planning instruments at "
        "national, provincial, district and municipal levels — including Urban "
        "Structure Plans (PEU), General/Partial Urbanisation Plans (PGU/PPU) "
        "and general instruments (soil qualification, cadastre, zoning). "
        "Plastic-relevant scope: PEU must define solid-waste treatment systems "
        "and reception/processing zones (Art. 42(2)(c)); GPU/PPU must define "
        "sanitation networks and locate drainage for surface and used waters "
        "(Arts. 4, 44); PEU maps polluting industrial zones and environmental "
        "protection areas (Art. 43). No explicit plastics (plástico) reference; "
        "resíduos sólidos includes plastic waste fractions."
    ),
    "policy_target": 0,
    "policy_target_text": "",
    "policy_type": 0.75,
    "policy_type_justification": (
        "Executive regulation (Regulamento) approved by Council of Ministers "
        "Decree n.º 23/2008 (approved 13 May 2008, published 1 July 2008, "
        "Boletim da República Series I No. 26, 3rd Supplement) under Law "
        "19/2007 Art. 30 — sub-legislative instrument implementing LOT "
        "planning instruments with binding effect once published. Art. 2 "
        "assigns expropriation procedures to Environment, Finance and Justice "
        "ministers. Operational waste collection rules remain in Decree 94/2014."
    ),
    "policy_integration": 0.75,
    "policy_sectors_list": (
        "municipalities, water, waste management, industry, environment"
    ),
    "policy_circularity": 0.50,
    "policy_lifecycle_phases_list": "disposal, environmental leakage",
    "policy_budget": 0,
    "policy_budget_text": "",
    "policy_score": "[auto]",
}

INSTRUMENTS = [
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Disposal",
        "instrument_description": (
            "Art. 42(2)(c) — Urban Structure Plan (PEU) materialisation: must "
            "define solid-waste treatment systems (sistemas de tratamento de "
            "resíduos sólidos) and zones for their reception and processing — "
            "the only explicit solid-waste reference in the regulation. Resíduos "
            "sólidos includes plastic waste fractions though plastics are not named."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            'Art. 42(2)(c): "Os sistemas de tratamento de resíduos sólidos e as '
            'zonas para a sua recepção e processamento"; progressive execution '
            "principles under PEU objectives (Art. 42(1)); municipal elaboration "
            "and approval with public consultation (Arts. 6, 10–11, 13)."
        ),
        "instrument_score": "[auto]",
        "comments": (
            "Primary solid-waste instrument — explicit resíduos sólidos in PEU. "
            "Not coded separately: cemetery location principles (Art. 42(2)(d))."
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Disposal",
        "instrument_description": (
            "Art. 4(5)(b)–(c) / GPU and PPU — General and Partial Urbanisation "
            "Plans must structure and qualify urban land, balancing uses and "
            "functions, and define networks for transport, communications, "
            "energy and sanitation (saneamento) plus social facilities — "
            "requiring municipal urbanisation plans to integrate sanitation "
            "infrastructure relevant to wastewater and municipal waste services."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            'Art. 4(5)(b): PGU "define as redes de transporte, comunicações, '
            'energia e saneamento"; Art. 44(a): GPU/PPU materialise PEU '
            "principles including sanitation infrastructure; public consultation "
            "and publication requirements (Arts. 10–11, 18)."
        ),
        "instrument_score": "[auto]",
        "comments": (
            "Sanitation-network planning instrument — complements Art. 42(2)(c) "
            "waste-treatment zoning. Not coded: traffic/pedestrianisation "
            "provisions (Art. 42(2)(g), Art. 44(i))."
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Disposal",
        "instrument_description": (
            "Art. 44(h) — GPU/PPU objectives: locate railways, high-voltage lines, "
            "aqueducts, surface-water drainage systems and used-water (águas usadas) "
            "drainage, and other public-interest infrastructure — spatial planning "
            "for wastewater and stormwater systems affecting waste and pollution "
            "pathways in urban territories."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            'Art. 44(h): "localização das vias férreas, linhas de alta tensão, '
            'aquedutos, sistemas de drenagem de águas superficiais e de águas '
            'usadas e de todo e qualquer outro sistema ou infra-estrutura para '
            'uso público e interesse colectivo".'
        ),
        "instrument_score": "[auto]",
        "comments": (
            "Wastewater/drainage infrastructure instrument — bundled with Art. 42(2)(b) "
            "surface-runoff control systems. Not coded: water capture/distribution "
            "in soil classification (Art. 52(f))."
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Environmental leakage",
        "instrument_description": (
            "Art. 43(f)–(g) — PEU content: detailed mapping of zones for polluting "
            "and/or incompatible industrial activities; identification of "
            "environmental protection zones and ecologically important areas — "
            "spatial planning to segregate polluting land uses including potential "
            "plastic manufacturing, processing or waste-handling facilities."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            'Art. 43(f): "zonas destinadas à implantação de actividades industriais '
            'poluentes e, ou incompatíveis com outras funções e usos do espaço '
            'urbano"; Art. 43(g): "zonas de protecção ambiental e... áreas de '
            'importância ecológica".'
        ),
        "instrument_score": "[auto]",
        "comments": (
            "Polluting-industry zoning instrument — no explicit plastic-sector "
            "targeting. Not coded: landscape/cultural heritage zones (Art. 44(e))."
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Environmental leakage",
        "instrument_description": (
            "Arts. 42(1)(a), 51, 60 — Environmental sustainability in territorial "
            "instruments: PEU establishes environmental sustainability principles "
            "and equitable distribution of infrastructure; soil classification "
            "preserves ecological structure; zoning safeguards regional ecological "
            "and environmental qualities and prevents environmental degradation — "
            "framework for sustainable urban waste-service siting."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            'Art. 42(1)(a): "princípios de sustentabilidade ambiental"; Art. 51(1)(a): '
            '"preservação da estrutura ecológica do território"; Art. 60: zoneamento '
            'to "impedir a sua degradação ambiental e a fomentar o seu uso sustentável".'
        ),
        "instrument_score": "[auto]",
        "comments": (
            "Environmental-sustainability framework — supports waste-infrastructure "
            "siting. Not coded: environmental inventories (Arts. 58–59) as data "
            "instruments without operational waste rules."
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Waste management",
        "instrument_description": (
            "Arts. 77–82 — Fiscalization, offences and sanctions: territorial "
            "planning authority inspects compliance with the Regulation; violations "
            "of planning instruments (including PEU waste-treatment and GPU "
            "sanitation provisions) are subject to administrative, civil, "
            "disciplinary and criminal liability; administrative fines for "
            "non-compliance with instrument elaboration and implementation deadlines."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            'Art. 77: fiscalização do cumprimento do Regulamento; Art. 79: '
            '"violations of territorial planning instrument provisions" subject to '
            "administrative, civil, disciplinary and criminal liability; Art. 82: "
            "administrative penalties for non-compliance including failure to "
            "elaborate/revise instruments within defined deadlines."
        ),
        "instrument_score": "[auto]",
        "comments": (
            "Enforcement instrument for planning-instrument compliance — applies "
            "to waste/sanitation zoning once plans are approved. Not coded: "
            "expropriation procedures (Arts. 70–76)."
        ),
    },
]

COLUMNS = [
    "policy_name",
    "policy_url",
    "policy_year",
    "policy_objective",
    "policy_target",
    "policy_target_text",
    "policy_type",
    "policy_type_justification",
    "policy_integration",
    "policy_sectors_list",
    "policy_circularity",
    "policy_lifecycle_phases_list",
    "policy_budget",
    "policy_budget_text",
    "policy_score",
    "instrument_type",
    "instrument_lifecycle_stage",
    "instrument_description",
    "instrument_in_force",
    "instrument_implementation",
    "instrument_implementation_text",
    "instrument_score",
    "comments",
]

COLUMN_LABELS = {
    "policy_name": "A — policy_name",
    "policy_url": "B — policy_url",
    "policy_year": "C — policy_year",
    "policy_objective": "D — policy_objective",
    "policy_target": "E — policy_target",
    "policy_target_text": "F — policy_target_text",
    "policy_type": "G — policy_type",
    "policy_type_justification": "H — policy_type_justification",
    "policy_integration": "I — policy_integration",
    "policy_sectors_list": "J — policy_sectors_list",
    "policy_circularity": "K — policy_circularity",
    "policy_lifecycle_phases_list": "L — policy_lifecycle_phases_list",
    "policy_budget": "M — policy_budget",
    "policy_budget_text": "N — policy_budget_text",
    "policy_score": "O — policy_score [auto]",
    "instrument_type": "P — instrument_type",
    "instrument_lifecycle_stage": "Q — instrument_lifecycle_stage",
    "instrument_description": "R — instrument_description",
    "instrument_in_force": "S — instrument_in_force",
    "instrument_implementation": "T — instrument_implementation",
    "instrument_implementation_text": "U — instrument_implementation_text",
    "instrument_score": "V — instrument_score [auto]",
    "comments": "W — comments",
}


def build_dataframe():
    return pd.DataFrame([{**POLICY_FIELDS, **i} for i in INSTRUMENTS], columns=COLUMNS)


def export_csv(df, basename):
    path = os.path.join(OUTPUT_DIR, f"{basename}.csv")
    df.rename(columns=COLUMN_LABELS).to_csv(path, index=False, encoding="utf-8-sig")
    return path


def export_xlsx(df, basename):
    path = os.path.join(OUTPUT_DIR, f"{basename}.xlsx")
    labelled = df.rename(columns=COLUMN_LABELS)
    labelled.to_excel(path, index=False, sheet_name="4P Index Coding")
    wb = load_workbook(path)
    ws = wb.active
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = min(
            45, max(12, len(str(ws.cell(1, col).value or "")) * 0.9)
        )
    ws.freeze_panes = "A2"
    wb.save(path)
    return path


def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_para(doc, text):
    doc.add_paragraph(text)


def create_translation_doc():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading(
        "Regulation of the Territorial Planning Law", 0
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("Decree No. 23/2008 of 1 July")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(
        "English translation of: Regulamento da Lei de Ordenamento do Território "
        "(Decreto n.º 23/2008)\n"
    ).italic = True
    meta.add_run(f"Translated for 4P Index policy analysis — {date.today().isoformat()}\n")
    meta.add_run(
        "Source: https://sibmoz.gov.mz/content/uploads/2025/11/"
        "Decreto-n.o-23_2008_Regulamento-da-Lei-de-Ordenamento-do-Territorio.pdf"
    )

    sections = [
        ("Decree No. 23/2008 — Enacting Provisions", (
            "Law No. 19/2007 of 18 July established the legal framework for "
            "Mozambique's Territorial Planning Policy and the legal bases for "
            "national territorial planning instruments.\n\n"
            "Regulatory measures and procedures are required to ensure rational "
            "and sustainable occupation and use of natural resources, valorisation "
            "of regional potential, infrastructure and urban systems, national "
            "cohesion and population security, under Law 19/2007 Art. 30.\n\n"
            "The Council of Ministers decrees:\n\n"
            "Article 1. The Regulation of the Territorial Planning Law, attached "
            "hereto, is approved as an integral part of this Decree.\n\n"
            "Article 2. The Ministers for Environmental Action Coordination, "
            "Finance and Justice shall define appropriate expropriation procedures "
            "under the Territorial Planning Law, this Decree and applicable legislation.\n\n"
            "Article 3. This Decree enters into force ninety days after publication.\n\n"
            "Approved by the Council of Ministers on 13 May 2008. Published "
            "1 July 2008.\n"
            "Signed: Luísa Dias Diogo, Prime Minister."
        )),
        ("Chapter I — General Provisions (summary)", (
            "Art. 1: Definitions (local community, sustainable development, "
            "planning instruments, rural/urban land, territory, etc.).\n\n"
            "Art. 2: Object — establish the legal regime of territorial planning "
            "instruments.\n\n"
            "Art. 3: Scope — entire national territory; regulates relations among "
            "administration levels and public/private subjects including communities.\n\n"
            "Art. 4: Intervention levels and instruments — national (PNDT, Special "
            "Plans), provincial (PPDT), district (PDUT), municipal (PEU, PGU, PPU, "
            "Detail Plans).\n\n"
            "Art. 5: General instruments — soil qualification, classification, "
            "National Land Cadastre, environmental/social/economic inventories, "
            "zoning, geological map, mining cadastre."
        )),
        ("Chapter II — Elaboration Process (summary)", (
            "Arts. 6–11: Minimum phases (objectives, inventory, diagnosis, "
            "alternatives, decision, monitoring, review); sectoral coordination; "
            "competences for elaboration and approval at each level; ratification; "
            "amendment; review; suspension; publicity requirements."
        )),
        ("Section II — Urban Structure Plan (plastic-relevant)", [
            ("Article 42 — PEU objectives", (
                "The Urban Structure Plan (PEU) shall:\n"
                "a) Establish environmental sustainability principles, main access "
                "networks, urban development priorities and general parameters for "
                "municipal territory occupation;\n"
                "b) Eliminate social asymmetries and privileges in locating "
                "infrastructure, services and social facilities;\n"
                "c) Define principles and models for municipal territorial ordering.\n\n"
                "Materialisation of (c) requires definition of:\n"
                "a) Primary accessibility network structure;\n"
                "b) Major surface-water runoff control systems and progressive "
                "execution principles;\n"
                "c) Solid-waste treatment systems and zones for reception and "
                "processing;\n"
                "d) Principles for cemetery construction and location;\n"
                "e) Multifunctional activity-centre network;\n"
                "f) General principles and parameters for public-space use;\n"
                "g) Principles for public/private motor transport and progressive "
                "pedestrianisation."
            )),
            ("Article 43 — PEU content", (
                "PEU elements include: land-occupation forms and rules; territorial "
                "ordering principles; biophysical/ecological characterisation; "
                "demographic structure; economic/social/cultural activities; "
                "detailed mapping of zones for polluting and/or incompatible "
                "industrial activities; environmental protection and ecologically "
                "important areas; road network and equipment distribution; projected "
                "financial needs; graphic maps and schemes."
            )),
        ]),
        ("Section III — General/Partial Urbanisation Plans (plastic-relevant)", [
            ("Article 44 — GPU/PPU objectives", (
                "GPU/PPU objectives include:\n"
                "a) Materialise PEU principles and parameters;\n"
                "b) Population evolution and urban occupation models;\n"
                "c) Public-space reserves;\n"
                "d) Subdivision geometry for urban uses;\n"
                "e) Exceptional landscape/cultural heritage areas;\n"
                "f) Urban areas to requalify with essential infrastructure;\n"
                "g) General and local road structure;\n"
                "h) Location of railways, high-voltage lines, aqueducts, surface-water "
                "and used-water drainage systems and other public-interest infrastructure;\n"
                "i) Pedestrianisation principles;\n"
                "j) Spatial units for partial/detail plans;\n"
                "k) Quantitative/qualitative indicators and urban parameters."
            )),
            ("Article 4 — GPU/PPU definitions", (
                "General Urbanisation Plan (PGU): structures and qualifies all urban "
                "land, defining transport, communications, energy and sanitation "
                "networks plus social facilities.\n\n"
                "Partial Urbanisation Plan (PPU): same requirements for partial urban "
                "territory."
            )),
        ]),
        ("Environmental Instruments and Enforcement (summary)", (
            "Arts. 51–52: Soil classification for ecological structure preservation.\n"
            "Arts. 58–59: Environmental, social and economic inventories as planning "
            "data bases.\n"
            "Art. 60: Zoning to prevent environmental degradation.\n\n"
            "Chapter XII — Arts. 77–82: Fiscalization of Regulation compliance; "
            "violations of planning instruments subject to administrative, civil, "
            "disciplinary and criminal liability; administrative fines."
        )),
        ("Note on Plastics and Solid Waste", (
            "Decree 23/2008 contains no explicit reference to plastics (plástico). "
            "Plastic-relevant content:\n\n"
            "— Art. 42(2)(c): explicit solid-waste treatment systems and "
            "reception/processing zones in PEU (resíduos sólidos);\n"
            "— Arts. 4, 44: sanitation networks and wastewater/stormwater drainage "
            "in GPU/PPU;\n"
            "— Art. 43(f)–(g): polluting industrial zones and environmental "
            "protection areas;\n"
            "— Arts. 42(1)(a), 51, 60: environmental sustainability framework;\n"
            "— Arts. 77–82: fiscalization and sanctions.\n\n"
            "No quantified plastic-reduction, recycling or EPR targets and no "
            "waste-specific budget lines. Operational collection, recycling and "
            "plastic-bag rules remain in Decree 94/2014, Decree 16/2015, etc. "
            "This regulation operationalises Law 19/2007 Art. 10 municipal instruments."
        )),
    ]

    for heading, content in sections:
        add_heading(doc, heading, level=1)
        if isinstance(content, str):
            add_para(doc, content)
        else:
            for subheading, body in content:
                add_heading(doc, subheading, level=2)
                add_para(doc, body)

    doc.add_page_break()
    add_heading(doc, "4P Index Coding Summary", level=1)
    add_para(
        doc,
        f"Policy: {POLICY_FIELDS['policy_name']}\n"
        f"Instruments coded: {len(INSTRUMENTS)}\n"
        f"E — policy_target: {POLICY_FIELDS['policy_target']}\n"
        f"K — policy_circularity: {POLICY_FIELDS['policy_circularity']} "
        f"({POLICY_FIELDS['policy_lifecycle_phases_list']})\n"
        f"See: 4P_Index_Mozambique_Decree_23_2008_LOT_Regulation.xlsx",
    )

    path = os.path.join(
        OUTPUT_DIR, "Mozambique_Decree_23_2008_LOT_Regulation_English_Translation.docx"
    )
    doc.save(path)
    return path


def create_index_html(basename):
    base = (
        "https://github.com/shivrain/Thinkswiss/raw/"
        "cursor/mozambique-decree-23-2008-lot-regulation-04b6/output/mozambique"
    )
    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <title>Mozambique Decree 23/2008 LOT Regulation — 4P Index</title>
  <style>
    body {{ font-family: system-ui, sans-serif; max-width: 800px; margin: 2rem auto; padding: 0 1rem; }}
    .card {{ border: 1px solid #ddd; border-radius: 8px; padding: 1rem; margin: 1rem 0; }}
    a.download {{ display: inline-block; background: #1F4E79; color: #fff; padding: 0.6rem 1.2rem;
      border-radius: 4px; text-decoration: none; margin-top: 0.5rem; }}
  </style>
</head>
<body>
  <h1>4P Index — Decree 23/2008 (LOT Regulation)</h1>
  <p>Regulamento da Lei de Ordenamento do Território. Generated {date.today().isoformat()}.
  Six instruments coded (solid waste, sanitation, pollution zoning).</p>
  <div class="card">
    <a class="download" href="{base}/{basename}.xlsx" download>Download Excel</a><br><br>
    <a class="download" href="{base}/{basename}.csv" download>Download CSV</a><br><br>
    <a class="download" href="{base}/Mozambique_Decree_23_2008_LOT_Regulation_English_Translation.docx" download>Download English translation</a>
  </div>
</body>
</html>"""
    path = os.path.join(OUTPUT_DIR, "index_decree_23_2008_lot_regulation.html")
    with open(path, "w", encoding="utf-8") as f:
        f.write(html)
    return path


def main():
    basename = "4P_Index_Mozambique_Decree_23_2008_LOT_Regulation"
    df = build_dataframe()
    paths = [
        export_csv(df, basename),
        export_xlsx(df, basename),
        create_translation_doc(),
        create_index_html(basename),
    ]
    print("Generated:")
    for p in paths:
        print(f"  {p}")


if __name__ == "__main__":
    main()
