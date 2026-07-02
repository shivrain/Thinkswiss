#!/usr/bin/env python3
"""Generate 4P Index Excel and English translation for Peru DS 009-2019-MINAM (WEEE/RAEE)."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

OUTPUT_DIR = Path("/workspace/output/peru-ds-009-2019-minam")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

BRANCH = "cursor/peru-ds-009-2019-4p-index-d350"
REPO_RAW = f"https://github.com/shivrain/Thinkswiss/raw/{BRANCH}"

POLICY_NAME_ES = (
    "Decreto Supremo N.° 009-2019-MINAM – Régimen Especial de Gestión y Manejo de Residuos "
    "de Aparatos Eléctricos y Electrónicos (RAEE)"
)
POLICY_NAME_EN = (
    "Supreme Decree No. 009-2019-MINAM – Special Regime for the Management and Handling of "
    "Waste Electrical and Electronic Equipment (WEEE)"
)
POLICY_URL = (
    "https://www.gob.pe/institucion/minam/normas-legales/354138-009-2019-minam"
)


def bi(es: str, en: str) -> str:
    return f"{es} / {en}"


POLICY = {
    "policy_name": bi(POLICY_NAME_ES, POLICY_NAME_EN),
    "policy_url": POLICY_URL,
    "policy_year": 2019,
    "policy_objective": bi(
        "Decreto supremo que establece el régimen especial de gestión y manejo de residuos de "
        "aparatos eléctricos y electrónicos (RAEE/WEEE) como bienes priorizados bajo la "
        "responsabilidad extendida del productor (REP) del DL 1278, asignando obligaciones de "
        "recolección, valorización y disposición final a productores mediante sistemas "
        "individuales o colectivos, Planes de Manejo aprobados por MINAM, metas anuales de "
        "recolección, puntos de acopio, operadores autorizados (EO-RS), segregación municipal "
        "y fiscalización OEFA — cubriendo equipos con carcasas y componentes poliméricos.",
        "Supreme decree establishing the special regime for waste electrical and electronic "
        "equipment (WEEE/RAEE) as prioritized goods under extended producer responsibility "
        "(EPR) of DL 1278, assigning collection, valorization and final-disposal obligations "
        "to producers through individual or collective management systems, MINAM-approved "
        "Management Plans, annual collection targets, drop-off points, authorized operators "
        "(EO-RS), municipal segregation and OEFA enforcement — covering equipment with polymer "
        "casings and components.",
    ),
    "policy_target": 1,
    "policy_target_text": bi(
        "Art. 18.2: metas mínimas anuales categorías 3–4 (IT/telecom y electrónica de consumo) "
        "2020–2024: 16%, 19%, 22%, 25%, 28% de línea base (peso). Art. 18.3: categorías 1–2 "
        "(electrodomésticos) 2020–2024: 4%, 7%, 10%, 13%, 16%. Art. 21.1: obligatoriedad "
        "recolección cat. 1–2 desde 1 ene. 2020. DCT Cuarta: PMRAEE obligatorio hasta último "
        "día hábil 2019. DCF Quinta: metas cat. 5 (alumbrado) y 8 (médicos) en 12 meses.",
        "Art. 18.2: minimum annual targets categories 3–4 (IT/telecom and consumer electronics) "
        "2020–2024: 16%, 19%, 22%, 25%, 28% of baseline (weight). Art. 18.3: categories 1–2 "
        "(appliances) 2020–2024: 4%, 7%, 10%, 13%, 16%. Art. 21.1: mandatory collection cat. "
        "1–2 from 1 Jan 2020. Fourth TCP: management plan mandatory by last business day 2019. "
        "Fifth FCP: targets for cat. 5 (lighting) and 8 (medical) within 12 months.",
    ),
    "policy_type": 0.75,
    "policy_type_justification": bi(
        "Decreto Supremo N.° 009-2019-MINAM aprobado por el Poder Ejecutivo (7 nov. 2019, "
        "publicado 8 nov. 2019 en El Peruano), refrendado por MINAM y MTC. Enmendado en "
        "disposiciones complementarias por DS N.° 035-2021-MINAM (2021). Régimen especial "
        "ejecutivo bajo DL 1278 Art. 13 y DS 014-2017-MINAM Art. 84; deroga DS 001-2012-MINAM. "
        "Clasificado 0.75 (decreto supremo/régimen especial).",
        "Supreme Decree No. 009-2019-MINAM approved by the Executive (7 Nov 2019, published "
        "8 Nov 2019 in El Peruano), countersigned by MINAM and MTC. Complementary provisions "
        "amended by Supreme Decree No. 035-2021-MINAM (2021). Executive special regime under "
        "DL 1278 Art. 13 and DS 014-2017-MINAM Art. 84; repeals DS 001-2012-MINAM. Classified "
        "0.75 (supreme decree/special regime).",
    ),
    "policy_integration": 1,
    "policy_sectors_list": bi(
        "manufactura e importación de electrónica, retail, ambiente, gobiernos locales, "
        "operadores de residuos, salud (DIGESA), transportes y comunicaciones (MTC), "
        "certificación ambiental (SENACE), sector público, reciclaje de plásticos y metales",
        "electronics manufacturing and imports, retail, environment, local governments, waste "
        "operators, health (DIGESA), transport and communications (MTC), environmental "
        "certification (SENACE), public sector, plastic and metal recycling",
    ),
    "policy_circularity": 1.0,
    "policy_lifecycle_phases_list": bi(
        "producción, diseño, consumo, recolección, reciclaje, valorización material, disposición",
        "production, design, consumption, collection, recycling, material valorization, disposal",
    ),
    "policy_budget": 0.5,
    "policy_budget_text": bi(
        "DS Art. 3: financiamiento con cargo al presupuesto institucional de los pliegos "
        "involucrados, sin demandar recursos adicionales al Tesoro Público. Art. 16.1m: "
        "presupuesto del PMRAEE a cargo de productores/sistemas colectivos.",
        "DS Art. 3: financing from institutional budgets of involved entities, without "
        "requiring additional Public Treasury resources. Art. 16.1m: WEEE Management Plan "
        "budget borne by producers/collective systems.",
    ),
}

INSTRUMENTS = [
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Collection",
        "instrument_description": bi(
            "DS Art. 1: aprueba el Régimen Especial de Gestión y Manejo de RAEE — 6 Títulos, "
            "35 Artículos, 4 DCT, 6 DCF, 1 DCD y 2 Anexos, parte integrante del decreto.",
            "DS Art. 1: approves the Special Regime for WEEE Management — 6 Titles, 35 Articles, "
            "4 transitory provisions, 6 final provisions, 1 repealing provision and 2 Annexes, "
            "integral part of the decree.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "'Apruébase' — vigencia desde publicación 8 nov. 2019. Deroga DS 001-2012-MINAM.",
            "'Is approved' — effective from publication 8 Nov 2019. Repeals DS 001-2012-MINAM.",
        ),
        "comments": bi(
            "Acto habilitante del régimen REP para residuos electrónicos (componentes plásticos).",
            "Enabling act for EPR regime on e-waste (plastic components).",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Collection",
        "instrument_description": bi(
            "Reglamento Art. 1–2: objeto — régimen especial RAEE como bienes priorizados; "
            "finalidad maximizar eficiencia con valorización como primera finalidad y "
            "disposición final como última.",
            "Regulation Arts. 1–2: purpose — special WEEE regime as prioritized goods; aim to "
            "maximize efficiency with valorization as first purpose and final disposal as last.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "Objeto y jerarquía de manejo declarativa. Marco DL 1278 REP.",
            "Declaratory purpose and management hierarchy. DL 1278 EPR framework.",
        ),
        "comments": bi(
            "Jerarquía valorización > disposición aplicable a flujos con polímeros.",
            "Valorization > disposal hierarchy applicable to polymer-containing streams.",
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": bi(
            "Reglamento Art. 3: ámbito obligatorio para generadores, productores, operadores, "
            "distribuidores y comercializadores de RAEE categorizados en Anexo II (11 categorías "
            "AEE: electrodomésticos, informática, consumo, alumbrado, herramientas, juguetes, "
            "médicos, paneles fotovoltaicos, etc.).",
            "Regulation Art. 3: mandatory scope for generators, producers, operators, "
            "distributors and retailers of WEEE categorized in Annex II (11 EEE categories: "
            "appliances, IT, consumer electronics, lighting, tools, toys, medical, PV panels, "
            "etc.).",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "Aplicación territorial nacional. Exclusiones DL 1278 Art. 27. Anexo II lista "
            "indicativa no excluyente.",
            "National territorial application. DL 1278 Art. 27 exclusions. Annex II indicative "
            "non-exclusive list.",
        ),
        "comments": bi(
            "Equipos con carcasas y componentes plásticos dentro del ámbito.",
            "Equipment with plastic casings and components within scope.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Collection",
        "instrument_description": bi(
            "Reglamento Art. 4: lineamientos — priorizar valorización; REP y responsabilidad "
            "compartida; involucrar comercializadores/distribuidores en recolección; "
            "sensibilización y minimización de RAEE.",
            "Regulation Art. 4: guidelines — prioritize valorization; EPR and shared "
            "responsibility; involve retailers/distributors in collection; awareness and WEEE "
            "minimization.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "Lineamientos orientadores para PMRAEE y sistemas de manejo.",
            "Guiding principles for WEEE Management Plans and management systems.",
        ),
        "comments": bi(
            "Principios REP aplicables a residuos con fracciones poliméricas recuperables.",
            "EPR principles applicable to waste with recoverable polymer fractions.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Collection",
        "instrument_description": bi(
            "Reglamento Art. 5–8: funciones institucionales — MINAM (normar, aprobar metas y "
            "PMRAEE, EO-RS, SINIA); SENACE (EIA plantas valorización); OEFA (fiscalizar "
            "productores, generadores y plantas); municipalidades (recolección selectiva, puntos "
            "de acopio).",
            "Regulation Arts. 5–8: institutional roles — MINAM (regulate, approve targets and "
            "plans, EO-RS, SINIA); SENACE (EIA for valorization plants); OEFA (supervise "
            "producers, generators and plants); municipalities (selective collection, drop-off "
            "points).",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "Roles designados multi-nivel. DCT Segunda: MTC fiscaliza hasta transferencia a "
            "OEFA.",
            "Designated multi-level roles. Second TCP: MTC enforces until transfer to OEFA.",
        ),
        "comments": bi(
            "Gobernanza coordinada del régimen RAEE.",
            "Coordinated WEEE regime governance.",
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": bi(
            "Reglamento Art. 9: definición de productor (fabricante, ensamblador, importador "
            "que pone AEE por primera vez en el mercado) con responsabilidad extendida sobre "
            "todo el ciclo de vida incluyendo posconsumo.",
            "Regulation Art. 9: producer definition (manufacturer, assembler, importer placing "
            "EEE on market for first time) with extended responsibility over full lifecycle "
            "including post-consumption.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "'Responsabilidad extendida' vinculante. Incluye venta a distancia y electrónica. "
            "Exención de daños si entrega a operador autorizado.",
            "Binding 'extended responsibility'. Includes distance and electronic sales. "
            "Liability exemption if delivered to authorized operator.",
        ),
        "comments": bi(
            "Instrumento central REP para e-waste (incl. carcasas plásticas).",
            "Central EPR instrument for e-waste (incl. plastic housings).",
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Collection",
        "instrument_description": bi(
            "Reglamento Art. 10: obligaciones del productor — diseñar/administrar sistemas de "
            "manejo; presentar PMRAEE; cumplir metas de recolección; recibir RAEE sin costo; "
            "informar clientes; entregar a operador autorizado; reportar declaración anual.",
            "Regulation Art. 10: producer obligations — design/administer management systems; "
            "submit WEEE Management Plan; meet collection targets; receive WEEE free of charge; "
            "inform customers; deliver to authorized operator; file annual declaration.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": bi(
            "Obligaciones operativas múltiples con lenguaje imperativo. Financiamiento a cargo "
            "del productor. SIGERSOL para reportes.",
            "Multiple operative obligations with imperative language. Producer-funded. SIGERSOL "
            "for reporting.",
        ),
        "comments": bi(
            "Obligaciones EPR de recolección y valorización de RAEE.",
            "EPR collection and valorization obligations for WEEE.",
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Collection",
        "instrument_description": bi(
            "Reglamento Art. 11: Declaración Anual del productor (individual o colectiva) con "
            "datos de AEE fabricados/importados, RAEE recolectado, % cumplimiento meta, "
            "valorización/disposición, operadores — presentación primeros 15 días hábiles de "
            "abril vía SIGERSOL.",
            "Regulation Art. 11: Annual Producer Declaration (individual or collective) with "
            "manufactured/imported EEE data, WEEE collected, target compliance %, "
            "valorization/disposal, operators — filed first 15 business days of April via "
            "SIGERSOL.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "Plazo anual fijo y contenido cuantificado. DCT Primera: remisión a MINAM mientras "
            "SIGERSOL no operativo.",
            "Fixed annual deadline and quantified content. First TCP: submit to MINAM while "
            "SIGERSOL not operational.",
        ),
        "comments": bi(
            "Monitoreo de cumplimiento de metas de recolección.",
            "Monitoring of collection-target compliance.",
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Collection",
        "instrument_description": bi(
            "Reglamento Art. 12: sistemas de manejo individual o colectivo de RAEE bajo REP; "
            "deben someter PMRAEE a MINAM, garantizar facilidades de entrega, cumplir metas y "
            "entregar RAEE a operador autorizado para valorización.",
            "Regulation Art. 12: individual or collective WEEE management systems under EPR; "
            "must submit Management Plan to MINAM, ensure delivery facilities, meet targets and "
            "deliver WEEE to authorized operator for valorization.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "'Debe' someter plan, garantizar entrega y cumplir metas. Sistemas colectivos con "
            "representante y financiamiento compartido.",
            "'Must' submit plan, ensure delivery and meet targets. Collective systems with "
            "representative and shared financing.",
        ),
        "comments": bi(
            "Mecanismo operativo REP individual/colectivo.",
            "Individual/collective EPR operational mechanism.",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Collection",
        "instrument_description": bi(
            "Reglamento Art. 15–16: Plan de Manejo de RAEE (PMRAEE) — instrumento con línea "
            "base, metas anuales, flujograma, estrategia de recolección, puntos de acopio, "
            "operadores, presupuesto; evaluación y aprobación MINAM.",
            "Regulation Arts. 15–16: WEEE Management Plan (PMRAEE) — instrument with baseline, "
            "annual targets, flowchart, collection strategy, drop-off points, operators, "
            "budget; MINAM evaluation and approval.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "Contenido detallado obligatorio Art. 16.1. Aprobación MINAM requerida. "
            "Actualización anual hasta septiembre (Art. 19).",
            "Mandatory detailed content Art. 16.1. MINAM approval required. Annual update by "
            "September (Art. 19).",
        ),
        "comments": bi(
            "Plan vinculante de recolección y valorización material (plásticos/metales).",
            "Binding plan for collection and material valorization (plastics/metals).",
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Collection",
        "instrument_description": bi(
            "Reglamento Art. 17–18: línea base (promedio 3 años AEE en peso, -10% empaque); "
            "metas mínimas 2020–2024 — cat. 3–4: 16–28%; cat. 1–2: 4–16%; cat. 5–8 y otras "
            "voluntarias salvo DCF Quinta.",
            "Regulation Arts. 17–18: baseline (3-year average EEE weight, -10% packaging); "
            "minimum 2020–2024 targets — cat. 3–4: 16–28%; cat. 1–2: 4–16%; cat. 5–8 and "
            "others voluntary except Fifth FCP.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": bi(
            "Metas cuantificadas porcentuales en texto. MINAM establece metas quinquenales. "
            "Art. 21.1: cat. 1–2 obligatorias desde 2020.",
            "Quantified percentage targets in text. MINAM sets five-year targets. Art. 21.1: "
            "cat. 1–2 mandatory from 2020.",
        ),
        "comments": bi(
            "Objetivos cuantitativos de recolección RAEE por categoría.",
            "Quantitative WEEE collection targets by category.",
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Collection",
        "instrument_description": bi(
            "Reglamento Art. 23: obligaciones distribuidor/comercializador — establecer "
            "gratuitamente puntos de acopio en coordinación con sistemas de manejo; entregar "
            "RAEE acopiados a sistemas; difundir segregación y entrega adecuada.",
            "Regulation Art. 23: distributor/retailer obligations — establish free drop-off "
            "points in coordination with management systems; deliver collected WEEE to systems; "
            "disseminate segregation and proper delivery.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "'Establecer de manera gratuita' puntos de acopio. Aliados estratégicos Art. 12.5.",
            "'Establish free of charge' drop-off points. Strategic allies Art. 12.5.",
        ),
        "comments": bi(
            "Infraestructura de recolección en retail para RAEE.",
            "Retail collection infrastructure for WEEE.",
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Collection",
        "instrument_description": bi(
            "Reglamento Art. 25: obligaciones del generador — minimizar, segregar y almacenar "
            "RAEE; entregar a sistemas de manejo u operadores sin costo; incluir RAEE en Plan "
            "de Minimización si tiene IGA; reportar vía SIGERSOL.",
            "Regulation Art. 25: generator obligations — minimize, segregate and store WEEE; "
            "deliver to management systems or operators free of charge; include WEEE in "
            "Minimization Plan if holding IGA; report via SIGERSOL.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "Obligaciones de segregación y entrega gratuita. Hogares pueden entregar a "
            "distribuidores del sistema.",
            "Segregation and free-delivery obligations. Households may deliver to system "
            "retailers.",
        ),
        "comments": bi(
            "Responsabilidad compartida del generador en cadena RAEE.",
            "Shared generator responsibility in WEEE chain.",
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": bi(
            "Reglamento Art. 27–28: operador de RAEE — persona jurídica con planta de "
            "valorización inscrita en EO-RS; obligaciones de recolectar, transportar, "
            "desmantelar/desensamblar, habilitar materiales, entregar componentes peligrosos a "
            "rellenos de seguridad y reportar declaración anual.",
            "Regulation Arts. 27–28: WEEE operator — legal entity with valorization plant "
            "registered in EO-RS; duties to collect, transport, dismantle/disassemble, enable "
            "materials, deliver hazardous components to security landfills and file annual "
            "declaration.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "Autorización EO-RS MINAM obligatoria. Valorización incluye recuperación de "
            "materiales (plásticos, metales). Solo recibe RAEE de sistemas de manejo.",
            "Mandatory MINAM EO-RS authorization. Valorization includes material recovery "
            "(plastics, metals). Receives WEEE only from management systems.",
        ),
        "comments": bi(
            "Operador de valorización material de fracciones poliméricas y otras.",
            "Material valorization operator for polymer and other fractions.",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Collection",
        "instrument_description": bi(
            "Reglamento Art. 30: operaciones autorizadas vía Registro Autoritativo EO-RS — "
            "recolección y transporte; valorización (desmantelamiento/desensamblaje); "
            "disposición final de componentes no aprovechables.",
            "Regulation Art. 30: operations authorized via EO-RS Authoritative Registry — "
            "collection and transport; valorization (dismantling/disassembly); final disposal "
            "of non-recoverable components.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Tres operaciones tipificadas. Autorización MINAM por registro.",
            "Three typified operations. MINAM authorization via registry.",
        ),
        "comments": bi(
            "Marco regulatorio de operaciones en cadena RAEE.",
            "Regulatory framework for WEEE-chain operations.",
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Collection",
        "instrument_description": bi(
            "Reglamento Art. 32–33: recolección selectiva por operadores EO-RS o logística del "
            "productor; puntos de acopio temporales (campañas) o permanentes con pisos "
            "impermeables y seguridad; coordinación municipal.",
            "Regulation Arts. 32–33: selective collection by EO-RS operators or producer "
            "logistics; temporary (campaign) or permanent drop-off points with impermeable "
            "floors and security; municipal coordination.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "'Se debe realizar' recolección selectiva. Requisitos técnicos acopio temporal/"
            "permanente. No son centros de acopio municipales DS 014-2017.",
            "'Must be carried out' selective collection. Technical requirements for temporary/"
            "permanent collection points. Not municipal collection centers under DS 014-2017.",
        ),
        "comments": bi(
            "Infraestructura de acopio y recolección de RAEE.",
            "WEEE collection and drop-off infrastructure.",
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": bi(
            "Reglamento Art. 34: planta de valorización de RAEE — infraestructura con "
            "desmantelamiento/desensamblaje y acondicionamiento; requiere EIA/IGA aprobado; "
            "MINAM determina características técnicas mínimas conforme NTP.",
            "Regulation Art. 34: WEEE valorization plant — infrastructure for "
            "dismantling/disassembly and conditioning; requires approved EIA/IGA; MINAM sets "
            "minimum technical characteristics per Peruvian technical standards.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "EIA SENACE obligatorio. MINAM 'determina' características mínimas — puede requerir "
            "actos complementarios.",
            "SENACE EIA mandatory. MINAM 'determines' minimum characteristics — may require "
            "complementary acts.",
        ),
        "comments": bi(
            "Instalaciones de reciclaje de componentes plásticos y otros materiales.",
            "Recycling facilities for plastic components and other materials.",
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": bi(
            "Reglamento Art. 35: OEFA fiscaliza y tipifica infracciones/sanciones de todos los "
            "actores RAEE; escala supletoria para otras entidades de fiscalización ambiental.",
            "Regulation Art. 35: OEFA supervises and typifies infractions/sanctions for all "
            "WEEE actors; supplementary scale for other environmental enforcement entities.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "OEFA competente principal. DCF Tercera: cuadro de infracciones en 60 días hábiles.",
            "OEFA main competent authority. Third FCP: infraction table within 60 business days.",
        ),
        "comments": bi(
            "Régimen sancionador del cumplimiento REP.",
            "Sanction regime for EPR compliance.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Collection",
        "instrument_description": bi(
            "Reglamento Art. 13–14: migración entre sistemas colectivos/individuales y cese de "
            "actividades del productor — solicitud MINAM, metas no inferiores al sistema "
            "anterior, documentación SUNAT/SUNARP.",
            "Regulation Arts. 13–14: migration between collective/individual systems and "
            "producer cessation — MINAM application, targets not lower than prior system, "
            "SUNAT/SUNARP documentation.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Procedimientos de transición entre sistemas. Meta mínima preservada Art. 13.4.",
            "Transition procedures between systems. Minimum target preserved Art. 13.4.",
        ),
        "comments": bi(
            "Gobernanza de sistemas colectivos REP.",
            "Governance of collective EPR systems.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Collection",
        "instrument_description": bi(
            "DCT Primera: mientras SIGERSOL no esté implementado, declaraciones anuales "
            "productor y operador se remiten a MINAM con copia a entidad de fiscalización.",
            "First TCP: while SIGERSOL is not implemented, annual producer and operator "
            "declarations are sent to MINAM with copy to enforcement entity.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Mecanismo transitorio de reporte. SIGERSOL es canal definitivo Art. 11/29.",
            "Transitional reporting mechanism. SIGERSOL is definitive channel Arts. 11/29.",
        ),
        "comments": bi(
            "Disposición transitoria de implementación.",
            "Transitory implementation provision.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Collection",
        "instrument_description": bi(
            "DCT Tercera–Cuarta: productores con PMRAEE bajo DS 001-2012 deben actualizar plan "
            "(60 días individual, 120 colectivo); productores sin plan deben presentarlo hasta "
            "último día hábil 2019.",
            "Third–Fourth TCPs: producers with Management Plan under DS 001-2012 must update "
            "(60 days individual, 120 collective); producers without plan must submit by last "
            "business day 2019.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": bi(
            "Plazos fijos de transición desde régimen anterior. 'Indefectiblemente' para "
            "productores sin plan.",
            "Fixed transition deadlines from prior regime. 'Mandatory' for producers without "
            "plan.",
        ),
        "comments": bi(
            "Transición desde DS 001-2012-MINAM.",
            "Transition from DS 001-2012-MINAM.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Collection",
        "instrument_description": bi(
            "DCF Primera: operadores DIGESA (EPS-RS/EC-RS) y MINAM (EO-RS) que recolectan/"
            "transportan RAEE comunican trimestralmente vía Informe de Operador el operador "
            "RAEE destinatario.",
            "First FCP: DIGESA (EPS-RS/EC-RS) and MINAM (EO-RS) operators collecting/"
            "transporting WEEE report quarterly via Operator Report the destination WEEE "
            "operator.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Reporte trimestral obligatorio de cadena de custodia.",
            "Mandatory quarterly chain-of-custody reporting.",
        ),
        "comments": bi(
            "Trazabilidad de RAEE entre operadores.",
            "WEEE traceability between operators.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Collection",
        "instrument_description": bi(
            "DCF Quinta: MINAM en 12 meses emite disposiciones sobre metas anuales de "
            "recolección y sistemas de manejo para categorías 5 (alumbrado) y 8 (aparatos "
            "médicos y equipos de laboratorio).",
            "Fifth FCP: MINAM within 12 months issues provisions on annual collection targets "
            "and management systems for categories 5 (lighting) and 8 (medical devices and "
            "laboratory equipment).",
        ),
        "instrument_in_force": 0,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "MINAM 'emite las disposiciones' en plazo — habilitación para metas cat. 5 y 8. "
            "Verificar DS 035-2021-MINAM y actos posteriores.",
            "MINAM 'issues provisions' within deadline — enabling for cat. 5 and 8 targets. "
            "Verify DS 035-2021-MINAM and subsequent acts.",
        ),
        "comments": bi(
            "Metas cat. 5/8 delegadas — revisar enmienda 2021 DS 035-2021-MINAM.",
            "Cat. 5/8 targets delegated — review 2021 amendment DS 035-2021-MINAM.",
        ),
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Disposal",
        "instrument_description": bi(
            "DCF Sexta: entidad competente regula procedimiento de baja de RAEE generados por "
            "entidades públicas.",
            "Sixth FCP: competent entity regulates decommissioning procedure for WEEE generated "
            "by public entities.",
        ),
        "instrument_in_force": 0,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "'Regula el procedimiento' — habilitación; procedimiento no contenido en DS 009-2019.",
            "'Regulates the procedure' — enabling; procedure not contained in DS 009-2019.",
        ),
        "comments": bi(
            "Baja patrimonial sector público — verificar norma emitida.",
            "Public-sector asset decommissioning — verify issued norm.",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": bi(
            "Anexo I: definiciones — valorización, reciclaje, desmantelamiento, componentes "
            "peligrosos, REP, responsabilidad compartida, etapas de manejo (segregación a "
            "disposición final).",
            "Annex I: definitions — valorization, recycling, dismantling, hazardous "
            "components, EPR, shared responsibility, management stages (segregation to final "
            "disposal).",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Glosario vinculante del régimen. Define reciclaje como transformación material.",
            "Binding regime glossary. Defines recycling as material transformation.",
        ),
        "comments": bi(
            "Definiciones técnicas del manejo RAEE y recuperación de materiales.",
            "Technical definitions of WEEE management and material recovery.",
        ),
    },
]

COLUMNS = [
    ("A", "policy_name"),
    ("B", "policy_url"),
    ("C", "policy_year"),
    ("D", "policy_objective"),
    ("E", "policy_target"),
    ("F", "policy_target_text"),
    ("G", "policy_type"),
    ("H", "policy_type_justification"),
    ("I", "policy_integration"),
    ("J", "policy_sectors_list"),
    ("K", "policy_circularity"),
    ("L", "policy_lifecycle_phases_list"),
    ("M", "policy_budget"),
    ("N", "policy_budget_text"),
    ("O", "policy_score"),
    ("P", "instrument_type"),
    ("Q", "instrument_lifecycle_stage"),
    ("R", "instrument_description"),
    ("S", "instrument_in_force"),
    ("T", "instrument_implementation"),
    ("U", "instrument_implementation_text"),
    ("V", "instrument_score"),
    ("W", "comments"),
]


def policy_score_row(instrument_type: float) -> float:
    return round(
        (
            POLICY["policy_type"]
            + POLICY["policy_integration"]
            + POLICY["policy_circularity"]
            + POLICY["policy_budget"]
            + instrument_type
        )
        / 5,
        3,
    )


def instrument_score_row(instrument_type: float, implementation: float) -> float:
    return round((instrument_type + implementation) / 2, 3)


def build_excel(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "4P Index Coding"
    ws.append([f"{col} — {name}" for col, name in COLUMNS])
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for inst in INSTRUMENTS:
        ws.append(
            [
                POLICY["policy_name"],
                POLICY["policy_url"],
                POLICY["policy_year"],
                POLICY["policy_objective"],
                POLICY["policy_target"],
                POLICY["policy_target_text"],
                POLICY["policy_type"],
                POLICY["policy_type_justification"],
                POLICY["policy_integration"],
                POLICY["policy_sectors_list"],
                POLICY["policy_circularity"],
                POLICY["policy_lifecycle_phases_list"],
                POLICY["policy_budget"],
                POLICY["policy_budget_text"],
                policy_score_row(inst["instrument_type"]),
                inst["instrument_type"],
                inst["instrument_lifecycle_stage"],
                inst["instrument_description"],
                inst["instrument_in_force"],
                inst["instrument_implementation"],
                inst["instrument_implementation_text"],
                instrument_score_row(inst["instrument_type"], inst["instrument_implementation"]),
                inst["comments"],
            ]
        )

    widths = {
        "A": 40, "B": 38, "C": 8, "D": 44, "E": 8, "F": 44, "G": 8, "H": 40,
        "I": 10, "J": 40, "K": 10, "L": 36, "M": 10, "N": 42, "O": 10, "P": 10,
        "Q": 18, "R": 50, "S": 10, "T": 12, "U": 44, "V": 10, "W": 40,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(path)


def add_article(doc: Document, es: str, en: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(es)
    r.bold = True
    doc.add_paragraph(en).paragraph_format.space_after = Pt(8)


def build_word(path: Path) -> None:
    doc = Document()
    for margin in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(doc.sections[0], margin, Inches(1))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t = title.add_run(f"{POLICY_NAME_EN}\n{POLICY_NAME_ES}")
    t.bold = True
    t.font.size = Pt(14)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(
        "Signed: 7 November 2019 | Published: 8 November 2019 (El Peruano)\n"
        "2021 amendment: DS No. 035-2021-MINAM (complementary provisions)\n"
        f"Source: {POLICY_URL}"
    )

    doc.add_heading("English Translation of Key Provisions", 1)
    doc.add_paragraph(
        "English translation of principal provisions of Supreme Decree No. 009-2019-MINAM "
        "on the special WEEE/RAEE management regime under extended producer responsibility."
    )

    sections = [
        (
            "Supreme Decree Art. 1 — Approval",
            "Apruébase el Régimen Especial de Gestión y Manejo de Residuos de Aparatos "
            "Eléctricos y Electrónicos... treinta y cinco (35) Artículos... y dos (2) Anexos.",
            "The Special Regime for WEEE Management is approved... thirty-five (35) Articles... "
            "and two (2) Annexes.",
        ),
        (
            "Regulation Arts. 1–2 — Purpose",
            "Establecer un régimen especial para la gestión y manejo de los RAEE... primera "
            "finalidad su valorización y como última, la disposición final.",
            "Establish a special regime for WEEE management... first purpose valorization and "
            "last, final disposal.",
        ),
        (
            "Regulation Art. 9–10 — Extended producer responsibility",
            "El productor tiene responsabilidad extendida sobre el AEE durante todo su ciclo "
            "de vida... diseñar, implementar y administrar sistemas de manejo... recibir sin "
            "costo los RAEE.",
            "The producer has extended responsibility over EEE throughout its lifecycle... "
            "design, implement and administer management systems... receive WEEE free of charge.",
        ),
        (
            "Regulation Art. 18 — Collection targets",
            "Metas mínimas 2020–2024: categorías 3–4 de 16% a 28%; categorías 1–2 de 4% a "
            "16% de la línea base en peso.",
            "Minimum 2020–2024 targets: categories 3–4 from 16% to 28%; categories 1–2 from "
            "4% to 16% of weight baseline.",
        ),
        (
            "Regulation Art. 23 — Retail collection points",
            "Establecer de manera gratuita... puntos de acopio de RAEE... entregar los RAEE "
            "acopiados a los sistemas de manejo.",
            "Establish free of charge... WEEE drop-off points... deliver collected WEEE to "
            "management systems.",
        ),
        (
            "Regulation Art. 27–28 — Valorization operators",
            "Recolectar, transportar y valorizar (desmantelar/desensamblar) los RAEE... "
            "habilitar materiales a partir de RAEE para su aprovechamiento.",
            "Collect, transport and valorize (dismantle/disassemble) WEEE... enable materials "
            "from WEEE for recovery.",
        ),
        (
            "Regulation Art. 35 — Enforcement",
            "El OEFA ejerce la fiscalización ambiental y tipifica... infracciones y establece "
            "las sanciones correspondientes.",
            "OEFA carries out environmental enforcement and typifies... infractions and "
            "establishes corresponding sanctions.",
        ),
    ]

    for heading, es, en in sections:
        doc.add_heading(heading, 2)
        add_article(doc, es, en)

    doc.add_paragraph()
    note = doc.add_paragraph(
        "Note: DS 009-2019-MINAM implements DL 1278 Art. 13 EPR for prioritized goods, "
        "replacing DS 001-2012-MINAM. WEEE streams commonly contain recoverable polymer "
        "casings and components. Complementary provisions were amended by DS 035-2021-MINAM."
    )
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(path)


def main() -> None:
    excel = OUTPUT_DIR / "Peru_DS_009_2019_MINAM_4P_Index_Coding.xlsx"
    word = OUTPUT_DIR / "Peru_DS_009_2019_MINAM_English_Translation.docx"
    build_excel(excel)
    build_word(word)
    print(f"Created: {excel}")
    print(f"Created: {word}")
    print(f"Instruments: {len(INSTRUMENTS)}")
    print()
    print("Download links:")
    print(
        f"  Excel: {REPO_RAW}/output/peru-ds-009-2019-minam/"
        "Peru_DS_009_2019_MINAM_4P_Index_Coding.xlsx"
    )
    print(
        f"  Word:  {REPO_RAW}/output/peru-ds-009-2019-minam/"
        "Peru_DS_009_2019_MINAM_English_Translation.docx"
    )


if __name__ == "__main__":
    main()
