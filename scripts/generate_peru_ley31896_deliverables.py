#!/usr/bin/env python3
"""Generate 4P Index Excel and English translation for Peru Law 31896."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

OUTPUT_DIR = Path("/workspace/output/peru-ley-31896")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

BRANCH = "cursor/peru-ley-31896-4p-index-d350"
REPO_RAW = f"https://github.com/shivrain/Thinkswiss/raw/{BRANCH}"

POLICY_NAME_ES = (
    "Ley N.° 31896 – Ley que modifica el Decreto Legislativo 1278 e introduce la "
    "industrialización del reciclaje"
)
POLICY_NAME_EN = (
    "Law No. 31896 – Law that Amends Legislative Decree No. 1278, Introducing "
    "Industrialization of Recycling"
)
POLICY_URL = (
    "https://www.gob.pe/institucion/minam/normas-legales/4728521-31896"
)


def bi(es: str, en: str) -> str:
    return f"{es} / {en}"


POLICY = {
    "policy_name": bi(POLICY_NAME_ES, POLICY_NAME_EN),
    "policy_url": POLICY_URL,
    "policy_year": 2023,
    "policy_objective": bi(
        "Ley del Congreso que modifica el DL 1278 (Ley de Gestión Integral de Residuos "
        "Sólidos) para introducir la industrialización del reciclaje: prioriza inversión "
        "pública, privada y mixta en infraestructura de valorización; refuerza funciones "
        "MINAM sobre formulación, monitoreo y evaluación del PLANRES con medidas "
        "correctivas; y habilita gobiernos regionales a elaborar programas de inversión "
        "para infraestructuras de valorización (incl. reciclaje industrial de plásticos y "
        "otros residuos aprovechables).",
        "Congress law amending DL 1278 (Comprehensive Solid Waste Management Law) to "
        "introduce recycling industrialization: prioritizes public, private and mixed "
        "investment in valorization infrastructure; strengthens MINAM functions on "
        "PLANRES formulation, monitoring and evaluation with corrective measures; and "
        "enables regional governments to develop investment programs for valorization "
        "infrastructure (incl. industrial recycling of plastics and other recoverables).",
    ),
    "policy_target": 1,
    "policy_target_text": bi(
        "DCF Única: Poder Ejecutivo adecúa Reglamento DL 1278 (DS 014-2017-MINAM) en 30 "
        "días calendario. Art. 15(b): MINAM informa anualmente al Congreso y publica en "
        "portal web resultados de implementación del PLANRES y medidas correctivas. "
        "Art. 21(a): gobiernos regionales elaboran y ponen en marcha programas de "
        "inversión para infraestructuras de valorización.",
        "DCF Única: Executive adapts DL 1278 Regulation (DS 014-2017-MINAM) within 30 "
        "calendar days. Art. 15(b): MINAM annually reports to Congress and publishes on "
        "its website PLANRES implementation results and corrective measures. "
        "Art. 21(a): regional governments develop and launch investment programs for "
        "valorization infrastructure.",
    ),
    "policy_type": 1.0,
    "policy_type_justification": bi(
        "Ley N.° 31896 aprobada por el Congreso de la República (21 set. 2023), promulgada "
        "10 oct. 2023 y publicada 11 oct. 2023 en El Peruano. Modifica artículos del DL "
        "1278 (decreto legislativo delegado). No enmendada. Clasificado 1.0 (legislación "
        "parlamentaria).",
        "Law No. 31896 enacted by Congress (21 Sep 2023), promulgated 10 Oct 2023 and "
        "published 11 Oct 2023 in El Peruano. Amends articles of DL 1278 (delegated "
        "legislative decree). Not amended. Classified 1.0 (parliamentary legislation).",
    ),
    "policy_integration": 0.75,
    "policy_sectors_list": bi(
        "ambiente, industria, gobiernos regionales, gobiernos locales, inversión pública/"
        "privada, reciclaje, valorización",
        "environment, industry, regional governments, local governments, public/private "
        "investment, recycling, valorization",
    ),
    "policy_circularity": 1,
    "policy_lifecycle_phases_list": bi(
        "reciclaje, valorización, disposición",
        "recycling, valorization, disposal",
    ),
    "policy_budget": 0.75,
    "policy_budget_text": bi(
        "Art. 6(e): promoción de inversión pública, privada y mixta en infraestructura de "
        "valorización. Art. 21(a): programas de inversión pública, mixta o privada "
        "regionales para infraestructuras de valorización. Sin asignación presupuestal "
        "específica en la ley.",
        "Art. 6(e): promotion of public, private and mixed investment in valorization "
        "infrastructure. Art. 21(a): regional public, mixed or private investment "
        "programs for valorization infrastructure. No specific budget allocation in the law.",
    ),
}

INSTRUMENTS = [
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": bi(
            "Artículo Único: modifica los artículos 6 (literal e), 15 (literal b) y 21 "
            "(literal a) del Decreto Legislativo 1278 para introducir la industrialización "
            "del reciclaje.",
            "Single Article: amends Articles 6 (lit. e), 15 (lit. b) and 21 (lit. a) of "
            "Legislative Decree 1278 to introduce recycling industrialization.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Acto habilitante de las tres modificaciones al DL 1278. Vigencia desde "
            "publicación 11 oct. 2023.",
            "Enabling act for the three DL 1278 amendments. Effective from publication "
            "11 Oct 2023.",
        ),
        "comments": bi(
            "Ley marco de industrialización del reciclaje en el DL 1278.",
            "Framework law for recycling industrialization in DL 1278.",
        ),
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": bi(
            "DL 1278 Art. 6(e) modificado: lineamiento de gestión integral — fomentar la "
            "valorización de residuos sólidos, priorizando promoción de inversión pública, "
            "privada y mixta en infraestructura de valorización, y adopción complementaria "
            "de prácticas de tratamiento y adecuada disposición final.",
            "Amended DL 1278 Art. 6(e): integrated-management guideline — foster solid-waste "
            "valorization, prioritizing promotion of public, private and mixed investment in "
            "valorization infrastructure, and complementary adoption of treatment practices "
            "and adequate final disposal.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "Lineamiento declarativo del DL 1278. Orienta política nacional hacia "
            "industrialización del reciclaje. Implementación vía PLANRES e inversión.",
            "Declaratory DL 1278 guideline. Orients national policy toward recycling "
            "industrialization. Implementation via PLANRES and investment.",
        ),
        "comments": bi(
            "Base legal para infraestructura industrial de reciclaje de plásticos y otros "
            "residuos.",
            "Legal basis for industrial recycling infrastructure for plastics and other waste.",
        ),
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": bi(
            "DL 1278 Art. 15(b) modificado (1ª parte): MINAM formula y aprueba el PLANRES "
            "de obligatorio cumplimiento, incluyendo metas, estrategias y acciones para "
            "universalización del servicio de limpieza pública, formalización de recicladores "
            "y promoción de minimización y valorización de residuos.",
            "Amended DL 1278 Art. 15(b) (1st part): MINAM formulates and approves the "
            "mandatory PLANRES, including targets, strategies and actions for universal "
            "public-cleaning service, recycler formalization and promotion of waste "
            "minimization and valorization.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Función rectora MINAM reforzada. PLANRES vinculado a industrialización del "
            "reciclaje. Coordinación con autoridades sectoriales.",
            "Strengthened MINAM steering function. PLANRES linked to recycling "
            "industrialization. Coordination with sector authorities.",
        ),
        "comments": bi(
            "PLANRES como instrumento de planificación de capacidad industrial de "
            "valorización (incl. plásticos).",
            "PLANRES as planning instrument for industrial valorization capacity (incl. plastics).",
        ),
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": bi(
            "DL 1278 Art. 15(b) modificado (2ª parte): MINAM monitorea y evalúa "
            "implementación del PLANRES y dispone medidas correctivas como ente rector; "
            "autoridades competentes obligadas a remitir información requerida bajo "
            "responsabilidad.",
            "Amended DL 1278 Art. 15(b) (2nd part): MINAM monitors and evaluates PLANRES "
            "implementation and issues corrective measures as national steering authority; "
            "competent authorities must submit required information under liability.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Mecanismo de cumplimiento y accountability sobre metas de valorización. "
            "Información obligatoria de autoridades.",
            "Compliance and accountability mechanism for valorization targets. "
            "Mandatory authority reporting.",
        ),
        "comments": bi(
            "Refuerza rol de supervisión MINAM sobre reciclaje/valorización a escala "
            "industrial.",
            "Strengthens MINAM oversight role over industrial-scale recycling/valorization.",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": bi(
            "DL 1278 Art. 15(b) modificado (3ª parte): MINAM informa anualmente a la "
            "Comisión de Pueblos Andinos, Amazónicos y Afroperuanos, Ambiente y Ecología "
            "del Congreso, y publica en portal web resultados de implementación del "
            "PLANRES y medidas correctivas dispuestas.",
            "Amended DL 1278 Art. 15(b) (3rd part): MINAM annually reports to the Congress "
            "Commission on Andean, Amazonian and Afro-Peruvian Peoples, Environment and "
            "Ecology, and publishes on its website PLANRES implementation results and "
            "corrective measures issued.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Reporte anual obligatorio. Transparencia pública de avances en "
            "industrialización del reciclaje.",
            "Mandatory annual reporting. Public transparency on recycling industrialization "
            "progress.",
        ),
        "comments": bi(
            "Rendición de cuentas parlamentaria sobre metas de valorización.",
            "Parliamentary accountability on valorization targets.",
        ),
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": bi(
            "DL 1278 Art. 21(a) modificado: gobiernos regionales elaboran y ponen en marcha "
            "programas de inversión pública, mixta o privada para implementación de "
            "infraestructuras de residuos sólidos, como infraestructuras de valorización, "
            "en coordinación con municipalidades provinciales.",
            "Amended DL 1278 Art. 21(a): regional governments develop and launch public, "
            "mixed or private investment programs for implementation of solid-waste "
            "infrastructure, such as valorization infrastructure, in coordination with "
            "provincial municipalities.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.25,
        "instrument_implementation_text": bi(
            "Competencia regional habilitada. Sin plazo específico en la ley. Depende de "
            "programas de inversión regionales.",
            "Regional competence enabled. No specific deadline in the law. Depends on "
            "regional investment programs.",
        ),
        "comments": bi(
            "Habilita plantas/industrias de valorización regional (reciclaje industrial de "
            "plásticos PET, PE, PP, etc.).",
            "Enables regional valorization plants/industries (industrial recycling of PET, "
            "PE, PP plastics, etc.).",
        ),
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": bi(
            "DCF Única: Poder Ejecutivo adecúa el Reglamento del DL 1278 (DS 014-2017-MINAM) "
            "y dicta disposiciones necesarias para implementación de la ley, en plazo no "
            "mayor de 30 días calendario desde su vigencia.",
            "DCF Única: Executive adapts the DL 1278 Regulation (DS 014-2017-MINAM) and "
            "issues necessary provisions for law implementation, within no more than 30 "
            "calendar days of its entry into force.",
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": bi(
            "Plazo vencido (nov. 2023). Adecuación reglamentaria pendiente de verificación "
            "de DS posterior específico.",
            "Deadline elapsed (Nov 2023). Regulatory adaptation subject to verification of "
            "subsequent specific supreme decree.",
        ),
        "comments": bi(
            "Mandato de desarrollo reglamentario para operacionalizar industrialización.",
            "Regulatory development mandate to operationalize industrialization.",
        ),
    },
]

COLUMNS = [
    ("A", "policy_name"),
    ("B", "policy_url"),
    ("C", "policy_year"),
    ("D", "policy_objective"),
    ("E", "policy_target"),
    ("F", "policy_target_text"),
    ("G", "policy_type"),
    ("H", "policy_type_justification"),
    ("I", "policy_integration"),
    ("J", "policy_sectors_list"),
    ("K", "policy_circularity"),
    ("L", "policy_lifecycle_phases_list"),
    ("M", "policy_budget"),
    ("N", "policy_budget_text"),
    ("O", "policy_score"),
    ("P", "instrument_type"),
    ("Q", "instrument_lifecycle_stage"),
    ("R", "instrument_description"),
    ("S", "instrument_in_force"),
    ("T", "instrument_implementation"),
    ("U", "instrument_implementation_text"),
    ("V", "instrument_score"),
    ("W", "comments"),
]


def policy_score_row(instrument_type: float) -> float:
    return round(
        (
            POLICY["policy_type"]
            + POLICY["policy_integration"]
            + POLICY["policy_circularity"]
            + POLICY["policy_budget"]
            + instrument_type
        )
        / 5,
        3,
    )


def instrument_score_row(instrument_type: float, implementation: float) -> float:
    return round((instrument_type + implementation) / 2, 3)


def build_excel(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "4P Index Coding"
    ws.append([f"{col} — {name}" for col, name in COLUMNS])
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for inst in INSTRUMENTS:
        ws.append(
            [
                POLICY["policy_name"],
                POLICY["policy_url"],
                POLICY["policy_year"],
                POLICY["policy_objective"],
                POLICY["policy_target"],
                POLICY["policy_target_text"],
                POLICY["policy_type"],
                POLICY["policy_type_justification"],
                POLICY["policy_integration"],
                POLICY["policy_sectors_list"],
                POLICY["policy_circularity"],
                POLICY["policy_lifecycle_phases_list"],
                POLICY["policy_budget"],
                POLICY["policy_budget_text"],
                policy_score_row(inst["instrument_type"]),
                inst["instrument_type"],
                inst["instrument_lifecycle_stage"],
                inst["instrument_description"],
                inst["instrument_in_force"],
                inst["instrument_implementation"],
                inst["instrument_implementation_text"],
                instrument_score_row(inst["instrument_type"], inst["instrument_implementation"]),
                inst["comments"],
            ]
        )

    widths = {
        "A": 40, "B": 38, "C": 8, "D": 44, "E": 8, "F": 44, "G": 8, "H": 40,
        "I": 10, "J": 40, "K": 10, "L": 36, "M": 10, "N": 42, "O": 10, "P": 10,
        "Q": 18, "R": 50, "S": 10, "T": 12, "U": 44, "V": 10, "W": 40,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(path)


def add_article(doc: Document, es: str, en: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(es)
    r.bold = True
    doc.add_paragraph(en).paragraph_format.space_after = Pt(8)


def build_word(path: Path) -> None:
    doc = Document()
    for margin in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(doc.sections[0], margin, Inches(1))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t = title.add_run(f"{POLICY_NAME_EN}\n{POLICY_NAME_ES}")
    t.bold = True
    t.font.size = Pt(14)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(
        "Congress enacted: 21 September 2023 | Promulgated: 10 October 2023 | "
        "Published: 11 October 2023 (El Peruano)\n"
        "Not amended\n"
        f"Source: {POLICY_URL}"
    )

    doc.add_heading("English Translation of Key Provisions", 1)
    doc.add_paragraph(
        "English translation of principal provisions of Law No. 31896 amending Legislative "
        "Decree No. 1278 to introduce recycling industrialization."
    )

    sections = [
        (
            "Single Article — Amendments",
            "Modifica los artículos 6 (e), 15 (b) y 21 (a) del Decreto Legislativo 1278.",
            "Amends Articles 6 (e), 15 (b) and 21 (a) of Legislative Decree 1278.",
        ),
        (
            "Art. 6(e) — Valorization guideline",
            "Fomentar la valorización de residuos sólidos, priorizando la promoción de la "
            "inversión pública, privada y mixta en infraestructura de valorización.",
            "Foster solid-waste valorization, prioritizing promotion of public, private and "
            "mixed investment in valorization infrastructure.",
        ),
        (
            "Art. 15(b) — PLANRES and MINAM oversight",
            "MINAM formula y aprueba PLANRES; monitorea y evalúa su implementación; dispone "
            "medidas correctivas; informa anualmente al Congreso y publica resultados.",
            "MINAM formulates and approves PLANRES; monitors and evaluates its implementation; "
            "issues corrective measures; annually reports to Congress and publishes results.",
        ),
        (
            "Art. 21(a) — Regional investment",
            "Gobiernos regionales elaboran y ponen en marcha programas de inversión pública, "
            "mixta o privada para infraestructuras de valorización.",
            "Regional governments develop and launch public, mixed or private investment "
            "programs for valorization infrastructure.",
        ),
        (
            "DCF Única — Regulatory adaptation",
            "Poder Ejecutivo adecúa el Reglamento del DL 1278 en 30 días calendario.",
            "Executive adapts the DL 1278 Regulation within 30 calendar days.",
        ),
    ]

    for heading, es, en in sections:
        doc.add_heading(heading, 2)
        add_article(doc, es, en)

    doc.add_paragraph()
    note = doc.add_paragraph(
        "Note: Law 31896 introduces recycling industrialization into Peru's flagship solid-waste "
        "law (DL 1278). It strengthens MINAM's PLANRES monitoring role and enables regional "
        "valorization infrastructure investment, relevant to industrial-scale plastics recycling."
    )
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(path)


def main() -> None:
    excel = OUTPUT_DIR / "Peru_Ley_31896_4P_Index_Coding.xlsx"
    word = OUTPUT_DIR / "Peru_Ley_31896_English_Translation.docx"
    build_excel(excel)
    build_word(word)
    print(f"Created: {excel}")
    print(f"Created: {word}")
    print(f"Instruments: {len(INSTRUMENTS)}")
    print()
    print("Download links:")
    print(
        f"  Excel: {REPO_RAW}/output/peru-ley-31896/"
        "Peru_Ley_31896_4P_Index_Coding.xlsx"
    )
    print(
        f"  Word:  {REPO_RAW}/output/peru-ley-31896/"
        "Peru_Ley_31896_English_Translation.docx"
    )


if __name__ == "__main__":
    main()
