#!/usr/bin/env python3
"""Generate 4P Index Excel for Avfallsforskriften Kap. 1 (WEEE/EE-avfall)."""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

POLICY = {
    "policy_name": (
        "Avfallsforskriften – Kap. 1: Kasserte elektriske og elektroniske produkter (EE-avfall) "
        "(Waste Regulation – Ch. 1: Discarded electrical and electronic equipment (WEEE))"
    ),
    "policy_url": "https://lovdata.no/dokument/SF/forskrift/2004-06-01-930/KAPITTEL_1",
    "policy_year": 2025,
    "policy_objective": (
        "To prevent and reduce environmental and health problems from discarded electrical and electronic "
        "equipment (WEEE), including plastic housings and components, through separate collection, sorting "
        "and treatment of hazardous materials, high recovery of remaining fractions, preparation for reuse and "
        "material recycling, and nationwide extended producer responsibility (EPR) financing."
    ),
    "policy_target": 1,
    "policy_target_text": (
        '"For EE waste in product group 1 and 4 at least 85 % of collected waste shall be recovered, of which '
        'at least 80 % shall be prepared for reuse or materially recovered." (§1-18a(a)); '
        '"For EE waste in product groups 5 and 6 at least 75 % shall be recovered, of which at least 55 % shall '
        'be prepared for reuse or materially recovered." (§1-18a(c))'
    ),
    "policy_type": 0.75,
    "policy_type_justification": (
        "Executive regulation (forskrift) issued 1 June 2004 under forurensningsloven; Chapter 1 implements "
        "EU WEEE Directive 2012/19/EU. Substantially revised 16 December 2015 (in force 1 January 2016). "
        "Parent Avfallsforskriften last amended December 2025 (FOR-2025-12-18-2839); §1-25a added November "
        "2024 incorporating EU Regulation 2019/290 on producer registration."
    ),
    "policy_integration": 1,
    "policy_sectors_list": (
        "electronics, retail, waste management, municipalities, households, producers, recycling, "
        "manufacturing, consumption"
    ),
    "policy_circularity": 1,
    "policy_lifecycle_phases_list": "consumption, recycling, disposal, end-of-life, environmental leakage",
    "policy_budget": 1,
    "policy_budget_text": (
        '"Producers shall finance collection, sorting and treatment of EE waste through membership in a '
        'collectively or individually financed return company approved by the Norwegian Environment Agency." '
        "(§1-10); approved return companies shall pay approval and annual fees to the Treasury (§1-13a, §1-13b)"
    ),
}

INSTRUMENTS = [
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§1-1: Scope—chapter regulates receipt, collection, recovery and other treatment of electrical and "
            "electronic waste (EE-avfall/WEEE), including plastic-containing consumer electronics, with "
            "exclusions for batteries, vehicles, military equipment, medical devices and other listed categories."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            '"The provisions of this chapter regulate receipt, collection, recovery and other treatment of '
            'electrical and electronic waste (EE waste)." (§1-1)'
        ),
        "comments": "Defines national WEEE EPR scope; operative framework for plastics in EEE end-of-life streams.",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§1-2: Purpose—to prevent and reduce environmental and health problems from WEEE through separate "
            "collection, sorting and treatment of hazardous-fraction materials/components, high recovery of "
            "remaining fractions, and preparation for reuse and material recycling where justified."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            '"The purpose … is to prevent and reduce environmental and health problems that EE waste causes. '
            'This shall be done through separate collection, sorting and treatment of materials and components '
            'that are hazardous waste, and a high degree of recovery of other parts of the waste." (§1-2)'
        ),
        "comments": "Policy objective instrument; frames circular-economy intent for WEEE including plastics.",
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": (
            "§1-3a: Prohibition—no person may collect WEEE from a collection point without written agreement "
            "with the responsible party, except actors collecting on behalf of approved return companies."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            '"No one may collect EE waste from a collection point without a written agreement with the person '
            'responsible for the collection point." (§1-3a)'
        ),
        "comments": "Regulatory ban (må ikke); protects authorised EPR collection chains.",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§1-4: Retailer take-back—retailers shall accept household WEEE free of charge in-store or nearby; "
            "accept commercial WEEE free on purchase of equivalent products; large retailers (>400 m²) must "
            "accept all small electronics <25 cm; distance-sales retailers must operate effective return systems."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            '"The retailer shall accept EE waste that is household waste free of charge in return at the shop '
            'premises or at an equivalent place in the immediate vicinity." (§1-4); '
            '"When EE products are sold or delivered outside shop premises … the retailer shall establish an '
            'effective system for dispatch and receipt of equivalent quantities of EE waste." (§1-4)'
        ),
        "comments": "Core EPR collection instrument at point of sale; mandatory (skal).",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": (
            "§1-5: Retailer duties on received WEEE—sort separately from other waste, store safely without "
            "reducing reuse/material-recovery potential, prevent illegal collection, and allow collection only "
            "by approved return companies or agreed preparers for reuse."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            '"The retailer shall ensure that received EE waste is sorted from other waste and stored safely." '
            "(§1-5); reuse and material-recovery potential shall not be reduced (§1-5)"
        ),
        "comments": "Operational waste-handling duty supporting plastic/component recovery downstream.",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§1-7: Municipal take-back—municipalities shall ensure sufficient WEEE reception capacity and "
            "accept household WEEE free of charge; may charge for commercial WEEE reception and storage."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            '"The municipality shall ensure that a sufficient offer exists for receipt of EE waste." (§1-7); '
            '"The municipality shall accept household waste free of charge." (§1-7)'
        ),
        "comments": "Nationwide municipal collection infrastructure; mandatory (skal).",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§1-10: Producer EPR financing—producers shall finance WEEE collection, sorting and treatment "
            "through membership in a collectively or individually financed return company approved by "
            "Miljødirektoratet, covering all EE products they import or manufacture in Norway."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            '"Producers shall finance collection, sorting and treatment of EE waste through membership in a '
            'collectively or individually financed return company that is approved by the Norwegian Environment '
            'Agency, cf. §1-13." (§1-10)'
        ),
        "comments": "Central EPR economic instrument; binding producer duty (skal).",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§1-13: Return-company approval—return companies must register as legal entities and be approved "
            "by Miljødirektoratet; approval requires demonstrating an EPR system capable of meeting chapter "
            "requirements, including certification under Annex 2."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            '"Return companies shall be approved by the Norwegian Environment Agency." (§1-13); '
            '"To be approved, return companies must demonstrate … that they will establish a return system '
            'capable of meeting the requirements in this chapter with annexes." (§1-13)'
        ),
        "comments": "Governance gatekeeper for EPR operators; Miljødirektoratet kan trekke godkjenning.",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": (
            "§1-14: Collective return-company collection duties—free nationwide collection from retailers and "
            "municipalities, free reception from business generators, fulfil Annex 1 collection shares, ensure "
            "proper treatment under §1-18, and maintain six months' financial reserves."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            '"Collectively financed return companies shall: a. ensure free collection of EE waste from retailers '
            'and municipalities … c. ensure free reception of EE waste from businesses … e. ensure that '
            'collected EE waste is treated properly in accordance with §1-18." (§1-14)'
        ),
        "comments": "Operative nationwide collection/recovery system; Miljødirektoratet kan pålegge henting.",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": (
            "§1-18: Return-company treatment chain—ensure licensed collection/transport/treatment, preserve "
            "reuse and material-recovery potential (including plastic fractions), meet §1-18a recovery rates, "
            "handle data-bearing equipment securely, and document treatment through final recovery."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            '"Return companies shall ensure: … b. that collection and transport do not reduce the possibility '
            'for reuse, material recovery and sorting out of components with hazardous substances in EE waste … '
            'd. that EE waste is recovered, preferably by preparation for reuse or material recovery … The '
            'requirements for recovery rates in §1-18a shall be met." (§1-18)'
        ),
        "comments": "Core treatment and recovery obligation along the WEEE value chain.",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": (
            "§1-18a: Quantitative recovery targets by product group—e.g. ≥85 % recovery (≥80 % reuse/material "
            "recovery) for groups 1 and 4; ≥75 % recovery (≥55 % reuse/material recovery) for groups 5 and 6; "
            "≥80 % material recovery for light sources (group 3)."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            '"Return companies shall annually meet the following requirements for recovery rates for the '
            'individual product groups: a. For EE waste in product group 1 and 4 at least 85 % … at least 80 % '
            '… prepared for reuse or materially recovered … c. For EE waste in product groups 5 and 6 at least '
            '75 % … at least 55 % … prepared for reuse or materially recovered." (§1-18a)'
        ),
        "comments": "Binding quantitative targets; includes small electronics and large appliances with plastics.",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": (
            "§1-22: Treatment requirements—WEEE shall be treated using best available techniques; fluids and "
            "listed components (including plastics with brominated flame retardants, circuit boards, external "
            "cables) shall be manually removed as a first step unless equally safe mechanical/chemical processes "
            "are documented; hazardous fractions handled under Ch. 11."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            '"EE waste shall be treated properly and using the best available techniques." (§1-22); '
            '"The following materials, mixtures and components shall be removed manually as a first step … '
            'k. Plastic with brominated flame retardants." (§1-22 fourth paragraph)'
        ),
        "comments": "Direct plastics-relevant treatment rule; mandatory component stripping including plastic fractions.",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§1-19: Reporting—return companies shall report semi-annually and annually to the Produsentansvar "
            "register on collected and treated WEEE quantities by municipality, product group, treatment method "
            "and country; annual auditor statements on total collected volumes."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            '"Return companies shall report semi-annually, by 15 February and 15 August, electronically to the '
            'Produsentansvar register." (§1-19); annual reporting on treated quantities and treatment plants '
            "(§1-19)"
        ),
        "comments": "Monitoring and transparency instrument; mandatory reporting (skal).",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Consumption",
        "instrument_description": (
            "§1-20: Public information—return companies shall inform households and businesses that WEEE must "
            "not be mixed with other waste, explain the WEEE symbol, free take-back points, and hazards from "
            "hazardous substances; conduct regular nationwide information campaigns."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            '"Return companies shall inform households and businesses that EE waste must not be discarded '
            'together with other waste … and that it is received free of charge." (§1-20); '
            '"Return companies shall carry out regular nationwide information campaigns." (§1-20)'
        ),
        "comments": "Behaviour-change instrument supporting separate collection of plastic-rich EEE.",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§1-25: Produsentansvar register—Miljødirektoratet operates a national producer register compiling "
            "import/production, membership, collection and treatment data; identifies non-compliant producers "
            "and representatives; calculates national product supply by product group."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            '"The Norwegian Environment Agency owns a register that shall comprise all producers of EE products." '
            "(§1-25); register shall receive and compile data on collected and treated EE waste by product "
            "group and municipality (§1-25(d))"
        ),
        "comments": "National EPR data infrastructure; Miljødirektoratet kan fastsette nærmere krav (enabling).",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§1-28: Administrative fines—from 1 January 2025 Statsforvalteren may impose infringement fines "
            "for violations of retailer sorting/storage (§1-5) and information duties (§1-6); governed by "
            "Chapter 18B."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            '"From 1 January 2025 the County Governor may, in the event of a breach of §1-5 and §1-6, impose '
            'an infringement fine on the person responsible for the breach." (§1-28)'
        ),
        "comments": "Enforcement instrument (kan ilegge); operative from 1 January 2025.",
    },
]

HEADERS = [
    ("A", "policy_name"),
    ("B", "policy_url"),
    ("C", "policy_year"),
    ("D", "policy_objective"),
    ("E", "policy_target"),
    ("F", "policy_target_text"),
    ("G", "policy_type"),
    ("H", "policy_type_justification"),
    ("I", "policy_integration"),
    ("J", "policy_sectors_list"),
    ("K", "policy_circularity"),
    ("L", "policy_lifecycle_phases_list"),
    ("M", "policy_budget"),
    ("N", "policy_budget_text"),
    ("O", "policy_score"),
    ("P", "instrument_type"),
    ("Q", "instrument_lifecycle_stage"),
    ("R", "instrument_description"),
    ("S", "instrument_in_force"),
    ("T", "instrument_implementation"),
    ("U", "instrument_implementation_text"),
    ("V", "instrument_score"),
    ("W", "comments"),
]


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "4P Index - Avfallsforskr Kap1"

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col_idx, (col_letter, header) in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=f"{col_letter} — {header}")
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, inst in enumerate(INSTRUMENTS, start=2):
        policy_score = (
            POLICY["policy_type"]
            + POLICY["policy_integration"]
            + POLICY["policy_circularity"]
            + POLICY["policy_budget"]
            + inst["instrument_type"]
        ) / 5
        instrument_score = (inst["instrument_type"] + inst["instrument_implementation"]) / 2

        row_values = [
            POLICY["policy_name"],
            POLICY["policy_url"],
            POLICY["policy_year"],
            POLICY["policy_objective"],
            POLICY["policy_target"],
            POLICY["policy_target_text"],
            POLICY["policy_type"],
            POLICY["policy_type_justification"],
            POLICY["policy_integration"],
            POLICY["policy_sectors_list"],
            POLICY["policy_circularity"],
            POLICY["policy_lifecycle_phases_list"],
            POLICY["policy_budget"],
            POLICY["policy_budget_text"],
            round(policy_score, 3),
            inst["instrument_type"],
            inst["instrument_lifecycle_stage"],
            inst["instrument_description"],
            inst["instrument_in_force"],
            inst["instrument_implementation"],
            inst["instrument_implementation_text"],
            round(instrument_score, 3),
            inst["comments"],
        ]

        for col_idx, value in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30

    widths = {
        "A": 52, "B": 42, "C": 12, "D": 45, "E": 12, "F": 40, "G": 12, "H": 40,
        "I": 16, "J": 42, "K": 16, "L": 35, "M": 14, "N": 45, "O": 12, "P": 14,
        "Q": 20, "R": 55, "S": 14, "T": 18, "U": 45, "V": 14, "W": 42,
    }
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width

    output_path = "/workspace/4P_Index_Avfallsforskriften_Kap1.xlsx"
    wb.save(output_path)
    print(f"Saved {output_path} ({len(INSTRUMENTS)} instrument rows)")


if __name__ == "__main__":
    main()
