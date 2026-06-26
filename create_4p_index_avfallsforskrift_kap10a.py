#!/usr/bin/env python3
"""Generate 4P Index Excel for Avfallsforskriften Kap. 10a (plastic waste sorting and collection)."""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

POLICY = {
    "policy_name": (
        "Avfallsforskriften – Kap. 10a: Utsortering og innsamling av plastavfall "
        "(Waste Regulation – Ch. 10a: Sorting and collection of plastic waste)"
    ),
    "policy_url": "https://lovdata.no/dokument/SF/forskrift/2004-06-01-930/kap10a",
    "policy_year": 2025,
    "policy_objective": (
        "Chapter 10a of the Norwegian Waste Regulation (FOR-2004-06-01-930) under forurensningsloven, "
        "implementing EU Waste Framework Directive (2008/98/EC, amended 2018/851) separate-collection "
        "obligations. Establishes municipal duties for household plastic waste source-sorting with binding "
        "share targets (50% by 2028, 60% by 2030, 70% by 2035), separate collection via henteordning, "
        "delivery to material recovery, documentation and supervision; parallel duties for businesses "
        "generating household-like waste and for agricultural plastic waste (silage wrap, PP sacks, fibre "
        "mesh, etc.). Chapter added January 2023 (FOR-2022-06-07-971); restructured January 2025 "
        "(FOR-2024-05-13-849)."
    ),
    "policy_target": 1,
    "policy_target_text": (
        'NO: "minst følgende andel … plastavfall … utsorteres ved kildesortering: 50 prosent fra og med 2028, '
        '60 prosent fra og med 2030 og 70 prosent fra og med 2035" (§10a-4 b) / '
        'EN: "at least the following share … plastic waste … source-sorted: 50 per cent from 2028, '
        '60 per cent from 2030 and 70 per cent from 2035" (§10a-4(b))'
    ),
    "policy_type": 0.75,
    "policy_type_justification": (
        "Executive regulation (forskrift) under forurensningsloven §§30, 32, 33, 49 and 81. Chapter 10a "
        "added 7 June 2022 (FOR-2022-06-07-971, in force 1 January 2023); major restructuring and municipal "
        "plastic-sorting targets 13 May 2024 (FOR-2024-05-13-849, in force 1 January 2025). Implements "
        "EEA Waste Framework Directive separate-collection requirements; cross-linked to packaging EPR "
        "cost-coverage in Kap. 7 §§7-8 and 7-5a."
    ),
    "policy_integration": 1,
    "policy_sectors_list": (
        "municipalities, agriculture, waste management, recycling, households, retail, food service, "
        "treatment plants, exporters"
    ),
    "policy_circularity": 1,
    "policy_lifecycle_phases_list": "end-of-life, recycling",
    "policy_budget": 0.5,
    "policy_budget_text": (
        'NO: "Produsent skal dekke nødvendige kostnader til separat innsamling og etterfølgende transport og '
        'behandling av avfall fra emballasjen" når kildesortering erstattes jf. §§10a-4, 10a-5, 10a-8, 10a-9 '
        '(kap. 7 §7-8); kommunen bærer plikt til innsamling etter §10a-5 / '
        'EN: "The producer shall cover necessary costs of separate collection and subsequent transport and '
        'treatment of packaging waste" when source sorting is replaced per §§10a-4, 10a-5, 10a-8, 10a-9 '
        "(Ch. 7 §7-8); municipality bears collection duties under §10a-5"
    ),
}

INSTRUMENTS = [
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§10a-1: Chapter purpose—increase preparation for reuse and material recovery of household and "
            "commercial waste for better resource utilisation, environmental protection and reduced greenhouse "
            "gas emissions."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            'NO: "Formålet med dette kapitlet er å øke forberedelse til ombruk og materialgjenvinning av '
            'husholdningsavfall og næringsavfall for å oppnå bedre ressursutnyttelse av avfall, beskytte miljøet '
            'og redusere klimagassutslipp." (§10a-1) / '
            'EN: "The purpose of this chapter is to increase preparation for reuse and material recovery of '
            'household and commercial waste to achieve better resource utilisation of waste, protect the '
            'environment and reduce greenhouse gas emissions." (§10a-1)'
        ),
        "comments": "Frames EU WFD separate-collection policy; added FOR-2022-06-07-971.",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§10a-2: Scope exclusions—chapter does not apply to WEEE (Kap. 1), leisure boats (Kap. 2), batteries "
            "(Kap. 3), ELVs (Kap. 4), tyres (Kap. 5), deposit beverage packaging (Kap. 6 DRS), hazardous waste "
            "(Kap. 11) or construction/demolition waste."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            'NO: "Bestemmelsene i dette kapitlet gjelder ikke for … f. drikkevareemballasje som inngår i '
            'panteordning … h. bygge- og riveavfall" (§10a-2) / '
            'EN: "The provisions of this chapter do not apply to … f. beverage packaging included in a '
            'deposit-return scheme … h. construction and demolition waste" (§10a-2)'
        ),
        "comments": "Plastic beverage bottles under Kap. 6 pant excluded; other household plastics in scope.",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§10a-3(e–f, i–k): Key definitions—plastavfall (discarded plastic packaging, SUP, fossil/bio-based, "
            "excluding plastic textiles); landbruksplastavfall (silage wrap, PP sacks, fibre mesh, hard plastic "
            "packaging); utsortering/kildesortering; separat innsamling (henteordning/bringeordning)."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            'NO: "plastavfall; plast som har blitt avfall … kassert plastemballasje, kasserte engangsprodukter '
            'av plast"; "landbruksplastavfall; plastavfall i form av folie (rundballeplast), PP-sekker …"; '
            '"separat innsamling; innsamling der en avfallsstrøm holdes atskilt" (§10a-3) / '
            'EN: "plastic waste; plastic that has become waste … discarded plastic packaging, discarded single-use '
            'plastic products"; "agricultural plastic waste; plastic waste in the form of film (silage wrap), '
            'PP sacks …"; "separate collection; collection where a waste stream is kept separate" (§10a-3)'
        ),
        "comments": "Legal definitions delimiting plastic streams covered by sorting obligations.",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§10a-4(b): Municipal household plastic source-sorting targets—municipality shall ensure recyclable "
            "household plastic waste is source-sorted at minimum shares: 50% from 2028, 60% from 2030, "
            "70% from 2035 (of total recyclable plastic collected from households per year)."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            'NO: "Kommunen skal sørge for … b. at plastavfall fra husholdninger, som kan forberedes til ombruk '
            'eller materialgjenvinnes, utsorteres ved kildesortering, og at minst følgende andel … utsorteres '
            'ved kildesortering: 50 prosent fra og med 2028, 60 prosent fra og med 2030 og 70 prosent fra og '
            'med 2035" (§10a-4 b) / '
            'EN: "The municipality shall ensure … b. that plastic waste from households that can be prepared for '
            'reuse or materially recovered is source-sorted, and that at least the following share … is source-sorted: '
            '50 per cent from 2028, 60 per cent from 2030 and 70 per cent from 2035" (§10a-4(b))'
        ),
        "comments": "Core binding plastic sorting target (skal); restructured FOR-2024-05-13-849.",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§10a-4(2): Alternative sorting for plastic—source-sorting of plastic and metal packaging waste may "
            "be replaced by other sorting if method achieves at least equal sorting share and material-recovery rate."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            'NO: "Kildesortering av plastavfall og metallemballasjeavfall kan erstattes av annen sortering dersom '
            'metoden gir minst like høy utsorteringsandel og andel avfall som forberedes til ombruk eller '
            'materialgjenvinnes som ved kildesortering" (§10a-4 andre ledd) / '
            'EN: "Source-sorting of plastic waste and metal packaging waste may be replaced by other sorting if '
            'the method gives at least as high a sorting share and share of waste prepared for reuse or material '
            'recovery as source-sorting" (§10a-4 second paragraph)'
        ),
        "comments": "Allows post-collection sorting (e.g. ROAF/IVAR); producer EPR covers costs per Kap. 7 §7-8.",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§10a-4(3): Miljødirektoratet guidelines—Director may issue guidelines on which plastic and textile "
            "waste can be prepared for reuse or material recovery (defining recyclable plastic scope for §10a-4(b))."
        ),
        "instrument_in_force": 0,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            'NO: "Miljødirektoratet kan gi nærmere retningslinjer om plastavfall og tekstilavfall som kan '
            'forberedes til ombruk eller materialgjenvinnes" (§10a-4 tredje ledd) / '
            'EN: "The Norwegian Environment Agency may issue further guidelines on plastic waste and textile '
            'waste that can be prepared for reuse or material recovery" (§10a-4 third paragraph)'
        ),
        "comments": "Enabling power (kan gi retningslinjer); operationalises recyclable plastic definition.",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§10a-4(4): Exemption from sorting-share targets—Miljødirektoratet or delegate may exempt municipalities "
            "from quantitative sorting targets if environmental/cost assessment shows targets are not justified."
        ),
        "instrument_in_force": 0,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            'NO: "Dersom kravene til utsorteringsandel … ikke er berettiget, kan Miljødirektoratet … gjøre unntak '
            'fra kravene til utsorteringsandel i første ledd" (§10a-4 fjerde ledd) / '
            'EN: "If the sorting-share requirements … are not justified, the Norwegian Environment Agency may … '
            'make exceptions from the sorting-share requirements in the first paragraph" (§10a-4 fourth paragraph)'
        ),
        "comments": "Discretionary exemption (kan gjøre unntak); reduces unconditional implementation score.",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§10a-5(1): Municipal separate collection of plastic waste—municipality shall ensure separate collection "
            "of source-sorted household plastic waste (with food, paper, glass/metal, garden and textile waste)."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            'NO: "Kommunen skal sørge for separat innsamling av utsortert … plastavfall … fra husholdningene '
            'jf. § 10a-4 første ledd" (§10a-5) / '
            'EN: "The municipality shall ensure separate collection of sorted … plastic waste … from households '
            'cf. §10a-4 first paragraph" (§10a-5)'
        ),
        "comments": "Mandatory separate collection (skal sørge for); prerequisite for material recovery.",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§10a-5(2): Henteordning for plastic—collection of sorted household plastic, food and paper waste "
            "shall be via door-to-door henteordning (with bringeordning exceptions for cabins, rural areas, "
            "exempt properties and oversized waste)."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            'NO: "Innsamling av utsortert … plastavfall … skal skje ved en henteordning"; bringeordning tillatt '
            'i hytteområder, utenfor tettbygd strøk m.m. (§10a-5 andre ledd) / '
            'EN: "Collection of sorted … plastic waste … shall take place via a door-to-door collection scheme"; '
            'bring scheme permitted in cabin areas, outside built-up areas etc. (§10a-5 second paragraph)'
        ),
        "comments": "Default kerbside collection for plastics; rural/cabin bring-scheme flexibility.",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§10a-6: Municipal delivery to material recovery—municipality shall ensure source-sorted waste "
            "under §10a-4 is delivered for preparation for reuse or material recovery (including plastic)."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            'NO: "Kommunen skal sørge for at utsortert avfall jf. § 10a-4 leveres til forberedelse til ombruk '
            'eller materialgjenvinning" (§10a-6) / '
            'EN: "The municipality shall ensure that sorted waste cf. §10a-4 is delivered for preparation for '
            'reuse or material recovery" (§10a-6)'
        ),
        "comments": "Closes loop from collection to recycling; added FOR-2024-05-13-849.",
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§10a-7: Municipal documentation—municipality must document annual plastic sorting share, quantities "
            "source-sorted/collected/delivered for recovery, and alternative-sorting equivalence; retain 5 years."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            'NO: "Kommunen skal ha kunnskap om og dokumentasjon på a. oppnådd utsorteringsandel per år jf. '
            '§ 10a-4 … b. årlig mengde … plastavfall … levert til forberedelse til ombruk og materialgjenvinning"; '
            'dokumentasjon … i minst 5 år (§10a-7) / '
            'EN: "The municipality shall have knowledge of and documentation on a. achieved sorting share per year '
            'cf. §10a-4 … b. annual quantity … plastic waste … delivered for preparation for reuse and material '
            'recovery"; documentation retained at least 5 years (§10a-7)'
        ),
        "comments": "Monitoring/accountability for 50%/70% plastic targets; mandatory (skal ha).",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§10a-8: Business source-sorting of plastic—enterprises generating household-like waste shall "
            "source-sort plastic waste that can be prepared for reuse or material recovery."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            'NO: "Virksomheter som genererer husholdningslignende avfall skal sørge for at … plastavfall … '
            'utsorteres ved kildesortering"; gjelder plastavfall som kan forberedes til ombruk eller '
            'materialgjenvinnes (§10a-8) / '
            'EN: "Enterprises generating household-like waste shall ensure that … plastic waste … is source-sorted"; '
            'applies to plastic waste that can be prepared for reuse or material recovery (§10a-8)'
        ),
        "comments": "Extends plastic sorting to commercial/household-like waste generators; added 2025.",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§10a-9: Business separate collection of plastic—enterprises shall ensure separate collection of "
            "source-sorted household-like plastic waste."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            'NO: "Virksomheter som genererer husholdningslignende avfall skal sørge for separat innsamling av '
            'utsortert … plastavfall … jf. § 10a-8" (§10a-9) / '
            'EN: "Enterprises generating household-like waste shall ensure separate collection of sorted … '
            'plastic waste … cf. §10a-8" (§10a-9)'
        ),
        "comments": "Business parallel to municipal §10a-5 plastic collection duty.",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§10a-10: Business delivery to material recovery—enterprises shall ensure source-sorted plastic waste "
            "is delivered for preparation for reuse or material recovery."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            'NO: "Virksomheter som genererer husholdningslignende avfall skal sørge for at utsortert avfall '
            'jf. § 10a-8 leveres til forberedelse til ombruk eller materialgjenvinning" (§10a-10) / '
            'EN: "Enterprises generating household-like waste shall ensure that sorted waste cf. §10a-8 is '
            'delivered for preparation for reuse or material recovery" (§10a-10)'
        ),
        "comments": "Business parallel to municipal §10a-6.",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§10a-11: Agricultural plastic waste—enterprises using agricultural plastic shall source-sort, "
            "separately collect and deliver recyclable agricultural plastic waste (silage wrap, PP sacks, fibre "
            "mesh, hard plastic packaging) for material recovery; alternative sorting permitted if equivalent."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            'NO: "Virksomheter som bruker landbruksplast skal sørge for at landbruksplastavfall som kan '
            'forberedes til ombruk eller materialgjenvinnes, utsorteres ved kildesortering"; "separat innsamling … '
            'og at utsortert landbruksplastavfall leveres til forberedelse til ombruk eller materialgjenvinning" '
            '(§10a-11) / '
            'EN: "Enterprises using agricultural plastic shall ensure that agricultural plastic waste that can be '
            'prepared for reuse or material recovery is source-sorted"; "separate collection … and delivery for '
            'preparation for reuse or material recovery" (§10a-11)'
        ),
        "comments": "Farm-level plastic collection duty; complements Grønt Punkt voluntary scheme.",
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§10a-12: Business documentation—enterprises with §10a-8–11 duties must document annual quantities "
            "of plastic and agricultural plastic waste sorted, collected and delivered; retain 5 years."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            'NO: "Virksomheter som har plikter etter § 10a-8 til § 10a-11 skal ha kunnskap om og dokumentasjon på '
            '… plastavfall … og landbruksplastavfall … levert til forberedelse til ombruk og materialgjenvinning"; '
            'dokumentasjon … i minst 5 år (§10a-12) / '
            'EN: "Enterprises with duties under §10a-8 to §10a-11 shall have knowledge of and documentation on … '
            'plastic waste … and agricultural plastic waste … delivered for preparation for reuse and material '
            'recovery"; documentation retained at least 5 years (§10a-12)'
        ),
        "comments": "Traceability for business and farm plastic waste streams.",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": (
            "§10a-13: Treatment plant material recovery—waste treatment facilities receiving sorted plastic waste "
            "(and other §10a streams) shall ensure preparation for reuse or material recovery; exemption for "
            "fractions unsuitable per BAT/environmental assessment."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            'NO: "Behandlingsanlegg for avfall som tar imot utsortert … plastavfall … skal sørge for forberedelse '
            'til ombruk eller materialgjenvinning av avfallet"; unntak for avfallsdeler uegnet etter '
            'miljø-/ressurshensyn og beste tilgjengelige teknikk (§10a-13) / '
            'EN: "Waste treatment facilities receiving sorted … plastic waste … shall ensure preparation for reuse '
            'or material recovery of the waste"; exemption for waste fractions unsuitable per environmental/resource '
            'considerations and best available techniques (§10a-13)'
        ),
        "comments": "Downstream recycling obligation on treatment plants and exporters.",
    },
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Recycling",
        "instrument_description": (
            "§10a-14: Treatment plant and exporter documentation—facilities and exporters must document "
            "quantities of plastic waste received, materially recovered, or exported; retain 5 years."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            'NO: "Behandlingsanlegg … skal ha kunnskap om og dokumentasjon på … plastavfall … og landbruksplastavfall '
            'som er mottatt … forberedt til ombruk og materialgjenvunnet"; eksportører skal dokumentere eksportert '
            'plastavfall (§10a-14) / '
            'EN: "Treatment facilities … shall have knowledge of and documentation on … plastic waste … and '
            'agricultural plastic waste … received … prepared for reuse and materially recovered"; exporters shall '
            'document exported plastic waste (§10a-14)'
        ),
        "comments": "Supports Basel/export traceability for plastic waste fractions.",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§10a-15: Reporting—Miljødirektoratet may set reporting requirements for all §10a-4–14 obligations "
            "(municipal plastic sorting performance, business and treatment-plant data)."
        ),
        "instrument_in_force": 0,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            'NO: "Miljødirektoratet kan fastsette krav til rapportering iht. pliktene i § 10a-4 til § 10a-14" '
            '(§10a-15) / '
            'EN: "The Norwegian Environment Agency may set requirements for reporting in accordance with the duties '
            'in §10a-4 to §10a-14" (§10a-15)'
        ),
        "comments": "Enabling reporting power (kan fastsette); national monitoring of plastic sorting targets.",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "End-of-life",
        "instrument_description": (
            "§10a-16: Supervision—Miljødirektoratet supervises licensed treatment plants and exporters; "
            "Statsforvalteren supervises other enterprises, treatment plants and municipal compliance "
            "with chapter duties (forurensningsloven §48a)."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            'NO: "Miljødirektoratet fører tilsyn med … behandlingsanlegg … og eksportører"; "Statsforvalteren fører '
            'tilsyn med kommunens oppfyllelse av plikter … jf. forurensningsloven § 48a" (§10a-16) / '
            'EN: "The Norwegian Environment Agency supervises … treatment facilities … and exporters"; '
            '"The County Governor supervises the municipality\'s fulfilment of duties … cf. Pollution Control Act '
            '§48a" (§10a-16)'
        ),
        "comments": "Multi-level enforcement; county governors oversee municipal 50%/70% plastic targets.",
    },
]

HEADERS = [
    ("A", "policy_name"),
    ("B", "policy_url"),
    ("C", "policy_year"),
    ("D", "policy_objective"),
    ("E", "policy_target"),
    ("F", "policy_target_text"),
    ("G", "policy_type"),
    ("H", "policy_type_justification"),
    ("I", "policy_integration"),
    ("J", "policy_sectors_list"),
    ("K", "policy_circularity"),
    ("L", "policy_lifecycle_phases_list"),
    ("M", "policy_budget"),
    ("N", "policy_budget_text"),
    ("O", "policy_score"),
    ("P", "instrument_type"),
    ("Q", "instrument_lifecycle_stage"),
    ("R", "instrument_description"),
    ("S", "instrument_in_force"),
    ("T", "instrument_implementation"),
    ("U", "instrument_implementation_text"),
    ("V", "instrument_score"),
    ("W", "comments"),
]


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "4P Index - Kap 10a"

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col_idx, (col_letter, header) in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=f"{col_letter} — {header}")
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, inst in enumerate(INSTRUMENTS, start=2):
        policy_score = (
            POLICY["policy_type"]
            + POLICY["policy_integration"]
            + POLICY["policy_circularity"]
            + POLICY["policy_budget"]
            + inst["instrument_type"]
        ) / 5
        instrument_score = (inst["instrument_type"] + inst["instrument_implementation"]) / 2

        row_values = [
            POLICY["policy_name"],
            POLICY["policy_url"],
            POLICY["policy_year"],
            POLICY["policy_objective"],
            POLICY["policy_target"],
            POLICY["policy_target_text"],
            POLICY["policy_type"],
            POLICY["policy_type_justification"],
            POLICY["policy_integration"],
            POLICY["policy_sectors_list"],
            POLICY["policy_circularity"],
            POLICY["policy_lifecycle_phases_list"],
            POLICY["policy_budget"],
            POLICY["policy_budget_text"],
            round(policy_score, 3),
            inst["instrument_type"],
            inst["instrument_lifecycle_stage"],
            inst["instrument_description"],
            inst["instrument_in_force"],
            inst["instrument_implementation"],
            inst["instrument_implementation_text"],
            round(instrument_score, 3),
            inst["comments"],
        ]

        for col_idx, value in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30

    widths = {
        "A": 52, "B": 42, "C": 12, "D": 45, "E": 12, "F": 42, "G": 12, "H": 40,
        "I": 16, "J": 42, "K": 16, "L": 38, "M": 14, "N": 45, "O": 12, "P": 14,
        "Q": 20, "R": 55, "S": 14, "T": 18, "U": 55, "V": 14, "W": 42,
    }
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width

    output_path = "/workspace/4P_Index_Avfallsforskrift_Kap10a.xlsx"
    wb.save(output_path)
    print(f"Saved {output_path} ({len(INSTRUMENTS)} instrument rows)")


if __name__ == "__main__":
    main()
