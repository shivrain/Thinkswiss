"""Generate interview coding spreadsheet for Nepal interview NPL_22.

Coding follows the Overview All Interviews codebook and is based on the
second ordinary-household respondent interview of 17 February 2026 —
interview guideline only (no recording; no separate transcript). Uses
expanded Impacts taxonomy.
"""

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUTPUT_PATH = "/workspace/interview_coding_NPL_22.xlsx"

DARK_BLUE = "1F3864"
MID_BLUE = "2E75B6"
LIGHT_BLUE = "D9E2F3"
LIGHT_GREEN = "E2EFDA"
ORANGE = "FCE4D6"
WHITE = "FFFFFF"

SOURCES = [
  ("Interview guideline", "Ordinary household respondent #2 — field notes, 17 Feb 2026"),
  ("Codebook", "Overview All Interviews — NEW Codebook (expanded Impacts)"),
  (
    "Coding note",
    "No audio recording (consent given, no recording). Codes verified against "
    "guideline/field notes only. Same fieldwork day as NPL_21 but distinct "
    "respondent. Several questions (Q3, Q6, Q9, Q10) not asked; relevant "
    "content integrated from other themes.",
  ),
]

STAKEHOLDER = {
  "Interview ID": "NPL_22",
  "Country": "NEPAL",
  "Interview number": "Household / citizen fieldwork series (#2, same day as NPL_21)",
  "Interview date": "17 February 2026",
  "Interviewee": "Unknown — ordinary household resident",
  "Affiliation / role": "Ordinary person / household member",
  "Organization": "NA",
  "Stakeholder list mapping": "household",
  "Coded actor type (codebook)": "household",
  "Actor classification rationale": (
    "Lay citizen/household respondent describing convenience of plastic use, "
    "municipal collection reliance, and low personal engagement with waste — "
    "not government, private sector, or CSO."
  ),
  "Perspective": (
    "Low-concern everyday user — plastic convenient (buy and bin); burning is "
    "the main worry (smell, coughing); municipal measures largely ineffective; "
    "segregation abolished; wants enforcement, production stop, and awareness"
  ),
  "Project context": (
    "Daily municipal waste pick-up; anti-plastic campaign ineffective; household "
    "segregation abolished; municipality produces unused compost (no nearby "
    "agricultural plots); residents avoid engaging with waste themselves"
  ),
}

CODING = {
  "Country": "NPL",
  "Country name": "NEPAL",
  "ID": "NPL_22",
  "Actortype": "household",
  "Problem_awareness_pop": "high",
  "Problem_awareness_pol": "medium",
  "Problem_severity": "medium",
  "Problem_littering": "yes",
  "Problem_consumption": "yes",
  "Problem_recycling": "yes",
  "Problem_waste_mgmt": "yes",
  "Problem_production": "no",
  "Problem_alternatives": "yes",
  "Problem_waste_segregation": "yes",
  "Problem_import": "no",
  "Impacts": (
    "wildlife, air pollution, health, agriculture, drainage blockage"
  ),
  "Governance": "yes",
  "Relevance_international_pol": "no",
  "Coordination_sectoral": "no",
  "Coordination_levels": "no",
  "Unclear_responsibilities": "no",
  "Actors": "yes",
  "Federal government": "no",
  "Provincial government": "no",
  "Local government": "yes",
  "Students": "no",
  "Private_sector": "no",
  "Civil_society": "no",
  "Science": "no",
  "Households": "yes",
  "Education institutions": "no",
  "Private_companies(hotels, shops, etc.)": "no",
  "Implementation issues": "yes",
  "Monitoring": "yes",
  "Financial_resources": "no",
  "Research": "no",
  "Infrastructure": "no",
  "Enforcement": "yes",
  "Policies in place": "yes",
  "Pol_epr": "no",
  "Pol_import": "no",
  "Pol_awarness": "yes",
  "Pol_education": "no",
  "Pol_capacity": "no",
  "Pol_RD": "no",
  "Pol_tax": "no",
  "Pol_ban": "no",
  "Pol_subsitutes": "yes",
  "Pol_clean_up": "no",
  "Pol_upcycling": "no",
  "Pol_recycling": "no",
  "Pol_waste_collection": "yes",
  "Solutions": "yes",
  "Sol_lead_agency": "no",
  "Sol_responsibilities": "yes",
  "Sol_epr": "no",
  "Sol_awarness": "yes",
  "Sol_segregation": "no",
  "Sol_upcycling": "no",
  "Sol_recycling": "no",
  "Sol_education": "yes",
  "Sol_capacity": "no",
  "Sol_RD": "no",
  "Sol_tax": "no",
  "Sol_ban": "yes",
  "Sol_finance": "no",
  "Sol_infrastructure": "no",
  "Sol_subsitutes": "yes",
  "Sol_clean_up": "no",
  "Sol_enforcement": "yes",
  "Sol_monitoring": "yes",
  "Traditions to build on (free-hand)": (
    "Traditional paper bags and bamboo leaves/baskets for carrying goods; cloth "
    "bags used today but not yet frequently (though increasing)"
  ),
}

CODEBOOK = [
  ("Country", "Country ISO code"),
  ("Country name", "Country name"),
  ("ID", "InterviewID: ISO_Nr"),
  ("Actortype", "governmental, pol_party, science, cso, igo, private_sector, shop_owner, hotel_owner, household"),
  ("Problem_awareness_pop", "high, medium, low, NA — lack of awareness among population mentioned"),
  ("Problem_awareness_pol", "high, medium, low, NA — lack of awareness among policy makers mentioned"),
  ("Problem_severity", "high, medium, low, NA"),
  ("Problem_littering", "Littering is a problem; was mentioned: yes/no"),
  ("Problem_consumption", "High consumption is a problem; was mentioned: yes/no"),
  ("Problem_recycling", "No or too little recycling is a problem; was mentioned: yes/no"),
  ("Problem_waste_mgmt", "Ill waste management is a problem; was mentioned: yes/no"),
  ("Problem_production", "High production rates; was mentioned: yes/no"),
  ("Problem_alternatives", "No good alternatives; was mentioned: yes/no"),
  ("Problem_waste_segregation", "Lack of segregation was mentioned: yes/no"),
  ("Problem_import", "Plastic imports are a problem; was mentioned: yes/no"),
  ("Impacts", "Impacts mentioned on water, cities, biodiversity, health, etc. or NA"),
  ("Governance", "Mentioned: yes/no"),
  ("Relevance_international_pol", "International treaties are relevant for national level policy"),
  ("Coordination_sectoral", "Lack of coordination across sectors"),
  ("Coordination_levels", "Lack of coordination across levels"),
  ("Unclear_responsibilities", "Unclear responsibilities among the involved actors"),
  ("Actors", "Mentioned: yes/no"),
  ("Federal government", "Mentioned: yes/no"),
  ("Provincial government", "Mentioned: yes/no"),
  ("Local government", "Mentioned: yes/no"),
  ("Students", "Mentioned: yes/no"),
  ("Private_sector", "Mentioned: yes/no"),
  ("Civil_society", "Mentioned: yes/no"),
  ("Science", "Mentioned: yes/no"),
  ("Households", "Mentioned: yes/no"),
  ("Education institutions", "Mentioned: yes/no"),
  ("Private_companies(hotels, shops, etc.)", "Mentioned: yes/no"),
  ("Implementation issues", "Mentioned: yes/no"),
  ("Monitoring", "Lack of monitoring"),
  ("Financial_resources", "Lack of financial resources"),
  ("Research", "Lack of research"),
  ("Infrastructure", "Lack of infrastructure"),
  ("Enforcement", "Lack of enforcement"),
  ("Policies in place", "Mentioned: yes/no"),
  ("Pol_epr", "Extended producer responsibility"),
  ("Pol_import", "Import regulations"),
  ("Pol_awarness", "Awareness campaigns"),
  ("Pol_education", "Education programs"),
  ("Pol_capacity", "Capacity building"),
  ("Pol_RD", "Research & development"),
  ("Pol_tax", "Levies or taxes on specific products (often bags)"),
  ("Pol_ban", "Bans of specific products"),
  ("Pol_subsitutes", "Alternatives like reusable bags, or banana leaf plates"),
  ("Pol_clean_up", "Clean-up campaigns"),
  ("Pol_upcycling", "Making a new product, like blacktop"),
  ("Pol_recycling", "Making a new raw material"),
  ("Pol_waste_collection", "Waste collection"),
  ("Solutions", "Mentioned: yes/no"),
  ("Sol_lead_agency", "Procedural measures to establish lead agency"),
  ("Sol_responsibilities", "Clear responsibilities"),
  ("Sol_epr", "Extended producer responsibility"),
  ("Sol_awarness", "Awareness raising measures"),
  ("Sol_segregation", "Segregation of waste"),
  ("Sol_upcycling", "Upcycling of plastics"),
  ("Sol_recycling", "New raw materials"),
  ("Sol_education", "Education programs at schools"),
  ("Sol_capacity", "Capacity-building programs"),
  ("Sol_RD", "Research programs"),
  ("Sol_tax", "Taxes or levies on products"),
  ("Sol_ban", "Ban of products"),
  ("Sol_finance", "Financial measures"),
  ("Sol_infrastructure", "Infrastructural measures"),
  ("Sol_subsitutes", "Alternatives"),
  ("Sol_clean_up", "Clean-up campaigns"),
  ("Sol_enforcement", "Enforcement procedures"),
  ("Sol_monitoring", "Monitoring procedures"),
  ("Traditions to build on (free-hand)", "Freehand or NA"),
]

NOTES = {
  "Actortype": (
    "Unknown ordinary household resident (second respondent, 17 Feb 2026). "
    "Stakeholder list: household. Codebook: household."
  ),
  "Problem_awareness_pop": (
    "Information about plastic impacts not taught to people; awareness needs "
    "increasing despite anti-plastic campaign."
  ),
  "Problem_awareness_pol": (
    "Government does not enforce regulations and provides no monitoring — "
    "campaigns ineffective."
  ),
  "Problem_severity": (
    "Not very concerned about pollution overall; burning is the main worry "
    "(smell, coughing) alongside livestock ingestion and agricultural contamination."
  ),
  "Problem_littering": (
    "Livestock eat discarded plastics; convenient buy-and-bin behaviour implies "
    "widespread plastic discard."
  ),
  "Problem_consumption": (
    "Plastic very convenient — buy plastics and throw in dustbin."
  ),
  "Problem_recycling": (
    "Segregation abolished; municipality compost unused; no household recycling "
    "or composting engagement."
  ),
  "Problem_waste_mgmt": (
    "Current measures not very effective; residents do not want to engage with "
    "waste themselves; municipal compost unused due to no nearby agricultural plots."
  ),
  "Problem_alternatives": (
    "Plastic more convenient than traditional paper bags and bamboo baskets; cloth "
    "bags increasing but not yet frequent."
  ),
  "Problem_waste_segregation": (
    "Household waste segregation was abolished — no source separation."
  ),
  "Impacts": (
    "Wildlife: livestock eat plastics. Air pollution and health: burning causes "
    "smell, coughing, atmospheric harm. Agriculture: contaminated agricultural "
    "plots. Drainage blockage: blocked pipes."
  ),
  "Local government": (
    "Responsible for waste/plastic management; daily pick-up; anti-plastic "
    "campaign; municipal compost production."
  ),
  "Households": (
    "Residents do not make compost and avoid engaging with waste; rely on "
    "municipal collection."
  ),
  "Monitoring": (
    "No monitoring of compliance with waste/plastic rules."
  ),
  "Enforcement": (
    "Government does not enforce regulations; campaign against plastic use "
    "did not work well."
  ),
  "Policies in place": (
    "Anti-plastic use campaign; daily municipal waste pick-up; segregation "
    "previously in place but abolished; municipal compost programme (unused)."
  ),
  "Pol_awarness": (
    "Campaign not to use plastics — did not work well."
  ),
  "Pol_subsitutes": (
    "Cloth bags promoted today but not yet widely used; traditional paper bags "
    "and bamboo leaf/basket carriers recalled."
  ),
  "Pol_waste_collection": (
    "Daily waste pick-up organised by municipality."
  ),
  "Sol_responsibilities": (
    "Local government must enforce regulation — primary responsible actor."
  ),
  "Sol_awarness": (
    "Increase public awareness — teach impacts of plastic pollution."
  ),
  "Sol_education": (
    "Information about impacts not currently taught; education needed."
  ),
  "Sol_ban": (
    "Stop plastic production; strict regulations for carrying plastic bags."
  ),
  "Sol_subsitutes": (
    "Promote cloth bags, paper bags, and bamboo baskets over single-use plastics."
  ),
  "Sol_enforcement": (
    "Local government must enforce existing regulations strictly."
  ),
  "Sol_monitoring": (
    "Monitoring absent — needed alongside enforcement."
  ),
  "Traditions to build on (free-hand)": (
    "Paper bags and bamboo leaves/baskets traditionally used to carry goods; "
    "cloth bags emerging today but adoption still limited."
  ),
}


def hfill(hex_):
  return PatternFill("solid", fgColor=hex_)


def hfont(bold=False, size=10, color="000000", italic=False):
  return Font(bold=bold, size=size, color=color, italic=italic)


thin = Side(style="thin", color="BFBFBF")


def thin_border():
  return Border(left=thin, right=thin, top=thin, bottom=thin)


WRAP_ALIGN = Alignment(wrap_text=True, vertical="top", horizontal="left")
CENTER_ALIGN = Alignment(wrap_text=True, vertical="center", horizontal="center")


def build_workbook():
  wb = openpyxl.Workbook()

  ws = wb.active
  ws.title = "Interview Coding"

  ws.merge_cells("A1:D1")
  title = ws["A1"]
  title.value = "Plastic Pollution Governance — Interview Coding (Codebook)"
  title.fill = hfill(DARK_BLUE)
  title.font = hfont(bold=True, size=13, color=WHITE)
  title.alignment = CENTER_ALIGN

  ws.merge_cells("A2:D2")
  meta = ws["A2"]
  meta.value = (
    "Interview: NPL_22  |  Country: NEPAL  |  Ordinary household resident (#2)  |  "
    "Actor: household  |  17 Feb 2026"
  )
  meta.fill = hfill(MID_BLUE)
  meta.font = hfont(size=9, color=WHITE)
  meta.alignment = CENTER_ALIGN

  headers = ["Variable", "Code", "Codebook definition", "Coding notes"]
  for col, header in enumerate(headers, start=1):
    cell = ws.cell(row=4, column=col, value=header)
    cell.fill = hfill(LIGHT_BLUE)
    cell.font = hfont(bold=True, size=10, color=DARK_BLUE)
    cell.alignment = CENTER_ALIGN
    cell.border = thin_border()

  row = 5
  for variable, definition in CODEBOOK:
    values = [
      variable,
      CODING[variable],
      definition,
      NOTES.get(variable, ""),
    ]
    fill = hfill(LIGHT_GREEN) if row % 2 == 0 else hfill(WHITE)
    for col, value in enumerate(values, start=1):
      cell = ws.cell(row=row, column=col, value=value)
      cell.fill = fill
      cell.font = hfont(size=10)
      cell.alignment = WRAP_ALIGN
      cell.border = thin_border()
    row += 1

  ws.column_dimensions["A"].width = 34
  ws.column_dimensions["B"].width = 28
  ws.column_dimensions["C"].width = 52
  ws.column_dimensions["D"].width = 58
  ws.freeze_panes = "A5"

  wide = wb.create_sheet("Wide Format")
  wide_headers = [variable for variable, _ in CODEBOOK]
  for col, header in enumerate(wide_headers, start=1):
    cell = wide.cell(row=1, column=col, value=header)
    cell.fill = hfill(LIGHT_BLUE)
    cell.font = hfont(bold=True, size=9, color=DARK_BLUE)
    cell.alignment = CENTER_ALIGN
    cell.border = thin_border()
  for col, header in enumerate(wide_headers, start=1):
    cell = wide.cell(row=2, column=col, value=CODING[header])
    cell.fill = hfill(ORANGE)
    cell.alignment = WRAP_ALIGN
    cell.border = thin_border()
    wide.column_dimensions[get_column_letter(col)].width = max(14, min(28, len(header) + 2))

  ref = wb.create_sheet("Codebook Reference")
  ref.merge_cells("A1:B1")
  ref["A1"].value = "Codebook Variable Definitions"
  ref["A1"].fill = hfill(DARK_BLUE)
  ref["A1"].font = hfont(bold=True, size=12, color=WHITE)
  ref["A1"].alignment = CENTER_ALIGN
  for idx, (variable, definition) in enumerate(CODEBOOK, start=3):
    ref.cell(row=idx, column=1, value=variable).font = hfont(bold=True)
    ref.cell(row=idx, column=2, value=definition).alignment = WRAP_ALIGN
  ref.column_dimensions["A"].width = 34
  ref.column_dimensions["B"].width = 70

  stakeholder = wb.create_sheet("Stakeholder Analysis")
  stakeholder.merge_cells("A1:B1")
  stakeholder["A1"].value = "Stakeholder Analysis — NPL_22"
  stakeholder["A1"].fill = hfill(DARK_BLUE)
  stakeholder["A1"].font = hfont(bold=True, size=12, color=WHITE)
  stakeholder["A1"].alignment = CENTER_ALIGN
  for idx, (key, value) in enumerate(STAKEHOLDER.items(), start=3):
    stakeholder.cell(row=idx, column=1, value=key).font = hfont(bold=True)
    cell = stakeholder.cell(row=idx, column=2, value=value)
    cell.alignment = WRAP_ALIGN
  stakeholder.column_dimensions["A"].width = 28
  stakeholder.column_dimensions["B"].width = 90

  sources = wb.create_sheet("Sources")
  sources.merge_cells("A1:B1")
  sources["A1"].value = "Source Documents Used for Verification"
  sources["A1"].fill = hfill(DARK_BLUE)
  sources["A1"].font = hfont(bold=True, size=12, color=WHITE)
  sources["A1"].alignment = CENTER_ALIGN
  for idx, (doc, detail) in enumerate(SOURCES, start=3):
    sources.cell(row=idx, column=1, value=doc).font = hfont(bold=True)
    sources.cell(row=idx, column=2, value=detail).alignment = WRAP_ALIGN
  sources.column_dimensions["A"].width = 42
  sources.column_dimensions["B"].width = 60

  wb.save(OUTPUT_PATH)
  print(f"Saved {OUTPUT_PATH}")


if __name__ == "__main__":
  build_workbook()
