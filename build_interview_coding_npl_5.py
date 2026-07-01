"""Generate interview coding spreadsheet for Nepal interview NPL_5.

Coding follows the Overview All Interviews codebook and is based on the
Doco Recyclers (private sector) interview of 23 June 2025 — transcript,
interview notes, and guideline.
"""

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUTPUT_PATH = "/workspace/interview_coding_NPL_5.xlsx"

DARK_BLUE = "1F3864"
MID_BLUE = "2E75B6"
LIGHT_BLUE = "D9E2F3"
LIGHT_GREEN = "E2EFDA"
ORANGE = "FCE4D6"
WHITE = "FFFFFF"

SOURCES = [
  ("Interview transcript", "Doco Recyclers — Private Sector, 23 June 2025"),
  ("Interview notes", "Guideline responses and structured bullet notes"),
  ("Interview guideline", "Plastic Pollution Governance in Nepal — Private Sector"),
  ("Codebook", "Overview All Interviews — NEW Codebook"),
]

STAKEHOLDER = {
  "Interview ID": "NPL_5",
  "Country": "NEPAL",
  "Interview date": "23 June 2025",
  "Organization": "Doco Recyclers (est. 2017)",
  "Affiliation / role": "Private sector — dry waste and e-waste management company",
  "Respondents": "Two representatives (Respondent 1 & Respondent 2)",
  "Stakeholder list mapping": "private_sector",
  "Coded actor type (codebook)": "private_sector",
  "Actor classification rationale": (
    "Registered private waste-management company operating MRFs, collection "
    "chains, consulting (UNDP, GIZ), and research partnerships. Distinct from "
    "informal scrap collectors (partially integrated) and government actors."
  ),
  "Role in plastic governance": (
    "Collection, sorting (10–12 plastic grades), processing, and passing "
    "recyclables to downstream processors; KMC Cluster 7 MoU partner; policy "
    "advocate for EPR and apex waste-governance body"
  ),
  "Operations": (
    "MRFs at Sanothimi and Salaghari (15–16 tonnes/day); serves Kathmandu "
    "Valley; expanding to Pokhara; clientele model for institutions/households"
  ),
  "Perspective": (
    "Private-sector implementer — critical of fragmented policy, weak "
    "enforcement, and lack of incentives for low-value/multilayer plastics"
  ),
}

CODING = {
  "Country": "NPL",
  "Country name": "NEPAL",
  "ID": "NPL_5",
  "Actortype": "private_sector",
  "Problem_awareness_pop": "high",
  "Problem_awareness_pol": "high",
  "Problem_severity": "high",
  "Problem_littering": "yes",
  "Problem_consumption": "yes",
  "Problem_recycling": "yes",
  "Problem_waste_mgmt": "yes",
  "Problem_production": "yes",
  "Problem_alternatives": "yes",
  "Problem_waste_segregation": "yes",
  "Problem_import": "yes",
  "Impacts": "water, cities, health, agriculture",
  "Governance": "yes",
  "Relevance_international_pol": "no",
  "Coordination_sectoral": "yes",
  "Coordination_levels": "yes",
  "Unclear_responsibilities": "yes",
  "Actors": "yes",
  "Federal government": "yes",
  "Provincial government": "no",
  "Local government": "yes",
  "Students": "yes",
  "Private_sector": "yes",
  "Civil_society": "yes",
  "Science": "yes",
  "Households": "yes",
  "Education institutions": "yes",
  "Private_companies(hotels, shops, etc.)": "yes",
  "Implementation issues": "yes",
  "Monitoring": "yes",
  "Financial_resources": "yes",
  "Research": "yes",
  "Infrastructure": "yes",
  "Enforcement": "yes",
  "Policies in place": "yes",
  "Pol_epr": "no",
  "Pol_import": "yes",
  "Pol_awarness": "yes",
  "Pol_education": "yes",
  "Pol_capacity": "yes",
  "Pol_RD": "yes",
  "Pol_tax": "no",
  "Pol_ban": "yes",
  "Pol_subsitutes": "yes",
  "Pol_clean_up": "no",
  "Pol_upcycling": "yes",
  "Pol_recycling": "yes",
  "Pol_waste_collection": "yes",
  "Solutions": "yes",
  "Sol_lead_agency": "yes",
  "Sol_responsibilities": "yes",
  "Sol_epr": "yes",
  "Sol_awarness": "yes",
  "Sol_segregation": "yes",
  "Sol_upcycling": "yes",
  "Sol_recycling": "yes",
  "Sol_education": "yes",
  "Sol_capacity": "yes",
  "Sol_RD": "yes",
  "Sol_tax": "no",
  "Sol_ban": "no",
  "Sol_finance": "yes",
  "Sol_infrastructure": "yes",
  "Sol_subsitutes": "yes",
  "Sol_clean_up": "no",
  "Sol_enforcement": "yes",
  "Sol_monitoring": "yes",
  "Traditions to build on (free-hand)": (
    "Banana leaves instead of plastic packaging; cotton or towel bags instead "
    "of plastic carry bags"
  ),
}

CODEBOOK = [
  ("Country", "Country ISO code"),
  ("Country name", "Country name"),
  ("ID", "InterviewID: ISO_Nr"),
  ("Actortype", "governmental, pol_party, science, cso, igo, private_sector, shop_owner, hotel_owner, household"),
  ("Problem_awareness_pop", "high, medium, low, NA — lack of awareness among population mentioned"),
  ("Problem_awareness_pol", "high, medium, low, NA — lack of awareness among policy makers mentioned"),
  ("Problem_severity", "high, medium, low, NA"),
  ("Problem_littering", "Littering is a problem; was mentioned: yes/no"),
  ("Problem_consumption", "High consumption is a problem; was mentioned: yes/no"),
  ("Problem_recycling", "No or too little recycling is a problem; was mentioned: yes/no"),
  ("Problem_waste_mgmt", "Ill waste management is a problem; was mentioned: yes/no"),
  ("Problem_production", "High production rates; was mentioned: yes/no"),
  ("Problem_alternatives", "No good alternatives; was mentioned: yes/no"),
  ("Problem_waste_segregation", "Lack of segregation was mentioned: yes/no"),
  ("Problem_import", "Plastic imports are a problem; was mentioned: yes/no"),
  ("Impacts", "Impacts mentioned on water, cities, biodiversity, health, etc. or NA"),
  ("Governance", "Mentioned: yes/no"),
  ("Relevance_international_pol", "International treaties are relevant for national level policy"),
  ("Coordination_sectoral", "Lack of coordination across sectors"),
  ("Coordination_levels", "Lack of coordination across levels"),
  ("Unclear_responsibilities", "Unclear responsibilities among the involved actors"),
  ("Actors", "Mentioned: yes/no"),
  ("Federal government", "Mentioned: yes/no"),
  ("Provincial government", "Mentioned: yes/no"),
  ("Local government", "Mentioned: yes/no"),
  ("Students", "Mentioned: yes/no"),
  ("Private_sector", "Mentioned: yes/no"),
  ("Civil_society", "Mentioned: yes/no"),
  ("Science", "Mentioned: yes/no"),
  ("Households", "Mentioned: yes/no"),
  ("Education institutions", "Mentioned: yes/no"),
  ("Private_companies(hotels, shops, etc.)", "Mentioned: yes/no"),
  ("Implementation issues", "Mentioned: yes/no"),
  ("Monitoring", "Lack of monitoring"),
  ("Financial_resources", "Lack of financial resources"),
  ("Research", "Lack of research"),
  ("Infrastructure", "Lack of infrastructure"),
  ("Enforcement", "Lack of enforcement"),
  ("Policies in place", "Mentioned: yes/no"),
  ("Pol_epr", "Extended producer responsibility"),
  ("Pol_import", "Import regulations"),
  ("Pol_awarness", "Awareness campaigns"),
  ("Pol_education", "Education programs"),
  ("Pol_capacity", "Capacity building"),
  ("Pol_RD", "Research & development"),
  ("Pol_tax", "Levies or taxes on specific products (often bags)"),
  ("Pol_ban", "Bans of specific products"),
  ("Pol_subsitutes", "Alternatives like reusable bags, or banana leaf plates"),
  ("Pol_clean_up", "Clean-up campaigns"),
  ("Pol_upcycling", "Making a new product, like blacktop"),
  ("Pol_recycling", "Making a new raw material"),
  ("Pol_waste_collection", "Waste collection"),
  ("Solutions", "Mentioned: yes/no"),
  ("Sol_lead_agency", "Procedural measures to establish lead agency"),
  ("Sol_responsibilities", "Clear responsibilities"),
  ("Sol_epr", "Extended producer responsibility"),
  ("Sol_awarness", "Awareness raising measures"),
  ("Sol_segregation", "Segregation of waste"),
  ("Sol_upcycling", "Upcycling of plastics"),
  ("Sol_recycling", "New raw materials"),
  ("Sol_education", "Education programs at schools"),
  ("Sol_capacity", "Capacity-building programs"),
  ("Sol_RD", "Research programs"),
  ("Sol_tax", "Taxes or levies on products"),
  ("Sol_ban", "Ban of products"),
  ("Sol_finance", "Financial measures"),
  ("Sol_infrastructure", "Infrastructural measures"),
  ("Sol_subsitutes", "Alternatives"),
  ("Sol_clean_up", "Clean-up campaigns"),
  ("Sol_enforcement", "Enforcement procedures"),
  ("Sol_monitoring", "Monitoring procedures"),
  ("Traditions to build on (free-hand)", "Freehand or NA"),
]

NOTES = {
  "Actortype": (
    "Doco Recyclers — private waste-management company (dry waste, e-waste). "
    "Stakeholder list: private_sector."
  ),
  "Problem_awareness_pop": (
    "People talk about pollution but rarely act; ignorance about personal "
    "responsibility; open burning despite knowing harms."
  ),
  "Problem_awareness_pol": (
    "Government uncertain which ministry leads; MOFAGA handed to Urban "
    "Development; DoE vs MoUD writing separate policies without synchronization."
  ),
  "Problem_severity": (
    "~600 tonnes plastic waste generated daily nationally; only 20–30 tonnes "
    "managed; carcinogenic dioxin/furan from open burning."
  ),
  "Problem_littering": (
    "Plastics on riverbanks, farmland, open dumping sites; plastic layers "
    "beneath soil in rivers across Nepal."
  ),
  "Problem_consumption": (
    "Excessive plastic dependency; single-use culture; problem is end-of-life "
    "not plastic itself."
  ),
  "Problem_recycling": (
    "Cherry-picking of high-value plastics; low-value/MLP fractions largely "
    "unmanaged; only ~3–5% of daily generation processed."
  ),
  "Problem_waste_mgmt": (
    "Fragmented management; insufficient processing capacity; no destination "
    "for low-value/non-recyclable fractions."
  ),
  "Problem_production": "Widely produced and discarded; large volumes of MLPs.",
  "Problem_alternatives": (
    "Slow shift to alternatives needed; cannot ban overnight; traditional "
    "leaf/cotton bag substitutes noted in notes."
  ),
  "Problem_waste_segregation": (
    "Segregation at source cited as key challenge; habit and awareness gaps."
  ),
  "Problem_import": (
    "Importers/brand owners part of EPR chain; some plastics passed to "
    "neighbouring countries through illegal channels."
  ),
  "Impacts": (
    "Agriculture: reduced yield, poor soil moisture/nutrient retention. Health: "
    "dioxin/furan from burning. Water: river contamination. Cities: sewage blocking."
  ),
  "Coordination_sectoral": (
    "Government, private sector, informal collectors, and donors working in "
    "different directions without synergy."
  ),
  "Coordination_levels": (
    "Fragmented federal/municipal policies; inter-municipality conflicts; no "
    "national plastic-specific policy."
  ),
  "Unclear_responsibilities": (
    "No apex body; uncertain ministry (Environment vs Urban Development vs "
    "DoE); producers not clearly held responsible."
  ),
  "Federal government": (
    "Ministry of Forest & Environment, Ministry of Urban Development, Department "
    "of Environment; draft SWM act under consultation."
  ),
  "Local government": (
    "753 local governments; municipalities write own guidelines; KMC partnership; "
    "should be held accountable for dumping."
  ),
  "Students": "Green school campaigns; students visit MRF facilities for learning.",
  "Private_sector": (
    "Doco Recyclers, private waste companies, producers/importers/brand owners, "
    "cement industry (RDF co-processing)."
  ),
  "Civil_society": (
    "NGOs run project-based campaigns; SAFED Centre microplastics study; "
    "World Bank-funded plastic recovery facility."
  ),
  "Science": (
    "University partnerships (Loughborough UK EPR research); UNDP/GIZ projects; "
    "KMC dry-waste inventory survey."
  ),
  "Households": "Clientele collection model from households and institutions.",
  "Education institutions": (
    "School learnings at MRF; waste management should be in course structure "
    "from young age."
  ),
  "Private_companies(hotels, shops, etc.)": (
    "Embassies, institutions, hotels as clients; thousand producers targetable "
    "under EPR vs 30 million consumers."
  ),
  "Monitoring": (
    "Municipalities do not adequately monitor private collectors; weak ban "
    "enforcement (<40 micron available everywhere)."
  ),
  "Financial_resources": (
    "Private entities lack investment confidence; need capital injection/funding "
    "for infrastructure; no plastic credits/carbon credits."
  ),
  "Research": (
    "Gaps on burning, microplastics, agriculture; some studies (SAFED Everest "
    "microplastics); EPR/e-waste policy research with UK university."
  ),
  "Infrastructure": (
    "Inadequate MRF/recycling/disposal capacity; Doco handles 15–16 t/day vs "
    "600 t/day national generation."
  ),
  "Enforcement": (
    "<40 micron ban not enforced nationwide; policies exist on paper only; "
    "penalties for open burning needed."
  ),
  "Policies in place": (
    "General SWM policy; ban on carry bags <40 microns; fragmented municipal "
    "SUP bans; no comprehensive national plastics policy."
  ),
  "Pol_epr": "EPR advocated and researched but not yet implemented; draft SWM act mentions it.",
  "Pol_import": (
    "Thin-plastic ban includes imported bags; importers should be EPR-liable; "
    "illegal cross-border flows for some fractions."
  ),
  "Pol_awarness": (
    "Project-based campaigns by private companies/NGOs; green school campaigns; "
    "facility visits — not continuous government effort."
  ),
  "Pol_education": (
    "School learnings at MRF; advocacy for mandatory waste education in curricula."
  ),
  "Pol_capacity": (
    "UNDP Pokhara skill-development training (400 participants); recycling/upcycling "
    "intensive training for 25 participants."
  ),
  "Pol_RD": (
    "MLP/low-value plastic projects; paper upcycling; EPR framework research; "
    "end-of-life household appliances study (4 regions)."
  ),
  "Pol_ban": "Ban on single-use plastic carry bags below 40 microns (weakly enforced).",
  "Pol_subsitutes": "Notes reference banana leaves and cotton/towel bags as traditions.",
  "Pol_upcycling": (
    "Paper-to-pencil upcycling; board from low-value plastic (World Bank/Salaghari); "
    "300–500 upcycled products."
  ),
  "Pol_recycling": (
    "Core operations: 10–12 plastic grades sorted, baled, passed to recyclers; "
    "PET, HDPE, LDPE fractions collected."
  ),
  "Pol_waste_collection": (
    "Proprietary collection chain; clientele model; KMC Cluster 7 MoU (NPR 62 lakh/year)."
  ),
  "Sol_lead_agency": "Apex coordinating/governing body for all waste and pollution management.",
  "Sol_responsibilities": (
    "EPR for producers/importers/brand owners; municipalities accountable for "
    "all plastic dumped within jurisdiction."
  ),
  "Sol_epr": "First priority policy — extended producer responsibility for packaging and plastics.",
  "Sol_awarness": (
    "Mandatory public education; information campaigns; life-cycle assessment literacy."
  ),
  "Sol_segregation": "Segregation at source as foundation of collection value chain.",
  "Sol_upcycling": "Community upcycling centres for MLP/low-value plastics.",
  "Sol_recycling": (
    "3R/4R; decentralized MRFs in high-generation clusters; RDF co-processing "
    "in cement industry; standardize recycled granules."
  ),
  "Sol_education": (
    "Waste management in school curricula from youngest age; carrot-and-stick approach."
  ),
  "Sol_capacity": (
    "Technology transfer from international companies; increase private-sector "
    "capital and infrastructure capacity."
  ),
  "Sol_RD": (
    "Research on burning/air pollution, microplastics in water (river, drinking, "
    "irrigation), plastics in agricultural soil."
  ),
  "Sol_ban": (
    "Banning plastic is temporary/impractical; focus on end-of-life management "
    "and gradual consumption reduction."
  ),
  "Sol_finance": (
    "Monetary incentives for private collectors; financing decentralized collection "
    "chains and cluster-based MRFs."
  ),
  "Sol_infrastructure": (
    "Decentralized collection, MRF, recycling, and disposal facilities in "
    "high-potential clusters; export/standardization policy for recycled products."
  ),
  "Sol_subsitutes": (
    "Revive banana-leaf packaging and cotton/towel bags; slow shift from single-use."
  ),
  "Sol_enforcement": (
    "Penalize open burning; enforce micron ban through field inspection; "
    "mandatory compliance not optional."
  ),
  "Sol_monitoring": (
    "Municipalities as monitoring partners holding private sector accountable; "
    "753 local governments responsible for in-jurisdiction dumping."
  ),
  "Traditions to build on (free-hand)": (
    "Banana leaves for packaging; cotton or towel bags for shopping — cultural "
    "heritage alternatives noted in guideline follow-up."
  ),
}


def hfill(hex_):
  return PatternFill("solid", fgColor=hex_)


def hfont(bold=False, size=10, color="000000", italic=False):
  return Font(bold=bold, size=size, color=color, italic=italic)


thin = Side(style="thin", color="BFBFBF")


def thin_border():
  return Border(left=thin, right=thin, top=thin, bottom=thin)


WRAP_ALIGN = Alignment(wrap_text=True, vertical="top", horizontal="left")
CENTER_ALIGN = Alignment(wrap_text=True, vertical="center", horizontal="center")


def build_workbook():
  wb = openpyxl.Workbook()

  ws = wb.active
  ws.title = "Interview Coding"

  ws.merge_cells("A1:D1")
  title = ws["A1"]
  title.value = "Plastic Pollution Governance — Interview Coding (Codebook)"
  title.fill = hfill(DARK_BLUE)
  title.font = hfont(bold=True, size=13, color=WHITE)
  title.alignment = CENTER_ALIGN

  ws.merge_cells("A2:D2")
  meta = ws["A2"]
  meta.value = (
    "Interview: NPL_5  |  Country: NEPAL  |  Doco Recyclers (Private Sector)  |  "
    "Actor: private_sector  |  23 June 2025"
  )
  meta.fill = hfill(MID_BLUE)
  meta.font = hfont(size=9, color=WHITE)
  meta.alignment = CENTER_ALIGN

  headers = ["Variable", "Code", "Codebook definition", "Coding notes"]
  for col, header in enumerate(headers, start=1):
    cell = ws.cell(row=4, column=col, value=header)
    cell.fill = hfill(LIGHT_BLUE)
    cell.font = hfont(bold=True, size=10, color=DARK_BLUE)
    cell.alignment = CENTER_ALIGN
    cell.border = thin_border()

  row = 5
  for variable, definition in CODEBOOK:
    values = [
      variable,
      CODING[variable],
      definition,
      NOTES.get(variable, ""),
    ]
    fill = hfill(LIGHT_GREEN) if row % 2 == 0 else hfill(WHITE)
    for col, value in enumerate(values, start=1):
      cell = ws.cell(row=row, column=col, value=value)
      cell.fill = fill
      cell.font = hfont(size=10)
      cell.alignment = WRAP_ALIGN
      cell.border = thin_border()
    row += 1

  ws.column_dimensions["A"].width = 34
  ws.column_dimensions["B"].width = 28
  ws.column_dimensions["C"].width = 52
  ws.column_dimensions["D"].width = 58
  ws.freeze_panes = "A5"

  wide = wb.create_sheet("Wide Format")
  wide_headers = [variable for variable, _ in CODEBOOK]
  for col, header in enumerate(wide_headers, start=1):
    cell = wide.cell(row=1, column=col, value=header)
    cell.fill = hfill(LIGHT_BLUE)
    cell.font = hfont(bold=True, size=9, color=DARK_BLUE)
    cell.alignment = CENTER_ALIGN
    cell.border = thin_border()
  for col, header in enumerate(wide_headers, start=1):
    cell = wide.cell(row=2, column=col, value=CODING[header])
    cell.fill = hfill(ORANGE)
    cell.alignment = WRAP_ALIGN
    cell.border = thin_border()
    wide.column_dimensions[get_column_letter(col)].width = max(14, min(28, len(header) + 2))

  ref = wb.create_sheet("Codebook Reference")
  ref.merge_cells("A1:B1")
  ref["A1"].value = "Codebook Variable Definitions"
  ref["A1"].fill = hfill(DARK_BLUE)
  ref["A1"].font = hfont(bold=True, size=12, color=WHITE)
  ref["A1"].alignment = CENTER_ALIGN
  for idx, (variable, definition) in enumerate(CODEBOOK, start=3):
    ref.cell(row=idx, column=1, value=variable).font = hfont(bold=True)
    ref.cell(row=idx, column=2, value=definition).alignment = WRAP_ALIGN
  ref.column_dimensions["A"].width = 34
  ref.column_dimensions["B"].width = 70

  stakeholder = wb.create_sheet("Stakeholder Analysis")
  stakeholder.merge_cells("A1:B1")
  stakeholder["A1"].value = "Stakeholder Analysis — NPL_5"
  stakeholder["A1"].fill = hfill(DARK_BLUE)
  stakeholder["A1"].font = hfont(bold=True, size=12, color=WHITE)
  stakeholder["A1"].alignment = CENTER_ALIGN
  for idx, (key, value) in enumerate(STAKEHOLDER.items(), start=3):
    stakeholder.cell(row=idx, column=1, value=key).font = hfont(bold=True)
    cell = stakeholder.cell(row=idx, column=2, value=value)
    cell.alignment = WRAP_ALIGN
  stakeholder.column_dimensions["A"].width = 28
  stakeholder.column_dimensions["B"].width = 90

  sources = wb.create_sheet("Sources")
  sources.merge_cells("A1:B1")
  sources["A1"].value = "Source Documents Used for Verification"
  sources["A1"].fill = hfill(DARK_BLUE)
  sources["A1"].font = hfont(bold=True, size=12, color=WHITE)
  sources["A1"].alignment = CENTER_ALIGN
  for idx, (doc, detail) in enumerate(SOURCES, start=3):
    sources.cell(row=idx, column=1, value=doc).font = hfont(bold=True)
    sources.cell(row=idx, column=2, value=detail).alignment = WRAP_ALIGN
  sources.column_dimensions["A"].width = 42
  sources.column_dimensions["B"].width = 60

  wb.save(OUTPUT_PATH)
  print(f"Saved {OUTPUT_PATH}")


if __name__ == "__main__":
  build_workbook()
