"""Generate interview coding spreadsheet for Nepal interview NPL_19.

Coding follows the Overview All Interviews codebook and is based on the
TAAN Pokhara president / Ethical Trekking founder interview — transcript
and interview guideline. Uses expanded Impacts taxonomy.
"""

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUTPUT_PATH = "/workspace/interview_coding_NPL_19.xlsx"

DARK_BLUE = "1F3864"
MID_BLUE = "2E75B6"
LIGHT_BLUE = "D9E2F3"
LIGHT_GREEN = "E2EFDA"
ORANGE = "FCE4D6"
WHITE = "FFFFFF"

SOURCES = [
  ("Interview transcript", "TAAN Trekking Association — respondent (Ethical Trekking / TAAN Pokhara)"),
  ("Interview guideline", "Plastic Pollution Governance in Nepal — TAAN president structured guideline"),
  ("Codebook", "Overview All Interviews — NEW Codebook (expanded Impacts)"),
  (
    "Coding note",
    "Extended NGO/orphanage backstory in transcript precedes plastic questions; "
    "plastic-governance codes verified against transcript and guideline Themes A–D.",
  ),
]

STAKEHOLDER = {
  "Interview ID": "NPL_19",
  "Country": "NEPAL",
  "Interview number": "TAAN / trekking sector series (Pokhara)",
  "Interview date": "Not recorded in transcript or guideline",
  "Interviewee": "Name not recorded — TAAN Pokhara chapter president",
  "Affiliation / role": (
    "President, Trekking Agencies' Association of Nepal (TAAN), Pokhara chapter "
    "(first president elected without contest); founder, Ethical Trekking NGO "
    "(est. 2009); trekking guide"
  ),
  "Organization": (
    "TAAN Pokhara (~180 member trekking companies; 42-year national association); "
    "Ethical Trekking — eco-friendly trekking NGO formerly sustaining orphanage; "
    "coordinates with ACAP, local governments, Nepal Army, NTNC"
  ),
  "Stakeholder list mapping": "cso",
  "Coded actor type (codebook)": "cso",
  "Actor classification rationale": (
    "Leads a civil-society industry association (TAAN) and founded Ethical Trekking "
    "NGO; also operates as trekking entrepreneur, but primary coded role is "
    "association/NGO leadership coordinating sector conservation measures."
  ),
  "Perspective": (
    "Trekking-sector advocate — highly concerned about trail and lakeside litter, "
    "domestic tourism as major polluter, weak enforcement of existing policies, "
    "and need for paid off-season clean campaigns, upstream factory bans, and "
    "traditional alternatives (Dal Bhat, metal bottles, woven bags)"
  ),
  "Project context": (
    "Guide/porter training (~150 students); trash collection above 3,000 m; "
    "garbage deposit checkpoints in Upper Mustang/Manaslu; lake clean-up with "
    "Tourism Minister upon election; Pokhara municipality recycling-factory plans; "
    "discourages plastic bottles — filters/tablets for local water"
  ),
}

CODING = {
  "Country": "NPL",
  "Country name": "NEPAL",
  "ID": "NPL_19",
  "Actortype": "cso",
  "Problem_awareness_pop": "high",
  "Problem_awareness_pol": "high",
  "Problem_severity": "high",
  "Problem_littering": "yes",
  "Problem_consumption": "yes",
  "Problem_recycling": "yes",
  "Problem_waste_mgmt": "yes",
  "Problem_production": "yes",
  "Problem_alternatives": "yes",
  "Problem_waste_segregation": "yes",
  "Problem_import": "no",
  "Impacts": (
    "tourism, wildlife, visual pollution, health, air pollution, cities, water, forests"
  ),
  "Governance": "yes",
  "Relevance_international_pol": "no",
  "Coordination_sectoral": "no",
  "Coordination_levels": "no",
  "Unclear_responsibilities": "no",
  "Actors": "yes",
  "Federal government": "yes",
  "Provincial government": "no",
  "Local government": "yes",
  "Students": "yes",
  "Private_sector": "yes",
  "Civil_society": "yes",
  "Science": "yes",
  "Households": "yes",
  "Education institutions": "yes",
  "Private_companies(hotels, shops, etc.)": "yes",
  "Implementation issues": "yes",
  "Monitoring": "yes",
  "Financial_resources": "yes",
  "Research": "yes",
  "Infrastructure": "yes",
  "Enforcement": "yes",
  "Policies in place": "yes",
  "Pol_epr": "no",
  "Pol_import": "no",
  "Pol_awarness": "yes",
  "Pol_education": "yes",
  "Pol_capacity": "yes",
  "Pol_RD": "yes",
  "Pol_tax": "no",
  "Pol_ban": "yes",
  "Pol_subsitutes": "yes",
  "Pol_clean_up": "yes",
  "Pol_upcycling": "no",
  "Pol_recycling": "yes",
  "Pol_waste_collection": "yes",
  "Solutions": "yes",
  "Sol_lead_agency": "no",
  "Sol_responsibilities": "no",
  "Sol_epr": "no",
  "Sol_awarness": "yes",
  "Sol_segregation": "yes",
  "Sol_upcycling": "no",
  "Sol_recycling": "yes",
  "Sol_education": "yes",
  "Sol_capacity": "yes",
  "Sol_RD": "yes",
  "Sol_tax": "no",
  "Sol_ban": "yes",
  "Sol_finance": "yes",
  "Sol_infrastructure": "yes",
  "Sol_subsitutes": "yes",
  "Sol_clean_up": "yes",
  "Sol_enforcement": "yes",
  "Sol_monitoring": "yes",
  "Traditions to build on (free-hand)": (
    "Gurung Bhangra woven carry bags; leaf plates and mud cups; copper tableware; "
    "Dal Bhat as local plastic-free food alternative to packaged noodles"
  ),
}

CODEBOOK = [
  ("Country", "Country ISO code"),
  ("Country name", "Country name"),
  ("ID", "InterviewID: ISO_Nr"),
  ("Actortype", "governmental, pol_party, science, cso, igo, private_sector, shop_owner, hotel_owner, household"),
  ("Problem_awareness_pop", "high, medium, low, NA — lack of awareness among population mentioned"),
  ("Problem_awareness_pol", "high, medium, low, NA — lack of awareness among policy makers mentioned"),
  ("Problem_severity", "high, medium, low, NA"),
  ("Problem_littering", "Littering is a problem; was mentioned: yes/no"),
  ("Problem_consumption", "High consumption is a problem; was mentioned: yes/no"),
  ("Problem_recycling", "No or too little recycling is a problem; was mentioned: yes/no"),
  ("Problem_waste_mgmt", "Ill waste management is a problem; was mentioned: yes/no"),
  ("Problem_production", "High production rates; was mentioned: yes/no"),
  ("Problem_alternatives", "No good alternatives; was mentioned: yes/no"),
  ("Problem_waste_segregation", "Lack of segregation was mentioned: yes/no"),
  ("Problem_import", "Plastic imports are a problem; was mentioned: yes/no"),
  ("Impacts", "Impacts mentioned on water, cities, biodiversity, health, etc. or NA"),
  ("Governance", "Mentioned: yes/no"),
  ("Relevance_international_pol", "International treaties are relevant for national level policy"),
  ("Coordination_sectoral", "Lack of coordination across sectors"),
  ("Coordination_levels", "Lack of coordination across levels"),
  ("Unclear_responsibilities", "Unclear responsibilities among the involved actors"),
  ("Actors", "Mentioned: yes/no"),
  ("Federal government", "Mentioned: yes/no"),
  ("Provincial government", "Mentioned: yes/no"),
  ("Local government", "Mentioned: yes/no"),
  ("Students", "Mentioned: yes/no"),
  ("Private_sector", "Mentioned: yes/no"),
  ("Civil_society", "Mentioned: yes/no"),
  ("Science", "Mentioned: yes/no"),
  ("Households", "Mentioned: yes/no"),
  ("Education institutions", "Mentioned: yes/no"),
  ("Private_companies(hotels, shops, etc.)", "Mentioned: yes/no"),
  ("Implementation issues", "Mentioned: yes/no"),
  ("Monitoring", "Lack of monitoring"),
  ("Financial_resources", "Lack of financial resources"),
  ("Research", "Lack of research"),
  ("Infrastructure", "Lack of infrastructure"),
  ("Enforcement", "Lack of enforcement"),
  ("Policies in place", "Mentioned: yes/no"),
  ("Pol_epr", "Extended producer responsibility"),
  ("Pol_import", "Import regulations"),
  ("Pol_awarness", "Awareness campaigns"),
  ("Pol_education", "Education programs"),
  ("Pol_capacity", "Capacity building"),
  ("Pol_RD", "Research & development"),
  ("Pol_tax", "Levies or taxes on specific products (often bags)"),
  ("Pol_ban", "Bans of specific products"),
  ("Pol_subsitutes", "Alternatives like reusable bags, or banana leaf plates"),
  ("Pol_clean_up", "Clean-up campaigns"),
  ("Pol_upcycling", "Making a new product, like blacktop"),
  ("Pol_recycling", "Making a new raw material"),
  ("Pol_waste_collection", "Waste collection"),
  ("Solutions", "Mentioned: yes/no"),
  ("Sol_lead_agency", "Procedural measures to establish lead agency"),
  ("Sol_responsibilities", "Clear responsibilities"),
  ("Sol_epr", "Extended producer responsibility"),
  ("Sol_awarness", "Awareness raising measures"),
  ("Sol_segregation", "Segregation of waste"),
  ("Sol_upcycling", "Upcycling of plastics"),
  ("Sol_recycling", "New raw materials"),
  ("Sol_education", "Education programs at schools"),
  ("Sol_capacity", "Capacity-building programs"),
  ("Sol_RD", "Research programs"),
  ("Sol_tax", "Taxes or levies on products"),
  ("Sol_ban", "Ban of products"),
  ("Sol_finance", "Financial measures"),
  ("Sol_infrastructure", "Infrastructural measures"),
  ("Sol_subsitutes", "Alternatives"),
  ("Sol_clean_up", "Clean-up campaigns"),
  ("Sol_enforcement", "Enforcement procedures"),
  ("Sol_monitoring", "Monitoring procedures"),
  ("Traditions to build on (free-hand)", "Freehand or NA"),
]

NOTES = {
  "Actortype": (
    "TAAN Pokhara president and Ethical Trekking NGO founder. Stakeholder list: cso. "
    "Codebook: cso (association/NGO leader; also trekking entrepreneur)."
  ),
  "Problem_awareness_pop": (
    "Domestic tourists throw plastics from buses; no culture of carrying shopping bags; "
    "tourism sector more conscious but general public behaviour lags."
  ),
  "Problem_awareness_pol": (
    "Policies exist but not implemented; lack of monitoring and willingness to implement; "
    "outdated act; government not strict enough — need research to affect policy."
  ),
  "Problem_severity": (
    "Significant problem — awful on trails, harms animals, damages national image, "
    "pains TAAN president personally; junk food, noodles, beer bottles massive problem."
  ),
  "Problem_littering": (
    "Trash and plastic everywhere on trails; domestic tourists litter from buses; "
    "Pokhara lakeside hotspot; lake full of garbage at election."
  ),
  "Problem_consumption": (
    "Junk food, noodles, beer bottles, plastic bottles; discourage trekkers from buying "
    "plastic bottles — use filters or tablets for local water."
  ),
  "Problem_recycling": (
    "Promote recycling with ACAP and local governments; Pokhara municipality planning "
    "recycling factory — but village-level bury/burn persists."
  ),
  "Problem_waste_mgmt": (
    "Need better garbage management; sometimes nowhere to throw trash at lakeside; "
    "villages bury or burn; hard to bring all waste down from altitude."
  ),
  "Problem_production": (
    "Advocates closing factories that make plastic bags; government should totally ban "
    "plastic factories — upstream production intervention."
  ),
  "Problem_alternatives": (
    "People lack shopping-bag culture; promote metal bottles, treated water, Dal Bhat "
    "over packaged noodles — alternatives exist but uptake weak."
  ),
  "Problem_waste_segregation": (
    "Need deposit places to separate trash with sheds so waste doesn't fly away."
  ),
  "Impacts": (
    "Tourism: bad country image for visitors to mountains. Wildlife: harmful if animals "
    "eat plastic. Visual pollution: looks awful on trails. Health/air pollution: burning "
    "plastic causes more problems. Cities: Pokhara lakeside. Water: lake garbage. Forests: "
    "trekking trail degradation above 3,000 m."
  ),
  "Federal government": (
    "Tourism Minister joined lake clean-up; government should be strict, close plastic-bag "
    "factories, update outdated act, implement policies."
  ),
  "Local government": (
    "Local governments trying their best; some prohibited black plastics; Pokhara "
    "municipality recycling-factory plans and twice-weekly truck collection; could reward "
    "those who minimize plastic."
  ),
  "Students": (
    "Nationwide school campaigns advocated; some schools banned junk food and noodles — "
    "kids bring food in proper tiffins."
  ),
  "Private_sector": (
    "180 TAAN member trekking companies; tourism entrepreneurs minimize plastic; Ethical "
    "Trekking eco-friendly model; guides/porters trained as responsible actors."
  ),
  "Civil_society": (
    "TAAN, Ethical Trekking, ACAP, NTNC, Nepal Army clean-up efforts; mother groups and "
    "youth clubs weekly cleanups."
  ),
  "Science": (
    "'We need research like yours to affect government policy' — research-to-policy "
    "pipeline sought."
  ),
  "Households": (
    "Domestic travellers and general public lack reusable-bag habit; shopping behaviour "
    "change needed."
  ),
  "Education institutions": (
    "School junk-food bans and tiffin programmes; nationwide school campaigns proposed."
  ),
  "Private_companies(hotels, shops, etc.)": (
    "Trekking agencies and teahouse supply chain — noodles, junk food, beer bottles "
    "identified as problem products on trails."
  ),
  "Monitoring": (
    "Explicit: lack of monitoring and lack of willingness to implement policies."
  ),
  "Financial_resources": (
    "Off-season porters lack paid work; need funding for paid 'clean campaigns' on trails; "
    "economic support for remote mountain areas."
  ),
  "Research": (
    "Research like PEGO study needed to influence government policy — gap between evidence "
    "and policy action."
  ),
  "Infrastructure": (
    "Deposit/segregation points with sheds; Pokhara trucks twice weekly vs villages with "
    "no collection; municipality seeking land for recycling plant."
  ),
  "Enforcement": (
    "Policies not implemented properly; no fines for plastic bags (only monthly trash fees); "
    "if law strictly implemented problem could be cured; deposit scheme hard to scale."
  ),
  "Policies in place": (
    "Black plastic local bans; garbage deposit at Upper Mustang/Manaslu checkpoints; "
    "guide/porter trash collection above 3,000 m; school junk-food bans; lake clean-up; "
    "weekly community cleanups; monthly trash-collection fees."
  ),
  "Pol_awarness": (
    "Sector campaigns with ACAP/local government; discourage plastic bottles; more "
    "campaigns and promotion about effects needed."
  ),
  "Pol_education": (
    "Guide/porter training ~150 students on trash responsibility; nationwide school "
    "campaigns."
  ),
  "Pol_capacity": (
    "TAAN training programmes for guides and porters on ecological responsibility and "
    "waste handling."
  ),
  "Pol_RD": (
    "Calls for research (PEGO) to inform and affect national government policy."
  ),
  "Pol_ban": (
    "Some local governments prohibited black plastics (already hard to find); advocates "
    "total ban on plastic-bag factories."
  ),
  "Pol_subsitutes": (
    "Metal bottles and water filters/tablets; woven Gurung Bhangra bags; leaf plates; "
    "Dal Bhat instead of packaged noodles."
  ),
  "Pol_clean_up": (
    "Lake clean-up with Tourism Minister; mother groups/youth clubs weekly cleanups; "
    "guides collect trash above 3,000 m; proposed paid off-season trail clean campaigns."
  ),
  "Pol_recycling": (
    "TAAN/ACAP promote recycling; Pokhara municipality plans recycling factory on "
    "allocated land."
  ),
  "Pol_waste_collection": (
    "Pokhara municipal trucks twice weekly; monthly trash-collection fees; garbage-list "
    "deposit scheme at restricted-area checkpoints."
  ),
  "Sol_awarness": (
    "More campaigns nationwide; promotion of plastic effects; school-level behaviour change."
  ),
  "Sol_segregation": (
    "Proper infrastructure — deposit places to separate trash with sheds preventing "
    "wind dispersal."
  ),
  "Sol_recycling": (
    "Municipal recycling-factory development in Pokhara; sector promotion of recycling."
  ),
  "Sol_education": (
    "Nationwide school campaigns; expanded guide/porter training on trail stewardship."
  ),
  "Sol_capacity": (
    "Scale guide/porter training; paid off-season clean campaigns employ trained workforce."
  ),
  "Sol_RD": (
    "Research-to-policy pipeline — academic/policy research to move government action."
  ),
  "Sol_ban": (
    "Close plastic-bag factories; government should totally ban plastic factories."
  ),
  "Sol_finance": (
    "Fund paid off-season clean campaigns for porters/guides; local governments reward "
    "those who minimize plastic; deposit scheme at checkpoints."
  ),
  "Sol_infrastructure": (
    "Segregation deposit points with sheds; municipal recycling plant; improved village "
    "collection vs bury/burn."
  ),
  "Sol_subsitutes": (
    "Metal bottles and treated local water; Dal Bhat (local, plastic-free) over noodles; "
    "traditional woven bags and leaf plates."
  ),
  "Sol_clean_up": (
    "Paid 'clean campaigns' in off-season — porters earn income while cleaning trails; "
    "reduces seasonal migration."
  ),
  "Sol_enforcement": (
    "Government must be strict; strict policy implementation; fines/enforcement currently "
    "absent for plastic bags."
  ),
  "Sol_monitoring": (
    "Address lack of monitoring identified as core implementation barrier."
  ),
  "Traditions to build on (free-hand)": (
    "Gurung Bhangra woven bags for carrying goods; leaf plates and mud cups; copper "
    "plates at home; Dal Bhat using local farm produce instead of imported packaged noodles."
  ),
}


def hfill(hex_):
  return PatternFill("solid", fgColor=hex_)


def hfont(bold=False, size=10, color="000000", italic=False):
  return Font(bold=bold, size=size, color=color, italic=italic)


thin = Side(style="thin", color="BFBFBF")


def thin_border():
  return Border(left=thin, right=thin, top=thin, bottom=thin)


WRAP_ALIGN = Alignment(wrap_text=True, vertical="top", horizontal="left")
CENTER_ALIGN = Alignment(wrap_text=True, vertical="center", horizontal="center")


def build_workbook():
  wb = openpyxl.Workbook()

  ws = wb.active
  ws.title = "Interview Coding"

  ws.merge_cells("A1:D1")
  title = ws["A1"]
  title.value = "Plastic Pollution Governance — Interview Coding (Codebook)"
  title.fill = hfill(DARK_BLUE)
  title.font = hfont(bold=True, size=13, color=WHITE)
  title.alignment = CENTER_ALIGN

  ws.merge_cells("A2:D2")
  meta = ws["A2"]
  meta.value = (
    "Interview: NPL_19  |  Country: NEPAL  |  TAAN president / Ethical Trekking  |  "
    "Actor: cso"
  )
  meta.fill = hfill(MID_BLUE)
  meta.font = hfont(size=9, color=WHITE)
  meta.alignment = CENTER_ALIGN

  headers = ["Variable", "Code", "Codebook definition", "Coding notes"]
  for col, header in enumerate(headers, start=1):
    cell = ws.cell(row=4, column=col, value=header)
    cell.fill = hfill(LIGHT_BLUE)
    cell.font = hfont(bold=True, size=10, color=DARK_BLUE)
    cell.alignment = CENTER_ALIGN
    cell.border = thin_border()

  row = 5
  for variable, definition in CODEBOOK:
    values = [
      variable,
      CODING[variable],
      definition,
      NOTES.get(variable, ""),
    ]
    fill = hfill(LIGHT_GREEN) if row % 2 == 0 else hfill(WHITE)
    for col, value in enumerate(values, start=1):
      cell = ws.cell(row=row, column=col, value=value)
      cell.fill = fill
      cell.font = hfont(size=10)
      cell.alignment = WRAP_ALIGN
      cell.border = thin_border()
    row += 1

  ws.column_dimensions["A"].width = 34
  ws.column_dimensions["B"].width = 28
  ws.column_dimensions["C"].width = 52
  ws.column_dimensions["D"].width = 58
  ws.freeze_panes = "A5"

  wide = wb.create_sheet("Wide Format")
  wide_headers = [variable for variable, _ in CODEBOOK]
  for col, header in enumerate(wide_headers, start=1):
    cell = wide.cell(row=1, column=col, value=header)
    cell.fill = hfill(LIGHT_BLUE)
    cell.font = hfont(bold=True, size=9, color=DARK_BLUE)
    cell.alignment = CENTER_ALIGN
    cell.border = thin_border()
  for col, header in enumerate(wide_headers, start=1):
    cell = wide.cell(row=2, column=col, value=CODING[header])
    cell.fill = hfill(ORANGE)
    cell.alignment = WRAP_ALIGN
    cell.border = thin_border()
    wide.column_dimensions[get_column_letter(col)].width = max(14, min(28, len(header) + 2))

  ref = wb.create_sheet("Codebook Reference")
  ref.merge_cells("A1:B1")
  ref["A1"].value = "Codebook Variable Definitions"
  ref["A1"].fill = hfill(DARK_BLUE)
  ref["A1"].font = hfont(bold=True, size=12, color=WHITE)
  ref["A1"].alignment = CENTER_ALIGN
  for idx, (variable, definition) in enumerate(CODEBOOK, start=3):
    ref.cell(row=idx, column=1, value=variable).font = hfont(bold=True)
    ref.cell(row=idx, column=2, value=definition).alignment = WRAP_ALIGN
  ref.column_dimensions["A"].width = 34
  ref.column_dimensions["B"].width = 70

  stakeholder = wb.create_sheet("Stakeholder Analysis")
  stakeholder.merge_cells("A1:B1")
  stakeholder["A1"].value = "Stakeholder Analysis — NPL_19"
  stakeholder["A1"].fill = hfill(DARK_BLUE)
  stakeholder["A1"].font = hfont(bold=True, size=12, color=WHITE)
  stakeholder["A1"].alignment = CENTER_ALIGN
  for idx, (key, value) in enumerate(STAKEHOLDER.items(), start=3):
    stakeholder.cell(row=idx, column=1, value=key).font = hfont(bold=True)
    cell = stakeholder.cell(row=idx, column=2, value=value)
    cell.alignment = WRAP_ALIGN
  stakeholder.column_dimensions["A"].width = 28
  stakeholder.column_dimensions["B"].width = 90

  sources = wb.create_sheet("Sources")
  sources.merge_cells("A1:B1")
  sources["A1"].value = "Source Documents Used for Verification"
  sources["A1"].fill = hfill(DARK_BLUE)
  sources["A1"].font = hfont(bold=True, size=12, color=WHITE)
  sources["A1"].alignment = CENTER_ALIGN
  for idx, (doc, detail) in enumerate(SOURCES, start=3):
    sources.cell(row=idx, column=1, value=doc).font = hfont(bold=True)
    sources.cell(row=idx, column=2, value=detail).alignment = WRAP_ALIGN
  sources.column_dimensions["A"].width = 42
  sources.column_dimensions["B"].width = 60

  wb.save(OUTPUT_PATH)
  print(f"Saved {OUTPUT_PATH}")


if __name__ == "__main__":
  build_workbook()
