"""
Builds the 4P Index (Plastic Pollution Policy Index) coding table for:

Loi n. 2022-18 du 23 mai 2022 autorisant la creation d'une societe denommee Societe nationale de
Gestion integree des Dechets (SONAGED S.A.) -- Law No. 2022-18 of 23 May 2022 authorising creation of
the National Integrated Waste Management Company (SONAGED S.A.).

Source: user-uploaded PDF (Journal Officiel de la Republique du Senegal, 04 juin 2022, pp. 759-760),
cross-checked against the official FAOLEX copy (https://faolex.fao.org/docs/pdf/sen216236.pdf, LEX-
FAOC216236) which is textually identical (including OCR artefacts), confirming completeness/accuracy.
The task brief's urbanisme.gouv.sn URL corresponds to a Ministry news article ("GESTION DES ORDURES:
SONAGED SA prend le relai de l'UCG") summarising the National Assembly debate and providing contextual
figures (jobs, tonnage, budget) not present in the law's own six operative articles -- it is cited
below as supplementary context only, not as a basis for scoring, consistent with Rule 7 (score only on
explicit textual evidence).

Produces:
  - senegal_loi_2022-18_4p_index.csv
  - senegal_loi_2022-18_4p_index.xlsx (formatted, wrapped, one sheet)
"""
import csv
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [
    "A_policy_name",
    "B_policy_url",
    "C_policy_year",
    "D_policy_objective",
    "E_policy_target",
    "F_policy_target_text",
    "G_policy_type",
    "H_policy_type_justification",
    "I_policy_integration",
    "J_policy_sectors_list",
    "K_policy_circularity",
    "L_policy_lifecycle_phases_list",
    "M_policy_budget",
    "N_policy_budget_text",
    "O_policy_score",
    "P_instrument_type",
    "Q_instrument_lifecycle_stage",
    "R_instrument_description",
    "S_instrument_in_force",
    "T_instrument_implementation",
    "U_instrument_implementation_text",
    "V_instrument_score",
    "W_comments",
]

# ---------------------------------------------------------------------------
# Section A - Policy-level fields (identical across every row)
# ---------------------------------------------------------------------------

POLICY_NAME = (
    "Loi n° 2022-18 du 23 mai 2022 autorisant la création d'une société dénommée Société nationale "
    "de Gestion intégrée des Déchets (SONAGED S.A.)"
)
POLICY_URL = (
    "https://faolex.fao.org/docs/pdf/sen216236.pdf (FAOLEX LEX-FAOC216236, full text, matches the "
    "user-uploaded PDF) ; https://www.fao.org/faolex/results/details/en/c/LEX-FAOC216236 (FAOLEX "
    "record page, as given in task brief) ; Journal Officiel de la République du Sénégal, 04 juin "
    "2022, pp. 759-760 (source of the uploaded PDF) ; "
    "https://urbanisme.gouv.sn/actualites/gestion-des-ordures-sonaged-sa-prend-le-relai-de-l%E2%80%99ucg "
    "(Ministry news article on the National Assembly debate -- supplementary context only, cited in "
    "Column W where relevant, not used for scoring; corresponds to the task brief's truncated URL)"
)
POLICY_YEAR = 2022
POLICY_OBJECTIVE = (
    "Authorises creation of SONAGED S.A., a State-owned company under the technical oversight of the "
    "Minister for Public Hygiene and financial oversight of the Minister for Finance (Art. 1), tasked "
    "with coordinating integrated solid-waste management nationwide (Art. 2): collection, transport, "
    "landfilling, treatment and valorisation of solid waste; managing waste treatment/valorisation "
    "equipment and infrastructure; acting as sector regulator; improving waste management across the "
    "value chain while promoting inter-municipal cooperation; developing a circular economy and "
    "valorising waste; proposing institutional/regulatory/financial governance reforms; supervising "
    "State waste programmes and projects; and maximising private-sector involvement. SONAGED absorbs "
    "the pre-existing Unité de Coordination de la Gestion des déchets solides (UCG) and 'les autres "
    "projets et programmes publics de gestion intégrée des déchets solides' (Art. 5) -- per the "
    "exposé des motifs, this includes PROMOGED (created by décret n° 2021-831 du 22 juin 2021). The "
    "law's six operative articles never mention 'plastique' or synthetic polymers explicitly and set "
    "no plastics-specific target; its relevance rests on Plastics Relevance Filter criterion (c) -- "
    "it is the foundational institutional/governance instrument determining how effectively Senegal's "
    "municipal solid waste (a large share of which is plastic packaging and bags) will be collected, "
    "treated and valorised (recycled) nationally, and it explicitly tasks the new company with "
    "developing a circular economy for waste (Art. 2)."
)
POLICY_TARGET = 0
POLICY_TARGET_TEXT = ""
POLICY_TYPE = 1
POLICY_TYPE_JUSTIFICATION = (
    "Legislation: 'L'Assemblée nationale a adopté, en sa séance du jeudi 05 mai 2022 ; Le Président "
    "de la République promulgue la loi dont la teneur suit...' Enacted by the National Assembly "
    "(voted unanimously, per the Ministry's news account) and promulgated by the President -> coded "
    "1 per Rule 6."
)
POLICY_INTEGRATION = 0.75
POLICY_SECTORS_LIST = (
    "waste management, local government/municipalities, private sector, public hygiene/health, "
    "finance, governance/institutional reform"
)
POLICY_CIRCULARITY = 0.50
POLICY_LIFECYCLE_PHASES_LIST = "recycling, disposal"
POLICY_BUDGET = 1
POLICY_BUDGET_TEXT = (
    "Art. 4: 'Les ressources de la SONAGED S.A proviennent d'une dotation de l'Etat, de toutes "
    "autres ressources autorisées par les lois et règlements lui permettant d'assurer la collecte "
    "des déchets sur tout le territoire national et de couvrir les charges de fonctionnement, ainsi "
    "que des produits tirés de ses activités d'exploitation.' Coded 1 (ring-fenced): beyond a State "
    "endowment, the law dedicates SONAGED's own operating revenue ('produits tirés de ses activités "
    "d'exploitation') specifically to funding its national waste-collection/management mandate -- "
    "the policy generates its own funding, per Column M's rule 1 criterion, even though (per the "
    "Ministry's news account, not the law itself) the operating budget itself is not stated in "
    "figures within the six operative articles."
)
POLICY_SCORE_FORMULA = "[auto] O = (G + I + K + M) / 4"

policy_fields = dict(
    A_policy_name=POLICY_NAME,
    B_policy_url=POLICY_URL,
    C_policy_year=POLICY_YEAR,
    D_policy_objective=POLICY_OBJECTIVE,
    E_policy_target=POLICY_TARGET,
    F_policy_target_text=POLICY_TARGET_TEXT,
    G_policy_type=POLICY_TYPE,
    H_policy_type_justification=POLICY_TYPE_JUSTIFICATION,
    I_policy_integration=POLICY_INTEGRATION,
    J_policy_sectors_list=POLICY_SECTORS_LIST,
    K_policy_circularity=POLICY_CIRCULARITY,
    L_policy_lifecycle_phases_list=POLICY_LIFECYCLE_PHASES_LIST,
    M_policy_budget=POLICY_BUDGET,
    N_policy_budget_text=POLICY_BUDGET_TEXT,
    O_policy_score=POLICY_SCORE_FORMULA,
)

# ---------------------------------------------------------------------------
# Section B - Instrument-level fields (one dict per row)
# ---------------------------------------------------------------------------

instruments = [
    dict(
        P_instrument_type=0.40,
        Q_instrument_lifecycle_stage="Waste management",
        R_instrument_description=(
            "Art. 1, 5, 6: authorises creation of SONAGED S.A. as a State-owned company under the "
            "technical oversight of the Minister for Public Hygiene and financial oversight of the "
            "Minister for Finance, with its mission (Art. 2) to coordinate integrated solid-waste "
            "management nationwide, act as sector regulator, and propose institutional/regulatory/"
            "financial governance reforms. Art. 5 transfers the pre-existing UCG and 'les autres "
            "projets et programmes publics de gestion intégrée des déchets solides' (per the exposé "
            "des motifs, including PROMOGED) to SONAGED, consolidating a sector previously fragmented "
            "across seven successive institutions since the late 1990s. Art. 6: SONAGED's internal "
            "organisation is fixed by bylaws approved by decree."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.50,
        U_instrument_implementation_text=(
            "Authority: 'La SONAGED S.A. est placée sous la tutelle technique du Ministre chargé de "
            "l'Hygiène publique et sous la tutelle financière du Ministre chargé des Finances' (Art. "
            "1) -- explicit designation of two supervising ministries (+0.25). Unconditional: the "
            "creation, mission and transfer of assets/programmes to SONAGED are stated without "
            "exemption or carve-out (+0.25). Monitoring: not awarded -- no inspection, audit or data-"
            "collection requirement is specified for SONAGED's own institutional creation/consolidation "
            "in this law (only that its bylaws will later be approved by decree, Art. 6). Enforcement: "
            "not awarded -- no fine or penalty is attached to this instrument."
        ),
        V_instrument_score=None,
        W_comments=(
            "Coded per Rule 11 as an institutional/multi-level coordination instrument -- SONAGED "
            "replaces a chain of seven prior, unstable waste-management bodies (CUD, HAPD, APRODAK, "
            "Entente CADAK-CAR, APROSEN, SOPROSEN, UCG) and absorbs PROMOGED, consolidating national "
            "waste governance. Plastics relevance rests on filter criterion (c): the institution's "
            "core mandate (national solid-waste coordination) commonly applies to plastic waste "
            "streams, and it is tasked with sector-reform proposals that could, in future, include "
            "plastics-specific measures (cf. Loi n° 2020-04's EPR/tax provisions)."
        ),
    ),
    dict(
        P_instrument_type=0.80,
        Q_instrument_lifecycle_stage="Waste management",
        R_instrument_description=(
            "Art. 2, 1st-2nd bullets: SONAGED is tasked with 'assurer la collecte, le transport, la "
            "mise en décharge, le traitement et la valorisation des déchets solides sur l'ensemble du "
            "territoire national' and 'assurer la gestion des équipements et infrastructures de "
            "traitement et de valorisation des déchets' -- i.e. direct national operational "
            "responsibility for the full solid-waste value chain and for managing (and, per the "
            "exposé des motifs and Ministry press account, building/perpetuating) waste treatment and "
            "valorisation infrastructure."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.50,
        U_instrument_implementation_text=(
            "Authority: SONAGED is the explicitly designated national operator ('elle est notamment "
            "chargée de...') (+0.25). Unconditional: the mandate applies 'sur l'ensemble du territoire "
            "national' without stated exemption (+0.25). Monitoring/Enforcement: not awarded -- the "
            "law itself specifies no inspection, reporting, or penalty regime attached to this "
            "operational mandate (further detail is left to bylaws approved by decree, Art. 6, not "
            "yet available)."
        ),
        V_instrument_score=None,
        W_comments=(
            "Coded as Planned government investment (0.80) given the explicit infrastructure-"
            "management mandate for waste treatment/valorisation facilities, matching the prompt's own "
            "example ('recycling infrastructure investment'). Per the Ministry's news account "
            "(supplementary context, not scored): SONAGED aims to guarantee 'la pérennité des "
            "infrastructures de traitement réalisées par l'Etat' and process 2.5 million tonnes of "
            "waste/year -- not stated in the law's own articles, so not used to inflate Column T."
        ),
    ),
    dict(
        P_instrument_type=0.40,
        Q_instrument_lifecycle_stage="Recycling",
        R_instrument_description=(
            "Art. 2, 3rd-4th bullets: SONAGED is tasked to 'améliorer la gestion des déchets solides "
            "sur toute la chaîne de valeur et de promouvoir une gestion intégrée dans toutes les "
            "communes en favorisant l'intercommunalité' and to 'développer une économie circulaire et "
            "valoriser les déchets en tenant compte de la dimension socio-économique.'"
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.50,
        U_instrument_implementation_text=(
            "Authority: SONAGED explicitly designated as responsible entity (+0.25). Unconditional: "
            "mandate stated for 'toutes les communes' without carve-out (+0.25). Monitoring/"
            "Enforcement: not awarded -- no inspection, reporting metric, or penalty specified for "
            "circular-economy development in the law's text."
        ),
        V_instrument_score=None,
        W_comments=(
            "Combines a circular-economy/valorisation development mandate with multi-level "
            "municipal coordination ('favorisant l'intercommunalité') -- coded per Rule 11 as "
            "Procedural (0.40), since no concrete regulatory tool (ban, standard, EPR) or economic "
            "incentive is specified for achieving circularity; the mandate is aspirational/framework-"
            "setting rather than self-executing. Directly plastics-relevant in substance (circular "
            "economy for waste commonly centres on plastics valorisation) though 'plastique' is not "
            "named."
        ),
    ),
    dict(
        P_instrument_type=0.60,
        Q_instrument_lifecycle_stage="Waste management",
        R_instrument_description=(
            "Art. 4: 'Les ressources de la SONAGED S.A proviennent d'une dotation de l'Etat, de "
            "toutes autres ressources autorisées par les lois et règlements lui permettant d'assurer "
            "la collecte des déchets sur tout le territoire national et de couvrir les charges de "
            "fonctionnement, ainsi que des produits tirés de ses activités d'exploitation.'"
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.50,
        U_instrument_implementation_text=(
            "Authority: SONAGED explicitly designated as the recipient/manager of these resources "
            "(+0.25). Unconditional: the funding-sources provision applies without stated exemption "
            "(+0.25). Monitoring/Enforcement: not awarded -- the law specifies no audit, reporting "
            "or penalty mechanism for use of these resources (financial oversight is only generally "
            "assigned to the Minister for Finance under Art. 1, not tied to a specific monitoring "
            "duty in Art. 4)."
        ),
        V_instrument_score=None,
        W_comments=(
            "Coded as an Economic instrument (0.60): a dedicated State endowment plus self-generated "
            "operating revenue ('produits tirés de ses activités d'exploitation') earmarked "
            "specifically to fund national waste-collection operations and running costs -- the "
            "policy 'generates its own funding' per Column M's ring-fencing criterion, hence M = 1 "
            "at policy level. No specific tax/levy rate or figure is stated in the law itself; per "
            "the Ministry's news account (not scored here), roughly 200 billion FCFA/year is deemed "
            "necessary to operate the company across 557 local authorities."
        ),
    ),
    dict(
        P_instrument_type=0.40,
        Q_instrument_lifecycle_stage="Waste management",
        R_instrument_description=(
            "Art. 2, 5th bullet: SONAGED is tasked to 'proposer des axes de réforme visant à "
            "améliorer la gouvernance du secteur sur le plan institutionnel, réglementaire et "
            "financier' -- a mandatory duty to formulate sector-governance reform proposals "
            "(institutional, regulatory and financial), separate from the institution-creation "
            "instrument itself."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.50,
        U_instrument_implementation_text=(
            "Authority: SONAGED explicitly designated as the body tasked with formulating reform "
            "proposals (+0.25). Unconditional: the duty is stated without carve-out (+0.25). "
            "Monitoring/Enforcement: not awarded -- no mechanism specified for reviewing, adopting, "
            "or enforcing the proposed reforms within this law."
        ),
        V_instrument_score=None,
        W_comments=(
            "In-force test (Rule 12): the duty to formulate proposals uses mandatory language ('elle "
            "est chargée de proposer') and is therefore coded in force (S=1) even though the *content* "
            "of any future regulatory/fiscal reform (which could include plastics-specific measures, "
            "as occurred via Loi n° 2020-04's implementing decrees) is not yet known and is not itself "
            "coded here. Coded as Procedural (0.40) -- a planning/agenda-setting instrument, not a "
            "self-executing regulatory or economic tool."
        ),
    ),
]

# ---------------------------------------------------------------------------
# Assemble rows, computing auto-fields O and V
# ---------------------------------------------------------------------------

O_score = round(
    (policy_fields["G_policy_type"] + policy_fields["I_policy_integration"]
     + policy_fields["K_policy_circularity"] + policy_fields["M_policy_budget"]) / 4,
    3,
)
policy_fields["O_policy_score"] = f"{O_score} [auto = (G+I+K+M)/4]"

rows = []
for inst in instruments:
    row = dict(policy_fields)
    row.update(inst)
    s = row["S_instrument_in_force"]
    p = row["P_instrument_type"]
    t = row["T_instrument_implementation"]
    v = round(s * (p + t) / 2, 3)
    row["V_instrument_score"] = f"{v} [auto = S*(P+T)/2]"
    rows.append(row)

with open("4p_index/senegal_loi_2022-18_4p_index.csv", "w", newline="", encoding="utf-8") as f:
    writer = csv.DictWriter(f, fieldnames=COLUMNS)
    writer.writeheader()
    for row in rows:
        writer.writerow(row)

df = pd.DataFrame(rows, columns=COLUMNS)
xlsx_path = "4p_index/senegal_loi_2022-18_4p_index.xlsx"
df.to_excel(xlsx_path, sheet_name="4P_Index_Loi2022-18", index=False)

wb = load_workbook(xlsx_path)
ws = wb["4P_Index_Loi2022-18"]

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF")
wrap_top = Alignment(wrap_text=True, vertical="top")

col_widths = {
    "A": 34, "B": 30, "C": 10, "D": 46, "E": 10, "F": 30, "G": 10, "H": 34,
    "I": 10, "J": 30, "K": 10, "L": 26, "M": 10, "N": 30, "O": 16,
    "P": 12, "Q": 16, "R": 46, "S": 10, "T": 12, "U": 40, "V": 16, "W": 40,
}
for letter in [get_column_letter(c) for c in range(1, len(COLUMNS) + 1)]:
    ws.column_dimensions[letter].width = col_widths.get(letter, 20)

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for cell in row:
        cell.alignment = wrap_top

ws.freeze_panes = "A2"
ws.row_dimensions[1].height = 30
for r in range(2, ws.max_row + 1):
    ws.row_dimensions[r].height = 130

wb.save(xlsx_path)
print("Wrote", xlsx_path, "and CSV with", len(rows), "instrument rows.")
print("Policy score O =", O_score)
for row in rows:
    print(row["Q_instrument_lifecycle_stage"], "->", row["V_instrument_score"])
