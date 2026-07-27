"""
Builds the 4P Index (Plastic Pollution Policy Index) coding table for:

Decret n. 2016-1804 du 22 novembre 2016 portant application de la loi n. 2015-18 du 13 juillet 2015
portant Code de la Peche maritime (Senegal).

Source: user-uploaded PDF (clean, machine-readable text; mirror hosted by droit-afrique.com),
consistent with the task-brief URL www.ditp.gouv.sn/download/file/fid/76. The decree repeals and
replaces Decret n. 98-498 du 10 juin 1998 (Art. 70) and covers: fisheries advisory bodies (national
and local), vessel/gear classification, fishing-authorisation procedures, conservation measures
(mesh sizes, gear-type and vessel-size bans, minimum catch sizes, fishing zones), monitoring/
surveillance (vessel marking, at-sea observers) and an infractions advisory commission.

KEY FINDING: an exhaustive full-text search for "plastique", "nylon", "monofilament", "dechet",
"immersion", "rejet", "abandon", "epave", "detritus", "residu", "ordure" and "pollution" returned NO
matches anywhere in this 26-page implementing decree. Unlike its parent law (Loi n. 2015-18, Art. 66,
which explicitly bans nylon monofilament/multimonofilament gillnets -- already coded as its own
policy entry in this series), this decree does not itself repeat, extend or further operationalise
that plastics-relevant gear ban; the ban remains self-executing directly from the law's own text.
Only two provisions in this decree bear even an indirect, borderline relationship to gear-based
marine-plastics governance, and both are coded here with explicit Rule 8 uncertainty flags.

Produces:
  - senegal_decret_2016-1804_4p_index.csv
  - senegal_decret_2016-1804_4p_index.xlsx (formatted, wrapped, one sheet)
"""
import csv
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [
    "A_policy_name",
    "B_policy_url",
    "C_policy_year",
    "D_policy_objective",
    "E_policy_target",
    "F_policy_target_text",
    "G_policy_type",
    "H_policy_type_justification",
    "I_policy_integration",
    "J_policy_sectors_list",
    "K_policy_circularity",
    "L_policy_lifecycle_phases_list",
    "M_policy_budget",
    "N_policy_budget_text",
    "O_policy_score",
    "P_instrument_type",
    "Q_instrument_lifecycle_stage",
    "R_instrument_description",
    "S_instrument_in_force",
    "T_instrument_implementation",
    "U_instrument_implementation_text",
    "V_instrument_score",
    "W_comments",
]

# ---------------------------------------------------------------------------
# Section A - Policy-level fields (identical across every row)
# ---------------------------------------------------------------------------

POLICY_NAME = (
    "Décret n° 2016-1804 du 22 novembre 2016 portant application de la loi n° 2015-18 du 13 juillet "
    "2015 portant Code de la Pêche maritime"
)
POLICY_URL = (
    "www.ditp.gouv.sn/download/file/fid/76 (as provided in task brief) ; user-uploaded PDF sourced "
    "from the droit-afrique.com legal-text mirror"
)
POLICY_YEAR = 2016
POLICY_OBJECTIVE = (
    "Implements Loi n° 2015-18 portant Code de la Pêche maritime, repealing and replacing the prior "
    "implementing decree (Décret n° 98-498 du 10 juin 1998, Art. 70). It establishes the national and "
    "local fisheries advisory councils, vessel/gear classification rules, fishing-licence procedures, "
    "conservation measures (minimum mesh sizes by gear type, gear-type and vessel-tonnage bans, "
    "minimum catch sizes/weights by species, delimited fishing zones), monitoring/surveillance "
    "provisions (vessel identification marking, at-sea observers) and an advisory commission on "
    "fishing infractions. It is not a plastics policy: an exhaustive full-text search found no "
    "mention of 'plastique', 'nylon', 'monofilament', 'déchet', 'immersion', 'rejet', 'abandon', "
    "'épave', 'détritus', 'résidu', 'ordure' or 'pollution' anywhere in its 26 pages. Critically, this "
    "decree does not itself repeat, extend or further operationalise its parent law's plastics-"
    "relevant provision (Loi n° 2015-18, Art. 66's ban on nylon monofilament/multimonofilament "
    "gillnets, already coded as its own policy entry in this series) -- that ban remains self-"
    "executing directly from the Code's own text, with no implementing detail added here. Only two "
    "provisions in this decree bear an indirect, borderline relationship to gear-based marine "
    "governance that could, in principle, extend to material-based (including plastic) gear "
    "restrictions in the future, and both are coded below with explicit uncertainty flagged."
)
POLICY_TARGET = 0
POLICY_TARGET_TEXT = ""
POLICY_TYPE = 0.75
POLICY_TYPE_JUSTIFICATION = (
    "A presidential implementing decree ('Décret'), issued under the executive branch's regulatory "
    "power to apply an existing Act of Parliament (Loi n° 2015-18), not itself an Act adopted by the "
    "National Assembly -> sub-legislative regulation, coded 0.75 per Rule 6, the same tier as the "
    "other implementing decrees in this series (e.g. Décret n° 2001-282)."
)
POLICY_INTEGRATION = 1
POLICY_SECTORS_LIST = (
    "fisheries, maritime/marine transport, environment, industry, defense/security, justice, "
    "finance, foreign affairs, local government, research"
)
POLICY_CIRCULARITY = 0.25
POLICY_LIFECYCLE_PHASES_LIST = "environmental leakage"
POLICY_BUDGET = 0
POLICY_BUDGET_TEXT = ""
POLICY_SCORE_FORMULA = "[auto] O = (G + I + K + M) / 4"

policy_fields = dict(
    A_policy_name=POLICY_NAME,
    B_policy_url=POLICY_URL,
    C_policy_year=POLICY_YEAR,
    D_policy_objective=POLICY_OBJECTIVE,
    E_policy_target=POLICY_TARGET,
    F_policy_target_text=POLICY_TARGET_TEXT,
    G_policy_type=POLICY_TYPE,
    H_policy_type_justification=POLICY_TYPE_JUSTIFICATION,
    I_policy_integration=POLICY_INTEGRATION,
    J_policy_sectors_list=POLICY_SECTORS_LIST,
    K_policy_circularity=POLICY_CIRCULARITY,
    L_policy_lifecycle_phases_list=POLICY_LIFECYCLE_PHASES_LIST,
    M_policy_budget=POLICY_BUDGET,
    N_policy_budget_text=POLICY_BUDGET_TEXT,
    O_policy_score=POLICY_SCORE_FORMULA,
)

# ---------------------------------------------------------------------------
# Section B - Instrument-level fields (one dict per row)
# ---------------------------------------------------------------------------

instruments = [
    dict(
        P_instrument_type=0.40,
        Q_instrument_lifecycle_stage="Environmental leakage",
        R_instrument_description=(
            "Art. 3-8: elaborates the 'Conseil national consultatif des Pêches maritimes' (created by "
            "Art. 22 of the parent law), chaired by the Director of Maritime Fisheries, with a ~20-"
            "member composition spanning fisheries, environment, defence, finance, local governance, "
            "research and industry representatives, tasked among other things with advising on "
            "fisheries-management measures; and empowers the Minister to institute regional 'conseils "
            "locaux de pêche artisanale', whose missions include participating in local monitoring, "
            "control and surveillance of fisheries (Art. 6) and organising fishers to prevent conflicts "
            "over fishing methods (which would include gear-type disputes)."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.50,
        U_instrument_implementation_text=(
            "Authority: 'Le Conseil national consultatif des Pêches maritimes est présidé par le "
            "Directeur des Pêches maritimes' (Art. 4) (+0.25). Monitoring: local councils are "
            "explicitly tasked to 'participer à l'élaboration et à l'exécution des plans d'aménagement "
            "locaux des pêcheries et au système de suivi, contrôle et surveillance des pêches au niveau "
            "local' (Art. 6) (+0.25). Enforcement: no penalty is tied to non-participation in or non-"
            "establishment of these bodies -- not awarded. Unconditional: not awarded -- Art. 5 states "
            "regional councils 'peut instituer', an optional/enabling formulation."
        ),
        W_comments=(
            "Coded per Rule 11 as a multi-level fisheries-governance coordination instrument. Plastics "
            "relevance is genuinely indirect/borderline (Rule 8 uncertainty): this decree's own text "
            "never mentions plastics or gear materials, but these bodies are the institutional channels "
            "through which gear-type rules -- including, at the parent-law level, the plastics-relevant "
            "nylon-net ban (Loi n° 2015-18, Art. 66) -- are advised upon, monitored and locally enforced "
            "in practice. A stricter reading could exclude this instrument entirely as not meeting any "
            "filter criterion on this decree's text alone; it is included here on the strength of "
            "criterion (c)'s explicit 'multi-level enforcement coordination' example, with this "
            "uncertainty flagged rather than resolved unilaterally."
        ),
    ),
    dict(
        P_instrument_type=1.0,
        Q_instrument_lifecycle_stage="Environmental leakage",
        R_instrument_description=(
            "Art. 37: 'Le Ministre chargé de la Pêche maritime est habilité à prendre les mesures "
            "nécessaires concernant l'utilisation de tout dispositif ou gréement de nature à détruire "
            "les habitats naturels des espèces en vue de garantir la préservation des ressources et de "
            "l'environnement marins. Il peut promouvoir, au besoin rendre obligatoire, l'utilisation de "
            "tout engin ou dispositif sélectif ayant pour finalité la préservation de la biodiversité "
            "marine, de l'équilibre des stocks ou la gestion rationnelle des ressources.' A general "
            "enabling power allowing the Minister to regulate or restrict any habitat-destroying gear/"
            "rigging and to mandate the use of selective gear."
        ),
        S_instrument_in_force=0,
        T_instrument_implementation=0.25,
        U_instrument_implementation_text=(
            "Authority: 'le Ministre chargé de la Pêche maritime' is explicitly designated as the "
            "holder of this power (+0.25). Enforcement/Monitoring: not awarded -- no fine, inspection or "
            "data-collection mechanism is specified for this general enabling clause itself. "
            "Unconditional: not awarded -- the power is explicitly discretionary ('est habilité à "
            "prendre... Il peut promouvoir, au besoin rendre obligatoire')."
        ),
        W_comments=(
            "In-force test (Rule 12): 'est habilité à' and 'peut' are classic enabling-power language, "
            "not yet exercised for any new gear-material restriction within this decree's own text, so "
            "S=0. This is the clause structurally analogous to how the parent law's own nylon-net ban "
            "(Art. 66) could in principle be extended or complemented (e.g. to other synthetic-polymer "
            "gear types) via a future ministerial order under this power -- but no such order is "
            "evidenced here. Plastics relevance is again indirect/borderline and flagged per Rule 8: "
            "'habitat-destroying gear' is not itself plastics-specific language, but includes the type "
            "of concern (persistent synthetic gear/lost nets) that plastics-focused gear restrictions "
            "typically address."
        ),
    ),
]

# ---------------------------------------------------------------------------
# Assemble rows, computing auto-fields O and V
# ---------------------------------------------------------------------------

O_score = round(
    (policy_fields["G_policy_type"] + policy_fields["I_policy_integration"]
     + policy_fields["K_policy_circularity"] + policy_fields["M_policy_budget"]) / 4,
    3,
)
policy_fields["O_policy_score"] = f"{O_score} [auto = (G+I+K+M)/4]"

rows = []
for inst in instruments:
    row = dict(policy_fields)
    row.update(inst)
    s = row["S_instrument_in_force"]
    p = row["P_instrument_type"]
    t = row["T_instrument_implementation"]
    v = round(s * (p + t) / 2, 3)
    row["V_instrument_score"] = f"{v} [auto = S*(P+T)/2]"
    rows.append(row)

with open("4p_index/senegal_decret_2016-1804_4p_index.csv", "w", newline="", encoding="utf-8") as f:
    writer = csv.DictWriter(f, fieldnames=COLUMNS)
    writer.writeheader()
    for row in rows:
        writer.writerow(row)

df = pd.DataFrame(rows, columns=COLUMNS)
xlsx_path = "4p_index/senegal_decret_2016-1804_4p_index.xlsx"
df.to_excel(xlsx_path, sheet_name="4P_Index_Decret2016-1804", index=False)

wb = load_workbook(xlsx_path)
ws = wb["4P_Index_Decret2016-1804"]

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF")
wrap_top = Alignment(wrap_text=True, vertical="top")

col_widths = {
    "A": 34, "B": 30, "C": 10, "D": 46, "E": 10, "F": 30, "G": 10, "H": 34,
    "I": 10, "J": 30, "K": 10, "L": 26, "M": 10, "N": 30, "O": 16,
    "P": 12, "Q": 16, "R": 46, "S": 10, "T": 12, "U": 40, "V": 16, "W": 40,
}
for letter in [get_column_letter(c) for c in range(1, len(COLUMNS) + 1)]:
    ws.column_dimensions[letter].width = col_widths.get(letter, 20)

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for cell in row:
        cell.alignment = wrap_top

ws.freeze_panes = "A2"
ws.row_dimensions[1].height = 30
for r in range(2, ws.max_row + 1):
    ws.row_dimensions[r].height = 120

wb.save(xlsx_path)
print("Wrote", xlsx_path, "and CSV with", len(rows), "instrument rows.")
print("Policy score O =", O_score)
for row in rows:
    print(row["Q_instrument_lifecycle_stage"], "->", row["V_instrument_score"])
