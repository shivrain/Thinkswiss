"""
Builds the 4P Index (Plastic Pollution Policy Index) coding table for:

Decret n. 2001-282 du 12 avril 2001 portant application du Code de l'Environnement (Senegal)
-- the implementing decree for the (now-repealed) Loi n. 2001-01 du 15 janvier 2001 --

Source verified: user-uploaded PDF (clean, machine-readable text; signed "Fait a Dakar le 12
avril 2001" by President Abdoulaye Wade and PM Mame Madior Boye), matching the FAOLEX mirror at
https://faolex.fao.org/docs/pdf/sen37192.pdf supplied in the task brief.

IMPORTANT DATE FINDING: the "12 avril 2001" date that appeared in earlier task briefs attached to
"the Environmental Code" in fact belongs to THIS implementing decree, not to the 2001 Law itself
(Loi n. 2001-01, which is dated 15 January 2001). This resolves the date-discrepancy flagged in the
earlier 2001-Law coding.

Produces:
  - senegal_decret_2001-282_4p_index.csv
  - senegal_decret_2001-282_4p_index.xlsx (formatted, wrapped, one sheet)
"""
import csv
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [
    "A_policy_name",
    "B_policy_url",
    "C_policy_year",
    "D_policy_objective",
    "E_policy_target",
    "F_policy_target_text",
    "G_policy_type",
    "H_policy_type_justification",
    "I_policy_integration",
    "J_policy_sectors_list",
    "K_policy_circularity",
    "L_policy_lifecycle_phases_list",
    "M_policy_budget",
    "N_policy_budget_text",
    "O_policy_score",
    "P_instrument_type",
    "Q_instrument_lifecycle_stage",
    "R_instrument_description",
    "S_instrument_in_force",
    "T_instrument_implementation",
    "U_instrument_implementation_text",
    "V_instrument_score",
    "W_comments",
]

# ---------------------------------------------------------------------------
# Section A - Policy-level fields (identical across every row)
# ---------------------------------------------------------------------------

POLICY_NAME = "Décret n° 2001-282 du 12 avril 2001 portant application du Code de l'Environnement"
POLICY_URL = "https://faolex.fao.org/docs/pdf/sen37192.pdf"
POLICY_YEAR = 2001
POLICY_OBJECTIVE = (
    "Implementing/operational decree for the classified-installations (ICPE), environmental impact "
    "assessment (EIA), water pollution, water policing, air pollution and noise pollution titles of the "
    "(now-repealed) Loi n° 2001-01 du 15 janvier 2001 portant Code de l'Environnement. It never names "
    "'plastics' or 'synthetic polymer materials', and it sets no plastics-specific target -- it is a "
    "purely procedural/technical decree (permit application requirements, public-inquiry procedures, EIA "
    "content/categories, effluent and emission standards, sampling/monitoring protocols, tax computation, "
    "and enforcement/seizure procedures) rather than a plastics policy. It is nonetheless highly relevant "
    "to plastics because it operationalises the day-to-day regulation of the industrial installations, "
    "waste-water discharges, waste incineration/combustion equipment and marine/coastal pollution "
    "enforcement that plastic production, processing and end-of-life management in Senegal are subject to. "
    "Loi n° 2001-01 was repealed and replaced by Loi n° 2023-15 du 2 août 2023, which does not itself "
    "repeal this decree; no replacement implementing decree for the 2023 Code was identified at the time "
    "of this coding, so this decree is treated as remaining in force only to the extent its provisions do "
    "not conflict with the 2023 Code (ordinary principle of transitional legal continuity for "
    "sub-legislative texts whose parent Act has been replaced), consistent with the task brief's framing."
)
POLICY_TARGET = 0
POLICY_TARGET_TEXT = ""
POLICY_TYPE = 0.75
POLICY_TYPE_JUSTIFICATION = (
    "Executive/presidential decree, not an Act of Parliament: 'LE PRESIDENT DE LA REPUBLIQUE...Le Conseil "
    "d'Etat entendu en sa séance du 13 octobre 2000...Sur le rapport du Ministre de l'environnement... "
    "DECRETE'; signed 'Fait à Dakar le 12 avril 2001' by the President and Prime Minister and countersigned "
    "by name by ten sectoral Ministers (Art. R86). It implements/operationalises the parent legislation "
    "(Loi n° 2001-01) rather than being enacted by the legislature itself -> sub-legislative regulation."
)
POLICY_INTEGRATION = 1
POLICY_SECTORS_LIST = (
    "environment, industry, agriculture, fisheries, health, mining, maritime transport, "
    "urban planning & construction, tourism, water resources, local government/municipalities, "
    "commerce/trade, energy, transport"
)
POLICY_CIRCULARITY = 0.75
POLICY_LIFECYCLE_PHASES_LIST = "production, recycling, disposal, environmental leakage"
POLICY_BUDGET = 0.5
POLICY_BUDGET_TEXT = (
    "Art. R26: taxes and fines for classified installations 'doivent être acquittés dans un délai de "
    "quarante cinq (45) jours après l'émission du bulletin de liquidation.' Art. R32: 'La taxe superficiaire "
    "est due par toute installation classée.' Art. R54: 'le degré de pollution sur la base duquel la taxe à "
    "payer par l'exploitant est fixé.' The decree operationalises collection of these parafiscal taxes "
    "(which, under the parent Code, fund the 'fonds pour la protection de l'Environnement'), but this "
    "decree's own text does not itself restate a ring-fencing/dedicated-fund provision -- coded 0.5 rather "
    "than 1 because the ring-fencing declaration lives in the parent law (Loi n° 2001-01, Art. L25/L27; "
    "carried forward in Loi n° 2023-15, Art. 16-19), not in this implementing text."
)
POLICY_SCORE_FORMULA = "[auto] O = (G + I + K + M) / 4"

policy_fields = dict(
    A_policy_name=POLICY_NAME,
    B_policy_url=POLICY_URL,
    C_policy_year=POLICY_YEAR,
    D_policy_objective=POLICY_OBJECTIVE,
    E_policy_target=POLICY_TARGET,
    F_policy_target_text=POLICY_TARGET_TEXT,
    G_policy_type=POLICY_TYPE,
    H_policy_type_justification=POLICY_TYPE_JUSTIFICATION,
    I_policy_integration=POLICY_INTEGRATION,
    J_policy_sectors_list=POLICY_SECTORS_LIST,
    K_policy_circularity=POLICY_CIRCULARITY,
    L_policy_lifecycle_phases_list=POLICY_LIFECYCLE_PHASES_LIST,
    M_policy_budget=POLICY_BUDGET,
    N_policy_budget_text=POLICY_BUDGET_TEXT,
    O_policy_score=POLICY_SCORE_FORMULA,
)

# ---------------------------------------------------------------------------
# Section B - Instrument-level fields (one dict per row)
# ---------------------------------------------------------------------------

instruments = [
    dict(
        P_instrument_type=1.0,
        Q_instrument_lifecycle_stage="Production",
        R_instrument_description=(
            "Art. R1-R37 (Titre I): the operational authorisation/declaration procedure for classified "
            "installations (ICPE), incl. plastics manufacturing/processing/recycling plants. Class-1 "
            "installations require a 5-copy application, a 15-day public inquiry led by the Governor, "
            "review by the 'Comité Régional de Développement' and municipal/rural council opinion, with a "
            "decision by the Minister of Environment (Art. R5-R8); Class-2 installations require a "
            "declaration and receive prescriptive conditions (Art. R16-R18); authorisations lapse after 2 "
            "years' non-use (Art. R15, R19) and existing/unclassified installations can be placed under "
            "notice to remedy dangers or face suspension (Art. R36)."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: 'Ministre chargé de l'environnement' decides on authorisations/declarations (Art. "
            "R3-R4, R17) (+0.25). Enforcement: 'Le fonctionnement de toute installation en infraction "
            "entraîne, après mise en demeure non suivi d'effet...l'application des sanctions pénales prévues "
            "au Chapitre I Titre IV de la loi portant Code de l'Environnement' (Art. R24); provisional "
            "closure power (Art. R23) (+0.25). Monitoring: 'L'inspection des installations classées est "
            "exercée sous l'autorité du Ministre...' by sworn, accredited agents with a right of access and "
            "no right of refusal by operators (Art. R20-R22) (+0.25). Unconditional: not awarded -- Art. R31 "
            "grandfathers installations that pre-date classification requirements, allowing them to keep "
            "operating without authorisation/declaration under conditions."
        ),
        W_comments=(
            "This is the decree-level operationalisation of the ICPE permitting regime already coded at the "
            "policy (law) level for both Loi n° 2001-01 and Loi n° 2023-15; it is coded here as a distinct "
            "instrument because it is this decree's own text (not the parent law's) that fixes the concrete "
            "procedural/technical requirements."
        ),
    ),
    dict(
        P_instrument_type=0.40,
        Q_instrument_lifecycle_stage="Production",
        R_instrument_description=(
            "Art. R38-R44 (Titre II): the operational environmental impact assessment (EIA/EIE) procedure -- "
            "project categorisation (Category 1 = full EIA; Category 2 = 'analyse environnementale "
            "initiale', Art. R40), mandatory EIA content (Art. R39), a 5-year renewable accreditation regime "
            "for EIA consulting firms ('bureaux d'étude agréés', Art. R42), and a monthly-meeting 'comité "
            "technique' that administers and validates the process (Art. R43-R44)."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.50,
        U_instrument_implementation_text=(
            "Authority: the 'comité technique', administered via the 'Direction de l'environnement et des "
            "établissements classés', and the Minister of Environment who signs the final decision (Art. "
            "R43-R44) (+0.25). Monitoring: accreditation of EIA firms can be withdrawn 'lorsque la qualité "
            "de trois études au maximum a été jugée médiocre' -- an explicit quality-control mechanism (Art. "
            "R42) (+0.25). Enforcement: no specific fine/penalty for bypassing the EIA procedure is stated "
            "within this Titre itself -- not awarded. Unconditional: not awarded -- EIA requirements are "
            "tiered by project category (Art. R40), not universal."
        ),
        W_comments=(
            "Mirrors the EIA instrument coded at the policy level for the parent laws; here the emphasis is "
            "on the decree's distinctive institutional machinery (comité technique, accredited consultants) "
            "rather than the underlying legal obligation to conduct an EIA."
        ),
    ),
    dict(
        P_instrument_type=1.0,
        Q_instrument_lifecycle_stage="Environmental leakage",
        R_instrument_description=(
            "Art. R45-R53 (Titre III, Ch. I-II): technical standard that 'L'effluent rejeté ne doit en aucun "
            "cas entraîner la détérioration du milieu récepteur' (Art. R49); an EIA is required before any "
            "operator may discharge effluent into a natural receiving environment (Art. R50); discharge "
            "authorisation is conditioned on EIA results and compliance with physical/chemical/biological/"
            "bacteriological standards (Art. R51); operators must install standardised sampling/flow-"
            "measurement devices before any discharge (Art. R53)."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.50,
        U_instrument_implementation_text=(
            "Authority: discharge conditions fixed by joint order of the Ministers of Environment, "
            "Hydraulics, Health, Agriculture or the Sea (Art. R49) (+0.25). Monitoring: 'Un dispositif "
            "normalisé pour l'échantillonnage et la mesure de débit doit être installé, avant tout rejet' "
            "and controls are performed by sworn, competent agents (Art. R52-R53) (+0.25). Enforcement: no "
            "specific fine for non-compliant discharge is stated within this Titre itself (the pollution tax "
            "in Art. R54 is a separate financial instrument, coded in its own row) -- not awarded. "
            "Unconditional: not awarded -- the EIA/authorisation requirement in Art. R50-R51 applies "
            "specifically to use of natural receiving environments, not e.g. discharges into engineered "
            "public sewer networks (Art. R48)."
        ),
        W_comments=(
            "Industrial wastewater from plastics manufacturing/washing/recycling operations would fall "
            "within this discharge-standard regime as a classified-installation effluent."
        ),
    ),
    dict(
        P_instrument_type=0.60,
        Q_instrument_lifecycle_stage="Waste management",
        R_instrument_description=(
            "Art. R26, R32, R54: parafiscal financial instruments operationalising the parent Code's "
            "pollution taxes -- a flat 'taxe superficiaire' due by every classified installation regardless "
            "of land tenure (Art. R32); a pollution charge ('taxe à la pollution') computed from the "
            "measured physical/chemical/bacteriological/biological characteristics of controlled effluents, "
            "which determine 'le degré de pollution sur la base duquel la taxe...est fixée' (Art. R54); and "
            "a strict 45-day payment deadline for both taxes and pecuniary penalties following an infraction "
            "(Art. R26)."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=1.0,
        U_instrument_implementation_text=(
            "Authority: taxes/fines are administered by the 'Ministère chargé de l'environnement' via a "
            "'bulletin de liquidation' (Art. R26) (+0.25). Enforcement: 'les droits et taxes...doivent être "
            "acquittés dans un délai de quarante cinq (45) jours' and 'les pénalités pécuniaires...doivent "
            "être acquittées dans un délai de quarante cinq (45) jours' -- explicit, binding payment "
            "deadlines tied to the infraction/penalty regime (Art. R26) (+0.25). Monitoring: the pollution "
            "charge is directly and explicitly computed from controlled/analysed effluent data (Art. R54) "
            "(+0.25). Unconditional: the superficiary tax applies to every classified installation with no "
            "exemption identified in the text (Art. R32) (+0.25)."
        ),
        W_comments=(
            "The strongest-implemented instrument in this decree on the evidence available; feeds the "
            "ring-fenced Environmental Protection Fund created at the parent-law level (see Column N)."
        ),
    ),
    dict(
        P_instrument_type=1.0,
        Q_instrument_lifecycle_stage="Environmental leakage",
        R_instrument_description=(
            "Art. R56-R70 (Titre IV, 'Police de l'eau'): absolute prohibitions -- '(a) tous déversements, "
            "écoulements, dépôts directs ou indirects...susceptible de polluer les eaux continentales ou "
            "marines; (b) tous rejets à partir de la côte d'eaux et de toutes substances usées, de déchets "
            "industriels, de toutes substances solides ou liquides toxiques pouvant entraîner la pollution "
            "des plages et des zones littorales' (Art. R56) -- together with a comprehensive multi-agency "
            "enforcement apparatus: named categories of sworn agents for maritime and inland waters (Art. "
            "R59-R60), onboard/at-port inspection powers including sampling of effluents and tanks and "
            "checks of hydrocarbon record books and pollution-prevention certificates (Art. R61), fixed "
            "45-day-style transaction/fine procedures (Art. R66-R67), and vessel detention pending payment "
            "of a bond/security for suspected pollution offences (Art. R69)."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: an extensive, explicitly enumerated list of competent agents by context (Marine "
            "Nationale, Marine marchande, Océanographie et Pêches maritimes, Mines, port authorities, "
            "Ministère de l'environnement) (Art. R59-R60) (+0.25). Enforcement: transaction-fine procedure "
            "with official reports ('procès-verbaux') transmitted to the Procureur de la République, and "
            "vessel detention/bond mechanisms (Art. R66, R69) (+0.25). Monitoring: laboratory-accredited "
            "sampling of effluents/tanks and inspection of onboard pollution-prevention documentation (Art. "
            "R61) (+0.25). Unconditional: not awarded -- Art. R57 explicitly allows authorised exceptions "
            "for vessel discharges 'dans des cas limitativement prévus par arrêté conjoint...conformes aux "
            "Conventions internationales'."
        ),
        W_comments=(
            "Art. R56(b) is the single most directly plastics/marine-litter-relevant provision in this "
            "decree: it explicitly bans coastal discharge of 'déchets industriels' and 'toutes substances "
            "solides...toxiques' causing beach/coastal pollution -- precisely the mechanism by which "
            "packaging/plastic waste dumped near the coast becomes marine litter."
        ),
    ),
    dict(
        P_instrument_type=1.0,
        Q_instrument_lifecycle_stage="Production",
        R_instrument_description=(
            "Art. R71-R74, R80-R82 (Titre V, Ch. I & III): fixed-installation operators must act to "
            "'supprimer ou réduire leurs émissions polluantes' when weather conditions create a pollution "
            "spike threatening people or property (Art. R72); interministerial orders can impose temporary "
            "(up to 48h, renewable in 24h increments) restrictions such as banning use of certain chemical "
            "products or slowing/stopping equipment (Art. R73); special air-quality protection zones can be "
            "created and delimited based on measured particulate/gas concentrations and local conditions "
            "(Art. R80-R81); non-compliance with either mechanism is a punishable contravention (Art. R82)."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: interministerial orders (Environnement, Intérieur, Industrie, Santé, Urbanisme, "
            "Agriculture) (Art. R73-R74, R80) (+0.25). Enforcement: 'Sont punies des peines prévues pour les "
            "contraventions' for non-compliance within special zones or with Chapter I prescriptions (Art. "
            "R82) (+0.25). Monitoring: the triggering and zoning criteria are explicitly measurement-based -- "
            "'conditions météorologiques constatées' (Art. R72) and 'concentration pondérale et qualitative "
            "des particules dans l'air...concentration...de tout gaz toxique' (Art. R80) (+0.25). "
            "Unconditional: not awarded -- the emergency-restriction duty (Art. R72) is explicitly "
            "conditional on specific meteorological/pollution-spike circumstances, and special zones must "
            "first be created by order (Art. R80, itself an enabling power)."
        ),
        W_comments=(
            "Relevant to plastics production facilities as fixed installations whose gaseous/particulate "
            "emissions are subject to this emergency-restriction and zoning regime."
        ),
    ),
    dict(
        P_instrument_type=1.0,
        Q_instrument_lifecycle_stage="End of life",
        R_instrument_description=(
            "Art. R75-R79 (Titre V, Ch. II): technical-standards and inspection regime specific to fixed "
            "'installations d'incinération, de combustion ou de chauffage' -- i.e. waste-incineration and "
            "combustion equipment. Joint ministerial orders may set technical specifications that such "
            "equipment must meet to be manufactured, imported or sold in Senegal, with homologation/"
            "conformity-control procedures and phase-in deadlines of up to two years (Art. R76); further "
            "orders may set emission-limit and operating-log requirements (Art. R77); all such installations "
            "'sont soumises à une visite périodique par un expert ou un organisme agréé' (Art. R78); sworn "
            "agents have a right of access to the equipment, its annexes and fuel stocks for sampling (Art. "
            "R79)."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: joint orders of the Ministers of Environment, Industry, Health, Interior, Commerce "
            "and Agriculture (Art. R76-R77) (+0.25). Enforcement: Art. R82 extends contravention penalties "
            "to 'l'inobservation des prescriptions édictées en application des dispositions du chapitre II "
            "du présent titre' -- i.e. this incineration/combustion chapter (+0.25). Monitoring: mandatory "
            "periodic inspection by an accredited expert/body, with defined frequency and accreditation "
            "conditions fixed by interministerial order (Art. R78) (+0.25). Unconditional: not awarded -- "
            "the specific technical specifications themselves depend on further ministerial orders yet to "
            "be issued ('peuvent fixer', Art. R76), an enabling-power layer within an otherwise mandatory "
            "inspection duty."
        ),
        W_comments=(
            "Directly relevant to plastic waste management: open/uncontrolled incineration of plastic waste "
            "is a major pollution source, and this is the decree's mechanism for setting technical/emission "
            "standards and mandating periodic inspection of the incineration/combustion equipment used for "
            "waste disposal. In-force test (Rule 12): the periodic-inspection duty itself (Art. R78, "
            "'sont soumises') is mandatory and scored S=1; the underlying technical specifications (Art. "
            "R76, 'peuvent fixer') remain an unexercised enabling power layer, noted here rather than split "
            "into a separate row given the shared statutory basis and inspection backstop."
        ),
    ),
]

# ---------------------------------------------------------------------------
# Assemble rows, computing auto-fields O and V
# ---------------------------------------------------------------------------

O_score = round(
    (policy_fields["G_policy_type"] + policy_fields["I_policy_integration"]
     + policy_fields["K_policy_circularity"] + policy_fields["M_policy_budget"]) / 4,
    3,
)
policy_fields["O_policy_score"] = f"{O_score} [auto = (G+I+K+M)/4]"

rows = []
for inst in instruments:
    row = dict(policy_fields)
    row.update(inst)
    s = row["S_instrument_in_force"]
    p = row["P_instrument_type"]
    t = row["T_instrument_implementation"]
    v = round(s * (p + t) / 2, 3)
    row["V_instrument_score"] = f"{v} [auto = S*(P+T)/2]"
    rows.append(row)

with open("4p_index/senegal_decret_2001-282_4p_index.csv", "w", newline="", encoding="utf-8") as f:
    writer = csv.DictWriter(f, fieldnames=COLUMNS)
    writer.writeheader()
    for row in rows:
        writer.writerow(row)

df = pd.DataFrame(rows, columns=COLUMNS)
xlsx_path = "4p_index/senegal_decret_2001-282_4p_index.xlsx"
df.to_excel(xlsx_path, sheet_name="4P_Index_Decret2001-282", index=False)

wb = load_workbook(xlsx_path)
ws = wb["4P_Index_Decret2001-282"]

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF")
wrap_top = Alignment(wrap_text=True, vertical="top")

col_widths = {
    "A": 34, "B": 30, "C": 10, "D": 46, "E": 10, "F": 30, "G": 10, "H": 34,
    "I": 10, "J": 30, "K": 10, "L": 26, "M": 10, "N": 30, "O": 16,
    "P": 12, "Q": 16, "R": 46, "S": 10, "T": 12, "U": 40, "V": 16, "W": 40,
}
for letter in [get_column_letter(c) for c in range(1, len(COLUMNS) + 1)]:
    ws.column_dimensions[letter].width = col_widths.get(letter, 20)

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for cell in row:
        cell.alignment = wrap_top

ws.freeze_panes = "A2"
ws.row_dimensions[1].height = 30
for r in range(2, ws.max_row + 1):
    ws.row_dimensions[r].height = 120

wb.save(xlsx_path)
print("Wrote", xlsx_path, "and CSV with", len(rows), "instrument rows.")
print("Policy score O =", O_score)
for row in rows:
    print(row["Q_instrument_lifecycle_stage"], "->", row["V_instrument_score"])
