"""
Builds the 4P Index (Plastic Pollution Policy Index) coding table for:

Arrete ministeriel n. 027932 du 11 decembre 2020 portant creation et fixant les regles
d'organisation et de fonctionnement du Comite de pilotage du projet de Promotion de la Gestion
integree et de l'Economie des Dechets solides au Senegal (PROMOGED).

Source: user-uploaded PDF -- Journal Officiel de la Republique du Senegal, 166e annee, n. 7392,
samedi 16 janvier 2021 (pp. 53-55). This is the specific, official Senegalese legal instrument
available for PROMOGED (a World Bank-supported national solid-waste programme): a ministerial order
creating and organising the project's Comite de pilotage (Steering Committee) and Comite technique
(Technical Committee). The task brief's other URL (www.urbanisme.gouv.sn/actualites/...) redirects to
a general news page rather than the legal text itself, and the "corrected" link supplied
(file:///C:/Users/alish/Downloads/JO-7392-du-16-janvier-2021.pdf) is a local file path on the user's
own computer, not a public web URL -- the uploaded PDF is therefore the authoritative source used
here.

IMPORTANT SCOPE NOTE: this arrete is purely institutional/governance in content (committee creation,
composition, missions, meeting rules) -- it does not itself contain substantive waste-management
standards, targets, bans or budget figures (those would appear in PROMOGED's World Bank Project
Appraisal Document / Financing Agreement, which is not a Senegalese domestic legal instrument and was
not provided). Coding here is therefore limited to what this specific official document actually
contains.

Produces:
  - senegal_promoged_4p_index.csv
  - senegal_promoged_4p_index.xlsx (formatted, wrapped, one sheet)
"""
import csv
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [
    "A_policy_name",
    "B_policy_url",
    "C_policy_year",
    "D_policy_objective",
    "E_policy_target",
    "F_policy_target_text",
    "G_policy_type",
    "H_policy_type_justification",
    "I_policy_integration",
    "J_policy_sectors_list",
    "K_policy_circularity",
    "L_policy_lifecycle_phases_list",
    "M_policy_budget",
    "N_policy_budget_text",
    "O_policy_score",
    "P_instrument_type",
    "Q_instrument_lifecycle_stage",
    "R_instrument_description",
    "S_instrument_in_force",
    "T_instrument_implementation",
    "U_instrument_implementation_text",
    "V_instrument_score",
    "W_comments",
]

# ---------------------------------------------------------------------------
# Section A - Policy-level fields (identical across every row)
# ---------------------------------------------------------------------------

POLICY_NAME = (
    "Arrêté ministériel n° 027932 du 11 décembre 2020 portant création et fixant les règles "
    "d'organisation et de fonctionnement du Comité de pilotage du projet de Promotion de la Gestion "
    "intégrée et de l'Économie des Déchets solides au Sénégal (PROMOGED)"
)
POLICY_URL = (
    "Journal Officiel de la République du Sénégal, 166e année, n° 7392, samedi 16 janvier 2021, "
    "pp. 53-55 (source of the user-uploaded PDF) ; www.urbanisme.gouv.sn/actualites/r%C3%A9union-du-"
    "comit%C3%A9-de-pil... (as given in task brief -- redirects to a general news listing page, not "
    "the legal text itself) ; NB: the 'corrected' link supplied in the task brief "
    "('file:///C:/Users/alish/Downloads/JO-7392-du-16-janvier-2021.pdf') is a local file path on the "
    "requester's own computer, not a publicly accessible web URL -- not usable as a citation for "
    "others, though it corresponds to the same Journal Officiel issue as the uploaded PDF used here"
)
POLICY_YEAR = 2020
POLICY_OBJECTIVE = (
    "Establishes the institutional governance architecture -- a Comité de pilotage (Steering "
    "Committee) and, within it, a Comité technique (Technical Committee) with four thematic working "
    "groups -- for PROMOGED, a World Bank-supported national programme to strengthen integrated "
    "solid-waste governance, modernise municipal collection/treatment infrastructure and develop "
    "waste-valorisation value chains across Senegal, per the task brief's own description. This "
    "specific legal instrument (a ministerial order) is exclusively institutional/procedural in "
    "content -- it creates committees, defines their composition and missions, and sets meeting "
    "rules -- and contains no plastics-specific provision, substantive waste standard, target or "
    "budget figure of its own (those would appear in PROMOGED's underlying World Bank Project "
    "Appraisal Document / Financing Agreement, which is not a Senegalese domestic legal act and was "
    "not available for this coding). It nonetheless qualifies under Plastics Relevance Filter "
    "criterion (c): PROMOGED's household solid-waste management and valorisation mandate (Art. 2, "
    "5) is a paradigmatic governance instrument that commonly applies to plastic waste streams in "
    "practice, since plastic packaging and bags constitute a large share of Senegalese municipal "
    "solid waste, and 'l'Économie des Déchets' (the waste economy) in the project's own name directly "
    "implicates plastic recycling/valorisation value chains."
)
POLICY_TARGET = 0
POLICY_TARGET_TEXT = ""
POLICY_TYPE = 0.75
POLICY_TYPE_JUSTIFICATION = (
    "A ministerial order ('Arrêté ministériel n° 027932'), issued by the Minister of Urban Planning, "
    "Housing and Public Hygiene under executive authority, published in the Journal Officiel -- not "
    "an Act adopted by the National Assembly -> sub-legislative regulation, coded 0.75 per Rule 6."
)
POLICY_INTEGRATION = 1
POLICY_SECTORS_LIST = (
    "urban planning & housing, sanitation/public hygiene, finance, economy & international "
    "cooperation, health, local government/municipalities, environment, social protection, "
    "interior/territorial administration, civil society, waste management"
)
POLICY_CIRCULARITY = 0.50
POLICY_LIFECYCLE_PHASES_LIST = "recycling, disposal"
POLICY_BUDGET = 0.5
POLICY_BUDGET_TEXT = (
    "Art. 2: the Steering Committee's missions include 'd'examiner et d'approuver le Plan de travail "
    "budgétisé et annuel (PTBA)...; d'examiner les rapports de vérification sur l'atteinte des "
    "Indicateurs liés aux décaissements (ILDs) ; d'examiner l'allocation des fonds entre les "
    "différentes parties concernées par le Mécanisme de Financement basé sur les Résultats (REP) ; "
    "d'examiner les rapports d'évaluation de performance, d'audits des dépenses éligibles...; "
    "d'examiner les rapports d'audit annuels des comptes du projet.' The arrêté confirms and "
    "oversees the existence of a results-based-financing budget/funding mechanism for PROMOGED, but "
    "does not itself create or ring-fence a fund, nor state a specific budget figure -- both are "
    "presumably fixed in a separate financing agreement not published as part of this domestic legal "
    "instrument. Coded 0.5 (funding source mentioned/overseen, not ring-fenced by this document "
    "itself)."
)
POLICY_SCORE_FORMULA = "[auto] O = (G + I + K + M) / 4"

policy_fields = dict(
    A_policy_name=POLICY_NAME,
    B_policy_url=POLICY_URL,
    C_policy_year=POLICY_YEAR,
    D_policy_objective=POLICY_OBJECTIVE,
    E_policy_target=POLICY_TARGET,
    F_policy_target_text=POLICY_TARGET_TEXT,
    G_policy_type=POLICY_TYPE,
    H_policy_type_justification=POLICY_TYPE_JUSTIFICATION,
    I_policy_integration=POLICY_INTEGRATION,
    J_policy_sectors_list=POLICY_SECTORS_LIST,
    K_policy_circularity=POLICY_CIRCULARITY,
    L_policy_lifecycle_phases_list=POLICY_LIFECYCLE_PHASES_LIST,
    M_policy_budget=POLICY_BUDGET,
    N_policy_budget_text=POLICY_BUDGET_TEXT,
    O_policy_score=POLICY_SCORE_FORMULA,
)

# ---------------------------------------------------------------------------
# Section B - Instrument-level fields (one dict per row)
# ---------------------------------------------------------------------------

instruments = [
    dict(
        P_instrument_type=0.40,
        Q_instrument_lifecycle_stage="Waste management",
        R_instrument_description=(
            "Art. 1-3: creates the PROMOGED 'Comité de pilotage' (Steering Committee) within the "
            "Ministry of Urban Planning, Housing and Public Hygiene, chaired by that Minister (or "
            "representative), with members drawn from the Ministries of Finance, Economy, Health, "
            "Local Government/Territorial Collectivities, Environment, Public Hygiene, plus the "
            "Association des Maires du Sénégal and civil society (CONGAD). Its missions (Art. 2) "
            "include steering/supervising project implementation, coordinating and validating the "
            "Technical Committee's work, approving the annual budgeted work plan (PTBA), reviewing "
            "disbursement-linked-indicator (ILD) verification reports, reviewing results-based-"
            "financing (REP) fund allocation, and reviewing performance/audit reports. It meets twice "
            "yearly (extraordinary sessions as needed) and is dissolved automatically at project "
            "closure."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: chaired by 'le Ministre de l'Urbanisme, du Logement et de l'Hygiène publique "
            "ou son représentant', with a defined multi-ministerial and civil-society membership (Art. "
            "3) (+0.25). Monitoring: an extensive, explicit set of financial/performance-oversight "
            "functions -- PTBA approval, ILD verification review, REP fund-allocation review, "
            "performance/expenditure-audit review, and annual project-accounts audit review (Art. 2) "
            "(+0.25). Unconditional: the Committee's mandate applies generally with no exemption "
            "identified (+0.25). Enforcement: not awarded -- no penalty or sanction is attached to "
            "non-compliance with the Committee's decisions within this arrêté."
        ),
        W_comments=(
            "Coded per Rule 11 as a multi-level, multi-sectoral governance-coordination instrument for "
            "national solid-waste management -- spanning six ministries, the national mayors' "
            "association and civil society. Plastics relevance is indirect (criterion (c)): this body "
            "oversees the national waste-management/valorisation programme within which plastic waste "
            "streams are a major, if unnamed, component."
        ),
    ),
    dict(
        P_instrument_type=0.40,
        Q_instrument_lifecycle_stage="Waste management",
        R_instrument_description=(
            "Art. 4-8: creates, within the Steering Committee, a 'Comité technique' (Technical "
            "Committee) to assist in implementing the Committee's missions, chaired by the "
            "Coordinator of the Waste Management Coordination Unit (UCG), with a detailed multi-"
            "ministerial membership (Finance/Budget, Interior, Economy/Plan/Cooperation, Health, "
            "Family/Gender/Child Protection, Local Government/Territorial Development, Environment, "
            "and the lead Ministry itself) plus a mayors'-association representative; it is organised "
            "into four thematic technical groups covering institutional aspects, financing mechanisms, "
            "public-private partnership (PPP), and technical/environmental/social aspects (Art. 7), "
            "and reports on strategic studies related to the household solid-waste sector and on its "
            "own actions to the Steering Committee."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: chaired by 'le Coordonnateur de l'Unité de Coordination de la Gestion des "
            "Déchets', with an extensively named membership by ministry and directorate (Art. 6) "
            "(+0.25). Monitoring: tasked to 'examiner les rapports relatifs aux études stratégiques en "
            "lien avec le développement du secteur des déchets solides ménagers' and to 'rendre compte "
            "au Comité de pilotage des actions menées' (Art. 5) (+0.25). Unconditional: the Committee's "
            "mandate and composition are stated generally with no exemption identified (+0.25). "
            "Enforcement: not awarded -- no penalty or sanction is attached to non-performance within "
            "this arrêté."
        ),
        W_comments=(
            "A second-tier, technically-focused coordination instrument (Rule 11) supporting the "
            "Steering Committee above, with dedicated PPP and financing-mechanism working groups that "
            "would be directly relevant to structuring plastic-waste valorisation/recycling "
            "partnerships, even though 'plastic' is not named in the text."
        ),
    ),
]

# ---------------------------------------------------------------------------
# Assemble rows, computing auto-fields O and V
# ---------------------------------------------------------------------------

O_score = round(
    (policy_fields["G_policy_type"] + policy_fields["I_policy_integration"]
     + policy_fields["K_policy_circularity"] + policy_fields["M_policy_budget"]) / 4,
    3,
)
policy_fields["O_policy_score"] = f"{O_score} [auto = (G+I+K+M)/4]"

rows = []
for inst in instruments:
    row = dict(policy_fields)
    row.update(inst)
    s = row["S_instrument_in_force"]
    p = row["P_instrument_type"]
    t = row["T_instrument_implementation"]
    v = round(s * (p + t) / 2, 3)
    row["V_instrument_score"] = f"{v} [auto = S*(P+T)/2]"
    rows.append(row)

with open("4p_index/senegal_promoged_4p_index.csv", "w", newline="", encoding="utf-8") as f:
    writer = csv.DictWriter(f, fieldnames=COLUMNS)
    writer.writeheader()
    for row in rows:
        writer.writerow(row)

df = pd.DataFrame(rows, columns=COLUMNS)
xlsx_path = "4p_index/senegal_promoged_4p_index.xlsx"
df.to_excel(xlsx_path, sheet_name="4P_Index_PROMOGED", index=False)

wb = load_workbook(xlsx_path)
ws = wb["4P_Index_PROMOGED"]

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF")
wrap_top = Alignment(wrap_text=True, vertical="top")

col_widths = {
    "A": 34, "B": 30, "C": 10, "D": 46, "E": 10, "F": 30, "G": 10, "H": 34,
    "I": 10, "J": 30, "K": 10, "L": 26, "M": 10, "N": 30, "O": 16,
    "P": 12, "Q": 16, "R": 46, "S": 10, "T": 12, "U": 40, "V": 16, "W": 40,
}
for letter in [get_column_letter(c) for c in range(1, len(COLUMNS) + 1)]:
    ws.column_dimensions[letter].width = col_widths.get(letter, 20)

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for cell in row:
        cell.alignment = wrap_top

ws.freeze_panes = "A2"
ws.row_dimensions[1].height = 30
for r in range(2, ws.max_row + 1):
    ws.row_dimensions[r].height = 130

wb.save(xlsx_path)
print("Wrote", xlsx_path, "and CSV with", len(rows), "instrument rows.")
print("Policy score O =", O_score)
for row in rows:
    print(row["Q_instrument_lifecycle_stage"], "->", row["V_instrument_score"])
