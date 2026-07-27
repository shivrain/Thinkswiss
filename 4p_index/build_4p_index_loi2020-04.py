"""
Builds the 4P Index (Plastic Pollution Policy Index) coding table for:

Loi n. 2020-04 du 8 janvier 2020 relative a la prevention et a la reduction de l'incidence sur
l'environnement des produits plastiques (Senegal) -- the flagship plastics law, repealing and
replacing Loi n. 2015-09 du 4 mai 2015 (Art. 40).

Source: fetched directly from the internet at https://faolex.fao.org/docs/pdf/sen200382.pdf (a
clean, machine-readable mirror of the official text), cross-consistent with the task-brief URLs
https://www.au-senegal.com/IMG/pdf/loi-plastique-senegal-2020-04.pdf and
https://www.fao.org/faolex/results/details/en/c/LEX-FAOC200382. Confirmed authentic: "L'Assemblee
nationale a adopte, en sa seance du lundi 30 decembre 2019 ; Le President de la Republique promulgue
la loi..."; "Fait a Dakar, le 08 Janvier 2020. Macky SALL."; Art. 40 explicitly repeals Loi n. 2015-09.

This is the richest, most explicitly and comprehensively plastics-focused document in this coding
series: every one of its 10 chapters and 42 articles is directly about plastic products and plastic
waste (bans, deposit-refund, EPR, recycled content, import/export controls, taxation, penalties).

Produces:
  - senegal_loi_2020-04_4p_index.csv
  - senegal_loi_2020-04_4p_index.xlsx (formatted, wrapped, one sheet)
"""
import csv
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [
    "A_policy_name",
    "B_policy_url",
    "C_policy_year",
    "D_policy_objective",
    "E_policy_target",
    "F_policy_target_text",
    "G_policy_type",
    "H_policy_type_justification",
    "I_policy_integration",
    "J_policy_sectors_list",
    "K_policy_circularity",
    "L_policy_lifecycle_phases_list",
    "M_policy_budget",
    "N_policy_budget_text",
    "O_policy_score",
    "P_instrument_type",
    "Q_instrument_lifecycle_stage",
    "R_instrument_description",
    "S_instrument_in_force",
    "T_instrument_implementation",
    "U_instrument_implementation_text",
    "V_instrument_score",
    "W_comments",
]

# ---------------------------------------------------------------------------
# Section A - Policy-level fields (identical across every row)
# ---------------------------------------------------------------------------

POLICY_NAME = (
    "Loi n° 2020-04 du 8 janvier 2020 relative à la prévention et à la réduction de l'incidence sur "
    "l'environnement des produits plastiques"
)
POLICY_URL = (
    "https://faolex.fao.org/docs/pdf/sen200382.pdf ; "
    "https://www.au-senegal.com/IMG/pdf/loi-plastique-senegal-2020-04.pdf ; "
    "https://www.fao.org/faolex/results/details/en/c/LEX-FAOC200382"
)
POLICY_YEAR = 2020
POLICY_OBJECTIVE = (
    "Senegal's flagship plastics law, adopted specifically to remedy the acknowledged failure of the "
    "narrower Loi n° 2015-09 (which banned only thin plastic bags below a micron threshold) to curb "
    "the accumulation and marine dispersal of plastic waste. Per its own exposé des motifs, it shifts "
    "from a linear to a circular-economy approach and, in a single integrated package: bans single-"
    "use/disposable plastic products and all plastic checkout bags (Art. 4-5); introduces a deposit-"
    "refund scheme for plastic bottles (Art. 6-10); establishes an extended producer responsibility "
    "(EPR) regime via individual programmes or collective eco-organismes (Art. 11-14); imposes waste-"
    "reduction-at-source/eco-design, recycled-content-integration and mandatory-marking duties on "
    "producers, and a take-back duty on consumers (Art. 15-18); bans plastic-waste imports and "
    "regulates plastic-waste exports (Art. 19-20); sets a minimum floor price for recycled plastic "
    "waste and a 'taxe plastique' on non-recyclable plastic products (Art. 21-22); and backs all of "
    "the above with seizure/settlement procedures and extensive criminal penalties (Art. 23-39). It "
    "repeals and replaces Loi n° 2015-09 (Art. 40)."
)
POLICY_TARGET = 0.5
POLICY_TARGET_TEXT = (
    "Art. 16: 'Les producteurs sont tenus, lorsque c'est techniquement faisable et économiquement "
    "viable, d'intégrer une part de plastique recyclé dans les produits plastiques neufs qu'ils "
    "mettent sur le marché. Un décret détermine les objectifs nationaux en matière d'intégration de "
    "plastique recyclé dans les produits plastiques neufs mis sur le marché et fixe des délais pour "
    "réaliser ces objectifs.' A plastic-specific target is explicitly mentioned (recycled-content "
    "integration), but the law itself does not state a quantified percentage or deadline -- these are "
    "deferred to a future implementing decree -> coded 0.5 (non-quantifiable within this document) "
    "rather than 1."
)
POLICY_TYPE = 1
POLICY_TYPE_JUSTIFICATION = (
    "Full Act of Parliament: 'L'Assemblée nationale a adopté, en sa séance du lundi 30 décembre 2019 ; "
    "Le Président de la République promulgue la loi dont la teneur suit.' Signed 'Fait à Dakar, le 08 "
    "Janvier 2020' by President Macky Sall -> legislation, not an executive/ministerial instrument."
)
POLICY_INTEGRATION = 1
POLICY_SECTORS_LIST = (
    "manufacturing/production, retail, waste management, environment, health, industry, "
    "commerce/trade, finance, recycling, packaging, food & beverage, tourism, fisheries, agriculture"
)
POLICY_CIRCULARITY = 1
POLICY_LIFECYCLE_PHASES_LIST = "production, consumption, recycling, disposal, environmental leakage"
POLICY_BUDGET = 0.5
POLICY_BUDGET_TEXT = (
    "Art. 21: 'Il est établi un prix plancher auquel les entreprises du recyclage sont tenus d'acheter "
    "le kilogramme de déchet plastique. Ce prix plancher est fixé par décret.' Art. 22: 'Il est "
    "institué une taxe sur les produits constitués ou fabriqués à partir de matières plastiques non "
    "recyclables dite «taxe plastique». La taxe plastique frappe les produits figurant sur une liste "
    "établie par un décret qui en fixe le tarif et les modalités de recouvrement.' Art. 12 also "
    "provides for a compliance security deposit ('consignation') at the Caisse des Dépôts et "
    "Consignations. Coded 0.5 rather than 1: the law creates funding/pricing mechanisms (a price "
    "floor, a dedicated tax, a security-deposit mechanism), but does not itself state that tax "
    "proceeds are ring-fenced into a dedicated environmental fund for this law's own purposes -- that "
    "destination is left to the implementing decree contemplated by Art. 22 and Art. 41."
)
POLICY_SCORE_FORMULA = "[auto] O = (G + I + K + M) / 4"

policy_fields = dict(
    A_policy_name=POLICY_NAME,
    B_policy_url=POLICY_URL,
    C_policy_year=POLICY_YEAR,
    D_policy_objective=POLICY_OBJECTIVE,
    E_policy_target=POLICY_TARGET,
    F_policy_target_text=POLICY_TARGET_TEXT,
    G_policy_type=POLICY_TYPE,
    H_policy_type_justification=POLICY_TYPE_JUSTIFICATION,
    I_policy_integration=POLICY_INTEGRATION,
    J_policy_sectors_list=POLICY_SECTORS_LIST,
    K_policy_circularity=POLICY_CIRCULARITY,
    L_policy_lifecycle_phases_list=POLICY_LIFECYCLE_PHASES_LIST,
    M_policy_budget=POLICY_BUDGET,
    N_policy_budget_text=POLICY_BUDGET_TEXT,
    O_policy_score=POLICY_SCORE_FORMULA,
)

# ---------------------------------------------------------------------------
# Section B - Instrument-level fields (one dict per row)
# ---------------------------------------------------------------------------

instruments = [
    dict(
        P_instrument_type=1.0,
        Q_instrument_lifecycle_stage="Consumption",
        R_instrument_description=(
            "Art. 4: a comprehensive ban -- covering production, import, possession for sale, sale, "
            "supply to users, and use, 'sous quelque forme que ce soit' -- on single-use/disposable "
            "plastic products, exhaustively defined as: cups, glasses and cup lids; cutlery and "
            "plates; straws and beverage stirring sticks; and sachets used to package water or any "
            "other beverage for market placement."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=1.0,
        U_instrument_implementation_text=(
            "Authority: enforced by the multi-ministry control-agent framework of Art. 25 (Environment, "
            "Health, Industry, Commerce, Finance) (+0.25). Enforcement: Art. 26 (manufacture/import: "
            "1-3 years' imprisonment + 5,000,000-10,000,000 FCFA) and Art. 27 (sale/use: 1-3 months + "
            "50,000-100,000 FCFA) explicitly penalise breaches of Art. 4 (+0.25). Monitoring: sworn "
            "agents under Art. 25 and the seizure mechanism of Art. 23 apply (+0.25). Unconditional: "
            "the banned-item list is exhaustive with no exemption identified (+0.25)."
        ),
        W_comments=(
            "The law's central regulatory ban, matching the task brief's description almost verbatim "
            "('bans single-use/disposable plastics')."
        ),
    ),
    dict(
        P_instrument_type=1.0,
        Q_instrument_lifecycle_stage="Consumption",
        R_instrument_description=(
            "Art. 5: a ban on plastic checkout bags ('sacs plastiques sortie de caisse') -- with or "
            "without handles/straps, regardless of thickness. A narrow exception exists for bags used "
            "at the point of sale to package food items for protection/handling/presentation; such "
            "bags must be transparent, made from recyclable plastic materials, and their import is "
            "subject to prior authorisation by the Minister of Environment."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: same Art. 25 control-agent framework (+0.25). Enforcement: Art. 26-27 also "
            "explicitly cover checkout-bag violations (+0.25). Monitoring: same Art. 23/25 mechanisms "
            "(+0.25). Unconditional: not awarded -- Art. 5 al. 2 explicitly carves out the food-"
            "packaging bag exception."
        ),
        W_comments=(
            "Notably, the exposé des motifs explains this ban deliberately covers biodegradable, oxo-"
            "biodegradable and oxo-fragmentable bags too, rejecting them as not genuinely protective of "
            "the environment -- a stronger, more technically-informed ban than many comparable laws."
        ),
    ),
    dict(
        P_instrument_type=0.60,
        Q_instrument_lifecycle_stage="Consumption",
        R_instrument_description=(
            "Art. 6-7: a mandatory deposit-refund scheme for plastic bottles -- 'Une consigne est "
            "exigée à l'achat de tout produit contenu dans des bouteilles en plastique', with the "
            "amount fixed by decree, collected by the seller at purchase and refunded on return of the "
            "empty bottle; every seller must accept returned bottles and convey them to the nearest "
            "collection point."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: the deposit amount is 'fixé par décret', i.e. set by the competent Ministry "
            "(+0.25). Enforcement: Art. 33 penalises a seller's refusal to accept bottle returns (15 "
            "days-1 month + 50,000-100,000 FCFA) (+0.25). Monitoring: not specifically evidenced for "
            "the deposit mechanism itself (as distinct from the producer-reporting duty coded "
            "separately below) -- not awarded. Unconditional: applies to 'tout produit contenu dans des "
            "bouteilles en plastique' with no exemption identified (+0.25)."
        ),
        W_comments=(
            "Matches the type table's own 'deposit-refund schemes' example exactly; matches the task "
            "brief's description ('establishes bottle deposit-return')."
        ),
    ),
    dict(
        P_instrument_type=1.0,
        Q_instrument_lifecycle_stage="Recycling",
        R_instrument_description=(
            "Art. 8-10: producers must establish plastic-bottle collection points at their premises or "
            "other appropriate locations, and must valorise (or arrange valorisation of) collected "
            "bottles, prioritising in order re-use, recycling, then other valorisation; producers must "
            "report to the Minister of Environment every six months on quantities placed on the market "
            "vs. collected, collection-point details, the percentage gap, and corrective measures for "
            "any negative gap; the Minister may prescribe additional corrective measures if those taken "
            "are insufficient."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=1.0,
        U_instrument_implementation_text=(
            "Authority: 'le Ministre chargé de l'Environnement' receives reports and may prescribe "
            "corrective measures (Art. 9-10) (+0.25). Enforcement: Art. 34 penalises producers failing "
            "to establish sufficient collection points (3-6 months + 5,000,000-10,000,000 FCFA) "
            "(+0.25). Monitoring: Art. 9's mandatory biannual electronic sectoral report with specific "
            "quantitative metrics is itself an explicit data-collection/monitoring requirement (+0.25). "
            "Unconditional: the collection-point and valorisation duties (Art. 8) are stated without "
            "exemption (+0.25)."
        ),
        W_comments=(
            "Matches the instrument-type table's 'mandatory take-back' example precisely."
        ),
    ),
    dict(
        P_instrument_type=1.0,
        Q_instrument_lifecycle_stage="Waste management",
        R_instrument_description=(
            "Art. 11-14: an extended producer responsibility (EPR) regime -- producers placing plastic-"
            "based products on the market are responsible for managing the resulting waste, discharged "
            "either via Ministry-approved individual collection/treatment programmes (3-year renewable "
            "approval, minimum requirements fixed by order, subject to periodic sworn-agent inspections, "
            "suspension/termination for non-compliance, with an optional compliance security deposit at "
            "the Caisse des Dépôts et Consignations) or by joining collective 'éco-organismes' (approved "
            "for up to 10 years under a binding 'cahier des charges', also subject to periodic "
            "inspection and 30-day-notice withdrawal for breaches); both pathways require an annual "
            "activity report to the Minister by 30 April."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: 'le Ministre chargé de l'Environnement' approves individual programmes and eco-"
            "organismes (Art. 12-13) (+0.25). Enforcement: extensive -- suspension/termination and "
            "security-deposit mechanisms (Art. 12), approval withdrawal after a 30-day notice (Art. "
            "13), and criminal penalties (Art. 35: 1-3 years + 10,000,000-20,000,000 FCFA for the basic "
            "EPR duty; Art. 36: 3-6 months + 5,000,000-10,000,000 FCFA for individual-programme non-"
            "compliance) (+0.25). Monitoring: 'contrôles périodiques effectués par des agents "
            "assermentés' apply to both individual programmes and eco-organismes (Art. 12-13), plus the "
            "mandatory annual activity report (Art. 14) (+0.25). Unconditional: not awarded -- Art. 12 "
            "conditions individual-programme approval on the producer proving adequate technical/"
            "financial capacity, a qualifying gate."
        ),
        W_comments=(
            "Matches the instrument-type table's 'EPR obligations' example directly, and the task "
            "brief's description ('extended producer responsibility'). Mixed instrument per Rule 10: "
            "combines a Regulatory obligation (Art. 11) with Economic (security-deposit) and Procedural "
            "(institutional accreditation of eco-organismes) elements; scored at the highest applicable "
            "value."
        ),
    ),
    dict(
        P_instrument_type=1.0,
        Q_instrument_lifecycle_stage="Production",
        R_instrument_description=(
            "Art. 15: 'Les producteurs sont tenus de réduire à la source les quantités de déchets qui "
            "peuvent résulter de leurs activités et de mettre sur le marché des produits susceptibles, "
            "après être devenus des déchets, de faire l'objet d'un recyclage ou d'une valorisation dans "
            "des conditions qui respectent l'environnement.' A waste-prevention/eco-design duty on "
            "producers."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.50,
        U_instrument_implementation_text=(
            "Authority: falls under the general Art. 25 control-agent enforcement framework applicable "
            "throughout the law (+0.25). Unconditional: the duty applies to 'les producteurs' generally "
            "with no exemption identified (+0.25). Enforcement: notably, unlike almost every other "
            "substantive duty in this law, no penal article (Art. 26-39) specifically penalises a "
            "breach of Art. 15 -- not awarded. Monitoring: no dedicated tracking mechanism is specified "
            "for this duty -- not awarded."
        ),
        W_comments=(
            "A genuine implementation gap flagged per Rule 8: this eco-design/waste-prevention duty is "
            "the only substantive producer obligation in Chapters II-VI without its own dedicated "
            "criminal sanction, which may weaken its practical enforceability relative to the law's "
            "other, more heavily-sanctioned duties."
        ),
    ),
    dict(
        P_instrument_type=1.0,
        Q_instrument_lifecycle_stage="Production",
        R_instrument_description=(
            "Art. 16: producers must, 'lorsque c'est techniquement faisable et économiquement viable', "
            "incorporate a share of recycled plastic in the new plastic products they place on the "
            "market; a decree is to set the national recycled-content-integration targets and "
            "deadlines for achieving them."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.50,
        U_instrument_implementation_text=(
            "Authority: 'Un décret détermine les objectifs nationaux' -- the decree-issuing mechanism is "
            "explicitly named (+0.25). Enforcement: Art. 32 penalises failure to integrate recycled "
            "content when technically/economically feasible (1-3 months + 5,000,000-10,000,000 FCFA) "
            "(+0.25). Monitoring: not specifically evidenced -- not awarded. Unconditional: not awarded "
            "-- explicitly conditioned on 'techniquement faisable et économiquement viable'."
        ),
        W_comments=(
            "In-force test (Rule 12) uncertainty flagged per Rule 8: the general duty to integrate "
            "recycled content is stated in mandatory terms ('sont tenus'), so S=1 is applied, but the "
            "concrete, quantified national targets and deadlines are entirely deferred to a future "
            "decree not yet shown to exist -- practical enforceability of specific numeric targets "
            "likely awaits that decree even though the general duty and its penalty are already in "
            "force. Matches the task brief's 'recycled-content targets' description."
        ),
    ),
    dict(
        P_instrument_type=1.0,
        Q_instrument_lifecycle_stage="Production",
        R_instrument_description=(
            "Art. 17: 'Les produits constitués ou fabriqués à partir de matières plastiques mis sur le "
            "marché portent un marquage visible, nettement lisible et indélébile...indiquant "
            "l'identité ou la raison sociale et l'adresse du producteur.' A mandatory traceability-"
            "labelling standard for all plastic products placed on the market."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: falls under the general Art. 25 enforcement framework (+0.25). Enforcement: "
            "Art. 30 explicitly penalises non-compliance (3-6 months + 2,000,000-5,000,000 FCFA) "
            "(+0.25). Monitoring: not specifically evidenced -- not awarded. Unconditional: applies to "
            "all plastic-based products placed on the market with no exemption identified (+0.25)."
        ),
        W_comments=(
            "A traceability/accountability instrument enabling enforcement of the EPR regime by "
            "identifying the responsible producer for any given plastic product found as waste."
        ),
    ),
    dict(
        P_instrument_type=1.0,
        Q_instrument_lifecycle_stage="Disposal",
        R_instrument_description=(
            "Art. 18: 'Les consommateurs et les utilisateurs finaux de produits constitués ou "
            "fabriqués à partir de matières plastiques sont tenus, lorsque ces produits deviennent des "
            "déchets, de les acheminer vers les points de collectes aménagés à cet effet.' A mandatory "
            "duty on consumers/end users to deposit plastic waste at designated collection points."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: falls under the general Art. 25 enforcement framework (+0.25). Enforcement: "
            "Art. 37 explicitly penalises abandoning plastic waste elsewhere than designated collection "
            "points (15 days-1 month + 20,000-50,000 FCFA) (+0.25). Monitoring: not specifically "
            "evidenced -- not awarded. Unconditional: applies to all consumers/end users with no "
            "exemption identified (+0.25)."
        ),
        W_comments=(
            "The consumer-facing counterpart to the producer collection-point obligations (Art. 8) -- "
            "together they are meant to close the collection loop for plastic waste."
        ),
    ),
    dict(
        P_instrument_type=1.0,
        Q_instrument_lifecycle_stage="Waste management",
        R_instrument_description=(
            "Art. 19: 'L'importation de déchets plastiques sur le territoire national est interdite.' "
            "In case of violation, the imported plastic waste is seized and re-exported to the country "
            "of origin/provenance at the importer's expense, without prejudice to criminal prosecution."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: falls under the general Art. 25 multi-ministry control framework, applicable at "
            "borders/customs (+0.25). Enforcement: the article's own seizure/re-export mechanism plus "
            "Art. 28's severe criminal penalty (3-5 years + 50,000,000-100,000,000 FCFA) (+0.25). "
            "Monitoring: not specifically evidenced beyond the general control-agent framework -- not "
            "awarded. Unconditional: the ban is absolute with no exemption identified (+0.25)."
        ),
        W_comments=(
            "Matches the task brief's 'plastic-waste import/export controls' description; the severity "
            "of the Art. 28 penalty (up to 100 million FCFA) signals this is treated as a serious "
            "environmental-trade offence."
        ),
    ),
    dict(
        P_instrument_type=1.0,
        Q_instrument_lifecycle_stage="Waste management",
        R_instrument_description=(
            "Art. 20: plastic waste produced in Senegal may only be exported after prior authorisation "
            "by the Minister of Environment, and only to countries that permit its import and possess "
            "adequate treatment facilities."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.50,
        U_instrument_implementation_text=(
            "Authority: 'le Ministre chargé de l'Environnement' explicitly grants export authorisation "
            "(+0.25). Enforcement: Art. 29 penalises unauthorised export (3-5 years + 50,000,000-"
            "100,000,000 FCFA) (+0.25). Monitoring: not specifically evidenced -- not awarded. "
            "Unconditional: not awarded -- export is explicitly conditional on prior authorisation and "
            "destination-country criteria."
        ),
        W_comments=(
            "The export-side counterpart to the Art. 19 import ban, both explicitly cited in the task "
            "brief's 'plastic-waste import/export controls'."
        ),
    ),
    dict(
        P_instrument_type=0.60,
        Q_instrument_lifecycle_stage="Recycling",
        R_instrument_description=(
            "Art. 21: 'Il est établi un prix plancher auquel les entreprises du recyclage sont tenus "
            "d'acheter le kilogramme de déchet plastique. Ce prix plancher est fixé par décret.' A "
            "minimum floor price that recycling companies must pay for plastic waste, intended to "
            "support and stabilise the plastic-recycling value chain."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: the floor price is 'fixé par décret' by the competent authority (+0.25). "
            "Enforcement: Art. 31 penalises recyclers buying below the floor price (2,000,000-"
            "5,000,000 FCFA fine) (+0.25). Monitoring: not specifically evidenced -- not awarded. "
            "Unconditional: applies to 'les entreprises du recyclage' generally with no exemption "
            "identified (+0.25)."
        ),
        W_comments=(
            "An economic-incentive instrument supporting the informal/formal waste-picking and "
            "recycling sector's viability by guaranteeing a minimum sale price for collected plastic "
            "waste."
        ),
    ),
    dict(
        P_instrument_type=0.60,
        Q_instrument_lifecycle_stage="Production",
        R_instrument_description=(
            "Art. 22: 'Il est institué une taxe sur les produits constitués ou fabriqués à partir de "
            "matières plastiques non recyclables dite «taxe plastique». La taxe plastique frappe les "
            "produits figurant sur une liste établie par un décret qui en fixe le tarif et les "
            "modalités de recouvrement.' A dedicated tax targeting non-recyclable plastic products, "
            "matching the task brief's description ('plastic tax on non-recyclable products')."
        ),
        S_instrument_in_force=0,
        T_instrument_implementation=0.25,
        U_instrument_implementation_text=(
            "Authority: 'un décret' is explicitly contemplated to fix the tariff and collection "
            "modalities (+0.25). Enforcement/Monitoring/Unconditional: not awarded -- none is evidenced "
            "within the law's own text for this specific tax."
        ),
        W_comments=(
            "In-force test (Rule 12): although 'il est institué' (it is established) reads as "
            "declarative, the tax has no operative content whatsoever without its implementing decree "
            "-- the very list of taxed products, the rate, and the collection procedure are all "
            "entirely absent from the law and expressly left to that decree. Unlike Art. 16's recycled-"
            "content duty (which has a self-executing general obligation even before its decree), this "
            "tax cannot be assessed or collected at all until the decree issues -- a clearer case for "
            "S=0 than most of the other decree-dependent provisions in this table, and coded "
            "accordingly, with the caveat that no verification was performed here on whether such a "
            "decree has since been adopted."
        ),
    ),
]

# ---------------------------------------------------------------------------
# Assemble rows, computing auto-fields O and V
# ---------------------------------------------------------------------------

O_score = round(
    (policy_fields["G_policy_type"] + policy_fields["I_policy_integration"]
     + policy_fields["K_policy_circularity"] + policy_fields["M_policy_budget"]) / 4,
    3,
)
policy_fields["O_policy_score"] = f"{O_score} [auto = (G+I+K+M)/4]"

rows = []
for inst in instruments:
    row = dict(policy_fields)
    row.update(inst)
    s = row["S_instrument_in_force"]
    p = row["P_instrument_type"]
    t = row["T_instrument_implementation"]
    v = round(s * (p + t) / 2, 3)
    row["V_instrument_score"] = f"{v} [auto = S*(P+T)/2]"
    rows.append(row)

with open("4p_index/senegal_loi_2020-04_4p_index.csv", "w", newline="", encoding="utf-8") as f:
    writer = csv.DictWriter(f, fieldnames=COLUMNS)
    writer.writeheader()
    for row in rows:
        writer.writerow(row)

df = pd.DataFrame(rows, columns=COLUMNS)
xlsx_path = "4p_index/senegal_loi_2020-04_4p_index.xlsx"
df.to_excel(xlsx_path, sheet_name="4P_Index_Loi2020-04", index=False)

wb = load_workbook(xlsx_path)
ws = wb["4P_Index_Loi2020-04"]

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF")
wrap_top = Alignment(wrap_text=True, vertical="top")

col_widths = {
    "A": 34, "B": 30, "C": 10, "D": 46, "E": 10, "F": 30, "G": 10, "H": 34,
    "I": 10, "J": 30, "K": 10, "L": 26, "M": 10, "N": 30, "O": 16,
    "P": 12, "Q": 16, "R": 46, "S": 10, "T": 12, "U": 40, "V": 16, "W": 40,
}
for letter in [get_column_letter(c) for c in range(1, len(COLUMNS) + 1)]:
    ws.column_dimensions[letter].width = col_widths.get(letter, 20)

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for cell in row:
        cell.alignment = wrap_top

ws.freeze_panes = "A2"
ws.row_dimensions[1].height = 30
for r in range(2, ws.max_row + 1):
    ws.row_dimensions[r].height = 130

wb.save(xlsx_path)
print("Wrote", xlsx_path, "and CSV with", len(rows), "instrument rows.")
print("Policy score O =", O_score)
for row in rows:
    print(row["Q_instrument_lifecycle_stage"], "->", row["V_instrument_score"])
