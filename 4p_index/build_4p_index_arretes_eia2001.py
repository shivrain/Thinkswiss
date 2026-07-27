"""
Builds the 4P Index (Plastic Pollution Policy Index) coding table for:

Arretes ministeriels n. 9468-9472 MJEHP-DEEC du 28 novembre 2001 relatifs aux etudes d'impact
environnemental (Senegal) -- a package of five ministerial orders implementing the EIA titles of
Decret n. 2001-282 (itself implementing Loi n. 2001-01 portant Code de l'Environnement):

  - n. 9468: reglementation de la participation du public a l'EIE
  - n. 9469: organisation et fonctionnement du Comite technique
  - n. 9470: conditions de delivrance de l'agrement des bureaux d'etude / consultants EIE
  - n. 9471: contenu des termes de reference des etudes d'impact  (user-uploaded PDF)
  - n. 9472: contenu du rapport de l'etude d'impact environnemental

Only n. 9471 was supplied by the user as an uploaded PDF; the full text of the other four orders
(n. 9468, 9469, 9470, 9472) was retrieved from their official mirrors (denv.gouv.sn / FAOLEX) to
code the complete "regulatory package" described in the task brief. All five orders share the same
date, reference prefix (MJEHP-DEEC) and Journal Officiel citation (J.O. n. 6025 du 12 janvier 2002).

Produces:
  - senegal_arretes_eia2001_4p_index.csv
  - senegal_arretes_eia2001_4p_index.xlsx (formatted, wrapped, one sheet)
"""
import csv
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [
    "A_policy_name",
    "B_policy_url",
    "C_policy_year",
    "D_policy_objective",
    "E_policy_target",
    "F_policy_target_text",
    "G_policy_type",
    "H_policy_type_justification",
    "I_policy_integration",
    "J_policy_sectors_list",
    "K_policy_circularity",
    "L_policy_lifecycle_phases_list",
    "M_policy_budget",
    "N_policy_budget_text",
    "O_policy_score",
    "P_instrument_type",
    "Q_instrument_lifecycle_stage",
    "R_instrument_description",
    "S_instrument_in_force",
    "T_instrument_implementation",
    "U_instrument_implementation_text",
    "V_instrument_score",
    "W_comments",
]

# ---------------------------------------------------------------------------
# Section A - Policy-level fields (identical across every row)
# ---------------------------------------------------------------------------

POLICY_NAME = (
    "Arrêtés ministériels du 28 novembre 2001 relatifs aux études d'impact environnemental "
    "(n° 9468, 9469, 9470, 9471 et 9472 MJEHP-DEEC)"
)
POLICY_URL = (
    "https://docs.google.com/viewer?url=https%3A%2F%2Fwww.denv.gouv.sn%2Fwp-admin%2Fadmin-ajax.php"
    "%3Fjuwpfisadmin%3Dfalse%26action%3Dwpfd%26task%3Dfile.download%26wpfd_category_id%3D81"
    "%26wpfd_file_id%3D19565%26token%3D%26preview%3D1&embedded=true (n° 9471, as pointed to in the "
    "task brief; user also uploaded this order's PDF directly) ; www.denv.gouv.sn/decrets/ (lists all "
    "five orders) ; full text of n° 9468/9469/9470/9472 cross-checked against "
    "https://faolex.fao.org/docs/pdf/sen54258.pdf (9468), sen54259.pdf (9469), sen54260.pdf (9470), "
    "sen54262.pdf (9472)"
)
POLICY_YEAR = 2001
POLICY_OBJECTIVE = (
    "A package of five ministerial orders, all signed the same day (28 November 2001) under the same "
    "reference prefix (MJEHP-DEEC) and published together in the same Journal Officiel (n° 6025 du 12 "
    "janvier 2002), that operationalise the environmental impact assessment (EIA/EIE) machinery created "
    "by Décret n° 2001-282 (Art. R38-R44) implementing Loi n° 2001-01 portant Code de l'Environnement: "
    "(i) n° 9468 regulates public participation in the EIA process; (ii) n° 9469 establishes the "
    "composition and functions of the inter-ministerial 'Comité technique' that validates EIA reports; "
    "(iii) n° 9470 sets the accreditation conditions for EIA consultants/firms; (iv) n° 9471 fixes the "
    "mandatory content of EIA terms of reference; and (v) n° 9472 fixes the mandatory content of the EIA "
    "report itself. None of the five orders mentions 'plastics' or 'synthetic polymer materials' -- this "
    "is a procedural/institutional framework, not a plastics-specific policy -- but it is the concrete "
    "gateway through which any plastic manufacturing plant, recycling facility, landfill or incinerator "
    "in Senegal must pass an EIA before being authorised, per the task brief."
)
POLICY_TARGET = 0
POLICY_TARGET_TEXT = ""
POLICY_TYPE = 0.75
POLICY_TYPE_JUSTIFICATION = (
    "Ministerial orders ('arrêtés ministériels') signed by the Minister responsible for the Environment "
    "(MJEHP, i.e. Ministre de la Jeunesse, de l'Environnement et de l'Hygiène publique) with the DEEC "
    "reference, implementing a decree that itself implements an Act of Parliament (Loi n° 2001-01). "
    "Executive-branch, sub-legislative instruments, not enacted by Parliament -> 0.75 per Rule 6, the "
    "same tier as the parent implementing decree."
)
POLICY_INTEGRATION = 1
POLICY_SECTORS_LIST = (
    "environment, industry, agriculture, mining, energy, livestock, urban planning & construction, "
    "commerce/trade, water resources, tourism, local government/municipalities, health, forestry"
)
POLICY_CIRCULARITY = 0.75
POLICY_LIFECYCLE_PHASES_LIST = "production, disposal, environmental leakage"
POLICY_BUDGET = 0
POLICY_BUDGET_TEXT = ""
POLICY_SCORE_FORMULA = "[auto] O = (G + I + K + M) / 4"

policy_fields = dict(
    A_policy_name=POLICY_NAME,
    B_policy_url=POLICY_URL,
    C_policy_year=POLICY_YEAR,
    D_policy_objective=POLICY_OBJECTIVE,
    E_policy_target=POLICY_TARGET,
    F_policy_target_text=POLICY_TARGET_TEXT,
    G_policy_type=POLICY_TYPE,
    H_policy_type_justification=POLICY_TYPE_JUSTIFICATION,
    I_policy_integration=POLICY_INTEGRATION,
    J_policy_sectors_list=POLICY_SECTORS_LIST,
    K_policy_circularity=POLICY_CIRCULARITY,
    L_policy_lifecycle_phases_list=POLICY_LIFECYCLE_PHASES_LIST,
    M_policy_budget=POLICY_BUDGET,
    N_policy_budget_text=POLICY_BUDGET_TEXT,
    O_policy_score=POLICY_SCORE_FORMULA,
)

# ---------------------------------------------------------------------------
# Section B - Instrument-level fields (one dict per row = one arrêté)
# ---------------------------------------------------------------------------

instruments = [
    dict(
        P_instrument_type=0.40,
        Q_instrument_lifecycle_stage="Production",
        R_instrument_description=(
            "Arrêté n° 9468: regulates public participation as a constitutive part of the EIA process -- "
            "announcement of the project by posting at the town hall/governance and/or press notice; "
            "deposit of documents with the local authority; an information meeting; collection of written/"
            "oral comments; negotiation if needed; and a public hearing on-site within 15 days of internal "
            "validation of the EIA report, chaired by the relevant technical ministry with the local "
            "authority as vice-chair and DEEC as secretariat, culminating in a public-hearing report that "
            "the promoter must address within two weeks (Art. 1-8)."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.50,
        U_instrument_implementation_text=(
            "Authority: the process explicitly involves the 'comité technique', the 'promoteur', the "
            "'collectivité décentralisée hôte' and the DEEC secretariat, with hearings chaired by the "
            "relevant technical ministry (Art. 4, 7) (+0.25). Monitoring: the local authority has 10 days "
            "to submit written observations on the deposited report, and the public hearing itself is "
            "structured as a formal consultation/comment-collection mechanism with a dedicated hearing "
            "report (Art. 5-6) (+0.25). Enforcement: no specific penalty for skipping/short-circuiting "
            "public participation is stated within this order itself -- not awarded. Unconditional: not "
            "awarded -- 'les modalités d'exécution de l'audience seront retenues d'un commun accord avec "
            "les différentes parties impliquées' (Art. 4), i.e. implementation details are negotiated "
            "case-by-case."
        ),
        W_comments=(
            "Coded as a Procedural measure per the multi-level/multi-stakeholder coordination logic of "
            "Rule 11 (Ministry, technical ministry, local authority and promoter each have defined roles), "
            "rather than as a purely voluntary measure, since the participation steps are mandatory "
            "('doit informer', 'sera présidée') rather than optional."
        ),
    ),
    dict(
        P_instrument_type=0.40,
        Q_instrument_lifecycle_stage="Production",
        R_instrument_description=(
            "Arrêté n° 9469: establishes the 'Comité technique' as the standing inter-ministerial "
            "institution that administers the EIA process and internally validates EIA reports before the "
            "Minister of Environment decides on a project (Art. 1, 3). Its ~25-member composition spans "
            "Environment, Planning, Agriculture, Industry, Mines, Energy, Livestock, Urban Planning, "
            "the Institut des Sciences de l'Environnement, APIX, Commerce, the National Assembly's "
            "environment committee, Water & Forests, industry associations (SPIDS), Public Works, "
            "National Meteorology, mayors'/rural-community associations, Plant Protection, Public Health, "
            "Territorial Planning, Hydraulics, National Parks and Tourism (Art. 2); its DEEC-run "
            "secretariat prepares public hearings, periodically inspects project sites, and coordinates "
            "between government and operators (Art. 5)."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: the Comité technique's composition, secretariat (DEEC) and functions are explicitly "
            "fixed (Art. 1-2, 5) (+0.25). Monitoring: the secretariat is explicitly tasked to 'inspecter "
            "périodiquement les sites des projets' and to keep the Comité informed of project developments "
            "via reports and meetings (Art. 5) (+0.25). Unconditional: the Comité's validation and advisory "
            "functions apply to 'tous les projets assujettis à l'étude d'impact sur l'environnement' with "
            "no exemption identified (Art. 1) (+0.25). Enforcement: no penalty is tied to the Committee's "
            "own functioning (as opposed to substantive EIA non-compliance, covered by other instruments) "
            "-- not awarded."
        ),
        W_comments=(
            "A clear example of a multi-level/multi-sectoral institutional-coordination instrument under "
            "Rule 11, spanning national ministries, a legislative committee, local-government associations "
            "and civil-society/industry bodies."
        ),
    ),
    dict(
        P_instrument_type=0.40,
        Q_instrument_lifecycle_stage="Production",
        R_instrument_description=(
            "Arrêté n° 9470: sets the accreditation ('agrément') regime for individuals/firms authorised to "
            "conduct EIAs in Senegal. Applications are reviewed by a technical commission chaired by the "
            "Minister of Environment or representative (Art. 1-2); eligibility requires an advanced degree "
            "and sufficient EIA experience for individuals, or a team of at least five senior EIA experts "
            "and adequate logistical/IT resources for firms, plus good administrative standing (Art. 3-4); "
            "accreditation is granted for a renewable 5-year term and can be withdrawn for serious "
            "professional misconduct, loss of civic rights, or loss of the required qualifications (Art. "
            "6)."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: a technical commission chaired by the Minister of Environment (or representative) "
            "reviews applications and prepares the Minister's decision (Art. 1-2) (+0.25). Enforcement: "
            "accreditation 'peut être retiré...pour...manquement grave aux obligations professionnelles "
            "(qualité des travaux)...perte de droits civiques...perte de qualité requise' (Art. 6) -- an "
            "explicit sanction (+0.25). Monitoring: withdrawal decisions are based on 'le rapport du "
            "secrétariat de la Direction de l'Environnement', implying an underlying oversight function "
            "over accredited consultants' work (Art. 6) (+0.25). Unconditional: not awarded -- "
            "accreditation is expressly conditional on meeting the qualification criteria in Art. 4."
        ),
        W_comments=(
            "Functions as a quality/competence gate for the EIA market rather than a product standard; "
            "coded as Procedural ('institutional setup') rather than Regulatory, since the instrument-type "
            "table's Regulatory examples (bans, technology/performance standards, take-back, EPR) concern "
            "products/materials, not professional accreditation -- flagged as a borderline case where a "
            "'Regulatory' (1.0) coding could also be argued given its binding qualification-standard "
            "character."
        ),
    ),
    dict(
        P_instrument_type=0.40,
        Q_instrument_lifecycle_stage="Production",
        R_instrument_description=(
            "Arrêté n° 9471 (the order supplied by the user): fixes the mandatory content of EIA terms of "
            "reference. Art. 1 lists 14 mandatory elements, including baseline environmental description, "
            "assessment of water/energy/raw-material supply impacts, social impact analysis, 'une "
            "évaluation des mesures envisagées pour l'évacuation des eaux usées, l'élimination des déchets "
            "solides et la réduction des émissions' (item 4), alternatives analysis, mitigation measures "
            "with costed implementation plans, and a non-technical summary of recommendations."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.50,
        U_instrument_implementation_text=(
            "Authority: 'Le Directeur de l'Environnement et des Etablissements classés...est chargé de "
            "l'exécution du présent arrêté' (Art. 3); the DEEC may also elaborate project-specific terms of "
            "reference (Art. 2) (+0.25). Unconditional: the 14-point content list applies to 'toute étude "
            "d'impact sur l'environnement' with no exemption identified (Art. 1) (+0.25). Enforcement: no "
            "penalty for non-compliant terms of reference is stated within this order itself -- not "
            "awarded. Monitoring: no distinct inspection/audit/data-collection mechanism beyond the content "
            "list itself was identified in this order -- not awarded (contrast with n° 9472's explicit "
            "rejection mechanism for non-compliant reports, coded separately below)."
        ),
        W_comments=(
            "Explicitly plastics-relevant despite never naming plastics: item 4 requires EIA terms of "
            "reference to assess 'l'élimination des déchets solides' -- directly applicable to plastic "
            "waste generated by any project subject to this requirement."
        ),
    ),
    dict(
        P_instrument_type=0.40,
        Q_instrument_lifecycle_stage="Production",
        R_instrument_description=(
            "Arrêté n° 9472: fixes the mandatory content of the EIA report (REIE) itself -- 16 required "
            "sections including project description, baseline environmental analysis, legal framework "
            "summary, alternatives analysis, impact assessment, technological-accident risk analysis, "
            "mitigation/compensation measures, and an environmental monitoring and surveillance plan "
            "('plan de surveillance et de suivi de l'environnement', PSE) with costed timelines and "
            "responsible bodies (Art. 1); the report must be entirely in French, submitted in 10 copies "
            "(Art. 2); and any report failing to meet these requirements 'sera déclaré irrecevable' (Art. "
            "3)."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: 'Le Directeur de l'Environnement et des Etablissements classés...est chargé de "
            "l'exécution du présent arrêté' (Art. 4) (+0.25). Enforcement: 'Tout rapport...qui ne satisfait "
            "pas aux dispositions des articles précédemment cités sera déclaré irrecevable' -- an explicit, "
            "self-contained rejection mechanism (Art. 3) (+0.25). Monitoring: the mandatory 'plan de "
            "surveillance et de suivi de l'environnement' (PSE) that every report must contain is itself an "
            "environmental monitoring instrument, and report compliance is reviewed by the Comité technique "
            "(Art. 1 §13, cross-referencing Arrêté n° 9469) (+0.25). Unconditional: not awarded -- Art. 1 "
            "§16 expressly allows industrial promoters to withhold confidential manufacturing-process "
            "information into a separate, non-public document, a defined carve-out from full disclosure."
        ),
        W_comments=(
            "The mandatory PSE (monitoring/surveillance plan) requirement is the clearest direct link "
            "between this order and ongoing environmental monitoring of a project once operational -- "
            "relevant to tracking plastic-waste and emissions management commitments made in a plastics-"
            "sector EIA."
        ),
    ),
]

# ---------------------------------------------------------------------------
# Assemble rows, computing auto-fields O and V
# ---------------------------------------------------------------------------

O_score = round(
    (policy_fields["G_policy_type"] + policy_fields["I_policy_integration"]
     + policy_fields["K_policy_circularity"] + policy_fields["M_policy_budget"]) / 4,
    3,
)
policy_fields["O_policy_score"] = f"{O_score} [auto = (G+I+K+M)/4]"

rows = []
for inst in instruments:
    row = dict(policy_fields)
    row.update(inst)
    s = row["S_instrument_in_force"]
    p = row["P_instrument_type"]
    t = row["T_instrument_implementation"]
    v = round(s * (p + t) / 2, 3)
    row["V_instrument_score"] = f"{v} [auto = S*(P+T)/2]"
    rows.append(row)

with open("4p_index/senegal_arretes_eia2001_4p_index.csv", "w", newline="", encoding="utf-8") as f:
    writer = csv.DictWriter(f, fieldnames=COLUMNS)
    writer.writeheader()
    for row in rows:
        writer.writerow(row)

df = pd.DataFrame(rows, columns=COLUMNS)
xlsx_path = "4p_index/senegal_arretes_eia2001_4p_index.xlsx"
df.to_excel(xlsx_path, sheet_name="4P_Index_ArretesEIA2001", index=False)

wb = load_workbook(xlsx_path)
ws = wb["4P_Index_ArretesEIA2001"]

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF")
wrap_top = Alignment(wrap_text=True, vertical="top")

col_widths = {
    "A": 34, "B": 30, "C": 10, "D": 46, "E": 10, "F": 30, "G": 10, "H": 34,
    "I": 10, "J": 30, "K": 10, "L": 26, "M": 10, "N": 30, "O": 16,
    "P": 12, "Q": 16, "R": 46, "S": 10, "T": 12, "U": 40, "V": 16, "W": 40,
}
for letter in [get_column_letter(c) for c in range(1, len(COLUMNS) + 1)]:
    ws.column_dimensions[letter].width = col_widths.get(letter, 20)

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for cell in row:
        cell.alignment = wrap_top

ws.freeze_panes = "A2"
ws.row_dimensions[1].height = 30
for r in range(2, ws.max_row + 1):
    ws.row_dimensions[r].height = 120

wb.save(xlsx_path)
print("Wrote", xlsx_path, "and CSV with", len(rows), "instrument rows.")
print("Policy score O =", O_score)
for row in rows:
    print(row["Q_instrument_lifecycle_stage"], "->", row["V_instrument_score"])
