"""
Builds the 4P Index (Plastic Pollution Policy Index) coding table for:

Norme senegalaise NS 05-061 -- Eaux usees : Normes de rejet (juillet 2001), together with its
implementing "Arrete Interministeriel fixant les conditions de rejets des eaux usees" (undated in
the extracted text, but citing decrees up to 30 August 2001, so signed later in 2001), which gives
the technical standard binding legal force.

Source: user-uploaded PDF (Source: Direction de l'Environnement et des Etablissements Classes /
Institut Senegalais de Normalisation), consistent with the task-brief URL
www.denv.gouv.sn/telechargement/74/normes-cadre-juridique/19105/nor...

Produces:
  - senegal_norme_ns05061_4p_index.csv
  - senegal_norme_ns05061_4p_index.xlsx (formatted, wrapped, one sheet)
"""
import csv
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [
    "A_policy_name",
    "B_policy_url",
    "C_policy_year",
    "D_policy_objective",
    "E_policy_target",
    "F_policy_target_text",
    "G_policy_type",
    "H_policy_type_justification",
    "I_policy_integration",
    "J_policy_sectors_list",
    "K_policy_circularity",
    "L_policy_lifecycle_phases_list",
    "M_policy_budget",
    "N_policy_budget_text",
    "O_policy_score",
    "P_instrument_type",
    "Q_instrument_lifecycle_stage",
    "R_instrument_description",
    "S_instrument_in_force",
    "T_instrument_implementation",
    "U_instrument_implementation_text",
    "V_instrument_score",
    "W_comments",
]

# ---------------------------------------------------------------------------
# Section A - Policy-level fields (identical across every row)
# ---------------------------------------------------------------------------

POLICY_NAME = (
    "Norme sénégalaise NS 05-061 — Eaux usées : Normes de rejet (juillet 2001), rendue obligatoire par "
    "l'Arrêté Interministériel fixant les conditions de rejets des eaux usées"
)
POLICY_URL = (
    "www.denv.gouv.sn/telechargement/74/normes-cadre-juridique/19105/nor... (as provided in task brief; "
    "truncated/unverified) ; full text sourced from user-uploaded PDF (Source: Direction de "
    "l'Environnement et des Etablissements Classés / Institut Sénégalais de Normalisation)"
)
POLICY_YEAR = 2001
POLICY_OBJECTIVE = (
    "Sets threshold/limit values (concentration and flux limits for suspended solids, COD, BOD5, pH, "
    "temperature, heavy metals -- chromium, arsenic, cyanides, fluorides --, total hydrocarbons, phenols, "
    "faecal coliforms, etc.) for industrial and municipal wastewater discharges into surface, groundwater "
    "and marine receiving environments, plus absolute prohibitions on certain discharges and a mandatory "
    "sampling/monitoring and pollution-charge regime. It never names 'plastics' or 'synthetic polymer "
    "materials' and sets no plastics-specific target, but it is the technical standard through which "
    "industrial effluent -- including washing/processing/production wastewater from plastics "
    "manufacturing, recycling and packaging operations -- is regulated and compliance-monitored in "
    "Senegal, and it directly complements the classified-installations (ICPE) and pollution-tax regimes "
    "created by Loi n° 2001-01 and operationalised by Décret n° 2001-282 (both already coded in this "
    "series): its own Avant-Propos states 'La présente norme vient compléter le décret n° 2001-282 du 12 "
    "avril 2001 portant application de la loi n° 2001-01 du 15 janvier 2001 portant Code de "
    "l'Environnement.'"
)
POLICY_TARGET = 0
POLICY_TARGET_TEXT = ""
POLICY_TYPE = 0.75
POLICY_TYPE_JUSTIFICATION = (
    "The underlying document (NS 05-061) is a national technical standard adopted by a normalisation "
    "technical committee (ISN/CT5, chaired by the Direction de l'Environnement et des Etablissements "
    "Classés) rather than a legal instrument in itself. However, it was made legally binding by a joint "
    "'Arrêté Interministériel' signed by the Minister of Youth, Environment and Public Hygiene and the "
    "Minister of Handicrafts and Industry: 'Le présent arrêté a pour objet d'appliquer la norme NS 05-061 "
    "...réglementant les rejets des eaux usées' (Art. 1), with a mandatory 6-month compliance deadline "
    "(Art. 8) and sanctions cross-referenced to the Code de l'Environnement and its implementing decree "
    "(Art. 7). A joint ministerial order is a sub-legislative, executive-branch instrument, not an Act of "
    "Parliament -> coded 0.75 per Rule 6, the same tier as the implementing decree it operationalises."
)
POLICY_INTEGRATION = 1
POLICY_SECTORS_LIST = (
    "environment, industry, sanitation/wastewater management, health, agriculture, "
    "fisheries & aquaculture, water resources, local government/municipalities, mining, energy"
)
POLICY_CIRCULARITY = 0.75
POLICY_LIFECYCLE_PHASES_LIST = "production, consumption, disposal, environmental leakage"
POLICY_BUDGET = 0.5
POLICY_BUDGET_TEXT = (
    "Art. 9 (Arrêté): 'Une redevance annuelle est exigible pour toute installation rejetant des effluents "
    "dans un milieu naturel pourvu ou non de station d'épuration. Elle est fixée à 180 F CFA par kg de "
    "charge polluante.' Art. 10 ties payment of 'les différentes taxes' to Art. L73 of Loi n° 2001-01 and "
    "to Décret n° 2001-282. Coded 0.5 rather than 1 because this document's own text establishes the "
    "funding mechanism (a mandatory pollution-load-based charge) but does not itself restate the "
    "ring-fencing declaration -- the dedicated 'fonds pour la protection de l'Environnement' into which "
    "these charges flow is created by the parent law (Art. L25/L27 of Loi n° 2001-01; Art. 16-19 of Loi "
    "n° 2023-15), consistent with the same convention used for the implementing decree in this series."
)
POLICY_SCORE_FORMULA = "[auto] O = (G + I + K + M) / 4"

policy_fields = dict(
    A_policy_name=POLICY_NAME,
    B_policy_url=POLICY_URL,
    C_policy_year=POLICY_YEAR,
    D_policy_objective=POLICY_OBJECTIVE,
    E_policy_target=POLICY_TARGET,
    F_policy_target_text=POLICY_TARGET_TEXT,
    G_policy_type=POLICY_TYPE,
    H_policy_type_justification=POLICY_TYPE_JUSTIFICATION,
    I_policy_integration=POLICY_INTEGRATION,
    J_policy_sectors_list=POLICY_SECTORS_LIST,
    K_policy_circularity=POLICY_CIRCULARITY,
    L_policy_lifecycle_phases_list=POLICY_LIFECYCLE_PHASES_LIST,
    M_policy_budget=POLICY_BUDGET,
    N_policy_budget_text=POLICY_BUDGET_TEXT,
    O_policy_score=POLICY_SCORE_FORMULA,
)

# ---------------------------------------------------------------------------
# Section B - Instrument-level fields (one dict per row)
# ---------------------------------------------------------------------------

instruments = [
    dict(
        P_instrument_type=1.0,
        Q_instrument_lifecycle_stage="Environmental leakage",
        R_instrument_description=(
            "Chapitre 1 §V.5.2 and Annexe II: detailed technology/performance discharge standards -- "
            "treated effluent 'doit respecter les valeurs indiquées à l'annexe II' before release into a "
            "receiving environment. Annexe II sets concentration/flux limits for suspended solids (50 mg/l), "
            "COD (100-200 mg/l), BOD5 (40-80 mg/l), nitrogen, phosphorus, phenols, hexavalent chromium, "
            "cyanides, arsenic, chromium, total hydrocarbons (15 mg/l), fluorides, temperature (<30°C) and "
            "pH (5.5-9.5), with stricter values for specially protected receiving environments, plus "
            "separate limits for connection to collective treatment plants."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: limit values are set/administered via 'l'arrêté d'autorisation des Etablissements "
            "Classés' issued by the Ministère chargé de l'Environnement (Annexe II §1) (+0.25). Enforcement: "
            "Art. 7 of the Arrêté: 'Toutes infractions aux dispositions normatives contenues dans la norme "
            "NS 05-061...sont passibles de sanctions définies aux articles L96, L97, L98, L100 de la loi "
            "n° 2001-01...et à l'article [R]51 du décret n° 2001-282' (+0.25). Monitoring: mandatory "
            "standardised sampling/flow-measurement devices before any discharge (Annexe I) and periodic "
            "laboratory analysis (Chapitre 3) (+0.25). Unconditional: not awarded -- 'des valeurs limites de "
            "concentration différentes peuvent être fixées par l'arrêté d'autorisation' in several places, "
            "so limits are case-specific rather than uniformly fixed."
        ),
        W_comments=(
            "The core technical standard operationalising discharge limits for classified installations, "
            "including plastics-sector effluent (e.g. washing/extrusion wastewater carrying suspended "
            "solids, COD/BOD load and, for some processes, hydrocarbons)."
        ),
    ),
    dict(
        P_instrument_type=1.0,
        Q_instrument_lifecycle_stage="Environmental leakage",
        R_instrument_description=(
            "Chapitre 1 §V.5.1 & 5.3: absolute prohibitions -- discharges causing stagnation, nuisance or "
            "pollution of surface/ground/marine waters are banned nationwide; separately banned are "
            "discharges of cyclic hydroxylated compounds and halogenated derivatives, substances causing "
            "abnormal odour/taste/colour in waters used for human or animal consumption, hydrocarbon/toxic "
            "chemical discharges from vessels or pipelines, septic-tank truck dumping outside authorised "
            "sites, raw-wastewater irrigation of food/feed crops, and any discharge into lakes, ponds or "
            "pools."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=1.0,
        U_instrument_implementation_text=(
            "Authority: Ministère chargé de l'Environnement / DEEC enforce via the same inspection regime "
            "as the ICPE framework (+0.25). Enforcement: Art. 7 of the Arrêté applies the same "
            "Loi-n°2001-01/Décret-n°2001-282 sanctions to violations of these prohibitions (+0.25). "
            "Monitoring: the sampling/analysis regime of Chapitre 3 covers detection of these prohibited "
            "discharges (+0.25). Unconditional: the prohibitions are stated in absolute terms -- 'est "
            "interdit sur toute l'étendue du territoire national', 'Sont aussi interdits' -- with no "
            "exemption identified (+0.25)."
        ),
        W_comments=(
            "The ban on discharge into lakes/ponds/pools and on dumping of septic/industrial waste at "
            "unauthorised sites is directly relevant to preventing plastic and other solid/liquid waste from "
            "entering standing water bodies used for human/animal consumption or aquaculture."
        ),
    ),
    dict(
        P_instrument_type=0.40,
        Q_instrument_lifecycle_stage="Waste management",
        R_instrument_description=(
            "Chapitre 3 §I and Arrêté Art. 5-6: a compliance-monitoring and accreditation system -- sampling "
            "and analysis performed by DEEC technicians or Ministry-designated entities (§1.1); mandatory "
            "minimum twice-yearly unannounced sampling per discharge point, with continuous flow "
            "measurement above 100 m³/day (§1.3-1.7); an annual reporting dossier required from "
            "installations discharging listed hazardous substances (Annexe II §2); and a formal "
            "accreditation ('agrément') regime for third-party sampling/testing bodies (Arrêté Art. 5), "
            "using the standardised methods listed in Annexe IV as the official reference (Arrêté Art. 6)."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: 'Ministère chargé de l'Environnement' designates/accredits monitoring entities "
            "(Chapitre 3 §1.1; Arrêté Art. 5) (+0.25). Monitoring: this instrument is itself a data-"
            "collection/inspection mechanism -- 'Chaque rejet doit au moins, faire l'objet d'un prélèvement "
            "suivi d'analyses, deux fois par an' (§1.5) (+0.25). Unconditional: the minimum sampling "
            "frequency applies to every discharge point with no exemption identified (+0.25). Enforcement: "
            "no penalty specific to a failure of the monitoring/accreditation process itself (as distinct "
            "from substantive discharge-standard violations, covered in the rows above) was identified -- "
            "not awarded."
        ),
        W_comments=(
            "Coded as a Procedural measure (0.40, 'monitoring...institutional setup') under Rule 11's logic, "
            "since its primary function is establishing the inspection/data-collection architecture rather "
            "than a substantive standard or economic incentive."
        ),
    ),
    dict(
        P_instrument_type=0.60,
        Q_instrument_lifecycle_stage="Waste management",
        R_instrument_description=(
            "Arrêté Art. 9-11 and Annexe I: a mandatory annual pollution charge ('redevance') on every "
            "installation discharging effluent into a natural environment, whether or not it has a "
            "treatment plant, set at 180 FCFA per kg of pollutant load ('charge polluante' = suspended "
            "solids + oxidisable matter), computed via a fixed formula in Annexe I from measured effluent "
            "characteristics; sampling/analysis costs are borne by the operator (Art. 11)."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: the charge is administered under the Ministère chargé de l'Environnement/DEEC "
            "framework (Art. 9-10) (+0.25). Monitoring: the charge is explicitly and directly computed from "
            "controlled/measured effluent data -- 'Elle dépend du degré de pollution exprimé par la somme "
            "des matières en suspension et des matières oxydables', with the formula given in Annexe I "
            "(+0.25). Unconditional: applies to 'toute installation rejetant des effluents dans un milieu "
            "naturel pourvu ou non de station d'épuration', with no exemption identified (+0.25). "
            "Enforcement: no explicit fine/surcharge for late or non-payment of this specific redevance was "
            "identified in the extracted text (contrast with the 45-day payment deadline/penalty in Décret "
            "n° 2001-282, Art. R26) -- not awarded."
        ),
        W_comments=(
            "This is the concrete rate-setting instrument for the pollution charge that Décret n° 2001-282 "
            "(Art. R54) and Loi n° 2001-01 (Art. L73) establish only in general/computational terms; feeds "
            "the same Environmental Protection Fund referenced in Column N."
        ),
    ),
    dict(
        P_instrument_type=1.0,
        Q_instrument_lifecycle_stage="Environmental leakage",
        R_instrument_description=(
            "Chapitre 3 §II and Arrêté Art. 2-4: an enhanced regime for water bodies under special "
            "protection. A zero-discharge objective applies to multi-use water bodies (lakes, ponds, "
            "marshes, water reserves) used for human/animal consumption, fishing or industrial water supply "
            "-- discharge into them is banned outright, and existing authorised dischargers must undergo "
            "case-by-case impact studies leading to a formal agreement ('protocole d'accord') with the "
            "Ministry and the local authority ('collectivité locale') and a phase-out schedule toward zero "
            "discharge; specifically designated sensitive zones (e.g. the Baie de Hann) are subject to "
            "stricter discharge limits and require polluters to establish pollution-reduction programmes "
            "with quarterly analytical reports to the Direction de l'Environnement over a four-year period."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: joint involvement of the 'Ministère Chargé de l'Environnement et des Etablissements "
            "Classés' and the 'Collectivité locale concernée' in the protocole d'accord (Arrêté Art. 2) "
            "(+0.25). Enforcement: the same Art. 7 sanctions apply to violations of this chapter's "
            "provisions (+0.25). Monitoring: explicit quarterly reporting duty for designated sensitive "
            "zones -- 'un rapport trimestriel de résultats d'analyses, soumis à la Direction de "
            "l'Environnement, sur une période de quatre ans' (§2.3) (+0.25). Unconditional: not awarded -- "
            "the phase-out pathway for existing dischargers is explicitly case-by-case ('au cas par cas') "
            "and the scheme applies only to specifically designated protected water bodies, not universally."
        ),
        W_comments=(
            "Mixed instrument per Rule 10: combines a hard regulatory endpoint (zero discharge / stricter "
            "limits) with a multi-level governance coordination mechanism (Ministry + local authority "
            "protocol) under Rule 11 -- scored at the higher Regulatory value (1.0), with the coordination "
            "aspect noted here. Uncertainty (Rule 8): the general case-by-case negotiation track for "
            "existing dischargers into multi-use water bodies uses softer language ('seront commanditées', "
            "'sera proposé') than the flagship Baie-de-Hann example ('doivent établir'); S=1 is applied to "
            "the row as a whole on the strength of the clearly mandatory language in the latter, but a "
            "reasonable alternative reading could split this into two rows with the general track scored "
            "S=0 as a not-yet-operationalised case-by-case power."
        ),
    ),
]

# ---------------------------------------------------------------------------
# Assemble rows, computing auto-fields O and V
# ---------------------------------------------------------------------------

O_score = round(
    (policy_fields["G_policy_type"] + policy_fields["I_policy_integration"]
     + policy_fields["K_policy_circularity"] + policy_fields["M_policy_budget"]) / 4,
    3,
)
policy_fields["O_policy_score"] = f"{O_score} [auto = (G+I+K+M)/4]"

rows = []
for inst in instruments:
    row = dict(policy_fields)
    row.update(inst)
    s = row["S_instrument_in_force"]
    p = row["P_instrument_type"]
    t = row["T_instrument_implementation"]
    v = round(s * (p + t) / 2, 3)
    row["V_instrument_score"] = f"{v} [auto = S*(P+T)/2]"
    rows.append(row)

with open("4p_index/senegal_norme_ns05061_4p_index.csv", "w", newline="", encoding="utf-8") as f:
    writer = csv.DictWriter(f, fieldnames=COLUMNS)
    writer.writeheader()
    for row in rows:
        writer.writerow(row)

df = pd.DataFrame(rows, columns=COLUMNS)
xlsx_path = "4p_index/senegal_norme_ns05061_4p_index.xlsx"
df.to_excel(xlsx_path, sheet_name="4P_Index_NS05-061", index=False)

wb = load_workbook(xlsx_path)
ws = wb["4P_Index_NS05-061"]

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF")
wrap_top = Alignment(wrap_text=True, vertical="top")

col_widths = {
    "A": 34, "B": 30, "C": 10, "D": 46, "E": 10, "F": 30, "G": 10, "H": 34,
    "I": 10, "J": 30, "K": 10, "L": 26, "M": 10, "N": 30, "O": 16,
    "P": 12, "Q": 16, "R": 46, "S": 10, "T": 12, "U": 40, "V": 16, "W": 40,
}
for letter in [get_column_letter(c) for c in range(1, len(COLUMNS) + 1)]:
    ws.column_dimensions[letter].width = col_widths.get(letter, 20)

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for cell in row:
        cell.alignment = wrap_top

ws.freeze_panes = "A2"
ws.row_dimensions[1].height = 30
for r in range(2, ws.max_row + 1):
    ws.row_dimensions[r].height = 120

wb.save(xlsx_path)
print("Wrote", xlsx_path, "and CSV with", len(rows), "instrument rows.")
print("Policy score O =", O_score)
for row in rows:
    print(row["Q_instrument_lifecycle_stage"], "->", row["V_instrument_score"])
