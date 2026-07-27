"""
Builds the 4P Index (Plastic Pollution Policy Index) coding table for:

Norme senegalaise NS 05-062 -- Pollution atmospherique : Norme de rejets (octobre 2003), together
with its implementing "Arrete Interministeriel n. 7358 du 5 novembre 2003 fixant les conditions
d'application de la norme NS 05-062", which gives the technical standard binding legal force.

Source: user-uploaded PDF (OCR'd with tesseract/fra, since the embedded text layer was a corrupted/
non-standard font encoding), cross-checked word-for-word against the clean Journal Officiel mirror
(http://www.jo.gouv.sn, archived via FAOLEX at https://faolex.fao.org/docs/pdf/Sen175486.pdf and
https://faolex.fao.org/docs/pdf/sen54266.pdf), consistent with the task-brief URL
www.denv.gouv.sn/telechargement/74/normes-cadre-juridique/19106/nor...

KEY FINDING: Chapitre II Section 8.2 of the standard explicitly names plastics -- "Le brulage a
l'air libre des pneumatiques, plastiques et tout autre compose renfermant des produits chimiques
est interdit." This is the first document in this coding series with a direct, explicit textual
reference to plastics (Plastics Relevance Filter criterion (a)), rather than only a general waste/
governance nexus (criterion (c)).

Produces:
  - senegal_norme_ns05062_4p_index.csv
  - senegal_norme_ns05062_4p_index.xlsx (formatted, wrapped, one sheet)
"""
import csv
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [
    "A_policy_name",
    "B_policy_url",
    "C_policy_year",
    "D_policy_objective",
    "E_policy_target",
    "F_policy_target_text",
    "G_policy_type",
    "H_policy_type_justification",
    "I_policy_integration",
    "J_policy_sectors_list",
    "K_policy_circularity",
    "L_policy_lifecycle_phases_list",
    "M_policy_budget",
    "N_policy_budget_text",
    "O_policy_score",
    "P_instrument_type",
    "Q_instrument_lifecycle_stage",
    "R_instrument_description",
    "S_instrument_in_force",
    "T_instrument_implementation",
    "U_instrument_implementation_text",
    "V_instrument_score",
    "W_comments",
]

# ---------------------------------------------------------------------------
# Section A - Policy-level fields (identical across every row)
# ---------------------------------------------------------------------------

POLICY_NAME = (
    "Norme sénégalaise NS 05-062 — Pollution atmosphérique : Norme de rejets (octobre 2003), rendue "
    "obligatoire par l'Arrêté Interministériel n° 7358 du 5 novembre 2003 fixant les conditions "
    "d'application de la norme NS 05-062"
)
POLICY_URL = (
    "www.denv.gouv.sn/telechargement/74/normes-cadre-juridique/19106/nor... (as provided in task "
    "brief; truncated/unverified) ; full text sourced from user-uploaded PDF (Association Sénégalaise "
    "de Normalisation / Direction de l'Environnement et des Etablissements Classés), cross-checked "
    "against https://faolex.fao.org/docs/pdf/Sen175486.pdf and https://faolex.fao.org/docs/pdf/"
    "sen54266.pdf and the Journal Officiel mirror at jo.gouv.sn"
)
POLICY_YEAR = 2003
POLICY_OBJECTIVE = (
    "Sets emission limit/technology-performance values for atmospheric pollutants (particulates, "
    "carbon monoxide, sulphur/nitrogen oxides, carcinogenic substances, odorous compounds, etc.) "
    "released by stationary industrial installations and vehicles, and mandates operator self-"
    "monitoring, local air-quality planning and an air-pollution tax. Chapitre II §8.2 explicitly and "
    "directly names plastics: 'Le brûlage à l'air libre des pneumatiques, plastiques et tout autre "
    "composé renfermant des produits chimiques est interdit' -- an outright ban on open-air burning of "
    "plastics (and tyres), which is otherwise a very common, highly polluting informal disposal "
    "practice for plastic waste in Senegal and across the region. §8.1 separately requires that waste "
    "incineration/thermal decomposition only occur in technologically dedicated installations -- "
    "directly applicable to plastic-waste incinerators. The standard complements NS 05-061 "
    "(wastewater) in the overall industrial pollution-control regime that governs plastics production, "
    "processing, recycling and waste-treatment facilities in Senegal, per its own Avant-Propos: 'La "
    "présente norme vient compléter le décret n° 2001-282 du 12 avril 2001 portant application de la "
    "loi n° 2001-01 du 15 janvier 2001 portant Code de l'Environnement.'"
)
POLICY_TARGET = 0
POLICY_TARGET_TEXT = ""
POLICY_TYPE = 0.75
POLICY_TYPE_JUSTIFICATION = (
    "The underlying document (NS 05-062) is a national technical standard adopted by a normalisation "
    "technical committee (ASN/CT5, chaired by the Direction de l'Environnement et des Etablissements "
    "Classés), not itself a legal instrument. It was made legally binding by the joint 'Arrêté "
    "Interministériel n° 7358 du 5 novembre 2003': 'Le présent arrêté a pour objet d'appliquer la norme "
    "NS 05-062 réglementant les conditions de rejets de polluants atmosphériques dans l'air ambiant' "
    "(Art. 1), tying compliance to Art. L78 of the Code de l'Environnement and sanctions to Art. L99 "
    "(Art. 2, 8). A ministerial/interministerial order is a sub-legislative, executive-branch "
    "instrument, not an Act of Parliament -> coded 0.75 per Rule 6, the same tier as NS 05-061 and its "
    "own implementing order."
)
POLICY_INTEGRATION = 1
POLICY_SECTORS_LIST = (
    "environment, industry, energy, transport, health, local government/municipalities, agriculture, "
    "chemicals, commerce/trade, meteorology"
)
POLICY_CIRCULARITY = 0.75
POLICY_LIFECYCLE_PHASES_LIST = "production, consumption, disposal, environmental leakage"
POLICY_BUDGET = 0.5
POLICY_BUDGET_TEXT = (
    "Art. 12 (Arrêté): 'La taxe sur la pollution de l'air est exigible pour toute installation "
    "stationnaire ou mobile et tout véhicule rejetant des polluants atmosphériques dépassant la "
    "norme.' Art. 13 ties payment to Art. L73 of Loi n° 2001-01 and to Décret n° 2001-282. Coded 0.5 "
    "rather than 1 for the same reason as the other 2001-282-linked instruments in this series: this "
    "document establishes/operationalises a mandatory pollution charge but does not itself restate the "
    "ring-fencing declaration for the 'fonds pour la protection de l'Environnement' created by the "
    "parent law. Government sourcing (denv.gouv.sn) further indicates the proceeds are paid into general "
    "Treasury accounts with only 'une partie' (a portion) subsequently allocated to air-pollution-"
    "control actions -- i.e. partial, not full, earmarking."
)
POLICY_SCORE_FORMULA = "[auto] O = (G + I + K + M) / 4"

policy_fields = dict(
    A_policy_name=POLICY_NAME,
    B_policy_url=POLICY_URL,
    C_policy_year=POLICY_YEAR,
    D_policy_objective=POLICY_OBJECTIVE,
    E_policy_target=POLICY_TARGET,
    F_policy_target_text=POLICY_TARGET_TEXT,
    G_policy_type=POLICY_TYPE,
    H_policy_type_justification=POLICY_TYPE_JUSTIFICATION,
    I_policy_integration=POLICY_INTEGRATION,
    J_policy_sectors_list=POLICY_SECTORS_LIST,
    K_policy_circularity=POLICY_CIRCULARITY,
    L_policy_lifecycle_phases_list=POLICY_LIFECYCLE_PHASES_LIST,
    M_policy_budget=POLICY_BUDGET,
    N_policy_budget_text=POLICY_BUDGET_TEXT,
    O_policy_score=POLICY_SCORE_FORMULA,
)

# ---------------------------------------------------------------------------
# Section B - Instrument-level fields (one dict per row)
# ---------------------------------------------------------------------------

instruments = [
    dict(
        P_instrument_type=1.0,
        Q_instrument_lifecycle_stage="Production",
        R_instrument_description=(
            "Chapitre II §1 and Annexes I-III (Arrêté Art. 2-4): technology/performance emission-limit "
            "standards -- 'Les installations existantes et nouvelles stationnaires doivent être équipées "
            "et exploitées de manière à respecter la limitation maximale des émissions fixée aux annexes "
            "I, II, III', covering particulates, gaseous pollutants, combustion equipment and stack-"
            "design requirements (capture as close to source as possible, chimney height/dispersion "
            "formulas, etc.)."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: 'l'autorité compétente' and the DEEC ensure existing installations are upgraded "
            "to the standard and fix compliance deadlines (Arrêté Art. 3) (+0.25). Enforcement: Arrêté "
            "Art. 8 -- 'Toute infraction aux dispositions contenues dans la norme NS 05-062 peut être "
            "passible de sanctions définies à l'article L99 de la loi n° 2001-01' (+0.25). Monitoring: "
            "extensive declaration, measurement and control regime for emissions (Chapitre II §3-6) "
            "(+0.25). Unconditional: not awarded -- Arrêté Art. 4 explicitly allows owners facing high "
            "compliance investment costs to negotiate a 'protocole d'accord' with the Ministry for "
            "deferred, progressive compliance."
        ),
        W_comments=(
            "The core technology-standard instrument operationalising air-emission limits for classified "
            "installations, including plastics manufacturing/processing/recycling facilities."
        ),
    ),
    dict(
        P_instrument_type=1.0,
        Q_instrument_lifecycle_stage="End of life",
        R_instrument_description=(
            "Chapitre II §8.1: 'L'incinération ou la décomposition thermique des déchets n'est autorisée "
            "que dans des installations technologiquement destinées à cet effet', with the specific "
            "technical provisions of Annexe II, lettre J applicable -- a technology-siting standard "
            "restricting waste incineration/thermal decomposition (including of plastic waste) to "
            "purpose-built facilities."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: the general 'autorité compétente' emissions-control framework applies to all "
            "regulated installations, including incinerators (Chapitre II §4.1) (+0.25). Enforcement: "
            "Arrêté Art. 8's sanctions cross-reference applies to all provisions of the norm, including "
            "this one (+0.25). Unconditional: 'n'est autorisée que dans des installations "
            "technologiquement destinées à cet effet' is stated categorically with no exemption "
            "identified (+0.25). Monitoring: no inspection/audit mechanism specific to verifying "
            "incinerator technology (as distinct from routine emissions monitoring, covered in a "
            "separate row below) was identified -- not awarded."
        ),
        W_comments=(
            "Directly applicable to any plant incinerating plastic waste: only technologically dedicated "
            "installations may lawfully incinerate or thermally decompose waste."
        ),
    ),
    dict(
        P_instrument_type=1.0,
        Q_instrument_lifecycle_stage="End of life",
        R_instrument_description=(
            "Chapitre II §8.2: 'Le brûlage à l'air libre des pneumatiques, plastiques et tout autre "
            "composé renfermant des produits chimiques est interdit.' An outright, nationwide ban on "
            "open-air burning of tyres, plastics and any other chemical-containing compound."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: enforced under the same 'autorité compétente' / sworn-agent framework as the "
            "rest of the norm (Arrêté Art. 6) (+0.25). Enforcement: Arrêté Art. 8's sanctions apply to "
            "'toute infraction aux dispositions contenues dans la norme NS 05-062' (+0.25). Unconditional: "
            "the ban is stated absolutely -- 'est interdit' -- with no exemption identified (+0.25). "
            "Monitoring: no dedicated inspection/detection mechanism specific to illegal open burning "
            "(as distinct from the stack-emissions monitoring regime aimed at authorised installations) "
            "was identified in the text -- not awarded."
        ),
        W_comments=(
            "The single most directly plastics-relevant provision found across this entire coding "
            "series: this is an explicit, named ban on burning plastics in the open air (Plastics "
            "Relevance Filter criterion (a), not merely (b)/(c)). Highly significant in practice because "
            "uncontrolled open burning of plastic waste and packaging is a common informal disposal "
            "method in the absence of formal collection/recycling infrastructure."
        ),
    ),
    dict(
        P_instrument_type=0.40,
        Q_instrument_lifecycle_stage="Environmental leakage",
        R_instrument_description=(
            "Chapitre V and Chapitre II §3-6 (Arrêté Art. 6-7): operators must implement a self-"
            "monitoring programme for their emissions at their own expense (Ch. V §1); for high-emission "
            "installations, continuous measurement/recording may be mandated (Ch. II §4.2); results must "
            "be logged in a report and transmitted to the competent authority at least quarterly, with "
            "at least one annual measurement by an accredited or approved third-party body (Ch. V §3-4); "
            "sampling/testing organisations must be accredited by the Ministry, or sworn agents may be "
            "used (Arrêté Art. 6); the norm's Annexe V methods are the official reference methods "
            "(Arrêté Art. 7)."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: 'l'autorité administrative compétente' and Ministry-accredited bodies/sworn "
            "agents (Arrêté Art. 6) (+0.25). Monitoring: this instrument is itself the data-collection/"
            "inspection architecture -- mandatory quarterly reporting and at least annual third-party "
            "measurement (Ch. V §3-4) (+0.25). Enforcement: the same Art. 8 sanctions apply to breaches of "
            "the monitoring obligations, which are themselves 'dispositions contenues dans la norme' "
            "(+0.25). Unconditional: not awarded -- Ch. V §2 allows the authorisation order to substitute "
            "continuous parameter tracking for some measurements 'pour certains polluants spécifiques et "
            "certains procédés', a case-by-case variation."
        ),
        W_comments=(
            "Coded as a Procedural measure (0.40, 'monitoring') consistent with the equivalent instrument "
            "coded for NS 05-061 (wastewater)."
        ),
    ),
    dict(
        P_instrument_type=0.40,
        Q_instrument_lifecycle_stage="Environmental leakage",
        R_instrument_description=(
            "Arrêté Art. 9: in any agglomeration where the limit values in Annexe I of NS 05-062 are "
            "exceeded or at risk of being exceeded, the local competent authority must draw up a local "
            "'plan de protection de l'atmosphère', submitted for opinion and approval to the ASN's "
            "environment technical committee; local air-protection plans must be evaluated and, if "
            "needed, revised every five years."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.50,
        U_instrument_implementation_text=(
            "Authority: 'l'autorité compétente locale' drafts the plan; the ASN's environment technical "
            "committee reviews/approves it (Arrêté Art. 9) (+0.25). Monitoring: mandatory five-yearly "
            "evaluation/revision cycle for every local plan (Arrêté Art. 9) (+0.25). Enforcement: no "
            "specific penalty for a local authority's failure to adopt a required plan was identified -- "
            "not awarded. Unconditional: not awarded -- the duty is explicitly triggered only where limit "
            "values 'sont dépassées ou risquent de l'être', not universally."
        ),
        W_comments=(
            "Coded per Rule 11 as a multi-level governance/coordination instrument -- a direct air-"
            "quality analogue to the municipal waste-management-planning obligations coded elsewhere in "
            "this series (e.g. Loi n° 2023-15, Art. 87)."
        ),
    ),
    dict(
        P_instrument_type=0.60,
        Q_instrument_lifecycle_stage="Environmental leakage",
        R_instrument_description=(
            "Arrêté Art. 12-15: a mandatory 'taxe sur la pollution de l'air', exigible from every "
            "stationary or mobile installation and every vehicle discharging atmospheric pollutants in "
            "excess of the norm; payment follows the general pollution-tax procedure of Art. L73 of Loi "
            "n° 2001-01 (charge computed from accredited-laboratory sampling of the pollutant load); the "
            "specific calculation method for the air-pollution charge, based on the nature and scale of "
            "the classified installation, is to be fixed by a further ministerial order (Art. 14); "
            "sampling/analysis costs are borne by the installation owner (Art. 15)."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.50,
        U_instrument_implementation_text=(
            "Authority: administered by the 'Ministre chargé de l'Environnement' under the general Art. "
            "L73 tax-collection framework (Arrêté Art. 13) (+0.25). Monitoring: the tax base ('charge "
            "polluante') is explicitly computed from accredited-laboratory sampling/measurement data "
            "(Art. L73; Arrêté Art. 6) (+0.25). Enforcement: no explicit fine/surcharge specific to "
            "non-payment of this air-pollution tax was identified in the text -- not awarded. "
            "Unconditional: not awarded -- the tax applies only to installations/vehicles 'rejetant des "
            "polluants atmosphériques dépassant la norme', a conditional trigger, not a blanket levy."
        ),
        W_comments=(
            "In-force test (Rule 12): the tax obligation itself is stated in mandatory terms ('est "
            "exigible') and rests on the already-operative general Art. L73 charge-computation "
            "mechanism, so S=1 is applied; however, Art. 14's air-pollution-specific calculation "
            "refinement ('Le Ministre...prendra un arrêté pour déterminer le mode de calcul') is itself "
            "an unexercised enabling power at the time of this arrêté -- flagged as a genuine Rule 8 "
            "uncertainty, since the tax's precise, air-pollution-specific quantum may not have been fully "
            "operational until that further order issued (not independently verified here). Third-party "
            "government sourcing indicates the tax is collected in practice, with proceeds paid into "
            "Treasury and only partially earmarked for air-pollution-control action."
        ),
    ),
    dict(
        P_instrument_type=0.20,
        Q_instrument_lifecycle_stage="Environmental leakage",
        R_instrument_description=(
            "Arrêté Art. 10: 'Le droit à l'information sur la qualité de l'air est reconnu à chaque "
            "citoyen sur l'ensemble du territoire' -- the State guarantees this right and the reliability/"
            "dissemination of the information; the Ministry of Environment must regularly publish a "
            "report on air quality and its likely evolution, and must immediately inform the public, "
            "with measured values and advice, whenever air-quality objectives are not met or are at risk "
            "of not being met."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: 'Le Ministère chargé de l'Environnement' is explicitly responsible for publishing "
            "reports and informing the public (Arrêté Art. 10) (+0.25). Monitoring: the duty to "
            "'régulièrement' publish an air-quality report is itself a recurring data-dissemination "
            "requirement (+0.25). Unconditional: the right is 'reconnu à chaque citoyen sur l'ensemble du "
            "territoire' with no exemption identified (+0.25). Enforcement: no penalty is attached to a "
            "failure by the State to provide this information -- not awarded."
        ),
        W_comments=(
            "Coded as a Voluntary/information-type instrument (0.20) since its substance is public "
            "disclosure/awareness rather than a binding standard on emitters, even though the duty on the "
            "State itself is mandatory ('doit', 'est le garant')."
        ),
    ),
]

# ---------------------------------------------------------------------------
# Assemble rows, computing auto-fields O and V
# ---------------------------------------------------------------------------

O_score = round(
    (policy_fields["G_policy_type"] + policy_fields["I_policy_integration"]
     + policy_fields["K_policy_circularity"] + policy_fields["M_policy_budget"]) / 4,
    3,
)
policy_fields["O_policy_score"] = f"{O_score} [auto = (G+I+K+M)/4]"

rows = []
for inst in instruments:
    row = dict(policy_fields)
    row.update(inst)
    s = row["S_instrument_in_force"]
    p = row["P_instrument_type"]
    t = row["T_instrument_implementation"]
    v = round(s * (p + t) / 2, 3)
    row["V_instrument_score"] = f"{v} [auto = S*(P+T)/2]"
    rows.append(row)

with open("4p_index/senegal_norme_ns05062_4p_index.csv", "w", newline="", encoding="utf-8") as f:
    writer = csv.DictWriter(f, fieldnames=COLUMNS)
    writer.writeheader()
    for row in rows:
        writer.writerow(row)

df = pd.DataFrame(rows, columns=COLUMNS)
xlsx_path = "4p_index/senegal_norme_ns05062_4p_index.xlsx"
df.to_excel(xlsx_path, sheet_name="4P_Index_NS05-062", index=False)

wb = load_workbook(xlsx_path)
ws = wb["4P_Index_NS05-062"]

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF")
wrap_top = Alignment(wrap_text=True, vertical="top")

col_widths = {
    "A": 34, "B": 30, "C": 10, "D": 46, "E": 10, "F": 30, "G": 10, "H": 34,
    "I": 10, "J": 30, "K": 10, "L": 26, "M": 10, "N": 30, "O": 16,
    "P": 12, "Q": 16, "R": 46, "S": 10, "T": 12, "U": 40, "V": 16, "W": 40,
}
for letter in [get_column_letter(c) for c in range(1, len(COLUMNS) + 1)]:
    ws.column_dimensions[letter].width = col_widths.get(letter, 20)

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for cell in row:
        cell.alignment = wrap_top

ws.freeze_panes = "A2"
ws.row_dimensions[1].height = 30
for r in range(2, ws.max_row + 1):
    ws.row_dimensions[r].height = 120

wb.save(xlsx_path)
print("Wrote", xlsx_path, "and CSV with", len(rows), "instrument rows.")
print("Policy score O =", O_score)
for row in rows:
    print(row["Q_instrument_lifecycle_stage"], "->", row["V_instrument_score"])
