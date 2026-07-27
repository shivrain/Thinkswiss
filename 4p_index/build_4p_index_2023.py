"""
Builds the 4P Index (Plastic Pollution Policy Index) coding table for:

Loi n° 2023-15 du 02 août 2023 portant Code de l'Environnement (Senegal)
-- repeals and replaces Loi n° 2001-01 du 15 janvier 2001 --

Source verified against:
  - user-uploaded PDF: Journal Officiel de la République du Sénégal, 88e année, n° 7660,
    samedi 16 septembre 2023, pp. 1111-1163 ("Loi n° 2023-15 du 02 août 2023 portant Code de
    l'Environnement") -- confirmed via OCR (tesseract, fra) to be the correct, official
    publication of this law.
  - clean full-text mirror: https://primature.sn/sites/default/files/2023-10/Loi%20n%C2%B02023-15%20du%2002%20ao%C3%BBt%202023%20portant%20Code%20de%20l%27Environnement.pdf
  - https://www.vie-publique.sn/documents/7949/loi-n-2023-15-du-02-aout-2023-portant-code-de-l-environnement-7949

Produces:
  - senegal_loi_2023-15_4p_index.csv
  - senegal_loi_2023-15_4p_index.xlsx (formatted, wrapped, one sheet)
"""
import csv
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [
    "A_policy_name",
    "B_policy_url",
    "C_policy_year",
    "D_policy_objective",
    "E_policy_target",
    "F_policy_target_text",
    "G_policy_type",
    "H_policy_type_justification",
    "I_policy_integration",
    "J_policy_sectors_list",
    "K_policy_circularity",
    "L_policy_lifecycle_phases_list",
    "M_policy_budget",
    "N_policy_budget_text",
    "O_policy_score",
    "P_instrument_type",
    "Q_instrument_lifecycle_stage",
    "R_instrument_description",
    "S_instrument_in_force",
    "T_instrument_implementation",
    "U_instrument_implementation_text",
    "V_instrument_score",
    "W_comments",
]

# ---------------------------------------------------------------------------
# Section A - Policy-level fields (identical across every row)
# ---------------------------------------------------------------------------

POLICY_NAME = "Loi n° 2023-15 du 02 août 2023 portant Code de l'Environnement"
POLICY_URL = (
    "https://primature.sn/sites/default/files/2023-10/Loi%20n%C2%B02023-15%20du%2002%20ao%C3%BBt"
    "%202023%20portant%20Code%20de%20l%27Environnement.pdf ; "
    "https://www.vie-publique.sn/documents/7949/loi-n-2023-15-du-02-aout-2023-portant-code-de-l-"
    "environnement-7949 ; Journal Officiel de la République du Sénégal, n° 7660 du 16 septembre 2023, "
    "pp. 1111-1163 (user-provided PDF, verified by OCR to match)"
)
POLICY_YEAR = 2023
POLICY_OBJECTIVE = (
    "Modernised, foundational environmental framework law replacing Loi n° 2001-01 (2001). Like its "
    "predecessor, it never names 'plastics' or 'synthetic polymers' anywhere in its 255 articles, but it "
    "substantially strengthens the general legal architecture that governs plastics indirectly: (i) it "
    "codifies the 'extended producer responsibility' principle and the waste hierarchy (réduction, "
    "réutilisation, recyclage, valorisation énergétique, élimination) directly in the general-principles "
    "and waste chapters (Art. 5, Art. 69) -- the legal basis most relevant to plastics circularity; (ii) "
    "it creates two ring-fenced funds, a general 'Fonds spécial pour la protection de l'Environnement' "
    "(Art. 18) and a dedicated marine/coastal fund, the 'Fonds national de Prévention et de Protection de "
    "l'Environnement marin et côtier' (FN-PEM, Art. 19) -- directly relevant given plastics dominate marine "
    "litter; (iii) it adds a mandatory periodic environmental-audit regime (Art. 42-43) on top of the "
    "existing EIA regime; (iv) it imposes a new mandatory duty on every household to sort waste at source "
    "('tri à la source', Art. 89) and requires every municipality to adopt a household-waste management "
    "plan (Art. 87); and (v) it criminalises the use of recycled waste-derived materials in the manufacture "
    "of food-contact packaging (Art. 210) -- a food-safety-driven restriction that constrains how recycled "
    "plastic can be used in packaging. There is still no plastic-specific reduction, recycling or EPR "
    "target in this text; Senegal's plastics-specific obligations remain in separate legislation (see "
    "comments)."
)
POLICY_TARGET = 0
POLICY_TARGET_TEXT = ""
POLICY_TYPE = 1
POLICY_TYPE_JUSTIFICATION = (
    "Full Act of Parliament: 'L'Assemblée nationale a adopté, en sa séance du mercredi 07 juin 2023, Le "
    "Conseil Constitutionnel ayant statué par sa décision n° 5/C/2023 du 12 juillet 2023, Le Président de "
    "la République promulgue la loi dont la teneur suit.' Adopted by the National Assembly, reviewed by "
    "the Constitutional Council, and promulgated by the President -> legislation, not an executive/"
    "ministerial instrument."
)
POLICY_INTEGRATION = 1
POLICY_SECTORS_LIST = (
    "environment, industry, agriculture, health, fisheries, water resources, mining, oil & gas/energy, "
    "maritime transport, road/land transport, urban planning & construction, commerce/trade, "
    "local government/municipalities, waste management, chemicals"
)
POLICY_CIRCULARITY = 1
POLICY_LIFECYCLE_PHASES_LIST = "production, consumption, recycling, disposal, environmental leakage"
POLICY_BUDGET = 1
POLICY_BUDGET_TEXT = (
    "Art. 16: 'Pour la protection de l'Environnement, il est institué des droits, redevances et taxes "
    "parafiscales supportés par les exploitants des installations classées ou toute personne qui mène une "
    "activité réglementée à incidence environnementale.' Art. 18: 'Il est créé un Fonds spécial pour la "
    "protection de l'Environnement.' Art. 19: 'Pour le suivi de l'Environnement marin et côtier, il est "
    "créé un Fonds spécial dénommé \"Fonds national de Prévention et de Protection de l'Environnement "
    "marin et côtier (FN-PEM)\".' Two dedicated, ring-fenced funds financed by the Code's own parafiscal "
    "levies."
)
POLICY_SCORE_FORMULA = "[auto] O = (G + I + K + M) / 4"

policy_fields = dict(
    A_policy_name=POLICY_NAME,
    B_policy_url=POLICY_URL,
    C_policy_year=POLICY_YEAR,
    D_policy_objective=POLICY_OBJECTIVE,
    E_policy_target=POLICY_TARGET,
    F_policy_target_text=POLICY_TARGET_TEXT,
    G_policy_type=POLICY_TYPE,
    H_policy_type_justification=POLICY_TYPE_JUSTIFICATION,
    I_policy_integration=POLICY_INTEGRATION,
    J_policy_sectors_list=POLICY_SECTORS_LIST,
    K_policy_circularity=POLICY_CIRCULARITY,
    L_policy_lifecycle_phases_list=POLICY_LIFECYCLE_PHASES_LIST,
    M_policy_budget=POLICY_BUDGET,
    N_policy_budget_text=POLICY_BUDGET_TEXT,
    O_policy_score=POLICY_SCORE_FORMULA,
)

# ---------------------------------------------------------------------------
# Section B - Instrument-level fields (one dict per row)
# ---------------------------------------------------------------------------

instruments = [
    dict(
        P_instrument_type=1.0,
        Q_instrument_lifecycle_stage="Production",
        R_instrument_description=(
            "Art. 44-62: 'Installations classées pour la protection de l'Environnement' (ICPE). Any "
            "factory, workshop, depot, worksite, mine/quarry or oil & gas installation (incl. plastics "
            "manufacturing/processing/recycling plants) presenting a danger is divided into two classes "
            "(Art. 45-46) requiring prior authorisation (Class 1) or declaration (Class 2); Class-1 "
            "facilities require an environmental-conformity certificate before construction (Art. 48); the "
            "Minister of Environment can order closure of non-compliant facilities (Art. 58)."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: 'Le Ministre chargé de l'Environnement' issues authorisations, prescriptions and "
            "closure orders (Art. 47-58) (+0.25). Enforcement: criminal penalties for operating without "
            "authorisation/declaration (Art. 189-190: up to 25,000,000 FCFA + 5 years) and for unauthorised "
            "modification (Art. 192-193) (+0.25). Monitoring: 'L'inspection des installations classées...est "
            "assurée par des agents assermentés habilités par le Ministre chargé de l'Environnement' (Art. "
            "61) (+0.25). Not unconditional: Art. 248 grants existing operators a 6-month transitional grace "
            "period to comply, so the +0.25 unconditional sub-score is not awarded."
        ),
        W_comments=(
            "Directly comparable to the 2001 Code's Titre II Ch.1 regime (Art. L9-L24), now with an added "
            "environmental-conformity certificate step (Art. 48) and an explicit oil & gas scope (Art. 44) "
            "reflecting the law's stated objective of covering offshore hydrocarbon activity."
        ),
    ),
    dict(
        P_instrument_type=0.60,
        Q_instrument_lifecycle_stage="Waste management",
        R_instrument_description=(
            "Art. 16-19: parafiscal 'droits, redevances et taxes' on classified installations and other "
            "regulated activities with environmental impact (Art. 16); a 3-year, non-renewable tax "
            "exoneration for firms investing in anti-pollution activities (Art. 17); a new general ring-"
            "fenced 'Fonds spécial pour la protection de l'Environnement' (Art. 18); and a second, dedicated "
            "'Fonds national de Prévention et de Protection de l'Environnement marin et côtier (FN-PEM)' "
            "for marine/coastal environmental monitoring (Art. 19) -- a marine fund is new relative to the "
            "2001 Code and is directly relevant given plastics dominate marine litter."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.50,
        U_instrument_implementation_text=(
            "Authority: taxes/funds administered by the 'Ministre chargé de l'Environnement' jointly with "
            "the 'Ministre chargé des Finances' for exemptions (Art. 17) (+0.25). Monitoring: the related "
            "pollution-charge ('taxe à la pollution') calculation methodology elsewhere in the Code is based "
            "on accredited-laboratory sampling and analysis of discharges (Art. 149-150) (+0.25). "
            "Enforcement: no explicit fine/surcharge for non-payment of the Art. 16 parafiscal taxes is "
            "stated in this chapter (unlike the 2001 Code's Art. L25 10% late-payment surcharge, which was "
            "not clearly re-stated here) -- not awarded. Unconditional: Art. 17 creates a conditional tax "
            "exoneration regime -- not awarded."
        ),
        W_comments=(
            "The FN-PEM marine/coastal fund (Art. 19) is a genuinely new instrument versus the 2001 Code and "
            "is highlighted because ocean/coastal plastic leakage is a core 4P Index concern. The decree(s) "
            "implementing the funds' 'modalités d'alimentation, de gestion et d'utilisation' (Art. 18-19) "
            "had not been identified/verified at the time of this coding -- flagged as an uncertainty."
        ),
    ),
    dict(
        P_instrument_type=0.40,
        Q_instrument_lifecycle_stage="Production",
        R_instrument_description=(
            "Art. 20-38 (Titre III): mandatory environmental evaluation regime -- 'évaluation environnementale "
            "stratégique' for sectoral policies/plans/programmes (Art. 30-32), 'étude d'impact environnemental "
            "et social' (EIES) for projects/activities likely to endanger the environment or human health "
            "(Art. 33-38), and 'analyse environnementale initiale' for lower-risk projects (Art. 39-41); a "
            "'Comité technique de Validation et de Suivi des Evaluations environnementales' reviews and "
            "validates these evaluations (Art. 14, Art. 25); public participation ('audience publique') is "
            "mandatory (Art. 24)."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: 'Comité technique de Validation et de Suivi des Evaluations environnementales' "
            "(Art. 14) and the Ministère chargé de l'Environnement (Art. 22, 25, 35) (+0.25). Enforcement: "
            "Art. 206 criminalises implementing a project without the required environmental evaluation, or "
            "non-compliant with its criteria/measures (5-10 years' imprisonment + 100,000,000-500,000,000 "
            "FCFA) (+0.25). Monitoring: evaluations must be carried out by Ministry-accredited firms "
            "('bureaux agréés', Art. 22) and validated by the Comité technique (Art. 25) (+0.25). "
            "Unconditional: not awarded -- the specific evaluation type required depends on categories/"
            "thresholds set by decree (Art. 28, Art. 40)."
        ),
        W_comments=(
            "Coded as Procedural (0.40, 'planning requirements') consistent with the 2001-Code EIA row for "
            "comparability; a reasonable alternative view would score it higher given the binding, "
            "approval-blocking effect and the Art. 206 criminal backstop -- flagged as a borderline case, "
            "as in the 2001 coding."
        ),
    ),
    dict(
        P_instrument_type=0.40,
        Q_instrument_lifecycle_stage="Waste management",
        R_instrument_description=(
            "Art. 42-43 (new Chapitre V, 'De l'audit environnemental'): projects and classified installations "
            "already in operation must periodically undergo an environmental audit -- 'pour certaines "
            "transformations/activités/opérations, la mise à niveau et la fin du projet' -- to verify "
            "ongoing compliance ('conformes à la réglementation environnementale'), in conditions fixed by "
            "decree."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.50,
        U_instrument_implementation_text=(
            "Authority: falls under the same Ministry/Comité technique oversight framework as the "
            "evaluation regime generally (+0.25). Monitoring: the audit is itself defined as a 'évaluation "
            "périodique et systématique, documentée et objective...de la performance des équipements' (Art. "
            "3 definition of 'audit environnemental') -- an explicit, self-contained monitoring mechanism "
            "(+0.25). Enforcement: no specific penal article requiring/sanctioning failure to conduct the "
            "audit was identified -- not awarded. Unconditional: the audit only applies to installations "
            "undergoing 'certaines transformations/activités/opérations' -- conditional, not universal -- "
            "not awarded."
        ),
        W_comments=(
            "This mandatory-audit chapter is entirely new relative to the 2001 Code (which had no equivalent "
            "'audit environnemental' title) and represents a genuine strengthening of ongoing-compliance "
            "monitoring for facilities that may process or recycle plastics."
        ),
    ),
    dict(
        P_instrument_type=1.0,
        Q_instrument_lifecycle_stage="Recycling",
        R_instrument_description=(
            "Art. 5, 65-71: the waste chapter now explicitly codifies the 'principe de la responsabilité "
            "élargie du producteur' (EPR) among the Code's general principles (Art. 5) and the 'hiérarchie "
            "des modes de traitement des déchets' -- réutilisation, recyclage, valorisation énergétique, "
            "élimination, in that order (Art. 69) -- among the principles of 'gestion écologiquement "
            "rationnelle des déchets'. Producers must reduce/prevent waste generation through clean "
            "production and eco-design (Art. 67); every waste holder must manage their waste 'en respectant "
            "l'ordre de priorité de traitement' (Art. 70); waste elimination/treatment requires prior "
            "Ministerial authorisation (Art. 71)."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: 'l'autorisation préalable du Ministre chargé de l'Environnement' governs elimination/"
            "treatment (Art. 71) (+0.25). Enforcement: Art. 71 provides for a formal notice ('mise en "
            "demeure') and mandatory deposit at the 'Caisse des Dépôts et Consignations' for remediation "
            "costs, and Art. 191 criminalises refusal to make that deposit (1,000,000-5,000,000 FCFA fine) "
            "(+0.25). Monitoring: waste treatment/elimination facilities are themselves classified "
            "installations subject to the sworn-inspector regime of Art. 61 (+0.25). Unconditional: not "
            "awarded -- Art. 70 al. 2 explicitly carves out household/assimilated waste from the priority-"
            "order duty and allows the order to be varied 'dans des conditions particulières'."
        ),
        W_comments=(
            "This is the single most important general-purpose instrument for plastics circularity in the "
            "Code: the explicit EPR principle (Art. 5) is the same legal concept underpinning the EPR "
            "obligations for plastic producers already created by the freestanding Loi n° 2020-04 (2020); "
            "here it is elevated to a codified general principle of environmental law rather than a "
            "sector-specific rule. Mixed nature per Rule 10: combines a regulatory duty (Art. 70) with a "
            "procedural principle-setting function (Art. 5, 69); scored at the higher Regulatory value."
        ),
    ),
    dict(
        P_instrument_type=1.0,
        Q_instrument_lifecycle_stage="End of life",
        R_instrument_description=(
            "Art. 73: 'Le brûlage à l'air libre des déchets, à l'exception du brulis, est interdit. Le "
            "brûlage à l'air libre des pneus usagés est interdit.' (Open-air burning of waste and of used "
            "tyres is prohibited, with a narrow exception for traditional agricultural stubble-burning "
            "('brulis') and for burning operations separately authorised by decree.)"
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.50,
        U_instrument_implementation_text=(
            "Authority: enforced by sworn environmental agents under the Titre VI framework (Art. 229-235) "
            "(+0.25). Enforcement: Art. 217 criminalises burying or burning waste or used tyres in violation "
            "of the Code (1 month-1 year imprisonment + 1,000,000-5,000,000 FCFA) (+0.25). Monitoring: no "
            "specific inspection/data-collection mechanism for open burning was identified beyond the "
            "general agent framework -- not awarded. Unconditional: not awarded -- 'certaines opérations de "
            "brûlage peuvent être autorisées dans des conditions prévues par décret', and the 'brulis' "
            "exception is explicit."
        ),
        W_comments=(
            "Highly plastics-relevant in practice: open burning of household/municipal waste (which in "
            "Senegal, as elsewhere, is dominated by plastic packaging and bags) is a major source of toxic "
            "emissions and is explicitly banned here."
        ),
    ),
    dict(
        P_instrument_type=1.0,
        Q_instrument_lifecycle_stage="End of life",
        R_instrument_description=(
            "Absolute prohibitions on uncontrolled disposal: Art. 72 bans immersion/incineration/elimination "
            "of waste in inland, marine or estuarine waters under Senegalese jurisdiction; Art. 80 bans "
            "import of hazardous waste (other than radioactive) onto Senegalese territory, subject to a "
            "conditional exception for prior-consent imports from African countries with adequate treatment "
            "technology; Art. 83 bans discharge of hazardous waste into waters under Senegalese "
            "jurisdiction; Art. 84 bans storage of hazardous waste outside designated sites and its burial "
            "in soil/subsoil; Art. 91 states 'Est interdit le dépôt des déchets sur le domaine public.'"
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: Ministère chargé de l'Environnement / collectivités territoriales enforce (Art. 92) "
            "(+0.25). Enforcement: Art. 197-198 (illegal hazardous-waste discharge/import: 5-10 years + "
            "100,000,000 FCFA-1 billion FCFA) and Art. 218 (illegal hazardous-waste storage: 1-2 years + "
            "5,000,000-10,000,000 FCFA) (+0.25). Monitoring: sworn agents/officers empowered under Art. "
            "229-235 to inspect, seize and draw up official reports (+0.25). Unconditional: not fully "
            "awarded -- Art. 80 explicitly carves out a conditional exception for imports from African "
            "countries with adequate local treatment technology."
        ),
        W_comments=(
            "Directly relevant to plastic litter/marine leakage even though 'plastic' is not named: illegal "
            "dumping of packaging/bags on public land or into water is the paradigmatic offence captured "
            "here. Cross-reference: Senegal's plastics-specific bans (Loi n° 2015-09 of 2015, replaced by "
            "Loi n° 2020-04 of 2020) operate alongside this general prohibition regime rather than through "
            "it."
        ),
    ),
    dict(
        P_instrument_type=0.40,
        Q_instrument_lifecycle_stage="Waste management",
        R_instrument_description=(
            "Art. 87: 'Chaque collectivité territoriale ou groupements constitués élabore, en rapport avec "
            "les services techniques de l'Environnement, un plan communal de prévention et de gestion des "
            "déchets ménagers et assimilés.' Modalities for drafting/approving the plan are set by joint "
            "order of the Ministers of Local Government and Environment; Art. 88 ties approval of new waste-"
            "treatment facilities to consistency with the applicable zonal plan."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.50,
        U_instrument_implementation_text=(
            "Authority: joint ministerial order (Ministre chargé des Collectivités territoriales + Ministre "
            "chargé de l'Environnement) governs the plan's modalities (Art. 87) (+0.25). Unconditional: the "
            "duty to elaborate a plan is stated for every local authority with no exemption identified "
            "(+0.25). Enforcement: no penal or administrative sanction specifically tied to a municipality's "
            "failure to adopt a plan was identified -- not awarded. Monitoring: no dedicated inspection/"
            "reporting mechanism for plan implementation itself (as distinct from Art. 88's facility-"
            "approval linkage, which is a planning-coherence rule rather than monitoring) -- not awarded."
        ),
        W_comments=(
            "Coded per Rule 11 as a multi-level governance/coordination instrument -- a textbook example of "
            "'waste management planning obligations imposed on sub-national authorities'. New relative to "
            "the 2001 Code, which only assigned a general household-waste elimination duty to collectivités "
            "locales (Art. L32) without a formal planning requirement."
        ),
    ),
    dict(
        P_instrument_type=1.0,
        Q_instrument_lifecycle_stage="Recycling",
        R_instrument_description=(
            "Art. 89: 'Chaque ménage est tenu d'assurer le tri à la source et le conditionnement de ses "
            "déchets, conformément à la réglementation en vigueur.' A new, mandatory source-separation duty "
            "imposed directly on every household -- a foundational enabler of household plastic-waste "
            "recycling that had no equivalent in the 2001 Code."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.25,
        U_instrument_implementation_text=(
            "Authority: municipalities are responsible for household waste generally (Art. 85) (+0.25). "
            "Enforcement: no specific fine or sanction for a household's failure to sort waste at source was "
            "identified in the penal Titre VI -- not awarded. Monitoring: no inspection/data-collection "
            "mechanism at household level is specified -- not awarded. Unconditional: the duty is expressly "
            "made contingent on further implementing regulation ('conformément à la réglementation en "
            "vigueur') which was not identified/verified as adopted -- not awarded."
        ),
        W_comments=(
            "In-force test (Rule 12): 'est tenu d'assurer' is mandatory language, so S=1, but weak "
            "implementation evidence (T=0.25) reflects the absence of any identified enforcement mechanism "
            "for household-level non-compliance -- flagged as a likely practical gap between the law's text "
            "and its enforceability."
        ),
    ),
    dict(
        P_instrument_type=1.0,
        Q_instrument_lifecycle_stage="Recycling",
        R_instrument_description=(
            "Art. 210: 'Est puni d'un emprisonnement d'un (01) mois à un (01) an et d'une amende de cinq "
            "cent mille (500 000) à un million (1 000 000) de francs CFA...toute personne qui utilise les "
            "produits du recyclage issus des déchets dans la fabrication des contenants des produits "
            "alimentaires.' A direct criminal prohibition on using recycled waste-derived materials to "
            "manufacture food-contact containers/packaging."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: enforced by sworn environmental/judicial agents under Titre VI (Art. 229-235) "
            "(+0.25). Enforcement: the criminal penalty is stated directly in the same article (+0.25). "
            "Monitoring: no dedicated inspection/testing regime for food-contact packaging composition was "
            "identified -- not awarded. Unconditional: the prohibition is stated categorically with no "
            "exemption identified (+0.25)."
        ),
        W_comments=(
            "IMPORTANT for plastics circularity coding: unlike most instruments in this table, this "
            "provision is a food-safety-driven RESTRICTION on recycled content rather than an incentive for "
            "recycling -- it directly limits the use of recycled (often plastic) materials in food-contact "
            "packaging manufacture, in tension with circular-economy objectives elsewhere in the Code (Art. "
            "5, 69). Flagged as a genuinely borderline/dual-nature case worth separate analytical attention "
            "in any cross-country 4P Index comparison of recycled-content rules."
        ),
    ),
    dict(
        P_instrument_type=0.40,
        Q_instrument_lifecycle_stage="Waste management",
        R_instrument_description=(
            "Art. 6, 13-15: multi-level governance/institutional coordination framework. The transfer of "
            "environment/natural-resource competencies to 'collectivités territoriales' does not displace "
            "State civil-defence powers (Art. 6); the Minister of Environment implements national policy "
            "and creates national monitoring committees for international environmental agreements, plus an "
            "'Organe national de Suivi de l'Environnement marin et côtier' (Art. 13); sectoral policies/"
            "plans/programmes must integrate environmentally sound management (Art. 15)."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.50,
        U_instrument_implementation_text=(
            "Authority: 'Ministre chargé de l'Environnement' explicitly designated (Art. 13) (+0.25). "
            "Monitoring: national follow-up committees and the marine/coastal monitoring body are explicitly "
            "created institutions with a monitoring mandate (Art. 13) (+0.25). Enforcement: no specific "
            "sanction mechanism tied to this coordination framework itself -- not awarded. Unconditional: "
            "Art. 6 explicitly conditions the competency transfer on the State's reserved civil-defence "
            "powers -- not awarded."
        ),
        W_comments=(
            "Coded per Rule 11 as a coordination/institutional-assignment instrument. The dedicated marine/"
            "coastal monitoring body (Art. 13) is new relative to the 2001 Code's simpler Secrétariat "
            "Permanent arrangement and is a positive development for tracking marine plastic pollution, "
            "though its concrete mandate/resourcing is fixed by a decree not independently verified here."
        ),
    ),
    dict(
        P_instrument_type=0.40,
        Q_instrument_lifecycle_stage="Environmental leakage",
        R_instrument_description=(
            "Art. 94-99: Class-1 (autorisation) classified-installation operators must establish an internal "
            "emergency-response plan ('plan d'opération interne', Art. 94); Class-2 (déclaration) operators "
            "may be required to do the same by order (Art. 95, an enabling power); a national environmental-"
            "emergency plan framework is established, linked to danger studies and triggered by specified "
            "risk thresholds (Art. 96); a dedicated emergency hotline and response service, 'Urgences "
            "Environnement' (call number '1221'), is created with priority vehicles (Art. 99)."
        ),
        S_instrument_in_force=1,
        T_instrument_implementation=0.75,
        U_instrument_implementation_text=(
            "Authority: joint order of the Ministers of Environment, Interior and Industry governs plan "
            "content/implementation (Art. 98) (+0.25). Enforcement: Art. 207-208 criminalise operating an "
            "authorised/declared installation in breach of the internal-operations-plan or particular-"
            "intervention-plan requirements (up to 15,000,000 FCFA + 2 years) (+0.25). Unconditional: the "
            "Class-1 internal-plan duty (Art. 94) is stated as mandatory ('est tenu d'établir') with no "
            "exemption identified (+0.25). Monitoring: no dedicated inspection/audit mechanism specific to "
            "emergency plans (as distinct from the general ICPE inspection regime) was identified -- not "
            "awarded."
        ),
        W_comments=(
            "The Class-2 extension in Art. 95 ('peut...être tenu') remains a separate enabling power and "
            "would be coded 0 (not in force) unless activated by ministerial order; only the mandatory "
            "Class-1 duty and the national emergency-plan/hotline framework are scored as in force here. "
            "Implementation is scored higher (0.75) than the equivalent 2001-Code provision (0.50) because "
            "the 2023 Code adds an explicit, on-point criminal penalty (Art. 207-208) that was not clearly "
            "present in the 2001 text."
        ),
    ),
]

# ---------------------------------------------------------------------------
# Assemble rows, computing auto-fields O and V
# ---------------------------------------------------------------------------

O_score = round(
    (policy_fields["G_policy_type"] + policy_fields["I_policy_integration"]
     + policy_fields["K_policy_circularity"] + policy_fields["M_policy_budget"]) / 4,
    3,
)
policy_fields["O_policy_score"] = f"{O_score} [auto = (G+I+K+M)/4]"

rows = []
for inst in instruments:
    row = dict(policy_fields)
    row.update(inst)
    s = row["S_instrument_in_force"]
    p = row["P_instrument_type"]
    t = row["T_instrument_implementation"]
    v = round(s * (p + t) / 2, 3)
    row["V_instrument_score"] = f"{v} [auto = S*(P+T)/2]"
    rows.append(row)

with open("4p_index/senegal_loi_2023-15_4p_index.csv", "w", newline="", encoding="utf-8") as f:
    writer = csv.DictWriter(f, fieldnames=COLUMNS)
    writer.writeheader()
    for row in rows:
        writer.writerow(row)

df = pd.DataFrame(rows, columns=COLUMNS)
xlsx_path = "4p_index/senegal_loi_2023-15_4p_index.xlsx"
df.to_excel(xlsx_path, sheet_name="4P_Index_Loi2023-15", index=False)

wb = load_workbook(xlsx_path)
ws = wb["4P_Index_Loi2023-15"]

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF")
wrap_top = Alignment(wrap_text=True, vertical="top")

col_widths = {
    "A": 34, "B": 30, "C": 10, "D": 46, "E": 10, "F": 30, "G": 10, "H": 34,
    "I": 10, "J": 30, "K": 10, "L": 26, "M": 10, "N": 30, "O": 16,
    "P": 12, "Q": 16, "R": 46, "S": 10, "T": 12, "U": 40, "V": 16, "W": 40,
}
for letter in [get_column_letter(c) for c in range(1, len(COLUMNS) + 1)]:
    ws.column_dimensions[letter].width = col_widths.get(letter, 20)

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for cell in row:
        cell.alignment = wrap_top

ws.freeze_panes = "A2"
ws.row_dimensions[1].height = 30
for r in range(2, ws.max_row + 1):
    ws.row_dimensions[r].height = 120

wb.save(xlsx_path)
print("Wrote", xlsx_path, "and CSV with", len(rows), "instrument rows.")
print("Policy score O =", O_score)
for row in rows:
    print(row["Q_instrument_lifecycle_stage"], "->", row["V_instrument_score"])
