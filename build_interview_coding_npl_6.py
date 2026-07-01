"""Generate interview coding spreadsheet for Nepal interview NPL_6.

Coding follows the Overview All Interviews codebook and is based on the
Swasti Byanju (hotel owner, Dhulikhel) interview of 22 June 2025 — transcript,
interview notes, and guideline. Uses expanded Impacts taxonomy.
"""

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUTPUT_PATH = "/workspace/interview_coding_NPL_6.xlsx"

DARK_BLUE = "1F3864"
MID_BLUE = "2E75B6"
LIGHT_BLUE = "D9E2F3"
LIGHT_GREEN = "E2EFDA"
ORANGE = "FCE4D6"
WHITE = "FFFFFF"

SOURCES = [
  ("Interview transcript", "Restaurant/Hotel Owner — 22 June 2025 (Part 1 ENG)"),
  ("Interview notes", "Kathmandu interview no. 3 — structured guideline notes"),
  ("Interview guideline", "22 June 2025 — Swasti Byanju, Dhulikhel"),
  ("Codebook", "Overview All Interviews — NEW Codebook (expanded Impacts)"),
]

STAKEHOLDER = {
  "Interview ID": "NPL_6",
  "Country": "NEPAL",
  "Interview number": "3 (Kathmandu series)",
  "Interview date": "22 June 2025 (Sunday)",
  "Interviewee": "Mr. Swasti Byanju",
  "Affiliation / role": "Hotel owner — Dhulikhel",
  "Stakeholder list mapping": "hotel_owner",
  "Coded actor type (codebook)": "hotel_owner",
  "Actor classification rationale": (
    "Hotel/restaurant business owner in Dhulikhel managing commercial plastic "
    "waste (shampoo sachets, packaging) alongside personal consumption. "
    "Distinct from shop_owner (retail) and household."
  ),
  "Perspective": (
    "Business + community member — personally reduces bag use but depends on "
    "municipal collection; advocates strict bans, budget allocation, and "
    "Dankhuta/Chitwan replication models"
  ),
  "Key local references": (
    "Dhulikhel dumping site (15 years), Panchkhal downstream conflict, "
    "ward 7 shopkeeper litter rule, Dankhuta 'Waste into Money' model, "
    "Chitwan recycling unit"
  ),
}

CODING = {
  "Country": "NPL",
  "Country name": "NEPAL",
  "ID": "NPL_6",
  "Actortype": "hotel_owner",
  "Problem_awareness_pop": "medium",
  "Problem_awareness_pol": "high",
  "Problem_severity": "high",
  "Problem_littering": "yes",
  "Problem_consumption": "yes",
  "Problem_recycling": "yes",
  "Problem_waste_mgmt": "yes",
  "Problem_production": "yes",
  "Problem_alternatives": "yes",
  "Problem_waste_segregation": "yes",
  "Problem_import": "no",
  "Impacts": (
    "water, rivers, cities, health, agriculture, soil, foul smell, dumping sites, "
    "visual pollution, aquatic life, wildlife, marine pollution"
  ),
  "Governance": "yes",
  "Relevance_international_pol": "no",
  "Coordination_sectoral": "yes",
  "Coordination_levels": "yes",
  "Unclear_responsibilities": "yes",
  "Actors": "yes",
  "Federal government": "yes",
  "Provincial government": "no",
  "Local government": "yes",
  "Students": "yes",
  "Private_sector": "yes",
  "Civil_society": "no",
  "Science": "no",
  "Households": "yes",
  "Education institutions": "yes",
  "Private_companies(hotels, shops, etc.)": "yes",
  "Implementation issues": "yes",
  "Monitoring": "yes",
  "Financial_resources": "yes",
  "Research": "no",
  "Infrastructure": "yes",
  "Enforcement": "yes",
  "Policies in place": "yes",
  "Pol_epr": "no",
  "Pol_import": "no",
  "Pol_awarness": "yes",
  "Pol_education": "yes",
  "Pol_capacity": "no",
  "Pol_RD": "no",
  "Pol_tax": "no",
  "Pol_ban": "yes",
  "Pol_subsitutes": "yes",
  "Pol_clean_up": "yes",
  "Pol_upcycling": "no",
  "Pol_recycling": "yes",
  "Pol_waste_collection": "yes",
  "Solutions": "yes",
  "Sol_lead_agency": "yes",
  "Sol_responsibilities": "yes",
  "Sol_epr": "no",
  "Sol_awarness": "yes",
  "Sol_segregation": "yes",
  "Sol_upcycling": "no",
  "Sol_recycling": "yes",
  "Sol_education": "yes",
  "Sol_capacity": "yes",
  "Sol_RD": "no",
  "Sol_tax": "no",
  "Sol_ban": "yes",
  "Sol_finance": "yes",
  "Sol_infrastructure": "yes",
  "Sol_subsitutes": "yes",
  "Sol_clean_up": "yes",
  "Sol_enforcement": "yes",
  "Sol_monitoring": "yes",
  "Traditions to build on (free-hand)": (
    "Thokodi (newspaper paper bags sold in shops); grandfathers' handmade paper "
    "bags; cloth bags; household bag collection and reuse before recycling"
  ),
}

CODEBOOK = [
  ("Country", "Country ISO code"),
  ("Country name", "Country name"),
  ("ID", "InterviewID: ISO_Nr"),
  ("Actortype", "governmental, pol_party, science, cso, igo, private_sector, shop_owner, hotel_owner, household"),
  ("Problem_awareness_pop", "high, medium, low, NA — lack of awareness among population mentioned"),
  ("Problem_awareness_pol", "high, medium, low, NA — lack of awareness among policy makers mentioned"),
  ("Problem_severity", "high, medium, low, NA"),
  ("Problem_littering", "Littering is a problem; was mentioned: yes/no"),
  ("Problem_consumption", "High consumption is a problem; was mentioned: yes/no"),
  ("Problem_recycling", "No or too little recycling is a problem; was mentioned: yes/no"),
  ("Problem_waste_mgmt", "Ill waste management is a problem; was mentioned: yes/no"),
  ("Problem_production", "High production rates; was mentioned: yes/no"),
  ("Problem_alternatives", "No good alternatives; was mentioned: yes/no"),
  ("Problem_waste_segregation", "Lack of segregation was mentioned: yes/no"),
  ("Problem_import", "Plastic imports are a problem; was mentioned: yes/no"),
  ("Impacts", "Impacts mentioned on water, cities, biodiversity, health, etc. or NA"),
  ("Governance", "Mentioned: yes/no"),
  ("Relevance_international_pol", "International treaties are relevant for national level policy"),
  ("Coordination_sectoral", "Lack of coordination across sectors"),
  ("Coordination_levels", "Lack of coordination across levels"),
  ("Unclear_responsibilities", "Unclear responsibilities among the involved actors"),
  ("Actors", "Mentioned: yes/no"),
  ("Federal government", "Mentioned: yes/no"),
  ("Provincial government", "Mentioned: yes/no"),
  ("Local government", "Mentioned: yes/no"),
  ("Students", "Mentioned: yes/no"),
  ("Private_sector", "Mentioned: yes/no"),
  ("Civil_society", "Mentioned: yes/no"),
  ("Science", "Mentioned: yes/no"),
  ("Households", "Mentioned: yes/no"),
  ("Education institutions", "Mentioned: yes/no"),
  ("Private_companies(hotels, shops, etc.)", "Mentioned: yes/no"),
  ("Implementation issues", "Mentioned: yes/no"),
  ("Monitoring", "Lack of monitoring"),
  ("Financial_resources", "Lack of financial resources"),
  ("Research", "Lack of research"),
  ("Infrastructure", "Lack of infrastructure"),
  ("Enforcement", "Lack of enforcement"),
  ("Policies in place", "Mentioned: yes/no"),
  ("Pol_epr", "Extended producer responsibility"),
  ("Pol_import", "Import regulations"),
  ("Pol_awarness", "Awareness campaigns"),
  ("Pol_education", "Education programs"),
  ("Pol_capacity", "Capacity building"),
  ("Pol_RD", "Research & development"),
  ("Pol_tax", "Levies or taxes on specific products (often bags)"),
  ("Pol_ban", "Bans of specific products"),
  ("Pol_subsitutes", "Alternatives like reusable bags, or banana leaf plates"),
  ("Pol_clean_up", "Clean-up campaigns"),
  ("Pol_upcycling", "Making a new product, like blacktop"),
  ("Pol_recycling", "Making a new raw material"),
  ("Pol_waste_collection", "Waste collection"),
  ("Solutions", "Mentioned: yes/no"),
  ("Sol_lead_agency", "Procedural measures to establish lead agency"),
  ("Sol_responsibilities", "Clear responsibilities"),
  ("Sol_epr", "Extended producer responsibility"),
  ("Sol_awarness", "Awareness raising measures"),
  ("Sol_segregation", "Segregation of waste"),
  ("Sol_upcycling", "Upcycling of plastics"),
  ("Sol_recycling", "New raw materials"),
  ("Sol_education", "Education programs at schools"),
  ("Sol_capacity", "Capacity-building programs"),
  ("Sol_RD", "Research programs"),
  ("Sol_tax", "Taxes or levies on products"),
  ("Sol_ban", "Ban of products"),
  ("Sol_finance", "Financial measures"),
  ("Sol_infrastructure", "Infrastructural measures"),
  ("Sol_subsitutes", "Alternatives"),
  ("Sol_clean_up", "Clean-up campaigns"),
  ("Sol_enforcement", "Enforcement procedures"),
  ("Sol_monitoring", "Monitoring procedures"),
  ("Traditions to build on (free-hand)", "Freehand or NA"),
]

NOTES = {
  "Actortype": (
    "Swasti Byanju — hotel owner, Dhulikhel. Stakeholder list: hotel_owner."
  ),
  "Problem_awareness_pop": (
    "People aware of pollution but do not act; behavioral change lacking; "
    "lack of practical awareness on how to manage beyond bags."
  ),
  "Problem_awareness_pol": (
    "No certain policy; national ban notice not implemented; budget not "
    "allocated; projects announced but no follow-up for years."
  ),
  "Problem_severity": (
    "Major problem; very concerned personally and for business; visible "
    "everywhere despite nobody liking plastic."
  ),
  "Problem_littering": (
    "Plastic visible everywhere on streets; litter in front of shops (ward 7 rule)."
  ),
  "Problem_consumption": (
    "Packaged goods in remote villages; shampoo sachets; clothes, electronics, "
    "packaging beyond bags — plastic in every daily-life step."
  ),
  "Problem_recycling": (
    "Recycling almost zero in Dhulikhel; only reusables picked at dump; contrast "
    "with Chitwan 10-truck unit and Dankhuta model."
  ),
  "Problem_waste_mgmt": (
    "No proper dumping sites; 15-year local dump; hospital/KU waste mixed in; "
    "vegetables still sold in plastic despite ban discussions."
  ),
  "Problem_production": "Widespread plastic use in all shops and daily consumption.",
  "Problem_alternatives": (
    "Thokodi/paper bags largely disappeared; need paper/cloth bag alternatives."
  ),
  "Problem_waste_segregation": (
    "No household or hotel segregation before municipal handover; key barrier."
  ),
  "Impacts": (
    "Water: drinking water affected wards 2–3; Panchkhal groundwater not potable. "
    "Rivers: tractor dumping Khurkot/Sindhuli scenic route. Agriculture/soil: "
    "leachate wards 2–3, downstream Panchkhal. Health: hospital needles/syringes "
    "at dump. Foul smell: daily on route to municipality/Panchkhal. Dumping sites: "
    "15-year accumulation. Visual pollution: litter, birds at waste, aesthetics. "
    "Aquatic life/marine pollution: heard ocean plastic India, fish problems. "
    "Wildlife: birds taking food from waste. Cities: street litter, inter-municipal conflict."
  ),
  "Coordination_sectoral": (
    "Waste collectors important but underintegrated; private recyclers need "
    "municipal coordination (Chitwan model)."
  ),
  "Coordination_levels": (
    "Central budget not reaching local; national policy without local implementation; "
    "Dhulikhel–Panchkhal inter-municipal conflict."
  ),
  "Unclear_responsibilities": (
    "Roles known in principle (mayor/municipality + citizens) but weak coordination "
    "and implementation; respondent unaware of specific national policies."
  ),
  "Federal government": (
    "Ministry of Forest & Environment / Department of Environment — policy, "
    "plastic bank, national ban notice."
  ),
  "Local government": (
    "Mayor holds budget and policy; Mirvan post/municipal collection; ward 7 "
    "shopkeeper litter rule."
  ),
  "Students": "KU students participated in Beat the Plastic (World Environment Day).",
  "Private_sector": (
    "Waste collectors/recyclers essential; Chitwan recycling unit; Dankhuta "
    "segregation buildings."
  ),
  "Households": (
    "Individual responsibility — alternative bags, home bag collection/reuse; "
    "housewives targeted for awareness."
  ),
  "Education institutions": "Kathmandu University students in campaigns; KU hospital waste at dump.",
  "Private_companies(hotels, shops, etc.)": (
    "Hotel waste handed to municipality; shops source of plastic bags; ward 7 "
    "shopkeeper dustbin rule."
  ),
  "Monitoring": (
    "Campaigns without lasting change; no long-term monitoring; Beat the Plastic "
    "had participation but no practice change."
  ),
  "Financial_resources": (
    "Budget exists but not allocated/accessible; central budget not brought to "
    "local level; 3 years waiting on promised projects."
  ),
  "Infrastructure": (
    "No segregation areas, transfer stations, or environmentally safe landfill; "
    "40 years of plans (Daika, ADB, World Bank) with same situation."
  ),
  "Enforcement": (
    "Weak enforcement of national/municipal rules; national ban not implemented "
    "in practice."
  ),
  "Policies in place": (
    "National ban notice (weak); ward 7 shopkeeper litter rule; Beat the Plastic; "
    "clean-up campaigns; plastic bank (heard)."
  ),
  "Pol_awarness": "Beat the Plastic program; clean-up campaigns; advocated household campaigns.",
  "Pol_education": (
    "Awareness to every household; especially women/housewives who shop."
  ),
  "Pol_ban": "National ban notice issued; advocate ban plastic in every shop.",
  "Pol_subsitutes": "Paper bags, thokodi, cloth bags; distribute to every household initially.",
  "Pol_clean_up": "Dhulikhel clean-up campaigns; ward 7 shopkeeper collection system.",
  "Pol_recycling": (
    "Dankhuta segregation/recycling buildings; Chitwan large unit referenced."
  ),
  "Pol_waste_collection": (
    "Municipal collection (Mirvan post); daily collection in market areas; hotel "
    "hands all waste in one bag."
  ),
  "Sol_lead_agency": "Mayor must prioritize; municipality designates segregation/landfill areas.",
  "Sol_responsibilities": (
    "Municipality: segregation, separate collection, transfer stations, safe landfill; "
    "hotels segregate before handover."
  ),
  "Sol_awarness": (
    "Most important need — campaigns more important than funding alone; show "
    "side effects to every household."
  ),
  "Sol_segregation": "Segregation at source — household and hotel level.",
  "Sol_recycling": (
    "Integrate waste collectors and recyclers; replicate Chitwan/Dankhuta models."
  ),
  "Sol_education": (
    "Large-scale awareness; target women/housewives; behavioral change not just knowledge."
  ),
  "Sol_capacity": (
    "Bring policy and trained manpower from well-managed municipalities (Dankhuta)."
  ),
  "Sol_ban": (
    "Ban plastic in every shop; mandate paper/cloth bags with initial household distribution."
  ),
  "Sol_finance": (
    "2–10% of municipal annual budget for plastic/waste; 5% of national waste "
    "budget; separate compulsory budget line."
  ),
  "Sol_infrastructure": (
    "Designated segregation areas, transfer stations, environmentally safe landfill; "
    "adapt Dankhuta 'Waste into Money' model."
  ),
  "Sol_subsitutes": (
    "Revive thokodi/paper bags; cloth bags; home reuse before recycling."
  ),
  "Sol_clean_up": "Replicate ward 7 shopkeeper dustbin model across wards.",
  "Sol_enforcement": "Strict policies with real enforcement from national to local level.",
  "Sol_monitoring": (
    "Policies need follow-up and monitoring — not just announcements; continuous "
    "campaigns not one-off events."
  ),
  "Traditions to build on (free-hand)": (
    "Thokodi (newspaper paper bags sold in shops — now largely gone); grandfathers "
    "making paper bags; cloth bags; family practice of collecting bags for reuse."
  ),
}


def hfill(hex_):
  return PatternFill("solid", fgColor=hex_)


def hfont(bold=False, size=10, color="000000", italic=False):
  return Font(bold=bold, size=size, color=color, italic=italic)


thin = Side(style="thin", color="BFBFBF")


def thin_border():
  return Border(left=thin, right=thin, top=thin, bottom=thin)


WRAP_ALIGN = Alignment(wrap_text=True, vertical="top", horizontal="left")
CENTER_ALIGN = Alignment(wrap_text=True, vertical="center", horizontal="center")


def build_workbook():
  wb = openpyxl.Workbook()

  ws = wb.active
  ws.title = "Interview Coding"

  ws.merge_cells("A1:D1")
  title = ws["A1"]
  title.value = "Plastic Pollution Governance — Interview Coding (Codebook)"
  title.fill = hfill(DARK_BLUE)
  title.font = hfont(bold=True, size=13, color=WHITE)
  title.alignment = CENTER_ALIGN

  ws.merge_cells("A2:D2")
  meta = ws["A2"]
  meta.value = (
    "Interview: NPL_6  |  Country: NEPAL  |  Swasti Byanju, Hotel Owner — Dhulikhel  |  "
    "Actor: hotel_owner  |  22 June 2025"
  )
  meta.fill = hfill(MID_BLUE)
  meta.font = hfont(size=9, color=WHITE)
  meta.alignment = CENTER_ALIGN

  headers = ["Variable", "Code", "Codebook definition", "Coding notes"]
  for col, header in enumerate(headers, start=1):
    cell = ws.cell(row=4, column=col, value=header)
    cell.fill = hfill(LIGHT_BLUE)
    cell.font = hfont(bold=True, size=10, color=DARK_BLUE)
    cell.alignment = CENTER_ALIGN
    cell.border = thin_border()

  row = 5
  for variable, definition in CODEBOOK:
    values = [
      variable,
      CODING[variable],
      definition,
      NOTES.get(variable, ""),
    ]
    fill = hfill(LIGHT_GREEN) if row % 2 == 0 else hfill(WHITE)
    for col, value in enumerate(values, start=1):
      cell = ws.cell(row=row, column=col, value=value)
      cell.fill = fill
      cell.font = hfont(size=10)
      cell.alignment = WRAP_ALIGN
      cell.border = thin_border()
    row += 1

  ws.column_dimensions["A"].width = 34
  ws.column_dimensions["B"].width = 28
  ws.column_dimensions["C"].width = 52
  ws.column_dimensions["D"].width = 58
  ws.freeze_panes = "A5"

  wide = wb.create_sheet("Wide Format")
  wide_headers = [variable for variable, _ in CODEBOOK]
  for col, header in enumerate(wide_headers, start=1):
    cell = wide.cell(row=1, column=col, value=header)
    cell.fill = hfill(LIGHT_BLUE)
    cell.font = hfont(bold=True, size=9, color=DARK_BLUE)
    cell.alignment = CENTER_ALIGN
    cell.border = thin_border()
  for col, header in enumerate(wide_headers, start=1):
    cell = wide.cell(row=2, column=col, value=CODING[header])
    cell.fill = hfill(ORANGE)
    cell.alignment = WRAP_ALIGN
    cell.border = thin_border()
    wide.column_dimensions[get_column_letter(col)].width = max(14, min(28, len(header) + 2))

  ref = wb.create_sheet("Codebook Reference")
  ref.merge_cells("A1:B1")
  ref["A1"].value = "Codebook Variable Definitions"
  ref["A1"].fill = hfill(DARK_BLUE)
  ref["A1"].font = hfont(bold=True, size=12, color=WHITE)
  ref["A1"].alignment = CENTER_ALIGN
  for idx, (variable, definition) in enumerate(CODEBOOK, start=3):
    ref.cell(row=idx, column=1, value=variable).font = hfont(bold=True)
    ref.cell(row=idx, column=2, value=definition).alignment = WRAP_ALIGN
  ref.column_dimensions["A"].width = 34
  ref.column_dimensions["B"].width = 70

  stakeholder = wb.create_sheet("Stakeholder Analysis")
  stakeholder.merge_cells("A1:B1")
  stakeholder["A1"].value = "Stakeholder Analysis — NPL_6"
  stakeholder["A1"].fill = hfill(DARK_BLUE)
  stakeholder["A1"].font = hfont(bold=True, size=12, color=WHITE)
  stakeholder["A1"].alignment = CENTER_ALIGN
  for idx, (key, value) in enumerate(STAKEHOLDER.items(), start=3):
    stakeholder.cell(row=idx, column=1, value=key).font = hfont(bold=True)
    cell = stakeholder.cell(row=idx, column=2, value=value)
    cell.alignment = WRAP_ALIGN
  stakeholder.column_dimensions["A"].width = 28
  stakeholder.column_dimensions["B"].width = 90

  sources = wb.create_sheet("Sources")
  sources.merge_cells("A1:B1")
  sources["A1"].value = "Source Documents Used for Verification"
  sources["A1"].fill = hfill(DARK_BLUE)
  sources["A1"].font = hfont(bold=True, size=12, color=WHITE)
  sources["A1"].alignment = CENTER_ALIGN
  for idx, (doc, detail) in enumerate(SOURCES, start=3):
    sources.cell(row=idx, column=1, value=doc).font = hfont(bold=True)
    sources.cell(row=idx, column=2, value=detail).alignment = WRAP_ALIGN
  sources.column_dimensions["A"].width = 42
  sources.column_dimensions["B"].width = 60

  wb.save(OUTPUT_PATH)
  print(f"Saved {OUTPUT_PATH}")


if __name__ == "__main__":
  build_workbook()
