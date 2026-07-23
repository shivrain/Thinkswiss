#!/usr/bin/env python3
"""Generate 4P Index coding spreadsheet for Loi 2015-09 (English only). Repealed 2020."""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT = "/workspace/4P_Index_Loi_2015-09_Sachets_Plastiques.xlsx"

COLUMNS = [
    ("A", "policy_name"), ("B", "policy_url"), ("C", "policy_year"),
    ("D", "policy_objective"), ("E", "policy_target"), ("F", "policy_target_text"),
    ("G", "policy_type"), ("H", "policy_type_justification"),
    ("I", "policy_integration"), ("J", "policy_sectors_list"),
    ("K", "policy_circularity"), ("L", "policy_lifecycle_phases_list"),
    ("M", "policy_budget"), ("N", "policy_budget_text"),
    ("O", "policy_score"), ("P", "instrument_type"),
    ("Q", "instrument_lifecycle_stage"), ("R", "instrument_description"),
    ("S", "instrument_in_force"), ("T", "instrument_implementation"),
    ("U", "instrument_implementation_text"), ("V", "instrument_score"),
    ("W", "comments"),
]

POLICY = {
    "policy_name": "Law No. 2015-09 of 4 May 2015 on prohibition of thin plastic bags (repealed)",
    "policy_url": (
        "https://www.dri.gouv.sn/sites/default/files/an-documents/"
        "LOI%20N%202015%2009%20DU%204%20MAI%202015.pdf"
    ),
    "policy_year": 2015,
    "policy_objective": (
        "First dedicated national plastic-bag law: ban thin plastic bags (<30 microns), end free distribution "
        "of bags ≥30 microns, standardise permitted bags and require rational plastic-waste management. "
        "Repealed and replaced by Law No. 2020-04."
    ),
    "policy_target": 1,
    "policy_target_text": (
        "Art. 2: thickness below 30 microns prohibited; Art. 3: bags ≥30 microns may not be distributed free; "
        "Art. 10: production offence — FCFA 10,000,000 to 20,000,000 fine and 3–6 months imprisonment."
    ),
    "policy_type": 1.0,
    "policy_type_justification": (
        "Law adopted by the National Assembly on 21 April 2015 and promulgated 4 May 2015 as Law No. 2015-09. "
        "Parliamentary legislation (≠0.75). Repealed by Law No. 2020-04."
    ),
    "policy_integration": 0.75,
    "policy_sectors_list": "retail, production, consumption, waste management, environment",
    "policy_circularity": 0.75,
    "policy_lifecycle_phases_list": "production, consumption, recycling, disposal",
    "policy_budget": 0,
    "policy_budget_text": "",
}

INSTRUMENTS = [
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": (
            "Art. 2: prohibits production, import, use, possession for sale, sale and free distribution of "
            "plastic bags under 30 microns thickness throughout national territory."
        ),
        "instrument_in_force": 0,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            "Art. 10: production offence — FCFA 10–20 million and 3–6 months imprisonment (+0.25 enforcement); "
            "Art. 11: import customs offence (+0.25 enforcement); Art. 9: control agents designated (+0.25 authority). "
            "Repealed 2020."
        ),
        "comments": "Core thin-bag ban. Repealed — instrument_in_force=0 on all rows.",
    },
    {
        "instrument_type": 0.60,
        "instrument_lifecycle_stage": "Consumption",
        "instrument_description": (
            "Art. 3: plastic bags ≥30 microns may not be distributed or offered free; joint Commerce/Environment "
            "ministerial order sets transfer price to users."
        ),
        "instrument_in_force": 0,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            "Art. 12: distribution offence — FCFA 20,000–50,000 (+0.25 enforcement); "
            "price via subordinate order (+0.25 authority). Repealed 2020."
        ),
        "comments": "Economic instrument ending free bag distribution.",
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Production",
        "instrument_description": (
            "Art. 5: plastics industrialists must reduce plastic waste from their activities and develop "
            "recovery activities for production-process waste where appropriate."
        ),
        "instrument_in_force": 0,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            "Art. 13: non-compliance — FCFA 5–10 million and 1–3 months imprisonment (+0.25 enforcement); "
            "Art. 7: operator register (+0.25 monitoring). Repealed 2020."
        ),
        "comments": "Precursor EPR/upstream waste-reduction obligation.",
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Waste management",
        "instrument_description": (
            "Art. 6: plastics-sector operators must offer households and users a collection or take-back system "
            "for plastic waste for recovery, recycling or disposal."
        ),
        "instrument_in_force": 0,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            "Environment Minister sets by order conditions for collection, storage, sorting, transport and "
            "recovery/disposal (+0.25 authority). Art. 13 penalties (+0.25 enforcement). Repealed 2020."
        ),
        "comments": "Mandatory plastic-waste collection/take-back system.",
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "End of life",
        "instrument_description": (
            "Art. 8: users must deposit plastic waste at designated collection/take-back points. "
            "Art. 14: littering punishable by FCFA 10,000–30,000."
        ),
        "instrument_in_force": 0,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            "Art. 14: littering penalty (+0.25 enforcement); Art. 9: control agents (+0.25 authority). Repealed 2020."
        ),
        "comments": "Consumer delivery obligation and littering penalty combined.",
    },
]


def build_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "4P Index Coding"
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    auto_fill = PatternFill("solid", fgColor="E2EFDA")

    for col_idx, (col_letter, col_name) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=f"{col_letter} — {col_name}")
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for row_idx, instrument in enumerate(INSTRUMENTS, start=2):
        row_data = {**POLICY, **instrument}
        for col_idx, (_, key) in enumerate(COLUMNS, start=1):
            if key in ("policy_score", "instrument_score"):
                continue
            ws.cell(row=row_idx, column=col_idx, value=row_data.get(key, ""))
        g_col, i_col, k_col, m_col = (get_column_letter(n) for n in (7, 9, 11, 13))
        p_col, t_col, o_col, v_col = (get_column_letter(n) for n in (16, 20, 15, 22))
        ws[f"{o_col}{row_idx}"] = f"=AVERAGE({g_col}{row_idx},{i_col}{row_idx},{k_col}{row_idx},{m_col}{row_idx},{p_col}{row_idx})"
        ws[f"{v_col}{row_idx}"] = f"=AVERAGE({p_col}{row_idx},{t_col}{row_idx})"
        ws[f"{o_col}{row_idx}"].fill = auto_fill
        ws[f"{v_col}{row_idx}"].fill = auto_fill

    widths = {"A": 48, "B": 52, "C": 10, "D": 52, "E": 12, "F": 60, "G": 10, "H": 52,
              "I": 14, "J": 40, "K": 14, "L": 36, "M": 12, "N": 52, "O": 12, "P": 14,
              "Q": 22, "R": 60, "S": 14, "T": 18, "U": 52, "V": 14, "W": 52}
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(OUTPUT)
    print(f"Saved {OUTPUT} ({len(INSTRUMENTS)} instrument rows)")


if __name__ == "__main__":
    build_workbook()
