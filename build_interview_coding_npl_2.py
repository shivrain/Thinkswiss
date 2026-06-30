"""Generate interview coding spreadsheet for Nepal interview NPL_2.

Coding follows the Overview All Interviews codebook and is based on the
Ganesh Shah (Former Minister) interview of 21 June 2025 — transcript,
interview notes, and guideline.
"""

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUTPUT_PATH = "/workspace/interview_coding_NPL_2.xlsx"

DARK_BLUE = "1F3864"
MID_BLUE = "2E75B6"
LIGHT_BLUE = "D9E2F3"
LIGHT_GREEN = "E2EFDA"
ORANGE = "FCE4D6"
WHITE = "FFFFFF"

SOURCES = [
  ("Interview transcript", "Nepal Former Minister — Ganesh Shah, 21 June 2025"),
  ("Interview notes", "Interviewer bullet notes and guideline responses"),
  ("Interview guideline", "Plastic Pollution Governance in Nepal — general form"),
  ("Codebook", "Overview All Interviews — NEW Codebook"),
]

STAKEHOLDER = {
  "Interview ID": "NPL_2",
  "Country": "NEPAL",
  "Interview date": "21 June 2025",
  "Interviewee": "Ganesh Shah",
  "Affiliation / role": "Former Minister (established Department of Environment ~15 years ago)",
  "Professional background": "Water-supply engineer; science diplomacy; PLEASE project involvement",
  "Coded actor type": "pol_party",
  "Actor classification rationale": (
    "Former Minister and political leader who created the Department of Environment. "
    "Speaks as senior policy architect and political advisor, not a current civil "
    "servant (governmental) or active researcher (science). Strong science background "
    "noted but primary coded role is political leadership."
  ),
  "Role in plastic governance": (
    "Policy advocate for UN plastics treaty, EPR, federal–provincial–local division; "
    "links informal waste sector, private companies, and government frameworks"
  ),
  "Perspective": (
    "Outsider-insider — former minister observing weak top-down governance while "
    "praising grassroots/project-led progress (PLEASE, Bio-Camp, Bagmati movement)"
  ),
  "Key networks mentioned": (
    "PLEASE project, Nepal Plastic Foundation, Solid Waste Management Association, "
    "Bio-Camp, Bagmati Cleaning Movement, UN plastics treaty delegation"
  ),
}

CODING = {
  "Country": "NPL",
  "Country name": "NEPAL",
  "ID": "NPL_2",
  "Actortype": "pol_party",
  "Problem_awareness_pop": "high",
  "Problem_awareness_pol": "high",
  "Problem_severity": "high",
  "Problem_littering": "yes",
  "Problem_consumption": "yes",
  "Problem_recycling": "yes",
  "Problem_waste_mgmt": "yes",
  "Problem_production": "yes",
  "Problem_alternatives": "yes",
  "Problem_waste_segregation": "yes",
  "Problem_import": "yes",
  "Impacts": "water, cities, health, agriculture",
  "Governance": "yes",
  "Relevance_international_pol": "yes",
  "Coordination_sectoral": "yes",
  "Coordination_levels": "yes",
  "Unclear_responsibilities": "yes",
  "Actors": "yes",
  "Federal government": "yes",
  "Provincial government": "yes",
  "Local government": "yes",
  "Students": "no",
  "Private_sector": "yes",
  "Civil_society": "yes",
  "Science": "yes",
  "Households": "yes",
  "Education institutions": "yes",
  "Private_companies(hotels, shops, etc.)": "yes",
  "Implementation issues": "yes",
  "Monitoring": "yes",
  "Financial_resources": "yes",
  "Research": "yes",
  "Infrastructure": "yes",
  "Enforcement": "yes",
  "Policies in place": "yes",
  "Pol_epr": "no",
  "Pol_import": "yes",
  "Pol_awarness": "no",
  "Pol_education": "no",
  "Pol_capacity": "no",
  "Pol_RD": "no",
  "Pol_tax": "yes",
  "Pol_ban": "yes",
  "Pol_subsitutes": "yes",
  "Pol_clean_up": "yes",
  "Pol_upcycling": "yes",
  "Pol_recycling": "yes",
  "Pol_waste_collection": "yes",
  "Solutions": "yes",
  "Sol_lead_agency": "yes",
  "Sol_responsibilities": "yes",
  "Sol_epr": "yes",
  "Sol_awarness": "yes",
  "Sol_segregation": "yes",
  "Sol_upcycling": "yes",
  "Sol_recycling": "yes",
  "Sol_education": "yes",
  "Sol_capacity": "yes",
  "Sol_RD": "yes",
  "Sol_tax": "no",
  "Sol_ban": "no",
  "Sol_finance": "yes",
  "Sol_infrastructure": "yes",
  "Sol_subsitutes": "yes",
  "Sol_clean_up": "yes",
  "Sol_enforcement": "yes",
  "Sol_monitoring": "yes",
  "Traditions to build on (free-hand)": (
    "Pat/banana leaf plates for community feasts; cotton or towel bags for shopping; "
    "reuse plastic bags before recycling; Bagmati Cleaning Movement as grassroots model"
  ),
}

CODEBOOK = [
  ("Country", "Country ISO code"),
  ("Country name", "Country name"),
  ("ID", "InterviewID: ISO_Nr"),
  ("Actortype", "governmental, pol_party, science, cso, igo, private_sector, shop_owner, hotel_owner, household"),
  ("Problem_awareness_pop", "high, medium, low, NA — lack of awareness among population mentioned"),
  ("Problem_awareness_pol", "high, medium, low, NA — lack of awareness among policy makers mentioned"),
  ("Problem_severity", "high, medium, low, NA"),
  ("Problem_littering", "Littering is a problem; was mentioned: yes/no"),
  ("Problem_consumption", "High consumption is a problem; was mentioned: yes/no"),
  ("Problem_recycling", "No or too little recycling is a problem; was mentioned: yes/no"),
  ("Problem_waste_mgmt", "Ill waste management is a problem; was mentioned: yes/no"),
  ("Problem_production", "High production rates; was mentioned: yes/no"),
  ("Problem_alternatives", "No good alternatives; was mentioned: yes/no"),
  ("Problem_waste_segregation", "Lack of segregation was mentioned: yes/no"),
  ("Problem_import", "Plastic imports are a problem; was mentioned: yes/no"),
  ("Impacts", "Impacts mentioned on water, cities, biodiversity, health, etc. or NA"),
  ("Governance", "Mentioned: yes/no"),
  ("Relevance_international_pol", "International treaties are relevant for national level policy"),
  ("Coordination_sectoral", "Lack of coordination across sectors"),
  ("Coordination_levels", "Lack of coordination across levels"),
  ("Unclear_responsibilities", "Unclear responsibilities among the involved actors"),
  ("Actors", "Mentioned: yes/no"),
  ("Federal government", "Mentioned: yes/no"),
  ("Provincial government", "Mentioned: yes/no"),
  ("Local government", "Mentioned: yes/no"),
  ("Students", "Mentioned: yes/no"),
  ("Private_sector", "Mentioned: yes/no"),
  ("Civil_society", "Mentioned: yes/no"),
  ("Science", "Mentioned: yes/no"),
  ("Households", "Mentioned: yes/no"),
  ("Education institutions", "Mentioned: yes/no"),
  ("Private_companies(hotels, shops, etc.)", "Mentioned: yes/no"),
  ("Implementation issues", "Mentioned: yes/no"),
  ("Monitoring", "Lack of monitoring"),
  ("Financial_resources", "Lack of financial resources"),
  ("Research", "Lack of research"),
  ("Infrastructure", "Lack of infrastructure"),
  ("Enforcement", "Lack of enforcement"),
  ("Policies in place", "Mentioned: yes/no"),
  ("Pol_epr", "Extended producer responsibility"),
  ("Pol_import", "Import regulations"),
  ("Pol_awarness", "Awareness campaigns"),
  ("Pol_education", "Education programs"),
  ("Pol_capacity", "Capacity building"),
  ("Pol_RD", "Research & development"),
  ("Pol_tax", "Levies or taxes on specific products (often bags)"),
  ("Pol_ban", "Bans of specific products"),
  ("Pol_subsitutes", "Alternatives like reusable bags, or banana leaf plates"),
  ("Pol_clean_up", "Clean-up campaigns"),
  ("Pol_upcycling", "Making a new product, like blacktop"),
  ("Pol_recycling", "Making a new raw material"),
  ("Pol_waste_collection", "Waste collection"),
  ("Solutions", "Mentioned: yes/no"),
  ("Sol_lead_agency", "Procedural measures to establish lead agency"),
  ("Sol_responsibilities", "Clear responsibilities"),
  ("Sol_epr", "Extended producer responsibility"),
  ("Sol_awarness", "Awareness raising measures"),
  ("Sol_segregation", "Segregation of waste"),
  ("Sol_upcycling", "Upcycling of plastics"),
  ("Sol_recycling", "New raw materials"),
  ("Sol_education", "Education programs at schools"),
  ("Sol_capacity", "Capacity-building programs"),
  ("Sol_RD", "Research programs"),
  ("Sol_tax", "Taxes or levies on products"),
  ("Sol_ban", "Ban of products"),
  ("Sol_finance", "Financial measures"),
  ("Sol_infrastructure", "Infrastructural measures"),
  ("Sol_subsitutes", "Alternatives"),
  ("Sol_clean_up", "Clean-up campaigns"),
  ("Sol_enforcement", "Enforcement procedures"),
  ("Sol_monitoring", "Monitoring procedures"),
  ("Traditions to build on (free-hand)", "Freehand or NA"),
]

NOTES = {
  "Actortype": (
    "Ganesh Shah, Former Minister. Created Department of Environment; water-supply "
    "engineer; science-diplomacy advocate. Coded pol_party (political leadership), "
    "not science, despite technical background."
  ),
  "Problem_awareness_pop": (
    "Minimal public awareness of plastic types, purposes, and impacts; people do not "
    "know what is eco-friendly; hot food in thin black bags; one plastic bag per shop item."
  ),
  "Problem_awareness_pol": (
    "Officials distribute synthetic khada and plastic folders at ceremonies; raised "
    "repeatedly with bureaucrats and politicians who still use plastic inappropriately."
  ),
  "Problem_severity": (
    "Very concerned; visible and invisible (micro) plastics; forecasts of >1 tonne "
    "plastic waste per person; global transboundary challenge."
  ),
  "Problem_littering": (
    "Visible litter in urban/rural areas; clogging drains and sewers; EPA fines for "
    "street/river dumping exist but enforcement weak."
  ),
  "Problem_consumption": (
    "Excessive single-use packaging; multilayer materials; single-use bottles, plates, "
    "glasses, and spoons at feasts."
  ),
  "Problem_recycling": (
    "Informal sector collects ~99% of valuable plastics but formal systems incomplete; "
    "need linkage between informal collectors and formal processors."
  ),
  "Problem_waste_mgmt": (
    "No specific plastics policy; general Environment Protection Act only; management "
    "evolved through projects rather than governance."
  ),
  "Problem_production": (
    "Imports of polymers, granules, pipes, bottles, telecom cables; rapid polymer "
    "technology evolution outpaces regulation."
  ),
  "Problem_alternatives": (
    "Traditions exist (leaf plates, cotton bags) but underused; public lacks clarity "
    "on eco-friendly products and government labelling standards."
  ),
  "Problem_waste_segregation": (
    "Households rarely sort everything; PLEASE introduced some segregation; people "
    "must learn correct separation."
  ),
  "Problem_import": (
    "Imports from abroad without polymer testing labs; federal government must know "
    "which plastics/granules are entering the country."
  ),
  "Impacts": (
    "Water: microplastics in rivers, irrigation, drinking water, fisheries, bottled/jar "
    "water safety. Cities: clogged sewers. Health: burning carcinogenic, low-grade "
    "plastics, chemical leaching. Agriculture: soil moisture, crop productivity decline."
  ),
  "Relevance_international_pol": (
    "UN plastics treaty (Korea meeting, Colombo regional forum); Nepal delegation; "
    "international pressure pushing domestic legislation."
  ),
  "Coordination_sectoral": (
    "Manufacturers and waste managers lack coordination; informal collectors vs formal "
    "processors need linkage; Nepal Plastic Foundation CSR consortium."
  ),
  "Coordination_levels": (
    "No clear division federal/provincial/local; upside-down sequence (informal sector "
    "first, federal framework only now emerging)."
  ),
  "Unclear_responsibilities": (
    "Main open question: which agency is lead body; debate on centralization vs "
    "decentralization of inspection and monitoring."
  ),
  "Federal government": (
    "Department of Environment (small unit), plastic-waste assessment led by government, "
    "import regulation, EPR, research labs."
  ),
  "Provincial government": (
    "Provincial assemblies debated plastic waste; provincial government should coordinate "
    "among municipalities."
  ),
  "Local government": (
    "6 metropolitan, 11 sub-metropolitan, ~293 municipalities; inspection, monitoring, "
    "evaluation, waste taxes, competitive tenders."
  ),
  "Private_sector": (
    "78+ registered waste companies; Bio-Camp factory; recycling/upcycling entrepreneurs; "
    "15-year growth of collectors and processors."
  ),
  "Civil_society": (
    "Bagmati Cleaning Movement (people's initiative); informal waste pickers; PLEASE project."
  ),
  "Science": (
    "Institute of Engineering, Central Dept. Environmental Science, polymer chemists, "
    "young EPR researcher, agricultural scientists."
  ),
  "Households": (
    "Separate recyclables when collectors arrive; pay collection companies; waste "
    "producers targeted in future policy."
  ),
  "Education institutions": (
    "Kathmandu University ceremony example; Institute of Engineering PLEASE component; "
    "Tribhuvan University."
  ),
  "Private_companies(hotels, shops, etc.)": (
    "Hotels store bottles for collectors; Kalishishi; tourist coaches; pharmacies "
    "giving plastic bags for few tablets."
  ),
  "Monitoring": (
    "Inadequate monitoring; DoE has very little manpower; local government should "
    "inspect and evaluate."
  ),
  "Financial_resources": (
    "Department of Environment established with very little manpower and few resources; "
    "lack of resources at local operational level."
  ),
  "Research": (
    "Rigorous scientific research still lacking; microplastics in organic fertilizer; "
    "life-cycle assessment needed."
  ),
  "Infrastructure": (
    "No testing laboratories for imported plastic materials, granules, or polymer types."
  ),
  "Enforcement": (
    "40-micron ban exists on paper but never enforced; officers cannot act without "
    "strict laws, bylaws, and guidelines."
  ),
  "Policies in place": (
    "Environment Protection Act (general); public announcements/bans; piecemeal notices; "
    "plastic assessment underway — no specific plastics act yet."
  ),
  "Pol_epr": "EPR modalities being drafted but not yet implemented nationwide.",
  "Pol_import": (
    "Ban on plastics <40 microns includes import restriction; future federal role to "
    "regulate polymer imports."
  ),
  "Pol_awarness": (
    "Nationwide campaign urgently needed but not in place; progress mainly from "
    "projects (PLEASE, Bio-Camp) not national policy."
  ),
  "Pol_tax": (
    "Metropolitan waste taxes/levies; EPA fines (5,000 NPR for river/street pollution); "
    "municipal fees for landfill dumping."
  ),
  "Pol_ban": "Ban on plastics thinner than 40 microns announced ~10 years ago (unenforced).",
  "Pol_subsitutes": (
    "Leaf-based biodegradable substitutes emerging; traditional pat-leaf plates still "
    "used in some community feasts."
  ),
  "Pol_clean_up": "Bagmati Cleaning Movement — riverbank plastic collection and processing.",
  "Pol_upcycling": (
    "Bio-Camp: multilayer pouches into plastic plywood boards and flowerpots."
  ),
  "Pol_recycling": (
    "Baling, processing, PLEASE catalogue of 8–9 plastic categories matched to products."
  ),
  "Pol_waste_collection": (
    "Informal then semi-formal collection networks; 99% of valuable plastics collected "
    "by informal sector; licensed companies with municipal contracts."
  ),
  "Sol_lead_agency": "Designate lead agency to oversee future plastics act.",
  "Sol_responsibilities": (
    "Federal: EPR, research, import regulation. Provincial: coordination. Local: "
    "inspection, monitoring, evaluation."
  ),
  "Sol_epr": "EPR critical; Nepal must begin formulating; UN should stress EPR.",
  "Sol_awarness": (
    "Sustained nationwide campaign on correct plastic use; citizens should ask whether "
    "plastic is necessary."
  ),
  "Sol_segregation": "Train workers and public to separate plastics correctly.",
  "Sol_upcycling": "Expand pilot projects like Bio-Camp and PLEASE product development.",
  "Sol_recycling": "Apply 3R/4R/5R principles; send surplus plastic for recycling.",
  "Sol_education": (
    "Public education first priority; behavioural/mindset change; life-cycle assessment literacy."
  ),
  "Sol_capacity": (
    "Human-resource development; skilled separation workers; federal polymer testing laboratories."
  ),
  "Sol_RD": (
    "Research on plastic burning and air pollution; plastics in water (river, drinking, "
    "irrigation); plastics in soil and agriculture."
  ),
  "Sol_ban": (
    "Existing ban failed due to no enforcement; forward solutions emphasize EPR, "
    "awareness, and alternatives over new bans."
  ),
  "Sol_finance": (
    "Waste-based entrepreneurship; transform informal sector into formal enterprises; "
    "incentives for alternatives."
  ),
  "Sol_infrastructure": "Establish federal testing laboratories for imported plastic materials.",
  "Sol_subsitutes": (
    "Return to pat/banana leaf plates, cotton/towel bags; reuse bags 2–3 times before recycling."
  ),
  "Sol_clean_up": "Learn from and expand Bagmati Cleaning Movement and PLEASE pilots.",
  "Sol_enforcement": "Strict enforcement of thin-plastic restrictions and labelling laws.",
  "Sol_monitoring": "Local-government inspection and evaluation; air/water quality monitoring.",
  "Traditions to build on (free-hand)": (
    "Pat/banana leaf plates at bhoj feasts; father's towel for market shopping; cotton "
    "bags; reuse-before-recycle habit; Bagmati people's initiative."
  ),
}


def hfill(hex_):
  return PatternFill("solid", fgColor=hex_)


def hfont(bold=False, size=10, color="000000", italic=False):
  return Font(bold=bold, size=size, color=color, italic=italic)


thin = Side(style="thin", color="BFBFBF")


def thin_border():
  return Border(left=thin, right=thin, top=thin, bottom=thin)


WRAP_ALIGN = Alignment(wrap_text=True, vertical="top", horizontal="left")
CENTER_ALIGN = Alignment(wrap_text=True, vertical="center", horizontal="center")


def build_workbook():
  wb = openpyxl.Workbook()

  ws = wb.active
  ws.title = "Interview Coding"

  ws.merge_cells("A1:D1")
  title = ws["A1"]
  title.value = "Plastic Pollution Governance — Interview Coding (Codebook)"
  title.fill = hfill(DARK_BLUE)
  title.font = hfont(bold=True, size=13, color=WHITE)
  title.alignment = CENTER_ALIGN

  ws.merge_cells("A2:D2")
  meta = ws["A2"]
  meta.value = (
    "Interview: NPL_2  |  Country: NEPAL  |  Ganesh Shah, Former Minister  |  "
    "Actor: pol_party  |  21 June 2025"
  )
  meta.fill = hfill(MID_BLUE)
  meta.font = hfont(size=9, color=WHITE)
  meta.alignment = CENTER_ALIGN

  headers = ["Variable", "Code", "Codebook definition", "Coding notes"]
  for col, header in enumerate(headers, start=1):
    cell = ws.cell(row=4, column=col, value=header)
    cell.fill = hfill(LIGHT_BLUE)
    cell.font = hfont(bold=True, size=10, color=DARK_BLUE)
    cell.alignment = CENTER_ALIGN
    cell.border = thin_border()

  row = 5
  for variable, definition in CODEBOOK:
    values = [
      variable,
      CODING[variable],
      definition,
      NOTES.get(variable, ""),
    ]
    fill = hfill(LIGHT_GREEN) if row % 2 == 0 else hfill(WHITE)
    for col, value in enumerate(values, start=1):
      cell = ws.cell(row=row, column=col, value=value)
      cell.fill = fill
      cell.font = hfont(size=10)
      cell.alignment = WRAP_ALIGN
      cell.border = thin_border()
    row += 1

  ws.column_dimensions["A"].width = 34
  ws.column_dimensions["B"].width = 28
  ws.column_dimensions["C"].width = 52
  ws.column_dimensions["D"].width = 58
  ws.freeze_panes = "A5"

  wide = wb.create_sheet("Wide Format")
  wide_headers = [variable for variable, _ in CODEBOOK]
  for col, header in enumerate(wide_headers, start=1):
    cell = wide.cell(row=1, column=col, value=header)
    cell.fill = hfill(LIGHT_BLUE)
    cell.font = hfont(bold=True, size=9, color=DARK_BLUE)
    cell.alignment = CENTER_ALIGN
    cell.border = thin_border()
  for col, header in enumerate(wide_headers, start=1):
    cell = wide.cell(row=2, column=col, value=CODING[header])
    cell.fill = hfill(ORANGE)
    cell.alignment = WRAP_ALIGN
    cell.border = thin_border()
    wide.column_dimensions[get_column_letter(col)].width = max(14, min(28, len(header) + 2))

  ref = wb.create_sheet("Codebook Reference")
  ref.merge_cells("A1:B1")
  ref["A1"].value = "Codebook Variable Definitions"
  ref["A1"].fill = hfill(DARK_BLUE)
  ref["A1"].font = hfont(bold=True, size=12, color=WHITE)
  ref["A1"].alignment = CENTER_ALIGN
  for idx, (variable, definition) in enumerate(CODEBOOK, start=3):
    ref.cell(row=idx, column=1, value=variable).font = hfont(bold=True)
    ref.cell(row=idx, column=2, value=definition).alignment = WRAP_ALIGN
  ref.column_dimensions["A"].width = 34
  ref.column_dimensions["B"].width = 70

  stakeholder = wb.create_sheet("Stakeholder Analysis")
  stakeholder.merge_cells("A1:B1")
  stakeholder["A1"].value = "Stakeholder Analysis — NPL_2"
  stakeholder["A1"].fill = hfill(DARK_BLUE)
  stakeholder["A1"].font = hfont(bold=True, size=12, color=WHITE)
  stakeholder["A1"].alignment = CENTER_ALIGN
  for idx, (key, value) in enumerate(STAKEHOLDER.items(), start=3):
    stakeholder.cell(row=idx, column=1, value=key).font = hfont(bold=True)
    cell = stakeholder.cell(row=idx, column=2, value=value)
    cell.alignment = WRAP_ALIGN
  stakeholder.column_dimensions["A"].width = 28
  stakeholder.column_dimensions["B"].width = 90

  sources = wb.create_sheet("Sources")
  sources.merge_cells("A1:B1")
  sources["A1"].value = "Source Documents Used for Verification"
  sources["A1"].fill = hfill(DARK_BLUE)
  sources["A1"].font = hfont(bold=True, size=12, color=WHITE)
  sources["A1"].alignment = CENTER_ALIGN
  for idx, (doc, detail) in enumerate(SOURCES, start=3):
    sources.cell(row=idx, column=1, value=doc).font = hfont(bold=True)
    sources.cell(row=idx, column=2, value=detail).alignment = WRAP_ALIGN
  sources.column_dimensions["A"].width = 42
  sources.column_dimensions["B"].width = 60

  wb.save(OUTPUT_PATH)
  print(f"Saved {OUTPUT_PATH}")


if __name__ == "__main__":
  build_workbook()
