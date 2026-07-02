"""Generate interview coding spreadsheet for Nepal interview NPL_29.

Coding follows the Overview All Interviews codebook and is based on the
former UNDP/UNEP member (ecology background) interview of 18 February 2026 —
interview guideline (recorded). Uses expanded Impacts taxonomy.
"""

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUTPUT_PATH = "/workspace/interview_coding_NPL_29.xlsx"

DARK_BLUE = "1F3864"
MID_BLUE = "2E75B6"
LIGHT_BLUE = "D9E2F3"
LIGHT_GREEN = "E2EFDA"
ORANGE = "FCE4D6"
WHITE = "FFFFFF"

SOURCES = [
  ("Interview guideline", "Former UNDP/UNEP member — structured guideline, 18 Feb 2026"),
  ("Codebook", "Overview All Interviews — NEW Codebook (expanded Impacts)"),
  (
    "Coding note",
    "Consent given with recording. Respondent name not recorded. Codes verified "
    "against interview guideline. Question 6 and Q9 not asked.",
  ),
]

STAKEHOLDER = {
  "Interview ID": "NPL_29",
  "Country": "NEPAL",
  "Interview number": "UNDP/UNEP / international-environment fieldwork series",
  "Interview date": "18 February 2026",
  "Interviewee": "Unknown — former UNDP and UNEP member",
  "Affiliation / role": (
    "Former international organisation staff (UNDP, UNEP); ecology background"
  ),
  "Organization": (
    "Former UNDP and UNEP; worked with two waste-segregation companies in Nepal"
  ),
  "Stakeholder list mapping": "igo",
  "Coded actor type (codebook)": "igo",
  "Actor classification rationale": (
    "Former UNDP and UNEP member — international governmental organisation "
    "perspective on plastic governance, supplemented by ecology expertise and "
    "private segregation-company experience."
  ),
  "Perspective": (
    "Highly concerned ecology/IGO practitioner — plastic cheapness drives use; "
    "bans without alternatives failed in 1990s–2000s; no Kathmandu segregation; "
    "CSO/youth action without government; municipalities must lead with tax "
    "incentives, awareness, education, and regulated waste-to-wealth private sector"
  ),
  "Project context": (
    "1990s–2000s plastic ban without alternatives; cloth bags NPR 40; two "
    "segregation companies; young-generation initiatives; ocean/soil/microplastic "
    "concerns; Kathmandu mixed-waste disposal"
  ),
}

CODING = {
  "Country": "NPL",
  "Country name": "NEPAL",
  "ID": "NPL_29",
  "Actortype": "igo",
  "Problem_awareness_pop": "medium",
  "Problem_awareness_pol": "medium",
  "Problem_severity": "high",
  "Problem_littering": "yes",
  "Problem_consumption": "yes",
  "Problem_recycling": "yes",
  "Problem_waste_mgmt": "yes",
  "Problem_production": "no",
  "Problem_alternatives": "yes",
  "Problem_waste_segregation": "yes",
  "Problem_import": "no",
  "Impacts": (
    "aquatic life, marine pollution, wildlife, soil, health, microplastics, agriculture"
  ),
  "Governance": "yes",
  "Relevance_international_pol": "yes",
  "Coordination_sectoral": "yes",
  "Coordination_levels": "yes",
  "Unclear_responsibilities": "yes",
  "Actors": "yes",
  "Federal government": "yes",
  "Provincial government": "no",
  "Local government": "yes",
  "Students": "yes",
  "Private_sector": "yes",
  "Civil_society": "yes",
  "Science": "yes",
  "Households": "yes",
  "Education institutions": "yes",
  "Private_companies(hotels, shops, etc.)": "yes",
  "Implementation issues": "yes",
  "Monitoring": "no",
  "Financial_resources": "no",
  "Research": "yes",
  "Infrastructure": "no",
  "Enforcement": "yes",
  "Policies in place": "yes",
  "Pol_epr": "no",
  "Pol_import": "no",
  "Pol_awarness": "yes",
  "Pol_education": "yes",
  "Pol_capacity": "yes",
  "Pol_RD": "no",
  "Pol_tax": "yes",
  "Pol_ban": "yes",
  "Pol_subsitutes": "yes",
  "Pol_clean_up": "no",
  "Pol_upcycling": "no",
  "Pol_recycling": "yes",
  "Pol_waste_collection": "yes",
  "Solutions": "yes",
  "Sol_lead_agency": "no",
  "Sol_responsibilities": "yes",
  "Sol_epr": "no",
  "Sol_awarness": "yes",
  "Sol_segregation": "yes",
  "Sol_upcycling": "no",
  "Sol_recycling": "yes",
  "Sol_education": "yes",
  "Sol_capacity": "yes",
  "Sol_RD": "no",
  "Sol_tax": "yes",
  "Sol_ban": "no",
  "Sol_finance": "no",
  "Sol_infrastructure": "yes",
  "Sol_subsitutes": "yes",
  "Sol_clean_up": "no",
  "Sol_enforcement": "yes",
  "Sol_monitoring": "no",
  "Traditions to build on (free-hand)": "NA",
}

CODEBOOK = [
  ("Country", "Country ISO code"),
  ("Country name", "Country name"),
  ("ID", "InterviewID: ISO_Nr"),
  ("Actortype", "governmental, pol_party, science, cso, igo, private_sector, shop_owner, hotel_owner, household"),
  ("Problem_awareness_pop", "high, medium, low, NA — lack of awareness among population mentioned"),
  ("Problem_awareness_pol", "high, medium, low, NA — lack of awareness among policy makers mentioned"),
  ("Problem_severity", "high, medium, low, NA"),
  ("Problem_littering", "Littering is a problem; was mentioned: yes/no"),
  ("Problem_consumption", "High consumption is a problem; was mentioned: yes/no"),
  ("Problem_recycling", "No or too little recycling is a problem; was mentioned: yes/no"),
  ("Problem_waste_mgmt", "Ill waste management is a problem; was mentioned: yes/no"),
  ("Problem_production", "High production rates; was mentioned: yes/no"),
  ("Problem_alternatives", "No good alternatives; was mentioned: yes/no"),
  ("Problem_waste_segregation", "Lack of segregation was mentioned: yes/no"),
  ("Problem_import", "Plastic imports are a problem; was mentioned: yes/no"),
  ("Impacts", "Impacts mentioned on water, cities, biodiversity, health, etc. or NA"),
  ("Governance", "Mentioned: yes/no"),
  ("Relevance_international_pol", "International treaties are relevant for national level policy"),
  ("Coordination_sectoral", "Lack of coordination across sectors"),
  ("Coordination_levels", "Lack of coordination across levels"),
  ("Unclear_responsibilities", "Unclear responsibilities among the involved actors"),
  ("Actors", "Mentioned: yes/no"),
  ("Federal government", "Mentioned: yes/no"),
  ("Provincial government", "Mentioned: yes/no"),
  ("Local government", "Mentioned: yes/no"),
  ("Students", "Mentioned: yes/no"),
  ("Private_sector", "Mentioned: yes/no"),
  ("Civil_society", "Mentioned: yes/no"),
  ("Science", "Mentioned: yes/no"),
  ("Households", "Mentioned: yes/no"),
  ("Education institutions", "Mentioned: yes/no"),
  ("Private_companies(hotels, shops, etc.)", "Mentioned: yes/no"),
  ("Implementation issues", "Mentioned: yes/no"),
  ("Monitoring", "Lack of monitoring"),
  ("Financial_resources", "Lack of financial resources"),
  ("Research", "Lack of research"),
  ("Infrastructure", "Lack of infrastructure"),
  ("Enforcement", "Lack of enforcement"),
  ("Policies in place", "Mentioned: yes/no"),
  ("Pol_epr", "Extended producer responsibility"),
  ("Pol_import", "Import regulations"),
  ("Pol_awarness", "Awareness campaigns"),
  ("Pol_education", "Education programs"),
  ("Pol_capacity", "Capacity building"),
  ("Pol_RD", "Research & development"),
  ("Pol_tax", "Levies or taxes on specific products (often bags)"),
  ("Pol_ban", "Bans of specific products"),
  ("Pol_subsitutes", "Alternatives like reusable bags, or banana leaf plates"),
  ("Pol_clean_up", "Clean-up campaigns"),
  ("Pol_upcycling", "Making a new product, like blacktop"),
  ("Pol_recycling", "Making a new raw material"),
  ("Pol_waste_collection", "Waste collection"),
  ("Solutions", "Mentioned: yes/no"),
  ("Sol_lead_agency", "Procedural measures to establish lead agency"),
  ("Sol_responsibilities", "Clear responsibilities"),
  ("Sol_epr", "Extended producer responsibility"),
  ("Sol_awarness", "Awareness raising measures"),
  ("Sol_segregation", "Segregation of waste"),
  ("Sol_upcycling", "Upcycling of plastics"),
  ("Sol_recycling", "New raw materials"),
  ("Sol_education", "Education programs at schools"),
  ("Sol_capacity", "Capacity-building programs"),
  ("Sol_RD", "Research programs"),
  ("Sol_tax", "Taxes or levies on products"),
  ("Sol_ban", "Ban of products"),
  ("Sol_finance", "Financial measures"),
  ("Sol_infrastructure", "Infrastructural measures"),
  ("Sol_subsitutes", "Alternatives"),
  ("Sol_clean_up", "Clean-up campaigns"),
  ("Sol_enforcement", "Enforcement procedures"),
  ("Sol_monitoring", "Monitoring procedures"),
  ("Traditions to build on (free-hand)", "Freehand or NA"),
]

NOTES = {
  "Actortype": (
    "Former UNDP and UNEP member; ecology background. Stakeholder list: igo. "
    "Codebook: igo."
  ),
  "Problem_awareness_pop": (
    "Family/friends aware but have no other options or choice; cloth bags "
    "expensive (NPR 40) — awareness without affordable alternatives."
  ),
  "Problem_awareness_pol": (
    "Organisations and youth run initiatives but government not involved; "
    "municipalities should lead but initiatives lacking."
  ),
  "Problem_severity": (
    "Very concerned — worries about family and children's health; plastics "
    "everywhere including garden soil; cheap plastic sustains high use."
  ),
  "Problem_littering": (
    "Plastics everywhere — garden pit contained 6–7 pieces; ocean and "
    "environment widely affected."
  ),
  "Problem_consumption": (
    "Plastic cheap and easy — as long as it remains cheap people will continue "
    "using it."
  ),
  "Problem_recycling": (
    "Waste-to-wealth potential via private companies; worked with two "
    "segregation companies but system not scaled."
  ),
  "Problem_waste_mgmt": (
    "Kathmandu: ordinary people throw waste together; no segregation at all."
  ),
  "Problem_alternatives": (
    "1990s–2000s ban attempted without providing alternatives; cloth bags "
    "exist but too expensive for ordinary people."
  ),
  "Problem_waste_segregation": (
    "No segregation at all in Kathmandu despite segregation-company presence."
  ),
  "Impacts": (
    "Aquatic life and marine pollution: ocean affected. Wildlife: cows eat "
    "plastic. Soil/agriculture: garden pits and soil contaminated. Health: "
    "family/children health worry. Microplastics: micro- and nano-plastic cited."
  ),
  "Relevance_international_pol": (
    "UNDP/UNEP background; global ocean pollution framing; international "
    "organisation experience on environmental governance."
  ),
  "Coordination_sectoral": (
    "Organisations and young generation active but government absent from "
    "awareness/campaign efforts."
  ),
  "Coordination_levels": (
    "Municipalities, local bodies, ruling bodies, and central government "
    "initiatives needed but poorly coordinated."
  ),
  "Unclear_responsibilities": (
    "Municipalities should come up with initiatives; government must regulate "
    "private sector and create mechanisms — roles not clearly fulfilled."
  ),
  "Federal government": (
    "Government initiatives needed to regulate private sector and motivate "
    "behaviour (tax); historically banned plastic without alternatives."
  ),
  "Local government": (
    "Municipalities and local municipal bodies should lead initiatives and "
    "create waste-management mechanisms."
  ),
  "Students": (
    "Young generation sometimes makes initiatives on plastic/waste issues."
  ),
  "Private_sector": (
    "Private sector must be regulated but needs profit pathway; two segregation "
    "companies; waste-to-wealth model for private companies."
  ),
  "Civil_society": (
    "Organisations involved in awareness and segregation efforts — not government."
  ),
  "Science": (
    "Ecology background; micro- and nano-plastic, soil, aquatic impacts discussed "
    "scientifically."
  ),
  "Households": (
    "Households and communities must act cautiously; Kathmandu households mix "
    "all waste."
  ),
  "Education institutions": (
    "Education and awareness campaigns advocated for household behaviour change."
  ),
  "Private_companies(hotels, shops, etc.)": (
    "Private companies can develop waste-to-wealth models under regulation."
  ),
  "Research": (
    "Support needed: knowledge and opportunities — knowledge gap for effective "
    "action."
  ),
  "Enforcement": (
    "Private sector needs regulation; tax motivation for public compliance."
  ),
  "Policies in place": (
    "1990s–2000s plastic-ban idea (without alternatives); cloth-bag alternative; "
    "two segregation companies operating; youth/CSO awareness initiatives."
  ),
  "Pol_awarness": (
    "Awareness campaigns by organisations and young generation; respondent "
    "worked with segregation companies."
  ),
  "Pol_education": (
    "Education programmes advocated as solution pathway."
  ),
  "Pol_capacity": (
    "Worked with two segregation companies — private-sector collection capacity."
  ),
  "Pol_tax": (
    "Tax proposed to motivate people toward cautious plastic use."
  ),
  "Pol_ban": (
    "1990s–2000s idea to ban plastic without providing alternatives — failed "
    "approach recalled."
  ),
  "Pol_subsitutes": (
    "Cloth bags as alternative but NPR 40 — too expensive for ordinary people."
  ),
  "Pol_recycling": (
    "Segregation companies and waste-to-wealth recycling model."
  ),
  "Pol_waste_collection": (
    "Two segregation companies worked with; Kathmandu mixed-waste disposal "
    "undermines collection."
  ),
  "Sol_responsibilities": (
    "Municipalities must create mechanisms and come up with initiatives; "
    "households and communities share responsibility."
  ),
  "Sol_awarness": (
    "Awareness campaigns — even small individual actions matter."
  ),
  "Sol_segregation": (
    "Household segregation needed — currently absent in Kathmandu."
  ),
  "Sol_recycling": (
    "Waste-to-wealth model for private companies under government regulation."
  ),
  "Sol_education": (
    "Education and awareness to change household and community behaviour."
  ),
  "Sol_capacity": (
    "Government initiatives to enable regulated private-sector waste-to-wealth "
    "operations."
  ),
  "Sol_tax": (
    "Motivate people with tax incentives/levies to reduce plastic use."
  ),
  "Sol_ban": (
    "Sol_ban coded no — respondent warns against bans without alternatives "
    "(historical lesson); solutions focus on tax, awareness, alternatives."
  ),
  "Sol_infrastructure": (
    "Municipalities can create mechanisms for waste management at local level."
  ),
  "Sol_subsitutes": (
    "Affordable alternatives to plastic needed — cloth bags too expensive at "
    "NPR 40."
  ),
  "Sol_enforcement": (
    "Regulate private sector while enabling profitable waste-to-wealth pathways."
  ),
  "Traditions to build on (free-hand)": (
    "NA — traditions follow-up not asked."
  ),
}


def hfill(hex_):
  return PatternFill("solid", fgColor=hex_)


def hfont(bold=False, size=10, color="000000", italic=False):
  return Font(bold=bold, size=size, color=color, italic=italic)


thin = Side(style="thin", color="BFBFBF")


def thin_border():
  return Border(left=thin, right=thin, top=thin, bottom=thin)


WRAP_ALIGN = Alignment(wrap_text=True, vertical="top", horizontal="left")
CENTER_ALIGN = Alignment(wrap_text=True, vertical="center", horizontal="center")


def build_workbook():
  wb = openpyxl.Workbook()

  ws = wb.active
  ws.title = "Interview Coding"

  ws.merge_cells("A1:D1")
  title = ws["A1"]
  title.value = "Plastic Pollution Governance — Interview Coding (Codebook)"
  title.fill = hfill(DARK_BLUE)
  title.font = hfont(bold=True, size=13, color=WHITE)
  title.alignment = CENTER_ALIGN

  ws.merge_cells("A2:D2")
  meta = ws["A2"]
  meta.value = (
    "Interview: NPL_29  |  Country: NEPAL  |  Former UNDP/UNEP (ecology)  |  "
    "Actor: igo  |  18 Feb 2026"
  )
  meta.fill = hfill(MID_BLUE)
  meta.font = hfont(size=9, color=WHITE)
  meta.alignment = CENTER_ALIGN

  headers = ["Variable", "Code", "Codebook definition", "Coding notes"]
  for col, header in enumerate(headers, start=1):
    cell = ws.cell(row=4, column=col, value=header)
    cell.fill = hfill(LIGHT_BLUE)
    cell.font = hfont(bold=True, size=10, color=DARK_BLUE)
    cell.alignment = CENTER_ALIGN
    cell.border = thin_border()

  row = 5
  for variable, definition in CODEBOOK:
    values = [
      variable,
      CODING[variable],
      definition,
      NOTES.get(variable, ""),
    ]
    fill = hfill(LIGHT_GREEN) if row % 2 == 0 else hfill(WHITE)
    for col, value in enumerate(values, start=1):
      cell = ws.cell(row=row, column=col, value=value)
      cell.fill = fill
      cell.font = hfont(size=10)
      cell.alignment = WRAP_ALIGN
      cell.border = thin_border()
    row += 1

  ws.column_dimensions["A"].width = 34
  ws.column_dimensions["B"].width = 28
  ws.column_dimensions["C"].width = 52
  ws.column_dimensions["D"].width = 58
  ws.freeze_panes = "A5"

  wide = wb.create_sheet("Wide Format")
  wide_headers = [variable for variable, _ in CODEBOOK]
  for col, header in enumerate(wide_headers, start=1):
    cell = wide.cell(row=1, column=col, value=header)
    cell.fill = hfill(LIGHT_BLUE)
    cell.font = hfont(bold=True, size=9, color=DARK_BLUE)
    cell.alignment = CENTER_ALIGN
    cell.border = thin_border()
  for col, header in enumerate(wide_headers, start=1):
    cell = wide.cell(row=2, column=col, value=CODING[header])
    cell.fill = hfill(ORANGE)
    cell.alignment = WRAP_ALIGN
    cell.border = thin_border()
    wide.column_dimensions[get_column_letter(col)].width = max(14, min(28, len(header) + 2))

  ref = wb.create_sheet("Codebook Reference")
  ref.merge_cells("A1:B1")
  ref["A1"].value = "Codebook Variable Definitions"
  ref["A1"].fill = hfill(DARK_BLUE)
  ref["A1"].font = hfont(bold=True, size=12, color=WHITE)
  ref["A1"].alignment = CENTER_ALIGN
  for idx, (variable, definition) in enumerate(CODEBOOK, start=3):
    ref.cell(row=idx, column=1, value=variable).font = hfont(bold=True)
    ref.cell(row=idx, column=2, value=definition).alignment = WRAP_ALIGN
  ref.column_dimensions["A"].width = 34
  ref.column_dimensions["B"].width = 70

  stakeholder = wb.create_sheet("Stakeholder Analysis")
  stakeholder.merge_cells("A1:B1")
  stakeholder["A1"].value = "Stakeholder Analysis — NPL_29"
  stakeholder["A1"].fill = hfill(DARK_BLUE)
  stakeholder["A1"].font = hfont(bold=True, size=12, color=WHITE)
  stakeholder["A1"].alignment = CENTER_ALIGN
  for idx, (key, value) in enumerate(STAKEHOLDER.items(), start=3):
    stakeholder.cell(row=idx, column=1, value=key).font = hfont(bold=True)
    cell = stakeholder.cell(row=idx, column=2, value=value)
    cell.alignment = WRAP_ALIGN
  stakeholder.column_dimensions["A"].width = 28
  stakeholder.column_dimensions["B"].width = 90

  sources = wb.create_sheet("Sources")
  sources.merge_cells("A1:B1")
  sources["A1"].value = "Source Documents Used for Verification"
  sources["A1"].fill = hfill(DARK_BLUE)
  sources["A1"].font = hfont(bold=True, size=12, color=WHITE)
  sources["A1"].alignment = CENTER_ALIGN
  for idx, (doc, detail) in enumerate(SOURCES, start=3):
    sources.cell(row=idx, column=1, value=doc).font = hfont(bold=True)
    sources.cell(row=idx, column=2, value=detail).alignment = WRAP_ALIGN
  sources.column_dimensions["A"].width = 42
  sources.column_dimensions["B"].width = 60

  wb.save(OUTPUT_PATH)
  print(f"Saved {OUTPUT_PATH}")


if __name__ == "__main__":
  build_workbook()
