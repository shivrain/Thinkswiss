#!/usr/bin/env python3
"""Generate 4P Index coding spreadsheet for Arrêté 07022/2009 POLMAR (English only)."""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT = "/workspace/4P_Index_Arrete_07022_POLMAR.xlsx"

COLUMNS = [
    ("A", "policy_name"), ("B", "policy_url"), ("C", "policy_year"),
    ("D", "policy_objective"), ("E", "policy_target"), ("F", "policy_target_text"),
    ("G", "policy_type"), ("H", "policy_type_justification"),
    ("I", "policy_integration"), ("J", "policy_sectors_list"),
    ("K", "policy_circularity"), ("L", "policy_lifecycle_phases_list"),
    ("M", "policy_budget"), ("N", "policy_budget_text"),
    ("O", "policy_score"), ("P", "instrument_type"),
    ("Q", "instrument_lifecycle_stage"), ("R", "instrument_description"),
    ("S", "instrument_in_force"), ("T", "instrument_implementation"),
    ("U", "instrument_implementation_text"), ("V", "instrument_score"),
    ("W", "comments"),
]

POLICY = {
    "policy_name": (
        "Order No. 07022 of 16 July 2009 on organisation and operation of the "
        "National Marine Pollution Response Plan (POLMAR)"
    ),
    "policy_url": "https://www.hassmar.gouv.sn/sites/default/files/reglementations/PLAN%20POLMAR.pdf",
    "policy_year": 2009,
    "policy_objective": (
        "Organise national marine pollution preparedness and response (oil, chemicals, recovered waste) "
        "through HASSMAR/MRCC. Indirectly covers marine plastic waste via MARPOL 73/78 ship garbage "
        "reception, shoreline cleaning and recovered-waste treatment — no plastic-specific measures."
    ),
    "policy_target": 1,
    "policy_target_text": (
        "Arts. 69–71: Tier 1 spill less than 7 tonnes; Tier 2 between 7 and 700 tonnes; "
        "Tier 3 greater than 700 tonnes of pollutant products in national waters."
    ),
    "policy_type": 0.75,
    "policy_type_justification": (
        "Prime Minister's order adopted 16 July 2009 operationalising Decree 2006-322 (HASSMAR). "
        "Executive implementing regulation (≠1.0 law)."
    ),
    "policy_integration": 0.75,
    "policy_sectors_list": "fisheries, industry, environment, municipalities, waste management",
    "policy_circularity": 0.50,
    "policy_lifecycle_phases_list": "disposal, environmental leakage",
    "policy_budget": 0.5,
    "policy_budget_text": (
        "Art. 80: POLMAR fund finances implementation costs; Art. 77: ship/shipowner bears assistance costs."
    ),
}

INSTRUMENTS = [
    {
        "instrument_type": 0.40,
        "instrument_lifecycle_stage": "Environmental leakage",
        "instrument_description": (
            "Art. 73: prevention measures include ratifying maritime conventions, ensuring ships apply "
            "environmental regulations, and inspecting maritime installations. Preamble cites MARPOL 73/78."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            "Art. 36: port reception of ship operational waste (+0.25 monitoring via inspections); "
            "HASSMAR National Coordinator designated (+0.25 authority). MARPOL Annex V covers ship garbage/plastics."
        ),
        "comments": "Most direct plastic link — MARPOL 73/78 ship garbage reception.",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Environmental leakage",
        "instrument_description": (
            "Arts. 11–12: port authorities and coastal local authorities must develop prevention and "
            "response plans for marine pollution; sectoral plans annexed to POLMAR."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            "Art. 16: sectoral plans submitted for National Coordinator approval (+0.25 authority); "
            "Art. 12: local authorities acquire Tier 1/2 response means (+0.25 monitoring). No enforcement fines in plan."
        ),
        "comments": "Local/port plans — framework for coastal litter management.",
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Environmental leakage",
        "instrument_description": (
            "Art. 84: response modes include recover spilled products, clean and restore the shoreline, "
            "and treat waste; terrestrial response on beaches and rivers by Regional Governor."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            "Art. 86: polluter must lead initial response (+0.25 enforcement via polluter responsibility); "
            "Arts. 88–95: escalation procedures (+0.25 monitoring); HASSMAR/MRCC coordination (+0.25 authority)."
        ),
        "comments": "Shoreline cleaning — relevant for coastal plastic litter.",
    },
    {
        "instrument_type": 0.80,
        "instrument_lifecycle_stage": "End of life",
        "instrument_description": (
            "Arts. 37, 79, 84: Environment Ministry coordinates waste treatment and site restoration; "
            "identifies provisional storage for pollutant products and polluted waste from response operations."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            "Art. 37: Environment Ministry designated for waste treatment coordination (+0.25 authority); "
            "Arts. 40, 46, 57, 61: coastal storage at ports/industry (+0.25 monitoring). No penalties in plan text."
        ),
        "comments": "Recovered waste management — includes absorbent/polluted materials post-intervention.",
    },
    {
        "instrument_type": 0.20,
        "instrument_lifecycle_stage": "Environmental leakage",
        "instrument_description": (
            "Arts. 17–28: HASSMAR Secretary-General as National Coordinator; MRCC Dakar and regional "
            "secondary centres; national and local coordination committees for marine pollution response."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.50,
        "instrument_implementation_text": (
            "Art. 29: National Coordinator approves decontamination companies, maintains early-warning system "
            "(+0.25 authority, +0.25 monitoring). Complements CENPOLMAR (2021)."
        ),
        "comments": "Institutional coordination — not plastic-specific but enables marine waste response.",
    },
]


def build_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "4P Index Coding"
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    auto_fill = PatternFill("solid", fgColor="E2EFDA")

    for col_idx, (col_letter, col_name) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=f"{col_letter} — {col_name}")
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for row_idx, instrument in enumerate(INSTRUMENTS, start=2):
        row_data = {**POLICY, **instrument}
        for col_idx, (_, key) in enumerate(COLUMNS, start=1):
            if key in ("policy_score", "instrument_score"):
                continue
            ws.cell(row=row_idx, column=col_idx, value=row_data.get(key, ""))
        g_col, i_col, k_col, m_col = (get_column_letter(n) for n in (7, 9, 11, 13))
        p_col, t_col, o_col, v_col = (get_column_letter(n) for n in (16, 20, 15, 22))
        ws[f"{o_col}{row_idx}"] = f"=AVERAGE({g_col}{row_idx},{i_col}{row_idx},{k_col}{row_idx},{m_col}{row_idx},{p_col}{row_idx})"
        ws[f"{v_col}{row_idx}"] = f"=AVERAGE({p_col}{row_idx},{t_col}{row_idx})"
        ws[f"{o_col}{row_idx}"].fill = auto_fill
        ws[f"{v_col}{row_idx}"].fill = auto_fill

    widths = {"A": 48, "B": 52, "C": 10, "D": 52, "E": 12, "F": 60, "G": 10, "H": 52,
              "I": 14, "J": 40, "K": 14, "L": 36, "M": 12, "N": 52, "O": 12, "P": 14,
              "Q": 22, "R": 60, "S": 14, "T": 18, "U": 52, "V": 14, "W": 52}
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(OUTPUT)
    print(f"Saved {OUTPUT} ({len(INSTRUMENTS)} instrument rows)")


if __name__ == "__main__":
    build_workbook()
