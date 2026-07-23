#!/usr/bin/env python3
"""Generate 4P Index coding spreadsheet for Décret 2021-1115 CENPOLMAR (English only).

NOTE: This policy fails the strict plastics relevance filter (hydrocarbon marine pollution only).
File retained with policy-level metadata and zero instrument rows per strict coding rules.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT = "/workspace/4P_Index_Decret_2021-1115_CENPOLMAR.xlsx"

COLUMNS = [
    ("A", "policy_name"), ("B", "policy_url"), ("C", "policy_year"),
    ("D", "policy_objective"), ("E", "policy_target"), ("F", "policy_target_text"),
    ("G", "policy_type"), ("H", "policy_type_justification"),
    ("I", "policy_integration"), ("J", "policy_sectors_list"),
    ("K", "policy_circularity"), ("L", "policy_lifecycle_phases_list"),
    ("M", "policy_budget"), ("N", "policy_budget_text"),
    ("O", "policy_score"), ("P", "instrument_type"),
    ("Q", "instrument_lifecycle_stage"), ("R", "instrument_description"),
    ("S", "instrument_in_force"), ("T", "instrument_implementation"),
    ("U", "instrument_implementation_text"), ("V", "instrument_score"),
    ("W", "comments"),
]

# Policy metadata only — no instruments pass strict plastics relevance filter.
# CENPOLMAR coordinates hydrocarbon marine pollution response; no plastics mention.
# Cross-reference: code plastic-relevant provisions in Arrêté 07022/2009 (POLMAR) instead.

POLICY = {
    "policy_name": (
        "Decree No. 2021-1115 of 25 August 2021 establishing the National Marine Pollution "
        "Response Coordination Centre (CENPOLMAR)"
    ),
    "policy_url": (
        "https://www.hassmar.gouv.sn/sites/default/files/reglementations/Decret%20CENPOLMAR_0.pdf"
    ),
    "policy_year": 2021,
    "policy_objective": (
        "Creates CENPOLMAR within HASSMAR as operational coordination body for marine hydrocarbon pollution "
        "response. Does not mention plastics or regulate plastic waste. EXCLUDED from instrument coding per "
        "strict plastics relevance filter — see Arrêté 07022/2009 (POLMAR) for marine waste provisions."
    ),
    "policy_target": 0,
    "policy_target_text": "",
    "policy_type": 0.75,
    "policy_type_justification": (
        "Presidential decree of 25 August 2021 establishing national operational centre within HASSMAR. "
        "Published Official Gazette No. 7470 of 13 November 2021. Executive regulation (≠1.0 law)."
    ),
    "policy_integration": 0.50,
    "policy_sectors_list": "fisheries, environment, industry, water",
    "policy_circularity": 0.25,
    "policy_lifecycle_phases_list": "disposal",
    "policy_budget": 0.5,
    "policy_budget_text": (
        "Art. 12: CENPOLMAR resources from HASSMAR budget, grants, private-sector services, "
        "exceptional contributions, donations and international cooperation."
    ),
}

INSTRUMENTS = []


def build_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "4P Index Coding"
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)

    for col_idx, (col_letter, col_name) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=f"{col_letter} — {col_name}")
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    # Row 2: exclusion note (not a scored instrument)
    note_row = 2
    ws.cell(row=note_row, column=1, value=POLICY["policy_name"])
    ws.cell(row=note_row, column=2, value=POLICY["policy_url"])
    ws.cell(row=note_row, column=3, value=POLICY["policy_year"])
    ws.cell(row=note_row, column=4, value=POLICY["policy_objective"])
    for col_idx in range(5, 15):
        key = COLUMNS[col_idx - 1][1]
        ws.cell(row=note_row, column=col_idx, value=POLICY.get(key, ""))
    ws.cell(row=note_row, column=16, value=0)
    ws.cell(row=note_row, column=23, value=(
        "EXCLUDED: No instruments coded. Policy fails strict plastics relevance filter "
        "(hydrocarbon marine pollution only; no plastics, EPR, bans, or municipal waste planning). "
        "Code marine plastic provisions under POLMAR (Order 07022/2009) instead."
    ))

    widths = {"A": 48, "B": 52, "C": 10, "D": 52, "E": 12, "F": 60, "G": 10, "H": 52,
              "I": 14, "J": 40, "K": 14, "L": 36, "M": 12, "N": 52, "O": 12, "P": 14,
              "Q": 22, "R": 60, "S": 14, "T": 18, "U": 52, "V": 14, "W": 52}
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(OUTPUT)
    print(f"Saved {OUTPUT} (0 instrument rows — excluded per plastics filter)")


if __name__ == "__main__":
    build_workbook()
