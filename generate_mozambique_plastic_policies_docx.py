#!/usr/bin/env python3
"""Generate Mozambique plastic pollution policy list as a Word document."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

OUTPUT = "/workspace/Mozambique_Plastic_Pollution_Policies.docx"

POLICIES = [
    {
        "title": "Política Nacional do Ambiente (Resolução n.º 5/95)",
        "type": "Policy",
        "year": "1995",
        "description": "Approves Mozambique's national environmental policy, establishing principles and guidelines for sustainable use of natural resources, pollution prevention, and environmental management across sectors.",
        "sectors": "Environment; all economic sectors",
        "lifecycle": "All stages (framework)",
        "link": "https://www.agricultura.gov.mz/governo-lanca-processo-de-revisao-da-politica-e-lei-do-ambiente-para-reforcar-a-gestao-ambiental-e-dos-recursos-naturais/",
    },
    {
        "title": "Lei do Ambiente (Lei n.º 20/97)",
        "type": "Law",
        "year": "1997",
        "description": "Framework environmental law prohibiting unauthorized release of pollutants, establishing bases for environmental management, licensing, and sustainable development.",
        "sectors": "Environment; industry; municipalities; agriculture",
        "lifecycle": "All stages (framework)",
        "link": "https://sibmoz.gov.mz/content/uploads/2022/01/1547563349-Lei-do-Ambiente.pdf",
    },
    {
        "title": "Regulamento sobre a Gestão e Controlo do Saco de Plástico (Decreto n.º 16/2015)",
        "type": "Regulation",
        "year": "2015",
        "description": "Regulates production, import, marketing, and use of plastic bags; prohibits thin plastic bags, sets labelling and thickness requirements, and phases in import/production bans to reduce health and environmental impacts.",
        "sectors": "Industry; commerce; municipalities; environment",
        "lifecycle": "Production; distribution; use; end-of-life",
        "link": "https://sibmoz.gov.mz/content/uploads/2022/01/Regulamento-Sobre-a-Gestao-e-Controlo-do-Saco-de-Plastico.pdf",
    },
    {
        "title": "Regulamento sobre a Responsabilidade Alargada dos Produtores e Importadores de Embalagens (Decreto n.º 79/2017)",
        "type": "Regulation",
        "year": "2017",
        "description": "Establishes extended producer responsibility (EPR) for packaging, including internal management systems, packaging environmental fee (TAE), and packaging standardisation to reduce pollution and fund waste management.",
        "sectors": "Industry; commerce; environment; municipalities",
        "lifecycle": "Production; distribution; end-of-life",
        "link": "https://www.agricultura.gov.mz/governo-mobiliza-cerca-de-18-4-milhoes-de-euros-para-viabilizar-o-programa-nacional-de-gestao-sustentavel-de-residuos/",
    },
    {
        "title": "Regulamento sobre a Gestão de Resíduos Sólidos Urbanos (Decreto n.º 94/2014)",
        "type": "Regulation",
        "year": "2014",
        "description": "National regulation on urban solid waste management: segregation (including plastics), collection, transport, treatment, valorisation, landfill rules, and municipal/district government responsibilities.",
        "sectors": "Municipalities; environment; public health; private waste operators",
        "lifecycle": "Collection; end-of-life",
        "link": "https://www.agricultura.gov.mz/wp-content/uploads/2022/11/PIDACC_ESMF_Eng.pdf",
    },
    {
        "title": "Regulamento sobre a Gestão de Resíduos Perigosos (Decreto n.º 83/2014)",
        "type": "Regulation",
        "year": "2014",
        "description": "Sets procedures for safe management, transport, storage, and disposal of hazardous wastes from industrial and other activities, including wastes with high environmental and health risk.",
        "sectors": "Industry; health; environment; municipalities",
        "lifecycle": "Production; end-of-life",
        "link": "https://sibmoz.gov.mz/content/uploads/2022/01/Aprova-o-Regulamento-sobre-Gestao-de-Residuos-perigosos-e-os-respectivos-anexos.pdf",
    },
    {
        "title": "Programa Nacional de Gestão Sustentável de Resíduos (ValoRe)",
        "type": "Program",
        "year": "2025",
        "description": "Government circular-economy programme to build sustainable waste-treatment infrastructure, strengthen recycling value chains, and pilot municipal waste management in Nampula, Nacala, and Pemba.",
        "sectors": "Municipalities; environment; industry; finance",
        "lifecycle": "Collection; end-of-life",
        "link": "https://www.agricultura.gov.mz/governo-mobiliza-cerca-de-18-4-milhoes-de-euros-para-viabilizar-o-programa-nacional-de-gestao-sustentavel-de-residuos/",
    },
    {
        "title": "Regulamento sobre Padrões de Qualidade Ambiental e de Emissão de Efluentes (Decreto n.º 18/2004, alterado pelo Decreto n.º 67/2010)",
        "type": "Regulation",
        "year": "2010",
        "description": "Sets national environmental quality and effluent/emission standards for air, water, and mobile sources; limits pollutant releases and provides basis for enforcement against industrial pollution.",
        "sectors": "Industry; environment; energy; transport",
        "lifecycle": "Production; end-of-life",
        "link": "https://sibmoz.gov.mz/content/uploads/2022/01/Regulamento-sobre-Padroes-de-Qualidade-Ambiental-e-de-Emissao-de-Efluentes.pdf",
    },
    {
        "title": "Regulamento para Prevenção da Poluição e Proteção do Ambiente Marinho e Costeiro (Decreto n.º 45/2006)",
        "type": "Regulation",
        "year": "2006",
        "description": "Requires measures to prevent, control, and combat marine pollution from ships and land-based sources in Mozambique's jurisdictional waters and coastal environment.",
        "sectors": "Maritime; fisheries; ports; environment",
        "lifecycle": "Distribution; end-of-life",
        "link": "https://sibmoz.gov.mz/content/uploads/2022/01/1547623391-Regulamento-sobre-a-Prevencao-da-Poluicao-ea-Proteccao-Ambiente-Marinho-e-Costeiro-Decreto-45-barra-2006.pdf",
    },
    {
        "title": "Regulamento sobre a Actividade de Fiscalização Ambiental (Decreto n.º 51/2024)",
        "type": "Regulation",
        "year": "2024",
        "description": "Approves rules for environmental inspection and enforcement of environmental legislation, replacing Decree 11/2006 and defining inspection types and competencies.",
        "sectors": "Environment; all regulated sectors",
        "lifecycle": "All stages (enforcement)",
        "link": "https://sibmoz.gov.mz/content/uploads/2025/11/Regulamento-de-fiscalizac%CC%A7a%CC%83o-ambiental.pdf",
    },
    {
        "title": "Regulamento da Avaliação do Impacto Ambiental (Decreto n.º 54/2015)",
        "type": "Regulation",
        "year": "2015",
        "description": "Regulates environmental impact assessment and licensing for projects that may cause significant environmental harm, including mitigation hierarchy and biodiversity safeguards.",
        "sectors": "Industry; infrastructure; mining; energy; tourism",
        "lifecycle": "Production; end-of-life (project planning)",
        "link": "https://sibmoz.gov.mz/content/uploads/2022/01/1547556163-Decreto-54-2015-Regulamento-da-AIA.pdf",
    },
    {
        "title": "Regulamento sobre o Processo de Auditoria Ambiental (Decreto n.º 45/2024)",
        "type": "Regulation",
        "year": "2024",
        "description": "Establishes rules for environmental audits of public and private activities that may impact the environment, including reporting and corrective action requirements.",
        "sectors": "Industry; environment; public administration",
        "lifecycle": "Production; end-of-life (compliance)",
        "link": "https://sibmoz.gov.mz/content/uploads/2025/11/Regulamento-sobre-o-Processo-de-Auditoria-Ambiental.pdf",
    },
    {
        "title": "Política e Estratégia do Mar – POLMAR (Resolução n.º 39/2017)",
        "type": "Policy / Strategy",
        "year": "2017",
        "description": "National sea policy and implementation strategy guiding sustainable, integrated use of maritime and coastal spaces and blue-economy development.",
        "sectors": "Maritime; fisheries; tourism; ports; environment",
        "lifecycle": "All stages (marine/coastal framework)",
        "link": "https://www.agricultura.gov.mz/maap-realiza-primeira-reuniao-da-comissao-consultiva-do-plano-de-ordenamento-do-espaco-maritimo-nacional/",
    },
    {
        "title": "Lei do Mar (Lei n.º 20/2019)",
        "type": "Law",
        "year": "2019",
        "description": "Legal regime for sovereignty, jurisdiction, exploitation of marine resources, and use of national maritime public domain.",
        "sectors": "Maritime; fisheries; mining; transport",
        "lifecycle": "All stages (marine framework)",
        "link": "https://sibmoz.gov.mz/content/uploads/2022/01/Lei-do-Mar.pdf",
    },
    {
        "title": "Regulamento de Gestão e Ordenamento da Zona Costeira e das Praias (Decreto n.º 97/2020)",
        "type": "Regulation",
        "year": "2020",
        "description": "Defines principles and norms for integrated, sustainable management and development of coastal zones and beaches nationwide.",
        "sectors": "Municipalities; tourism; fisheries; environment",
        "lifecycle": "Distribution; end-of-life",
        "link": "https://sibmoz.gov.mz/content/uploads/2022/06/Regulamento-Zonas-costeiras.pdf",
    },
    {
        "title": "Plano Nacional de Ordenamento do Espaço Marítimo – POEM (Resolução n.º 63/2024)",
        "type": "Plan",
        "year": "2024",
        "description": "Approves national maritime spatial plan to coordinate competing sea uses, minimise environmental impacts, and support sustainable maritime and coastal economies.",
        "sectors": "Maritime; fisheries; energy; transport; environment",
        "lifecycle": "All stages (marine planning)",
        "link": "https://sibmoz.gov.mz/content/uploads/2025/11/Resolucao-n.o-63_2024-que-aprova-o-Plano-de-Ordenamento-do-Espaco-Maritimo.pdf",
    },
    {
        "title": "Estratégia de Desenvolvimento da Economia Azul – EDEA (Resolução n.º 53/2024)",
        "type": "Strategy",
        "year": "2024",
        "description": "National blue-economy strategy (2024–2033) promoting sustainable use of marine and aquatic resources, circular economy, waste valorisation, and explicit targets to reduce plastic production, import, and use.",
        "sectors": "Maritime; fisheries; industry; tourism; environment",
        "lifecycle": "Production; distribution; end-of-life",
        "link": "https://www.proazul.gov.mz/wp-content/uploads/2024/11/EDEA-ESTRATEGIA-DE-DESENVOLVIMENTO-DA-ECONOMIA-AZUL-_FINAL_TYPO_PRINT_WEB.pdf",
    },
    {
        "title": "Estratégia de Gestão Integrada das Zonas Costeiras (2016–2025)",
        "type": "Strategy",
        "year": "2016",
        "description": "Integrated coastal zone management action plan addressing erosion, climate vulnerability, and sustainable use of coastal resources linked to marine pollution and livelihood pressures.",
        "sectors": "Municipalities; fisheries; tourism; environment",
        "lifecycle": "Distribution; end-of-life",
        "link": "https://sibmoz.gov.mz/specific-biodiversity-strategies-and-action-plans/",
    },
    {
        "title": "Estratégia Nacional de Gestão e Conservação dos Corais (2020–2030; Resolução n.º 51/2022)",
        "type": "Strategy",
        "year": "2022",
        "description": "National strategy to protect coral reef ecological integrity and resilience by reducing anthropogenic degradation, including marine litter and pollution pressures.",
        "sectors": "Fisheries; environment; tourism; maritime",
        "lifecycle": "End-of-life (marine leakage)",
        "link": "https://sibmoz.gov.mz/specific-biodiversity-strategies-and-action-plans/",
    },
    {
        "title": "Política de Saúde e Estratégia de Implementação (Resolução n.º 13/2021)",
        "type": "Policy",
        "year": "2021",
        "description": "National health policy including strategies to expand basic sanitation and manage solid waste pressures linked to rapid urbanisation and public health.",
        "sectors": "Health; municipalities; environment",
        "lifecycle": "End-of-life (sanitation/waste)",
        "link": "https://sisma.misau.gov.mz/docs/Politica%20Nacional%20de%20Saude.pdf",
    },
    {
        "title": "Plano Nacional de Desenvolvimento Territorial (Resolução n.º 7/2021)",
        "type": "Plan",
        "year": "2021",
        "description": "Approves national territorial development plan and action plan, guiding sustainable land use, urban growth, and environmental balance at national and subnational levels.",
        "sectors": "Urban planning; municipalities; environment; infrastructure",
        "lifecycle": "All stages (spatial planning)",
        "link": "https://sibmoz.gov.mz/content/uploads/2022/06/Plano-Nacional-de-Desenvolvimento-Territorial-1.pdf",
    },
    {
        "title": "Lei do Ordenamento do Território (Lei n.º 19/2007)",
        "type": "Law",
        "year": "2007",
        "description": "Territorial planning law promoting rational and sustainable use of natural resources, environmental balance, and quality of life in national development.",
        "sectors": "Urban planning; municipalities; environment; infrastructure",
        "lifecycle": "All stages (planning framework)",
        "link": "https://sibmoz.gov.mz/content/uploads/2022/01/Lei-do-ordenamento-territorial.pdf",
    },
    {
        "title": "Regulamento da Lei de Ordenamento do Território (Decreto n.º 23/2008)",
        "type": "Regulation",
        "year": "2008",
        "description": "Regulates territorial occupation and sustainable land use through national, provincial, district, and municipal planning instruments coordinating economic, social, and environmental development.",
        "sectors": "Municipalities; urban planning; environment",
        "lifecycle": "All stages (subnational planning)",
        "link": "https://sibmoz.gov.mz/content/uploads/2025/11/Decreto-n.o-23_2008_Regulamento-da-Lei-de-Ordenamento-do-Territorio.pdf",
    },
]

COLUMNS = [
    "Policy Title",
    "Policy Type",
    "Last amendment year",
    "Short description",
    "Sectors",
    "Life-Cycle stage",
    "Link",
]


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    new_run.append(r_pr)
    text_elem = OxmlElement("w:t")
    text_elem.text = text
    new_run.append(text_elem)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    title = doc.add_heading("Mozambique: National Policies Addressing Plastic Pollution", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph(
        "National and national-subnational instruments (laws, regulations, policies, strategies, plans, and programmes) "
        "that directly or indirectly address plastic pollution across the life cycle. International treaties excluded. "
        "Links are official Government of Mozambique sources (.gov.mz), verified June 2026."
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT

    table = doc.add_table(rows=1, cols=len(COLUMNS))
    table.style = "Table Grid"
    table.autofit = False
    widths = [Inches(1.55), Inches(0.75), Inches(0.55), Inches(1.8), Inches(1.0), Inches(0.95), Inches(1.4)]
    for idx, width in enumerate(widths):
        table.columns[idx].width = width

    hdr_cells = table.rows[0].cells
    for idx, col in enumerate(COLUMNS):
        p = hdr_cells[idx].paragraphs[0]
        p.text = col
        for run in p.runs:
            run.bold = True
            run.font.size = Pt(9)

    for policy in POLICIES:
        row = table.add_row().cells
        values = [
            policy["title"],
            policy["type"],
            policy["year"],
            policy["description"],
            policy["sectors"],
            policy["lifecycle"],
            policy["link"],
        ]
        for idx, value in enumerate(values):
            cell = row[idx]
            p = cell.paragraphs[0]
            if idx == 6:
                add_hyperlink(p, value, value)
            else:
                p.text = value
            for run in p.runs:
                run.font.size = Pt(8)

    note = doc.add_paragraph(
        "Note: Where the full legal text is not published as a standalone file on a .gov.mz portal, the link points to the "
        "official government page or document that identifies and describes the instrument (e.g., MAAP news release or "
        "official environmental assessment citing the decree). Full-text PDFs hosted on sibmoz.gov.mz are official "
        "government publications from the national biodiversity portal."
    )
    for run in note.runs:
        run.font.size = Pt(8)

    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT} with {len(POLICIES)} policies")


if __name__ == "__main__":
    build_document()
