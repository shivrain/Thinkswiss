#!/usr/bin/env python3
"""Generate 4P Index coding spreadsheet for NS 05-061 (English only).

Plastic-industry effluent provisions only; strict plastics-relevance filter.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT = "/workspace/4P_Index_NS_05-061_Eaux_Usees.xlsx"

COLUMNS = [
    ("A", "policy_name"), ("B", "policy_url"), ("C", "policy_year"),
    ("D", "policy_objective"), ("E", "policy_target"), ("F", "policy_target_text"),
    ("G", "policy_type"), ("H", "policy_type_justification"),
    ("I", "policy_integration"), ("J", "policy_sectors_list"),
    ("K", "policy_circularity"), ("L", "policy_lifecycle_phases_list"),
    ("M", "policy_budget"), ("N", "policy_budget_text"),
    ("O", "policy_score"), ("P", "instrument_type"),
    ("Q", "instrument_lifecycle_stage"), ("R", "instrument_description"),
    ("S", "instrument_in_force"), ("T", "instrument_implementation"),
    ("U", "instrument_implementation_text"), ("V", "instrument_score"),
    ("W", "comments"),
]

POLICY = {
    "policy_name": "Senegalese Standard NS 05-061 — Wastewater discharges (July 2001)",
    "policy_url": (
        "https://www.denv.gouv.sn/telechargement/74/normes-cadre-juridique/"
        "19105/norme-rejets-ns-05-061-eaux-usees-juillet-2001.pdf"
    ),
    "policy_year": 2001,
    "policy_objective": (
        "Set limit values, prohibitions, monitoring modalities and discharge criteria for industrial and "
        "domestic wastewater into receiving environments, complementing Decree No. 2001-282 — indirectly "
        "controlling plastic-related industrial effluents through physico-chemical parameters (TSS, COD, BOD5, "
        "TOC, micropollutants) without plastic-specific measures."
    ),
    "policy_target": 1,
    "policy_target_text": (
        "Total suspended solids: 50 mg/l; BOD5: 80 mg/l (daily load ≤30 kg/d) or 40 mg/l beyond; "
        "COD: 200 mg/l (daily load ≤100 kg/d) or 100 mg/l beyond; "
        "total hydrocarbons: 15 mg/l if discharge exceeds 150 g/d."
    ),
    "policy_type": 0.75,
    "policy_type_justification": (
        "Official technical standard adopted in July 2001 by the ISN/CT5 standardisation committee, published "
        "by the Senegalese Standardisation Institute, complementing Decree No. 2001-282. Operative technical "
        "regulation of reference for effluent discharges (≠1.0 law)."
    ),
    "policy_integration": 0.75,
    "policy_sectors_list": (
        "industry, sanitation, environment, water, fisheries, agriculture, waste management"
    ),
    "policy_circularity": 0.75,
    "policy_lifecycle_phases_list": "production, disposal, environmental leakage",
    "policy_budget": 0.5,
    "policy_budget_text": (
        "Ch. 3 §1.4-1.5: determination of effluent pollution degree and tax rate payable by operator; "
        "fee calculation set out in interministerial order on application of this standard."
    ),
}

INSTRUMENTS = [
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Environmental leakage",
        "instrument_description": (
            "§5.1: any discharge of liquid effluents causing stagnation, neighbourhood nuisance, or pollution "
            "of surface, groundwater or marine waters is prohibited throughout the national territory."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            "Ch. 3 §1.1-1.2: sampling and analyses by Environment Directorate technicians or designated entity "
            "(+0.25 authority); §1.7: enhanced monitoring when Annex II thresholds exceeded (+0.25 monitoring); "
            "Ch. 3 §1.4: results determine tax rate payable by operator (+0.25 enforcement)."
        ),
        "comments": (
            "General prohibition applicable to plastic-industry effluent discharges. No explicit plastics mention."
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Environmental leakage",
        "instrument_description": (
            "§5.2 and Annex II §2: any treated effluent discharged into a receiving environment must comply "
            "with Annex II values, including total suspended solids 50 mg/l, BOD5 80/40 mg/l and COD 200/100 mg/l."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 1.0,
        "instrument_implementation_text": (
            "Ch. 3 §1.3-1.5: mandatory sampling before discharge, minimum twice per year (+0.25 monitoring); "
            "§1.7: continuous flow measurement when daily flow >100 m³ (+0.25 authority); "
            "§1.4: results determine tax rate payable by operator (+0.25 enforcement); "
            "Annex II §2 iii: annual discharge report for ICPE operators (+0.25 monitoring)."
        ),
        "comments": (
            "Core instrument. TSS may capture suspended plastic particles in industrial effluents, without specific targeting."
        ),
    },
    {
        "instrument_type": 1.0,
        "instrument_lifecycle_stage": "Waste management",
        "instrument_description": (
            "Ch. 2 §II: effluent discharged via public evacuation channels without treatment plant must be "
            "cleared of all products likely to harm infrastructure integrity, as well as floating, settleable "
            "or precipitable matter."
        ),
        "instrument_in_force": 1,
        "instrument_implementation": 0.75,
        "instrument_implementation_text": (
            "Ch. 2 §II: applies to non-domestic effluent connections (+0.25 authority); "
            "connection protocol required between generator and plant operator (+0.25 monitoring); "
            "more stringent values where discharge contains toxic chemicals (+0.25 enforcement)."
        ),
        "comments": (
            "Relevant for floating or precipitable plastic particles/matter in industrial effluents."
        ),
    },
]


def build_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "4P Index Coding"
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    auto_fill = PatternFill("solid", fgColor="E2EFDA")

    for col_idx, (col_letter, col_name) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=f"{col_letter} — {col_name}")
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for row_idx, instrument in enumerate(INSTRUMENTS, start=2):
        row_data = {**POLICY, **instrument}
        for col_idx, (_, key) in enumerate(COLUMNS, start=1):
            if key in ("policy_score", "instrument_score"):
                continue
            ws.cell(row=row_idx, column=col_idx, value=row_data.get(key, ""))

        g_col, i_col, k_col, m_col = (get_column_letter(n) for n in (7, 9, 11, 13))
        p_col, t_col, o_col, v_col = (get_column_letter(n) for n in (16, 20, 15, 22))
        ws[f"{o_col}{row_idx}"] = f"=AVERAGE({g_col}{row_idx},{i_col}{row_idx},{k_col}{row_idx},{m_col}{row_idx},{p_col}{row_idx})"
        ws[f"{v_col}{row_idx}"] = f"=AVERAGE({p_col}{row_idx},{t_col}{row_idx})"
        ws[f"{o_col}{row_idx}"].fill = auto_fill
        ws[f"{v_col}{row_idx}"].fill = auto_fill

    widths = {
        "A": 48, "B": 52, "C": 10, "D": 52, "E": 12, "F": 60,
        "G": 10, "H": 52, "I": 14, "J": 40, "K": 14, "L": 36,
        "M": 12, "N": 52, "O": 12, "P": 14, "Q": 22, "R": 60,
        "S": 14, "T": 18, "U": 52, "V": 14, "W": 52,
    }
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = "A2"
    wb.save(OUTPUT)
    print(f"Saved {OUTPUT} ({len(INSTRUMENTS)} instrument rows)")


if __name__ == "__main__":
    build_workbook()
